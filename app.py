"""CombiTrace-MS desktop workbench for enumeration, XIC analysis and response modelling."""

from __future__ import annotations

import threading
import shutil
from datetime import datetime
from pathlib import Path
from queue import Queue
from typing import Dict, List, Optional, Tuple

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from core.spectrum_processing import SpectrumSettings
from core.targets_csv import TargetSpec, load_targets_mapping
from core.abc_enumerator import (
    enumerate_abc,
    enumerate_abc_detailed,
    load_component_csv,
    write_results_csv,
    write_results_csv_detailed,
)
from core.csv_formula_filter import (
    filter_csv_remove_matching_formulas,
    guess_formula_column,
    read_csv_header,
)
from core.excel_sheet_xic import (
    load_excel_sheet_plans,
    normalize_internal_standard_formula,
    parse_replicate_suffixes,
    read_quant_csv,
    resolve_sheet_raws,
    safe_sheet_folder_name,
    targets_for_raw,
    write_all_targets_csv,
    write_parallel_summary_xlsx,
    write_plan_csv,
)
from core.summary_patcher import patch_summary_workbook
from core.response_predictor import run_response_prediction, MODEL_LABEL_TO_CODE
from core.combitrace_ie import run_combitrace_ie
from core.esi_descriptors import (
    default_gradient_text,
    parse_gradient_text,
    run_esi_descriptor_export,
)
from core.esi_model_benchmark import (
    CALIBRATION_MODE_LABELS as ESI_CALIBRATION_MODE_LABELS,
    CALIBRATION_MODE_MAP as ESI_CALIBRATION_MODE_MAP,
    MODEL_OBJECTIVE_LABELS as ESI_MODEL_OBJECTIVE_LABELS,
    MODEL_OBJECTIVE_MAP as ESI_MODEL_OBJECTIVE_MAP,
    OUTLIER_MODE_LABELS as ESI_OUTLIER_MODE_LABELS,
    OUTLIER_MODE_MAP as ESI_OUTLIER_MODE_MAP,
)
from core.esi_descriptor_meta import (
    LANGUAGE_LABELS as ESI_LANGUAGE_LABELS,
    LANGUAGE_MAP as ESI_LANGUAGE_MAP,
    display_name as esi_display_name,
)
from core.descriptor_review import load_descriptor_review, detect_previous_best_model
from core.fixed_feature_search import (
    FIXED_MODEL_LABELS as ESI_FIXED_MODEL_LABELS,
    FIXED_MODEL_MAP as ESI_FIXED_MODEL_MAP,
)
from core.simple_prediction_outputs import replot_all_prediction_plots_from_workbook
from core.published_ie_benchmark import (
    PUBLISHED_MODE_LABELS as IE_PUBLISHED_MODE_LABELS,
    PUBLISHED_MODE_MAP as IE_PUBLISHED_MODE_MAP,
    ION_MODE_LABELS as IE_ION_MODE_LABELS,
    ION_MODE_MAP as IE_ION_MODE_MAP,
    ADDUCT_LABELS as IE_ADDUCT_LABELS,
    ADDUCT_MAP as IE_ADDUCT_MAP,
    ORGANIC_MODIFIER_LABELS as IE_ORGANIC_MODIFIER_LABELS,
    ORGANIC_MODIFIER_MAP as IE_ORGANIC_MODIFIER_MAP,
)
from core.ms2quant_official import (
    CONCENTRATION_UNIT_LABELS as IE_CONCENTRATION_UNIT_LABELS,
    CONCENTRATION_UNIT_MAP as IE_CONCENTRATION_UNIT_MAP,
    check_ms2quant_environment,
    detect_rscript,
)
from core.negative_esi_literature import (
    DYNAMIC_RANGE_MODE_LABELS as IE_DYNAMIC_RANGE_MODE_LABELS,
    DYNAMIC_RANGE_MODE_MAP as IE_DYNAMIC_RANGE_MODE_MAP,
)
from core.smiles_builder import (
    BUILD_MODE_LABELS as SMILES_BUILD_MODE_LABELS,
    BUILD_MODE_MAP as SMILES_BUILD_MODE_MAP,
    preview_first_product,
    run_smiles_builder,
)
from core.simple_combo_smiles import (
    diagnose_simple_combo_inputs,
    preview_first_simple_combo,
    run_simple_combo_smiles,
)
from core.smiles_diagnostics import diagnostics_text
from core.xic_quant import quant_xic_for_raw
from core.negative_ion_channels import CHANNEL_BY_ID
from core.negative_ion_panel_review import (
    apply_panel_to_existing_triplicate,
    load_panel_review_rows,
    merge_existing_actions,
    refresh_discovery_summary,
    save_panel_review_rows,
)
from core.utils import format_yyyymmdd, open_in_os, read_json, safe_stem_from_raw_path
# =====================
# Adduct matching modes (UI -> internal code)
# =====================
ADDUCT_MODE_OPTIONS = [
    ('Both: negative [M-H]-', "posneg_neghonly"),
    ('Both: all ions', "posneg_all"),
    ('Positive ions (all)', "pos_all"),
    ('Negative ions (all)', "neg_all"),
    ('Use one specified ion', "force"),
]
ADDUCT_MODE_LABELS = [x[0] for x in ADDUCT_MODE_OPTIONS]
ADDUCT_MODE_MAP = {x[0]: x[1] for x in ADDUCT_MODE_OPTIONS}

# =====================
# XIC multi-peak (duplicates) detection modes
# =====================
MULTIPEAK_MODE_OPTIONS = [
    ('Automatic', "auto"),
    ('Conservative', "conservative"),
    ('Balanced', "balanced"),
    ('Sensitive', "sensitive"),
]
MULTIPEAK_MODE_LABELS = [x[0] for x in MULTIPEAK_MODE_OPTIONS]
MULTIPEAK_MODE_MAP = {x[0]: x[1] for x in MULTIPEAK_MODE_OPTIONS}

NEG_CHANNEL_MODE_OPTIONS = [
    ('Legacy: [M-H]- with optional formate', "legacy"),
    ('Discover channels, then review', "discover"),
    ('Use the selected channel panel', "panel"),
]
NEG_CHANNEL_MODE_LABELS = [x[0] for x in NEG_CHANNEL_MODE_OPTIONS]
NEG_CHANNEL_MODE_MAP = {x[0]: x[1] for x in NEG_CHANNEL_MODE_OPTIONS}

INTERNAL_STANDARD_ADDUCT_OPTIONS = [
    ('[M-H]-', "M-H"),
    ('Determine from RAW polarity', "auto_raw"),
    ("[M+H]+", "M+H"),
]
INTERNAL_STANDARD_ADDUCT_LABELS = [x[0] for x in INTERNAL_STANDARD_ADDUCT_OPTIONS]
INTERNAL_STANDARD_ADDUCT_MAP = {x[0]: x[1] for x in INTERNAL_STANDARD_ADDUCT_OPTIONS}


# =====================
# Excel sheet layout modes
# =====================
EXCEL_LAYOUT_OPTIONS = [
    ('Automatic (prefer four-column layout)', "auto"),
    ('Name / formula / mass / SMILES', "named_quartets"),
    ('Name / formula / mass', "named_triplets"),
    ('Formula / mass', "legacy_pairs"),
]
EXCEL_LAYOUT_LABELS = [x[0] for x in EXCEL_LAYOUT_OPTIONS]
EXCEL_LAYOUT_MAP = {x[0]: x[1] for x in EXCEL_LAYOUT_OPTIONS}

def _find_raw_paths(root: Path, *, recursive: bool = False) -> List[Path]:
    'Find RAW files and RAW directories under the selected directory.'

    if not root.exists() or not root.is_dir():
        return []

    patterns = ["*.raw", "*.RAW"]
    out: List[Path] = []
    for pat in patterns:
        it = root.rglob(pat) if recursive else root.glob(pat)
        for p in it:
            out.append(p)
    
    uniq = {str(p.resolve()): p for p in out}
    return [uniq[k] for k in sorted(uniq.keys())]


class ThermoBatchReportApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("CombiTrace-MS 18.48.2-en.1")
        self.geometry("1280x850")
        self.minsize(1100, 740)
        from core.ui_support import setup_style
        setup_style(self)
        self._build_help_menu()

        self._last_out_dir: Optional[Path] = None
        self._log_q_abc: "Queue[str]" = Queue()
        self._log_q_excel: "Queue[str]" = Queue()
        self._log_q_patch: "Queue[str]" = Queue()
        self._log_q_filter: "Queue[str]" = Queue()
        self._log_q_resp: "Queue[str]" = Queue()
        self._log_q_ie: "Queue[str]" = Queue()
        self._log_q_esi: "Queue[str]" = Queue()
        self._log_q_smiles: "Queue[str]" = Queue()
        self._log_q_xic: "Queue[str]" = Queue()

        # ESI manual descriptor candidate pool / fixed-model feature search.
        self.esi_manual_features: List[str] = []
        self.esi_descriptor_review_rows: List[Dict[str, object]] = []
        self.esi_previous_best_model: str = ""

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)

        self.tab_abc = ttk.Frame(self.notebook)
        self.tab_excel = ttk.Frame(self.notebook)
        self.tab_patch = ttk.Frame(self.notebook)
        self.tab_filter = ttk.Frame(self.notebook)
        self.tab_resp = ttk.Frame(self.notebook)
        self.tab_ie = ttk.Frame(self.notebook)
        self.tab_esi = ttk.Frame(self.notebook)
        self.tab_smiles = ttk.Frame(self.notebook)
        self.tab_xic = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_abc, text='Enumeration')
        self.notebook.add(self.tab_excel, text='Triplicate XIC')
        self.notebook.add(self.tab_patch, text='Table tools')
        self.notebook.add(self.tab_filter, text='CSV filter')
        self.notebook.add(self.tab_resp, text='Response')
        self.notebook.add(self.tab_smiles, text='SMILES')
        self.notebook.add(self.tab_esi, text='ESI models')
        self.notebook.add(self.tab_ie, text='Local RRF')
        self.notebook.add(self.tab_xic, text='XIC')

        # Append-only final-table processing: independent of modelling/RAW extraction.
        from core.final_table_lab import FinalTablePanel
        self.tab_final_tables = FinalTablePanel(self.notebook)
        self.notebook.insert(1, self.tab_final_tables, text='Final tables')

        self._build_abc_tab()
        self._build_excel_sheet_tab()
        self._build_summary_patch_tab()
        self._build_filter_tab()
        self._build_response_prediction_tab()
        self._build_smiles_builder_tab()
        self._build_esi_descriptor_tab()
        self._build_combitrace_ie_tab()
        self._build_xic_tab()
        from core.ui_support import wrap_labels
        wrap_labels(self)

        self.after(100, self._poll_log)

    # =====================
    # Generate tab
    # =====================



    def _build_help_menu(self):
        import webbrowser
        bar = tk.Menu(self)
        help_menu = tk.Menu(bar, tearoff=False)
        def open_guide():
            path = Path(__file__).resolve().parent / 'docs' / 'USER_GUIDE.html'
            if path.exists():
                webbrowser.open(path.as_uri())
            else:
                messagebox.showinfo('User guide', 'See docs/USER_GUIDE.md in the source release.')
        help_menu.add_command(label='User guide', command=open_guide)
        help_menu.add_command(label='About CombiTrace-MS', command=lambda: messagebox.showinfo(
            'CombiTrace-MS', 'CombiTrace-MS 18.48.2-en.1\nEnumeration, triplicate XIC and relative-response-factor modelling.\n\n'
            'RawFileReader reading tool. Copyright © 2016 by Thermo Fisher Scientific, Inc. All rights reserved.\n'
            'Independent software; not a Thermo Fisher Scientific product.'))
        bar.add_cascade(label='Help', menu=help_menu)
        self.config(menu=bar)

    def _update_xic_adduct_ui(self):
        'Enable the specified-ion control when single-ion matching is selected.'
        try:
            label = (self.var_xic_adduct_mode.get() or "").strip()
            code = ADDUCT_MODE_MAP.get(label, "posneg_neghonly")
            if code == "force":
                self.cmb_xic_forced_adduct.configure(state="readonly")
            else:
                self.cmb_xic_forced_adduct.configure(state="disabled")
        except Exception:
            pass







    def _log_abc(self, msg: str) -> None:
        self._log_q_abc.put(msg)

    def _log_excel(self, msg: str) -> None:
        self._log_q_excel.put(msg)

    def _log_patch(self, msg: str) -> None:
        self._log_q_patch.put(msg)

    def _log_filter(self, msg: str) -> None:
        self._log_q_filter.put(msg)

    def _log_resp(self, msg: str) -> None:
        self._log_q_resp.put(msg)

    def _log_ie(self, msg: str) -> None:
        self._log_q_ie.put(msg)

    def _log_esi(self, msg: str) -> None:
        self._log_q_esi.put(msg)

    def _log_smiles(self, msg: str) -> None:
        self._log_q_smiles.put(msg)

    def _log_xic(self, msg: str) -> None:
        self._log_q_xic.put(msg)


    def _poll_log(self):
        # ABC tab log
        while True:
            try:
                msg = self._log_q_abc.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_abc") and self.txt_log_abc is not None:
                self.txt_log_abc.insert("end", msg + "\n")
                self.txt_log_abc.see("end")

        # Excel sheet batch tab log
        while True:
            try:
                msg = self._log_q_excel.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_excel") and self.txt_log_excel is not None:
                self.txt_log_excel.insert("end", msg + "\n")
                self.txt_log_excel.see("end")

        # Summary patch tab log
        while True:
            try:
                msg = self._log_q_patch.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_patch") and self.txt_log_patch is not None:
                self.txt_log_patch.insert("end", msg + "\n")
                self.txt_log_patch.see("end")

        # CSV filter tab log
        while True:
            try:
                msg = self._log_q_filter.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_filter") and self.txt_log_filter is not None:
                self.txt_log_filter.insert("end", msg + "\n")
                self.txt_log_filter.see("end")

        # Response prediction tab log
        while True:
            try:
                msg = self._log_q_resp.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_resp") and self.txt_log_resp is not None:
                self.txt_log_resp.insert("end", msg + "\n")
                self.txt_log_resp.see("end")

        # SMILES builder tab log
        while True:
            try:
                msg = self._log_q_smiles.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_smiles") and self.txt_log_smiles is not None:
                self.txt_log_smiles.insert("end", msg + "\n")
                self.txt_log_smiles.see("end")

        # ESI descriptor / gradient tab log
        while True:
            try:
                msg = self._log_q_esi.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_esi") and self.txt_log_esi is not None:
                self.txt_log_esi.insert("end", msg + "\n")
                self.txt_log_esi.see("end")

        # CombiTrace-IE tab log
        while True:
            try:
                msg = self._log_q_ie.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_ie") and self.txt_log_ie is not None:
                self.txt_log_ie.insert("end", msg + "\n")
                self.txt_log_ie.see("end")

        # XIC tab log
        while True:
            try:
                msg = self._log_q_xic.get_nowait()
            except Exception:
                break
            if hasattr(self, "txt_log_xic") and self.txt_log_xic is not None:
                self.txt_log_xic.insert("end", msg + "\n")
                self.txt_log_xic.see("end")

        self.after(120, self._poll_log)


    # =====================
    # A+B+C enumerator tab
    # =====================
    def _build_abc_tab(self):
        frm = self.tab_abc
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)

        r = 0

        ttk.Label(frm, text="CSV A\uFF1A").grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_csv_a = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_csv_a).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_csv_a).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text="CSV B\uFF1A").grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_csv_b = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_csv_b).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_csv_b).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text="CSV C\uFF1A").grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_csv_c = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_csv_c).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_csv_c).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Output CSV:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_abc_out = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_abc_out).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_abc_out).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        opt = ttk.LabelFrame(frm, text='Enumeration settings')
        opt.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        opt.columnconfigure(1, weight=1)

        self.var_abc_mode = tk.StringVar(value="formula")  # formula | all
        
        self.var_abc_dedup = tk.BooleanVar(value=False)

        ttk.Label(opt, text='Element policy:').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        mode_box = ttk.Frame(opt)
        mode_box.grid(row=0, column=1, sticky="w", padx=8, pady=4)
        ttk.Radiobutton(
            mode_box,
            text='Keep C/H/N/O/P only',
            variable=self.var_abc_mode,
            value="formula",
        ).pack(side="left")
        ttk.Radiobutton(
            mode_box,
            text='Keep all elements (including Br, I, Cl, S)',
            variable=self.var_abc_mode,
            value="all",
        ).pack(side="left", padx=14)

        ttk.Checkbutton(opt, text='Deduplicate by product formula', variable=self.var_abc_dedup).grid(
            row=1, column=0, columnspan=2, sticky="w", padx=8, pady=4
        )
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)
        self.btn_run_abc = ttk.Button(btns, text='Enumerate', command=self._on_run_abc)
        self.btn_run_abc.pack(side="left")

        self.lbl_status_abc = ttk.Label(btns, text="Ready")
        self.lbl_status_abc.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_log_abc = tk.Text(frm, height=18, wrap="word")
        self.txt_log_abc.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=6)
        frm.rowconfigure(r, weight=1)

    def _browse_csv_a(self):
        path = filedialog.askopenfilename(title='Select component A CSV', filetypes=[("CSV", "*.csv"), ("All", "*")])
        if path:
            self.var_csv_a.set(path)
            
            if not self.var_abc_out.get().strip():
                p = Path(path)
                self.var_abc_out.set(str(p.with_name(f"{p.stem}__AplusBplusC.csv")))

    def _browse_csv_b(self):
        path = filedialog.askopenfilename(title='Select component B CSV', filetypes=[("CSV", "*.csv"), ("All", "*")])
        if path:
            self.var_csv_b.set(path)

    def _browse_csv_c(self):
        path = filedialog.askopenfilename(title='Select component C CSV', filetypes=[("CSV", "*.csv"), ("All", "*")])
        if path:
            self.var_csv_c.set(path)

    def _browse_abc_out(self):
        path = filedialog.asksaveasfilename(
            title='Save output CSV',
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("All", "*")],
        )
        if path:
            self.var_abc_out.set(path)

    def _on_run_abc(self):
        a_path = self.var_csv_a.get().strip()
        b_path = self.var_csv_b.get().strip()
        c_path = self.var_csv_c.get().strip()
        out_path = self.var_abc_out.get().strip()

        if not a_path or not b_path or not c_path:
            messagebox.showerror('Error', 'Select the component A, B and C CSV files.')
            return
        if not out_path:
            messagebox.showerror('Error', 'Select an output CSV file.')
            return

        only_formula = (self.var_abc_mode.get().strip() == "formula")
        dedup = bool(self.var_abc_dedup.get())

        
        self.btn_run_abc.config(state="disabled")
        self.lbl_status_abc.config(text="Running...")
        self.txt_log_abc.delete("1.0", "end")

        def worker():
            try:
                self._log_abc(f"Loading A: {a_path}")
                a_rows = load_component_csv(Path(a_path))
                self._log_abc(f"A rows: {len(a_rows)}")

                self._log_abc(f"Loading B: {b_path}")
                b_rows = load_component_csv(Path(b_path))
                self._log_abc(f"B rows: {len(b_rows)}")

                self._log_abc(f"Loading C: {c_path}")
                c_rows = load_component_csv(Path(c_path))
                self._log_abc(f"C rows: {len(c_rows)}")

                total = len(a_rows) * len(b_rows) * len(c_rows)
                self._log_abc(f"Total combinations: {total}")
                self._log_abc(
                    f"Mode: {'ONLY_CHNOP' if only_formula else 'ALL_ELEMENTS'} | Dedup: {dedup}"
                )

                n_written = 0
                if dedup:
                    results, skipped_invalid = enumerate_abc(
                        a_rows,
                        b_rows,
                        c_rows,
                        only_formula=only_formula,
                        deduplicate=True,
                    )
                    self._log_abc(f"Valid results: {len(results)}")
                    n_written = len(results)
                    if skipped_invalid:
                        self._log_abc(f"Skipped invalid combos (negative counts): {skipped_invalid}")
                    write_results_csv(results, Path(out_path))
                    self._log_abc(f"Written: {out_path}")
                else:
                    results_d, skipped_invalid = enumerate_abc_detailed(
                        a_rows,
                        b_rows,
                        c_rows,
                        only_formula=only_formula,
                    )
                    self._log_abc(f"Valid results (with duplicates): {len(results_d)}")
                    n_written = len(results_d)
                    if skipped_invalid:
                        self._log_abc(f"Skipped invalid combos (negative counts): {skipped_invalid}")
                    write_results_csv_detailed(results_d, Path(out_path))
                    self._log_abc(f"Written: {out_path}")

                self.after(
                    0,
                    lambda: messagebox.showinfo(
                        'Complete',
                        f'Enumeration complete\nResults={n_written}, SkippedInvalid={skipped_invalid}\nOutput: {out_path}',
                    ),
                )
            except Exception as e:
                self._log_abc(f"ERROR: {e}")
                self.after(0, lambda: messagebox.showerror('Failed', str(e)))
            finally:
                self.after(0, lambda: self.btn_run_abc.config(state="normal"))
                self.after(0, lambda: self.lbl_status_abc.config(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    # =====================
    # Excel worksheets -> A/C/E enumeration -> triplicate XIC
    # =====================
    def _build_excel_sheet_tab(self):
        from core.ui_support import scroll_body
        frm = scroll_body(self.tab_excel)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)
        r = 0

        ttk.Label(frm, text='Component workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_excel_workbook = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_excel_workbook).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_excel_workbook).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='RAW directory:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_excel_raw_dir = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_excel_raw_dir).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_excel_raw_dir).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Output directory:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_excel_out_dir = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_excel_out_dir).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_excel_out_dir).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        note = ttk.LabelFrame(frm, text='Workbook layout')
        note.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=5)
        note.columnconfigure(1, weight=1)
        ttk.Label(
            note,
            text='Each sheet contains components A, B and C. Select the layout used in each component block. Original component names are retained in Combo.',
        ).grid(row=0, column=0, columnspan=4, sticky="w", padx=8, pady=4)

        self.var_excel_layout_label = tk.StringVar(value=EXCEL_LAYOUT_LABELS[0])
        ttk.Label(note, text='Layout:').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(
            note,
            textvariable=self.var_excel_layout_label,
            values=EXCEL_LAYOUT_LABELS,
            state="readonly",
            width=38,
        ).grid(row=1, column=1, columnspan=3, sticky="w", padx=8, pady=4)

        self.var_excel_suffixes = tk.StringVar(value="-1,-2,-3")
        ttk.Label(note, text='Replicate RAW suffixes:').grid(row=2, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(note, textvariable=self.var_excel_suffixes, width=22).grid(row=2, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(note, text='Example: sheet 1-12-x maps to 1-12-x-1, -2 and -3.RAW.').grid(
            row=2, column=2, columnspan=2, sticky="w", padx=8, pady=4
        )

        self.var_excel_mode = tk.StringVar(value="all")
        ttk.Label(note, text='Element policy:').grid(row=3, column=0, sticky="w", padx=8, pady=4)
        mode_fr = ttk.Frame(note)
        mode_fr.grid(row=3, column=1, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Radiobutton(
            mode_fr,
            text='Keep all elements (including Cl, Br, S)',
            variable=self.var_excel_mode,
            value="all",
        ).pack(side="left")
        ttk.Radiobutton(
            mode_fr,
            text='Keep C/H/N/O/P only',
            variable=self.var_excel_mode,
            value="formula",
        ).pack(side="left", padx=(16, 0))
        r += 1

        opt = ttk.LabelFrame(frm, text='Run options')
        opt.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=5)
        for c in range(6):
            opt.columnconfigure(c, weight=1, uniform="options")

        self.var_excel_run_xic = tk.BooleanVar(value=True)
        self.var_excel_recursive = tk.BooleanVar(value=False)
        self.var_excel_export_ms2 = tk.BooleanVar(value=False)
        self.var_excel_ppm = tk.DoubleVar(value=5.0)
        self.var_excel_min_peak_height = tk.DoubleVar(value=1e4)
        self.var_excel_dup_min_rel_height = tk.DoubleVar(value=0.10)
        self.var_excel_multipeak_mode = tk.StringVar(value=MULTIPEAK_MODE_LABELS[0])
        self.var_excel_adduct_mode = tk.StringVar(value=ADDUCT_MODE_LABELS[0])
        self.var_excel_forced_adduct = tk.StringVar(value="M+H")
        self.var_excel_combine_formate = tk.BooleanVar(value=True)
        self.var_excel_allow_formate_only = tk.BooleanVar(value=False)
        self.var_excel_formate_include_is = tk.BooleanVar(value=True)
        self.var_excel_formate_rt_tolerance = tk.DoubleVar(value=0.10)
        self.var_excel_formate_min_rel_pct = tk.DoubleVar(value=1.0)
        self.var_excel_formate_min_corr = tk.DoubleVar(value=0.30)

        ttk.Checkbutton(opt, text='Extract triplicate XIC after enumeration', variable=self.var_excel_run_xic).grid(
            row=0, column=0, columnspan=2, sticky="w", padx=8, pady=4
        )
        ttk.Checkbutton(opt, text='Search RAW subdirectories', variable=self.var_excel_recursive).grid(
            row=0, column=2, columnspan=2, sticky="w", padx=8, pady=4
        )
        ttk.Checkbutton(opt, text='Export MS2 for DDA files', variable=self.var_excel_export_ms2).grid(
            row=0, column=4, columnspan=2, sticky="w", padx=8, pady=4
        )

        ttk.Label(opt, text='Extraction tolerance (ppm)').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_ppm, width=10).grid(row=1, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Minimum peak height').grid(row=1, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_min_peak_height, width=12).grid(row=1, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Duplicate relative height').grid(row=1, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_dup_min_rel_height, width=10).grid(row=1, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Duplicate-peak mode').grid(row=2, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(
            opt,
            textvariable=self.var_excel_multipeak_mode,
            values=MULTIPEAK_MODE_LABELS,
            state="readonly",
            width=31,
        ).grid(row=2, column=1, columnspan=2, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Ion matching').grid(row=2, column=3, sticky="w", padx=8, pady=4)
        self.cmb_excel_adduct_mode = ttk.Combobox(
            opt,
            textvariable=self.var_excel_adduct_mode,
            values=ADDUCT_MODE_LABELS,
            state="readonly",
            width=28,
        )
        self.cmb_excel_adduct_mode.grid(row=2, column=4, sticky="w", padx=8, pady=4)
        self.cmb_excel_adduct_mode.bind("<<ComboboxSelected>>", lambda e: self._update_excel_adduct_ui())
        self.cmb_excel_forced_adduct = ttk.Combobox(
            opt,
            textvariable=self.var_excel_forced_adduct,
            values=["M+H", "M-H", "M+HCOO", "M+Na", "M+K", "M+NH4"],
            state="disabled",
            width=10,
        )
        self.cmb_excel_forced_adduct.grid(row=2, column=5, sticky="w", padx=8, pady=4)

        self.var_excel_internal_standard_formula = tk.StringVar(value="")
        self.var_excel_is_adduct = tk.StringVar(value=INTERNAL_STANDARD_ADDUCT_LABELS[0])
        self.var_excel_is_ppm = tk.DoubleVar(value=10.0)
        self.var_excel_is_adaptive = tk.BooleanVar(value=True)
        self.var_excel_is_min_snr = tk.DoubleVar(value=5.0)
        self.var_excel_is_min_fraction = tk.DoubleVar(value=0.10)
        self.var_excel_is_expected_rt = tk.StringVar(value="")
        self.var_excel_is_rt_tolerance = tk.DoubleVar(value=0.30)
        ttk.Label(opt, text='Common internal-standard formula').grid(row=3, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_internal_standard_formula, width=28).grid(
            row=3, column=1, columnspan=2, sticky="w", padx=8, pady=4
        )
        ttk.Label(
            opt,
            text='One common IS across sheets and replicates. Ratios are Area/IS x 100%.',
        ).grid(row=3, column=3, columnspan=3, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Internal-standard ion').grid(row=4, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(
            opt, textvariable=self.var_excel_is_adduct, values=INTERNAL_STANDARD_ADDUCT_LABELS,
            state="readonly", width=43,
        ).grid(row=4, column=1, columnspan=2, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='IS tolerance (ppm)').grid(row=4, column=3, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_is_ppm, width=10).grid(row=4, column=4, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(
            opt,
            text='Adaptive IS QC',
            variable=self.var_excel_is_adaptive,
        ).grid(row=4, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Expected IS RT (min; optional)').grid(row=5, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_is_expected_rt, width=10).grid(row=5, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text="RT tolerance (min)").grid(row=5, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_is_rt_tolerance, width=10).grid(row=5, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='IS SNR / height fraction').grid(row=5, column=4, sticky="w", padx=8, pady=4)
        is_adapt_fr = ttk.Frame(opt)
        is_adapt_fr.grid(row=5, column=5, sticky="w", padx=8, pady=4)
        ttk.Entry(is_adapt_fr, textvariable=self.var_excel_is_min_snr, width=6).pack(side="left")
        ttk.Label(is_adapt_fr, text=" / ").pack(side="left")
        ttk.Entry(is_adapt_fr, textvariable=self.var_excel_is_min_fraction, width=6).pack(side="left")
        ttk.Label(
            opt,
            text="IS XIC, raw trace CSV and diagnostic CSV are always written, even when Found=False.",
        ).grid(row=12, column=0, columnspan=6, sticky="w", padx=8, pady=(3, 6))

        ttk.Checkbutton(
            opt,
            text='Sum coeluting [M-H]- and formate',
            variable=self.var_excel_combine_formate,
        ).grid(row=6, column=0, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(
            opt,
            text='Allow formate-only peak',
            variable=self.var_excel_allow_formate_only,
        ).grid(row=6, column=3, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text="Formate RT tolerance (min)").grid(row=7, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_formate_rt_tolerance, width=10).grid(row=7, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Formate min. height (%)').grid(row=7, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_formate_min_rel_pct, width=10).grid(row=7, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Min. shape correlation').grid(row=7, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_formate_min_corr, width=10).grid(row=7, column=5, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(
            opt,
            text='Apply selected ion panel to IS',
            variable=self.var_excel_formate_include_is,
        ).grid(row=8, column=0, columnspan=3, sticky="w", padx=8, pady=4)
        self.var_excel_neg_channel_mode = tk.StringVar(value=NEG_CHANNEL_MODE_LABELS[1])
        self.var_excel_neg_panel_file = tk.StringVar(value="")
        self.var_excel_panel_prevalence = tk.DoubleVar(value=20.0)
        self.var_excel_panel_isotope_support = tk.DoubleVar(value=50.0)
        self.var_excel_panel_min_area_pct = tk.DoubleVar(value=0.5)
        ttk.Label(opt, text="Negative-ion channel mode").grid(row=9, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(opt, textvariable=self.var_excel_neg_channel_mode, values=NEG_CHANNEL_MODE_LABELS, state="readonly", width=46).grid(row=9, column=1, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text="Selected panel CSV").grid(row=10, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_excel_neg_panel_file).grid(row=10, column=1, columnspan=4, sticky="ew", padx=8, pady=4)
        ttk.Button(opt, text="Browse...", command=self._browse_excel_neg_panel).grid(row=10, column=5, padx=8, pady=4)

        review_box = ttk.LabelFrame(opt, text='Channel discovery and review')
        review_box.grid(row=11, column=0, columnspan=6, sticky="ew", padx=8, pady=(6, 8))
        for cc in range(8):
            review_box.columnconfigure(cc, weight=1 if cc in (1, 3, 5) else 0)
        ttk.Label(review_box, text="Present threshold (%)").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(review_box, textvariable=self.var_excel_panel_prevalence, width=8).grid(row=0, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(review_box, text="Isotope support (%)").grid(row=0, column=2, sticky="w", padx=6, pady=4)
        ttk.Entry(review_box, textvariable=self.var_excel_panel_isotope_support, width=8).grid(row=0, column=3, sticky="w", padx=6, pady=4)
        ttk.Label(review_box, text="Median area vs [M-H] (%)").grid(row=0, column=4, sticky="w", padx=6, pady=4)
        ttk.Entry(review_box, textvariable=self.var_excel_panel_min_area_pct, width=8).grid(row=0, column=5, sticky="w", padx=6, pady=4)
        self.btn_excel_review_neg = ttk.Button(
            review_box,
            text='Review discovered channels...',
            command=self._on_review_excel_discovery,
        )
        self.btn_excel_review_neg.grid(row=1, column=0, columnspan=3, sticky="ew", padx=6, pady=5)
        self.btn_excel_apply_cached = ttk.Button(
            review_box,
            text='Apply panel to existing results',
            command=self._on_apply_excel_selected_panel_cached,
        )
        self.btn_excel_apply_cached.grid(row=1, column=3, columnspan=3, sticky="ew", padx=6, pady=5)
        ttk.Label(
            review_box,
            text=("Run Discover all once. Then review channel prevalence/area/RT/shape evidence, mark which "
                  "channels are included in quantitative Area, and recalculate the triplicate summary from cached CSVs."),
            wraplength=980,
            justify="left",
        ).grid(row=2, column=0, columnspan=8, sticky="w", padx=6, pady=(2, 5))
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=8)
        self.btn_run_excel = ttk.Button(btns, text='Run enumeration and XIC', command=self._on_run_excel_sheet)
        self.btn_run_excel.pack(side="left")
        ttk.Button(btns, text='Open output directory', command=self._open_excel_out_dir).pack(side="left", padx=10)
        self.lbl_status_excel = ttk.Label(btns, text="Ready")
        self.lbl_status_excel.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=5)
        self.txt_log_excel = tk.Text(frm, height=13, wrap="word")
        self.txt_log_excel.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=5)
        frm.rowconfigure(r, weight=1)
        self._update_excel_adduct_ui()

    def _update_excel_adduct_ui(self):
        try:
            code = ADDUCT_MODE_MAP.get((self.var_excel_adduct_mode.get() or "").strip(), "posneg_neghonly")
            self.cmb_excel_forced_adduct.configure(state="readonly" if code == "force" else "disabled")
        except Exception:
            pass

    def _browse_excel_workbook(self):
        path = filedialog.askopenfilename(
            title='Select component workbook',
            filetypes=[('Component workbook', "*.xlsx *.xlsm"), ("All", "*")],
        )
        if path:
            self.var_excel_workbook.set(path)
            if not self.var_excel_out_dir.get().strip():
                p = Path(path)
                self.var_excel_out_dir.set(str(p.with_name(f"{p.stem}__sheet_xic_output")))

    def _browse_excel_raw_dir(self):
        path = filedialog.askdirectory(title='Select triplicate RAW directory')
        if path:
            self.var_excel_raw_dir.set(path)

    def _browse_excel_out_dir(self):
        path = filedialog.askdirectory(title='Select output directory')
        if path:
            self.var_excel_out_dir.set(path)

    def _browse_excel_neg_panel(self):
        path = filedialog.askopenfilename(title="Select negative-ion panel", filetypes=[("Panel", "*.csv *.json"), ("CSV", "*.csv"), ("JSON", "*.json"), ("All", "*")])
        if path:
            self.var_excel_neg_panel_file.set(path)

    def _on_review_excel_discovery(self):
        """Aggregate and review discover-all evidence in the triplicate output folder."""
        out_dir = self.var_excel_out_dir.get().strip()
        if not out_dir:
            messagebox.showerror("Error", "Select the triplicate output folder first.")
            return
        root = Path(out_dir)
        if not root.exists():
            messagebox.showerror("Error", f"Output folder does not exist:\n{root}")
            return
        try:
            prevalence = float(self.var_excel_panel_prevalence.get())
            isotope_support = float(self.var_excel_panel_isotope_support.get())
            min_area = float(self.var_excel_panel_min_area_pct.get())
        except Exception as exc:
            messagebox.showerror("Invalid thresholds", str(exc))
            return

        self.btn_excel_review_neg.configure(state="disabled")
        self.btn_excel_apply_cached.configure(state="disabled")
        self.lbl_status_excel.configure(text="Building discovery review...")
        self._log_excel("Refreshing discover-all channel summary from existing evidence CSV files...")

        def worker():
            try:
                summary_xlsx, suggested_panel, summary_rows, evidence_files = refresh_discovery_summary(
                    root,
                    prevalence_threshold_pct=prevalence,
                    min_isotope_support_pct=isotope_support,
                    min_median_area_fraction_pct=min_area,
                )
                selected_panel = root / "negative_ion_selected_panel.csv"
                current_panel_text = self.var_excel_neg_panel_file.get().strip()
                existing_panel = Path(current_panel_text) if current_panel_text and Path(current_panel_text).exists() else None
                if existing_panel is None and selected_panel.exists():
                    existing_panel = selected_panel
                suggested_rows = load_panel_review_rows(suggested_panel)
                review_rows = merge_existing_actions(suggested_rows, existing_panel)
                self._log_excel(
                    f"Discovery review ready: evidence_files={len(evidence_files)}, channels={len(review_rows)}"
                )
                self._log_excel(f"Summary workbook: {summary_xlsx}")
                self.after(
                    0,
                    lambda: self._show_excel_negative_ion_review(
                        review_rows,
                        summary_xlsx=summary_xlsx,
                        selected_panel_path=selected_panel,
                        evidence_count=len(evidence_files),
                    ),
                )
            except Exception as exc:
                msg = str(exc)
                self._log_excel("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror("Discovery review failed", m))
            finally:
                self.after(0, lambda: self.btn_excel_review_neg.configure(state="normal"))
                self.after(0, lambda: self.btn_excel_apply_cached.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_excel.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    def _show_excel_negative_ion_review(
        self,
        review_rows,
        *,
        summary_xlsx: Path,
        selected_panel_path: Path,
        evidence_count: int,
    ):
        """Show an intuitive channel-level review table and save a fixed panel."""
        top = tk.Toplevel(self)
        top.title('Negative-ion channel review')
        top.geometry("1320x780")
        top.minsize(1050, 620)
        top.columnconfigure(0, weight=1)
        top.rowconfigure(1, weight=1)

        action_text = {
            "sum": "SUM into quantitative area",
            "search": "Report only (do not sum)",
            "evidence": 'Evidence only',
            "exclude": "Exclude",
        }
        rows_by_id = {}
        for raw in review_rows:
            row = {str(k): ("" if v is None else str(v)) for k, v in dict(raw).items()}
            cid = str(row.get("Channel_ID", "") or "").strip().upper()
            if not cid:
                continue
            row.setdefault("Auto_Suggested_Action", row.get("Quant_Action", "search"))
            rows_by_id[cid] = row

        hdr = ttk.Frame(top)
        hdr.grid(row=0, column=0, sticky="ew", padx=10, pady=8)
        hdr.columnconfigure(0, weight=1)
        ttk.Label(
            hdr,
            text=(
                f"Evidence files: {evidence_count}. Green rows will be summed into Area; blue/yellow rows are "
                "reported as evidence only; grey rows are ignored. Select rows and use the buttons below."
            ),
            wraplength=1120,
            justify="left",
        ).grid(row=0, column=0, sticky="w")
        ttk.Button(hdr, text="Open detailed summary Excel", command=lambda: open_in_os(summary_xlsx)).grid(
            row=0, column=1, sticky="e", padx=(10, 0)
        )

        table_frame = ttk.Frame(top)
        table_frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 6))
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)
        columns = (
            "Action", "Channel", "Category", "Detected", "DetectionRate", "MedianArea",
            "MedianRT", "ShapeCorr", "Isotope", "Assessment", "Suggested",
        )
        tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="extended")
        headings = {
            "Action": 'Action',
            "Channel": 'Ion channel',
            "Category": 'Type',
            "Detected": 'Detected / eligible',
            "DetectionRate": 'Detection (%)',
            "MedianArea": 'Median area vs [M-H] (%)',
            "MedianRT": 'Median RT difference',
            "ShapeCorr": 'Median shape correlation',
            "Isotope": 'Isotope support (%)',
            "Assessment": 'Presence',
            "Suggested": 'Suggested action',
        }
        widths = {
            "Action": 180, "Channel": 220, "Category": 90, "Detected": 110,
            "DetectionRate": 90, "MedianArea": 135, "MedianRT": 105,
            "ShapeCorr": 105, "Isotope": 105, "Assessment": 125, "Suggested": 120,
        }
        for col in columns:
            tree.heading(col, text=headings[col])
            tree.column(col, width=widths[col], minwidth=75, anchor="center" if col not in {"Action", "Channel"} else "w")
        ybar = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        xbar = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        tree.tag_configure("sum", background="#DDF2D8")
        tree.tag_configure("search", background="#DCEBFA")
        tree.tag_configure("evidence", background="#FFF1C7")
        tree.tag_configure("exclude", foreground="#777777", background="#EEEEEE")

        def ftext(value, decimals=2):
            try:
                if value is None or str(value).strip() in {"", "NA", "nan"}:
                    return "-"
                return f"{float(value):.{decimals}f}"
            except Exception:
                return str(value or "-")

        def refresh_item(cid):
            row = rows_by_id[cid]
            action = str(row.get("Quant_Action", "search") or "search").strip().lower()
            if action not in action_text:
                action = "search"
                row["Quant_Action"] = action
            detected = f"{row.get('Detected_Count', '0')} / {row.get('Eligible_Count', '0')}"
            values = (
                action_text[action],
                row.get("Display", row.get("Channel", cid)),
                row.get("Category", ""),
                detected,
                ftext(row.get("Detection_Rate_Pct", ""), 1),
                ftext(row.get("Median_Area_Fraction_vs_MH_Pct", ""), 2),
                ftext(row.get("Median_RT_Delta_Min", ""), 4),
                ftext(row.get("Median_Shape_Correlation", ""), 3),
                ftext(row.get("Isotope_Support_Rate_Pct", ""), 1),
                row.get("Suggested_Status", ""),
                action_text.get(str(row.get("Auto_Suggested_Action", "search")).lower(), row.get("Auto_Suggested_Action", "")),
            )
            if tree.exists(cid):
                tree.item(cid, values=values, tags=(action,))
            else:
                tree.insert("", "end", iid=cid, values=values, tags=(action,))

        for cid in rows_by_id:
            refresh_item(cid)

        detail_var = tk.StringVar(value="Select a channel to view its source hypothesis and interpretation note.")
        detail = ttk.Label(top, textvariable=detail_var, wraplength=1260, justify="left")
        detail.grid(row=2, column=0, sticky="ew", padx=12, pady=(0, 6))

        def show_detail(_event=None):
            selected = tree.selection()
            if not selected:
                return
            cid = selected[0]
            row = rows_by_id[cid]
            defn = CHANNEL_BY_ID.get(cid)
            quantifiable = bool(defn.quantifiable) if defn is not None else False
            detail_var.set(
                f"{cid} | {row.get('Display', '')} | source: {row.get('Source_Hypothesis', '')} | "
                f"priority: {row.get('Priority_Stars', '')} | quantifiable: {quantifiable}. "
                f"{row.get('Panel_Note', row.get('Note', ''))}"
            )
        tree.bind("<<TreeviewSelect>>", show_detail)

        controls = ttk.Frame(top)
        controls.grid(row=3, column=0, sticky="ew", padx=10, pady=6)

        def set_action(action):
            selected = tree.selection()
            if not selected:
                messagebox.showinfo("Select channels", "Select one or more channel rows first.", parent=top)
                return
            changed = 0
            downgraded = []
            for cid in selected:
                row = rows_by_id[cid]
                defn = CHANNEL_BY_ID.get(cid)
                chosen = action
                if action == "sum" and defn is not None and not defn.quantifiable:
                    chosen = "evidence"
                    downgraded.append(cid)
                row["Quant_Action"] = chosen
                row["Enabled"] = "False" if chosen == "exclude" else "True"
                refresh_item(cid)
                changed += 1
            if downgraded:
                messagebox.showwarning(
                    "Evidence-only channels",
                    "These channels cannot be summed and were set to evidence-only:\n" + ", ".join(downgraded),
                    parent=top,
                )
            detail_var.set(f"Updated {changed} selected channel(s).")

        def use_suggestions():
            for cid, row in rows_by_id.items():
                action = str(row.get("Auto_Suggested_Action", "search") or "search").lower()
                defn = CHANNEL_BY_ID.get(cid)
                if action == "sum" and defn is not None and not defn.quantifiable:
                    action = "evidence"
                row["Quant_Action"] = action
                row["Enabled"] = "False" if action == "exclude" else "True"
                refresh_item(cid)
            detail_var.set("Automatic suggestions restored.")

        ttk.Button(controls, text="Use auto suggestions", command=use_suggestions).pack(side="left")
        ttk.Button(controls, text='Sum into quantitative area', command=lambda: set_action("sum")).pack(side="left", padx=(8, 0))
        ttk.Button(controls, text='Report only', command=lambda: set_action("search")).pack(side="left", padx=(8, 0))
        ttk.Button(controls, text='Evidence only', command=lambda: set_action("evidence")).pack(side="left", padx=(8, 0))
        ttk.Button(controls, text='Exclude', command=lambda: set_action("exclude")).pack(side="left", padx=(8, 0))

        def cycle_action(_event=None):
            selected = tree.selection()
            if not selected:
                return
            cid = selected[0]
            current = str(rows_by_id[cid].get("Quant_Action", "search") or "search").lower()
            sequence = ["sum", "search", "evidence", "exclude"]
            next_action = sequence[(sequence.index(current) + 1) % len(sequence)] if current in sequence else "sum"
            set_action(next_action)
        tree.bind("<Double-1>", cycle_action)

        footer = ttk.Frame(top)
        footer.grid(row=4, column=0, sticky="ew", padx=10, pady=(2, 10))
        footer.columnconfigure(0, weight=1)
        status_var = tk.StringVar(value=f"Selected panel will be saved as: {selected_panel_path}")
        ttk.Label(footer, textvariable=status_var, wraplength=760, justify="left").grid(row=0, column=0, sticky="w")

        def save_panel(show_message=True):
            saved = save_panel_review_rows(
                list(rows_by_id.values()),
                selected_panel_path,
                panel_name="selected_negative_ion_experiment_panel",
            )
            self.var_excel_neg_panel_file.set(str(saved))
            panel_label = next((label for label, code in NEG_CHANNEL_MODE_OPTIONS if code == "panel"), NEG_CHANNEL_MODE_LABELS[-1])
            self.var_excel_neg_channel_mode.set(panel_label)
            n_sum = sum(1 for r in rows_by_id.values() if str(r.get("Quant_Action", "")).lower() == "sum")
            status_var.set(f"Saved: {saved} | channels summed into Area: {n_sum}")
            self._log_excel(f"Selected negative-ion panel saved: {saved} | SUM channels={n_sum}")
            if show_message:
                messagebox.showinfo(
                    "Panel saved",
                    f"Selected panel saved.\nSUM channels: {n_sum}\n\n{saved}",
                    parent=top,
                )
            return saved

        def save_and_apply():
            save_panel(show_message=False)
            top.destroy()
            self.after(100, self._on_apply_excel_selected_panel_cached)

        ttk.Button(footer, text='Save panel', command=save_panel).grid(row=0, column=1, padx=6)
        ttk.Button(
            footer,
            text='Save and apply',
            command=save_and_apply,
        ).grid(row=0, column=2, padx=6)
        ttk.Button(footer, text="Close", command=top.destroy).grid(row=0, column=3, padx=(6, 0))

    def _on_apply_excel_selected_panel_cached(self):
        """Apply a selected panel to existing discover-all CSVs without re-reading RAW."""
        workbook = self.var_excel_workbook.get().strip()
        discovery_root_text = self.var_excel_out_dir.get().strip()
        panel_text = self.var_excel_neg_panel_file.get().strip()
        if not workbook:
            messagebox.showerror("Error", "Select the original A/B/C Excel workbook first.")
            return
        if not discovery_root_text:
            messagebox.showerror("Error", "Select the existing triplicate discover-all output folder first.")
            return
        discovery_root = Path(discovery_root_text)
        if not discovery_root.exists():
            messagebox.showerror("Error", f"Discovery output folder does not exist:\n{discovery_root}")
            return
        if not panel_text or not Path(panel_text).exists():
            candidate = discovery_root / "negative_ion_selected_panel.csv"
            if candidate.exists():
                panel_text = str(candidate)
                self.var_excel_neg_panel_file.set(panel_text)
            else:
                messagebox.showerror(
                    "No selected panel",
                    "Review the discover-all results and save a selected panel first.",
                )
                return
        try:
            suffixes = parse_replicate_suffixes(self.var_excel_suffixes.get())
            if len(suffixes) != 3:
                raise ValueError("Three replicate suffixes are required, for example -1,-2,-3")
            layout = EXCEL_LAYOUT_MAP.get((self.var_excel_layout_label.get() or "").strip(), "auto")
            internal_formula = self.var_excel_internal_standard_formula.get().strip()
        except Exception as exc:
            messagebox.showerror("Invalid settings", str(exc))
            return

        output_root = discovery_root / "selected_panel_results"
        if output_root.exists():
            try:
                shutil.rmtree(output_root)
            except Exception as exc:
                messagebox.showerror("Cannot refresh selected-panel results", str(exc))
                return

        self.btn_excel_review_neg.configure(state="disabled")
        self.btn_excel_apply_cached.configure(state="disabled")
        self.btn_run_excel.configure(state="disabled")
        self.lbl_status_excel.configure(text="Applying selected panel to cached results...")
        self._log_excel(f"Applying selected panel without RAW re-read: {panel_text}")
        self._log_excel(f"Cached discovery root: {discovery_root}")

        def worker():
            try:
                report = apply_panel_to_existing_triplicate(
                    workbook_path=Path(workbook),
                    discovery_root=discovery_root,
                    selected_panel_file=Path(panel_text),
                    output_root=output_root,
                    suffixes=suffixes,
                    layout=layout,
                    only_formula=(self.var_excel_mode.get().strip() == "formula"),
                    internal_standard_formula=internal_formula,
                    include_for_internal_standard=bool(self.var_excel_formate_include_is.get()),
                    progress=self._log_excel,
                )
                self._last_out_dir = report.output_root
                self._log_excel(f"Selected-panel triplicate summary: {report.summary_xlsx}")
                self._log_excel(f"Selected-panel triplicate average: {report.average_csv}")
                self._log_excel(f"Application audit: {report.audit_csv}")
                if report.warnings:
                    for warning in report.warnings[:30]:
                        self._log_excel("WARN: " + str(warning))
                self.after(
                    0,
                    lambda: messagebox.showinfo(
                        "Cached panel application complete",
                        f"Updated RAW result CSVs: {report.processed_raws}\n"
                        f"Missing quant: {len(report.missing_quant_files)}\n"
                        f"Missing evidence: {len(report.missing_evidence_files)}\n\n"
                        f"Summary:\n{report.summary_xlsx}\n\n"
                        "This operation reused the discover-all CSV evidence and did not open RAW files.",
                    ),
                )
            except Exception as exc:
                msg = str(exc)
                self._log_excel("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror("Cached panel application failed", m))
            finally:
                self.after(0, lambda: self.btn_excel_review_neg.configure(state="normal"))
                self.after(0, lambda: self.btn_excel_apply_cached.configure(state="normal"))
                self.after(0, lambda: self.btn_run_excel.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_excel.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()


    def _open_excel_out_dir(self):
        p = self.var_excel_out_dir.get().strip()
        if p:
            open_in_os(Path(p))

    def _on_run_excel_sheet(self):
        workbook = self.var_excel_workbook.get().strip()
        raw_dir = self.var_excel_raw_dir.get().strip()
        out_dir = self.var_excel_out_dir.get().strip()
        run_xic = bool(self.var_excel_run_xic.get())

        if not workbook:
            messagebox.showerror('Error', 'Select the component workbook.')
            return
        if Path(workbook).suffix.lower() not in (".xlsx", ".xlsm"):
            messagebox.showerror('Error', 'Supported formats: .xlsx and .xlsm. Convert .xls files to .xlsx first.')
            return
        if run_xic and not raw_dir:
            messagebox.showerror('Error', 'Select a RAW directory to run XIC extraction.')
            return
        if not out_dir:
            messagebox.showerror('Error', 'Select an output directory.')
            return

        try:
            suffixes = parse_replicate_suffixes(self.var_excel_suffixes.get())
            if len(suffixes) != 3:
                raise ValueError('Enter three replicate suffixes, for example -1,-2,-3.')
            ppm = float(self.var_excel_ppm.get())
            min_peak_height = float(self.var_excel_min_peak_height.get())
            dup_min_rel = float(self.var_excel_dup_min_rel_height.get())
            excel_combine_formate = bool(self.var_excel_combine_formate.get())
            excel_allow_formate_only = bool(self.var_excel_allow_formate_only.get())
            excel_formate_include_is = bool(self.var_excel_formate_include_is.get())
            excel_formate_rt_tolerance = float(self.var_excel_formate_rt_tolerance.get())
            excel_formate_min_rel_pct = float(self.var_excel_formate_min_rel_pct.get())
            excel_formate_min_corr = float(self.var_excel_formate_min_corr.get())
            excel_neg_channel_mode = NEG_CHANNEL_MODE_MAP.get(self.var_excel_neg_channel_mode.get(), "legacy")
            excel_neg_panel_file = self.var_excel_neg_panel_file.get().strip()
            excel_panel_prevalence = float(self.var_excel_panel_prevalence.get())
            excel_panel_isotope_support = float(self.var_excel_panel_isotope_support.get())
            excel_panel_min_area_pct = float(self.var_excel_panel_min_area_pct.get())
            excel_is_adduct = INTERNAL_STANDARD_ADDUCT_MAP.get(self.var_excel_is_adduct.get(), "M-H")
            excel_is_ppm = float(self.var_excel_is_ppm.get())
            excel_is_adaptive = bool(self.var_excel_is_adaptive.get())
            excel_is_min_snr = float(self.var_excel_is_min_snr.get())
            excel_is_min_fraction = float(self.var_excel_is_min_fraction.get())
            excel_is_rt_tolerance = float(self.var_excel_is_rt_tolerance.get())
            is_rt_text = self.var_excel_is_expected_rt.get().strip()
            excel_is_expected_rt = float(is_rt_text) if is_rt_text else None
            if excel_is_ppm <= 0 or excel_is_min_snr < 0 or not (0 <= excel_is_min_fraction <= 1) or excel_is_rt_tolerance <= 0:
                raise ValueError("Invalid internal-standard diagnostic settings")
            internal_standard_formula, internal_standard_mass = normalize_internal_standard_formula(
                self.var_excel_internal_standard_formula.get()
            )
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return

        self.btn_run_excel.configure(state="disabled")
        self.lbl_status_excel.configure(text="Running...")
        self.txt_log_excel.delete("1.0", "end")

        def worker():
            try:
                out_root = Path(out_dir)
                out_root.mkdir(parents=True, exist_ok=True)
                self._last_out_dir = out_root

                self._log_excel(f'Reading workbook: {workbook}')
                if internal_standard_formula:
                    self._log_excel(
                        f'Common internal standard: {internal_standard_formula} | exact mass={internal_standard_mass:.6f} | Shared across sheets and replicates; Area/IS x 100%'
                    )
                else:
                    self._log_excel('No common internal standard: ratios will be percentages of total area.')
                excel_layout = EXCEL_LAYOUT_MAP.get(
                    (self.var_excel_layout_label.get() or "").strip(),
                    "auto",
                )
                self._log_excel(f'Workbook layout: {self.var_excel_layout_label.get()} ({excel_layout})')
                plans, warnings = load_excel_sheet_plans(
                    Path(workbook),
                    formula_columns=(1, 3, 5),
                    header_row=1,
                    data_start_row=2,
                    only_formula=(self.var_excel_mode.get().strip() == "formula"),
                    layout=excel_layout,
                )
                for w in warnings:
                    self._log_excel("WARN: " + w)
                if not plans:
                    raise RuntimeError('No valid component sheets were found. Check the selected layout and the three component column blocks.')
                self._log_excel(f'Valid sheets: {len(plans)}')

                raw_paths: List[Path] = []
                if raw_dir:
                    raw_paths = _find_raw_paths(Path(raw_dir), recursive=bool(self.var_excel_recursive.get()))
                    self._log_excel(f'RAW files: {len(raw_paths)}')
                    if run_xic and not raw_paths:
                        raise RuntimeError('No .raw files were found in the selected directory.')

                plans_and_raws = []
                run_records = []
                raw_ok = 0
                raw_missing = 0
                raw_failed = 0
                learned_internal_standard_rts: List[float] = []

                for si, plan in enumerate(plans, start=1):
                    self._log_excel(
                        f'\n[{si}/{len(plans)}] Sheet={plan.sheet_name} | A={len(plan.a_rows)}, B={len(plan.b_rows)}, C={len(plan.c_rows)}; combinations={len(plan.targets)} (duplicates retained)'
                    )
                    resolutions = resolve_sheet_raws(plan.sheet_name, raw_paths, suffixes=suffixes)
                    plans_and_raws.append((plan, resolutions))
                    for rr in resolutions:
                        if rr.raw_path:
                            self._log_excel(
                                f'  Replicate {rr.replicate_label}: {rr.raw_path.name} [{rr.match_mode}]'
                            )
                        else:
                            raw_missing += 1
                            self._log_excel(
                                f'  Replicate {rr.replicate_label}: MISSING ({rr.expected_stem}.RAW) {rr.note}'
                            )

                    sheet_dir = out_root / safe_sheet_folder_name(plan.sheet_name)
                    sheet_dir.mkdir(parents=True, exist_ok=True)
                    enum_csv = sheet_dir / f"{safe_sheet_folder_name(plan.sheet_name)}__enumerated.csv"
                    write_plan_csv(plan, enum_csv, resolutions)
                    self._log_excel(f'  Enumeration table: {enum_csv.name}')

                    quant_rows = {}
                    notes = []
                    if run_xic:
                        for rr in resolutions:
                            if rr.raw_path is None:
                                notes.append(f'Replicate{rr.replicate_label}Missing')
                                continue
                            try:
                                tlist = targets_for_raw(
                                    plan,
                                    rr.raw_path.stem,
                                    internal_standard_formula=internal_standard_formula,
                                )
                                qcsv, pdir = quant_xic_for_raw(
                                    rr.raw_path,
                                    tlist,
                                    sheet_dir,
                                    ppm=ppm,
                                    ms_filter="ms",
                                    avg_scans=8,
                                    bin_decimals=4,
                                    min_peak_height=min_peak_height,
                                    use_observed_mz=False,
                                    use_last_as_internal_standard=bool(internal_standard_formula),
                                    internal_standard_adduct=excel_is_adduct,
                                    internal_standard_ppm=excel_is_ppm,
                                    internal_standard_adaptive_detection=excel_is_adaptive,
                                    internal_standard_min_snr_proxy=excel_is_min_snr,
                                    internal_standard_min_height_fraction=excel_is_min_fraction,
                                    internal_standard_expected_rt=(
                                        excel_is_expected_rt
                                        if excel_is_expected_rt is not None
                                        else (
                                            float(__import__("numpy").median(learned_internal_standard_rts))
                                            if learned_internal_standard_rts else None
                                        )
                                    ),
                                    internal_standard_rt_tolerance_min=excel_is_rt_tolerance,
                                    export_ms2_if_dda=bool(self.var_excel_export_ms2.get()),
                                    ms2_min_rel=1.0,
                                    relax_min_peak_height_for_duplicates=False,
                                    dup_min_rel_height=dup_min_rel,
                                    adduct_mode=ADDUCT_MODE_MAP.get(
                                        (self.var_excel_adduct_mode.get() or "").strip(),
                                        "posneg_neghonly",
                                    ),
                                    forced_adduct=(self.var_excel_forced_adduct.get() or ""),
                                    multi_peak_mode=MULTIPEAK_MODE_MAP.get(
                                        (self.var_excel_multipeak_mode.get() or "").strip(),
                                        "auto",
                                    ),
                                    combine_negative_formate=excel_combine_formate,
                                    formate_rt_tolerance_min=excel_formate_rt_tolerance,
                                    formate_min_rel_height_pct=excel_formate_min_rel_pct,
                                    formate_min_shape_correlation=excel_formate_min_corr,
                                    allow_formate_only=excel_allow_formate_only,
                                    include_formate_for_internal_standard=excel_formate_include_is,
                                    negative_channel_mode=excel_neg_channel_mode,
                                    negative_channel_panel_file=excel_neg_panel_file,
                                    negative_channel_rt_tolerance_min=excel_formate_rt_tolerance,
                                    negative_channel_min_rel_height_pct=excel_formate_min_rel_pct,
                                    negative_channel_min_shape_correlation=excel_formate_min_corr,
                                )
                                qrows = read_quant_csv(qcsv)
                                quant_rows[rr.replicate_label] = qrows
                                is_rows_here = [q for q in qrows if str(q.get("Is_internal_standard", "")).strip().lower() in {'1', 'true', 'yes', 'y', 'Yes', '是'}]
                                if is_rows_here:
                                    isr = is_rows_here[-1]
                                    try:
                                        isrt = float(isr.get("Apex_RT", ""))
                                        if str(isr.get("Found", "")).strip().lower() in {'1', 'true', 'yes', 'y', 'Yes', '是'} and __import__("math").isfinite(isrt):
                                            learned_internal_standard_rts.append(isrt)
                                    except Exception:
                                        pass
                                    self._log_excel(
                                        "  IS diagnostic: "
                                        f"status={isr.get('IS_Diagnostic_Status', '')}, found={isr.get('Found', '')}, "
                                        f"adduct={isr.get('Matched_Adduct', isr.get('Adduct', ''))}, "
                                        f"RT={isr.get('Apex_RT', '')}, area={isr.get('Area', '')}, "
                                        f"SNR={isr.get('IS_SNR_Proxy', '')}, plot={isr.get('IS_Diagnostic_PNG', isr.get('XIC_PNG', ''))}"
                                    )
                                raw_ok += 1
                                self._log_excel(
                                    f'  XIC replicate {rr.replicate_label}: OK -> {Path(qcsv).name} | {Path(pdir).name}'
                                )
                            except Exception as e:
                                raw_failed += 1
                                notes.append(f'Replicate{rr.replicate_label}Failed:{e}')
                                self._log_excel(f'  XIC replicate {rr.replicate_label}: FAIL -> {e}')

                    run_records.append(
                        {
                            "plan": plan,
                            "resolutions": resolutions,
                            "quant_rows": quant_rows,
                            "note": "; ".join(notes),
                        }
                    )

                targets_csv = write_all_targets_csv(
                    plans_and_raws,
                    out_root / "00_all_targets_for_xic.csv",
                    internal_standard_formula=internal_standard_formula,
                )
                average_csv = out_root / '00_Triplicate_Average.csv'
                summary_xlsx = write_parallel_summary_xlsx(
                    out_root / '00_Enumeration_Triplicate_XIC_Summary.xlsx',
                    run_records,
                    internal_standard_formula=internal_standard_formula,
                    average_csv_path=average_csv,
                )
                self._log_excel(f'\nCombined targets CSV: {targets_csv}')
                self._log_excel(f'Triplicate summary workbook: {summary_xlsx}')
                self._log_excel(f'Triplicate average CSV: {average_csv}')

                # Discover-all is reviewed directly in this triplicate page. Build a
                # channel summary automatically so a fixed experiment panel can be
                # selected without using a separate discovery tab.
                if run_xic and excel_neg_channel_mode == "discover":
                    try:
                        disc_xlsx, suggested_panel, disc_summary, disc_files = refresh_discovery_summary(
                            out_root,
                            prevalence_threshold_pct=excel_panel_prevalence,
                            min_isotope_support_pct=excel_panel_isotope_support,
                            min_median_area_fraction_pct=excel_panel_min_area_pct,
                        )
                        self._log_excel(f"Discover-all evidence files: {len(disc_files)}")
                        self._log_excel(f"Discovery review workbook: {disc_xlsx}")
                        self._log_excel(f"Suggested panel: {suggested_panel}")
                        selected_panel = out_root / "negative_ion_selected_panel.csv"
                        panel_for_ui = selected_panel if selected_panel.exists() else suggested_panel
                        self.after(0, lambda p=str(panel_for_ui): self.var_excel_neg_panel_file.set(p))
                        self._log_excel(
                            "Next: click 'Review discover-all results / select channels...' in this page. "
                            "After saving, use the no-RAW cached apply button to regenerate the final triplicate summary."
                        )
                    except Exception as disc_exc:
                        self._log_excel(f"WARN: discovery summary could not be built automatically: {disc_exc}")

                self._log_excel(
                    f'Complete: sheets={len(plans)}, XIC_OK={raw_ok}, Missing={raw_missing}, Failed={raw_failed}'
                )
                self.after(
                    0,
                    lambda: messagebox.showinfo(
                        'Complete',
                        f'Enumeration/XIC complete\nSheets={len(plans)}\nXIC_OK={raw_ok}\nMissing={raw_missing}\nFailed={raw_failed}\n\nSummary: {summary_xlsx}\nTriplicate average: {average_csv}',
                    ),
                )
            except Exception as e:
                msg = str(e)
                self._log_excel("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Failed', m))
            finally:
                self.after(0, lambda: self.btn_run_excel.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_excel.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()


    # =====================
    # Patch existing Excel summary workbook without re-running XIC
    # =====================
    def _build_summary_patch_tab(self):
        frm = self.tab_patch
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)
        r = 0

        ttk.Label(frm, text='Existing summary workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_patch_summary_xlsx = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_patch_summary_xlsx).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_patch_summary_xlsx).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Component workbook (name / formula / mass):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_patch_source_workbook = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_patch_source_workbook).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_patch_source_workbook).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Output workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_patch_out_xlsx = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_patch_out_xlsx).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Save as...', command=self._browse_patch_out_xlsx).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        box = ttk.LabelFrame(frm, text='Update options')
        box.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=5)
        self.var_patch_combo = tk.BooleanVar(value=True)
        self.var_patch_english = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            box,
            text='Update Combo labels from component names in the source workbook',
            variable=self.var_patch_combo,
        ).grid(row=0, column=0, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(
            box,
            text='Translate legacy sheet names, headers and standard notes to English',
            variable=self.var_patch_english,
        ).grid(row=1, column=0, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Label(
            box,
            text='This tool updates an existing summary without reading RAW files or repeating integration. Updating Combo labels requires the original component workbook.',
        ).grid(row=2, column=0, columnspan=3, sticky="w", padx=8, pady=4)
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=8)
        self.btn_run_patch = ttk.Button(btns, text='Update existing summary', command=self._on_run_summary_patch)
        self.btn_run_patch.pack(side="left")
        ttk.Button(btns, text='Open output directory', command=self._open_patch_out_dir).pack(side="left", padx=10)
        self.lbl_status_patch = ttk.Label(btns, text="Ready")
        self.lbl_status_patch.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=5)
        self.txt_log_patch = tk.Text(frm, height=15, wrap="word")
        self.txt_log_patch.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=5)
        frm.rowconfigure(r, weight=1)

    def _browse_patch_summary_xlsx(self):
        path = filedialog.askopenfilename(
            title='Select summary workbook',
            filetypes=[('Component workbook', "*.xlsx *.xlsm"), ("All", "*")],
        )
        if path:
            self.var_patch_summary_xlsx.set(path)
            if not self.var_patch_out_xlsx.get().strip():
                p = Path(path)
                self.var_patch_out_xlsx.set(str(p.with_name(f"{p.stem}__combo_english.xlsx")))

    def _browse_patch_source_workbook(self):
        path = filedialog.askopenfilename(
            title='Select component workbook (name / formula / mass)',
            filetypes=[('Component workbook', "*.xlsx *.xlsm"), ("All", "*")],
        )
        if path:
            self.var_patch_source_workbook.set(path)

    def _browse_patch_out_xlsx(self):
        path = filedialog.asksaveasfilename(
            title='Save updated workbook',
            defaultextension=".xlsx",
            filetypes=[('Component workbook', "*.xlsx"), ("All", "*")],
        )
        if path:
            self.var_patch_out_xlsx.set(path)

    def _open_patch_out_dir(self):
        p = self.var_patch_out_xlsx.get().strip()
        if p:
            open_in_os(Path(p).parent)

    def _on_run_summary_patch(self):
        summary = self.var_patch_summary_xlsx.get().strip()
        source = self.var_patch_source_workbook.get().strip()
        out_xlsx = self.var_patch_out_xlsx.get().strip()
        patch_combo = bool(self.var_patch_combo.get())
        patch_english = bool(self.var_patch_english.get())

        if not summary:
            messagebox.showerror('Error', 'Select an existing summary workbook.')
            return
        if patch_combo and not source:
            messagebox.showerror('Error', 'Select the component workbook to update Combo labels.')
            return
        if not out_xlsx:
            p = Path(summary)
            out_xlsx = str(p.with_name(f"{p.stem}__combo_english.xlsx"))
            self.var_patch_out_xlsx.set(out_xlsx)

        self.btn_run_patch.configure(state="disabled")
        self.lbl_status_patch.configure(text="Running...")
        self.txt_log_patch.delete("1.0", "end")

        def worker():
            try:
                self._log_patch(f'Summary workbook: {summary}')
                if patch_combo:
                    self._log_patch(f'Component workbook: {source}')
                self._log_patch(f'Output file: {out_xlsx}')
                result = patch_summary_workbook(
                    Path(summary),
                    Path(source) if source else Path(summary),
                    Path(out_xlsx),
                    patch_combo=patch_combo,
                    translate_to_english=patch_english,
                )
                for w in result.warnings:
                    self._log_patch("WARN: " + str(w))
                self._log_patch(f'Component labels loaded: {result.reagent_labels_loaded}')
                self._log_patch(f'Combo cells: {result.combo_cells_seen}; updated: {result.combo_cells_changed}')
                self._log_patch(f'Sheet names translated: {result.sheet_names_changed}')
                self._log_patch(f'Headers translated: {result.header_cells_translated}')
                self._log_patch(f'Values translated: {result.value_cells_translated}')
                self._log_patch('Complete.')
                self.after(0, lambda: messagebox.showinfo('Complete', f'Generated:\n{result.output_path}'))
            except Exception as e:
                msg = str(e)
                self._log_patch("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Failed', m))
            finally:
                self.after(0, lambda: self.btn_run_patch.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_patch.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    # =====================
    # CSV formula filter tab
    # =====================
    def _build_filter_tab(self):
        frm = self.tab_filter
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)

        r = 0

        ttk.Label(frm, text='Reference CSV (formulas to remove):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_filter_ref_csv = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_filter_ref_csv).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_filter_ref_csv).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Input CSV:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_filter_target_csv = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_filter_target_csv).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_filter_target_csv).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Output CSV (retained rows):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_filter_out_csv = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_filter_out_csv).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_filter_out_csv).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        # removed
        self.var_filter_write_removed = tk.BooleanVar(value=True)
        self.var_filter_removed_csv = tk.StringVar()

        rem_fr = ttk.Frame(frm)
        rem_fr.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=(0, 6))
        rem_fr.columnconfigure(1, weight=1)

        ttk.Checkbutton(
            rem_fr,
            text='Also save removed rows:',
            variable=self.var_filter_write_removed,
            command=self._update_filter_removed_ui,
        ).grid(row=0, column=0, sticky="w")
        self.ent_filter_removed = ttk.Entry(rem_fr, textvariable=self.var_filter_removed_csv)
        self.ent_filter_removed.grid(row=0, column=1, sticky="ew", padx=10)
        self.btn_filter_removed = ttk.Button(rem_fr, text='Browse...', command=self._browse_filter_removed_csv)
        self.btn_filter_removed.grid(row=0, column=2)
        r += 1

        ttk.Separator(frm).grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)
        r += 1

        opt = ttk.LabelFrame(frm, text='Matching settings')
        opt.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(6):
            opt.columnconfigure(c, weight=1, uniform="options")

        # mode
        self.var_filter_mode_label = tk.StringVar(value='Match normalized molecular formulas')
        self.var_filter_mode_code = tk.StringVar(value="canonical")

        ttk.Label(opt, text='Matching mode:').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        self.cmb_filter_mode = ttk.Combobox(
            opt,
            textvariable=self.var_filter_mode_label,
            values=['Match normalized molecular formulas', 'Exact text match'],
            state="readonly",
            width=24,
        )
        self.cmb_filter_mode.grid(row=0, column=1, sticky="w", padx=8, pady=4)
        self.cmb_filter_mode.bind("<<ComboboxSelected>>", lambda e: self._sync_filter_mode())

        # columns
        ttk.Label(opt, text='Reference formula column:').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        self.var_filter_ref_col = tk.StringVar(value="formula")
        self.cmb_filter_ref_col = ttk.Combobox(opt, textvariable=self.var_filter_ref_col, values=["formula"], state="readonly")
        self.cmb_filter_ref_col.grid(row=1, column=1, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Input formula column:').grid(row=1, column=2, sticky="w", padx=8, pady=4)
        self.var_filter_target_col = tk.StringVar(value="formula")
        self.cmb_filter_target_col = ttk.Combobox(opt, textvariable=self.var_filter_target_col, values=["formula"], state="readonly")
        self.cmb_filter_target_col.grid(row=1, column=3, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Normalized matching parses each molecular formula and uses Hill notation.').grid(
            row=2, column=0, columnspan=6, sticky="w", padx=8, pady=(6, 4)
        )
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)
        self.btn_run_filter = ttk.Button(btns, text='Filter CSV', command=self._on_run_filter)
        self.btn_run_filter.pack(side="left")
        self.lbl_status_filter = ttk.Label(btns, text="Ready")
        self.lbl_status_filter.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_log_filter = tk.Text(frm, height=18, wrap="word")
        self.txt_log_filter.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=6)
        frm.rowconfigure(r, weight=1)

        self._sync_filter_mode()
        self._update_filter_removed_ui()

    def _sync_filter_mode(self):
        label = (self.var_filter_mode_label.get() or "").strip()
        code = "canonical" if 'normalized' in label.casefold() else "exact"
        self.var_filter_mode_code.set(code)

    def _update_filter_removed_ui(self):
        try:
            enabled = bool(self.var_filter_write_removed.get())
            self.ent_filter_removed.configure(state="normal" if enabled else "disabled")
            self.btn_filter_removed.configure(state="normal" if enabled else "disabled")
        except Exception:
            pass

    def _browse_filter_ref_csv(self):
        path = filedialog.askopenfilename(title='Select reference CSV', filetypes=[("CSV", "*.csv"), ("All", "*")])
        if not path:
            return
        self.var_filter_ref_csv.set(path)
        self._refresh_filter_columns(which="ref")

    def _browse_filter_target_csv(self):
        path = filedialog.askopenfilename(title='Select input CSV', filetypes=[("CSV", "*.csv"), ("All", "*")])
        if not path:
            return
        self.var_filter_target_csv.set(path)
        p = Path(path)
        if not self.var_filter_out_csv.get().strip():
            self.var_filter_out_csv.set(str(p.with_name(f"{p.stem}__filtered.csv")))
        if not self.var_filter_removed_csv.get().strip():
            self.var_filter_removed_csv.set(str(p.with_name(f"{p.stem}__removed.csv")))
        self._refresh_filter_columns(which="target")

    def _browse_filter_out_csv(self):
        path = filedialog.asksaveasfilename(
            title='Save retained rows',
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("All", "*")],
        )
        if path:
            self.var_filter_out_csv.set(path)

    def _browse_filter_removed_csv(self):
        path = filedialog.asksaveasfilename(
            title='Save removed rows',
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv"), ("All", "*")],
        )
        if path:
            self.var_filter_removed_csv.set(path)

    def _refresh_filter_columns(self, which: str):
        'Read CSV headers and update formula-column choices.'
        try:
            if which == "ref":
                p = Path(self.var_filter_ref_csv.get().strip())
                if not p.exists():
                    return
                header = read_csv_header(p)
                if header:
                    self.cmb_filter_ref_col.configure(values=header)
                    guess = guess_formula_column(header)
                    if guess:
                        self.var_filter_ref_col.set(guess)
                    else:
                        self.var_filter_ref_col.set(header[0])
            else:
                p = Path(self.var_filter_target_csv.get().strip())
                if not p.exists():
                    return
                header = read_csv_header(p)
                if header:
                    self.cmb_filter_target_col.configure(values=header)
                    guess = guess_formula_column(header)
                    if guess:
                        self.var_filter_target_col.set(guess)
                    else:
                        self.var_filter_target_col.set(header[0])
        except Exception as e:
            self._log_filter(f'WARN: Could not read headers: {e}')

    def _on_run_filter(self):
        ref_csv = self.var_filter_ref_csv.get().strip()
        target_csv = self.var_filter_target_csv.get().strip()
        out_csv = self.var_filter_out_csv.get().strip()

        if not ref_csv or not target_csv:
            messagebox.showerror('Error', 'Select the reference and input CSV files.')
            return
        if not out_csv:
            messagebox.showerror('Error', 'Select an output CSV file.')
            return

        ref_col = (self.var_filter_ref_col.get() or "").strip()
        tgt_col = (self.var_filter_target_col.get() or "").strip()
        if not ref_col or not tgt_col:
            messagebox.showerror('Error', 'Select the formula column in each CSV.')
            return

        mode = (self.var_filter_mode_code.get() or "canonical").strip()

        removed_csv = None
        if bool(self.var_filter_write_removed.get()):
            rc = self.var_filter_removed_csv.get().strip()
            if rc:
                removed_csv = Path(rc)

        self.btn_run_filter.config(state="disabled")
        self.lbl_status_filter.config(text="Running...")
        self.txt_log_filter.delete("1.0", "end")

        def worker():
            try:
                self._log_filter(f"Reference: {ref_csv}")
                self._log_filter(f"Target:    {target_csv}")
                self._log_filter(f"Mode:      {mode}")
                self._log_filter(f"Ref col:   {ref_col} | Target col: {tgt_col}")
                self._log_filter(f"Output:    {out_csv}")
                if removed_csv is not None:
                    self._log_filter(f"Removed:   {removed_csv}")

                stats = filter_csv_remove_matching_formulas(
                    ref_csv=Path(ref_csv),
                    target_csv=Path(target_csv),
                    out_csv=Path(out_csv),
                    removed_csv=removed_csv,
                    ref_formula_col=ref_col,
                    target_formula_col=tgt_col,
                    mode=mode,
                )

                self._log_filter(
                    f"Done. TargetRows={stats.total_target_rows}, Kept={stats.kept_rows}, Removed={stats.removed_rows}"
                )
                self._log_filter(
                    f"RefFormulas={stats.ref_unique_formulas} (rows with formula={stats.ref_rows_with_formula}, parseErr={stats.ref_parse_errors})"
                )
                if mode != "exact":
                    self._log_filter(
                        f"Target blank formula rows={stats.target_blank_formula}, parseErr={stats.target_parse_errors}"
                    )

                self.after(
                    0,
                    lambda: messagebox.showinfo(
                        'Complete',
                        f'Filtering complete\nInput rows: {stats.total_target_rows}\nRetained: {stats.kept_rows}\nRemoved: {stats.removed_rows}\nOutput: {out_csv}',
                    ),
                )
            except Exception as e:
                self._log_filter(f"ERROR: {e}")
                self.after(0, lambda: messagebox.showerror('Failed', str(e)))
            finally:
                self.after(0, lambda: self.btn_run_filter.config(state="normal"))
                self.after(0, lambda: self.lbl_status_filter.config(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    # =====================
    # Response correction / concentration prediction tab
    # =====================
    def _build_response_prediction_tab(self):
        frm = self.tab_resp
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)
        r = 0

        ttk.Label(frm, text='Calibration table (up to 100 known standards):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_resp_cal_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_resp_cal_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_resp_cal_file).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Target table (Combo and area ratio):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_resp_target_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_resp_target_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_resp_target_file).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Output workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_resp_out_xlsx = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_resp_out_xlsx).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Save as...', command=self._browse_resp_out_xlsx).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        files_note = ttk.Label(
            frm,
            text='Calibration and target files may be CSV or XLSX. Use the same ratio scale in both files (for example, Area/IS x 100%). Combo identifies component combinations for response-factor transfer.',
        )
        files_note.grid(row=r, column=0, columnspan=3, sticky="w", padx=10, pady=(0, 6))
        r += 1

        box = ttk.LabelFrame(frm, text='Sheets and columns (blank = automatic)')
        box.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(6):
            box.columnconfigure(c, weight=1)

        self.var_resp_cal_sheet = tk.StringVar(value="")
        self.var_resp_target_sheet = tk.StringVar(value="")
        self.var_resp_combo_col = tk.StringVar(value="")
        self.var_resp_formula_col = tk.StringVar(value="")
        self.var_resp_ratio_col = tk.StringVar(value="")
        self.var_resp_target_ratio_col = tk.StringVar(value="")
        self.var_resp_conc_col = tk.StringVar(value="")

        ttk.Label(box, text='Calibration sheet').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_cal_sheet, width=18).grid(row=0, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Target sheet').grid(row=0, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_target_sheet, width=18).grid(row=0, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Blank = automatic; prefer triplicate averages').grid(row=0, column=4, columnspan=2, sticky="w", padx=8, pady=4)

        ttk.Label(box, text='Combo column').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_combo_col, width=18).grid(row=1, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Formula column').grid(row=1, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_formula_col, width=18).grid(row=1, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Known concentration column').grid(row=1, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_conc_col, width=18).grid(row=1, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(box, text='Calibration ratio column').grid(row=2, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_ratio_col, width=18).grid(row=2, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Target ratio column').grid(row=2, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_resp_target_ratio_col, width=18).grid(row=2, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Leave target ratio blank when both tables use the same column name.').grid(row=2, column=4, columnspan=2, sticky="w", padx=8, pady=4)
        r += 1

        opt = ttk.LabelFrame(frm, text='Model settings')
        opt.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(6):
            opt.columnconfigure(c, weight=1, uniform="options")

        self.var_resp_model_label = tk.StringVar(value='Hybrid (KNN + ridge)')
        self.var_resp_k = tk.IntVar(value=15)
        self.var_resp_lambda = tk.DoubleVar(value=1.0)
        self.var_resp_output_csv = tk.BooleanVar(value=True)

        ttk.Label(opt, text='Prediction model').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(
            opt,
            textvariable=self.var_resp_model_label,
            values=list(MODEL_LABEL_TO_CODE.keys()),
            state="readonly",
            width=32,
        ).grid(row=0, column=1, columnspan=2, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Neighbors (K)').grid(row=0, column=3, sticky="w", padx=8, pady=4)
        ttk.Spinbox(opt, from_=1, to=100, textvariable=self.var_resp_k, width=8).grid(row=0, column=4, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(opt, text='Also export predictions as CSV', variable=self.var_resp_output_csv).grid(row=0, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(opt, text='Ridge penalty').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_resp_lambda, width=10).grid(row=1, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(
            opt,
            text='This transfer model uses an inverse response factor: known concentration / measured ratio. Predicted concentration = target ratio x predicted inverse response factor. LOOCV results are reported separately.',
        ).grid(row=1, column=2, columnspan=4, sticky="w", padx=8, pady=4)
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)
        self.btn_run_resp = ttk.Button(btns, text='Fit and predict', command=self._on_run_response_prediction)
        self.btn_run_resp.pack(side="left")
        ttk.Button(btns, text='Open output directory', command=self._open_resp_out_dir).pack(side="left", padx=10)
        self.lbl_status_resp = ttk.Label(btns, text="Ready")
        self.lbl_status_resp.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_log_resp = tk.Text(frm, height=17, wrap="word")
        self.txt_log_resp.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=6)
        frm.rowconfigure(r, weight=1)

    def _browse_resp_cal_file(self):
        path = filedialog.askopenfilename(
            title='Select calibration table (known concentration and area ratio)',
            filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")],
        )
        if path:
            self.var_resp_cal_file.set(path)
            if not self.var_resp_out_xlsx.get().strip():
                p = Path(path)
                self.var_resp_out_xlsx.set(str(p.with_name("response_concentration_prediction.xlsx")))

    def _browse_resp_target_file(self):
        path = filedialog.askopenfilename(
            title='Select target table',
            filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")],
        )
        if path:
            self.var_resp_target_file.set(path)
            if not self.var_resp_out_xlsx.get().strip():
                p = Path(path)
                self.var_resp_out_xlsx.set(str(p.with_name(f"{p.stem}__predicted_concentration.xlsx")))

    def _browse_resp_out_xlsx(self):
        path = filedialog.asksaveasfilename(
            title='Save concentration predictions',
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx"), ("All", "*")],
        )
        if path:
            self.var_resp_out_xlsx.set(path)

    def _open_resp_out_dir(self):
        p = self.var_resp_out_xlsx.get().strip()
        if p:
            open_in_os(Path(p).parent)

    def _on_run_response_prediction(self):
        cal_file = self.var_resp_cal_file.get().strip()
        target_file = self.var_resp_target_file.get().strip()
        out_xlsx = self.var_resp_out_xlsx.get().strip()
        if not cal_file or not target_file:
            messagebox.showerror('Error', 'Select calibration and target tables.')
            return
        if not out_xlsx:
            messagebox.showerror('Error', 'Select an output Excel file.')
            return

        try:
            k = int(self.var_resp_k.get())
            lam = float(self.var_resp_lambda.get())
            if k <= 0:
                raise ValueError('The number of neighbors must be greater than zero.')
            if lam < 0:
                raise ValueError('The ridge penalty cannot be negative.')
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return

        self.btn_run_resp.configure(state="disabled")
        self.lbl_status_resp.configure(text="Running...")
        self.txt_log_resp.delete("1.0", "end")

        def worker():
            try:
                self._log_resp(f"Calibration: {cal_file}")
                self._log_resp(f"Targets:     {target_file}")
                self._log_resp(f"Output:      {out_xlsx}")
                self._log_resp(f"Model:       {self.var_resp_model_label.get()} | K={k} | lambda={lam}")
                report = run_response_prediction(
                    Path(cal_file),
                    Path(target_file),
                    Path(out_xlsx),
                    calibration_sheet=self.var_resp_cal_sheet.get().strip(),
                    target_sheet=self.var_resp_target_sheet.get().strip(),
                    combo_col=self.var_resp_combo_col.get().strip(),
                    formula_col=self.var_resp_formula_col.get().strip(),
                    ratio_col=self.var_resp_ratio_col.get().strip(),
                    concentration_col=self.var_resp_conc_col.get().strip(),
                    target_ratio_col=self.var_resp_target_ratio_col.get().strip(),
                    model_mode=self.var_resp_model_label.get().strip(),
                    k_neighbors=k,
                    ridge_lambda=lam,
                    output_csv=bool(self.var_resp_output_csv.get()),
                )
                for w in report.warnings[:30]:
                    self._log_resp("WARN: " + str(w))
                if len(report.warnings) > 30:
                    self._log_resp(f'WARN: {len(report.warnings) - 30} additional warnings are recorded in the output.')
                self._log_resp(f"Training rows: total={report.n_train_total}, used={report.n_train_used}")
                self._log_resp(f"Target rows:   total={report.n_target_total}, predicted={report.n_predicted}")
                if report.qc:
                    self._log_resp("LOOCV QC:")
                    for key in ("N", "MAPE_%", "Median_APE_%", "P90_APE_%", "R2_log", "Within_2x_%", "Within_5x_%"):
                        if key in report.qc:
                            self._log_resp(f"  {key}: {report.qc[key]}")
                    if "Model_rating" in report.qc:
                        self._log_resp(f"Model rating: {report.qc.get('Model_rating')}")
                    if "Model_conclusion" in report.qc:
                        self._log_resp(f"Conclusion: {report.qc.get('Model_conclusion')}")
                    if "Recommendation" in report.qc:
                        self._log_resp(f"Recommendation: {report.qc.get('Recommendation')}")
                self._log_resp(f"Excel: {report.output_xlsx}")
                if report.output_csv:
                    self._log_resp(f"CSV:   {report.output_csv}")
                self.after(0, lambda: messagebox.showinfo('Complete', f'Prediction complete\nCalibration rows: {report.n_train_used}\nTarget predictions: {report.n_predicted}\n\nOutput: {report.output_xlsx}'))
            except Exception as e:
                msg = str(e)
                self._log_resp("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Failed', m))
            finally:
                self.after(0, lambda: self.btn_run_resp.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_resp.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()



    # =====================
    # Simplified A/B/C + fixed CORE SMILES master builder
    # =====================
    def _build_smiles_builder_tab(self):
        from core.ui_support import scroll_body
        frm = scroll_body(self.tab_smiles)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)
        r = 0

        ttk.Label(frm, text='Component workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_smiles_combo_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_smiles_combo_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        bf = ttk.Frame(frm)
        bf.grid(row=r, column=2, sticky="e", padx=10, pady=6)
        ttk.Button(bf, text='Browse...', command=self._browse_smiles_combo_file).pack(side="left")
        ttk.Button(bf, text='Export template', command=self._export_smiles_builder_templates).pack(side="left", padx=(6, 0))
        r += 1

        ttk.Label(frm, text='Mapped core SMILES:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_smiles_core = tk.Text(frm, height=3, wrap="word")
        self.txt_smiles_core.grid(row=r, column=1, columnspan=2, sticky="ew", padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Product SMILES output:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_smiles_out_xlsx = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_smiles_out_xlsx).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Save as...', command=self._browse_smiles_out_xlsx).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(
            frm,
            text=(
                'Columns A:D, E:H and I:L contain the name, formula, mass and SMILES of components A, B and C. Enter the common core once. Components and core use paired mapped dummy atoms, for example [*:1], [*:2] and [*:3]. The product table retains every successful combination. Downstream matching uses the ordered A/B/C formula key and checks the product formula; local component numbers are not treated as global identifiers.'
            ),
            wraplength=900,
        ).grid(row=r, column=0, columnspan=3, sticky="w", padx=10, pady=(0, 6))
        r += 1

        cfg = ttk.LabelFrame(frm, text='Input and output')
        cfg.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(8):
            cfg.columnconfigure(c, weight=1)
        self.var_smiles_combo_sheet = tk.StringVar(value="")
        self.var_smiles_data_start = tk.IntVar(value=2)
        self.var_smiles_only_formula = tk.BooleanVar(value=False)
        self.var_smiles_png = tk.BooleanVar(value=True)
        self.var_smiles_embed = tk.BooleanVar(value=True)
        self.var_smiles_include_inputs = tk.BooleanVar(value=False)
        self.var_smiles_output_csv = tk.BooleanVar(value=True)
        self.var_smiles_image_limit = tk.IntVar(value=100)
        self.var_smiles_embed_limit = tk.IntVar(value=30)

        ttk.Label(cfg, text='Sheet (blank = all visible sheets)').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(cfg, textvariable=self.var_smiles_combo_sheet, width=18).grid(row=0, column=1, sticky="ew", padx=8, pady=4)
        ttk.Label(cfg, text='First data row').grid(row=0, column=2, sticky="w", padx=8, pady=4)
        ttk.Spinbox(cfg, from_=2, to=1000, textvariable=self.var_smiles_data_start, width=8).grid(row=0, column=3, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(cfg, text='Keep C/H/N/O/P only', variable=self.var_smiles_only_formula).grid(row=0, column=4, columnspan=2, sticky="w", padx=8, pady=4)
        ttk.Label(cfg, text='All elements retained by default').grid(row=0, column=6, columnspan=2, sticky="w", padx=8, pady=4)

        ttk.Checkbutton(cfg, text='Generate structure PNGs', variable=self.var_smiles_png).grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(cfg, text='Embed previews in Excel', variable=self.var_smiles_embed).grid(row=1, column=1, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(cfg, text='Include component SMILES columns', variable=self.var_smiles_include_inputs).grid(row=1, column=2, columnspan=2, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(cfg, text='Also export CSV', variable=self.var_smiles_output_csv).grid(row=1, column=4, sticky="w", padx=8, pady=4)
        ttk.Label(cfg, text='Maximum PNGs').grid(row=1, column=5, sticky="e", padx=4, pady=4)
        ttk.Spinbox(cfg, from_=0, to=100000, textvariable=self.var_smiles_image_limit, width=8).grid(row=1, column=6, sticky="w", padx=4, pady=4)
        ttk.Label(cfg, text='0 = all').grid(row=1, column=7, sticky="w", padx=4, pady=4)

        ttk.Label(cfg, text='Maximum embedded images').grid(row=2, column=0, sticky="e", padx=4, pady=4)
        ttk.Spinbox(cfg, from_=0, to=10000, textvariable=self.var_smiles_embed_limit, width=8).grid(row=2, column=1, sticky="w", padx=4, pady=4)
        ttk.Label(
            cfg,
            text='Defaults: export up to 100 structure PNGs and embed up to 30 previews. Product SMILES are still exported for all successful combinations.',
            wraplength=700,
        ).grid(row=2, column=2, columnspan=6, sticky="w", padx=8, pady=4)
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=8)
        self.btn_diagnose_smiles = ttk.Button(btns, text='Check structures and attachment points', command=self._on_diagnose_smiles)
        self.btn_diagnose_smiles.pack(side="left")
        self.btn_preview_smiles = ttk.Button(btns, text='Preview first product', command=self._on_preview_smiles)
        self.btn_preview_smiles.pack(side="left", padx=(10, 0))
        self.btn_run_smiles = ttk.Button(btns, text='Build product SMILES table', command=self._on_run_smiles_builder)
        self.btn_run_smiles.pack(side="left", padx=(10, 0))
        ttk.Button(btns, text='Open output directory', command=self._open_smiles_out_dir).pack(side="left", padx=10)
        self.lbl_status_smiles = ttk.Label(btns, text="Ready")
        self.lbl_status_smiles.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_log_smiles = tk.Text(frm, height=12, wrap="word")
        self.txt_log_smiles.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=6)
        frm.rowconfigure(r, weight=1)

    def _browse_smiles_combo_file(self):
        path = filedialog.askopenfilename(
            title='Select component workbook (four columns per component)',
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*")],
        )
        if path:
            self.var_smiles_combo_file.set(path)
            if not self.var_smiles_out_xlsx.get().strip():
                self.var_smiles_out_xlsx.set(str(Path(path).with_name(f"{Path(path).stem}__1000_Combo_SMILES_Master.xlsx")))

    def _browse_smiles_out_xlsx(self):
        path = filedialog.asksaveasfilename(
            title='Save product SMILES table',
            defaultextension=".xlsx",
            initialfile="1000_Combo_SMILES_Master.xlsx",
            filetypes=[("Excel", "*.xlsx"), ("All", "*")],
        )
        if path:
            self.var_smiles_out_xlsx.set(path)

    def _export_smiles_builder_templates(self):
        src = Path(__file__).resolve().parent / "templates" / "ABC_components_with_SMILES_template.xlsx"
        if not src.exists():
            messagebox.showerror('Error', f'Template not found: {src}')
            return
        dst = filedialog.asksaveasfilename(
            title='Export component template',
            defaultextension=".xlsx",
            initialfile="ABC_components_with_SMILES_template.xlsx",
            filetypes=[("Excel", "*.xlsx"), ("All", "*")],
        )
        if not dst:
            return
        try:
            shutil.copy2(str(src), str(dst))
            self.var_smiles_combo_file.set(str(dst))
            messagebox.showinfo('Complete', f'Template exported:\n{dst}')
        except Exception as e:
            messagebox.showerror('Failed', str(e))

    def _open_smiles_out_dir(self):
        p = self.var_smiles_out_xlsx.get().strip()
        if p:
            open_in_os(Path(p).parent)

    def _smiles_builder_args(self):
        component_file = self.var_smiles_combo_file.get().strip()
        out_xlsx = self.var_smiles_out_xlsx.get().strip()
        core_smiles = self.txt_smiles_core.get("1.0", "end").strip()
        if not component_file:
            raise ValueError('Select a component workbook.')
        if not core_smiles:
            raise ValueError('Enter the mapped SMILES of the common core.')
        if not out_xlsx:
            raise ValueError('Select an output Excel file.')
        data_start = int(self.var_smiles_data_start.get())
        image_limit = int(self.var_smiles_image_limit.get())
        embed_limit = int(self.var_smiles_embed_limit.get())
        if data_start < 2 or image_limit < 0 or embed_limit < 0:
            raise ValueError('The first data row must be at least 2. Image counts cannot be negative.')
        return {
            "component_workbook": Path(component_file),
            "output_xlsx": Path(out_xlsx),
            "core_smiles": core_smiles,
            "sheet_name": self.var_smiles_combo_sheet.get().strip(),
            "data_start_row": data_start,
            "only_formula": bool(self.var_smiles_only_formula.get()),
            "generate_png": bool(self.var_smiles_png.get()),
            "image_limit": image_limit,
            "embed_images": bool(self.var_smiles_embed.get()),
            "embed_image_limit": embed_limit,
            "include_fragment_smiles": bool(self.var_smiles_include_inputs.get()),
            "output_csv": bool(self.var_smiles_output_csv.get()),
        }

    def _on_diagnose_smiles(self):
        try:
            args = self._smiles_builder_args()
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return
        self.btn_diagnose_smiles.configure(state="disabled")
        self.btn_preview_smiles.configure(state="disabled")
        self.btn_run_smiles.configure(state="disabled")
        self.lbl_status_smiles.configure(text="Diagnosing...")
        self._log_smiles('Checking the core, component SMILES, attachment points and valences locally...')

        def worker():
            try:
                base = args["output_xlsx"]
                dx = base.with_name(base.stem + "__SMILES_Input_Diagnostics.xlsx")
                dt = base.with_name(base.stem + "__SMILES_Input_Diagnostics.txt")
                report = diagnose_simple_combo_inputs(
                    args["component_workbook"], core_smiles=args["core_smiles"],
                    sheet_name=args["sheet_name"], data_start_row=args["data_start_row"],
                    output_xlsx=dx, output_txt=dt, include_raw_smiles=True,
                )
                text = diagnostics_text(report, include_raw=False)
                self._log_smiles(
                    f"SMILES diagnostics: {report.overall_status} | checked={len(report.diagnostics)}, "
                    f"errors={report.error_count}, warnings={report.warning_count}"
                )
                for d in report.diagnostics:
                    if d.fatal:
                        self._log_smiles("ERROR: " + d.one_line())
                        for fix in d.suggestions[:3]:
                            self._log_smiles("  FIX: " + fix)
                for warning in report.global_warnings[:30]:
                    self._log_smiles("AUDIT: " + str(warning))
                self._log_smiles(f"Diagnostic Excel: {dx}")
                self._log_smiles(f"Diagnostic TXT: {dt}")
                self.after(0, lambda r=report, t=text, x=dx, q=dt: self._show_smiles_diagnostics(r, t, x, q))
            except Exception as e:
                msg = str(e)
                self._log_smiles("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('SMILES check failed', m))
            finally:
                self.after(0, lambda: self.btn_diagnose_smiles.configure(state="normal"))
                self.after(0, lambda: self.btn_preview_smiles.configure(state="normal"))
                self.after(0, lambda: self.btn_run_smiles.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_smiles.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    def _show_smiles_diagnostics(self, report, text: str, xlsx_path: Path, txt_path: Path):
        top = tk.Toplevel(self)
        top.title(f'SMILES diagnostics - {report.overall_status}')
        top.geometry("1080x760")
        top.columnconfigure(0, weight=1)
        top.rowconfigure(1, weight=1)
        ttk.Label(
            top,
            text=(f"Status: {report.overall_status}    Checked: {len(report.diagnostics)}    "
                  f"Errors: {report.error_count}    Warnings: {report.warning_count}"),
            font=("TkDefaultFont", 10, "bold"),
        ).grid(row=0, column=0, sticky="w", padx=10, pady=10)
        frame = ttk.Frame(top)
        frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 10))
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)
        box = tk.Text(frame, wrap="word")
        bar = ttk.Scrollbar(frame, orient="vertical", command=box.yview)
        box.configure(yscrollcommand=bar.set)
        box.grid(row=0, column=0, sticky="nsew")
        bar.grid(row=0, column=1, sticky="ns")
        box.insert("1.0", text)
        box.configure(state="disabled")
        bf = ttk.Frame(top)
        bf.grid(row=2, column=0, sticky="ew", padx=10, pady=(0, 10))
        ttk.Button(bf, text='Open diagnostic workbook', command=lambda: open_in_os(xlsx_path)).pack(side="left")
        ttk.Button(bf, text='Open diagnostic text', command=lambda: open_in_os(txt_path)).pack(side="left", padx=8)
        ttk.Button(bf, text='Close', command=top.destroy).pack(side="right")

    def _on_preview_smiles(self):
        try:
            args = self._smiles_builder_args()
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return
        self.btn_diagnose_smiles.configure(state="disabled")
        self.btn_preview_smiles.configure(state="disabled")
        self.btn_run_smiles.configure(state="disabled")
        self.lbl_status_smiles.configure(text="Previewing...")
        self._log_smiles("Previewing first A\u00D7B\u00D7C combination...")

        def worker():
            try:
                preview_dir = args["output_xlsx"].parent / "_simple_combo_smiles_preview"
                result = preview_first_simple_combo(
                    args["component_workbook"],
                    core_smiles=args["core_smiles"],
                    sheet_name=args["sheet_name"],
                    data_start_row=args["data_start_row"],
                    only_formula=args["only_formula"],
                    preview_dir=preview_dir,
                )
                self._log_smiles(f"Preview status: {result.status}")
                self._log_smiles(f"Combo: {result.combo}")
                self._log_smiles(f"Product: {result.product_smiles or '-'}")
                self._log_smiles(f"Expected formula: {result.expected_formula} | Product formula: {result.product_formula or '-'} | match={result.formula_match or '-'}")
                for w in result.warnings:
                    self._log_smiles("WARN: " + str(w))
                self.after(0, lambda rr=result: self._show_smiles_preview(rr))
            except Exception as e:
                msg = str(e)
                self._log_smiles("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Preview failed', m))
            finally:
                self.after(0, lambda: self.btn_diagnose_smiles.configure(state="normal"))
                self.after(0, lambda: self.btn_preview_smiles.configure(state="normal"))
                self.after(0, lambda: self.btn_run_smiles.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_smiles.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    def _show_smiles_preview(self, result):
        top = tk.Toplevel(self)
        top.title('Product SMILES preview')
        top.geometry("980x720")
        top.columnconfigure(0, weight=1)
        top.rowconfigure(1, weight=1)

        info = tk.Text(top, height=20, wrap="word")
        info.grid(row=0, column=0, sticky="ew", padx=10, pady=10)
        lines = [
            f"Status: {result.status}",
            f"Source Sheet: {result.source_sheet}",
            f"Combo: {result.combo}",
            f"A/B/C: {result.a.item_id} | {result.b.item_id} | {result.c.item_id}",
            f"Expected Formula: {result.expected_formula}",
            f"Expected Exact Mass: {result.expected_exact_mass:.6f}" if result.expected_exact_mass is not None else "Expected Exact Mass: -",
            f"Product SMILES: {result.product_smiles or '-'}",
            f"Product Formula: {result.product_formula or '-'}",
            f"Formula Match: {result.formula_match or '-'}",
            f"Product Exact Mass: {result.product_exact_mass:.6f}" if result.product_exact_mass is not None else "Product Exact Mass: -",
            f"InChIKey: {result.inchikey or '-'}",
        ]
        if result.warnings:
            lines.append("Warnings / exact fixes:")
            lines.extend("  - " + str(w) for w in result.warnings)
        info.insert("1.0", "\n".join(lines))
        info.configure(state="disabled")

        frame = ttk.Frame(top)
        frame.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 10))
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)
        if result.structure_png and Path(result.structure_png).exists():
            try:
                from PIL import Image, ImageTk
                img = Image.open(result.structure_png)
                img.thumbnail((920, 480))
                photo = ImageTk.PhotoImage(img)
                lbl = ttk.Label(frame, image=photo)
                lbl.image = photo
                lbl.grid(row=0, column=0, sticky="nsew")
            except Exception as e:
                ttk.Label(frame, text=f'Could not load the 2D structure image: {e}').grid(row=0, column=0)
        else:
            ttk.Label(frame, text='No 2D structure image was generated.').grid(row=0, column=0)

    def _on_run_smiles_builder(self):
        try:
            args = self._smiles_builder_args()
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return
        self.btn_diagnose_smiles.configure(state="disabled")
        self.btn_run_smiles.configure(state="disabled")
        self.btn_preview_smiles.configure(state="disabled")
        self.lbl_status_smiles.configure(text="Running...")
        self.txt_log_smiles.delete("1.0", "end")

        def worker():
            try:
                self._log_smiles(f"ABC component workbook: {args['component_workbook']}")
                self._log_smiles(f"Output master table: {args['output_xlsx']}")
                self._log_smiles("Join rule: ordered A/B/C component formulas are primary; IDs and A#/B#/C# numbers are ignored; final product Formula is used for verification")
                report = run_simple_combo_smiles(progress=self._log_smiles, **args)
                for w in report.warnings[:40]:
                    self._log_smiles("WARN: " + str(w))
                self.after(0, lambda: messagebox.showinfo(
                    'Complete',
                    f"Product SMILES table complete\nSheets: {report.n_sheets}\nCombinations: {report.n_total}\nSuccessful: {report.n_success}\nFormula matches: {report.n_formula_match}\nFailed: {report.n_failed}\nInput diagnostics: {(report.diagnostics.overall_status if report.diagnostics else '-')} (errors={(report.diagnostics.error_count if report.diagnostics else 0)}, warnings={(report.diagnostics.warning_count if report.diagnostics else 0)})\n\nOutput: {report.output_xlsx}\nSee the SMILES_Input_Diagnostics sheet for details.",
                ))
            except Exception as e:
                msg = str(e)
                self._log_smiles("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Failed', m))
            finally:
                self.after(0, lambda: self.btn_diagnose_smiles.configure(state="normal"))
                self.after(0, lambda: self.btn_run_smiles.configure(state="normal"))
                self.after(0, lambda: self.btn_preview_smiles.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_smiles.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()



    # =====================
    # ESI descriptors + LC gradient features
    # =====================
    def _build_esi_descriptor_tab(self):
        outer = self.tab_esi
        outer.columnconfigure(0, weight=1)
        outer.rowconfigure(0, weight=1)
        canvas = tk.Canvas(outer, highlightthickness=0, borderwidth=0)
        vbar = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vbar.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        vbar.grid(row=0, column=1, sticky="ns")
        frm = ttk.Frame(canvas)
        window_id = canvas.create_window((0, 0), window=frm, anchor="nw")
        self._esi_scroll_canvas = canvas
        self._esi_scroll_frame = frm

        def _sync_scrollregion(_event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))
        def _sync_width(event):
            canvas.itemconfigure(window_id, width=max(1, event.width))
        def _wheel(event):
            if getattr(event, "num", None) == 4:
                canvas.yview_scroll(-3, "units")
            elif getattr(event, "num", None) == 5:
                canvas.yview_scroll(3, "units")
            else:
                delta = getattr(event, "delta", 0)
                if delta:
                    canvas.yview_scroll(int(-delta / 120) * 3, "units")
            return "break"
        def _bind_wheel(_event=None):
            canvas.bind_all("<MouseWheel>", _wheel)
            canvas.bind_all("<Button-4>", _wheel)
            canvas.bind_all("<Button-5>", _wheel)
        def _unbind_wheel(_event=None):
            canvas.unbind_all("<MouseWheel>")
            canvas.unbind_all("<Button-4>")
            canvas.unbind_all("<Button-5>")
        frm.bind("<Configure>", _sync_scrollregion)
        canvas.bind("<Configure>", _sync_width)
        canvas.bind("<Enter>", _bind_wheel)
        frm.bind("<Enter>", _bind_wheel)
        canvas.bind("<Leave>", _unbind_wheel)

        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)
        r = 0

        ttk.Label(frm, text='Product SMILES master table:').grid(row=r, column=0, sticky="w", padx=10, pady=5)
        self.var_esi_master_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_esi_master_file).grid(row=r, column=1, sticky="ew", padx=10, pady=5)
        ttk.Button(frm, text='Browse...', command=self._browse_esi_master_file).grid(row=r, column=2, padx=10, pady=5)
        r += 1

        ttk.Label(frm, text='Known-concentration table (optional; up to 100 rows):').grid(row=r, column=0, sticky="w", padx=10, pady=5)
        self.var_esi_cal_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_esi_cal_file).grid(row=r, column=1, sticky="ew", padx=10, pady=5)
        ttk.Button(frm, text='Browse...', command=self._browse_esi_cal_file).grid(row=r, column=2, padx=10, pady=5)
        r += 1

        ttk.Label(frm, text='Target mixture table (optional):').grid(row=r, column=0, sticky="w", padx=10, pady=5)
        self.var_esi_target_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_esi_target_file).grid(row=r, column=1, sticky="ew", padx=10, pady=5)
        ttk.Button(frm, text='Browse...', command=self._browse_esi_target_file).grid(row=r, column=2, padx=10, pady=5)
        r += 1

        ttk.Label(frm, text='Output workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=5)
        self.var_esi_out_xlsx = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_esi_out_xlsx).grid(row=r, column=1, sticky="ew", padx=10, pady=5)
        ttk.Button(frm, text='Save as...', command=self._browse_esi_out_xlsx).grid(row=r, column=2, padx=10, pady=5)
        r += 1

        ttk.Label(frm, text="Reuse previous ESI descriptor workbook (skip RDKit/3D):").grid(row=r, column=0, sticky="w", padx=10, pady=5)
        self.var_esi_reuse_workbook = tk.StringVar(value="")
        ttk.Entry(frm, textvariable=self.var_esi_reuse_workbook).grid(row=r, column=1, sticky="ew", padx=10, pady=5)
        reuse_btns = ttk.Frame(frm)
        reuse_btns.grid(row=r, column=2, sticky="e", padx=10, pady=5)
        ttk.Button(reuse_btns, text="Select...", command=self._browse_esi_reuse_workbook).pack(side="left")
        ttk.Button(reuse_btns, text="Suggest poor groups", command=self._suggest_bad_esi_groups).pack(side="left", padx=(5, 0))
        r += 1

        columns = ttk.LabelFrame(frm, text='Sheets and columns (blank = automatic)')
        columns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=5)
        for c in range(8):
            columns.columnconfigure(c, weight=1)
        self.var_esi_master_sheet = tk.StringVar(value="Combo_SMILES_Master")
        self.var_esi_master_combo_col = tk.StringVar(value="Combo")
        self.var_esi_master_smiles_col = tk.StringVar(value="Product_SMILES")
        self.var_esi_cal_sheet = tk.StringVar(value="")
        self.var_esi_target_sheet = tk.StringVar(value="")
        self.var_esi_combo_col = tk.StringVar(value="")
        self.var_esi_formula_col = tk.StringVar(value="")
        self.var_esi_cal_rt_col = tk.StringVar(value="")
        self.var_esi_target_rt_col = tk.StringVar(value="")
        self.var_esi_cal_group_col = tk.StringVar(value="")

        labels_vars = [
            ('Master sheet', self.var_esi_master_sheet), ('Master Combo column', self.var_esi_master_combo_col),
            ('Master SMILES column', self.var_esi_master_smiles_col), ('Calibration sheet', self.var_esi_cal_sheet),
            ('Target sheet', self.var_esi_target_sheet), ('Combo column', self.var_esi_combo_col),
            ('Formula column', self.var_esi_formula_col), ('Calibration RT column', self.var_esi_cal_rt_col),
            ('Target RT column', self.var_esi_target_rt_col), ('Injection group / RAW column', self.var_esi_cal_group_col),
        ]
        for i, (lab, var) in enumerate(labels_vars):
            rr, cc = divmod(i, 3)
            base = cc * 2
            ttk.Label(columns, text=lab).grid(row=rr, column=base, sticky="w", padx=6, pady=3)
            ttk.Entry(columns, textvariable=var, width=16).grid(row=rr, column=base + 1, sticky="ew", padx=6, pady=3)
        r += 1

        grad = ttk.LabelFrame(frm, text='LC gradient (editable; linear interpolation)')
        grad.grid(row=r, column=0, columnspan=3, sticky="nsew", padx=10, pady=5)
        grad.columnconfigure(0, weight=1)
        grad.columnconfigure(1, weight=1)
        grad.rowconfigure(0, weight=1)
        self.txt_esi_gradient = tk.Text(grad, height=10, width=38, wrap="none")
        self.txt_esi_gradient.grid(row=0, column=0, rowspan=5, sticky="nsew", padx=8, pady=6)
        self.txt_esi_gradient.insert("1.0", default_gradient_text())

        self.var_esi_delay = tk.DoubleVar(value=0.0)
        self.var_esi_phase_a_name = tk.StringVar(value="Water + 0.1% formic acid")
        self.var_esi_phase_b_name = tk.StringVar(value="Acetonitrile")
        self.var_esi_use_3d = tk.BooleanVar(value=False)
        self.var_esi_extended = tk.BooleanVar(value=False)
        self.var_esi_engineered = tk.BooleanVar(value=True)
        self.var_esi_privacy = tk.BooleanVar(value=True)
        self.var_esi_output_csv = tk.BooleanVar(value=True)
        self.var_esi_language_label = tk.StringVar(value=ESI_LANGUAGE_LABELS[-1])
        default_unit_label = next((x for x in IE_CONCENTRATION_UNIT_LABELS if "mol/mL" in x), IE_CONCENTRATION_UNIT_LABELS[0])
        self.var_esi_concentration_unit_label = tk.StringVar(value=default_unit_label)

        ttk.Label(grad, text='Gradient delay (min)').grid(row=0, column=1, sticky="w", padx=8, pady=3)
        ttk.Entry(grad, textvariable=self.var_esi_delay, width=12).grid(row=0, column=2, sticky="w", padx=8, pady=3)
        ttk.Label(grad, text='Effective time = apex RT - delay. A zero delay leaves RT unchanged.', wraplength=420).grid(row=0, column=3, sticky="w", padx=8, pady=3)
        ttk.Label(grad, text='Mobile phase A').grid(row=1, column=1, sticky="w", padx=8, pady=3)
        ttk.Entry(grad, textvariable=self.var_esi_phase_a_name, width=24).grid(row=1, column=2, columnspan=2, sticky="ew", padx=8, pady=3)
        ttk.Label(grad, text='Mobile phase B').grid(row=2, column=1, sticky="w", padx=8, pady=3)
        ttk.Entry(grad, textvariable=self.var_esi_phase_b_name, width=24).grid(row=2, column=2, columnspan=2, sticky="ew", padx=8, pady=3)
        ttk.Checkbutton(grad, text='Compute 3D volume and shape descriptors', variable=self.var_esi_use_3d).grid(row=3, column=1, columnspan=2, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(grad, text='Include extended VSA/BCUT descriptors', variable=self.var_esi_extended).grid(row=3, column=3, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(grad, text='Privacy mode: omit product SMILES', variable=self.var_esi_privacy).grid(row=4, column=1, columnspan=2, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(grad, text='Also export CSV', variable=self.var_esi_output_csv).grid(row=4, column=3, sticky="w", padx=8, pady=3)
        ttk.Label(grad, text='Output language').grid(row=5, column=1, sticky="w", padx=8, pady=3)
        ttk.Combobox(
            grad, textvariable=self.var_esi_language_label,
            values=ESI_LANGUAGE_LABELS, state="readonly", width=16,
        ).grid(row=5, column=2, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(
            grad,
            text='Include derived interaction descriptors',
            variable=self.var_esi_engineered,
        ).grid(row=5, column=3, sticky="w", padx=8, pady=3)
        r += 1

        model_frame = ttk.LabelFrame(frm, text='ESI response models (target: log10 RRF)')
        model_frame.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=5)
        for c in range(9):
            model_frame.columnconfigure(c, weight=1)
        self.var_esi_run_models = tk.BooleanVar(value=True)
        self.var_esi_model_calibration_label = tk.StringVar(value=ESI_CALIBRATION_MODE_LABELS[0])
        self.var_esi_model_objective_label = tk.StringVar(value=ESI_MODEL_OBJECTIVE_LABELS[0])
        self.var_esi_model_cv_splits = tk.IntVar(value=5)
        self.var_esi_model_cv_repeats = tk.IntVar(value=3)
        self.var_esi_model_max_features = tk.IntVar(value=40)
        self.var_esi_model_seed = tk.IntVar(value=42)
        self.var_esi_outlier_mode_label = tk.StringVar(value=ESI_OUTLIER_MODE_LABELS[2])
        self.var_esi_outlier_min_fold = tk.DoubleVar(value=8.0)
        self.var_esi_outlier_max_pct = tk.DoubleVar(value=5.0)
        self.var_esi_outlier_consensus_pct = tk.DoubleVar(value=75.0)
        self.var_esi_auto_feature_select = tk.BooleanVar(value=True)
        self.var_esi_feature_stability_pct = tk.DoubleVar(value=50.0)
        self.var_esi_min_selected_features = tk.IntVar(value=6)
        self.var_esi_corr_threshold = tk.DoubleVar(value=0.97)
        self.var_esi_use_categorical = tk.BooleanVar(value=False)
        self.var_esi_deep_tuning = tk.BooleanVar(value=True)
        self.var_esi_trend_optimizer = tk.BooleanVar(value=True)
        self.var_esi_trend_trials = tk.IntVar(value=30)
        self.var_esi_trend_pair_min_fold = tk.DoubleVar(value=1.5)
        self.var_esi_trend_min_features = tk.IntVar(value=6)
        self.var_esi_trend_max_features = tk.IntVar(value=24)
        self.var_esi_trend_permutations = tk.IntVar(value=500)
        ttk.Checkbutton(
            model_frame, text='Compare Bayesian ridge, SVR, RF, ExtraTrees, HGB, GP, KNN and PLS',
            variable=self.var_esi_run_models,
        ).grid(row=0, column=0, columnspan=8, sticky="w", padx=8, pady=4)
        ttk.Label(model_frame, text='Calibration ratio source').grid(row=1, column=0, sticky="w", padx=8, pady=3)
        ttk.Combobox(
            model_frame, textvariable=self.var_esi_model_calibration_label,
            values=ESI_CALIBRATION_MODE_LABELS, state="readonly", width=42,
        ).grid(row=1, column=1, columnspan=3, sticky="ew", padx=8, pady=3)
        ttk.Label(model_frame, text='CV folds').grid(row=1, column=4, sticky="e", padx=4, pady=3)
        ttk.Spinbox(model_frame, from_=2, to=10, textvariable=self.var_esi_model_cv_splits, width=7).grid(row=1, column=5, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Repeats').grid(row=1, column=6, sticky="e", padx=4, pady=3)
        ttk.Spinbox(model_frame, from_=1, to=10, textvariable=self.var_esi_model_cv_repeats, width=7).grid(row=1, column=7, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Maximum features').grid(row=2, column=0, sticky="w", padx=8, pady=3)
        ttk.Spinbox(model_frame, from_=5, to=200, textvariable=self.var_esi_model_max_features, width=8).grid(row=2, column=1, sticky="w", padx=8, pady=3)
        ttk.Label(model_frame, text='Random seed').grid(row=2, column=2, sticky="e", padx=8, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_model_seed, width=9).grid(row=2, column=3, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(
            model_frame,
            text='Tune supported models within training folds',
            variable=self.var_esi_deep_tuning,
        ).grid(row=2, column=4, columnspan=2, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(
            model_frame,
            text='Include categorical component/gradient features',
            variable=self.var_esi_use_categorical,
        ).grid(row=2, column=6, columnspan=2, sticky="w", padx=8, pady=3)

        ttk.Label(model_frame, text='Outlier handling').grid(row=3, column=0, sticky="w", padx=8, pady=3)
        ttk.Combobox(
            model_frame, textvariable=self.var_esi_outlier_mode_label,
            values=ESI_OUTLIER_MODE_LABELS, state="readonly", width=42,
        ).grid(row=3, column=1, columnspan=3, sticky="ew", padx=8, pady=3)
        ttk.Label(model_frame, text='Fold-error threshold').grid(row=3, column=4, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_outlier_min_fold, width=8).grid(row=3, column=5, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Maximum excluded (%)').grid(row=3, column=6, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_outlier_max_pct, width=8).grid(row=3, column=7, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Directional consensus (%)').grid(row=4, column=0, sticky="w", padx=8, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_outlier_consensus_pct, width=8).grid(row=4, column=1, sticky="w", padx=8, pady=3)
        ttk.Label(
            model_frame,
            text=(
                'Threshold-consensus mode applies the specified fold-error and agreement thresholds. Ranked-CV mode excludes at most the specified fraction among records above the threshold. All excluded records remain in Outlier_Audit for review.'
            ),
            wraplength=760,
        ).grid(row=4, column=2, columnspan=6, sticky="w", padx=8, pady=3)

        ttk.Checkbutton(
            model_frame,
            text='Select descriptor count and evaluate selection stability',
            variable=self.var_esi_auto_feature_select,
        ).grid(row=5, column=0, columnspan=3, sticky="w", padx=8, pady=3)
        ttk.Label(model_frame, text='Minimum stability (%)').grid(row=5, column=3, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_feature_stability_pct, width=8).grid(row=5, column=4, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Minimum features').grid(row=5, column=5, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_min_selected_features, width=8).grid(row=5, column=6, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Correlation threshold').grid(row=6, column=0, sticky="w", padx=8, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_corr_threshold, width=8).grid(row=6, column=1, sticky="w", padx=8, pady=3)
        ttk.Label(
            model_frame,
            text=(
                'Numeric features are imputed and scaled where required; categorical inputs are one-hot encoded. Derived interaction descriptors are computational proxies, not measured physical properties.'
            ),
            wraplength=760,
        ).grid(row=6, column=2, columnspan=6, sticky="w", padx=8, pady=3)
        ttk.Label(model_frame, text='Selection objective').grid(row=7, column=0, sticky="w", padx=8, pady=3)
        ttk.Combobox(
            model_frame, textvariable=self.var_esi_model_objective_label,
            values=ESI_MODEL_OBJECTIVE_LABELS, state="readonly", width=42,
        ).grid(row=7, column=1, columnspan=3, sticky="ew", padx=8, pady=3)
        ttk.Label(
            model_frame,
            text='Use up to 100 known standards. Unknown mixture concentrations are not used for training; predictions remain response-corrected estimates.',
            wraplength=430,
        ).grid(row=7, column=4, columnspan=4, sticky="w", padx=8, pady=3)
        ttk.Checkbutton(
            model_frame,
            text='Run the trend optimizer (pairwise response model and RRF ExtraTrees)',
            variable=self.var_esi_trend_optimizer,
        ).grid(row=8, column=0, columnspan=3, sticky="w", padx=8, pady=3)
        ttk.Label(model_frame, text='Subset trials per fold').grid(row=8, column=3, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_trend_trials, width=7).grid(row=8, column=4, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Minimum pairwise concentration ratio').grid(row=8, column=5, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_trend_pair_min_fold, width=7).grid(row=8, column=6, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Permutations').grid(row=8, column=7, sticky="e", padx=4, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_trend_permutations, width=7).grid(row=8, column=8, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text='Trend-model feature range').grid(row=9, column=0, sticky="w", padx=8, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_trend_min_features, width=7).grid(row=9, column=1, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text="to").grid(row=9, column=2, sticky="e", padx=2, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_trend_max_features, width=7).grid(row=9, column=3, sticky="w", padx=4, pady=3)
        ttk.Label(
            model_frame,
            text='The trend optimizer selects descriptors and parameters within outer-training folds. Unknown target concentrations are not used.',
            wraplength=600,
        ).grid(row=9, column=4, columnspan=5, sticky="w", padx=8, pady=3)
        ttk.Label(model_frame, text="Known-standard concentration unit").grid(row=10, column=0, sticky="w", padx=8, pady=3)
        ttk.Combobox(
            model_frame, textvariable=self.var_esi_concentration_unit_label,
            values=IE_CONCENTRATION_UNIT_LABELS, state="readonly", width=30,
        ).grid(row=10, column=1, columnspan=3, sticky="ew", padx=8, pady=3)
        self.var_esi_excluded_groups = tk.StringVar(value="")
        ttk.Label(model_frame, text="Exclude standard injection groups").grid(row=10, column=4, sticky="e", padx=6, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_excluded_groups).grid(row=10, column=5, columnspan=3, sticky="ew", padx=6, pady=3)
        ttk.Label(model_frame, text="Comma-separated; e.g. Standard-6, Standard-9").grid(row=11, column=4, columnspan=4, sticky="w", padx=6, pady=2)
        ttk.Label(model_frame, text="Group exclusion is model-only; the source data remain in the audit sheet.").grid(row=11, column=0, columnspan=4, sticky="w", padx=8, pady=2)

        # Ten-level concentration classification is an additional supervised diagnostic.
        self.var_esi_level_classification = tk.BooleanVar(value=True)

        # Manual bilingual descriptor review and fixed-model random subset search.
        self.var_esi_manual_feature_summary = tk.StringVar(value="Manual pool: not set (all available descriptors)")
        self.var_esi_fixed_search_enabled = tk.BooleanVar(value=True)
        self.var_esi_fixed_model_label = tk.StringVar(value=ESI_FIXED_MODEL_LABELS[0])
        self.var_esi_fixed_trials = tk.IntVar(value=100)
        self.var_esi_fixed_min_features = tk.IntVar(value=4)
        self.var_esi_fixed_max_features = tk.IntVar(value=20)
        ttk.Label(model_frame, text='Descriptor pool for fixed-model search').grid(row=12, column=0, sticky="w", padx=8, pady=4)
        ttk.Button(model_frame, text='Select descriptors...', command=self._open_esi_descriptor_selector).grid(row=12, column=1, columnspan=2, sticky="ew", padx=6, pady=4)
        ttk.Button(model_frame, text="Reset", command=self._reset_esi_manual_features).grid(row=12, column=3, sticky="w", padx=4, pady=4)
        ttk.Label(model_frame, textvariable=self.var_esi_manual_feature_summary, wraplength=520).grid(row=12, column=4, columnspan=5, sticky="w", padx=6, pady=4)

        ttk.Checkbutton(
            model_frame,
            text='Run fixed-model random-subset search',
            variable=self.var_esi_fixed_search_enabled,
        ).grid(row=13, column=0, columnspan=3, sticky="w", padx=8, pady=4)
        ttk.Label(model_frame, text="Fixed model").grid(row=13, column=3, sticky="e", padx=4, pady=4)
        ttk.Combobox(
            model_frame, textvariable=self.var_esi_fixed_model_label,
            values=ESI_FIXED_MODEL_LABELS, state="readonly", width=28,
        ).grid(row=13, column=4, columnspan=2, sticky="ew", padx=4, pady=4)
        ttk.Label(model_frame, text="Trials").grid(row=13, column=6, sticky="e", padx=4, pady=4)
        ttk.Entry(model_frame, textvariable=self.var_esi_fixed_trials, width=8).grid(row=13, column=7, sticky="w", padx=4, pady=4)
        ttk.Label(model_frame, text="Feature count").grid(row=14, column=0, sticky="w", padx=8, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_fixed_min_features, width=7).grid(row=14, column=1, sticky="w", padx=4, pady=3)
        ttk.Label(model_frame, text="to").grid(row=14, column=2, sticky="e", padx=2, pady=3)
        ttk.Entry(model_frame, textvariable=self.var_esi_fixed_max_features, width=7).grid(row=14, column=3, sticky="w", padx=4, pady=3)
        ttk.Label(
            model_frame,
            text=(
                "The manual selector shows Chinese/English names, previous positive/negative effects, availability and mechanistic rationale. "
                "The fixed model is kept unchanged while descriptor subsets are randomly varied. Every trial is saved; separate top-5 tables are produced for absolute prediction and trend recovery."
            ),
            wraplength=720,
        ).grid(row=14, column=4, columnspan=5, sticky="w", padx=6, pady=3)
        ttk.Checkbutton(
            model_frame,
            text='Run ten-level concentration classification with OOF evaluation',
            variable=self.var_esi_level_classification,
        ).grid(row=15, column=0, columnspan=5, sticky="w", padx=8, pady=4)
        ttk.Label(
            model_frame,
            text="Outputs exact-level accuracy, +/-1-level accuracy, a confusion matrix, concise known-standard results and unknown-sample level predictions.",
            wraplength=520,
        ).grid(row=15, column=5, columnspan=4, sticky="w", padx=6, pady=4)
        r += 1

        dual_frame = ttk.LabelFrame(frm, text='Comparison with published ionization-efficiency methods')
        dual_frame.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=5)
        dual_frame.columnconfigure(1, weight=1)
        self.var_esi_dual_enabled = tk.BooleanVar(value=True)
        self.var_esi_published_mode_label = tk.StringVar(value=IE_PUBLISHED_MODE_LABELS[0])
        self.var_esi_ion_mode_label = tk.StringVar(value=IE_ION_MODE_LABELS[0])
        self.var_esi_adduct_label = tk.StringVar(value=IE_ADDUCT_LABELS[0])
        self.var_esi_organic_modifier_label = tk.StringVar(value=IE_ORGANIC_MODIFIER_LABELS[0])
        self.var_esi_formic_acid_pct = tk.DoubleVar(value=0.1)
        self.var_esi_organic_formic_acid_pct = tk.DoubleVar(value=0.0)
        self.var_esi_aqueous_ph = tk.DoubleVar(value=2.7)
        self.var_esi_nh4_present = tk.BooleanVar(value=False)
        self.var_esi_negative_proxies = tk.BooleanVar(value=True)
        self.var_esi_dynamic_range_label = tk.StringVar(value=IE_DYNAMIC_RANGE_MODE_LABELS[0])
        self.var_esi_external_logie_file = tk.StringVar(value="")
        self.var_esi_external_logie_sheet = tk.StringVar(value="")
        self.var_esi_external_combo_col = tk.StringVar(value="")
        self.var_esi_external_logie_col = tk.StringVar(value="")
        self.var_esi_rscript_path = tk.StringVar(value=detect_rscript(""))
        self.var_esi_instrument_name = tk.StringVar(value="Thermo Scientific Q Exactive HF")
        ttk.Checkbutton(
            dual_frame, text='Compare response models and export agreement results', variable=self.var_esi_dual_enabled,
        ).grid(row=0, column=0, sticky="w", padx=8, pady=5)
        ttk.Label(
            dual_frame,
            text="Default: negative-ESI literature RF framework, Q Exactive HF, aqueous phase with 0.1% formic acid.",
            wraplength=630,
        ).grid(row=0, column=1, sticky="w", padx=8, pady=5)
        ttk.Button(dual_frame, text='Comparison settings...', command=self._open_esi_dual_settings).grid(row=0, column=2, padx=8, pady=5)
        r += 1

        ttk.Label(
            frm,
            text=(
                'Acid/base and charge-derived fields are SMARTS/Gasteiger proxies, not pKa predictions. 3D descriptors describe calculated conformers. Known standards provide training labels; matching identities in the target mixture are used only for structure/RT diagnostics, never as known target concentrations.'
            ),
            wraplength=1000,
        ).grid(row=r, column=0, columnspan=3, sticky="w", padx=10, pady=(2, 5))
        r += 1

        audit_bar = ttk.Frame(frm)
        audit_bar.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        ttk.Button(audit_bar, text='Continuous concentration audit...',
                   command=self._open_continuous_audit).pack(side="left")
        ttk.Button(audit_bar, text='Multi-view response models...',
                   command=self._open_multiview_audit).pack(side="left", padx=12)
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        self.btn_run_esi = ttk.Button(btns, text='Generate descriptors and compare models', command=self._on_run_esi_descriptors)
        self.btn_run_esi.pack(side="left")
        ttk.Button(btns, text='Restore default gradient', command=self._reset_esi_gradient).pack(side="left", padx=8)
        ttk.Button(btns, text='Open output directory', command=self._open_esi_out_dir).pack(side="left", padx=8)
        ttk.Button(btns, text='Replot existing results', command=self._replot_esi_known_standards).pack(side="left", padx=8)
        self.lbl_status_esi = ttk.Label(btns, text="Ready")
        self.lbl_status_esi.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=5)
        self.txt_log_esi = tk.Text(frm, height=9, wrap="word")
        self.txt_log_esi.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=5)
        frm.rowconfigure(r, weight=1)

    def _open_continuous_audit(self):
        from core.continuous_lab import ContinuousLab
        cache = self.var_esi_reuse_workbook.get().strip() if hasattr(self, "var_esi_reuse_workbook") else ""
        if not cache and hasattr(self, "var_esi_out_xlsx"):
            candidate = self.var_esi_out_xlsx.get().strip()
            if candidate and Path(candidate).is_file():
                cache = candidate
        # Current table is visible and optional; it is not silently ignored on reuse.
        training = self.var_esi_cal_file.get().strip() if hasattr(self, "var_esi_cal_file") else ""
        ContinuousLab(self, cache_path=cache, training_path=training)

    def _open_multiview_audit(self):
        from core.multiview_lab import MultiviewLab
        cache = self.var_esi_reuse_workbook.get().strip() if hasattr(self, "var_esi_reuse_workbook") else ""
        if not cache and hasattr(self, "var_esi_out_xlsx"):
            candidate = self.var_esi_out_xlsx.get().strip()
            if candidate and Path(candidate).is_file():
                cache = candidate
        training = self.var_esi_cal_file.get().strip() if hasattr(self, "var_esi_cal_file") else ""
        MultiviewLab(self, cache_path=cache, training_path=training)

    def _browse_esi_reuse_workbook(self):
        path = filedialog.askopenfilename(
            title="Select previous ESI descriptor workbook",
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*")],
        )
        if path:
            self.var_esi_reuse_workbook.set(path)
            if not self.var_esi_out_xlsx.get().strip():
                p = Path(path)
                self.var_esi_out_xlsx.set(str(p.with_name(p.stem + "__model_rerun.xlsx")))
            self.esi_previous_best_model = detect_previous_best_model(Path(path))
            self.esi_manual_features = []
            self.esi_descriptor_review_rows = []
            suffix = f"; previous best={self.esi_previous_best_model}" if self.esi_previous_best_model else ""
            self.var_esi_manual_feature_summary.set("Manual pool: not set (all available descriptors)" + suffix)

    def _reset_esi_manual_features(self):
        self.esi_manual_features = []
        suffix = f"; previous best={self.esi_previous_best_model}" if self.esi_previous_best_model else ""
        self.var_esi_manual_feature_summary.set("Manual pool: not set (all available descriptors)" + suffix)

    def _open_esi_descriptor_selector(self):
        workbook = self.var_esi_reuse_workbook.get().strip()
        if not workbook:
            current_out = self.var_esi_out_xlsx.get().strip()
            if current_out and Path(current_out).exists():
                workbook = current_out
            else:
                workbook = filedialog.askopenfilename(
                    title="Select a previous ESI result workbook for descriptor review",
                    filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*")],
                )
                if workbook:
                    self.var_esi_reuse_workbook.set(workbook)
        if not workbook:
            messagebox.showinfo(
                "Descriptor selector",
                "Select a previous ESI result workbook first. The selector uses its hidden ESI_Features sheet and prior model outputs to show availability and positive/negative effects.",
            )
            return
        try:
            rows, previous = load_descriptor_review(Path(workbook), include_advanced_external=True)
        except Exception as exc:
            messagebox.showerror("Descriptor selector", str(exc))
            return
        if not rows:
            messagebox.showerror("Descriptor selector", "No descriptor catalog could be read from the selected workbook.")
            return
        self.esi_descriptor_review_rows = rows
        self.esi_previous_best_model = detect_previous_best_model(Path(workbook)) or self.esi_previous_best_model

        selected = set(self.esi_manual_features)
        if not selected:
            selected = {str(r.get("Feature")) for r in rows if bool(r.get("Use")) and str(r.get("Availability")) == "Available"}
        row_by_feature = {str(r.get("Feature")): r for r in rows}

        top = tk.Toplevel(self)
        top.title('Descriptor selection')
        top.geometry("1320x820")
        top.minsize(1000, 620)
        top.transient(self)
        top.grab_set()
        top.columnconfigure(0, weight=1)
        top.rowconfigure(1, weight=1)

        controls = ttk.Frame(top, padding=(10, 8))
        controls.grid(row=0, column=0, sticky="ew")
        controls.columnconfigure(1, weight=1)
        search_var = tk.StringVar(value="")
        group_values = ["All"] + sorted({str(r.get("Group")) for r in rows if str(r.get("Group"))})
        group_var = tk.StringVar(value="All")
        effect_var = tk.StringVar(value="All")
        availability_var = tk.StringVar(value="All")
        ttk.Label(controls, text='Search').grid(row=0, column=0, sticky="w", padx=4)
        ttk.Entry(controls, textvariable=search_var).grid(row=0, column=1, sticky="ew", padx=4)
        ttk.Label(controls, text='Group').grid(row=0, column=2, sticky="e", padx=4)
        ttk.Combobox(controls, textvariable=group_var, values=group_values, state="readonly", width=23).grid(row=0, column=3, sticky="w", padx=4)
        ttk.Label(controls, text='Previous effect').grid(row=0, column=4, sticky="e", padx=4)
        ttk.Combobox(
            controls, textvariable=effect_var,
            values=["All", "Positive", "Weak positive", "Mixed", "Weak negative", "Negative", "Unknown"],
            state="readonly", width=16,
        ).grid(row=0, column=5, sticky="w", padx=4)
        ttk.Label(controls, text='Availability').grid(row=0, column=6, sticky="e", padx=4)
        ttk.Combobox(
            controls, textvariable=availability_var,
            values=["All", "Available", "External input required", "Unavailable", "Too sparse", "Constant"],
            state="readonly", width=20,
        ).grid(row=0, column=7, sticky="w", padx=4)

        body = ttk.Panedwindow(top, orient="vertical")
        body.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 8))
        table_frame = ttk.Frame(body)
        details_frame = ttk.LabelFrame(body, text='Descriptor definition and previous-run evidence')
        body.add(table_frame, weight=4)
        body.add(details_frame, weight=2)
        table_frame.columnconfigure(0, weight=1)
        table_frame.rowconfigure(0, weight=1)

        columns = ("use", "technical", "zh", "en", "group", "level", "effect", "availability", "sel", "trend", "absolute")
        tree = ttk.Treeview(table_frame, columns=columns, show="headings", selectmode="browse")
        headings = {
            "use": "Use", "technical": "Technical field", "zh": 'Display name', "en": "English name",
            "group": "Mechanistic group", "level": "Mechanistic level", "effect": "Previous effect",
            "availability": "Availability", "sel": "Prev selected %", "trend": "Trend delta",
            "absolute": "Absolute benefit",
        }
        widths = {"use": 55, "technical": 190, "zh": 170, "en": 190, "group": 155, "level": 115, "effect": 120,
                  "availability": 145, "sel": 90, "trend": 90, "absolute": 95}
        for col in columns:
            tree.heading(col, text=headings[col])
            tree.column(col, width=widths[col], minwidth=50, stretch=(col in {"technical", "zh", "en", "group"}))
        yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        details_frame.columnconfigure(0, weight=1)
        details_frame.rowconfigure(0, weight=1)
        details = tk.Text(details_frame, height=10, wrap="word")
        details.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        details.configure(state="disabled")

        iid_to_feature = {}

        def fmt_num(value, digits=3):
            try:
                return f"{float(value):.{digits}f}"
            except Exception:
                return ""

        def populate(*_args):
            for iid in tree.get_children():
                tree.delete(iid)
            iid_to_feature.clear()
            query = search_var.get().strip().lower()
            for idx, row in enumerate(rows):
                feature = str(row.get("Feature", ""))
                blob = " ".join(str(row.get(k, "")) for k in ("Feature", "Name_zh", "Name_en", "Group", "Theory_zh", "Theory_en")).lower()
                if query and query not in blob:
                    continue
                if group_var.get() != "All" and str(row.get("Group")) != group_var.get():
                    continue
                if effect_var.get() != "All" and str(row.get("Previous_effect")) != effect_var.get():
                    continue
                if availability_var.get() != "All" and str(row.get("Availability")) != availability_var.get():
                    continue
                iid = f"d{idx}"
                iid_to_feature[iid] = feature
                sel_pct = row.get("Previous_outer_selected_pct", "")
                tree.insert("", "end", iid=iid, values=(
                    "Yes" if feature in selected else "No", feature, row.get("Name_zh", ""), row.get("Name_en", ""),
                    row.get("Group", ""), row.get("Mechanistic_level", ""), row.get("Previous_effect", ""),
                    row.get("Availability", ""), fmt_num(sel_pct, 1), fmt_num(row.get("Previous_trend_score_delta"), 3),
                    fmt_num(row.get("Previous_absolute_improvement"), 3),
                ))

        def show_details(_event=None):
            item = tree.focus()
            feature = iid_to_feature.get(item, "")
            row = row_by_feature.get(feature)
            if not row:
                return
            text = f"Technical field: {feature}\nDisplay name: {row.get('Name_zh', '')}\nEnglish name: {row.get('Name_en', '')}\nMechanistic group: {row.get('Group', '')}\nMechanistic level: {row.get('Mechanistic_level', '')}\nAvailability: {row.get('Availability', '')} | Valid N={row.get('Valid_N', '')} | Missing={row.get('Missing_pct', '')}%\nPrevious effect: {row.get('Previous_effect', '')} — {row.get('Previous_effect_summary', '')}\nPrevious selection frequency: {row.get('Previous_outer_selected_pct', '')}%\nPrevious trend delta: {row.get('Previous_trend_score_delta', '')}\nPrevious absolute benefit: {row.get('Previous_absolute_improvement', '')}\nPrevious permutation importance: {row.get('Previous_permutation_importance', '')}\n\nDefinition:\n{row.get('Theory_zh', '')}\n\nEnglish rationale:\n{row.get('Theory_en', '')}\n\nExpected direction:\n{row.get('Expected_direction_zh', '')}\n{row.get('Expected_direction_en', '')}\n\nSource: {row.get('Source', '')} | Unit/type: {row.get('Unit_or_type', '')} | Requires 3D: {row.get('Requires_3D', '')}\n\nNote: previous positive/negative effects are empirical and model-dependent; they are not causal proof."
            details.configure(state="normal")
            details.delete("1.0", "end")
            details.insert("1.0", text)
            details.configure(state="disabled")

        def toggle(_event=None):
            item = tree.focus()
            feature = iid_to_feature.get(item, "")
            row = row_by_feature.get(feature)
            if not row:
                return
            if str(row.get("Availability")) != "Available":
                messagebox.showinfo("Descriptor selector", "This descriptor is not available in the selected workbook. External pKa/WAPS/COSMO descriptors can be added later as imported numeric columns.")
                return
            if feature in selected:
                selected.remove(feature)
            else:
                selected.add(feature)
            values = list(tree.item(item, "values"))
            if values:
                values[0] = "Yes" if feature in selected else "No"
                tree.item(item, values=values)

        def apply_preset(mode):
            selected.clear()
            for row in rows:
                feature = str(row.get("Feature", ""))
                if str(row.get("Availability")) != "Available":
                    continue
                if mode == "mechanistic":
                    if str(row.get("Mechanistic_level")) not in {"Coarse covariate"}:
                        selected.add(feature)
                elif mode == "positive":
                    if str(row.get("Previous_effect")) in {"Positive", "Weak positive"}:
                        selected.add(feature)
                elif mode == "previous":
                    try:
                        if float(row.get("Previous_outer_selected_pct") or 0) >= 50:
                            selected.add(feature)
                    except Exception:
                        pass
                elif mode == "all":
                    selected.add(feature)
            populate()

        tree.bind("<<TreeviewSelect>>", show_details)
        tree.bind("<Double-1>", toggle)
        search_var.trace_add("write", populate)
        group_var.trace_add("write", populate)
        effect_var.trace_add("write", populate)
        availability_var.trace_add("write", populate)
        populate()

        actions = ttk.Frame(top, padding=(10, 0, 10, 10))
        actions.grid(row=2, column=0, sticky="ew")
        ttk.Button(actions, text='Mechanistic core', command=lambda: apply_preset("mechanistic")).pack(side="left", padx=3)
        ttk.Button(actions, text='Previously positive', command=lambda: apply_preset("positive")).pack(side="left", padx=3)
        ttk.Button(actions, text='Previously stable', command=lambda: apply_preset("previous")).pack(side="left", padx=3)
        ttk.Button(actions, text='All available', command=lambda: apply_preset("all")).pack(side="left", padx=3)
        ttk.Button(actions, text='Clear', command=lambda: (selected.clear(), populate(), update_count())).pack(side="left", padx=3)

        count_var = tk.StringVar(value="")
        def update_count(*_args):
            count_var.set(f'Selected {len(selected)} descriptors selected: {len(selected)} fields')
        update_count()
        ttk.Label(actions, textvariable=count_var).pack(side="left", padx=12)

        def save_and_close():
            self.esi_manual_features = sorted(selected)
            self.esi_previous_best_model = detect_previous_best_model(Path(workbook)) or self.esi_previous_best_model
            if self.esi_manual_features:
                preview = ", ".join(self.esi_manual_features[:4])
                more = "..." if len(self.esi_manual_features) > 4 else ""
                self.var_esi_manual_feature_summary.set(
                    f"Manual pool: {len(self.esi_manual_features)} descriptors ({preview}{more}); previous best={self.esi_previous_best_model or '-'}"
                )
            else:
                self.var_esi_manual_feature_summary.set(
                    f"Manual pool: empty (fixed search will use all available); previous best={self.esi_previous_best_model or '-'}"
                )
            top.destroy()

        ttk.Button(actions, text='Save selection', command=save_and_close).pack(side="right", padx=3)
        ttk.Button(actions, text='Cancel', command=top.destroy).pack(side="right", padx=3)

    def _suggest_bad_esi_groups(self):
        path = self.var_esi_reuse_workbook.get().strip()
        if not path:
            path = filedialog.askopenfilename(
                title="Select previous result workbook with Injection_Group_Summary",
                filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*")],
            )
        if not path:
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            if "Injection_Group_Summary" not in wb.sheetnames:
                wb.close()
                raise ValueError("Injection_Group_Summary was not found in the selected workbook")
            ws = wb["Injection_Group_Summary"]
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
            if not rows:
                raise ValueError("Injection_Group_Summary is empty")
            headers = [str(x or "").strip() for x in rows[0]]
            index = {h: i for i, h in enumerate(headers)}
            gcol = index.get("Injection_Group")
            scol = index.get("Trend_Spearman_r")
            dcol = index.get("Delta_Spearman_vs_raw")
            ncol = index.get("Trend_N")
            if gcol is None or scol is None:
                raise ValueError("Injection_Group or Trend_Spearman_r column is missing")
            candidates = []
            for values in rows[1:]:
                group = str(values[gcol] or "").strip() if gcol < len(values) else ""
                if not group or group == "ALL_GROUPS":
                    continue
                try:
                    score = float(values[scol])
                except Exception:
                    continue
                try:
                    delta = float(values[dcol]) if dcol is not None and dcol < len(values) and values[dcol] not in (None, "") else 0.0
                except Exception:
                    delta = 0.0
                try:
                    count = int(float(values[ncol])) if ncol is not None and ncol < len(values) and values[ncol] not in (None, "") else 0
                except Exception:
                    count = 0
                candidates.append((score, delta, -count, group))
            if not candidates:
                raise ValueError("No usable group validation rows were found")
            severe = [x for x in candidates if x[0] < 0]
            chosen = sorted(severe or candidates)[:2]
            groups = [x[3] for x in chosen]
            self.var_esi_excluded_groups.set(", ".join(groups))
            self.var_esi_reuse_workbook.set(path)
            detail = "\n".join(f"{g}: corrected Spearman={s:.3f}, delta={d:+.3f}" for s, d, _n, g in chosen)
            messagebox.showinfo(
                "Suggested groups",
                "The following groups were filled into the exclusion box for manual review:\n\n" + detail +
                "\n\nThe original rows will remain in Group_Exclusion_Audit.",
            )
        except Exception as exc:
            messagebox.showerror("Cannot suggest groups", str(exc))

    def _browse_esi_master_file(self):
        path = filedialog.askopenfilename(title='Select product SMILES master table', filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")])
        if path:
            self.var_esi_master_file.set(path)
            if not self.var_esi_out_xlsx.get().strip():
                self.var_esi_out_xlsx.set(str(Path(path).with_name(Path(path).stem + "__ESI_descriptors.xlsx")))

    def _browse_esi_cal_file(self):
        path = filedialog.askopenfilename(title='Select known-concentration table', filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")])
        if path:
            self.var_esi_cal_file.set(path)

    def _browse_esi_target_file(self):
        path = filedialog.askopenfilename(title='Select target mixture table', filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")])
        if path:
            self.var_esi_target_file.set(path)

    def _open_esi_dual_settings(self):
        top = tk.Toplevel(self)
        top.title("Dual IE validation / negative-ESI literature settings")
        top.geometry("1020x760")
        top.transient(self)
        top.grab_set()
        frm = ttk.Frame(top, padding=12)
        frm.pack(fill="both", expand=True)
        for c in range(4):
            frm.columnconfigure(c, weight=1)

        ttk.Checkbutton(
            frm, text="Run dual-model validation and consensus output",
            variable=self.var_esi_dual_enabled,
        ).grid(row=0, column=0, columnspan=4, sticky="w", pady=5)

        ttk.Label(frm, text="Literature benchmark mode").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Combobox(
            frm, textvariable=self.var_esi_published_mode_label,
            values=IE_PUBLISHED_MODE_LABELS, state="readonly", width=54,
        ).grid(row=1, column=1, columnspan=3, sticky="ew", pady=4)

        ttk.Label(frm, text="Instrument / method").grid(row=2, column=0, sticky="w", pady=4)
        ttk.Entry(frm, textvariable=self.var_esi_instrument_name).grid(row=2, column=1, columnspan=3, sticky="ew", pady=4)
        ttk.Label(frm, text="Known-standard concentration unit").grid(row=3, column=0, sticky="w", pady=4)
        ttk.Combobox(
            frm, textvariable=self.var_esi_concentration_unit_label,
            values=IE_CONCENTRATION_UNIT_LABELS, state="readonly", width=32,
        ).grid(row=3, column=1, sticky="ew", pady=4)
        ttk.Label(
            frm,
            text="The selected unit is used to build the response factor. A wrong unit produces a systematic calibration error.",
            wraplength=500,
        ).grid(row=3, column=2, columnspan=2, sticky="w", padx=8, pady=4)

        ttk.Label(frm, text="Ion mode").grid(row=4, column=0, sticky="w", pady=4)
        ttk.Combobox(
            frm, textvariable=self.var_esi_ion_mode_label,
            values=IE_ION_MODE_LABELS, state="readonly", width=18,
        ).grid(row=4, column=1, sticky="w", pady=4)
        ttk.Label(frm, text="Primary ion form").grid(row=4, column=2, sticky="e", padx=8, pady=4)
        ttk.Combobox(
            frm, textvariable=self.var_esi_adduct_label,
            values=IE_ADDUCT_LABELS, state="readonly", width=30,
        ).grid(row=4, column=3, sticky="ew", pady=4)

        ttk.Label(frm, text="Organic modifier").grid(row=5, column=0, sticky="w", pady=4)
        ttk.Combobox(
            frm, textvariable=self.var_esi_organic_modifier_label,
            values=IE_ORGANIC_MODIFIER_LABELS, state="readonly", width=34,
        ).grid(row=5, column=1, columnspan=3, sticky="ew", pady=4)

        ttk.Label(frm, text="Aqueous-phase formic acid (%)").grid(row=6, column=0, sticky="w", pady=4)
        ttk.Entry(frm, textvariable=self.var_esi_formic_acid_pct, width=10).grid(row=6, column=1, sticky="w", pady=4)
        ttk.Label(frm, text="Organic-phase formic acid (%)").grid(row=6, column=2, sticky="e", padx=8, pady=4)
        ttk.Entry(frm, textvariable=self.var_esi_organic_formic_acid_pct, width=10).grid(row=6, column=3, sticky="w", pady=4)

        ttk.Label(frm, text="Aqueous pH proxy").grid(row=7, column=0, sticky="w", pady=4)
        ttk.Entry(frm, textvariable=self.var_esi_aqueous_ph, width=10).grid(row=7, column=1, sticky="w", pady=4)
        ttk.Checkbutton(
            frm, text="Mobile phase contains NH4 / ammonium",
            variable=self.var_esi_nh4_present,
        ).grid(row=7, column=2, columnspan=2, sticky="w", pady=4)

        ttk.Checkbutton(
            frm, text="Use negative-ion acidity / ionization / charge-delocalization proxies",
            variable=self.var_esi_negative_proxies,
        ).grid(row=8, column=0, columnspan=2, sticky="w", pady=4)
        ttk.Label(frm, text="High-range response correction").grid(row=8, column=2, sticky="e", padx=8, pady=4)
        ttk.Combobox(
            frm, textvariable=self.var_esi_dynamic_range_label,
            values=IE_DYNAMIC_RANGE_MODE_LABELS, state="readonly", width=34,
        ).grid(row=8, column=3, sticky="ew", pady=4)

        ttk.Label(
            frm,
            text=(
                "The negative-ESI literature mode uses molecular descriptors plus apex viscosity, surface tension, polarity, pH, "
                "formic-acid fraction and transparent anion-formation proxies. Dynamic correction is cross-fitted and does not prove analyte-specific saturation."
            ),
            wraplength=900,
        ).grid(row=9, column=0, columnspan=4, sticky="w", pady=6)

        official = ttk.LabelFrame(frm, text="Optional official MS2Quant route (positive ions only)")
        official.grid(row=10, column=0, columnspan=4, sticky="ew", pady=8)
        official.columnconfigure(1, weight=1)
        ttk.Label(official, text="Rscript.exe").grid(row=0, column=0, sticky="w", padx=6, pady=5)
        ttk.Entry(official, textvariable=self.var_esi_rscript_path).grid(row=0, column=1, sticky="ew", padx=6, pady=5)
        ttk.Button(official, text="Browse", command=self._browse_esi_rscript).grid(row=0, column=2, padx=4, pady=5)
        ttk.Button(official, text="Check R / MS2Quant", command=self._check_esi_ms2quant).grid(row=0, column=3, padx=4, pady=5)
        ttk.Button(official, text="Install / update official package", command=self._install_esi_ms2quant).grid(row=0, column=4, padx=4, pady=5)
        ttk.Label(
            official,
            text="The public pretrained MS2Quant route is retained for positive ESI [M+H]+ or [M]+ only.",
            wraplength=850,
        ).grid(row=1, column=0, columnspan=5, sticky="w", padx=6, pady=5)

        ext = ttk.LabelFrame(frm, text="Optional external predicted-logIE table")
        ext.grid(row=11, column=0, columnspan=4, sticky="ew", pady=6)
        for c in range(4):
            ext.columnconfigure(c, weight=1)
        ttk.Entry(ext, textvariable=self.var_esi_external_logie_file).grid(row=0, column=0, columnspan=3, sticky="ew", padx=6, pady=5)
        ttk.Button(ext, text="Browse", command=self._browse_esi_external_logie_file).grid(row=0, column=3, padx=6, pady=5)
        ttk.Label(ext, text="Sheet").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(ext, textvariable=self.var_esi_external_logie_sheet).grid(row=1, column=1, sticky="ew", padx=6, pady=4)
        ttk.Label(ext, text="Combo column").grid(row=1, column=2, sticky="e", padx=6, pady=4)
        ttk.Entry(ext, textvariable=self.var_esi_external_combo_col).grid(row=1, column=3, sticky="ew", padx=6, pady=4)
        ttk.Label(ext, text="logIE column").grid(row=2, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(ext, textvariable=self.var_esi_external_logie_col).grid(row=2, column=1, sticky="ew", padx=6, pady=4)

        btn = ttk.Frame(frm)
        btn.grid(row=12, column=0, columnspan=4, sticky="e", pady=8)
        ttk.Button(btn, text="Close", command=top.destroy).pack(side="right", padx=5)

    def _browse_esi_rscript(self):
        path = filedialog.askopenfilename(
            title='Select Rscript.exe',
            filetypes=[("Rscript", "Rscript.exe"), ("Executable", "*.exe"), ("All", "*")],
        )
        if path:
            self.var_esi_rscript_path.set(path)

    def _check_esi_ms2quant(self):
        env = check_ms2quant_environment(self.var_esi_rscript_path.get().strip())
        if env.rscript_path and not self.var_esi_rscript_path.get().strip():
            self.var_esi_rscript_path.set(env.rscript_path)
        details = (
            f"Available: {env.available}\n"
            f"Rscript: {env.rscript_path or '-'}\n"
            f"R: {env.r_version or '-'}\n"
            f"MS2Quant: {env.package_version or '-'}\n"
            f"Function: {env.function_available}\n"
            f"JAVA_HOME: {env.java_home or '-'}\n\n"
            f"{env.message}"
        )
        if env.available:
            messagebox.showinfo('MS2Quant environment', details)
        else:
            messagebox.showwarning('MS2Quant is not ready', details)

    def _install_esi_ms2quant(self):
        rscript = detect_rscript(self.var_esi_rscript_path.get().strip())
        if not rscript:
            messagebox.showerror(
                'R was not found',
                'Rscript.exe was not found. Install 64-bit R and select Rscript.exe in the comparison settings.',
            )
            return
        self.var_esi_rscript_path.set(rscript)
        script = Path(__file__).resolve().parent / "r_scripts" / "install_ms2quant.R"
        if not script.exists():
            messagebox.showerror('Error', f'Installation script not found: {script}')
            return
        if not messagebox.askyesno(
            'Install MS2Quant',
            'This will install or update the KruveLab MS2Quant R package from its official GitHub repository.\nThe package requires rJava/rcdk; 64-bit Java is usually required on Windows.\n\nContinue?',
        ):
            return

        def worker():
            import subprocess
            try:
                proc = subprocess.run(
                    [rscript, "--vanilla", str(script)],
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                    text=True, encoding="utf-8", errors="replace",
                    timeout=3600, check=False,
                )
                details = (proc.stdout or "")[-8000:]
                if proc.stderr:
                    details += "\n\nSTDERR:\n" + (proc.stderr or "")[-5000:]
                if proc.returncode == 0:
                    self.after(0, lambda: messagebox.showinfo(
                        'MS2Quant installation completed',
                        'The MS2Quant installation command completed. Run the environment check to verify the installation.\n\n' + details[-3500:],
                    ))
                else:
                    self.after(0, lambda: messagebox.showerror(
                        'MS2Quant installation failed',
                        'The R installation script returned a nonzero exit code. Check the 64-bit Java/rJava installation.\n\n' + details[-4500:],
                    ))
            except Exception as exc:
                self.after(0, lambda msg=str(exc): messagebox.showerror('MS2Quant installation failed', msg))

        threading.Thread(target=worker, daemon=True).start()

    def _open_ms2quant_install_script(self):
        path = Path(__file__).resolve().parent / "r_scripts" / "install_ms2quant.R"
        if not path.exists():
            messagebox.showerror('Error', f'Installation script not found: {path}')
            return
        open_in_os(path)

    def _browse_esi_external_logie_file(self):
        path = filedialog.askopenfilename(
            title='Select external logIE predictions (optional)',
            filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")],
        )
        if path:
            self.var_esi_external_logie_file.set(path)

    def _browse_esi_out_xlsx(self):
        path = filedialog.asksaveasfilename(title='Save ESI descriptor workbook', defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("All", "*")])
        if path:
            self.var_esi_out_xlsx.set(path)

    def _open_esi_out_dir(self):
        p = self.var_esi_out_xlsx.get().strip()
        if p:
            open_in_os(Path(p).parent)

    def _replot_esi_known_standards(self):
        candidate = self.var_esi_out_xlsx.get().strip()
        if not candidate or not Path(candidate).exists():
            candidate = filedialog.askopenfilename(
                title="Select an existing ESI model result workbook",
                filetypes=[("Excel", "*.xlsx *.xlsm"), ("All", "*")],
            )
        if not candidate:
            return
        try:
            language = ESI_LANGUAGE_MAP.get(self.var_esi_language_label.get(), "en")
            outputs = replot_all_prediction_plots_from_workbook(Path(candidate), output_language=language)
            messagebox.showinfo(
                "Replot completed",
                "Plots saved:\n" + "\n".join(str(path) for path in outputs),
            )
            open_in_os(Path(outputs[0]).parent)
        except Exception as exc:
            messagebox.showerror("Replot failed", str(exc))

    def _reset_esi_gradient(self):
        self.txt_esi_gradient.delete("1.0", "end")
        self.txt_esi_gradient.insert("1.0", default_gradient_text())

    def _on_run_esi_descriptors(self):
        master = self.var_esi_master_file.get().strip()
        reuse_workbook = self.var_esi_reuse_workbook.get().strip()
        out = self.var_esi_out_xlsx.get().strip()
        if not master and not reuse_workbook:
            messagebox.showerror('Error', "Select a Combo-SMILES master or an existing ESI descriptor workbook to reuse")
            return
        if not out:
            messagebox.showerror('Error', 'Select an output workbook.')
            return
        try:
            points, parse_warnings = parse_gradient_text(self.txt_esi_gradient.get("1.0", "end"))
            delay = float(self.var_esi_delay.get())
            if not (-30.0 <= delay <= 30.0):
                raise ValueError('Gradient delay must be between -30 and 30 min.')
        except Exception as e:
            messagebox.showerror('Invalid gradient settings', str(e))
            return
        try:
            model_splits = int(self.var_esi_model_cv_splits.get())
            model_repeats = int(self.var_esi_model_cv_repeats.get())
            model_max_features = int(self.var_esi_model_max_features.get())
            model_seed = int(self.var_esi_model_seed.get())
            outlier_min_fold = float(self.var_esi_outlier_min_fold.get())
            outlier_max_pct = float(self.var_esi_outlier_max_pct.get())
            outlier_consensus_pct = float(self.var_esi_outlier_consensus_pct.get())
            feature_stability_pct = float(self.var_esi_feature_stability_pct.get())
            min_selected_features = int(self.var_esi_min_selected_features.get())
            corr_threshold = float(self.var_esi_corr_threshold.get())
            trend_trials = int(self.var_esi_trend_trials.get())
            trend_pair_min_fold = float(self.var_esi_trend_pair_min_fold.get())
            trend_min_features = int(self.var_esi_trend_min_features.get())
            trend_max_features = int(self.var_esi_trend_max_features.get())
            trend_permutations = int(self.var_esi_trend_permutations.get())
            fixed_trials = int(self.var_esi_fixed_trials.get())
            fixed_min_features = int(self.var_esi_fixed_min_features.get())
            fixed_max_features = int(self.var_esi_fixed_max_features.get())
            fixed_model_code = ESI_FIXED_MODEL_MAP.get(self.var_esi_fixed_model_label.get(), "auto_previous")
            formic_acid_pct = float(self.var_esi_formic_acid_pct.get())
            organic_formic_acid_pct = float(self.var_esi_organic_formic_acid_pct.get())
            aqueous_ph = float(self.var_esi_aqueous_ph.get())
            if not (2 <= model_splits <= 10):
                raise ValueError('CV folds must be between 2 and 10.')
            if not (1 <= model_repeats <= 10):
                raise ValueError('CV repeats must be between 1 and 10.')
            if not (5 <= model_max_features <= 200):
                raise ValueError('Maximum selected features must be between 5 and 200.')
            if not (1.01 <= outlier_min_fold <= 1000):
                raise ValueError('The outlier fold-error threshold must be between 1.01 and 1000.')
            if not (0 <= outlier_max_pct <= 25):
                raise ValueError('The maximum exclusion fraction must be between 0 and 25%.')
            if not (0 <= outlier_consensus_pct <= 100):
                raise ValueError('Directional model consensus must be between 0 and 100%.')
            if not (0 <= feature_stability_pct <= 100):
                raise ValueError('Selection stability must be between 0 and 100%.')
            if not (2 <= min_selected_features <= model_max_features):
                raise ValueError('The minimum feature count must be at least 2 and no greater than the maximum.')
            if not (0.80 <= corr_threshold < 1.0):
                raise ValueError('The correlation threshold must be between 0.80 and 0.999.')
            if not (8 <= trend_trials <= 300):
                raise ValueError('Random-subset trials must be between 8 and 300.')
            if not (1.01 <= trend_pair_min_fold <= 10.0):
                raise ValueError('The pairwise concentration-ratio threshold must be between 1.01 and 10.')
            if not (2 <= trend_min_features <= trend_max_features <= 100):
                raise ValueError('Invalid feature-count range for the trend model.')
            if not (20 <= trend_permutations <= 5000):
                raise ValueError('The number of permutations must be between 20 and 5000.')
            if not (5 <= fixed_trials <= 5000):
                raise ValueError("Fixed-model random trials must be between 5 and 5000")
            if not (1 <= fixed_min_features <= fixed_max_features <= 200):
                raise ValueError("Fixed-model feature-count range is invalid")
            if not (0.0 <= formic_acid_pct <= 10.0):
                raise ValueError("Aqueous-phase formic acid must be between 0 and 10%")
            if not (0.0 <= organic_formic_acid_pct <= 10.0):
                raise ValueError("Organic-phase formic acid must be between 0 and 10%")
            if not (0.0 <= aqueous_ph <= 14.0):
                raise ValueError('Aqueous-phase pH must be between 0 and 14.')
            selected_published_mode = IE_PUBLISHED_MODE_MAP.get(
                self.var_esi_published_mode_label.get(), "negative_rf_literature"
            )
            selected_concentration_unit = IE_CONCENTRATION_UNIT_MAP.get(
                self.var_esi_concentration_unit_label.get(), ""
            )
            if bool(self.var_esi_dual_enabled.get()) and selected_published_mode == "official_ms2quant_r" and not selected_concentration_unit:
                raise ValueError('Select the concentration unit before running MS2Quant.')
        except Exception as e:
            messagebox.showerror('Invalid model settings', str(e))
            return

        self.btn_run_esi.configure(state="disabled")
        self.lbl_status_esi.configure(text="Running...")
        self.txt_log_esi.delete("1.0", "end")

        def worker():
            try:
                if reuse_workbook:
                    self._log_esi(f"Reuse descriptor workbook (model-only): {reuse_workbook}")
                else:
                    self._log_esi(f"Product master: {master}")
                if self.var_esi_cal_file.get().strip():
                    self._log_esi(f"Training table: {self.var_esi_cal_file.get().strip()}")
                if self.var_esi_target_file.get().strip():
                    self._log_esi(f"Target table:   {self.var_esi_target_file.get().strip()}")
                self._log_esi(f"Gradient points: {len(points)} | delay={delay:g} min")
                for w in parse_warnings:
                    self._log_esi("WARN: " + w)
                self._log_esi(
                    f"Descriptors: 3D={bool(self.var_esi_use_3d.get())} | extended={bool(self.var_esi_extended.get())} | privacy={bool(self.var_esi_privacy.get())}"
                )
                model_mode = "training_original"
                model_objective = ESI_MODEL_OBJECTIVE_MAP.get(
                    self.var_esi_model_objective_label.get(), "trend"
                )
                outlier_mode = ESI_OUTLIER_MODE_MAP.get(
                    self.var_esi_outlier_mode_label.get(), "threshold_consensus"
                )
                output_language = ESI_LANGUAGE_MAP.get(
                    self.var_esi_language_label.get(), "zh"
                )
                published_mode = IE_PUBLISHED_MODE_MAP.get(
                    self.var_esi_published_mode_label.get(), "negative_rf_literature"
                )
                published_ion_mode = IE_ION_MODE_MAP.get(
                    self.var_esi_ion_mode_label.get(), "negative"
                )
                published_adduct = IE_ADDUCT_MAP.get(
                    self.var_esi_adduct_label.get(), "[M-H]-"
                )
                published_organic_modifier = IE_ORGANIC_MODIFIER_MAP.get(
                    self.var_esi_organic_modifier_label.get(), "acetonitrile"
                )
                published_concentration_unit = IE_CONCENTRATION_UNIT_MAP.get(
                    self.var_esi_concentration_unit_label.get(), ""
                )
                published_dynamic_range_mode = IE_DYNAMIC_RANGE_MODE_MAP.get(
                    self.var_esi_dynamic_range_label.get(), "auto"
                )
                published_instrument_name = self.var_esi_instrument_name.get().strip() or "Thermo Scientific Q Exactive HF"
                published_rscript_path = self.var_esi_rscript_path.get().strip()
                self._log_esi(
                    f"Multi-model: enabled={bool(self.var_esi_run_models.get())} | calibration=known_standards_only | objective={model_objective} | "
                    f"CV={model_splits}-fold x {model_repeats} | max_features={model_max_features} | "
                    f"deep_tuning={bool(self.var_esi_deep_tuning.get())}"
                )
                self._log_esi(
                    f"Training QC: mode={outlier_mode} | min_fold={outlier_min_fold:g}x | "
                    f"max_remove={outlier_max_pct:g}% | consensus={outlier_consensus_pct:g}%"
                )
                self._log_esi(
                    f"Descriptor selection: auto={bool(self.var_esi_auto_feature_select.get())} | "
                    f"stability>={feature_stability_pct:g}% | min={min_selected_features} | corr<{corr_threshold:g} | "
                    f"engineered={bool(self.var_esi_engineered.get())} | categorical_onehot={bool(self.var_esi_use_categorical.get())} | "
                    f"language={output_language}"
                )
                self._log_esi(
                    f"Trend optimizer: enabled={bool(self.var_esi_trend_optimizer.get())} | models=PhysicalPairwiseRanker+ResponseFactorExtraTrees | "
                    f"trials/fold={trend_trials} | pair_min={trend_pair_min_fold:g}x | features={trend_min_features}-{trend_max_features} | "
                    f"permutations={trend_permutations}"
                )
                self._log_esi(
                    f"Fixed feature search: enabled={bool(self.var_esi_fixed_search_enabled.get())} | "
                    f"model={fixed_model_code} | previous_best={self.esi_previous_best_model or '-'} | "
                    f"trials={fixed_trials} | features={fixed_min_features}-{fixed_max_features} | "
                    f"manual_pool={len(self.esi_manual_features) if self.esi_manual_features else 'all'}"
                )
                self._log_esi(
                    f"Dual IE: enabled={bool(self.var_esi_dual_enabled.get())} | mode={published_mode} | "
                    f"ion={published_ion_mode} | adduct={published_adduct} | organic={published_organic_modifier} | "
                    f"instrument={published_instrument_name} | conc_unit={published_concentration_unit} | "
                    f"A_FA={formic_acid_pct:g}%, B_FA={organic_formic_acid_pct:g}%, pH={aqueous_ph:g}, "
                    f"NH4={bool(self.var_esi_nh4_present.get())}, negative_proxies={bool(self.var_esi_negative_proxies.get())}, "
                    f"dynamic_range={published_dynamic_range_mode}"
                )
                report = run_esi_descriptor_export(
                    Path(master or reuse_workbook),
                    Path(out),
                    calibration_file=Path(self.var_esi_cal_file.get().strip()) if self.var_esi_cal_file.get().strip() else None,
                    target_file=Path(self.var_esi_target_file.get().strip()) if self.var_esi_target_file.get().strip() else None,
                    product_master_sheet=self.var_esi_master_sheet.get().strip(),
                    product_master_combo_col=self.var_esi_master_combo_col.get().strip(),
                    product_master_smiles_col=self.var_esi_master_smiles_col.get().strip(),
                    calibration_sheet=self.var_esi_cal_sheet.get().strip(),
                    target_sheet=self.var_esi_target_sheet.get().strip(),
                    combo_col=self.var_esi_combo_col.get().strip(),
                    formula_col=self.var_esi_formula_col.get().strip(),
                    calibration_rt_col=self.var_esi_cal_rt_col.get().strip(),
                    target_rt_col=self.var_esi_target_rt_col.get().strip(),
                    calibration_group_col=self.var_esi_cal_group_col.get().strip(),
                    gradient_points=points,
                    gradient_delay_min=delay,
                    mobile_phase_a_name=self.var_esi_phase_a_name.get().strip() or "Water + 0.1% formic acid",
                    mobile_phase_b_name=self.var_esi_phase_b_name.get().strip() or "Acetonitrile",
                    use_3d=bool(self.var_esi_use_3d.get()),
                    include_extended=bool(self.var_esi_extended.get()),
                    include_engineered=bool(self.var_esi_engineered.get()),
                    privacy_mode=bool(self.var_esi_privacy.get()),
                    output_language=output_language,
                    output_csv=bool(self.var_esi_output_csv.get()),
                    run_models=bool(self.var_esi_run_models.get()),
                    model_calibration_mode="training_original",
                    model_objective=model_objective,
                    model_cv_splits=model_splits,
                    model_cv_repeats=model_repeats,
                    model_max_features=model_max_features,
                    model_random_state=model_seed,
                    model_auto_remove_outliers=(outlier_mode != "off"),
                    model_outlier_mode=outlier_mode,
                    model_outlier_min_fold_error=outlier_min_fold,
                    model_outlier_max_fraction_pct=outlier_max_pct,
                    model_outlier_consensus_pct=outlier_consensus_pct,
                    model_auto_select_features=bool(self.var_esi_auto_feature_select.get()),
                    model_use_categorical_features=bool(self.var_esi_use_categorical.get()),
                    model_deep_tuning=bool(self.var_esi_deep_tuning.get()),
                    model_feature_stability_threshold_pct=feature_stability_pct,
                    model_min_selected_features=min_selected_features,
                    model_correlation_threshold=corr_threshold,
                    trend_optimizer_enabled=bool(self.var_esi_trend_optimizer.get()),
                    trend_optimizer_trials=trend_trials,
                    trend_optimizer_pair_min_fold=trend_pair_min_fold,
                    trend_optimizer_min_features=trend_min_features,
                    trend_optimizer_max_features=trend_max_features,
                    trend_optimizer_permutations=trend_permutations,
                    fixed_feature_search_enabled=bool(self.var_esi_fixed_search_enabled.get()),
                    fixed_feature_model=fixed_model_code,
                    fixed_feature_previous_best_model=self.esi_previous_best_model,
                    fixed_feature_trials=fixed_trials,
                    fixed_feature_min_features=fixed_min_features,
                    fixed_feature_max_features=fixed_max_features,
                    manual_feature_pool=list(self.esi_manual_features),
                    level_classification_enabled=bool(self.var_esi_level_classification.get()),
                    dual_ie_enabled=bool(self.var_esi_dual_enabled.get()),
                    published_ie_mode=published_mode,
                    published_ion_mode=published_ion_mode,
                    published_adduct=published_adduct,
                    published_organic_modifier=published_organic_modifier,
                    published_aqueous_phase_name=self.var_esi_phase_a_name.get().strip() or "Water + 0.1% formic acid",
                    published_organic_phase_name=self.var_esi_phase_b_name.get().strip() or "Acetonitrile",
                    published_formic_acid_pct=formic_acid_pct,
                    published_organic_formic_acid_pct=organic_formic_acid_pct,
                    published_aqueous_pH=aqueous_ph,
                    published_nh4_present=bool(self.var_esi_nh4_present.get()),
                    published_negative_proxies_enabled=bool(self.var_esi_negative_proxies.get()),
                    published_dynamic_range_mode=published_dynamic_range_mode,
                    published_external_logie_file=(
                        Path(self.var_esi_external_logie_file.get().strip())
                        if self.var_esi_external_logie_file.get().strip() else None
                    ),
                    published_external_logie_sheet=self.var_esi_external_logie_sheet.get().strip(),
                    published_external_combo_col=self.var_esi_external_combo_col.get().strip(),
                    published_external_logie_col=self.var_esi_external_logie_col.get().strip(),
                    published_rscript_path=published_rscript_path,
                    published_instrument_name=published_instrument_name,
                    published_concentration_unit=published_concentration_unit,
                    reuse_descriptor_workbook=Path(reuse_workbook) if reuse_workbook else None,
                    excluded_injection_groups=self.var_esi_excluded_groups.get().strip(),
                )
                for w in report.warnings[:30]:
                    self._log_esi("WARN: " + str(w))
                self._log_esi(f"Master descriptors: {report.n_master}")
                self._log_esi(f"Training: {report.n_training}, structure matched={report.n_training_matched}, RT found={report.n_rt_training}")
                self._log_esi(f"Target:   {report.n_target}, structure matched={report.n_target_matched}, RT found={report.n_rt_target}")
                if report.model_benchmark_enabled:
                    self._log_esi(
                        f"Best model: {report.best_model} | rating={report.best_model_rating} | "
                        f"calibration={report.model_calibration_rows} | target={report.model_target_rows}"
                    )
                    metric_parts = []
                    if report.best_median_fold_error is not None:
                        metric_parts.append(f"median fold={report.best_median_fold_error:.2f}x")
                    if report.best_p80_fold_error is not None:
                        metric_parts.append(f"P80={report.best_p80_fold_error:.2f}x")
                    if report.best_within_2x_pct is not None:
                        metric_parts.append(f"Within2x={report.best_within_2x_pct:.1f}%")
                    if report.best_within_5x_pct is not None:
                        metric_parts.append(f"Within5x={report.best_within_5x_pct:.1f}%")
                    if report.best_r2_log is not None:
                        metric_parts.append(f"R2log={report.best_r2_log:.3f}")
                    if report.improvement_vs_global_pct is not None:
                        metric_parts.append(f"vs GlobalMedian={report.improvement_vs_global_pct:+.1f}%")
                    if report.corrected_spearman is not None:
                        metric_parts.append(f"corrected Spearman={report.corrected_spearman:.3f}")
                    if report.raw_spearman is not None:
                        metric_parts.append(f"raw Spearman={report.raw_spearman:.3f}")
                    if report.delta_spearman is not None:
                        metric_parts.append(f"delta Spearman={report.delta_spearman:+.3f}")
                    if report.pairwise_concordance_pct is not None:
                        metric_parts.append(f"pairwise={report.pairwise_concordance_pct:.1f}%")
                    if report.top20_overlap_pct is not None:
                        metric_parts.append(f"Top20={report.top20_overlap_pct:.1f}%")
                    if report.injection_group_spearman is not None:
                        metric_parts.append(f"group-out Spearman={report.injection_group_spearman:.3f}")
                    metric_parts.append(f"outliers_removed={report.model_outliers_removed}")
                    metric_parts.append(f"selected_descriptors={report.model_selected_descriptors}")
                    if metric_parts:
                        self._log_esi("Repeated-CV: " + " | ".join(metric_parts))
                    self._log_esi(
                        f"Descriptor QC: selected={report.model_selected_descriptors} | "
                        f"unavailable/too-sparse={report.model_unavailable_descriptors}"
                    )
                    if report.model_decision:
                        self._log_esi("MODEL DECISION: " + report.model_decision)
                    if report.model_decision_summary:
                        self._log_esi(report.model_decision_summary)
                    if report.trend_optimizer_enabled:
                        trend_parts = [
                            f"model={report.trend_optimizer_best_model}",
                            f"decision={report.trend_optimizer_decision}",
                        ]
                        if report.trend_optimizer_spearman is not None:
                            trend_parts.append(f"OOF Spearman={report.trend_optimizer_spearman:.3f}")
                        if report.trend_optimizer_delta_spearman is not None:
                            trend_parts.append(f"delta vs raw={report.trend_optimizer_delta_spearman:+.3f}")
                        if report.trend_optimizer_permutation_p is not None:
                            trend_parts.append(f"permutation p={report.trend_optimizer_permutation_p:.4g}")
                        self._log_esi("TREND OPTIMIZER: " + " | ".join(trend_parts))
                        if report.trend_optimizer_summary:
                            self._log_esi(report.trend_optimizer_summary)
                    if getattr(report, "fixed_feature_search_enabled", False):
                        self._log_esi(
                            "FIXED FEATURE SEARCH: "
                            f"model={getattr(report, 'fixed_feature_model', '')} | "
                            f"best absolute trial={getattr(report, 'fixed_feature_best_absolute_trial', '')} | "
                            f"best trend trial={getattr(report, 'fixed_feature_best_trend_trial', '')}"
                        )
                    if getattr(report, "level_classification_enabled", False):
                        self._log_esi(
                            "TEN-LEVEL CLASSIFICATION: "
                            f"model={getattr(report, 'level_classification_best_model', '')} | "
                            f"exact={getattr(report, 'level_classification_exact_pct', 0.0):.1f}% | "
                            f"within +/-1={getattr(report, 'level_classification_within1_pct', 0.0):.1f}% | "
                            f"all-hit={getattr(report, 'level_classification_all_hit', False)}"
                        )
                    if report.dual_ie_enabled:
                        dual_parts = [
                            f"published={report.dual_ie_published_method}",
                            f"decision={report.dual_ie_decision}",
                        ]
                        if report.published_ie_spearman is not None:
                            dual_parts.append(f"published Spearman={report.published_ie_spearman:.3f}")
                        if report.published_ie_median_fold is not None:
                            dual_parts.append(f"published median fold={report.published_ie_median_fold:.2f}x")
                        if report.dual_ie_target_rank_spearman is not None:
                            dual_parts.append(f"target rank agreement={report.dual_ie_target_rank_spearman:.3f}")
                        self._log_esi("DUAL IE: " + " | ".join(dual_parts))
                        if report.dual_ie_summary:
                            self._log_esi(report.dual_ie_summary)
                elif bool(self.var_esi_run_models.get()):
                    self._log_esi("WARN: Multi-model benchmark did not complete; see Diagnostics sheet.")
                self._log_esi(f"Excel: {report.output_xlsx}")
                if report.output_csv:
                    self._log_esi(f"CSV:   {report.output_csv}")
                self.after(0, lambda: messagebox.showinfo('Complete', f'ESI descriptors and gradient features generated\n\n{report.output_xlsx}'))
            except Exception as e:
                msg = str(e)
                self._log_esi("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Failed', m))
            finally:
                self.after(0, lambda: self.btn_run_esi.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_esi.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()


    # =====================
    # CombiTrace-IE structural/physical response model tab
    # =====================
    def _build_combitrace_ie_tab(self):
        from core.ui_support import scroll_body
        frm = scroll_body(self.tab_ie)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)
        r = 0

        ttk.Label(frm, text='Calibration table (known concentration and area ratio):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_ie_cal_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_ie_cal_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_ie_cal_file).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Target table (Combo and ratio):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_ie_target_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_ie_target_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_ie_target_file).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Product SMILES master table:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_ie_product_master_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_ie_product_master_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_ie_product_master_file).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Additional structure library (optional):').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_ie_structure_file = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_ie_structure_file).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        bf = ttk.Frame(frm)
        bf.grid(row=r, column=2, sticky="e", padx=10, pady=6)
        ttk.Button(bf, text='Browse...', command=self._browse_ie_structure_file).pack(side="left")
        ttk.Button(bf, text='Export template', command=self._export_ie_structure_template).pack(side="left", padx=(6, 0))
        r += 1

        ttk.Label(frm, text='Output workbook:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_ie_out_xlsx = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_ie_out_xlsx).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Save as...', command=self._browse_ie_out_xlsx).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(
            frm,
            text=(
                'Select the product SMILES table generated in the SMILES tab. Matching uses the ordered A/B/C component-formula key, not local identifiers. The product formula is checked independently. Review mode flags a product-formula mismatch without rejecting a unique key; conflicting structures are not resolved automatically. The optional library supplies reaction rules or missing structures. All processing is local.'
            ),
            wraplength=900,
        ).grid(row=r, column=0, columnspan=3, sticky="w", padx=10, pady=(0, 6))
        r += 1

        box = ttk.LabelFrame(frm, text='Sheets and columns (blank = automatic)')
        box.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(6):
            box.columnconfigure(c, weight=1)
        self.var_ie_cal_sheet = tk.StringVar(value="")
        self.var_ie_target_sheet = tk.StringVar(value="")
        self.var_ie_structure_sheet = tk.StringVar(value="Structures")
        self.var_ie_reaction_sheet = tk.StringVar(value="ReactionRules")
        self.var_ie_combo_col = tk.StringVar(value="")
        self.var_ie_formula_col = tk.StringVar(value="")
        self.var_ie_ratio_col = tk.StringVar(value="")
        self.var_ie_target_ratio_col = tk.StringVar(value="")
        self.var_ie_conc_col = tk.StringVar(value="")
        self.var_ie_product_smiles_col = tk.StringVar(value="")
        self.var_ie_product_master_sheet = tk.StringVar(value="Combo_SMILES_Master")
        self.var_ie_product_master_combo_col = tk.StringVar(value="Combo")
        self.var_ie_product_master_smiles_col = tk.StringVar(value="Product_SMILES")

        ttk.Label(box, text='Calibration sheet').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_cal_sheet, width=16).grid(row=0, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Target sheet').grid(row=0, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_target_sheet, width=16).grid(row=0, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Structure sheet').grid(row=0, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_structure_sheet, width=16).grid(row=0, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(box, text='Reaction-rules sheet').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_reaction_sheet, width=16).grid(row=1, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Combo column').grid(row=1, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_combo_col, width=16).grid(row=1, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Formula column').grid(row=1, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_formula_col, width=16).grid(row=1, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(box, text='Known concentration column').grid(row=2, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_conc_col, width=16).grid(row=2, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Calibration ratio column').grid(row=2, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_ratio_col, width=16).grid(row=2, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Target ratio column').grid(row=2, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_target_ratio_col, width=16).grid(row=2, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(box, text='Row-level product SMILES column').grid(row=3, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_product_smiles_col, width=16).grid(row=3, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Master sheet').grid(row=3, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_product_master_sheet, width=16).grid(row=3, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Master Combo column').grid(row=3, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_product_master_combo_col, width=16).grid(row=3, column=5, sticky="w", padx=8, pady=4)

        ttk.Label(box, text='Master product SMILES column').grid(row=4, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(box, textvariable=self.var_ie_product_master_smiles_col, width=18).grid(row=4, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(box, text='Structure priority: row-level product SMILES; A/B/C formula-key match; product-formula check; optional reaction library; formula/Combo fallback.').grid(row=4, column=2, columnspan=4, sticky="w", padx=8, pady=4)
        r += 1

        opt = ttk.LabelFrame(frm, text='Local structure-based response transfer')
        opt.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(8):
            opt.columnconfigure(c, weight=1, uniform="options")
        self.var_ie_k = tk.IntVar(value=15)
        self.var_ie_structure_threshold = tk.DoubleVar(value=0.60)
        self.var_ie_use_rt = tk.BooleanVar(value=True)
        self.var_ie_use_3d = tk.BooleanVar(value=False)
        self.var_ie_privacy = tk.BooleanVar(value=True)
        self.var_ie_output_csv = tk.BooleanVar(value=True)
        self.var_ie_formula_verify_mode = tk.StringVar(value='Flag mismatches without rejecting rows')
        self.var_ie_lambda = tk.DoubleVar(value=10.0)
        self.var_ie_cat_dim = tk.IntVar(value=96)

        ttk.Label(opt, text='Local neighbors (K)').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        ttk.Spinbox(opt, from_=1, to=100, textvariable=self.var_ie_k, width=8).grid(row=0, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Structural similarity threshold').grid(row=0, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(opt, textvariable=self.var_ie_structure_threshold, width=10).grid(row=0, column=3, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(opt, text='Use apex RT', variable=self.var_ie_use_rt).grid(row=0, column=4, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(opt, text='Compute 3D descriptors', variable=self.var_ie_use_3d).grid(row=0, column=5, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(opt, text='Privacy mode', variable=self.var_ie_privacy).grid(row=0, column=6, sticky="w", padx=8, pady=4)
        ttk.Checkbutton(opt, text='Export CSV', variable=self.var_ie_output_csv).grid(row=0, column=7, sticky="w", padx=8, pady=4)
        ttk.Label(opt, text='Product-formula policy').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(
            opt, textvariable=self.var_ie_formula_verify_mode, state="readonly", width=26,
            values=['Flag mismatches without rejecting rows', 'Strict: reject formula mismatches'],
        ).grid(row=1, column=1, columnspan=2, sticky="w", padx=8, pady=4)
        ttk.Label(
            opt,
            text='Start with 2D fingerprints. The formula policy controls whether mismatches are flagged or rejected; every mismatch remains in Structure_Coverage and Structure_Join_Summary. Component formula keys are computational identity links, not independent structural confirmation.',
            wraplength=880,
        ).grid(row=2, column=0, columnspan=8, sticky="w", padx=8, pady=4)
        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)
        self.btn_run_ie = ttk.Button(btns, text='Run local response transfer', command=self._on_run_combitrace_ie)
        self.btn_run_ie.pack(side="left")
        ttk.Button(btns, text='Open output directory', command=self._open_ie_out_dir).pack(side="left", padx=10)
        self.lbl_status_ie = ttk.Label(btns, text="Ready")
        self.lbl_status_ie.pack(side="left", padx=10)
        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_log_ie = tk.Text(frm, height=12, wrap="word")
        self.txt_log_ie.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=6)
        frm.rowconfigure(r, weight=1)

    def _browse_ie_cal_file(self):
        path = filedialog.askopenfilename(title='Select calibration table', filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")])
        if path:
            self.var_ie_cal_file.set(path)
            if not self.var_ie_out_xlsx.get().strip():
                self.var_ie_out_xlsx.set(str(Path(path).with_name("CombiTrace_IE_response_model.xlsx")))

    def _browse_ie_target_file(self):
        path = filedialog.askopenfilename(title='Select target table', filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")])
        if path:
            self.var_ie_target_file.set(path)
            if not self.var_ie_out_xlsx.get().strip():
                self.var_ie_out_xlsx.set(str(Path(path).with_name(f"{Path(path).stem}__CombiTrace_IE.xlsx")))

    def _browse_ie_product_master_file(self):
        path = filedialog.askopenfilename(
            title='Select product SMILES master table',
            filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")],
        )
        if path:
            self.var_ie_product_master_file.set(path)

    def _browse_ie_structure_file(self):
        path = filedialog.askopenfilename(
            title='Select structure library (CORE / A / B / C / ReactionRules)',
            filetypes=[("Table", "*.xlsx *.xlsm *.csv"), ("All", "*")],
        )
        if path:
            self.var_ie_structure_file.set(path)

    def _export_ie_structure_template(self):
        src = Path(__file__).resolve().parent / "templates" / "CombiTrace_IE_private_structure_template.xlsx"
        if not src.exists():
            messagebox.showerror('Error', f'Template not found: {src}')
            return
        dst = filedialog.asksaveasfilename(
            title='Export structure-library template',
            defaultextension=".xlsx",
            initialfile="CombiTrace_IE_private_structure_template.xlsx",
            filetypes=[("Excel", "*.xlsx"), ("All", "*")],
        )
        if not dst:
            return
        try:
            shutil.copy2(str(src), str(dst))
            self.var_ie_structure_file.set(str(dst))
            messagebox.showinfo('Complete', f'Structure-library template exported:\n{dst}')
        except Exception as e:
            messagebox.showerror('Failed', str(e))

    def _browse_ie_out_xlsx(self):
        path = filedialog.asksaveasfilename(title='Save CombiTrace-IE workbook', defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("All", "*")])
        if path:
            self.var_ie_out_xlsx.set(path)

    def _open_ie_out_dir(self):
        p = self.var_ie_out_xlsx.get().strip()
        if p:
            open_in_os(Path(p).parent)

    def _on_run_combitrace_ie(self):
        cal_file = self.var_ie_cal_file.get().strip()
        target_file = self.var_ie_target_file.get().strip()
        out_xlsx = self.var_ie_out_xlsx.get().strip()
        structure_file = self.var_ie_structure_file.get().strip()
        product_master_file = self.var_ie_product_master_file.get().strip()
        if not cal_file or not target_file:
            messagebox.showerror('Error', 'Select calibration and target tables.')
            return
        if not out_xlsx:
            messagebox.showerror('Error', 'Select an output Excel file.')
            return
        try:
            k = int(self.var_ie_k.get())
            structure_threshold = float(self.var_ie_structure_threshold.get())
            if k <= 0:
                raise ValueError('The local neighbor count K must be greater than zero.')
            if not (0.0 <= structure_threshold <= 1.0):
                raise ValueError('The structural-similarity threshold must be between 0 and 1.')
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return

        formula_verify_mode = (
            "strict" if self.var_ie_formula_verify_mode.get().strip().startswith('Strict') else "review"
        )

        self.btn_run_ie.configure(state="disabled")
        self.lbl_status_ie.configure(text="Running...")
        self.txt_log_ie.delete("1.0", "end")

        def worker():
            try:
                self._log_ie(f"Calibration: {cal_file}")
                self._log_ie(f"Targets:     {target_file}")
                if product_master_file:
                    self._log_ie(f"Combo-SMILES master: {product_master_file}")
                else:
                    self._log_ie("Combo-SMILES master: not provided")
                if structure_file:
                    self._log_ie(f"Advanced private structure library: {structure_file}")
                else:
                    self._log_ie("Advanced private structure library: not provided")
                self._log_ie(f"Output:      {out_xlsx}")
                self._log_ie(
                    "Model:       structure-aware local RF transfer | "
                    f"K={k} | structure threshold={structure_threshold:.2f} | "
                    f"Product Formula review={formula_verify_mode} | "
                    f"3D={bool(self.var_ie_use_3d.get())} | privacy={bool(self.var_ie_privacy.get())}"
                )
                self._log_ie("Stoichiometry check: Product = A + B + C + CORE contribution - N2 - H")

                report = run_combitrace_ie(
                    Path(cal_file),
                    Path(target_file),
                    Path(out_xlsx),
                    structure_file=Path(structure_file) if structure_file else None,
                    product_master_file=Path(product_master_file) if product_master_file else None,
                    product_master_sheet=self.var_ie_product_master_sheet.get().strip(),
                    product_master_combo_col=self.var_ie_product_master_combo_col.get().strip(),
                    product_master_smiles_col=self.var_ie_product_master_smiles_col.get().strip(),
                    calibration_sheet=self.var_ie_cal_sheet.get().strip(),
                    target_sheet=self.var_ie_target_sheet.get().strip(),
                    structure_sheet=self.var_ie_structure_sheet.get().strip(),
                    reaction_sheet=self.var_ie_reaction_sheet.get().strip(),
                    combo_col=self.var_ie_combo_col.get().strip(),
                    formula_col=self.var_ie_formula_col.get().strip(),
                    ratio_col=self.var_ie_ratio_col.get().strip(),
                    concentration_col=self.var_ie_conc_col.get().strip(),
                    target_ratio_col=self.var_ie_target_ratio_col.get().strip(),
                    product_smiles_col=self.var_ie_product_smiles_col.get().strip(),
                    formula_verification_mode=formula_verify_mode,
                    k_neighbors=k,
                    use_rt=bool(self.var_ie_use_rt.get()),
                    use_3d=bool(self.var_ie_use_3d.get()),
                    privacy_mode=bool(self.var_ie_privacy.get()),
                    structure_threshold=structure_threshold,
                    output_csv=bool(self.var_ie_output_csv.get()),
                )
                for w in report.warnings[:40]:
                    self._log_ie("WARN: " + str(w))
                if len(report.warnings) > 40:
                    self._log_ie(f'WARN: {len(report.warnings) - 40} warnings; see the output workbook.')

                self._log_ie(f"Training rows: total={report.n_train_total}, used={report.n_train_used}")
                self._log_ie(f"Target rows:   total={report.n_target_total}, predicted={report.n_predicted}")
                if report.qc:
                    self._log_ie("Structure coverage/QC:")
                    for key in (
                        "RDKit_available",
                        "Structure_rows_ready_training",
                        "Structure_rows_ready_target",
                        "Product_structures_training",
                        "Product_structures_target",
                        "Formula_mismatch_training",
                        "Formula_mismatch_target",
                        "Product_master_rows",
                        "Product_master_usable",
                        "Product_master_match_training",
                        "Product_master_match_target",
                        "Product_master_formula_verification_mode",
                        "Product_master_formula_flagged_training",
                        "Product_master_formula_flagged_target",
                        "Product_master_ambiguous_training",
                        "Product_master_ambiguous_target",
                        "Product_master_formula_key_match_training",
                        "Product_master_formula_key_match_target",
                        "Product_master_formula_keys",
                        "Product_master_duplicate_formula_keys",
                        "Product_master_lookup_keys",
                    ):
                        if key in report.qc:
                            self._log_ie(f"  {key}: {report.qc[key]}")
                    self._log_ie("LOOCV QC:")
                    for key in (
                        "N", "MAPE_%", "Median_APE_%", "P90_APE_%", "R2_log",
                        "Within_2x_%", "Within_5x_%", "Prediction_interval_fold_P80",
                    ):
                        if key in report.qc:
                            self._log_ie(f"  {key}: {report.qc[key]}")
                    if "Model_rating" in report.qc:
                        self._log_ie(f"Model rating: {report.qc.get('Model_rating')}")
                    if "Model_conclusion" in report.qc:
                        self._log_ie(f"Conclusion: {report.qc.get('Model_conclusion')}")
                    if "Recommendation" in report.qc:
                        self._log_ie(f"Recommendation: {report.qc.get('Recommendation')}")
                self._log_ie(f"Excel: {report.output_xlsx}")
                if report.output_csv:
                    self._log_ie(f"CSV:   {report.output_csv}")
                self.after(
                    0,
                    lambda: messagebox.showinfo(
                        'Complete',
                        f'Local response transfer complete\nCalibration rows: {report.n_train_used}\nTarget predictions: {report.n_predicted}\n\nOutput: {report.output_xlsx}',
                    ),
                )
            except Exception as e:
                msg = str(e)
                self._log_ie("ERROR: " + msg)
                self.after(0, lambda m=msg: messagebox.showerror('Failed', m))
            finally:
                self.after(0, lambda: self.btn_run_ie.configure(state="normal"))
                self.after(0, lambda: self.lbl_status_ie.configure(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()


    # =====================
    # XIC quant tab
    # =====================
    def _build_xic_tab(self):
        from core.ui_support import scroll_body
        frm = scroll_body(self.tab_xic)
        frm.columnconfigure(1, weight=1)
        frm.rowconfigure(99, weight=1)

        r = 0

        ttk.Label(frm, text="Targets CSV\uFF1A").grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_xic_targets_csv = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_xic_targets_csv).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_xic_targets).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='RAW directory:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_xic_raw_dir = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_xic_raw_dir).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_xic_raw_dir).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        ttk.Label(frm, text='Output directory:').grid(row=r, column=0, sticky="w", padx=10, pady=6)
        self.var_xic_out_dir = tk.StringVar()
        ttk.Entry(frm, textvariable=self.var_xic_out_dir).grid(row=r, column=1, sticky="ew", padx=10, pady=6)
        ttk.Button(frm, text='Browse...', command=self._browse_xic_out_dir).grid(row=r, column=2, padx=10, pady=6)
        r += 1

        opt_frame = ttk.Frame(frm)
        opt_frame.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=(0, 6))
        opt_frame.columnconfigure(0, weight=1)

        self.var_xic_recursive = tk.BooleanVar(value=False)
        self.var_xic_only_with_targets = tk.BooleanVar(value=True)
        self.var_xic_use_observed = tk.BooleanVar(value=False)
        self.var_xic_use_internal_std = tk.BooleanVar(value=False)
        self.var_xic_export_ms2_dda = tk.BooleanVar(value=False)
        self.var_xic_combine_formate = tk.BooleanVar(value=True)
        self.var_xic_allow_formate_only = tk.BooleanVar(value=False)
        self.var_xic_formate_include_is = tk.BooleanVar(value=True)
        self.var_xic_is_adduct = tk.StringVar(value=INTERNAL_STANDARD_ADDUCT_LABELS[1])
        self.var_xic_is_ppm = tk.DoubleVar(value=10.0)
        self.var_xic_is_adaptive = tk.BooleanVar(value=True)
        self.var_xic_is_min_snr = tk.DoubleVar(value=5.0)
        self.var_xic_is_min_fraction = tk.DoubleVar(value=0.10)
        self.var_xic_is_expected_rt = tk.StringVar(value="")
        self.var_xic_is_rt_tolerance = tk.DoubleVar(value=0.30)

        ttk.Checkbutton(opt_frame, text='Search subdirectories', variable=self.var_xic_recursive).grid(row=0, column=0, sticky="w")
        ttk.Checkbutton(opt_frame, text='Only process files listed in the CSV', variable=self.var_xic_only_with_targets).grid(
            row=0, column=1, sticky="w", padx=(18, 0)
        )
        ttk.Checkbutton(opt_frame, text='Prefer a matched observed m/z for extraction', variable=self.var_xic_use_observed).grid(
            row=0, column=2, sticky="w", padx=(18, 0)
        )

        ttk.Checkbutton(opt_frame, text='Use the last target row as IS (Area/IS x 100%)', variable=self.var_xic_use_internal_std).grid(
            row=1, column=0, sticky="w"
        )
        ttk.Checkbutton(opt_frame, text='Export MS2 spectra for DDA files', variable=self.var_xic_export_ms2_dda).grid(
            row=1, column=1, sticky="w", padx=(18, 0)
        )
        ttk.Checkbutton(
            opt_frame,
            text='Sum coeluting [M-H]- and formate',
            variable=self.var_xic_combine_formate,
        ).grid(row=2, column=0, columnspan=2, sticky="w")
        ttk.Checkbutton(
            opt_frame,
            text="Allow formate-only peak when [M-H]- is absent (default off)",
            variable=self.var_xic_allow_formate_only,
        ).grid(row=2, column=2, sticky="w", padx=(18, 0))
        ttk.Checkbutton(
            opt_frame,
            text='Apply selected ion panel to IS',
            variable=self.var_xic_formate_include_is,
        ).grid(row=3, column=0, columnspan=2, sticky="w")

        is_fr = ttk.LabelFrame(opt_frame, text='Internal-standard diagnostics')
        is_fr.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(6, 0))
        for cc in range(8):
            is_fr.columnconfigure(cc, weight=1 if cc in (1, 3, 5, 7) else 0)
        ttk.Label(is_fr, text="IS ion").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(is_fr, textvariable=self.var_xic_is_adduct, values=INTERNAL_STANDARD_ADDUCT_LABELS, state="readonly", width=40).grid(row=0, column=1, columnspan=2, sticky="w", padx=6, pady=4)
        ttk.Label(is_fr, text="IS ppm").grid(row=0, column=3, sticky="w", padx=6, pady=4)
        ttk.Entry(is_fr, textvariable=self.var_xic_is_ppm, width=8).grid(row=0, column=4, sticky="w", padx=6, pady=4)
        ttk.Checkbutton(is_fr, text="Adaptive background/SNR rule", variable=self.var_xic_is_adaptive).grid(row=0, column=5, columnspan=3, sticky="w", padx=6, pady=4)
        ttk.Label(is_fr, text="Expected RT (blank=auto-learn)").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(is_fr, textvariable=self.var_xic_is_expected_rt, width=8).grid(row=1, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(is_fr, text="RT tol").grid(row=1, column=2, sticky="w", padx=6, pady=4)
        ttk.Entry(is_fr, textvariable=self.var_xic_is_rt_tolerance, width=8).grid(row=1, column=3, sticky="w", padx=6, pady=4)
        ttk.Label(is_fr, text="SNR proxy").grid(row=1, column=4, sticky="w", padx=6, pady=4)
        ttk.Entry(is_fr, textvariable=self.var_xic_is_min_snr, width=8).grid(row=1, column=5, sticky="w", padx=6, pady=4)
        ttk.Label(is_fr, text="Min height fraction").grid(row=1, column=6, sticky="w", padx=6, pady=4)
        ttk.Entry(is_fr, textvariable=self.var_xic_is_min_fraction, width=8).grid(row=1, column=7, sticky="w", padx=6, pady=4)
        ttk.Label(is_fr, text="A dedicated IS XIC/diagnostic CSV is always written, even when Found=False.").grid(row=2, column=0, columnspan=8, sticky="w", padx=6, pady=4)

        self.var_xic_neg_channel_mode = tk.StringVar(value=NEG_CHANNEL_MODE_LABELS[0])
        self.var_xic_neg_panel_file = tk.StringVar(value="")
        panel_fr = ttk.LabelFrame(opt_frame, text="Extended negative-ion channels")
        panel_fr.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(6, 0))
        panel_fr.columnconfigure(2, weight=1)
        ttk.Label(panel_fr, text="Mode:").grid(row=0, column=0, sticky="w", padx=6, pady=4)
        ttk.Combobox(panel_fr, textvariable=self.var_xic_neg_channel_mode, values=NEG_CHANNEL_MODE_LABELS, state="readonly", width=48).grid(row=0, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(panel_fr, text="Fixed panel CSV/JSON:").grid(row=1, column=0, sticky="w", padx=6, pady=4)
        ttk.Entry(panel_fr, textvariable=self.var_xic_neg_panel_file).grid(row=1, column=1, columnspan=2, sticky="ew", padx=6, pady=4)
        ttk.Button(panel_fr, text="Browse...", command=self._browse_xic_neg_panel).grid(row=1, column=3, padx=6, pady=4)
        ttk.Label(panel_fr, text="Discovery searches all channels without changing Area. A fixed panel sums only channels marked 'sum'.").grid(row=2, column=0, columnspan=4, sticky="w", padx=6, pady=4)
        r += 1

        ttk.Separator(frm).grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)
        r += 1

        params = ttk.LabelFrame(frm, text='XIC settings')
        params.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=6)
        for c in range(6):
            params.columnconfigure(c, weight=1)

        self.var_xic_ppm = tk.DoubleVar(value=5.0)
        self.var_xic_ms_filter = tk.StringVar(value="ms")
        self.var_xic_min_peak_height = tk.DoubleVar(value=1e4)
        self.var_xic_avg_scans = tk.IntVar(value=8)
        self.var_xic_bin_decimals = tk.IntVar(value=4)

        # MS2 plot threshold (relative % of base peak); 0 means no filtering
        self.var_xic_ms2_min_rel = tk.DoubleVar(value=1.0)

        # Multi-peak detection mode for duplicated targets (same m/z)
        self.var_xic_multipeak_mode = tk.StringVar(value=MULTIPEAK_MODE_LABELS[0])

        # For duplicated targets: whether to relax min_peak_height to try to fill multiple peaks
        # (default False: min_peak_height applies to ALL selected peaks)
        self.var_xic_relax_min_peak_height = tk.BooleanVar(value=False)

        # For duplicated targets: relative threshold vs the highest peak (e.g. 0.10 means keep peaks >=10% of max)
        # Set 0 to disable; values >1 will be treated as percent (e.g. 10 -> 0.10).
        self.var_xic_dup_min_rel_height = tk.DoubleVar(value=0.10)
        self.var_xic_formate_rt_tolerance = tk.DoubleVar(value=0.10)
        self.var_xic_formate_min_rel_pct = tk.DoubleVar(value=1.0)
        self.var_xic_formate_min_corr = tk.DoubleVar(value=0.30)

        # adduct matching mode (auto vs force)
        self.var_xic_adduct_mode = tk.StringVar(value=ADDUCT_MODE_LABELS[0])
        self.var_xic_forced_adduct = tk.StringVar(value="M+H")

        ttk.Label(params, text='Extraction tolerance (ppm)').grid(row=0, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_ppm, width=10).grid(row=0, column=1, sticky="w", padx=8, pady=4)

        ttk.Label(params, text="MS Filter").grid(row=0, column=2, sticky="w", padx=8, pady=4)
        ttk.Combobox(params, textvariable=self.var_xic_ms_filter, values=["ms", "ms2"], state="readonly", width=8).grid(
            row=0, column=3, sticky="w", padx=8, pady=4
        )

        ttk.Label(params, text='Minimum peak height').grid(row=0, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_min_peak_height, width=10).grid(
            row=0, column=5, sticky="w", padx=8, pady=4
        )

        ttk.Label(params, text='Ion inference: averaged scans').grid(row=1, column=0, sticky="w", padx=8, pady=4)
        ttk.Spinbox(params, from_=1, to=200, textvariable=self.var_xic_avg_scans, width=8).grid(
            row=1, column=1, sticky="w", padx=8, pady=4
        )
        ttk.Label(params, text='Ion inference: m/z bin decimals').grid(row=1, column=2, sticky="w", padx=8, pady=4)
        ttk.Spinbox(params, from_=1, to=6, textvariable=self.var_xic_bin_decimals, width=8).grid(
            row=1, column=3, sticky="w", padx=8, pady=4
        )

        # Row 2: adduct selection
        xic_adduct_fr = ttk.Frame(params)
        xic_adduct_fr.grid(row=2, column=0, columnspan=6, sticky="w", padx=8, pady=4)

        ttk.Label(xic_adduct_fr, text='Ion matching mode:').pack(side="left")
        self.cmb_xic_adduct_mode = ttk.Combobox(
            xic_adduct_fr,
            textvariable=self.var_xic_adduct_mode,
            values=ADDUCT_MODE_LABELS,
            state="readonly",
            width=28,
        )
        self.cmb_xic_adduct_mode.pack(side="left", padx=(6, 0))
        self.cmb_xic_adduct_mode.bind("<<ComboboxSelected>>", lambda e: self._update_xic_adduct_ui())

        ttk.Label(xic_adduct_fr, text='Specified ion:').pack(side="left", padx=(12, 0))
        self.cmb_xic_forced_adduct = ttk.Combobox(
            xic_adduct_fr,
            textvariable=self.var_xic_forced_adduct,
            values=["M+H", "M-H", "M+HCOO", "M+Na", "M+K", "M+NH4"],
            state="disabled",
            width=10,
        )
        self.cmb_xic_forced_adduct.pack(side="left", padx=(6, 0))


        # Row 3: MS2 threshold (only used when exporting MS2 for DDA files)
        ttk.Label(params, text='MS2 relative-intensity threshold (%, 0 = off)').grid(row=3, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_ms2_min_rel, width=10).grid(row=3, column=1, sticky="w", padx=8, pady=4)

        # Row 4: multi-peak mode for duplicated targets
        ttk.Label(params, text='Duplicate-peak mode').grid(row=4, column=0, sticky="w", padx=8, pady=4)
        ttk.Combobox(
            params,
            textvariable=self.var_xic_multipeak_mode,
            values=MULTIPEAK_MODE_LABELS,
            state="readonly",
            width=32,
        ).grid(row=4, column=1, columnspan=3, sticky="w", padx=8, pady=4)

        ttk.Label(params, text='Duplicate-peak relative height (0-1; 0 = off)').grid(row=5, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_dup_min_rel_height, width=10).grid(row=5, column=1, sticky="w", padx=8, pady=4)

        ttk.Checkbutton(
            params,
            text='Allow a lower height threshold for minor duplicate peaks',
            variable=self.var_xic_relax_min_peak_height,
        ).grid(row=6, column=0, columnspan=6, sticky="w", padx=8, pady=4)

        ttk.Label(params, text="Formate coelution RT tolerance (min)").grid(row=7, column=0, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_formate_rt_tolerance, width=10).grid(row=7, column=1, sticky="w", padx=8, pady=4)
        ttk.Label(params, text="Formate minimum height vs [M-H]- (%)").grid(row=7, column=2, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_formate_min_rel_pct, width=10).grid(row=7, column=3, sticky="w", padx=8, pady=4)
        ttk.Label(params, text="Formate minimum shape correlation").grid(row=7, column=4, sticky="w", padx=8, pady=4)
        ttk.Entry(params, textvariable=self.var_xic_formate_min_corr, width=10).grid(row=7, column=5, sticky="w", padx=8, pady=4)

        r += 1

        btns = ttk.Frame(frm)
        btns.grid(row=r, column=0, columnspan=3, sticky="ew", padx=10, pady=10)

        self.btn_run_xic = ttk.Button(btns, text='Extract and integrate', command=self._on_run_xic)
        self.btn_run_xic.pack(side="left")

        ttk.Button(btns, text='Open output directory', command=self._open_xic_out_dir).pack(side="left", padx=10)

        self.lbl_status_xic = ttk.Label(btns, text="Ready")
        self.lbl_status_xic.pack(side="left", padx=10)

        r += 1

        ttk.Label(frm, text='Log:').grid(row=r, column=0, sticky="nw", padx=10, pady=6)
        self.txt_log_xic = tk.Text(frm, height=18, wrap="word")
        self.txt_log_xic.grid(row=r, column=1, columnspan=2, sticky="nsew", padx=10, pady=6)
        frm.rowconfigure(r, weight=1)

        # init state
        self._update_xic_adduct_ui()

    def _browse_xic_targets(self):
        path = filedialog.askopenfilename(title='Select targets CSV', filetypes=[("CSV", "*.csv"), ("All", "*")])
        if path:
            self.var_xic_targets_csv.set(path)

    def _browse_xic_raw_dir(self):
        path = filedialog.askdirectory(title='Select RAW directory')
        if path:
            self.var_xic_raw_dir.set(path)
            if not self.var_xic_out_dir.get().strip():
                self.var_xic_out_dir.set(path)

    def _browse_xic_out_dir(self):
        path = filedialog.askdirectory(title='Select output directory')
        if path:
            self.var_xic_out_dir.set(path)

    def _browse_xic_neg_panel(self):
        path = filedialog.askopenfilename(title="Select negative-ion panel", filetypes=[("Panel", "*.csv *.json"), ("CSV", "*.csv"), ("JSON", "*.json"), ("All", "*")])
        if path:
            self.var_xic_neg_panel_file.set(path)

    def _open_xic_out_dir(self):
        p = self.var_xic_out_dir.get().strip()
        if p:
            open_in_os(Path(p))

    def _on_run_xic(self):
        csv_path = self.var_xic_targets_csv.get().strip()
        raw_dir = self.var_xic_raw_dir.get().strip()
        out_dir = self.var_xic_out_dir.get().strip()

        if not csv_path:
            messagebox.showerror('Error', 'Select a targets CSV file.')
            return
        if not raw_dir:
            messagebox.showerror('Error', 'Select a RAW directory.')
            return
        if not out_dir:
            messagebox.showerror('Error', 'Select an output directory.')
            return

        try:
            ppm = float(self.var_xic_ppm.get())
            ms_filter = self.var_xic_ms_filter.get().strip() or "ms"
            min_peak_height = float(self.var_xic_min_peak_height.get())
            avg_scans = int(self.var_xic_avg_scans.get())
            bin_decimals = int(self.var_xic_bin_decimals.get())
            use_observed = bool(self.var_xic_use_observed.get())
            use_internal_std = bool(self.var_xic_use_internal_std.get())
            export_ms2_dda = bool(self.var_xic_export_ms2_dda.get())
            ms2_min_rel = float(self.var_xic_ms2_min_rel.get())
            relax_min_height = bool(self.var_xic_relax_min_peak_height.get())
            dup_min_rel_height = float(self.var_xic_dup_min_rel_height.get())
            combine_formate = bool(self.var_xic_combine_formate.get())
            allow_formate_only = bool(self.var_xic_allow_formate_only.get())
            formate_include_is = bool(self.var_xic_formate_include_is.get())
            formate_rt_tolerance = float(self.var_xic_formate_rt_tolerance.get())
            formate_min_rel_pct = float(self.var_xic_formate_min_rel_pct.get())
            formate_min_corr = float(self.var_xic_formate_min_corr.get())
            neg_channel_mode = NEG_CHANNEL_MODE_MAP.get(self.var_xic_neg_channel_mode.get(), "legacy")
            neg_panel_file = self.var_xic_neg_panel_file.get().strip()
            xic_is_adduct = INTERNAL_STANDARD_ADDUCT_MAP.get(self.var_xic_is_adduct.get(), "auto_raw")
            xic_is_ppm = float(self.var_xic_is_ppm.get())
            xic_is_adaptive = bool(self.var_xic_is_adaptive.get())
            xic_is_min_snr = float(self.var_xic_is_min_snr.get())
            xic_is_min_fraction = float(self.var_xic_is_min_fraction.get())
            xic_is_rt_tolerance = float(self.var_xic_is_rt_tolerance.get())
            xic_is_rt_text = self.var_xic_is_expected_rt.get().strip()
            xic_is_expected_rt = float(xic_is_rt_text) if xic_is_rt_text else None
            if xic_is_ppm <= 0 or xic_is_min_snr < 0 or not (0 <= xic_is_min_fraction <= 1) or xic_is_rt_tolerance <= 0:
                raise ValueError("Invalid internal-standard diagnostic settings")
        except Exception as e:
            messagebox.showerror('Invalid settings', str(e))
            return

        self.btn_run_xic.config(state="disabled")
        self.lbl_status_xic.config(text="Running...")
        self.txt_log_xic.delete("1.0", "end")

        def worker():
            try:
                mapping = load_targets_mapping(Path(csv_path))
                self._log_xic(f"Targets loaded: {sum(len(v) for v in mapping.values())} rows, {len(mapping)} files")
                all_targets = list(mapping.get("all", []) or [])
                if all_targets:
                    self._log_xic(f"Global targets detected (raw=all): {len(all_targets)} rows -> apply to ALL RAW files")

                raw_paths = _find_raw_paths(Path(raw_dir), recursive=bool(self.var_xic_recursive.get()))
                if not raw_paths:
                    raise RuntimeError('No .raw files were found in the selected directory.')
                self._log_xic(f"RAW files found: {len(raw_paths)}")

                out_root = Path(out_dir)
                self._last_out_dir = out_root

                processed = 0
                skipped = 0
                failed = 0
                learned_internal_standard_rts: List[float] = []

                for idx, rp in enumerate(raw_paths, start=1):
                    stem = safe_stem_from_raw_path(rp)
                    key = stem.strip().lower()
                    tlist = list(mapping.get(key, []) or [])
                    targets = tlist + list(all_targets)  # raw=all targets appended at end
                    if not targets and bool(self.var_xic_only_with_targets.get()):
                        skipped += 1
                        self._log_xic(f"[{idx}/{len(raw_paths)}] SKIP (no targets): {stem}")
                        continue
                    try:
                        out_csv, plots_dir = quant_xic_for_raw(
                            rp,
                            targets,
                            out_root,
                            ppm=ppm,
                            ms_filter=ms_filter,
                            avg_scans=avg_scans,
                            bin_decimals=bin_decimals,
                            min_peak_height=min_peak_height,
                            use_observed_mz=use_observed,
                            use_last_as_internal_standard=use_internal_std,
                            internal_standard_adduct=xic_is_adduct,
                            internal_standard_ppm=xic_is_ppm,
                            internal_standard_adaptive_detection=xic_is_adaptive,
                            internal_standard_min_snr_proxy=xic_is_min_snr,
                            internal_standard_min_height_fraction=xic_is_min_fraction,
                            internal_standard_expected_rt=(
                                xic_is_expected_rt
                                if xic_is_expected_rt is not None
                                else (
                                    float(__import__("numpy").median(learned_internal_standard_rts))
                                    if learned_internal_standard_rts else None
                                )
                            ),
                            internal_standard_rt_tolerance_min=xic_is_rt_tolerance,
                            export_ms2_if_dda=export_ms2_dda,
                            ms2_min_rel=ms2_min_rel,
                            relax_min_peak_height_for_duplicates=relax_min_height,
                            dup_min_rel_height=dup_min_rel_height,
                            adduct_mode=ADDUCT_MODE_MAP.get((self.var_xic_adduct_mode.get() or "").strip(), "posneg_neghonly"),
                            forced_adduct=(self.var_xic_forced_adduct.get() or ""),
                            multi_peak_mode=MULTIPEAK_MODE_MAP.get(
                                (self.var_xic_multipeak_mode.get() or "").strip(),
                                "auto",
                            ),
                            combine_negative_formate=combine_formate,
                            formate_rt_tolerance_min=formate_rt_tolerance,
                            formate_min_rel_height_pct=formate_min_rel_pct,
                            formate_min_shape_correlation=formate_min_corr,
                            allow_formate_only=allow_formate_only,
                            include_formate_for_internal_standard=formate_include_is,
                            negative_channel_mode=neg_channel_mode,
                            negative_channel_panel_file=neg_panel_file,
                            negative_channel_rt_tolerance_min=formate_rt_tolerance,
                            negative_channel_min_rel_height_pct=formate_min_rel_pct,
                            negative_channel_min_shape_correlation=formate_min_corr,
                        )
                        try:
                            qrows = read_quant_csv(out_csv)
                            isrows = [q for q in qrows if str(q.get("Is_internal_standard", "")).strip().lower() in {'1', 'true', 'yes', 'y', 'Yes', '是'}]
                            if isrows:
                                isr = isrows[-1]
                                try:
                                    isrt = float(isr.get("Apex_RT", ""))
                                    if str(isr.get("Found", "")).strip().lower() in {'1', 'true', 'yes', 'y', 'Yes', '是'} and __import__("math").isfinite(isrt):
                                        learned_internal_standard_rts.append(isrt)
                                except Exception:
                                    pass
                                self._log_xic(
                                    "  IS diagnostic: "
                                    f"status={isr.get('IS_Diagnostic_Status', '')}, found={isr.get('Found', '')}, "
                                    f"adduct={isr.get('Matched_Adduct', isr.get('Adduct', ''))}, RT={isr.get('Apex_RT', '')}, "
                                    f"area={isr.get('Area', '')}, SNR={isr.get('IS_SNR_Proxy', '')}, "
                                    f"plot={isr.get('IS_Diagnostic_PNG', isr.get('XIC_PNG', ''))}"
                                )
                        except Exception as diag_exc:
                            self._log_xic(f"  WARN: could not read IS diagnostics: {diag_exc}")
                        processed += 1
                        self._log_xic(f"[{idx}/{len(raw_paths)}] OK: {stem} -> {Path(out_csv).name} | {plots_dir.name}")
                    except Exception as e:
                        failed += 1
                        self._log_xic(f"[{idx}/{len(raw_paths)}] FAIL: {stem} -> {e}")

                self._log_xic("\n=== Summary ===")
                self._log_xic(f"Processed: {processed}")
                self._log_xic(f"Skipped:   {skipped}")
                self._log_xic(f"Failed:    {failed}")
                self.after(0, lambda: messagebox.showinfo('Complete', f'XIC extraction complete\nProcessed={processed}, Skipped={skipped}, Failed={failed}'))

            except Exception as e:
                self._log_xic(f"ERROR: {e}")
                self.after(0, lambda: messagebox.showerror('Failed', str(e)))
            finally:
                self.after(0, lambda: self.btn_run_xic.config(state="normal"))
                self.after(0, lambda: self.lbl_status_xic.config(text="Ready"))

        threading.Thread(target=worker, daemon=True).start()

    # =====================
    # Deuteration tab
    # =====================






    # =====================
    # Review tab
    # =====================






if __name__ == "__main__":
    app = ThermoBatchReportApp()
    app.mainloop()
