"""CombiTrace-IE private structure-aware local response-factor transfer model (v18.26).

This module replaces the previous global regression / hierarchical component model
with a local response-factor migration strategy.  It is intended for a small
calibration set (known concentration + measured area/internal-standard ratio) and
a larger candidate table addressed by Combo strings.

Core idea:
    RF = actual_concentration / measured_ratio
    predicted_concentration = measured_ratio * transferred_RF

For each target, the transferred RF is borrowed from calibration rows by a strict
evidence hierarchy:
    1. exact ordered A/B/C component-Formula combination
    2. identical locally generated product structure
    3. high product/fragment/core structural similarity
    4. same A+B / A+C / B+C component pair
    5. nearest structural neighbours
    6. same A/B/C single component(s)
    7. formula/mass/RT neighbours
    8. global median RF fallback

The output workbook prints all used/skipped training and target row positions,
all prediction evidence, and empirical LOOCV error by evidence level.
"""
from __future__ import annotations

import csv
import math
import re
import statistics
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np

from .chemistry import parse_formula
from .structure_features import (
    StructureLibrary,
    audit_row as structure_audit_row,
    enrich_rows as enrich_rows_with_structures,
    load_structure_library as load_private_structure_library,
    same_core,
    structure_similarity,
)
from .response_predictor import (
    PreparedRow,
    RunReport,
    _as_float_or_none,
    _metrics,
    _norm_header,
    guess_columns,
    parse_number,
    parse_combo,
    prepare_rows,
    read_table,
)


# -----------------------------------------------------------------------------
# Local/private structure support
# -----------------------------------------------------------------------------
# Implemented in core.structure_features.  Confidential SMILES and private
# reaction SMARTS remain on the user's local computer.


# -----------------------------------------------------------------------------
# Optional Combo -> Product_SMILES master table
# -----------------------------------------------------------------------------


@dataclass(frozen=True)
class ProductMasterEntry:
    """One Combo-SMILES master row.

    The authoritative join identity is the ordered triplet of A/B/C component
    molecular formulas.  Component IDs and the numbers after A#/B#/C# are kept
    only for human-readable diagnostics because they may change between tables.
    """

    smiles: str
    combo: str
    formula_key: str
    exact_key: str
    ordinal_key: str = ""
    expected_formula: str = ""
    product_formula: str = ""
    a_formula: str = ""
    b_formula: str = ""
    c_formula: str = ""
    source_file: str = ""
    source_sheet: str = ""
    source_row: object = ""


def _clean_combo_text(text: object) -> str:
    """Normalize harmless formatting differences without exposing structure data."""
    s = unicodedata.normalize("NFKC", str(text or ""))
    s = s.replace("\\|", "|").replace("\\:", ":")
    s = s.replace("\uFF5C", "|").replace("\u2223", "|").replace("\uFF1A", ":")
    s = re.sub(r"[\u200b-\u200d\ufeff]", "", s)
    return s.strip()


def _combo_exact_key(text: object) -> str:
    """Legacy display key retained only for diagnostics; it is not used to join."""
    s = _clean_combo_text(text)
    if not s:
        return ""
    return "combo|" + re.sub(r"\s+", "", s).lower()


def _norm_index(value: object) -> str:
    """Normalize A#/B#/C# numbers for diagnostics only."""
    s = str(value or "").strip()
    if not s:
        return ""
    try:
        return str(int(s))
    except Exception:
        return s.lower()


def _combo_index_key(text: object) -> str:
    """Legacy A#/B#/C# key retained for audit output; never used for matching."""
    parts = parse_combo(_clean_combo_text(text))
    vals = [_norm_index(parts.get(role, {}).get("index", "")) for role in ("A", "B", "C")]
    if not all(vals):
        return ""
    return "abc_index|" + "|".join(
        f"{role.lower()}={value}" for role, value in zip(("A", "B", "C"), vals)
    )


def _combo_component_formulas(text: object) -> Dict[str, str]:
    parts = parse_combo(_clean_combo_text(text))
    return {role: str(parts.get(role, {}).get("formula", "") or "").strip() for role in ("A", "B", "C")}


def _formula_signature(value: object) -> Optional[Tuple[Tuple[str, int], ...]]:
    s = str(value or "").strip()
    if not s:
        return None
    try:
        counts = parse_formula(s)
        return tuple(sorted((str(k), int(v)) for k, v in counts.items() if int(v) != 0))
    except Exception:
        return None


def _formula_key_token(value: object) -> str:
    """Canonical, parse-based formula token used in the structure join key."""
    sig = _formula_signature(value)
    if sig is None:
        return ""
    # Explicit counts avoid ambiguity and make CH3O / H3CO equivalent.
    return ",".join(f"{element}={count}" for element, count in sig)


def _abc_formula_key(a_formula: object, b_formula: object, c_formula: object) -> str:
    tokens = [_formula_key_token(x) for x in (a_formula, b_formula, c_formula)]
    if not all(tokens):
        return ""
    return "abc_formula|" + "|".join(
        f"{role.lower()}={token}" for role, token in zip(("A", "B", "C"), tokens)
    )


def _combo_formula_key(text: object) -> str:
    formulas = _combo_component_formulas(text)
    return _abc_formula_key(formulas.get("A", ""), formulas.get("B", ""), formulas.get("C", ""))


def _formula_compare(a: object, b: object) -> Optional[bool]:
    """Return True/False when both formulas are present, otherwise None."""
    sa = str(a or "").strip()
    sb = str(b or "").strip()
    if not sa or not sb:
        return None
    ka = _formula_signature(sa)
    kb = _formula_signature(sb)
    if ka is not None and kb is not None:
        return ka == kb
    na = re.sub(r"\s+", "", unicodedata.normalize("NFKC", sa))
    nb = re.sub(r"\s+", "", unicodedata.normalize("NFKC", sb))
    return na == nb


def _component_formula_check(target_combo: object, entry: ProductMasterEntry) -> Tuple[str, str]:
    target = _combo_component_formulas(target_combo)
    master = {"A": entry.a_formula, "B": entry.b_formula, "C": entry.c_formula}
    compared: List[str] = []
    mismatched: List[str] = []
    unavailable: List[str] = []
    for role in ("A", "B", "C"):
        result = _formula_compare(target.get(role, ""), master.get(role, ""))
        if result is True:
            compared.append(role)
        elif result is False:
            mismatched.append(role)
        else:
            unavailable.append(role)
    if mismatched:
        return "no", "component Formula mismatch: " + ", ".join(mismatched)
    if len(compared) == 3:
        return "yes", "A/B/C component formulas all agree"
    if compared:
        return "partial", f"component Formula agrees for {','.join(compared)}; unavailable for {','.join(unavailable)}"
    return "not_available", "component Formula unavailable for verification"


def _product_formula_check(target_formula: object, entry: ProductMasterEntry) -> Tuple[str, str, str]:
    """Compare analytical product Formula with the Product-SMILES master Formula."""
    master_formula = entry.product_formula or entry.expected_formula
    result = _formula_compare(target_formula, master_formula)
    if result is True:
        return "yes", master_formula, "target Formula agrees with master product Formula"
    if result is False:
        return "no", master_formula, f"target Formula={target_formula} differs from master Formula={master_formula}"
    return "not_available", master_formula, "target or master product Formula unavailable for verification"


def _master_internal_formula_check(entry: ProductMasterEntry) -> Tuple[str, str]:
    result = _formula_compare(entry.expected_formula, entry.product_formula)
    if result is True:
        return "yes", "master Expected Formula agrees with Product_SMILES Formula"
    if result is False:
        return "no", (
            f"master Expected Formula={entry.expected_formula} differs from "
            f"Product_SMILES Formula={entry.product_formula}"
        )
    return "not_available", "master internal Formula comparison unavailable"


def _guess_master_column(headers: Sequence[str], candidates: Sequence[str]) -> str:
    norm = {_norm_header(h): h for h in headers}
    for candidate in candidates:
        key = _norm_header(candidate)
        if key in norm:
            return norm[key]
    for h in headers:
        nh = _norm_header(h)
        for candidate in candidates:
            key = _norm_header(candidate)
            if key and key in nh:
                return h
    return ""


def load_product_smiles_master(
    master_file: Optional[Path],
    *,
    sheet_name: str = "",
    combo_col: str = "",
    smiles_col: str = "",
) -> Tuple[Dict[str, List[ProductMasterEntry]], List[str], Dict[str, object]]:
    """Load Product SMILES keyed exclusively by A/B/C component formulas."""
    if not master_file:
        return {}, [], {"rows": 0, "usable": 0, "sheet": ""}
    master_file = Path(master_file)
    if not master_file.exists():
        raise FileNotFoundError(str(master_file))
    table = read_table(master_file, sheet_name=sheet_name)
    ccol = combo_col or _guess_master_column(table.headers, ["Combo", 'Combo', 'Source_Combination', "Source Combo"])
    scol = smiles_col or _guess_master_column(
        table.headers,
        ["Product_SMILES", "Selected_Product_SMILES", "Product SMILES", 'Product_SMILES'],
    )
    expected_formula_col = _guess_master_column(
        table.headers,
        ["Formula", "Expected_Formula", "Expected Formula", 'Enumerated_Formula', 'Formula'],
    )
    product_formula_col = _guess_master_column(
        table.headers,
        ["Product_Formula", "Product Formula", 'Product_Formula', 'Product_Formula'],
    )
    a_formula_col = _guess_master_column(table.headers, ["A_Formula", "A Formula", 'A_Formula'])
    b_formula_col = _guess_master_column(table.headers, ["B_Formula", "B Formula", 'B_Formula'])
    c_formula_col = _guess_master_column(table.headers, ["C_Formula", "C Formula", 'C_Formula'])
    if not scol:
        raise ValueError("Product_SMILES column could not be detected in the Combo-SMILES master table")
    if not ccol and not (a_formula_col and b_formula_col and c_formula_col):
        raise ValueError(
            "The master table needs either a Combo column containing A/B/C formulas or explicit "
            "A_Formula, B_Formula and C_Formula columns"
        )

    mapping: Dict[str, List[ProductMasterEntry]] = {}
    warnings: List[str] = []
    blank_smiles = 0
    missing_formula_key = 0
    usable_rows = 0
    formula_keys = 0
    duplicate_formula_keys = 0

    def add_key(key: str, entry: ProductMasterEntry) -> None:
        nonlocal formula_keys, duplicate_formula_keys
        if not key:
            return
        bucket = mapping.setdefault(key, [])
        if bucket and not any(x.smiles == entry.smiles and x.combo == entry.combo for x in bucket):
            duplicate_formula_keys += 1
        if not any(x.smiles == entry.smiles and x.combo == entry.combo for x in bucket):
            bucket.append(entry)
        if len(bucket) == 1:
            formula_keys += 1

    for raw in table.rows:
        combo = str(raw.get(ccol, "") or "").strip() if ccol else ""
        smiles = str(raw.get(scol, "") or "").strip()
        if not smiles:
            blank_smiles += 1
            continue
        parts = parse_combo(_clean_combo_text(combo)) if combo else {}
        a_formula = (
            str(raw.get(a_formula_col, "") or "").strip()
            if a_formula_col else str(parts.get("A", {}).get("formula", "") or "").strip()
        )
        b_formula = (
            str(raw.get(b_formula_col, "") or "").strip()
            if b_formula_col else str(parts.get("B", {}).get("formula", "") or "").strip()
        )
        c_formula = (
            str(raw.get(c_formula_col, "") or "").strip()
            if c_formula_col else str(parts.get("C", {}).get("formula", "") or "").strip()
        )
        formula_key = _abc_formula_key(a_formula, b_formula, c_formula)
        if not formula_key:
            missing_formula_key += 1
            continue
        entry = ProductMasterEntry(
            smiles=smiles,
            combo=combo,
            formula_key=formula_key,
            exact_key=_combo_exact_key(combo),
            ordinal_key=_combo_index_key(combo),
            expected_formula=str(raw.get(expected_formula_col, "") or "").strip() if expected_formula_col else "",
            product_formula=str(raw.get(product_formula_col, "") or "").strip() if product_formula_col else "",
            a_formula=a_formula,
            b_formula=b_formula,
            c_formula=c_formula,
            source_file=str(raw.get("__source_file__", master_file.name) or master_file.name),
            source_sheet=str(raw.get("__source_sheet__", table.sheet_name or "") or ""),
            source_row=raw.get("__source_row__", ""),
        )
        usable_rows += 1
        add_key(formula_key, entry)

    if blank_smiles:
        warnings.append(f"Combo-SMILES master skipped {blank_smiles} rows with blank Product_SMILES")
    if missing_formula_key:
        warnings.append(
            f"Combo-SMILES master skipped {missing_formula_key} rows because A/B/C component formulas were missing or invalid"
        )
    if duplicate_formula_keys:
        warnings.append(
            f"Combo-SMILES master contains {duplicate_formula_keys} repeated A/B/C Formula keys; "
            "product Formula and identical-SMILES checks will prevent ambiguous attachment"
        )
    return mapping, warnings, {
        "rows": len(table.rows),
        "usable": usable_rows,
        "sheet": table.sheet_name or "",
        "combo_col": ccol,
        "smiles_col": scol,
        "expected_formula_col": expected_formula_col,
        "product_formula_col": product_formula_col,
        "a_formula_col": a_formula_col,
        "b_formula_col": b_formula_col,
        "c_formula_col": c_formula_col,
        "formula_keys": formula_keys,
        "lookup_keys": len(mapping),
        "duplicate_formula_keys": duplicate_formula_keys,
        "blank_smiles_rows": blank_smiles,
        "missing_formula_key_rows": missing_formula_key,
    }


def _select_master_entry(
    row: PreparedRow,
    candidates: Sequence[ProductMasterEntry],
    *,
    formula_verification_mode: str = "review",
) -> Tuple[Optional[ProductMasterEntry], Dict[str, str]]:
    """Select one Product-SMILES row after an A/B/C-Formula-key match.

    Component formulas are the identity and therefore must already agree.  The
    analytical product Formula and the Product-SMILES-derived Formula remain a
    review/strict QC guardrail.  If a formula key unexpectedly maps to multiple
    different structures, no arbitrary first-row selection is made.
    """
    mode = str(formula_verification_mode or "review").strip().lower()
    if mode not in {"review", "strict"}:
        mode = "review"

    diagnostics = {
        "product_check": "not_available",
        "component_check": "not_available",
        "internal_check": "not_available",
        "master_formula": "",
        "master_product_formula": "",
        "join_status": "unmatched",
        "formula_flag": "No",
        "reason": "",
    }
    if not candidates:
        diagnostics["reason"] = "no master candidate for A/B/C component Formula key"
        return None, diagnostics

    evaluated: List[Tuple[int, ProductMasterEntry, Dict[str, str], bool]] = []
    for entry in candidates:
        product_check, master_formula, product_note = _product_formula_check(row.formula, entry)
        component_check, component_note = _component_formula_check(row.combo, entry)
        internal_check, internal_note = _master_internal_formula_check(entry)
        mismatch = any(x == "no" for x in (product_check, component_check, internal_check))
        diag = {
            "product_check": product_check,
            "component_check": component_check,
            "internal_check": internal_check,
            "master_formula": master_formula,
            "master_product_formula": entry.product_formula,
            "join_status": "candidate",
            "formula_flag": "Yes" if mismatch else "No",
            "reason": "; ".join([product_note, component_note, internal_note]),
        }
        score = 0
        if product_check == "yes":
            score += 8
        elif product_check == "not_available":
            score += 1
        if component_check == "yes":
            score += 6
        elif component_check == "partial":
            score += 2
        if internal_check == "yes":
            score += 1
        evaluated.append((score, entry, diag, mismatch))

    if len(evaluated) == 1:
        _score, entry, diag, mismatch = evaluated[0]
        if mode == "strict" and mismatch:
            diag["join_status"] = "rejected_product_formula_mismatch"
            diag["reason"] = "A/B/C Formula key matched, but strict product Formula verification rejected it: " + diag["reason"]
            return None, diag
        diag["join_status"] = "matched_formula_warning" if mismatch else "matched_verified"
        if mismatch:
            diag["reason"] = (
                "A/B/C component Formula key matched uniquely; Product_SMILES was retained in review mode. "
                + diag["reason"]
            )
        return entry, diag

    compatible = [x for x in evaluated if not x[3]]
    ranked = compatible if compatible else []
    if mode == "strict" and not ranked:
        diagnostics.update(evaluated[0][2])
        diagnostics["join_status"] = "rejected_product_formula_mismatch"
        diagnostics["reason"] = (
            "A/B/C Formula key found, but strict Formula verification rejected all candidates: "
            + " | ".join(x[2]["reason"] for x in evaluated[:5])
        )
        return None, diagnostics

    if not ranked:
        unique_smiles = {x[1].smiles for x in evaluated}
        if len(unique_smiles) == 1 and mode == "review":
            best = sorted(evaluated, key=lambda x: x[0], reverse=True)[0]
            diag = dict(best[2])
            diag["join_status"] = "matched_formula_warning"
            diag["reason"] = (
                "Repeated A/B/C Formula-key rows all contain the same Product_SMILES; retained with Formula warning. "
                + diag["reason"]
            )
            return best[1], diag
        diagnostics.update(evaluated[0][2])
        diagnostics["join_status"] = "ambiguous_formula_key"
        diagnostics["reason"] = (
            f"ambiguous A/B/C Formula key: {len(evaluated)} candidates have different Product_SMILES "
            "and product Formula did not disambiguate them"
        )
        return None, diagnostics

    ranked.sort(key=lambda x: x[0], reverse=True)
    best_score = ranked[0][0]
    best = [x for x in ranked if x[0] == best_score]
    unique_smiles = {x[1].smiles for x in best}
    if len(unique_smiles) > 1:
        diagnostics.update(best[0][2])
        diagnostics["join_status"] = "ambiguous_formula_key"
        diagnostics["reason"] = (
            f"ambiguous A/B/C Formula key: {len(best)} product-Formula-compatible candidates "
            "have different Product_SMILES"
        )
        return None, diagnostics

    diag = dict(best[0][2])
    diag["join_status"] = "matched_verified"
    return best[0][1], diag


def apply_product_smiles_master(
    rows: Sequence[PreparedRow],
    mapping: Dict[str, List[ProductMasterEntry]],
    *,
    formula_verification_mode: str = "review",
) -> Tuple[int, int]:
    """Attach Product_SMILES using only the A/B/C component Formula triplet."""
    matched = 0
    filled = 0
    candidate_headers = {
        _norm_header("Product_SMILES"),
        _norm_header("Selected_Product_SMILES"),
        _norm_header("Product SMILES"),
        _norm_header('Product_SMILES'),
    }
    for row in rows:
        formula_key = _combo_formula_key(row.combo)
        exact_key = _combo_exact_key(row.combo)
        index_key = _combo_index_key(row.combo)
        candidates = list(mapping.get(formula_key, [])) if formula_key else []
        method = "abc_component_formulas" if candidates else "unmatched"
        used_key = formula_key if candidates else ""

        entry, diag = _select_master_entry(row, candidates, formula_verification_mode=formula_verification_mode)
        smi = entry.smiles if entry is not None else ""

        row.raw["__product_master_match__"] = "Yes" if smi else "No"
        row.raw["__product_master_match_method__"] = method
        row.raw["__product_master_exact_key__"] = exact_key
        row.raw["__product_master_index_key__"] = index_key  # diagnostic only
        row.raw["__product_master_formula_key__"] = formula_key
        row.raw["__product_master_abc_key__"] = formula_key
        row.raw["__product_master_used_key__"] = used_key if smi else ""
        row.raw["__product_master_candidate_count__"] = len(candidates)
        row.raw["__product_master_verification_mode__"] = str(formula_verification_mode or "review")
        row.raw["__product_master_join_status__"] = diag.get("join_status", "")
        row.raw["__product_master_formula_flag__"] = diag.get("formula_flag", "No")
        row.raw["__product_master_formula_check__"] = diag.get("product_check", "")
        row.raw["__product_master_component_formula_check__"] = diag.get("component_check", "")
        row.raw["__product_master_internal_formula_check__"] = diag.get("internal_check", "")
        row.raw["__product_master_target_formula__"] = row.formula
        row.raw["__product_master_master_formula__"] = diag.get("master_formula", "")
        row.raw["__product_master_master_product_formula__"] = diag.get("master_product_formula", "")
        row.raw["__product_master_formula_note__"] = diag.get("reason", "")

        if not smi:
            if not row.combo:
                reason = "blank Combo"
            elif not formula_key:
                reason = "Combo is missing a valid A, B or C component Formula"
            elif candidates:
                reason = diag.get("reason", "product Formula verification failed or Formula-key candidates are ambiguous")
            else:
                reason = "no matching A/B/C component Formula key found in master"
            row.raw["__product_master_unmatched_reason__"] = reason
            continue

        row.raw["__product_master_unmatched_reason__"] = ""
        matched += 1
        existing = ""
        for k, v in row.raw.items():
            if _norm_header(str(k)) in candidate_headers and str(v or "").strip():
                existing = str(v).strip()
                break
        if not existing:
            row.raw["Product_SMILES"] = smi
            filled += 1
    return matched, filled


# -----------------------------------------------------------------------------
# Row helpers
# -----------------------------------------------------------------------------


def _row_formula_flagged(row: PreparedRow) -> bool:
    return any(
        str(row.raw.get(k, "") or "").strip().lower() == "no"
        for k in (
            "__product_master_formula_check__",
            "__product_master_component_formula_check__",
            "__product_master_internal_formula_check__",
        )
    )


def _join_component_summary(rows: Sequence[PreparedRow], dataset: str) -> List[Dict[str, object]]:
    """Aggregate structure-join QC by A/B/C component Formula.

    Because each component library is formula-unique, a single incorrect or
    missing formula normally affects a complete 10\u00D710 slice.  Grouping by the
    formula itself exposes that pattern without relying on mutable IDs or row
    numbers and without exposing confidential SMILES.
    """
    groups: Dict[Tuple[str, str], Dict[str, object]] = {}
    for row in rows:
        for role in ("A", "B", "C"):
            raw_formula = _component_formula(row, role)
            token = _formula_key_token(raw_formula)
            if not token:
                continue
            key = (role, token)
            g = groups.setdefault(key, {
                "Dataset": dataset,
                "Role": role,
                "Component_Formula": raw_formula,
                "Formula_Key_Token": token,
                "Rows": 0,
                "Formula_key_candidate_found": 0,
                "Structure_join_matched": 0,
                "Formula_flagged": 0,
                "Product_formula_no": 0,
                "Component_formula_no": 0,
                "Master_internal_formula_no": 0,
                "Product_structure_ready": 0,
                "Unmatched_reasons": {},
            })
            g["Rows"] = int(g["Rows"]) + 1
            if int(row.raw.get("__product_master_candidate_count__", 0) or 0) > 0:
                g["Formula_key_candidate_found"] = int(g["Formula_key_candidate_found"]) + 1
            if str(row.raw.get("__product_master_match__", "")).lower() == "yes":
                g["Structure_join_matched"] = int(g["Structure_join_matched"]) + 1
            if _row_formula_flagged(row):
                g["Formula_flagged"] = int(g["Formula_flagged"]) + 1
            if str(row.raw.get("__product_master_formula_check__", "")).lower() == "no":
                g["Product_formula_no"] = int(g["Product_formula_no"]) + 1
            if str(row.raw.get("__product_master_component_formula_check__", "")).lower() == "no":
                g["Component_formula_no"] = int(g["Component_formula_no"]) + 1
            if str(row.raw.get("__product_master_internal_formula_check__", "")).lower() == "no":
                g["Master_internal_formula_no"] = int(g["Master_internal_formula_no"]) + 1
            if getattr(row, "structure_status", "") == "product_structure_ready":
                g["Product_structure_ready"] = int(g["Product_structure_ready"]) + 1
            reason = str(row.raw.get("__product_master_unmatched_reason__", "") or "").strip()
            if reason:
                reasons = g["Unmatched_reasons"]
                reasons[reason] = reasons.get(reason, 0) + 1

    out: List[Dict[str, object]] = []
    max_group_rows = max((int(g["Rows"]) for g in groups.values()), default=0)
    pattern_min_rows = max(10, int(math.ceil(max_group_rows * 0.50))) if max_group_rows else 10
    for (_role, _token), g in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        n = max(1, int(g["Rows"]))
        g["Match_rate_%"] = 100.0 * int(g["Structure_join_matched"]) / n
        g["Formula_flag_rate_%"] = 100.0 * int(g["Formula_flagged"]) / n
        reasons = g.pop("Unmatched_reasons")
        g["Top_unmatched_reason"] = max(reasons.items(), key=lambda x: x[1])[0] if reasons else ""
        if int(g["Structure_join_matched"]) == 0 and int(g["Rows"]) >= pattern_min_rows:
            g["Likely_issue"] = (
                f"No rows containing {g['Role']} component Formula {g['Component_Formula']} obtained Product_SMILES; "
                "check that the same Formula is present in the Combo-SMILES master."
            )
        elif int(g["Formula_flagged"]) == int(g["Rows"]) and int(g["Rows"]) >= pattern_min_rows:
            g["Likely_issue"] = (
                f"All rows containing {g['Role']} component Formula {g['Component_Formula']} are product-Formula flagged; "
                "inspect the generated Product_SMILES and final Formula."
            )
        else:
            g["Likely_issue"] = ""
        out.append(g)
    return out


def _source_row_no(row: PreparedRow) -> object:
    return row.raw.get("__source_row__", row.source_index)


def _source_sheet(row: PreparedRow) -> str:
    return str(row.raw.get("__source_sheet__", "") or "")


def _source_file(row: PreparedRow) -> str:
    return str(row.raw.get("__source_file__", "") or "")


def _is_internal_standard_row(row: PreparedRow) -> bool:
    combo = str(row.combo or "").strip().lower()
    name = str(row.name or "").strip().lower()
    formula = str(row.formula or "").strip().lower()
    no = str(row.raw.get("No", row.raw.get("no", "")) or "").strip().lower()
    text = " ".join([combo, name, formula, no])
    return (
        combo in {'internal standard', 'is', 'internal standard', '内标'}
        or no == "is"
        or "\u5185\u6807" in text
        or "internal standard" in text
        or 'internal standard' in text
    )


def valid_training_rows(rows: Sequence[PreparedRow]) -> List[PreparedRow]:
    out: List[PreparedRow] = []
    for r in rows:
        if _is_internal_standard_row(r):
            continue
        if (
            r.ratio is not None
            and r.ratio > 0
            and r.concentration is not None
            and r.concentration > 0
            and r.log_correction_factor is not None
            and math.isfinite(float(r.log_correction_factor))
        ):
            out.append(r)
    return out


def training_skip_reason(row: PreparedRow) -> str:
    reasons: List[str] = []
    if _is_internal_standard_row(row):
        reasons.append("internal standard row; not used for response model")
    if row.ratio is None:
        reasons.append("missing measured ratio")
    elif row.ratio <= 0:
        reasons.append("measured ratio <= 0")
    if row.concentration is None:
        reasons.append("missing actual concentration")
    elif row.concentration <= 0:
        reasons.append("actual concentration <= 0")
    if row.log_correction_factor is None and not reasons:
        reasons.append("cannot compute logRF = ln(concentration/ratio)")
    return "; ".join(reasons)


def target_skip_reason(row: PreparedRow, pred: Optional["LocalRFPrediction"] = None) -> str:
    reasons: List[str] = []
    if _is_internal_standard_row(row):
        reasons.append("internal standard row; excluded from concentration prediction")
    if row.ratio is None:
        reasons.append("missing measured ratio")
    elif row.ratio <= 0:
        reasons.append("measured ratio <= 0")
    if pred is not None and pred.pred_conc is None:
        if pred.evidence_level and pred.evidence_level not in {"internal_standard_row", "no_valid_ratio"}:
            reasons.append(str(pred.evidence_level))
        if pred.applicability_reason:
            ar = str(pred.applicability_reason)
            if "ratio" not in ar.lower() and ('internal standard' not in ar and '\u5185\u6807' not in ar):
                reasons.append(ar)
    return "; ".join(dict.fromkeys([r for r in reasons if r]))


def _component_formula(row: PreparedRow, group: str) -> str:
    formula = str(row.combo_parts.get(group, {}).get("formula", "") or "").strip()
    if formula:
        return formula
    # PreparedRow may have been created from a Combo containing escaped/full-width
    # punctuation. Reparse the normalized Combo before declaring Formula missing.
    return str(_combo_component_formulas(row.combo).get(group, "") or "").strip()


def _component_label(row: PreparedRow, group: str) -> str:
    """Stable component identity used by the RF model: molecular Formula only."""
    token = _formula_key_token(_component_formula(row, group))
    return f"formula:{token}" if token else ""


def _component_labels(row: PreparedRow) -> Dict[str, str]:
    return {g: _component_label(row, g) for g in ("A", "B", "C")}


def _component_matches(row: PreparedRow, other: PreparedRow, group: str) -> bool:
    """Match one component exclusively by its normalized molecular Formula."""
    fa = _component_formula(row, group)
    fb = _component_formula(other, group)
    if not fa or not fb:
        return False
    return _formula_compare(fa, fb) is True


def _same_combo_by_formulas(row: PreparedRow, other: PreparedRow) -> bool:
    return all(_component_matches(row, other, g) for g in ("A", "B", "C"))


def _pair_match(row: PreparedRow, other: PreparedRow, pair: Tuple[str, str]) -> bool:
    return all(_component_matches(row, other, g) for g in pair)


def _component_match_count(row: PreparedRow, other: PreparedRow) -> int:
    return sum(1 for g in ("A", "B", "C") if _component_matches(row, other, g))


def _element_keys(rows: Sequence[PreparedRow]) -> List[str]:
    keys = set()
    for r in rows:
        keys.update(r.element_counts.keys())
        for p in r.combo_parts.values():
            f = p.get("formula", "")
            if f:
                try:
                    keys.update(parse_formula(f).keys())
                except Exception:
                    pass
    preferred = ["C", "H", "D", "N", "O", "P", "S", "Na", "Cl", "Br", "I", "F", "K"]
    return [k for k in preferred if k in keys] + sorted(k for k in keys if k not in preferred)


def _formula_distance(a: Dict[str, int], b: Dict[str, int], keys: Sequence[str]) -> float:
    if not a and not b:
        return 0.0
    denom = 1.0 + sum(abs(a.get(k, 0)) + abs(b.get(k, 0)) for k in keys)
    return float(sum(abs(a.get(k, 0) - b.get(k, 0)) for k in keys) / max(denom, 1e-12))


def _guess_rt(row: PreparedRow) -> Optional[float]:
    candidates = [
        "Average_Apex_RT", "Avg_Apex_RT", "Mean_Apex_RT", "Apex_RT", "apex_rt",
        'Mean_Apex_RT', 'Mean_Apex_RT', "RT", "rt", "Retention_time", "Retention Time",
        "Rep1_Apex_RT", "Rep2_Apex_RT", "Rep3_Apex_RT",
    ]
    for c in candidates:
        if c in row.raw:
            x = parse_number(row.raw.get(c))
            if x is not None:
                return x
    for k, v in row.raw.items():
        nk = _norm_header(str(k))
        if ("apex" in nk and "rt" in nk) or nk in {"rt", "retentiontime"}:
            x = parse_number(v)
            if x is not None:
                return x
    return None


def _rt_similarity(a: PreparedRow, b: PreparedRow) -> float:
    ra = _guess_rt(a)
    rb = _guess_rt(b)
    if ra is None or rb is None:
        return 0.5
    # 0.3 min difference is already fairly weak, but not impossible.
    return float(math.exp(-abs(float(ra) - float(rb)) / 0.35))


def _formula_similarity(a: PreparedRow, b: PreparedRow, element_keys: Sequence[str]) -> float:
    return float(math.exp(-7.0 * _formula_distance(a.element_counts, b.element_counts, element_keys)))


def _mass_similarity(a: PreparedRow, b: PreparedRow) -> float:
    if a.exact_mass is None or b.exact_mass is None:
        return 0.5
    return float(math.exp(-abs(float(a.exact_mass) - float(b.exact_mass)) / 80.0))


def _row_similarity(a: PreparedRow, b: PreparedRow, element_keys: Sequence[str], *, use_rt: bool = True) -> float:
    """Combined structure/provenance similarity.

    Product or role-aware A/B/C/CORE fingerprints dominate when available.
    Formula, exact mass, retention time and component identity remain as support
    and as a fallback for rows whose private structures cannot be resolved.
    """
    if _same_combo_by_formulas(a, b):
        return 1.0
    comp = _component_match_count(a, b) / 3.0
    fs = _formula_similarity(a, b, element_keys)
    ms = _mass_similarity(a, b)
    rs = _rt_similarity(a, b) if use_rt else 0.5
    ss = structure_similarity(a, b)
    core = 1.0 if same_core(a, b) else (0.5 if not getattr(a, "core_id", "") or not getattr(b, "core_id", "") else 0.0)
    if ss is not None:
        score = 0.58 * float(ss) + 0.16 * comp + 0.09 * fs + 0.06 * ms + 0.06 * rs + 0.05 * core
    else:
        score = 0.58 * comp + 0.22 * fs + 0.10 * ms + 0.10 * rs
    return float(max(0.0, min(1.0, score)))


def _short_combo(combo: str, max_len: int = 110) -> str:
    s = str(combo or "").replace("\n", " ").strip()
    return s[: max_len - 1] + "\u2026" if len(s) > max_len else s


def _weighted_median(values: Sequence[float], weights: Sequence[float]) -> Optional[float]:
    pairs = []
    for v, w in zip(values, weights):
        try:
            vv = float(v)
            ww = float(w)
            if math.isfinite(vv) and math.isfinite(ww) and ww > 0:
                pairs.append((vv, ww))
        except Exception:
            pass
    if not pairs:
        return None
    pairs.sort(key=lambda t: t[0])
    total = sum(w for _, w in pairs)
    acc = 0.0
    for v, w in pairs:
        acc += w
        if acc >= total / 2.0:
            return float(v)
    return float(pairs[-1][0])


# -----------------------------------------------------------------------------
# Local RF model
# -----------------------------------------------------------------------------


@dataclass
class LocalRFModel:
    train_rows: List[PreparedRow]
    element_keys: List[str]
    global_log_rf: float
    evidence_interval_log: Dict[str, float]
    evidence_stats: Dict[str, Dict[str, object]]
    k_neighbors: int = 15
    min_single_count: int = 3
    use_rt: bool = True
    structure_threshold: float = 0.60


@dataclass
class LocalRFPrediction:
    row: PreparedRow
    pred_log_rf: Optional[float]
    pred_rf: Optional[float]
    pred_conc: Optional[float]
    lower_conc: Optional[float]
    upper_conc: Optional[float]
    evidence_level: str
    rf_source_level: str
    matched_count: int
    matched_rf_median: Optional[float]
    matched_rf_iqr: Optional[float]
    nearest_similarity: Optional[float]
    match_summary: str
    applicability: str
    applicability_reason: str
    warnings: List[str]


_EVIDENCE_PRIORITY = [
    "exact_ABC_formula",
    "same_product_structure",
    "structure_high",
    "same_AB",
    "same_AC",
    "same_BC",
    "same_pair_any",
    "structure_local",
    "same_component",
    "nearest_local",
    "global_median",
]


def _rf_values(rows: Sequence[PreparedRow]) -> List[float]:
    return [float(r.log_correction_factor) for r in rows if r.log_correction_factor is not None and math.isfinite(float(r.log_correction_factor))]


def _rf_stats(rows: Sequence[PreparedRow]) -> Tuple[Optional[float], Optional[float]]:
    vals = _rf_values(rows)
    if not vals:
        return None, None
    med = float(statistics.median(vals))
    if len(vals) >= 4:
        q75, q25 = np.percentile(np.array(vals, dtype=float), [75, 25])
        iqr = float(q75 - q25)
    elif len(vals) >= 2:
        iqr = float(max(vals) - min(vals))
    else:
        iqr = 0.0
    return med, iqr


def _estimate_from_rows(target: PreparedRow, rows: Sequence[PreparedRow], element_keys: Sequence[str], *, level: str, use_rt: bool = True) -> Tuple[Optional[float], Optional[float], Optional[float], str, Optional[float]]:
    if not rows:
        return None, None, None, "", None
    vals = []
    weights = []
    scored: List[Tuple[float, PreparedRow]] = []
    for r in rows:
        y = r.log_correction_factor
        if y is None:
            continue
        sim = _row_similarity(target, r, element_keys, use_rt=use_rt)
        # Evidence-specific weights.  Median/weighted median is intentionally robust.
        if level in {"exact_ABC_formula", "same_product_structure"}:
            w = 1.0
        elif level in {"structure_high", "structure_local"}:
            w = max(sim, 1e-6) ** 3
        elif level.startswith("same_"):
            w = 0.70 + 0.30 * sim
        else:
            w = max(sim, 1e-6) ** 2
        vals.append(float(y))
        weights.append(float(w))
        scored.append((sim, r))
    if not vals:
        return None, None, None, "", None
    y = _weighted_median(vals, weights)
    med, iqr = _rf_stats(rows)
    scored.sort(key=lambda t: t[0], reverse=True)
    summary = "; ".join([f"{_short_combo(r.combo, 75)} (RF={math.exp(float(r.log_correction_factor)):.4g}, sim={s:.2f})" for s, r in scored[:5] if r.log_correction_factor is not None])
    sim_mean = float(sum(s for s, _ in scored[: max(1, min(15, len(scored)))]) / max(1, min(15, len(scored)))) if scored else None
    return y, med, iqr, summary, sim_mean


def _select_local_log_rf(
    target: PreparedRow,
    train: Sequence[PreparedRow],
    element_keys: Sequence[str],
    *,
    k: int = 15,
    min_single_count: int = 3,
    use_rt: bool = True,
    structure_threshold: float = 0.60,
) -> Tuple[Optional[float], str, str, int, Optional[float], Optional[float], Optional[float], str]:
    """Select a local response factor using structural evidence first."""
    train = list(train)
    if not train:
        return None, "no_training", "no_training", 0, None, None, None, ""

    # 1. Exact ordered A/B/C component-Formula combination.
    exact = [r for r in train if _same_combo_by_formulas(target, r)]
    if exact:
        y, med, iqr, summary, sim = _estimate_from_rows(target, exact, element_keys, level="exact_ABC_formula", use_rt=use_rt)
        return y, "exact_ABC_formula", "exact_ABC_formula", len(exact), med, iqr, 1.0, summary

    # 2. Identical product structure, compared by privacy-safe hash.
    target_hash = str(getattr(target, "structure_hash", "") or "")
    if target_hash and getattr(target, "structure_status", "") == "product_structure_ready":
        same_product = [
            r for r in train
            if str(getattr(r, "structure_hash", "") or "") == target_hash
            and getattr(r, "structure_status", "") == "product_structure_ready"
        ]
        if same_product:
            y, med, iqr, summary, sim = _estimate_from_rows(target, same_product, element_keys, level="same_product_structure", use_rt=use_rt)
            return y, "same_product_structure", "same_product_structure", len(same_product), med, iqr, 1.0, summary

    # 3. High structural similarity within the same fixed core.
    structural: List[Tuple[float, PreparedRow]] = []
    for r in train:
        ss = structure_similarity(target, r)
        if ss is None:
            continue
        tc = str(getattr(target, "core_id", "") or "")
        rc = str(getattr(r, "core_id", "") or "")
        if tc and rc and tc != rc:
            continue
        structural.append((float(ss), r))
    structural.sort(key=lambda t: t[0], reverse=True)
    high = [(s, r) for s, r in structural if s >= float(structure_threshold)]
    if high:
        rows = [r for _, r in high[: max(1, int(k))]]
        y, med, iqr, summary, sim = _estimate_from_rows(target, rows, element_keys, level="structure_high", use_rt=use_rt)
        return y, "structure_high", "structure_high", len(rows), med, iqr, sim, summary

    # 4. Same component pair.
    pair_defs = [("same_AB", ("A", "B")), ("same_AC", ("A", "C")), ("same_BC", ("B", "C"))]
    pair_candidates: List[Tuple[int, float, str, List[PreparedRow]]] = []
    for level, pair in pair_defs:
        rows = [r for r in train if _pair_match(target, r, pair)]
        if not rows:
            continue
        sim = float(sum(_row_similarity(target, r, element_keys, use_rt=use_rt) for r in rows) / len(rows))
        pair_candidates.append((len(rows), sim, level, rows))
    if pair_candidates:
        pair_candidates.sort(key=lambda t: (t[0], t[1]), reverse=True)
        _, sim, level, rows = pair_candidates[0]
        y, med, iqr, summary, sim2 = _estimate_from_rows(target, rows, element_keys, level=level, use_rt=use_rt)
        return y, level, level, len(rows), med, iqr, sim2 if sim2 is not None else sim, summary

    # 5. Nearest structural neighbours even if below the high threshold.
    if structural:
        rows = [r for _, r in structural[: max(1, int(k))]]
        y, med, iqr, summary, sim = _estimate_from_rows(target, rows, element_keys, level="structure_local", use_rt=use_rt)
        return y, "structure_local", "structure_local", len(rows), med, iqr, sim, summary

    # 6. Same single component.
    single_rows: List[Tuple[float, PreparedRow]] = []
    for r in train:
        cm = _component_match_count(target, r)
        if cm <= 0:
            continue
        sim = _row_similarity(target, r, element_keys, use_rt=use_rt)
        single_rows.append((cm + sim, r))
    if single_rows:
        single_rows.sort(key=lambda t: t[0], reverse=True)
        rows = [r for _, r in single_rows[: max(int(k) * 2, min_single_count)]]
        if len(rows) >= min_single_count:
            y, med, iqr, summary, sim = _estimate_from_rows(target, rows, element_keys, level="same_component", use_rt=use_rt)
            return y, "same_component", "same_component", len(rows), med, iqr, sim, summary

    # 7. Formula/mass/RT/provenance nearest-neighbour fallback.
    scored = [(_row_similarity(target, r, element_keys, use_rt=use_rt), r) for r in train]
    scored.sort(key=lambda t: t[0], reverse=True)
    top = [(s, r) for s, r in scored[: max(1, int(k))] if s > 0]
    if top:
        rows = [r for _, r in top]
        y, med, iqr, summary, sim = _estimate_from_rows(target, rows, element_keys, level="nearest_local", use_rt=use_rt)
        return y, "nearest_local", "nearest_local", len(rows), med, iqr, sim, summary

    # 8. Global median.
    y, med, iqr, summary, sim = _estimate_from_rows(target, train, element_keys, level="global_median", use_rt=use_rt)
    return y, "global_median", "global_median", len(train), med, iqr, sim, summary


def _prediction_applicability(level: str, matched_count: int, sim: Optional[float], evidence_stats: Dict[str, Dict[str, object]]) -> Tuple[str, str]:
    stats = evidence_stats.get(level, {}) if evidence_stats else {}
    medape = _as_float_or_none(stats.get("Median_APE_%"))
    within2 = _as_float_or_none(stats.get("Within_2x_%"))
    ncv = int(_as_float_or_none(stats.get("N")) or 0)

    if level == "exact_ABC_formula":
        base = "High"
        reason = "The exact ordered A/B/C component-Formula combination exists in the calibration set."
    elif level == "same_product_structure":
        base = "High"
        reason = "The same locally resolved product structure exists in the calibration set."
    elif level == "structure_high":
        base = "Medium-high" if matched_count >= 2 else "Medium"
        reason = f"RF borrowed from high-similarity structures sharing the same fixed core (n={matched_count}, similarity={sim:.2f})." if sim is not None else f"RF borrowed from high-similarity structures sharing the same fixed core (n={matched_count})."
    elif level in {"same_AB", "same_AC", "same_BC"}:
        base = "Medium-high" if matched_count >= 2 else "Medium"
        reason = f"Local RF borrowed from calibration rows sharing the same {level[-2:]} component pair (n={matched_count})."
    elif level == "structure_local":
        base = "Medium" if sim is not None and sim >= 0.50 else "Low-medium"
        reason = f"RF borrowed from the nearest local structures (n={matched_count}, similarity={sim:.2f})." if sim is not None else f"RF borrowed from the nearest local structures (n={matched_count})."
    elif level == "same_component":
        base = "Medium"
        reason = f"Local RF borrowed from rows sharing at least one A/B/C component (n={matched_count})."
    elif level == "nearest_local":
        if sim is not None and sim >= 0.55:
            base = "Low-medium"
        else:
            base = "Low"
        reason = f"No shared component pair; nearest local RF was used (mean similarity={sim:.2f})." if sim is not None else "No shared component pair; nearest local RF was used."
    elif level == "global_median":
        base = "Very low"
        reason = "No useful local calibration support; global median RF was used."
    else:
        base = "No prediction"
        reason = "No valid prediction."

    # Evidence-level empirical QC can downgrade the reportability.
    if ncv >= 5 and medape is not None:
        if medape > 100 or (within2 is not None and within2 < 40):
            if base in {"High", "Medium-high"}:
                base = "Medium"
            elif base == "Medium":
                base = "Low-medium"
            reason += f" LOOCV for this evidence level is weak (Median_APE={medape:.1f}%, N={ncv})."
        elif medape <= 50 and within2 is not None and within2 >= 50:
            reason += f" Evidence-level LOOCV is acceptable for screening (Median_APE={medape:.1f}%, Within2x={within2:.1f}%, N={ncv})."
    elif level not in {"global_median", "no_training", "no_valid_ratio"}:
        reason += " Evidence-level LOOCV has few samples; interpret cautiously."
    return base, reason


def fit_local_rf_model(train_rows: Sequence[PreparedRow], element_keys: Sequence[str], *, k_neighbors: int = 15, evidence_interval_log: Optional[Dict[str, float]] = None, evidence_stats: Optional[Dict[str, Dict[str, object]]] = None, use_rt: bool = True, structure_threshold: float = 0.60) -> LocalRFModel:
    train = valid_training_rows(train_rows)
    vals = _rf_values(train)
    global_log = float(statistics.median(vals)) if vals else 0.0
    return LocalRFModel(
        train_rows=list(train),
        element_keys=list(element_keys),
        global_log_rf=global_log,
        evidence_interval_log=dict(evidence_interval_log or {}),
        evidence_stats=dict(evidence_stats or {}),
        k_neighbors=int(k_neighbors),
        min_single_count=3,
        use_rt=bool(use_rt),
        structure_threshold=float(structure_threshold),
    )


def predict_one_local_rf(row: PreparedRow, model: LocalRFModel) -> LocalRFPrediction:
    warnings: List[str] = []
    if _is_internal_standard_row(row):
        return LocalRFPrediction(row, None, None, None, None, None, "internal_standard_row", "internal_standard_row", 0, None, None, None, "", "No prediction", "Internal standard row is excluded from concentration prediction.", ['Internal-standard rows are not concentration prediction targets.'])
    if row.ratio is None or row.ratio <= 0:
        return LocalRFPrediction(row, None, None, None, None, None, "no_valid_ratio", "no_valid_ratio", 0, None, None, None, "", "No prediction", "No valid measured ratio.", ['No valid target-to-internal-standard area ratio.'])
    if not model.train_rows:
        return LocalRFPrediction(row, None, None, None, None, None, "no_training", "no_training", 0, None, None, None, "", "No prediction", "No usable calibration rows.", ['No usable calibration records.'])

    y, level, src, n, med, iqr, sim, summary = _select_local_log_rf(
        row, model.train_rows, model.element_keys,
        k=model.k_neighbors, min_single_count=model.min_single_count,
        use_rt=model.use_rt, structure_threshold=model.structure_threshold,
    )
    if y is None or not math.isfinite(float(y)):
        return LocalRFPrediction(row, None, None, None, None, None, "failed", "failed", n, med, iqr, sim, summary, "No prediction", "Invalid transferred response factor.", ['Response-factor prediction failed.'])

    rf = float(math.exp(float(y)))
    pred = float(row.ratio * rf)
    # Evidence-specific interval if available; otherwise fallback to global local interval.
    interval_log = model.evidence_interval_log.get(level)
    if interval_log is None:
        interval_log = model.evidence_interval_log.get("__global__", math.log(5.0))
    lower = float(pred / math.exp(interval_log)) if interval_log and interval_log > 0 else None
    upper = float(pred * math.exp(interval_log)) if interval_log and interval_log > 0 else None
    app, reason = _prediction_applicability(level, n, sim, model.evidence_stats)
    return LocalRFPrediction(row, float(y), rf, pred, lower, upper, level, src, n, math.exp(med) if med is not None else None, math.exp(iqr) if iqr is not None else None, sim, summary, app, reason, warnings)


# -----------------------------------------------------------------------------
# Evaluation
# -----------------------------------------------------------------------------


def _augment_cv_row(row: Dict[str, object]) -> Dict[str, object]:
    actual = _as_float_or_none(row.get("Actual_concentration"))
    pred = _as_float_or_none(row.get("Predicted_concentration_LOOCV"))
    out = dict(row)
    if actual is not None and pred is not None and actual > 0 and pred > 0:
        err_pct = (pred - actual) / actual * 100.0
        ape = abs(err_pct)
        pa = pred / actual
        fold = max(pa, 1.0 / pa)
        if fold <= 2:
            band = "within_2x"
        elif fold <= 5:
            band = "within_5x"
        elif fold <= 10:
            band = "within_10x"
        else:
            band = ">10x"
        out.update({
            "Abs_Error_%": float(ape),
            "Pred/Actual": float(pa),
            "Fold_error": float(fold),
            "Accuracy_band": band,
            "log10_actual": float(math.log10(actual)),
            "log10_predicted": float(math.log10(pred)),
        })
    return out


def evaluate_local_qc(qc: Dict[str, object], n_train: int) -> Dict[str, object]:
    medape = _as_float_or_none(qc.get("Median_APE_%"))
    within2 = _as_float_or_none(qc.get("Within_2x_%"))
    within5 = _as_float_or_none(qc.get("Within_5x_%"))
    r2 = _as_float_or_none(qc.get("R2_log"))
    if n_train < 20:
        rating = "Insufficient calibration"
        conclusion = "Calibration set is too small for robust local RF transfer."
        rec = "Increase known-concentration calibration compounds and cover more A/B/C component pairs."
    elif medape is not None and within2 is not None and r2 is not None and medape <= 35 and within2 >= 65 and r2 >= 0.25:
        rating = "Good for local semi-quantitation"
        conclusion = "Local response-factor transfer is reasonably supported for candidates within the calibration domain."
        rec = "Report predictions together with evidence level and empirical prediction intervals."
    elif within5 is not None and within5 >= 85 and medape is not None and medape <= 80:
        rating = "Screening / range estimation"
        conclusion = "Most predictions are within a few-fold range, but exact concentration accuracy is limited."
        rec = "Use for prioritization and concentration ranges; report only high/medium-evidence candidates as semi-quantitative."
    else:
        rating = "Poor / screening only"
        conclusion = "LOOCV does not support reliable concentration prediction."
        rec = "Use as rough screening only; add calibrants for poorly supported component pairs or structure classes."
    out = dict(qc)
    out["Model_rating"] = rating
    out["Model_conclusion"] = conclusion
    out["Recommendation"] = rec
    return out


def loocv_local_rf(train_rows: Sequence[PreparedRow], element_keys: Sequence[str], *, k: int = 15, use_rt: bool = True, structure_threshold: float = 0.60) -> Tuple[Dict[str, object], List[Dict[str, object]], Dict[str, Dict[str, object]], Dict[str, float]]:
    train = valid_training_rows(train_rows)
    cv_rows: List[Dict[str, object]] = []
    y_true: List[float] = []
    y_pred: List[float] = []
    if len(train) < 3:
        return {"N": len(train), "warning": 'Fewer than three calibration records; LOOCV was not performed.'}, [], {}, {"__global__": math.log(5.0)}

    # First pass with default interval, collecting evidence-level errors.
    dummy_intervals = {"__global__": math.log(5.0)}
    for idx, row in enumerate(train):
        others = [r for j, r in enumerate(train) if j != idx]
        m = fit_local_rf_model(others, element_keys, k_neighbors=k, evidence_interval_log=dummy_intervals, use_rt=use_rt, structure_threshold=structure_threshold)
        pr = predict_one_local_rf(row, m)
        actual = float(row.concentration) if row.concentration is not None else None
        pred = pr.pred_conc
        if actual is not None and pred is not None and actual > 0 and pred > 0:
            y_true.append(actual)
            y_pred.append(pred)
            err_pct = (pred - actual) / actual * 100.0
            cv_rows.append({
                "Source_file": _source_file(row),
                "Source_sheet": _source_sheet(row),
                "Source_row": _source_row_no(row),
                "Source_index": row.source_index,
                "Combo": row.combo,
                "Formula": row.formula,
                "Ratio": row.ratio,
                "Actual_concentration": actual,
                "Predicted_concentration_LOOCV": pred,
                "Predicted_RF_LOOCV": pr.pred_rf,
                "Transferred_logRF_LOOCV": pr.pred_log_rf,
                "Error_%": err_pct,
                "Evidence": pr.evidence_level,
                "RF_source_level": pr.rf_source_level,
                "Matched_count": pr.matched_count,
                "Matched_RF_median": pr.matched_rf_median,
                "Nearest_similarity": pr.nearest_similarity,
                "Applicability": pr.applicability,
                "Match_summary": pr.match_summary,
            })

    metrics = _metrics(y_true, y_pred)
    cv_aug = [_augment_cv_row(r) for r in cv_rows]

    # Evidence-level metrics and empirical intervals.
    evidence_stats: Dict[str, Dict[str, object]] = {}
    intervals: Dict[str, float] = {}
    all_fold = [float(r.get("Fold_error")) for r in cv_aug if _as_float_or_none(r.get("Fold_error")) is not None]
    if all_fold:
        global_p80 = float(np.percentile(np.array(all_fold, dtype=float), 80))
        intervals["__global__"] = float(math.log(max(global_p80, 1.01)))
        metrics["Prediction_interval_fold_P80"] = global_p80
    else:
        intervals["__global__"] = math.log(5.0)
        metrics["Prediction_interval_fold_P80"] = 5.0

    levels = sorted(set(str(r.get("Evidence", "")) for r in cv_aug if r.get("Evidence")), key=lambda x: _EVIDENCE_PRIORITY.index(x) if x in _EVIDENCE_PRIORITY else 999)
    for lev in levels:
        rows = [r for r in cv_aug if r.get("Evidence") == lev]
        yt = [_as_float_or_none(r.get("Actual_concentration")) for r in rows]
        yp = [_as_float_or_none(r.get("Predicted_concentration_LOOCV")) for r in rows]
        st = _metrics([x for x in yt if x is not None], [x for x in yp if x is not None])
        folds = [float(r.get("Fold_error")) for r in rows if _as_float_or_none(r.get("Fold_error")) is not None]
        if folds:
            p80 = float(np.percentile(np.array(folds, dtype=float), 80))
            st["Prediction_interval_fold_P80"] = p80
            # Use level-specific interval only when enough LOOCV examples exist.
            if int(st.get("N") or 0) >= 5:
                intervals[lev] = float(math.log(max(p80, 1.01)))
        evidence_stats[lev] = st

    return metrics, cv_aug, evidence_stats, intervals


def component_leaveout_cv(*args, **kwargs) -> List[Dict[str, object]]:
    # Local transfer evidence-level LOOCV is more informative than component leaveout
    # for this model.  Keep a stub for compatibility if external code imports it.
    return []


# -----------------------------------------------------------------------------
# Plots and workbook output
# -----------------------------------------------------------------------------


def _make_plots(cv_rows: Sequence[Dict[str, object]], out_xlsx: Path) -> List[Tuple[str, Path]]:
    if not cv_rows:
        return []
    try:
        import matplotlib.pyplot as plt
    except Exception:
        return []
    plot_dir = Path(out_xlsx).with_suffix("").parent / (Path(out_xlsx).stem + "__diagnostics")
    plot_dir.mkdir(parents=True, exist_ok=True)
    yt = np.array([float(r.get("Actual_concentration")) for r in cv_rows if _as_float_or_none(r.get("Actual_concentration")) is not None and _as_float_or_none(r.get("Predicted_concentration_LOOCV")) is not None and float(r.get("Actual_concentration")) > 0 and float(r.get("Predicted_concentration_LOOCV")) > 0], dtype=float)
    yp = np.array([float(r.get("Predicted_concentration_LOOCV")) for r in cv_rows if _as_float_or_none(r.get("Actual_concentration")) is not None and _as_float_or_none(r.get("Predicted_concentration_LOOCV")) is not None and float(r.get("Actual_concentration")) > 0 and float(r.get("Predicted_concentration_LOOCV")) > 0], dtype=float)
    if len(yt) < 2:
        return []
    outs: List[Tuple[str, Path]] = []

    # Observed vs predicted log plot.
    fig = plt.figure(figsize=(6.0, 4.8), dpi=160)
    ax = fig.add_subplot(111)
    ax.scatter(np.log10(yt), np.log10(yp), s=18, alpha=0.8)
    mn = float(min(np.min(np.log10(yt)), np.min(np.log10(yp))))
    mx = float(max(np.max(np.log10(yt)), np.max(np.log10(yp))))
    pad = max(0.05, (mx - mn) * 0.08)
    xs = np.linspace(mn - pad, mx + pad, 100)
    ax.plot(xs, xs, linewidth=1.0, label="1:1")
    ax.plot(xs, xs + math.log10(2), linestyle="--", linewidth=0.9, label="2x")
    ax.plot(xs, xs - math.log10(2), linestyle="--", linewidth=0.9)
    ax.plot(xs, xs + math.log10(5), linestyle=":", linewidth=0.9, label="5x")
    ax.plot(xs, xs - math.log10(5), linestyle=":", linewidth=0.9)
    ax.set_xlabel("log10(actual concentration)")
    ax.set_ylabel("log10(LOOCV predicted concentration)")
    ax.set_title("Local RF transfer LOOCV: observed vs predicted")
    ax.grid(True, alpha=0.25)
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    p1 = plot_dir / "local_rf_loocv_observed_vs_predicted.png"
    fig.savefig(p1)
    plt.close(fig)
    outs.append(("LOOCV observed vs predicted", p1))

    # Fold error profile.
    fold = np.maximum(yp / yt, yt / yp)
    fig = plt.figure(figsize=(6.0, 4.2), dpi=160)
    ax = fig.add_subplot(111)
    xs = np.arange(1, len(fold) + 1)
    ax.plot(xs, np.sort(fold), marker="o", markersize=3, linewidth=1.0)
    ax.axhline(2.0, linestyle="--", linewidth=0.9, label="2x")
    ax.axhline(5.0, linestyle=":", linewidth=0.9, label="5x")
    ax.set_xlabel("Calibration compounds sorted by fold error")
    ax.set_ylabel("Fold error")
    ax.set_yscale("log")
    ax.set_title("LOOCV fold-error profile")
    ax.grid(True, alpha=0.25)
    ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    p2 = plot_dir / "local_rf_loocv_fold_error_profile.png"
    fig.savefig(p2)
    plt.close(fig)
    outs.append(("LOOCV fold error", p2))

    # Evidence level median fold.
    evidence = {}
    for r in cv_rows:
        lev = str(r.get("Evidence", ""))
        fe = _as_float_or_none(r.get("Fold_error"))
        if lev and fe is not None:
            evidence.setdefault(lev, []).append(float(fe))
    if evidence:
        levels = sorted(evidence.keys(), key=lambda x: _EVIDENCE_PRIORITY.index(x) if x in _EVIDENCE_PRIORITY else 999)
        meds = [float(np.median(evidence[lev])) for lev in levels]
        fig = plt.figure(figsize=(7.0, 4.2), dpi=160)
        ax = fig.add_subplot(111)
        ax.bar(np.arange(len(levels)), meds)
        ax.axhline(2.0, linestyle="--", linewidth=0.9, label="2x")
        ax.axhline(5.0, linestyle=":", linewidth=0.9, label="5x")
        ax.set_xticks(np.arange(len(levels)))
        ax.set_xticklabels(levels, rotation=30, ha="right")
        ax.set_ylabel("Median fold error")
        ax.set_yscale("log")
        ax.set_title("LOOCV median fold error by RF evidence level")
        ax.grid(True, axis="y", alpha=0.25)
        ax.legend(frameon=False, fontsize=8)
        fig.tight_layout()
        p3 = plot_dir / "local_rf_evidence_level_fold_error.png"
        fig.savefig(p3)
        plt.close(fig)
        outs.append(("Evidence-level fold error", p3))

    return outs


def _style_workbook(wb) -> None:
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        return
    fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for col_idx in range(1, min(ws.max_column, 40) + 1):
            letter = get_column_letter(col_idx)
            max_len = 10
            for cell in ws[letter][: min(ws.max_row, 200)]:
                s = str(cell.value or "")
                if len(s) > max_len:
                    max_len = min(len(s), 60)
            ws.column_dimensions[letter].width = max(10, min(max_len + 2, 60))


def write_local_rf_workbook(
    out_xlsx: Path,
    predictions: Sequence[LocalRFPrediction],
    train_rows: Sequence[PreparedRow],
    qc: Dict[str, object],
    cv_rows: Sequence[Dict[str, object]],
    evidence_stats: Dict[str, Dict[str, object]],
    *,
    target_rows: Sequence[PreparedRow],
    structure_library: StructureLibrary,
    privacy_mode: bool,
    model_note: str,
    output_csv: bool = True,
) -> Tuple[Path, Optional[Path]]:
    try:
        from openpyxl import Workbook
    except Exception as e:
        raise RuntimeError('Excel export requires openpyxl. Install it in the active environment: python -m pip install openpyxl') from e

    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_pred = wb.active
    ws_pred.title = "Predictions"
    ws_cal = wb.create_sheet("Calibration")
    ws_qc = wb.create_sheet("Model_QC")
    ws_cv = wb.create_sheet("LOOCV")
    ws_evl = wb.create_sheet("LOOCV_by_Evidence")
    ws_train_diag = wb.create_sheet("Training_Row_Check")
    ws_target_diag = wb.create_sheet("Target_Row_Check")
    ws_struct = wb.create_sheet("Structure_Coverage")
    ws_join = wb.create_sheet("Structure_Join_Summary")
    ws_lib = wb.create_sheet("Structure_Library_Audit")
    ws_rules = wb.create_sheet("Reaction_Rules_Audit")
    ws_readme = wb.create_sheet("Readme")

    structure_headers = [
        "Core_ID", "Structure_status", "Structure_method", "Structure_hash",
        "Product_formula_calc", "Formula_match", "3D_status", "Descriptor_count",
        "Product_fingerprint", "Component_fingerprint_count",
    ]
    if not privacy_mode:
        structure_headers.append("Product_SMILES")

    pred_headers = [
        "Source_file", "Source_sheet", "Source_row", "Source_index", "Name", "Formula", "Combo", "Measured_ratio",
        "Predicted_concentration", "Prediction_lower", "Prediction_upper", "Predicted_RF", "Predicted_logRF",
        "RF_source_level", "Matched_count", "Matched_RF_median", "Matched_RF_IQR_fold", "Nearest_similarity",
        "Match_summary", "Applicability", "Applicability_reason", "A_label", "B_label", "C_label", "Exact_mass",
    ] + structure_headers + ["Warnings"]
    def _privacy_sensitive_header(header: object) -> bool:
        n = _norm_header(str(header or ""))
        return any(token in n for token in (
            "smiles", "smarts", "inchi", "molblock", "molfile", "sdfstructure",
            'Structure', 'Reaction_Structure', 'Product_Structure',
        ))

    orig_headers: List[str] = []
    for pr in predictions:
        for h in pr.row.raw.keys():
            if str(h).startswith("__source_"):
                continue
            if privacy_mode and _privacy_sensitive_header(h):
                continue
            if h not in orig_headers:
                orig_headers.append(h)
    ws_pred.append(pred_headers + [f"orig__{h}" for h in orig_headers])
    csv_rows: List[List[object]] = []
    for pr in predictions:
        p = pr.row.combo_parts
        struct_vals: List[object] = [
            getattr(pr.row, "core_id", ""),
            getattr(pr.row, "structure_status", ""),
            getattr(pr.row, "structure_method", ""),
            getattr(pr.row, "structure_hash", ""),
            getattr(pr.row, "product_formula_calc", ""),
            getattr(pr.row, "formula_match", ""),
            getattr(pr.row, "three_d_status", ""),
            len(getattr(pr.row, "structure_features", {}) or {}),
            "Yes" if getattr(pr.row, "structure_fingerprint", None) is not None else "No",
            len(getattr(pr.row, "component_fingerprints", {}) or {}),
        ]
        if not privacy_mode:
            struct_vals.append(getattr(pr.row, "product_smiles", ""))
        vals = [
            _source_file(pr.row), _source_sheet(pr.row), _source_row_no(pr.row), pr.row.source_index,
            pr.row.name, pr.row.formula, pr.row.combo, pr.row.ratio,
            pr.pred_conc, pr.lower_conc, pr.upper_conc, pr.pred_rf, pr.pred_log_rf,
            pr.rf_source_level, pr.matched_count, pr.matched_rf_median, pr.matched_rf_iqr,
            pr.nearest_similarity, pr.match_summary, pr.applicability, pr.applicability_reason,
            p.get("A", {}).get("label", ""), p.get("B", {}).get("label", ""), p.get("C", {}).get("label", ""),
            pr.row.exact_mass,
        ] + struct_vals + ["; ".join(pr.warnings + pr.row.warnings + (getattr(pr.row, "structure_warnings", []) or []))]
        ws_pred.append(vals + [pr.row.raw.get(h, "") for h in orig_headers])
        csv_rows.append(vals)

    cal_headers = [
        "Source_file", "Source_sheet", "Source_row", "Source_index", "Name", "Formula", "Combo", "Measured_ratio",
        "Actual_concentration", "RF(conc/ratio)", "logRF", "A_label", "B_label", "C_label", "Exact_mass",
    ] + structure_headers + ["Warnings"]
    ws_cal.append(cal_headers)
    for r in train_rows:
        p = r.combo_parts
        struct_vals = [
            getattr(r, "core_id", ""), getattr(r, "structure_status", ""), getattr(r, "structure_method", ""),
            getattr(r, "structure_hash", ""), getattr(r, "product_formula_calc", ""), getattr(r, "formula_match", ""),
            getattr(r, "three_d_status", ""), len(getattr(r, "structure_features", {}) or {}),
            "Yes" if getattr(r, "structure_fingerprint", None) is not None else "No",
            len(getattr(r, "component_fingerprints", {}) or {}),
        ]
        if not privacy_mode:
            struct_vals.append(getattr(r, "product_smiles", ""))
        ws_cal.append([
            _source_file(r), _source_sheet(r), _source_row_no(r), r.source_index, r.name, r.formula, r.combo,
            r.ratio, r.concentration, r.correction_factor, r.log_correction_factor,
            p.get("A", {}).get("label", ""), p.get("B", {}).get("label", ""), p.get("C", {}).get("label", ""),
            r.exact_mass,
        ] + struct_vals + ["; ".join(r.warnings + (getattr(r, "structure_warnings", []) or []))])

    ws_qc.append(["Metric", "Value"])
    for k, v in qc.items():
        ws_qc.append([k, v])

    cv_headers = [
        "Source_file", "Source_sheet", "Source_row", "Source_index", "Combo", "Formula", "Ratio",
        "Actual_concentration", "Predicted_concentration_LOOCV", "Predicted_RF_LOOCV", "Transferred_logRF_LOOCV",
        "Error_%", "Abs_Error_%", "Pred/Actual", "Fold_error", "Accuracy_band", "log10_actual", "log10_predicted",
        "Evidence", "RF_source_level", "Matched_count", "Matched_RF_median", "Nearest_similarity", "Applicability", "Match_summary",
    ]
    ws_cv.append(cv_headers)
    for r in cv_rows:
        ws_cv.append([r.get(h, "") for h in cv_headers])

    ev_headers = ["Evidence_level", "N", "Median_APE_%", "MAPE_%", "P90_APE_%", "Within_2x_%", "Within_5x_%", "R2_log", "Median_fold_error", "P80_fold_interval", "Use_in_prediction_interval"]
    ws_evl.append(ev_headers)
    ordered = sorted(evidence_stats.keys(), key=lambda x: _EVIDENCE_PRIORITY.index(x) if x in _EVIDENCE_PRIORITY else 999)
    for lev in ordered:
        st = evidence_stats.get(lev, {})
        ws_evl.append([
            lev, st.get("N", ""), st.get("Median_APE_%", ""), st.get("MAPE_%", ""), st.get("P90_APE_%", ""),
            st.get("Within_2x_%", ""), st.get("Within_5x_%", ""), st.get("R2_log", ""), st.get("Median_fold_error", ""),
            st.get("Prediction_interval_fold_P80", ""), "Yes if N>=5 else global interval",
        ])

    train_diag_headers = [
        "Status", "Skip_reason", "Source_file", "Source_sheet", "Source_row", "Source_index", "Name", "Formula", "Combo",
        "Measured_ratio", "Actual_concentration", "logRF",
        "Product_Master_Match", "Product_Master_Match_Method", "Product_Master_Formula_Key",
        "Product_Master_Verification_Mode", "Product_Master_Join_Status", "Product_Master_Formula_Flag",
        "Product_Master_Formula_Check", "Product_Master_Component_Formula_Check",
        "Product_Master_Internal_Formula_Check", "Product_Master_Master_Formula",
        "Product_Master_Candidate_Count", "Product_Master_Formula_Note",
        "Product_Master_Unmatched_Reason",
        "Core_ID", "Structure_status", "Formula_match", "Warnings",
    ]
    ws_train_diag.append(train_diag_headers)
    for r in train_rows:
        reason = training_skip_reason(r)
        ws_train_diag.append([
            "USED" if not reason else "SKIPPED", reason, _source_file(r), _source_sheet(r), _source_row_no(r), r.source_index,
            r.name, r.formula, r.combo, r.ratio, r.concentration, r.log_correction_factor,
            r.raw.get("__product_master_match__", ""), r.raw.get("__product_master_match_method__", ""),
            r.raw.get("__product_master_formula_key__", ""),
            r.raw.get("__product_master_verification_mode__", ""),
            r.raw.get("__product_master_join_status__", ""),
            r.raw.get("__product_master_formula_flag__", ""),
            r.raw.get("__product_master_formula_check__", ""),
            r.raw.get("__product_master_component_formula_check__", ""),
            r.raw.get("__product_master_internal_formula_check__", ""),
            r.raw.get("__product_master_master_formula__", ""),
            r.raw.get("__product_master_candidate_count__", ""),
            r.raw.get("__product_master_formula_note__", ""),
            r.raw.get("__product_master_unmatched_reason__", ""),
            getattr(r, "core_id", ""), getattr(r, "structure_status", ""), getattr(r, "formula_match", ""),
            "; ".join(r.warnings + (getattr(r, "structure_warnings", []) or [])),
        ])

    target_diag_headers = [
        "Status", "Skip_reason", "Source_file", "Source_sheet", "Source_row", "Source_index", "Name", "Formula", "Combo",
        "Measured_ratio", "Predicted_concentration", "RF_source_level", "Matched_count", "Applicability",
        "Product_Master_Match", "Product_Master_Match_Method", "Product_Master_Formula_Key",
        "Product_Master_Verification_Mode", "Product_Master_Join_Status", "Product_Master_Formula_Flag",
        "Product_Master_Formula_Check", "Product_Master_Component_Formula_Check",
        "Product_Master_Internal_Formula_Check", "Product_Master_Master_Formula",
        "Product_Master_Candidate_Count", "Product_Master_Formula_Note",
        "Product_Master_Unmatched_Reason",
        "Core_ID", "Structure_status", "Formula_match", "Warnings",
    ]
    ws_target_diag.append(target_diag_headers)
    for pr in predictions:
        reason = target_skip_reason(pr.row, pr)
        ws_target_diag.append([
            "PREDICTED" if pr.pred_conc is not None else "SKIPPED", reason, _source_file(pr.row), _source_sheet(pr.row),
            _source_row_no(pr.row), pr.row.source_index, pr.row.name, pr.row.formula, pr.row.combo, pr.row.ratio,
            pr.pred_conc, pr.rf_source_level, pr.matched_count, pr.applicability,
            pr.row.raw.get("__product_master_match__", ""), pr.row.raw.get("__product_master_match_method__", ""),
            pr.row.raw.get("__product_master_formula_key__", ""),
            pr.row.raw.get("__product_master_verification_mode__", ""),
            pr.row.raw.get("__product_master_join_status__", ""),
            pr.row.raw.get("__product_master_formula_flag__", ""),
            pr.row.raw.get("__product_master_formula_check__", ""),
            pr.row.raw.get("__product_master_component_formula_check__", ""),
            pr.row.raw.get("__product_master_internal_formula_check__", ""),
            pr.row.raw.get("__product_master_master_formula__", ""),
            pr.row.raw.get("__product_master_candidate_count__", ""),
            pr.row.raw.get("__product_master_formula_note__", ""),
            pr.row.raw.get("__product_master_unmatched_reason__", ""), getattr(pr.row, "core_id", ""),
            getattr(pr.row, "structure_status", ""), getattr(pr.row, "formula_match", ""),
            "; ".join(pr.warnings + pr.row.warnings + (getattr(pr.row, "structure_warnings", []) or [])),
        ])

    audit_headers = [
        "Dataset", "Source_file", "Source_sheet", "Source_row", "Source_index", "Name", "Formula", "Combo",
        "Product_Master_Match", "Product_Master_Match_Method",
        "Product_Master_Formula_Key", "Product_Master_Used_Key",
        "Product_Master_Candidate_Count", "Product_Master_Verification_Mode",
        "Product_Master_Join_Status", "Product_Master_Formula_Flag", "Product_Master_Formula_Check",
        "Product_Master_Component_Formula_Check", "Product_Master_Internal_Formula_Check",
        "Product_Master_Target_Formula", "Product_Master_Master_Formula",
        "Product_Master_Master_Product_Formula", "Product_Master_Formula_Note",
        "Product_Master_Unmatched_Reason",
        "Core_ID", "Structure_status", "Structure_method", "Structure_hash", "Product_formula_calc", "Formula_match", "3D_status",
        "Descriptor_count", "Product_fingerprint", "Component_fingerprint_count",
    ]
    if not privacy_mode:
        audit_headers.append("Product_SMILES")
    audit_headers.append("Warnings")
    ws_struct.append(audit_headers)
    for dataset, rows in (("Training", train_rows), ("Target", target_rows)):
        for r in rows:
            a = structure_audit_row(r, privacy_mode=privacy_mode)
            ws_struct.append([dataset] + [a.get(h, "") for h in audit_headers[1:]])

    join_headers = [
        "Dataset", "Role", "Component_Formula", "Formula_Key_Token", "Rows", "Formula_key_candidate_found",
        "Structure_join_matched", "Match_rate_%", "Formula_flagged",
        "Formula_flag_rate_%", "Product_formula_no", "Component_formula_no",
        "Master_internal_formula_no", "Product_structure_ready",
        "Top_unmatched_reason", "Likely_issue",
    ]
    ws_join.append(join_headers)
    for item in _join_component_summary(train_rows, "Training") + _join_component_summary(target_rows, "Target"):
        ws_join.append([item.get(h, "") for h in join_headers])

    lib_headers = [
        "Role", "ID", "Core_ID", "Formula", "Net_Formula", "Has_SMILES", "Has_Product_SMILES", "Combo", "Numeric_descriptor_count",
        "Source_file", "Source_sheet", "Source_row", "Notes",
    ]
    if not privacy_mode:
        lib_headers.extend(["SMILES", "Product_SMILES"])
    ws_lib.append(lib_headers)
    for key in sorted(structure_library.entries.keys()):
        e = structure_library.entries[key]
        vals = [
            e.role, e.entry_id, e.core_id, e.formula, e.net_formula, "Yes" if e.smiles else "No", "Yes" if e.product_smiles else "No",
            e.combo, len(e.numeric), e.source_file, e.source_sheet, e.source_row, e.notes,
        ]
        if not privacy_mode:
            vals.extend([e.smiles, e.product_smiles])
        ws_lib.append(vals)

    rule_headers = ["Rule_ID", "Core_ID", "Reactant_Order", "Has_Reaction_SMARTS", "A_Loss", "B_Loss", "C_Loss", "Enabled", "Notes"]
    if not privacy_mode:
        rule_headers.append("Reaction_SMARTS")
    ws_rules.append(rule_headers)
    for core_id in sorted(structure_library.rules.keys()):
        rule = structure_library.rules[core_id]
        vals = [
            rule.rule_id, rule.core_id, ",".join(rule.reactant_order), "Yes" if rule.reaction_smarts else "No",
            rule.a_loss, rule.b_loss, rule.c_loss, "Yes" if rule.enabled else "No", rule.notes,
        ]
        if not privacy_mode:
            vals.append(rule.reaction_smarts)
        ws_rules.append(vals)

    ws_readme.append(["Item", "Description"])
    lines = [
        ("Name", "CombiTrace-IE structure-aware local response-factor transfer model."),
        ("Core formula", "RF = actual_concentration / measured_ratio; predicted_concentration = measured_ratio * transferred_RF."),
        ("Model", model_note),
        ("Fixed core", "Each A/B/C fragment is linked to the fixed middle structure through Core_ID."),
        ("Reaction stoichiometry", "Default formula validation: product = A + B + C + CORE contribution - N2 - H. Use Net_Formula when the isolated CORE formula differs from its net elemental contribution."),
        ("Product structure", "Priority: row Product_SMILES > PRODUCT mapping > private Reaction SMARTS > role-aware A/B/C/CORE feature fusion."),
        ("Privacy", "When Privacy mode is enabled, raw SMILES and Reaction SMARTS are not written to the result workbook."),
        ("Interpretation", "This remains exploratory semi-quantitation and does not replace authentic-standard calibration."),
        ("Structure audit", "Structure_Coverage reports formula validation, structure method and failure reasons for every row."),
        ("Structure join summary", "Structure_Join_Summary groups join status by A/B/C component Formula; one missing or inconsistent component Formula in a 10\u00D710\u00D710 library often appears as a 100-row slice."),
        ("Formula-key join", "A/B/C component formulas are the authoritative join identity. IDs and A#/B#/C# numbers are ignored. The final product Formula is retained as review-only or strict QC, and ambiguous repeated Formula keys are never chosen arbitrarily."),
    ]
    for k, v in lines:
        ws_readme.append([k, v])

    plots = _make_plots(cv_rows, out_xlsx)
    if plots:
        ws_plot = wb.create_sheet("Diagnostic_Plots")
        ws_plot.append(["Plot", "PNG file"])
        for title, path in plots:
            ws_plot.append([title, str(path)])
        try:
            from openpyxl.drawing.image import Image as XLImage
            anchors = ["A4", "K4", "A32", "K32"]
            for (title, path), anchor in zip(plots, anchors):
                img = XLImage(str(path))
                img.width = 520
                img.height = 380
                ws_plot.add_image(img, anchor)
        except Exception:
            pass

    _style_workbook(wb)
    try:
        wb.active = wb.sheetnames.index("Predictions")
    except Exception:
        pass
    wb.save(str(out_xlsx))

    out_csv: Optional[Path] = None
    if output_csv:
        out_csv = out_xlsx.with_suffix(".csv")
        with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(pred_headers)
            w.writerows(csv_rows)
    return out_xlsx, out_csv


# -----------------------------------------------------------------------------
# Public entry point
# -----------------------------------------------------------------------------


def run_combitrace_ie(
    calibration_file: Path,
    target_file: Path,
    out_xlsx: Path,
    *,
    structure_file: Optional[Path] = None,
    product_master_file: Optional[Path] = None,
    product_master_sheet: str = "",
    product_master_combo_col: str = "",
    product_master_smiles_col: str = "",
    calibration_sheet: str = "",
    target_sheet: str = "",
    structure_sheet: str = "",
    reaction_sheet: str = "",
    combo_col: str = "",
    formula_col: str = "",
    ratio_col: str = "",
    concentration_col: str = "",
    target_ratio_col: str = "",
    product_smiles_col: str = "",
    formula_verification_mode: str = "review",
    k_neighbors: int = 15,
    ridge_lambda: float = 10.0,
    cat_dim: int = 96,
    use_rt: bool = True,
    use_3d: bool = False,
    privacy_mode: bool = True,
    structure_threshold: float = 0.60,
    output_csv: bool = True,
) -> RunReport:
    warnings: List[str] = []
    formula_verification_mode = str(formula_verification_mode or "review").strip().lower()
    if formula_verification_mode not in {"review", "strict"}:
        formula_verification_mode = "review"
    cal_table = read_table(Path(calibration_file), sheet_name=calibration_sheet)
    tar_table = read_table(Path(target_file), sheet_name=target_sheet)
    struct_lib = load_private_structure_library(
        Path(structure_file) if structure_file else None,
        structure_sheet=structure_sheet,
        reaction_sheet=reaction_sheet,
    )
    warnings.extend(struct_lib.warnings)
    product_master, master_warnings, master_meta = load_product_smiles_master(
        Path(product_master_file) if product_master_file else None,
        sheet_name=product_master_sheet,
        combo_col=product_master_combo_col,
        smiles_col=product_master_smiles_col,
    )
    warnings.extend(master_warnings)

    cal_cfg = guess_columns(cal_table, need_concentration=True)
    tar_cfg = guess_columns(tar_table, need_concentration=False)
    if combo_col:
        cal_cfg.combo_col = combo_col
        tar_cfg.combo_col = combo_col
    if formula_col:
        cal_cfg.formula_col = formula_col
        tar_cfg.formula_col = formula_col
    if ratio_col:
        cal_cfg.ratio_col = ratio_col
    if target_ratio_col:
        tar_cfg.ratio_col = target_ratio_col
    elif ratio_col:
        tar_cfg.ratio_col = ratio_col
    if concentration_col:
        cal_cfg.concentration_col = concentration_col

    train_rows, w1 = prepare_rows(cal_table, cal_cfg, is_training=True)
    target_rows, w2 = prepare_rows(tar_table, tar_cfg, is_training=False)
    warnings.extend(w1)
    warnings.extend(w2)

    master_match_train, master_fill_train = apply_product_smiles_master(
        train_rows, product_master, formula_verification_mode=formula_verification_mode
    )
    master_match_target, master_fill_target = apply_product_smiles_master(
        target_rows, product_master, formula_verification_mode=formula_verification_mode
    )
    master_methods_train: Dict[str, int] = {}
    master_methods_target: Dict[str, int] = {}
    for rr in train_rows:
        m = str(rr.raw.get("__product_master_match_method__", "") or "")
        master_methods_train[m] = master_methods_train.get(m, 0) + 1
    for rr in target_rows:
        m = str(rr.raw.get("__product_master_match_method__", "") or "")
        master_methods_target[m] = master_methods_target.get(m, 0) + 1
    if product_master:
        warnings = [
            ('No component library was supplied; complete structures from the product SMILES master table take priority.'
             if w == 'No structure library was supplied; using the Combo/formula/RT fallback.' else w)
            for w in warnings
        ]
        warnings.append(
            f"Combo-SMILES master: usable={int(master_meta.get('usable', 0) or 0)}, "
            f"training matched={master_match_train}/{len(train_rows)} "
            f"(A/B/C Formula={master_methods_train.get('abc_component_formulas', 0)}), "
            f"target matched={master_match_target}/{len(target_rows)} "
            f"(A/B/C Formula={master_methods_target.get('abc_component_formulas', 0)})"
        )
        warnings.append(
            "Structure join identity: ordered A/B/C component formulas. Component IDs and A#/B#/C# numbers are ignored. "
            + (
                "Final product Formula is reviewed and discrepancies are flagged without blocking a unique Formula-key join."
                if formula_verification_mode == "review"
                else "Strict mode rejects an explicit final product Formula discrepancy."
            )
        )
        if master_match_train == 0 or master_match_target == 0:
            warnings.append(
                "STRUCTURE JOIN FAILED for one or both tables. The model will not use Product_SMILES for unmatched rows. "
                "See Structure_Coverage and Structure_Join_Summary for the A/B/C Formula key and exact failure reason."
            )

    enrich_rows_with_structures(
        train_rows, struct_lib, use_3d=bool(use_3d), privacy_mode=bool(privacy_mode),
        product_smiles_col=product_smiles_col,
    )
    enrich_rows_with_structures(
        target_rows, struct_lib, use_3d=bool(use_3d), privacy_mode=bool(privacy_mode),
        product_smiles_col=product_smiles_col,
    )

    train_valid = valid_training_rows(train_rows)
    all_rows = list(train_rows) + list(target_rows)
    element_keys = _element_keys(all_rows)

    qc_raw, cv_rows, evidence_stats, intervals = loocv_local_rf(
        train_valid, element_keys, k=k_neighbors, use_rt=use_rt,
        structure_threshold=float(structure_threshold),
    )
    qc = evaluate_local_qc(qc_raw, len(train_valid))
    qc["RDKit_available"] = bool(struct_lib.rdkit_available)
    qc["Structure_rows_ready_training"] = sum(1 for r in train_rows if getattr(r, "structure_status", "") in {"product_structure_ready", "fragment_core_features_ready"})
    qc["Structure_rows_ready_target"] = sum(1 for r in target_rows if getattr(r, "structure_status", "") in {"product_structure_ready", "fragment_core_features_ready"})
    qc["Product_structures_training"] = sum(1 for r in train_rows if getattr(r, "structure_status", "") == "product_structure_ready")
    qc["Product_structures_target"] = sum(1 for r in target_rows if getattr(r, "structure_status", "") == "product_structure_ready")
    qc["Formula_mismatch_training"] = sum(1 for r in train_rows if getattr(r, "formula_match", "") == "no")
    qc["Formula_mismatch_target"] = sum(1 for r in target_rows if getattr(r, "formula_match", "") == "no")
    qc["Structure_similarity_threshold"] = float(structure_threshold)
    qc["3D_descriptors_enabled"] = bool(use_3d)
    qc["Privacy_mode"] = bool(privacy_mode)
    qc["Product_master_rows"] = int(master_meta.get("rows", 0) or 0)
    qc["Product_master_usable"] = int(master_meta.get("usable", 0) or 0)
    qc["Product_master_match_training"] = int(master_match_train)
    qc["Product_master_match_target"] = int(master_match_target)
    qc["Product_master_formula_key_match_training"] = int(master_methods_train.get("abc_component_formulas", 0))
    qc["Product_master_formula_key_match_target"] = int(master_methods_target.get("abc_component_formulas", 0))
    qc["Product_master_formula_verification_mode"] = formula_verification_mode
    qc["Product_master_formula_flagged_training"] = sum(1 for r in train_rows if _row_formula_flagged(r))
    qc["Product_master_formula_flagged_target"] = sum(1 for r in target_rows if _row_formula_flagged(r))
    qc["Product_master_ambiguous_training"] = sum(
        1 for r in train_rows if str(r.raw.get("__product_master_join_status__", "")).startswith("ambiguous")
    )
    qc["Product_master_ambiguous_target"] = sum(
        1 for r in target_rows if str(r.raw.get("__product_master_join_status__", "")).startswith("ambiguous")
    )
    # Backward-compatible names.  In review mode these are flagged rows rather
    # than rejected joins; the new *_flagged_* fields should be preferred.
    qc["Product_master_formula_rejected_training"] = (
        qc["Product_master_formula_flagged_training"] if formula_verification_mode == "strict" else 0
    )
    qc["Product_master_formula_rejected_target"] = (
        qc["Product_master_formula_flagged_target"] if formula_verification_mode == "strict" else 0
    )
    qc["Product_master_formula_keys"] = int(master_meta.get("formula_keys", 0) or 0)
    qc["Product_master_duplicate_formula_keys"] = int(master_meta.get("duplicate_formula_keys", 0) or 0)
    qc["Product_master_lookup_keys"] = int(master_meta.get("lookup_keys", 0) or 0)
    qc["Product_master_filled_training"] = int(master_fill_train)
    qc["Product_master_filled_target"] = int(master_fill_target)

    model = fit_local_rf_model(
        train_valid, element_keys, k_neighbors=k_neighbors,
        evidence_interval_log=intervals, evidence_stats=evidence_stats,
        use_rt=use_rt, structure_threshold=float(structure_threshold),
    )
    preds = [predict_one_local_rf(r, model) for r in target_rows]
    n_pred = sum(1 for p in preds if p.pred_conc is not None)
    n_train_skipped = sum(1 for r in train_rows if training_skip_reason(r))
    n_target_skipped = sum(1 for p in preds if p.pred_conc is None)
    warnings.append(f"Training row diagnostics: used={len(train_valid)}, skipped={n_train_skipped}; see Training_Row_Check")
    warnings.append(f"Target row diagnostics: predicted={n_pred}, skipped={n_target_skipped}; see Target_Row_Check")
    warnings.append(
        f"Structure coverage: training={qc['Structure_rows_ready_training']}/{len(train_rows)}, "
        f"target={qc['Structure_rows_ready_target']}/{len(target_rows)}; see Structure_Coverage"
    )
    if qc.get("Product_master_formula_flagged_training", 0) or qc.get("Product_master_formula_flagged_target", 0):
        action_note = (
            "Rows with a unique A/B/C Formula key remain joined; inspect Structure_Join_Summary before interpreting the model."
            if formula_verification_mode == "review"
            else "Strict mode rejects explicitly inconsistent candidates; inspect Structure_Join_Summary before interpreting the model."
        )
        warnings.append(
            f"Formula review flagged training={qc.get('Product_master_formula_flagged_training', 0)}, "
            f"target={qc.get('Product_master_formula_flagged_target', 0)} rows. " + action_note
        )
    # Point to component-level patterns such as one bad A/B/C entry affecting 100 combinations.
    for item in _join_component_summary(target_rows, "Target") + _join_component_summary(train_rows, "Training"):
        if item.get("Likely_issue"):
            warnings.append(str(item["Likely_issue"]))
    if qc["Formula_mismatch_training"] or qc["Formula_mismatch_target"]:
        warnings.append('Some structure-derived formulas differ from enumerated formulas. The default stoichiometric check is A+B+C+CORE-N2-H.')

    model_note = (
        "Structure-aware local RF transfer using optional Combo-to-Product_SMILES master: ordered A/B/C component-Formula match > identical product structure > high structural similarity > "
        "same component pair > nearest structure > same component > formula/RT neighbours > global median RF. "
        f"Fixed core is linked by Core_ID; default stoichiometry A+B+C+CORE-N2-H; "
        f"final product Formula verification={formula_verification_mode}; K={k_neighbors}; threshold={float(structure_threshold):.2f}; "
        f"3D={bool(use_3d)}; privacy={bool(privacy_mode)}."
    )
    out_xlsx, out_csv = write_local_rf_workbook(
        Path(out_xlsx), preds, train_rows, qc, cv_rows, evidence_stats,
        target_rows=target_rows, structure_library=struct_lib, privacy_mode=bool(privacy_mode),
        model_note=model_note, output_csv=output_csv,
    )
    return RunReport(
        output_xlsx=out_xlsx,
        output_csv=out_csv,
        n_train_total=len(train_rows),
        n_train_used=len(train_valid),
        n_target_total=len(target_rows),
        n_predicted=n_pred,
        warnings=warnings,
        qc=qc,
    )

