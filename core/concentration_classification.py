"""Ten-level concentration classification for known standards and unknown targets.

This branch is deliberately separate from continuous response-factor regression.
It asks a simpler question: after using the measured area/internal-standard ratio
and molecular/mobile-phase descriptors, can the model place each compound into
one of ten empirical concentration levels?

All reported standard-set predictions are out-of-fold (OOF).  The unknown sample
is never used to define labels or tune the classifier.
"""
from __future__ import annotations

import math
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np

SKLEARN_AVAILABLE = False
SKLEARN_ERROR = ""
try:
    from sklearn.base import clone
    from sklearn.ensemble import ExtraTreesClassifier, ExtraTreesRegressor
    from sklearn.feature_selection import SelectKBest, f_classif, f_regression
    from sklearn.impute import SimpleImputer
    from sklearn.metrics import (
        accuracy_score,
        balanced_accuracy_score,
        cohen_kappa_score,
        confusion_matrix,
        f1_score,
    )
    from sklearn.model_selection import RepeatedStratifiedKFold, LeaveOneGroupOut
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import StandardScaler
    from sklearn.svm import SVC, SVR
    SKLEARN_AVAILABLE = True
except Exception as exc:  # pragma: no cover
    SKLEARN_ERROR = str(exc)


@dataclass
class ConcentrationLevelResult:
    enabled: bool
    best_model: str = ""
    decision: str = ""
    decision_summary: str = ""
    level_count_requested: int = 10
    level_count_observed: int = 0
    level_definitions: List[Dict[str, object]] = field(default_factory=list)
    comparison: List[Dict[str, object]] = field(default_factory=list)
    known_predictions: List[Dict[str, object]] = field(default_factory=list)
    target_predictions: List[Dict[str, object]] = field(default_factory=list)
    confusion_rows: List[Dict[str, object]] = field(default_factory=list)
    group_summary: List[Dict[str, object]] = field(default_factory=list)
    feature_rows: List[Dict[str, object]] = field(default_factory=list)
    plot_paths: List[Path] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    concentration_unit: str = ""


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _num(value: object) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def _identity(record: Dict[str, object]) -> str:
    parts = [
        _text(record.get("Source_file")),
        _text(record.get("Source_sheet")),
        _text(record.get("Source_row")),
    ]
    if any(parts):
        return "|".join(parts)
    return _text(record.get("Product_Master_Formula_Key") or record.get("ABC_Formula_Key") or record.get("Combo"))


def _spearman(x: Sequence[float], y: Sequence[float]) -> float:
    a = np.asarray(x, dtype=float)
    b = np.asarray(y, dtype=float)
    if len(a) < 2 or float(np.std(a)) <= 0 or float(np.std(b)) <= 0:
        return float("nan")
    def ranks(v: np.ndarray) -> np.ndarray:
        order = np.argsort(v, kind="mergesort")
        out = np.empty(len(v), dtype=float)
        i = 0
        while i < len(v):
            j = i + 1
            while j < len(v) and v[order[j]] == v[order[i]]:
                j += 1
            out[order[i:j]] = (i + 1 + j) / 2.0
            i = j
        return out
    return float(np.corrcoef(ranks(a), ranks(b))[0, 1])


def _assign_empirical_levels(values: Sequence[float], n_levels: int = 10) -> Tuple[np.ndarray, List[Dict[str, object]], Dict[int, float]]:
    """Assign approximately equal-frequency ordinal levels while keeping ties together."""
    arr = np.asarray(values, dtype=float)
    n = len(arr)
    order = np.argsort(arr, kind="mergesort")
    raw = np.empty(n, dtype=int)
    for rank0, idx in enumerate(order.tolist()):
        raw[int(idx)] = min(n_levels, max(1, int(math.ceil((rank0 + 1) * n_levels / max(1, n)))))

    # Keep equal concentrations in the same level.  The median rank-bin of the
    # tie group is used; this may leave an empty level, which is preferable to
    # assigning identical concentrations to different classes.
    levels = raw.copy()
    for value in np.unique(arr):
        idx = np.where(arr == value)[0]
        levels[idx] = int(round(float(np.median(raw[idx]))))
    levels = np.clip(levels, 1, n_levels)

    definitions: List[Dict[str, object]] = []
    centers: Dict[int, float] = {}
    for level in range(1, n_levels + 1):
        vals = arr[levels == level]
        if len(vals):
            center = float(np.median(vals))
            centers[level] = center
            definitions.append({
                "Level": level,
                "Label": f"L{level:02d}",
                "N": int(len(vals)),
                "Minimum_concentration": float(np.min(vals)),
                "Median_concentration": center,
                "Maximum_concentration": float(np.max(vals)),
                "Definition": "empirical equal-frequency level; identical concentrations kept together",
            })
        else:
            definitions.append({
                "Level": level,
                "Label": f"L{level:02d}",
                "N": 0,
                "Minimum_concentration": "",
                "Median_concentration": "",
                "Maximum_concentration": "",
                "Definition": "empty level in this calibration set",
            })
    # Fill empty representative centers by interpolation in log space.
    known = sorted(centers)
    if known:
        for level in range(1, n_levels + 1):
            if level in centers:
                continue
            lower = max((k for k in known if k < level), default=None)
            upper = min((k for k in known if k > level), default=None)
            if lower is not None and upper is not None:
                frac = (level - lower) / (upper - lower)
                centers[level] = 10.0 ** ((1-frac) * math.log10(centers[lower]) + frac * math.log10(centers[upper]))
            elif lower is not None:
                centers[level] = centers[lower]
            elif upper is not None:
                centers[level] = centers[upper]
    return levels.astype(int), definitions, centers


def _available_numeric_features(
    calibration: Sequence[Dict[str, object]],
    descriptor_names: Sequence[str],
    *,
    manual_feature_pool: Sequence[str] = (),
    selected_feature_names: Sequence[str] = (),
) -> List[str]:
    blocked = {
        "Actual_concentration", "Measured_ratio", "Log_Measured_Ratio",
        "Predicted_concentration", "Estimated_concentration",
    }
    if manual_feature_pool:
        candidates = list(manual_feature_pool)
    elif selected_feature_names:
        candidates = list(selected_feature_names)
    else:
        candidates = list(descriptor_names)
    out: List[str] = []
    seen = set()
    minimum = max(5, int(math.ceil(len(calibration) * 0.50)))
    for name in candidates:
        name = str(name or "").strip()
        if not name or name in seen or name in blocked or name.startswith("_"):
            continue
        seen.add(name)
        vals = [_num(r.get(name)) for r in calibration]
        finite = [x for x in vals if x is not None]
        if len(finite) < minimum:
            continue
        if float(np.var(np.asarray(finite, dtype=float))) <= 1e-14:
            continue
        out.append(name)
    return out


def _matrix(records: Sequence[Dict[str, object]], features: Sequence[str]) -> np.ndarray:
    rows: List[List[float]] = []
    for record in records:
        ratio = _num(record.get("Measured_ratio"))
        row = [math.log10(ratio) if ratio is not None and ratio > 0 else np.nan]
        for name in features:
            value = _num(record.get(name))
            row.append(np.nan if value is None else float(value))
        rows.append(row)
    return np.asarray(rows, dtype=float)


def _build_models(n_features: int, k: int, random_state: int):
    k = max(1, min(int(k), int(n_features)))
    return {
        "SVC_RBF": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("scale", StandardScaler()),
            ("select", SelectKBest(score_func=f_classif, k=k)),
            ("model", SVC(C=4.0, gamma="scale", class_weight="balanced")),
        ]),
        "ExtraTreesClassifier": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("select", SelectKBest(score_func=f_classif, k=k)),
            ("model", ExtraTreesClassifier(
                n_estimators=250, max_features="sqrt", min_samples_leaf=2,
                class_weight="balanced", random_state=int(random_state), n_jobs=-1,
            )),
        ]),
        "OrdinalExtraTrees": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("select", SelectKBest(score_func=f_regression, k=k)),
            ("model", ExtraTreesRegressor(
                n_estimators=250, max_features="sqrt", min_samples_leaf=2,
                random_state=int(random_state), n_jobs=-1,
            )),
        ]),
        "OrdinalSVR": Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("scale", StandardScaler()),
            ("select", SelectKBest(score_func=f_regression, k=k)),
            ("model", SVR(C=8.0, gamma="scale", epsilon=0.15)),
        ]),
    }


def _predict_levels(model_name: str, estimator, X: np.ndarray, n_levels: int) -> np.ndarray:
    pred = np.asarray(estimator.predict(X), dtype=float)
    if model_name.startswith("Ordinal"):
        pred = np.rint(pred)
    return np.clip(pred.astype(int), 1, n_levels)


def _level_metrics(actual: np.ndarray, predicted: np.ndarray) -> Dict[str, object]:
    actual = np.asarray(actual, dtype=int)
    predicted = np.asarray(predicted, dtype=int)
    diff = np.abs(predicted - actual)
    exact = float(np.mean(diff == 0) * 100.0)
    within1 = float(np.mean(diff <= 1) * 100.0)
    within2 = float(np.mean(diff <= 2) * 100.0)
    macro_f1 = float(f1_score(actual, predicted, average="macro", zero_division=0))
    weighted_f1 = float(f1_score(actual, predicted, average="weighted", zero_division=0))
    balanced = float(balanced_accuracy_score(actual, predicted))
    qwk = float(cohen_kappa_score(actual, predicted, weights="quadratic"))
    rho = _spearman(actual, predicted)
    mae = float(np.mean(diff))
    score = (
        exact / 100.0 * 0.45
        + within1 / 100.0 * 0.25
        + weighted_f1 * 0.12
        + max(-1.0, min(1.0, 0.0 if not math.isfinite(rho) else rho)) * 0.10
        + max(-1.0, min(1.0, 0.0 if not math.isfinite(qwk) else qwk)) * 0.08
    )
    return {
        "N": int(len(actual)),
        "Exact_hits_N": int(np.sum(diff == 0)),
        "Exact_level_accuracy_pct": exact,
        "Within_1_level_pct": within1,
        "Within_2_levels_pct": within2,
        "Mean_absolute_level_error": mae,
        "Macro_F1": macro_f1,
        "Weighted_F1": weighted_f1,
        "Balanced_accuracy": balanced,
        "Quadratic_weighted_kappa": qwk,
        "Level_Spearman": rho,
        "All_exactly_hit": bool(np.all(diff == 0)),
        "Selection_score": score,
    }


def _majority_vote(values: Sequence[int]) -> int:
    counts: Dict[int, int] = {}
    for value in values:
        counts[int(value)] = counts.get(int(value), 0) + 1
    best_count = max(counts.values())
    tied = sorted(k for k, v in counts.items() if v == best_count)
    return int(round(float(np.median(tied))))


def _make_plots(
    output_xlsx: Path,
    known_rows: Sequence[Dict[str, object]],
    confusion: np.ndarray,
    n_levels: int,
    best_model: str,
    language: str,
) -> Tuple[List[Path], List[str]]:
    paths: List[Path] = []
    warnings: List[str] = []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception as exc:
        return [], [f"Level-classification plots skipped: {exc}"]
    plot_dir = Path(output_xlsx).with_name("_ctlevel_" + str(abs(hash(str(output_xlsx))))[-8:])
    plot_dir.mkdir(parents=True, exist_ok=True)
    zh = str(language or "en").lower().startswith("zh")

    try:
        fig, ax = plt.subplots(figsize=(7.4, 6.4))
        image = ax.imshow(confusion, aspect="auto")
        ax.set_xticks(np.arange(n_levels))
        ax.set_yticks(np.arange(n_levels))
        ax.set_xticklabels([f"L{i:02d}" for i in range(1, n_levels + 1)])
        ax.set_yticklabels([f"L{i:02d}" for i in range(1, n_levels + 1)])
        ax.set_xlabel("Predicted level" if not zh else 'Predicted level')
        ax.set_ylabel("Actual level" if not zh else 'Actual level')
        ax.set_title((best_model + ": ten-level confusion matrix") if not zh else (best_model + ': ten-level concentration confusion matrix'))
        for i in range(n_levels):
            for j in range(n_levels):
                ax.text(j, i, str(int(confusion[i, j])), ha="center", va="center")
        fig.colorbar(image, ax=ax, label="Count" if not zh else 'Count')
        fig.tight_layout()
        path = plot_dir / "30_level_confusion.png"
        fig.savefig(path, dpi=190, bbox_inches="tight")
        plt.close(fig)
        paths.append(path)
    except Exception as exc:
        warnings.append(f"Confusion-matrix plot skipped: {exc}")

    try:
        rows = sorted(known_rows, key=lambda r: float(r.get("Actual_concentration", 0.0)))
        x = np.arange(1, len(rows) + 1)
        actual = np.asarray([float(r["Actual_level_1_to_10"]) for r in rows], dtype=float)
        predicted = np.asarray([float(r["Predicted_level_1_to_10"]) for r in rows], dtype=float)
        fig, ax = plt.subplots(figsize=(max(12.0, len(rows) * 0.16), 5.8))
        ax.scatter(x, actual, label="Actual level" if not zh else 'Actual level', s=22)
        ax.scatter(x, predicted, label="Predicted level" if not zh else 'Predicted level', s=22)
        ax.set_ylim(0.5, n_levels + 0.5)
        ax.set_yticks(range(1, n_levels + 1))
        ax.set_xlabel("Known standards sorted by actual concentration" if not zh else 'Known standards sorted by actual concentration')
        ax.set_ylabel("Concentration level" if not zh else 'Concentration level')
        ax.set_title((best_model + ": actual and OOF-predicted levels") if not zh else (best_model + ': actual and out-of-fold concentration levels'))
        ax.legend()
        ax.grid(axis="y", alpha=0.25)
        fig.tight_layout()
        path = plot_dir / "31_level_actual_vs_predicted.png"
        fig.savefig(path, dpi=190, bbox_inches="tight")
        plt.close(fig)
        paths.append(path)
    except Exception as exc:
        warnings.append(f"Actual-vs-predicted level plot skipped: {exc}")
    return paths, warnings


def run_concentration_level_classification(
    calibration_records: Sequence[Dict[str, object]],
    target_records: Sequence[Dict[str, object]],
    descriptor_names: Sequence[str],
    output_xlsx: Path,
    *,
    selected_feature_names: Sequence[str] = (),
    manual_feature_pool: Sequence[str] = (),
    cv_splits: int = 5,
    cv_repeats: int = 3,
    max_features: int = 24,
    random_state: int = 42,
    level_count: int = 10,
    output_language: str = "en",
    concentration_unit: str = "",
) -> ConcentrationLevelResult:
    if not SKLEARN_AVAILABLE:
        raise RuntimeError("Ten-level classification needs scikit-learn: " + SKLEARN_ERROR)

    calibration: List[Dict[str, object]] = []
    for row in calibration_records:
        ratio = _num(row.get("Measured_ratio"))
        concentration = _num(row.get("Actual_concentration"))
        if ratio is not None and ratio > 0 and concentration is not None and concentration > 0:
            calibration.append(dict(row))
    if len(calibration) < 30:
        raise ValueError(f"Ten-level classification requires at least 30 valid standards; got {len(calibration)}")

    features = _available_numeric_features(
        calibration, descriptor_names,
        manual_feature_pool=manual_feature_pool,
        selected_feature_names=selected_feature_names,
    )
    if len(features) < 2:
        raise ValueError("Too few numeric descriptors for ten-level classification")

    concentrations = np.asarray([float(r["Actual_concentration"]) for r in calibration], dtype=float)
    y, level_definitions, centers = _assign_empirical_levels(concentrations, int(level_count))
    observed_classes, counts = np.unique(y, return_counts=True)
    min_class_count = int(np.min(counts)) if len(counts) else 0
    warnings: List[str] = []
    if len(observed_classes) < int(level_count):
        warnings.append(
            f"Only {len(observed_classes)} of {int(level_count)} empirical levels are populated because tied concentrations were kept together."
        )
    if min_class_count < 2:
        raise ValueError("At least one concentration level contains fewer than two standards; classification CV is not possible")

    X = _matrix(calibration, features)
    n_splits = max(2, min(int(cv_splits), min_class_count))
    repeats = max(1, int(cv_repeats))
    k = max(2, min(int(max_features), X.shape[1]))
    models = _build_models(X.shape[1], k, int(random_state))
    splitter = RepeatedStratifiedKFold(n_splits=n_splits, n_repeats=repeats, random_state=int(random_state))

    model_predictions: Dict[str, Dict[int, List[int]]] = {
        name: {i: [] for i in range(len(calibration))} for name in models
    }
    failures: Dict[str, int] = {name: 0 for name in models}
    for split_no, (train_idx, test_idx) in enumerate(splitter.split(X, y), start=1):
        for model_no, (name, base) in enumerate(models.items(), start=1):
            try:
                est = clone(base)
                # Make stochastic tree models reproducible but different across folds.
                try:
                    est.set_params(model__random_state=int(random_state) + split_no * 101 + model_no)
                except Exception:
                    pass
                est.fit(X[train_idx], y[train_idx])
                pred = _predict_levels(name, est, X[test_idx], int(level_count))
                for pos, idx in enumerate(test_idx.tolist()):
                    model_predictions[name][int(idx)].append(int(pred[pos]))
            except Exception as exc:
                failures[name] += 1
                if failures[name] <= 2:
                    warnings.append(f"{name} classification fold failed: {exc}")

    comparison: List[Dict[str, object]] = []
    aggregated_by_model: Dict[str, np.ndarray] = {}
    for name in models:
        pred = np.asarray([
            _majority_vote(model_predictions[name][i]) if model_predictions[name][i] else 0
            for i in range(len(calibration))
        ], dtype=int)
        valid = pred > 0
        if int(np.sum(valid)) == 0:
            continue
        metrics = _level_metrics(y[valid], pred[valid])
        # Full-fit accuracy is retained only as a separability/memorisation diagnostic.
        # The OOF metrics above remain the validation result used for decisions.
        full_fit_exact = float("nan")
        full_fit_within1 = float("nan")
        try:
            full_est = clone(models[name])
            full_est.fit(X, y)
            full_pred = _predict_levels(name, full_est, X, int(level_count))
            full_diff = np.abs(full_pred - y)
            full_fit_exact = float(np.mean(full_diff == 0) * 100.0)
            full_fit_within1 = float(np.mean(full_diff <= 1) * 100.0)
        except Exception as exc:
            warnings.append(f"{name} full-fit diagnostic failed: {exc}")
        metrics.update({
            "Model": name,
            "CV_scheme": f"Repeated stratified {n_splits}-fold x {repeats}",
            "Feature_candidates_including_log_ratio": int(X.shape[1]),
            "SelectKBest_k": int(k),
            "CV_prediction_coverage_pct": float(np.mean(valid) * 100.0),
            "Failed_folds": int(failures.get(name, 0)),
            "Full_fit_exact_accuracy_pct_diagnostic_only": full_fit_exact,
            "Full_fit_within_1_level_pct_diagnostic_only": full_fit_within1,
        })
        comparison.append(metrics)
        aggregated_by_model[name] = pred
    if not comparison:
        raise RuntimeError("All ten-level classification models failed")
    comparison.sort(key=lambda r: (
        -float(r.get("Selection_score", -1e9)),
        -float(r.get("Exact_level_accuracy_pct", 0.0)),
        -float(r.get("Within_1_level_pct", 0.0)),
    ))
    for rank, row in enumerate(comparison, start=1):
        row["Rank"] = rank
    best_model = str(comparison[0]["Model"])
    best_pred = aggregated_by_model[best_model]
    valid = best_pred > 0
    best_metrics = _level_metrics(y[valid], best_pred[valid])

    if bool(best_metrics.get("All_exactly_hit")):
        decision = "LEVEL-PERFECT"
    elif float(best_metrics["Exact_level_accuracy_pct"]) >= 60.0 and float(best_metrics["Within_1_level_pct"]) >= 90.0:
        decision = "LEVEL-GO"
    elif float(best_metrics["Exact_level_accuracy_pct"]) >= 25.0 and float(best_metrics["Within_1_level_pct"]) >= 65.0:
        decision = "LEVEL-CAUTION"
    else:
        decision = "LEVEL-NO-GO"
    summary = (
        f"{best_model}: exact level hits={best_metrics['Exact_hits_N']}/{best_metrics['N']} "
        f"({best_metrics['Exact_level_accuracy_pct']:.1f}%), within +/-1 level={best_metrics['Within_1_level_pct']:.1f}%, "
        f"mean absolute level error={best_metrics['Mean_absolute_level_error']:.2f}, "
        f"quadratic weighted kappa={best_metrics['Quadratic_weighted_kappa']:.3f}."
    )

    known_rows: List[Dict[str, object]] = []
    for i, record in enumerate(calibration):
        actual_level = int(y[i])
        predicted_level = int(best_pred[i]) if int(best_pred[i]) > 0 else 0
        representative = centers.get(predicted_level, "") if predicted_level else ""
        known_rows.append({
            "Order_by_actual_concentration": 0,
            "Name": record.get("Name", ""),
            "Formula": record.get("Formula", ""),
            "Combo": record.get("Combo", ""),
            "ABC_Formula_Key": record.get("Product_Master_Formula_Key", record.get("ABC_Formula_Key", "")),
            "Injection_Group": record.get("Injection_Group", ""),
            "Actual_concentration": float(record["Actual_concentration"]),
            "Concentration_unit": str(concentration_unit or ""),
            "Actual_level_1_to_10": actual_level,
            "Actual_level_label": f"L{actual_level:02d}",
            "Predicted_level_1_to_10": predicted_level if predicted_level else "",
            "Predicted_level_label": f"L{predicted_level:02d}" if predicted_level else "",
            "Predicted_level_representative_concentration": representative,
            "Exact_level_hit": "Yes" if predicted_level == actual_level else "No",
            "Level_difference": (predicted_level - actual_level) if predicted_level else "",
            "Absolute_level_error": abs(predicted_level - actual_level) if predicted_level else "",
            "Within_1_level": "Yes" if predicted_level and abs(predicted_level - actual_level) <= 1 else "No",
            "Level_model": best_model,
            "Prediction_type": "out-of-fold empirical concentration-level classification",
            "Source_file": record.get("Source_file", ""),
            "Source_sheet": record.get("Source_sheet", ""),
            "Source_row": record.get("Source_row", ""),
            "Record_ID": _identity(record),
        })
    known_rows.sort(key=lambda r: (float(r["Actual_concentration"]), _text(r.get("Name") or r.get("Formula") or r.get("Combo"))))
    for i, row in enumerate(known_rows, start=1):
        row["Order_by_actual_concentration"] = i

    target_valid: List[Dict[str, object]] = []
    target_invalid: List[Dict[str, object]] = []
    for record in target_records:
        ratio = _num(record.get("Measured_ratio"))
        if ratio is not None and ratio > 0:
            target_valid.append(dict(record))
        else:
            target_invalid.append(dict(record))
    fitted = clone(models[best_model])
    fitted.fit(X, y)
    target_rows: List[Dict[str, object]] = []
    if target_valid:
        X_target = _matrix(target_valid, features)
        target_levels = _predict_levels(best_model, fitted, X_target, int(level_count))
        for record, level in zip(target_valid, target_levels.tolist()):
            representative = centers.get(int(level), "")
            target_rows.append({
                "Name": record.get("Name", ""),
                "Formula": record.get("Formula", ""),
                "Combo": record.get("Combo", ""),
                "ABC_Formula_Key": record.get("Product_Master_Formula_Key", record.get("ABC_Formula_Key", "")),
                "Measured_ratio": record.get("Measured_ratio", ""),
                "Predicted_level_1_to_10": int(level),
                "Predicted_level_label": f"L{int(level):02d}",
                "Predicted_level_representative_concentration": representative,
                "Concentration_unit": str(concentration_unit or ""),
                "Level_model": best_model,
                "Prediction_status": "PREDICTED",
                "Source_file": record.get("Source_file", ""),
                "Source_sheet": record.get("Source_sheet", ""),
                "Source_row": record.get("Source_row", ""),
                "Record_ID": _identity(record),
            })
    for record in target_invalid:
        target_rows.append({
            "Name": record.get("Name", ""),
            "Formula": record.get("Formula", ""),
            "Combo": record.get("Combo", ""),
            "ABC_Formula_Key": record.get("Product_Master_Formula_Key", record.get("ABC_Formula_Key", "")),
            "Measured_ratio": record.get("Measured_ratio", ""),
            "Predicted_level_1_to_10": "",
            "Predicted_level_label": "",
            "Predicted_level_representative_concentration": "",
            "Concentration_unit": str(concentration_unit or ""),
            "Level_model": best_model,
            "Prediction_status": "SKIPPED: missing or nonpositive measured ratio",
            "Source_file": record.get("Source_file", ""),
            "Source_sheet": record.get("Source_sheet", ""),
            "Source_row": record.get("Source_row", ""),
            "Record_ID": _identity(record),
        })
    target_rows.sort(key=lambda r: (
        _text(r.get("Source_file")), _text(r.get("Source_sheet")),
        float(_num(r.get("Source_row")) or 1e18),
    ))

    matrix = confusion_matrix(y[valid], best_pred[valid], labels=list(range(1, int(level_count) + 1)))
    confusion_rows: List[Dict[str, object]] = []
    for i in range(int(level_count)):
        row = {"Actual_level": f"L{i+1:02d}"}
        for j in range(int(level_count)):
            row[f"Predicted_L{j+1:02d}"] = int(matrix[i, j])
        confusion_rows.append(row)

    group_summary: List[Dict[str, object]] = []
    group_values = np.asarray([_text(r.get("Injection_Group")) for r in calibration], dtype=object)
    groups = sorted({x for x in group_values.tolist() if x})
    if len(groups) >= 2:
        group_pred = np.zeros(len(calibration), dtype=int)
        logo = LeaveOneGroupOut()
        for train_idx, test_idx in logo.split(X, y, groups=group_values):
            try:
                est = clone(models[best_model])
                est.fit(X[train_idx], y[train_idx])
                group_pred[test_idx] = _predict_levels(best_model, est, X[test_idx], int(level_count))
            except Exception as exc:
                warnings.append(f"Level group-out fold failed: {exc}")
        for group in groups + ["ALL_GROUPS"]:
            idx = np.where(group_values == group)[0] if group != "ALL_GROUPS" else np.where(group_pred > 0)[0]
            idx = idx[group_pred[idx] > 0]
            if len(idx):
                metrics = _level_metrics(y[idx], group_pred[idx])
                metrics["Injection_Group"] = group
                group_summary.append(metrics)

    feature_rows = [{
        "Feature": "Log_Measured_Ratio",
        "Role": "mandatory analytical signal",
        "Used": "Yes",
    }] + [{
        "Feature": name,
        "Role": "molecular/mobile-phase descriptor candidate",
        "Used": "Yes",
    } for name in features]

    plot_paths, plot_warnings = _make_plots(
        Path(output_xlsx), known_rows, matrix, int(level_count), best_model, str(output_language or "en"),
    )
    warnings.extend(plot_warnings)

    return ConcentrationLevelResult(
        enabled=True,
        best_model=best_model,
        decision=decision,
        decision_summary=summary,
        level_count_requested=int(level_count),
        level_count_observed=int(len(observed_classes)),
        level_definitions=level_definitions,
        comparison=comparison,
        known_predictions=known_rows,
        target_predictions=target_rows,
        confusion_rows=confusion_rows,
        group_summary=group_summary,
        feature_rows=feature_rows,
        plot_paths=plot_paths,
        warnings=warnings,
        concentration_unit=str(concentration_unit or ""),
    )


def _append_dict_sheet(wb, title: str, rows: Sequence[Dict[str, object]]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    if rows:
        headers: List[str] = []
        seen = set()
        for row in rows:
            for key in row.keys():
                if key not in seen:
                    seen.add(key)
                    headers.append(str(key))
    else:
        headers = ["Message"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def append_concentration_level_to_workbook(wb, result: ConcentrationLevelResult) -> None:
    decision_rows = [
        {"Metric": "Decision", "Value": result.decision},
        {"Metric": "Best_level_model", "Value": result.best_model},
        {"Metric": "Summary", "Value": result.decision_summary},
        {"Metric": "Requested_levels", "Value": result.level_count_requested},
        {"Metric": "Populated_levels", "Value": result.level_count_observed},
        {"Metric": "Concentration_unit", "Value": result.concentration_unit},
        {"Metric": "OOF_validation_note", "Value": "Use out-of-fold exact/within-one-level accuracy to judge generalisation."},
        {"Metric": "Full_fit_note", "Value": "Full-fit accuracy is diagnostic only and may reflect memorisation; it is not validation."},
    ] + [{"Metric": "Warning", "Value": x} for x in result.warnings]
    _append_dict_sheet(wb, "Level_Model_Decision", decision_rows)
    _append_dict_sheet(wb, "Level_Model_Comparison", result.comparison)
    _append_dict_sheet(wb, "Level_Known_Standards", result.known_predictions)
    _append_dict_sheet(wb, "Level_Target_Predictions", result.target_predictions)
    _append_dict_sheet(wb, "Level_Confusion_Matrix", result.confusion_rows)
    _append_dict_sheet(wb, "Level_Definitions", result.level_definitions)
    _append_dict_sheet(wb, "Level_Group_Validation", result.group_summary)
    _append_dict_sheet(wb, "Level_Features", result.feature_rows)
    if result.plot_paths:
        try:
            from openpyxl.drawing.image import Image as XLImage
            if "Level_Plots" in wb.sheetnames:
                del wb["Level_Plots"]
            ws = wb.create_sheet("Level_Plots")
            row = 1
            for path in result.plot_paths:
                if not Path(path).exists():
                    continue
                ws.cell(row=row, column=1, value=Path(path).name)
                image = XLImage(str(path))
                image.width = min(image.width, 950)
                image.height = min(image.height, 680)
                ws.add_image(image, f"A{row+1}")
                row += 36
        except Exception:
            pass
