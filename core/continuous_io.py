"""Read-only input adapters and isolated output export for the v18.45 lab.

Uses the workbench's existing openpyxl dependency; never saves to an input path.
A current table, when supplied, is authoritative: deleted cached rows stay deleted.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import platform
import re
import uuid
from datetime import datetime
from pathlib import Path

import numpy as np

from .continuous_rrf import (run_audit, text, number, norm, row_info, VERSION, allowed_feature)
from .model_rerun import load_descriptor_records
from .response_predictor import read_table, guess_columns, parse_combo
from .training_limits import validate_training_count, is_nonblank_record


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_combo(value):
    parts = parse_combo(text(value))
    if set(parts) == {'A', 'B', 'C'}:
        # Preserve component labels/formulas for an additional hard consistency check.
        return tuple(parts[g]['index'] for g in ('A', 'B', 'C'))
    return None


def pick_header(headers, aliases):
    for alias in aliases:
        found = [h for h in headers if norm(h) == norm(alias)]
        if found:
            return found[0]
    return ''


def refresh_records(cached, path, sheet='', *, is_training=True, column_overrides=None):
    """Match by full Combo or A#/B#/C# + formula checks, never by row number.

    Raises on unmatched/ambiguous identities or changed cached RT. Formula alone
    is never an identity key. Blank optional paths should be handled by caller.
    """
    table = read_table(Path(path), sheet_name=sheet)
    table.rows = [r for r in table.rows if is_nonblank_record(r)]
    if is_training:
        validate_training_count(table.rows, label=str(Path(path).name))
    config = guess_columns(table, need_concentration=is_training)
    ratio = pick_header(table.headers, ['Measured_ratio']) or config.ratio_col
    conc = pick_header(table.headers, ['Actual_concentration']) or config.concentration_col
    overrides = column_overrides or {}
    fields = {'combo': config.combo_col, 'formula': config.formula_col,
              'name': config.name_col, 'ratio': ratio, 'concentration': conc,
              'group': pick_header(table.headers, ['Injection_Group', 'Injection group', 'Injection group', 'Injection_Group']),
              'rt': pick_header(table.headers, ['Apex_RT_min', 'Apex_RT', 'Apex RT', 'RT', 'RT'])}
    for field, header in overrides.items():
        if header:
            if header not in table.headers:
                raise ValueError('Column not found: ' + header)
            fields[field] = header
    if not fields['combo'] or not fields['ratio'] or (is_training and not fields['concentration']):
        raise ValueError('The current table must provide Combo and a measured Area/IS ratio'
                         + ('' if is_training else '')
                         + '. Set these in the column-mapping dialog. Predicted concentrations cannot replace known labels.')
    if is_training and any(term in norm(fields['concentration']) for term in ('predicted', 'estimated', 'oof', 'prediction', 'estimate', 'estimated')):
        raise ValueError('Predicted/estimated concentrations cannot be training labels.')
    exact_index, abc_index = {}, {}
    for i, row in enumerate(cached):
        key = re.sub(r'\s+', '', text(row.get('Combo')))
        if key:
            exact_index.setdefault(key, []).append(i)
        abc = canonical_combo(row.get('Combo'))
        if abc:
            abc_index.setdefault(abc, []).append(i)
    output, audit = [], []
    for raw in table.rows:
        combo = text(raw.get(fields['combo']))
        exact = re.sub(r'\s+', '', combo)
        hits = exact_index.get(exact, []) if exact else []
        mode = 'Exact_Combo'
        if not hits:
            key = canonical_combo(combo)
            hits = abc_index.get(key, []) if key else []
            mode = 'ABC_indices_plus_formula_check'
        group = text(raw.get(fields['group'])) if fields['group'] else ''
        if group:
            grouped = [j for j in hits if text(cached[j].get('Injection_Group')).casefold() == group.casefold()]
            # A different injection may have different ion/gradient descriptors: do not reuse it silently.
            hits = grouped
        source_row = raw.get('__source_row__', '')
        if len(hits) != 1:
            raise ValueError('Row %s matches %d cached Combo records: %s. Check Combo and Injection_Group or regenerate descriptors. Row order and formula alone are not used to guess identity.'
                             % (source_row, len(hits), combo))
        old = cached[hits[0]]
        formula = text(raw.get(fields['formula'])) if fields['formula'] else ''
        if formula and text(old.get('Formula')) and formula.replace(' ', '') != text(old['Formula']).replace(' ', ''):
            raise ValueError('Formula mismatch at current input row %s: %s vs cached %s'
                             % (source_row, formula, old.get('Formula')))
        if mode != 'Exact_Combo':
            a, b = parse_combo(combo), parse_combo(text(old.get('Combo')))
            for g in ('A', 'B', 'C'):
                for key in ('formula', 'label'):
                    av, bv = a[g].get(key), b[g].get(key)
                    if av and bv and av != bv:
                        raise ValueError('ABC index identity conflict at row %s (%s %s): regenerate descriptors.'
                                         % (source_row, g, key))
            if not formula:
                raise ValueError('ABC-index fallback needs a product Formula for hard verification.')
        if fields['rt']:
            new_rt, old_rt = number(raw.get(fields['rt'])), number(old.get('Apex_RT_min'))
            if np.isfinite(new_rt) and (not np.isfinite(old_rt) or abs(new_rt - old_rt) > 1e-6):
                raise ValueError('Row %s has a different RT from the cache. Apex-gradient descriptors may be outdated; regenerate descriptors first.'
                                 % source_row)
        # A refreshed table is not permission to silently retain changed descriptors.
        for header in table.headers:
            if header in old and allowed_feature(header):
                new_value, old_value = number(raw.get(header)), number(old.get(header))
                if np.isfinite(new_value) and (not np.isfinite(old_value) or
                        not np.isclose(new_value, old_value, rtol=1e-8, atol=1e-8)):
                    raise ValueError('Cached descriptor differs at current row %s / %s; regenerate ESI descriptors.'
                                     % (source_row, header))
        record = dict(old)
        record.update(Source_file=str(Path(path).name), Source_sheet=table.sheet_name,
                      Source_row=source_row, Source_index=len(output) + 1,
                      Combo=combo, Measured_ratio=number(raw.get(fields['ratio'])))
        if formula:
            record['Formula'] = formula
        if fields['name'] and text(raw.get(fields['name'])):
            record['Name'] = text(raw.get(fields['name']))
        if is_training:
            record['Actual_concentration'] = number(raw.get(fields['concentration']))
        else:
            # Never import possible target labels, even if present in the current table.
            record['Actual_concentration'] = None
        output.append(record)
        audit.append({'Dataset': 'Training' if is_training else 'Target', 'Current_source_row': source_row,
                      'Cached_source_row': old.get('Source_row'), 'Combo': combo, 'Match_method': mode,
                      'Ratio_column': fields['ratio'], 'Concentration_column': fields['concentration'] if is_training else '',
                      'Old_ratio': old.get('Measured_ratio'), 'New_ratio': record['Measured_ratio'],
                      'Old_known_concentration': old.get('Actual_concentration') if is_training else '',
                      'New_known_concentration': record.get('Actual_concentration') if is_training else ''})
    return output, audit, fields


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [clean(v) for v in value]
    if isinstance(value, (float, np.floating)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, Path):
        return str(value)
    return value


def write_json(path, value):
    Path(path).write_text(json.dumps(clean(value), ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def tabular_value(value):
    value = clean(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    # Prevent formula interpretation when user-entered text is exported.
    if isinstance(value, str) and value.startswith(('=', '+', '@')):
        return "'" + value
    return value


def columns(rows):
    return list(dict.fromkeys(k for row in rows for k in row))


def export_tables(folder, tables):
    """Application export using its existing Excel dependency, not a source edit."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in tables.items():
        headers = columns(rows) or ['Status']
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for r in rows:
            ws.append([tabular_value(r.get(h)) for h in headers])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24435B')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 32
        for j, h in enumerate(headers, start=1):
            width = min(48, max(15, len(h) + 2))
            ws.column_dimensions[get_column_letter(j)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Calibri', size=11)
                if isinstance(cell.value, float):
                    cell.number_format = '0.000000E+00'
                cell.alignment = Alignment(vertical='top')
        with (folder / (title + '.csv')).open('w', newline='', encoding='utf-8-sig') as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            writer.writeheader()
            writer.writerows({k: tabular_value(r.get(k)) for k in headers} for r in rows)
    out = folder / 'continuous_concentration_audit.xlsx'
    wb.save(out)
    return out


def make_plots(folder, result):
    # Figure objects + Agg canvas avoid modifying the host application's GUI backend.
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    y, p = result['actual_log'], result['predicted_log']
    files = []

    def save(fig, name):
        FigureCanvasAgg(fig)
        fig.savefig(str(folder / name), dpi=180, bbox_inches='tight')
        files.append(name)
        fig.clear()

    fig = Figure(figsize=(7, 6))
    ax = fig.add_subplot(111)
    ax.scatter(y, p, alpha=.8, label='Nested out-of-fold prediction')
    lo, hi = min(y.min(), p.min()), max(y.max(), p.max())
    ax.plot([lo, hi], [lo, hi], '--', label='Identity: predicted = actual')
    ax.set(xlabel='Actual concentration (log10)', ylabel='Predicted concentration (log10)',
           title='Held-out concentration agreement')
    ax.legend(fontsize=8)
    save(fig, '01_actual_vs_oof.png')

    fig = Figure(figsize=(9, 5))
    ax = fig.add_subplot(111)
    order = np.argsort(y, kind='stable')
    x = np.arange(1, len(y) + 1)
    ax.plot(x, y[order], '.-', label='Actual')
    ax.plot(x, p[order], '.', label='Nested OOF prediction')
    ax.plot(x, result['baseline_log'][order], '--', label='Constant-concentration null baseline')
    ax.set(xlabel='Standards sorted by actual concentration (same row alignment)',
           ylabel='Concentration (log10)', title='Trend recovery, not just average error')
    ax.legend(fontsize=8)
    save(fig, '02_known_sorted.png')

    fig = Figure(figsize=(7, 5))
    ax = fig.add_subplot(111)
    ax.scatter(y, p - y, alpha=.8)
    ax.axhline(0, linestyle='--')
    ax.set(xlabel='Actual concentration (log10)', ylabel='Predicted minus actual (log10)',
           title='Low/high concentration bias')
    save(fig, '03_concentration_residuals.png')

    rows = [r for r in result['comparison'] if 'Log_RMSE' in r]
    fig = Figure(figsize=(9, max(5, len(rows) * .32)))
    ax = fig.add_subplot(111)
    ax.barh([r['Model'] for r in rows], [r['Log_RMSE'] for r in rows])
    ax.invert_yaxis()
    ax.set(xlabel='Held-out log10 RMSE (lower is better)',
           title='Nested strategy, null baselines and exploratory candidates')
    ax.tick_params(axis='y', labelsize=8)
    save(fig, '04_model_comparison.png')
    return files


METRIC_GUIDE = [
    {'Metric': 'Log_RMSE', 'Meaning': 'Root mean squared log10 concentration error. Lower is better. Model selection uses inner validation only.'},
    {'Metric': 'Log_R2', 'Meaning': 'Squared-error performance relative to the evaluation-set mean. A value <=0 does not improve on this reference.'},
    {'Metric': 'Spearman', 'Meaning': 'Rank correlation of actual and predicted values; rank agreement is not concentration accuracy.'},
    {'Metric': 'Pred_on_actual_slope', 'Meaning': 'Slope of predicted vs actual log10 concentration. The identity slope is 1; a value near zero indicates weak trend recovery. Predictions are not rescaled to force a slope of 1.'},
    {'Metric': 'Log_spread_ratio', 'Meaning': 'Standard deviation of predicted log10 concentration divided by the actual standard deviation. A small ratio indicates compression; a ratio near 1 does not establish accuracy.'},
    {'Metric': 'Low20_log_bias / High20_log_bias', 'Meaning': 'Mean signed log10 error for the lowest and highest 20% of actual concentrations. Positive indicates overestimation.'},
    {'Metric': 'Empirical_P80_lower/upper', 'Meaning': 'Empirical range based on the 80th percentile of absolute nested-CV log residuals. It is not a compound-specific confidence interval and does not guarantee external coverage.'},
    {'Metric': 'NOT_VALIDATED', 'Meaning': 'Basic diagnostic checks failed, or only a response baseline was selected. Predictions are retained for review, not certified as reliable quantitative results.'},
    {'Metric': 'PRELIMINARY_ONLY_external_validation_needed', 'Meaning': 'Internal preliminary checks passed. This is not analytical-method validation or a guarantee for each unknown compound.'},
    {'Metric': 'Feature missing / range', 'Meaning': 'Range check using final model descriptors. Being within range does not establish full applicability; predictions are not clipped to the training range.'},
]


def execute(cache_path, output_parent, *, training_path='', target_path='', training_sheet='',
            target_sheet='', training_columns=None, target_columns=None,
            concentration_unit='same as known training table', seed=42, folds=5, progress=None,
            specs=None):
    emit = progress or (lambda message: None)
    paths = [Path(cache_path).resolve()] + [Path(p).resolve() for p in (training_path, target_path) if p]
    inputs = [{'path': str(p), 'sha256': sha256(p)} for p in paths]
    parent = Path(output_parent).expanduser().resolve()
    parent.mkdir(parents=True, exist_ok=True)
    folder = parent / ('v18_45_continuous_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:6])
    folder.mkdir(exist_ok=False)
    try:
        emit('Reading the ESI_Features cache; input files remain read-only.')
        _, training, targets, features, messages = load_descriptor_records(paths[0])
        matching, used_columns = [], {}
        cached_n = len(training)
        if training_path:
            training, rows, used_columns['training'] = refresh_records(training, training_path, training_sheet,
                                                        is_training=True, column_overrides=training_columns)
            matching.extend(rows)
            emit('Current calibration table replaces the cached row list: %d cached rows -> %d current records.' % (cached_n, len(training)))
        else:
            messages.append('No current training table supplied: labels/ratios and row list are from the cached ESI_Features sheet.')
        if target_path:
            targets, rows, used_columns['target'] = refresh_records(targets, target_path, target_sheet,
                                                    is_training=False, column_overrides=target_columns)
            matching.extend(rows)
        validate_training_count(training)
        result = run_audit(training, targets, features, seed=seed, folds=folds, specs=specs, progress=emit)
        for r in result['known'] + result['unknown']:
            r['Concentration_unit'] = concentration_unit
        summary = result['summary']
        summary['Notes'].extend(messages)
        summary['Concentration_unit'] = concentration_unit
        summary['Training_records_in_cache'] = cached_n
        summary['Input_column_mapping'] = used_columns
        import sklearn
        import scipy
        summary['Runtime'] = {'python': platform.python_version(), 'numpy': np.__version__,
                              'scikit_learn': sklearn.__version__, 'scipy': scipy.__version__}
        # Every direct input is fingerprinted before AND after evaluation/export.
        write_json(folder / 'input_manifest.json', {'version': VERSION, 'inputs': inputs,
                    'training_sheet': training_sheet, 'target_sheet': target_sheet, 'columns': used_columns,
                    'concentration_unit': concentration_unit, 'seed': seed, 'folds': folds,
                    'response_equation': 'log10(C) = log10(Area/IS) - predicted_log10_RRF'})
        summary_rows = [{'Setting': k, 'Value': v} for k, v in summary.items()]
        tables = {'Run_Summary': summary_rows, 'Model_Comparison': result['comparison'],
                  'Known_Standards_Predictions': result['known'], 'Unknown_Sample_Predictions': result['unknown'],
                  'Row_Audit': result['row_audit'], 'OOF_Folds': result['cv_folds'],
                  'Feature_Audit': result['feature_audit'], 'Inner_Search': result['inner_search'],
                  'Input_Matching': matching, 'Metric_Guide': METRIC_GUIDE}
        workbook = export_tables(folder, tables)
        images = make_plots(folder, result)
        write_json(folder / 'run_summary.json', summary)
        import joblib
        joblib.dump({'model': result['model'], 'feature_names': result['candidate_features'],
                     'unit': concentration_unit, 'version': VERSION}, folder / 'fitted_model.joblib')
        for item in inputs:
            if sha256(item['path']) != item['sha256']:
                raise RuntimeError('An input file changed during this run: ' + item['path'])
        write_json(folder / 'completion.json', {'status': 'completed', 'input_hashes_unchanged': True,
                   'workbook': workbook.name, 'plots': images, 'real_or_synthetic': 'depends on user-supplied inputs'})
        emit('Complete: ' + str(workbook) + '\nValidation status: ' + summary['Validation_status'])
        return folder, result
    except Exception as exc:
        write_json(folder / 'FAILED.json', {'status': 'failed', 'error': type(exc).__name__ + ': ' + str(exc),
                                          'inputs': inputs})
        raise
