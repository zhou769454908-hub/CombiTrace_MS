"""Descriptor catalog and previous-run effect review for the Tk feature selector.

The module reads the immutable ``ESI_Features`` cache and optional model sheets
from a previous result workbook.  It never evaluates unknown target labels and
never alters the source workbook.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .esi_descriptor_meta import (
    ADVANCED_EXTERNAL_DESCRIPTORS,
    descriptor_meta,
    mechanistic_profile,
)


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _num(value: object) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def _sheet_dict_rows(ws) -> List[Dict[str, object]]:
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_text(x) for x in values[0]]
    rows: List[Dict[str, object]] = []
    for line in values[1:]:
        if not line or not any(x not in (None, "") for x in line):
            continue
        rows.append({headers[i]: line[i] if i < len(line) else "" for i in range(len(headers)) if headers[i]})
    return rows


def _descriptor_names_from_esi_features(wb) -> Tuple[List[str], Dict[str, Dict[str, float]]]:
    names: List[str] = []
    stats: Dict[str, Dict[str, float]] = {}
    if "ESI_Features" not in wb.sheetnames:
        return names, stats
    ws = wb["ESI_Features"]
    rows = ws.iter_rows(values_only=True)
    try:
        header_values = next(rows)
    except StopIteration:
        return names, stats
    headers = [_text(x) for x in header_values]
    try:
        start = headers.index("Product_Master_Formula_Key") + 1
    except ValueError:
        start = 0
    try:
        end = headers.index("Warnings")
    except ValueError:
        end = len(headers)
    names = [x for x in headers[start:end] if x]
    base_numeric = [
        "Exact_mass", "DBE", "Apex_RT_min", "Effective_gradient_time_min",
        "Mobile_phase_A_pct", "Mobile_phase_B_pct", "B_slope_pct_per_min",
    ]
    names = [x for x in base_numeric if x in headers] + [x for x in names if x not in base_numeric]
    indices = {name: headers.index(name) for name in names}
    total_training = 0
    valid_counts = {name: 0 for name in names}
    values_by_name: Dict[str, List[float]] = {name: [] for name in names}
    dataset_idx = headers.index("Dataset") if "Dataset" in headers else -1
    for line in rows:
        dataset = _text(line[dataset_idx] if dataset_idx >= 0 and dataset_idx < len(line) else "").lower()
        if dataset != "training":
            continue
        total_training += 1
        for name, idx in indices.items():
            value = _num(line[idx] if idx < len(line) else None)
            if value is not None:
                valid_counts[name] += 1
                values_by_name[name].append(value)
    for name in names:
        vals = values_by_name[name]
        if vals:
            mean = sum(vals) / len(vals)
            variance = sum((x - mean) ** 2 for x in vals) / len(vals)
        else:
            variance = 0.0
        stats[name] = {
            "Training_N": float(total_training),
            "Valid_N": float(valid_counts[name]),
            "Missing_pct": (100.0 * (total_training - valid_counts[name]) / total_training) if total_training else 100.0,
            "Variance": variance,
        }
    return names, stats


def _read_previous_effects(wb) -> Dict[str, Dict[str, object]]:
    effects: Dict[str, Dict[str, object]] = {}

    def bucket(name: str) -> Dict[str, object]:
        return effects.setdefault(name, {})

    if "Trend_Feature_Stability" in wb.sheetnames:
        for row in _sheet_dict_rows(wb["Trend_Feature_Stability"]):
            name = _text(row.get("Feature") or row.get("Descriptor"))
            if not name:
                continue
            b = bucket(name)
            for src, dst in (
                ("Outer_selected_pct", "Previous_outer_selected_pct"),
                ("Trial_inclusion_pct", "Previous_trial_inclusion_pct"),
                ("Top10_trial_inclusion_pct", "Previous_top10_inclusion_pct"),
                ("Inclusion_score_delta", "Previous_trend_score_delta"),
                ("Mean_score_when_included", "Previous_mean_score_included"),
                ("Mean_score_when_excluded", "Previous_mean_score_excluded"),
            ):
                val = _num(row.get(src))
                if val is not None:
                    b[dst] = val
            if _text(row.get("Interpretation")):
                b["Previous_interpretation"] = _text(row.get("Interpretation"))

    for sheet_name in ("Feature_Importance", "Selected_Descriptors"):
        if sheet_name not in wb.sheetnames:
            continue
        for row in _sheet_dict_rows(wb[sheet_name]):
            name = _text(row.get("Feature") or row.get("Descriptor"))
            if not name:
                continue
            b = bucket(name)
            for key in (
                "Permutation_Importance_Mean", "Permutation_Importance", "Importance_Mean",
                "Selected_Frequency_pct", "Outer_selected_pct",
            ):
                val = _num(row.get(key))
                if val is not None:
                    if "Importance" in key:
                        b["Previous_permutation_importance"] = val
                    else:
                        b["Previous_selection_frequency_pct"] = val
            if _text(row.get("Final_Status")):
                b["Previous_status"] = _text(row.get("Final_Status"))
            elif _text(row.get("Status")):
                b["Previous_status"] = _text(row.get("Status"))

    # v18.37+ fixed-model feature search.
    if "Fixed_Feature_Effects" in wb.sheetnames:
        for row in _sheet_dict_rows(wb["Fixed_Feature_Effects"]):
            name = _text(row.get("Feature"))
            if not name:
                continue
            b = bucket(name)
            for src, dst in (
                ("Trend_Spearman_delta_included_minus_excluded", "Previous_fixed_trend_delta"),
                ("Absolute_MedianFold_improvement_included", "Previous_fixed_absolute_improvement"),
                ("Top5_Trend_frequency_pct", "Previous_top5_trend_pct"),
                ("Top5_Absolute_frequency_pct", "Previous_top5_absolute_pct"),
            ):
                val = _num(row.get(src))
                if val is not None:
                    b[dst] = val
            if _text(row.get("Effect_class")):
                b["Previous_fixed_effect_class"] = _text(row.get("Effect_class"))
    return effects


def _effect_class(effect: Dict[str, object]) -> Tuple[str, str]:
    trend = _num(effect.get("Previous_fixed_trend_delta"))
    if trend is None:
        trend = _num(effect.get("Previous_trend_score_delta"))
    absolute = _num(effect.get("Previous_fixed_absolute_improvement"))
    importance = _num(effect.get("Previous_permutation_importance"))
    selected = _num(effect.get("Previous_outer_selected_pct"))
    if selected is None:
        selected = _num(effect.get("Previous_selection_frequency_pct"))

    positive_votes = 0
    negative_votes = 0
    if trend is not None:
        if trend > 0.01:
            positive_votes += 2
        elif trend < -0.01:
            negative_votes += 2
    if absolute is not None:
        if absolute > 0.05:
            positive_votes += 2
        elif absolute < -0.05:
            negative_votes += 2
    if importance is not None:
        if importance > 0.005:
            positive_votes += 1
        elif importance < -0.002:
            negative_votes += 1
    if selected is not None and selected >= 50:
        positive_votes += 1
    if positive_votes >= 2 and negative_votes == 0:
        return "Positive", "Good / beneficial in previous run"
    if negative_votes >= 2 and positive_votes == 0:
        return "Negative", "Adverse in previous run"
    if positive_votes and negative_votes:
        return "Mixed", "Mixed objective/model-dependent effect"
    if positive_votes:
        return "Weak positive", "Weak positive evidence"
    if negative_votes:
        return "Weak negative", "Weak negative evidence"
    return "Unknown", "No reliable previous effect"


def load_descriptor_review(
    workbook_path: Optional[Path],
    *,
    include_advanced_external: bool = True,
) -> Tuple[List[Dict[str, object]], Dict[str, object]]:
    """Build a descriptor catalog for the bilingual manual selector."""
    names: List[str] = []
    stats: Dict[str, Dict[str, float]] = {}
    effects: Dict[str, Dict[str, object]] = {}
    previous: Dict[str, object] = {}
    if workbook_path and Path(workbook_path).exists():
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(workbook_path), read_only=True, data_only=True)
            names, stats = _descriptor_names_from_esi_features(wb)
            effects = _read_previous_effects(wb)
            # Detect previous best model for the fixed-model feature search.
            for sheet_name in ("Trend_Optimizer_Decision", "Model_Rerun_Diagnostics", "Diagnostics", "Descriptor_Run_Diagnostics"):
                if sheet_name not in wb.sheetnames:
                    continue
                for row in _sheet_dict_rows(wb[sheet_name]):
                    key = _text(row.get("Metric") or row.get("Setting") or row.get("Item"))
                    value = row.get("Value")
                    if not _text(value):
                        continue
                    if sheet_name == "Trend_Optimizer_Decision" and key in {"Best_model", "Trend_optimizer_best_model", "Trend_Best_model"}:
                        previous["Trend_Best_model"] = _text(value)
                    elif key in {"Best_model", "Trend_optimizer_best_model", "Trend_Best_model"}:
                        previous[key] = _text(value)
            wb.close()
        except Exception as exc:
            previous["Warning"] = str(exc)
    if include_advanced_external:
        for name in ADVANCED_EXTERNAL_DESCRIPTORS:
            if name not in names:
                names.append(name)
    rows: List[Dict[str, object]] = []
    for name in names:
        meta = descriptor_meta(name)
        mech = mechanistic_profile(name)
        stat = stats.get(name, {})
        eff = effects.get(name, {})
        cls, summary = _effect_class(eff)
        valid_n = int(stat.get("Valid_N", 0.0))
        missing_pct = float(stat.get("Missing_pct", 100.0))
        variance = float(stat.get("Variance", 0.0))
        available = valid_n > 0 and missing_pct < 50.0 and variance > 1e-14
        if name in ADVANCED_EXTERNAL_DESCRIPTORS and valid_n == 0:
            availability = "External input required"
        elif valid_n == 0:
            availability = "Unavailable"
        elif missing_pct >= 50.0:
            availability = "Too sparse"
        elif variance <= 1e-14:
            availability = "Constant"
        else:
            availability = "Available"
        default_use = available and mech.get("Mechanistic_level") not in {"Coarse covariate"}
        if cls in {"Positive", "Weak positive"}:
            default_use = available
        if cls == "Negative":
            default_use = False
        row: Dict[str, object] = {
            "Use": bool(default_use),
            "Feature": name,
            "Name_zh": meta.get("Name_zh", name),
            "Name_en": meta.get("Name_en", name),
            "Group": mech.get("Mechanistic_group", ""),
            "Mechanistic_level": mech.get("Mechanistic_level", ""),
            "Availability": availability,
            "Valid_N": valid_n,
            "Missing_pct": missing_pct,
            "Variance": variance,
            "Previous_effect": cls,
            "Previous_effect_summary": summary,
            "Previous_outer_selected_pct": eff.get("Previous_outer_selected_pct", ""),
            "Previous_trend_score_delta": eff.get("Previous_fixed_trend_delta", eff.get("Previous_trend_score_delta", "")),
            "Previous_absolute_improvement": eff.get("Previous_fixed_absolute_improvement", ""),
            "Previous_permutation_importance": eff.get("Previous_permutation_importance", ""),
            "Theory_zh": mech.get("Theory_zh", ""),
            "Theory_en": mech.get("Theory_en", ""),
            "Expected_direction_zh": mech.get("Expected_direction_zh", ""),
            "Expected_direction_en": mech.get("Expected_direction_en", ""),
            "Unit_or_type": meta.get("Unit_or_type", ""),
            "Source": meta.get("Source", ""),
            "Requires_3D": meta.get("Requires_3D", "No"),
        }
        rows.append(row)
    rows.sort(key=lambda r: (
        0 if r["Availability"] == "Available" else 1,
        0 if r["Previous_effect"] in {"Positive", "Weak positive"} else 1,
        str(r["Mechanistic_level"]), str(r["Feature"]),
    ))
    return rows, previous


def detect_previous_best_model(workbook_path: Optional[Path]) -> str:
    if not workbook_path or not Path(workbook_path).exists():
        return ""
    _rows, previous = load_descriptor_review(workbook_path, include_advanced_external=False)
    # Prefer the trend-specific model when available because this workflow is
    # normally run with the trend objective.
    return _text(
        previous.get("Trend_optimizer_best_model")
        or previous.get("Trend_Best_model")
        or previous.get("Best_model")
    )
