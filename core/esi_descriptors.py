"""ESI structure descriptors and LC-gradient features (v18.26).

This module is intentionally separate from concentration prediction.  It uses the
existing Combo -> Product_SMILES master table to generate auditable molecular
features and linearly interpolates the programmed mobile-phase composition at
each chromatographic apex.

Important limitations
---------------------
* Acid/base columns are SMARTS/Gasteiger structural proxies, not predicted pKa.
* The programmed gradient at Apex RT is an approximation.  A user-supplied
  gradient delay can be subtracted from RT, but dwell volume and column transit
  are not inferred automatically.
* 3D features are calculated from one deterministic RDKit conformer and are not
  a solution-phase conformational ensemble.
"""
from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from .chemistry import monoisotopic_mass, parse_formula
from .isotope_theory import calc_rdb
from .response_predictor import (
    PreparedRow,
    _norm_header,
    guess_columns,
    parse_combo,
    prepare_rows,
    read_table,
)
from .structure_features import (
    RDKIT_AVAILABLE,
    StructureLibrary,
    enrich_rows,
    extended_descriptors_from_smiles,
)
from .combitrace_ie import (
    ProductMasterEntry,
    _combo_formula_key,
    _guess_rt,
    apply_product_smiles_master,
    load_product_smiles_master,
)
from .esi_descriptor_meta import descriptor_meta


@dataclass(frozen=True)
class GradientPoint:
    time_min: float
    a_pct: float
    b_pct: float


@dataclass(frozen=True)
class GradientAtRT:
    apex_rt_min: Optional[float]
    delay_min: float
    effective_time_min: Optional[float]
    a_pct: Optional[float]
    b_pct: Optional[float]
    b_slope_pct_per_min: Optional[float]
    segment_start_min: Optional[float]
    segment_end_min: Optional[float]
    state: str
    range_status: str


@dataclass
class EsiDescriptorReport:
    output_xlsx: Path
    output_csv: Optional[Path]
    n_master: int
    n_training: int
    n_target: int
    n_training_matched: int
    n_target_matched: int
    n_rt_training: int
    n_rt_target: int
    warnings: List[str]
    model_benchmark_enabled: bool = False
    model_calibration_rows: int = 0
    model_target_rows: int = 0
    best_model: str = ""
    best_model_rating: str = ""
    model_decision: str = ""
    model_decision_summary: str = ""
    best_median_fold_error: Optional[float] = None
    best_p80_fold_error: Optional[float] = None
    best_within_2x_pct: Optional[float] = None
    best_within_5x_pct: Optional[float] = None
    best_r2_log: Optional[float] = None
    best_rmse_log10: Optional[float] = None
    improvement_vs_global_pct: Optional[float] = None
    model_outliers_removed: int = 0
    model_selected_descriptors: int = 0
    model_unavailable_descriptors: int = 0
    model_objective: str = ""
    corrected_spearman: Optional[float] = None
    raw_spearman: Optional[float] = None
    delta_spearman: Optional[float] = None
    pairwise_concordance_pct: Optional[float] = None
    top20_overlap_pct: Optional[float] = None
    injection_group_spearman: Optional[float] = None
    dual_ie_enabled: bool = False
    dual_ie_decision: str = ""
    dual_ie_summary: str = ""
    dual_ie_published_method: str = ""
    dual_ie_target_rank_spearman: Optional[float] = None
    published_ie_spearman: Optional[float] = None
    published_ie_median_fold: Optional[float] = None
    trend_optimizer_enabled: bool = False
    trend_optimizer_best_model: str = ""
    trend_optimizer_decision: str = ""
    trend_optimizer_summary: str = ""
    trend_optimizer_spearman: Optional[float] = None
    trend_optimizer_delta_spearman: Optional[float] = None
    trend_optimizer_permutation_p: Optional[float] = None
    fixed_feature_search_enabled: bool = False
    fixed_feature_model: str = ""
    fixed_feature_best_absolute_trial: Optional[int] = None
    fixed_feature_best_trend_trial: Optional[int] = None
    level_classification_enabled: bool = False
    level_classification_best_model: str = ""
    level_classification_decision: str = ""
    level_classification_exact_pct: Optional[float] = None
    level_classification_within1_pct: Optional[float] = None
    level_classification_all_hit: bool = False


def _optional_float(value: object) -> Optional[float]:
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


DEFAULT_GRADIENT: Tuple[GradientPoint, ...] = (
    GradientPoint(0.0, 80.0, 20.0),
    GradientPoint(0.5, 80.0, 20.0),
    GradientPoint(3.0, 60.0, 40.0),
    GradientPoint(6.0, 30.0, 70.0),
    GradientPoint(9.0, 5.0, 95.0),
    GradientPoint(10.0, 30.0, 70.0),
    GradientPoint(11.0, 80.0, 20.0),
    GradientPoint(15.0, 80.0, 20.0),
)


def default_gradient_text() -> str:
    lines = ["Time(min),A(%),B(%)"]
    lines.extend(f"{p.time_min:g},{p.a_pct:g},{p.b_pct:g}" for p in DEFAULT_GRADIENT)
    return "\n".join(lines)


def parse_gradient_text(text: str) -> Tuple[List[GradientPoint], List[str]]:
    """Parse comma/tab/space separated gradient rows."""
    points: List[GradientPoint] = []
    warnings: List[str] = []
    for line_no, raw in enumerate(str(text or "").splitlines(), start=1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        chunks = [x.strip() for x in re.split(r"[,\t; ]+", line) if x.strip()]
        if len(chunks) < 3:
            warnings.append(f"Gradient line {line_no}: expected Time,A,B; ignored: {raw!r}")
            continue
        try:
            t, a, b = float(chunks[0]), float(chunks[1]), float(chunks[2])
        except Exception:
            # Permit one header row.
            if any(ch.isalpha() for ch in line):
                continue
            warnings.append(f"Gradient line {line_no}: non-numeric values; ignored: {raw!r}")
            continue
        if not all(math.isfinite(x) for x in (t, a, b)):
            warnings.append(f"Gradient line {line_no}: non-finite value; ignored")
            continue
        if t < 0:
            warnings.append(f"Gradient line {line_no}: time < 0; ignored")
            continue
        if not (0 <= a <= 100 and 0 <= b <= 100):
            warnings.append(f"Gradient line {line_no}: A/B should be within 0-100%")
        if abs((a + b) - 100.0) > 0.5:
            warnings.append(f"Gradient line {line_no}: A+B={a+b:.3g}% (not approximately 100%)")
        points.append(GradientPoint(float(t), float(a), float(b)))

    # Sort and retain the last definition at a repeated time.
    by_time: Dict[float, GradientPoint] = {}
    for p in points:
        by_time[p.time_min] = p
    points = [by_time[t] for t in sorted(by_time)]
    if len(points) < 2:
        raise ValueError("Gradient program needs at least two valid time points")
    return points, warnings


def interpolate_gradient(
    points: Sequence[GradientPoint],
    apex_rt_min: Optional[float],
    *,
    delay_min: float = 0.0,
) -> GradientAtRT:
    if apex_rt_min is None or not math.isfinite(float(apex_rt_min)):
        return GradientAtRT(None, float(delay_min), None, None, None, None, None, None, "missing_rt", "missing_rt")
    pts = sorted(points, key=lambda p: p.time_min)
    if len(pts) < 2:
        raise ValueError("Gradient program needs at least two points")
    rt = float(apex_rt_min)
    eff = rt - float(delay_min)
    status = "inside"
    if eff <= pts[0].time_min:
        p = pts[0]
        return GradientAtRT(rt, float(delay_min), eff, p.a_pct, p.b_pct, 0.0, p.time_min, p.time_min, "isocratic", "before_or_at_start_clamped")
    if eff >= pts[-1].time_min:
        p = pts[-1]
        return GradientAtRT(rt, float(delay_min), eff, p.a_pct, p.b_pct, 0.0, p.time_min, p.time_min, "isocratic", "after_or_at_end_clamped")

    left = pts[0]
    right = pts[1]
    for i in range(len(pts) - 1):
        if pts[i].time_min <= eff <= pts[i + 1].time_min:
            left, right = pts[i], pts[i + 1]
            break
    dt = right.time_min - left.time_min
    frac = 0.0 if dt <= 0 else (eff - left.time_min) / dt
    a = left.a_pct + frac * (right.a_pct - left.a_pct)
    b = left.b_pct + frac * (right.b_pct - left.b_pct)
    slope = 0.0 if dt <= 0 else (right.b_pct - left.b_pct) / dt
    if abs(slope) < 1e-12:
        state = "isocratic"
    elif slope > 0:
        state = "B_increasing"
    else:
        state = "B_decreasing"
    return GradientAtRT(rt, float(delay_min), eff, float(a), float(b), float(slope), left.time_min, right.time_min, state, status)


EXPERIMENTAL_ION_BEHAVIOR_NAMES: Tuple[str, ...] = (
    "Observed_Fragility_Index",
    "Adduct_Cluster_Proneness_Index",
    "Primary_Ion_Fraction",
    "Ion_Form_Diversity_Count",
    "Accepted_Channel_Count",
    "Summed_Channel_Count",
    "Additional_Summed_Area",
    "Br_Fragment_Fraction",
)

# The first groups are deliberately compact and interpretable.  Extended VSA /
# BCUT features are appended after them when requested.
DESCRIPTOR_GROUPS: Tuple[Tuple[str, Tuple[str, ...]], ...] = (
    ("Experimental_ion_behavior", EXPERIMENTAL_ION_BEHAVIOR_NAMES),
    ("Acid_base_proxies", (
        "FormalCharge", "AbsoluteFormalChargeSum", "PositiveFormalChargeAtomCount", "NegativeFormalChargeAtomCount",
        "HBD", "HBA", "NHOHCount", "NOCount",
        "AcidicSiteCount_proxy", "BasicSiteCount_proxy", "IonizableSiteCount_proxy",
        "BasicMinusAcidic_proxy", "ZwitterionPotential_proxy",
        "CarboxylicAcidCount", "SulfonicAcidCount", "PhosphoricOHCount", "PhenolCount", "ThiolCount",
        "ImideNHCount", "AliphaticAmineCount", "AromaticBasicNCount", "AmidineGuanidineCount",
        "ImineNCount", "QuaternaryAmmoniumCount",
        "GasteigerChargeMax", "GasteigerChargeMin", "GasteigerChargeRange", "GasteigerAbsChargeMean",
        "GasteigerPositiveChargeSum", "GasteigerNegativeChargeAbsSum", "GasteigerChargeSeparation",
    )),
    ("Hydrophobic_polarity", (
        "MolLogP", "MolLogP_per_HeavyAtom", "TPSA", "TPSA_per_HeavyAtom", "PolarSurfaceFraction_proxy",
        "MolMR", "LabuteASA", "FractionCSP3", "AromaticAtomFraction", "HeteroAtomFraction", "CarbonAtomFraction",
        "AromaticRingCount", "AliphaticRingCount", "RingCount", "RotatableBonds",
    )),
    ("Size_volume_shape", (
        "MolWt", "HeavyAtomMolWt", "ExactMolWt", "HeavyAtomCount", "HeteroAtomCount",
        "MolVolume3D", "RadiusOfGyration", "Asphericity", "Eccentricity", "SpherocityIndex",
        "PMI1", "PMI2", "PMI3", "NPR1", "NPR2", "InertialShapeFactor", "PBF",
        "AmideBondCount", "BridgeheadAtomCount", "SpiroAtomCount", "BertzCT", "BalabanJ",
        "NumValenceElectrons", "NumRadicalElectrons",
    )),
)


ENGINEERED_DESCRIPTOR_NAMES: Tuple[str, ...] = (
    "HydrophobicGradientExposure_proxy",
    "PolarAqueousExposure_proxy",
    "IonizableAqueousExposure_proxy",
    "BasicAqueousExposure_proxy",
    "AcidicAqueousExposure_proxy",
    "ChargeDensitySurface_proxy",
    "PolarMassDensity_proxy",
    "HydrophobicSurface_proxy",
    "AromaticHydrophobicity_proxy",
    "IonizableSiteDensity_proxy",
    "DBE_per_C",
    "Hetero_to_C_ratio",
    "Mass_per_IonizableSite_proxy",
    "GradientChangeExposure_proxy",
    "LogP_TPSA_balance_proxy",
)


def _safe_ratio(numerator: object, denominator: object, *, minimum_denominator: float = 1e-12) -> Optional[float]:
    try:
        a = float(numerator)
        b = float(denominator)
        if not (math.isfinite(a) and math.isfinite(b)) or abs(b) < float(minimum_denominator):
            return None
        return a / b
    except Exception:
        return None


def _add_engineered_features(record: Dict[str, object]) -> None:
    """Add a compact set of physically motivated numeric interaction proxies.

    These are deterministic combinations of already generated descriptors and
    gradient conditions.  They do not claim to be mechanistic ESI simulations;
    they make key nonlinear/interaction relationships accessible to linear and
    PLS-type models without generating an uncontrolled polynomial feature set.
    """
    def n(name: str) -> Optional[float]:
        return _optional_float(record.get(name))

    a_pct = n("Mobile_phase_A_pct")
    b_pct = n("Mobile_phase_B_pct")
    a_frac = None if a_pct is None else a_pct / 100.0
    b_frac = None if b_pct is None else b_pct / 100.0
    logp = n("MolLogP")
    tpsa = n("TPSA")
    ion = n("IonizableSiteCount_proxy")
    basic = n("BasicSiteCount_proxy")
    acidic = n("AcidicSiteCount_proxy")
    charge_sep = n("GasteigerChargeSeparation")
    asa = n("LabuteASA")
    mw = n("MolWt") or n("Exact_mass")
    aromatic_frac = n("AromaticAtomFraction")
    heavy = n("HeavyAtomCount")
    dbe = n("DBE")
    carbon = n("Formula_C")
    exact = n("Exact_mass")
    slope = n("B_slope_pct_per_min")
    rt = n("Apex_RT_min")

    record["HydrophobicGradientExposure_proxy"] = None if logp is None or b_frac is None else logp * b_frac
    record["PolarAqueousExposure_proxy"] = None if tpsa is None or a_frac is None else tpsa * a_frac
    record["IonizableAqueousExposure_proxy"] = None if ion is None or a_frac is None else ion * a_frac
    record["BasicAqueousExposure_proxy"] = None if basic is None or a_frac is None else basic * a_frac
    record["AcidicAqueousExposure_proxy"] = None if acidic is None or a_frac is None else acidic * a_frac
    record["ChargeDensitySurface_proxy"] = _safe_ratio(charge_sep, asa)
    record["PolarMassDensity_proxy"] = _safe_ratio(tpsa, mw)
    record["HydrophobicSurface_proxy"] = None if logp is None or asa is None else logp * asa
    record["AromaticHydrophobicity_proxy"] = None if logp is None or aromatic_frac is None else logp * aromatic_frac
    record["IonizableSiteDensity_proxy"] = _safe_ratio(ion, heavy)
    record["DBE_per_C"] = _safe_ratio(dbe, carbon)
    hetero = sum(float(n(f"Formula_{el}") or 0.0) for el in ("N", "O", "P", "S", "F", "Cl", "Br", "I"))
    record["Hetero_to_C_ratio"] = _safe_ratio(hetero, carbon)
    record["Mass_per_IonizableSite_proxy"] = None if exact is None else exact / max(float(ion or 0.0), 1.0)
    record["GradientChangeExposure_proxy"] = None if slope is None or rt is None else slope * rt
    record["LogP_TPSA_balance_proxy"] = None if logp is None or tpsa is None else logp / (1.0 + tpsa / 100.0)


def _key_descriptor_names() -> List[str]:
    out: List[str] = []
    for _group, names in DESCRIPTOR_GROUPS:
        for name in names:
            if name not in out:
                out.append(name)
    for name in ENGINEERED_DESCRIPTOR_NAMES:
        if name not in out:
            out.append(name)
    return out


def _is_extended_descriptor(name: str) -> bool:
    return str(name).startswith(("PEOE_VSA", "SlogP_VSA", "SMR_VSA", "EState_VSA", "BCUT2D_"))


def _source_file(row: PreparedRow) -> str:
    return str(row.raw.get("__source_file__", "") or "")


def _source_sheet(row: PreparedRow) -> str:
    return str(row.raw.get("__source_sheet__", "") or "")


def _source_row(row: PreparedRow) -> object:
    return row.raw.get("__source_row__", row.source_index)


def _formula_mass(formula: str) -> Optional[float]:
    try:
        return float(monoisotopic_mass(parse_formula(formula)))
    except Exception:
        return None


def _formula_dbe(formula: str) -> Optional[float]:
    try:
        return calc_rdb(parse_formula(formula))
    except Exception:
        return None


def _flatten_master_entries(mapping: Dict[str, List[ProductMasterEntry]]) -> List[ProductMasterEntry]:
    seen = set()
    out: List[ProductMasterEntry] = []
    for key in sorted(mapping):
        for entry in mapping[key]:
            marker = (entry.source_file, entry.source_sheet, str(entry.source_row), entry.smiles, entry.combo)
            if marker in seen:
                continue
            seen.add(marker)
            out.append(entry)
    return out


def _master_prepared_rows(entries: Sequence[ProductMasterEntry]) -> List[PreparedRow]:
    rows: List[PreparedRow] = []
    for i, entry in enumerate(entries, start=1):
        formula = entry.expected_formula or entry.product_formula
        try:
            counts = {str(k): int(v) for k, v in parse_formula(formula).items()} if formula else {}
        except Exception:
            counts = {}
        row = PreparedRow(
            source_index=i,
            raw={
                "__source_file__": entry.source_file,
                "__source_sheet__": entry.source_sheet,
                "__source_row__": entry.source_row,
                "Product_SMILES": entry.smiles,
                "Expected_Formula": entry.expected_formula,
                "Product_Formula": entry.product_formula,
            },
            combo=entry.combo,
            formula=formula,
            name="",
            ratio=None,
            concentration=None,
            exact_mass=_formula_mass(formula),
            correction_factor=None,
            log_correction_factor=None,
            combo_parts=parse_combo(entry.combo),
            element_counts=counts,
        )
        rows.append(row)
    return rows



def _guess_injection_group_column(headers: Sequence[str]) -> str:
    """Find a standard-mixture injection/batch/RAW identifier column."""
    candidates = [
        "Injection_Group", "InjectionGroup", "Injection group", "Injection", "Batch", "Run_Group",
        "RAW", "Raw", "RAW file", "Raw file", 'RAW_file', 'Raw_file', 'Filename',
        'Injection group', 'Injection_Batch', 'Batch', 'Standard_Group', 'Sample_Group', "Group",
    ]
    norm_map = {_norm_header(h): h for h in headers}
    for candidate in candidates:
        key = _norm_header(candidate)
        if key in norm_map:
            return norm_map[key]
    # Conservative contains matching; generic 'group' is considered last.
    priority_tokens = ["injectiongroup", 'Injection group', 'Injection_Batch', "rawfile", 'raw_file', 'Batch', "rungroup"]
    for token in priority_tokens:
        nt = _norm_header(token)
        for header in headers:
            if nt and nt in _norm_header(header):
                return header
    return ""

def _prepare_dataset(
    path: Optional[Path],
    *,
    sheet_name: str,
    is_training: bool,
    combo_col: str = "",
    formula_col: str = "",
    rt_col: str = "",
    group_col: str = "",
) -> Tuple[List[PreparedRow], List[str]]:
    if not path:
        return [], []
    table = read_table(Path(path), sheet_name=sheet_name)
    cfg = guess_columns(table, need_concentration=is_training)
    if combo_col:
        cfg.combo_col = combo_col
    if formula_col:
        cfg.formula_col = formula_col
    rows, warnings = prepare_rows(table, cfg, is_training=is_training)
    # Retention-time override is copied to a canonical name so _guess_rt finds it.
    if rt_col:
        for row in rows:
            row.raw["Apex_RT"] = row.raw.get(rt_col, "")
    selected_group_col = str(group_col or "").strip()
    if is_training and not selected_group_col:
        selected_group_col = _guess_injection_group_column(table.headers)
        if selected_group_col:
            warnings.append(f'{table.path.name}: detected injection-group column: {selected_group_col!r}')
    if is_training:
        for row in rows:
            value = row.raw.get(selected_group_col, "") if selected_group_col else ""
            row.raw["__injection_group__"] = str(value or "").strip()
            row.raw["__injection_group_col__"] = selected_group_col
        if not selected_group_col:
            warnings.append(
                f'{table.path.name}: no injection-group/RAW column detected; group-held-out evaluation will be skipped. Specify the column in the interface if available.'
            )
    return rows, warnings


def _gradient_values(row: PreparedRow, points: Sequence[GradientPoint], delay_min: float) -> GradientAtRT:
    return interpolate_gradient(points, _guess_rt(row), delay_min=delay_min)


def _descriptor_header(rows: Sequence[PreparedRow], *, include_extended: bool) -> List[str]:
    present = set()
    for row in rows:
        present.update((getattr(row, "structure_features", {}) or {}).keys())
        present.update(str(k) for k in (getattr(row, "raw", {}) or {}).keys())
    out = [name for name in _key_descriptor_names() if name in present]
    # Include formula-derived descriptors after the physically interpretable block.
    formula_names = sorted(name for name in present if str(name).startswith(("Formula_", "FormulaFrac_")))
    out.extend(name for name in formula_names if name not in out)
    if include_extended:
        ext = sorted(name for name in present if _is_extended_descriptor(name))
        out.extend(name for name in ext if name not in out)
    return out


def _row_to_record(
    dataset: str,
    row: PreparedRow,
    descriptor_names: Sequence[str],
    points: Sequence[GradientPoint],
    delay_min: float,
    *,
    privacy_mode: bool,
    include_engineered: bool = True,
) -> Dict[str, object]:
    g = _gradient_values(row, points, delay_min)
    parts = row.combo_parts
    record: Dict[str, object] = {
        "Dataset": dataset,
        "Source_file": _source_file(row),
        "Source_sheet": _source_sheet(row),
        "Source_row": _source_row(row),
        "Source_index": row.source_index,
        "Name": row.name,
        "Formula": row.formula,
        "Combo": row.combo,
        "A_Formula": parts.get("A", {}).get("formula", ""),
        "B_Formula": parts.get("B", {}).get("formula", ""),
        "C_Formula": parts.get("C", {}).get("formula", ""),
        "Injection_Group": row.raw.get("__injection_group__", ""),
        "Injection_Group_Column": row.raw.get("__injection_group_col__", ""),
        "Measured_ratio": row.ratio,
        "Actual_concentration": row.concentration,
        "Exact_mass": row.exact_mass,
        "DBE": _formula_dbe(row.formula),
        "Apex_RT_min": g.apex_rt_min,
        "Gradient_delay_min": g.delay_min,
        "Effective_gradient_time_min": g.effective_time_min,
        "Mobile_phase_A_pct": g.a_pct,
        "Mobile_phase_B_pct": g.b_pct,
        "B_slope_pct_per_min": g.b_slope_pct_per_min,
        "Gradient_segment_start_min": g.segment_start_min,
        "Gradient_segment_end_min": g.segment_end_min,
        "Gradient_state": g.state,
        "Gradient_range_status": g.range_status,
        "Structure_status": getattr(row, "structure_status", ""),
        "Structure_method": getattr(row, "structure_method", ""),
        "Structure_hash": getattr(row, "structure_hash", ""),
        "Product_formula_calc": getattr(row, "product_formula_calc", ""),
        "Formula_match": getattr(row, "formula_match", ""),
        "3D_status": getattr(row, "three_d_status", ""),
        "Descriptor_count": len(getattr(row, "structure_features", {}) or {}),
        "Product_Master_Match": row.raw.get("__product_master_match__", ""),
        "Product_Master_Formula_Key": row.raw.get("__product_master_formula_key__", _combo_formula_key(row.combo)),
        "Warnings": "; ".join(row.warnings + (getattr(row, "structure_warnings", []) or [])),
        # Internal-only structure string used by the official MS2Quant R bridge.
        # It is intentionally omitted from every output header in privacy mode.
        "__Product_SMILES_Private": getattr(row, "product_smiles", ""),
    }
    if not privacy_mode:
        record["Product_SMILES"] = getattr(row, "product_smiles", "")
    features = getattr(row, "structure_features", {}) or {}
    raw_values = getattr(row, "raw", {}) or {}
    for name in descriptor_names:
        if name in ENGINEERED_DESCRIPTOR_NAMES:
            continue
        if name in EXPERIMENTAL_ION_BEHAVIOR_NAMES:
            record[name] = raw_values.get(name, "")
        else:
            record[name] = features.get(name, "")
    if include_engineered:
        _add_engineered_features(record)
    for name in ENGINEERED_DESCRIPTOR_NAMES:
        if name in descriptor_names and name not in record:
            record[name] = ""
    return record


def _overlap_records(
    train_rows: Sequence[PreparedRow],
    target_rows: Sequence[PreparedRow],
    points: Sequence[GradientPoint],
    delay_min: float,
) -> List[Dict[str, object]]:
    train_by_key: Dict[str, List[PreparedRow]] = {}
    target_by_key: Dict[str, List[PreparedRow]] = {}
    for row in train_rows:
        key = _combo_formula_key(row.combo)
        if key:
            train_by_key.setdefault(key, []).append(row)
    for row in target_rows:
        key = _combo_formula_key(row.combo)
        if key:
            target_by_key.setdefault(key, []).append(row)
    out: List[Dict[str, object]] = []
    for key in sorted(set(train_by_key).intersection(target_by_key)):
        # Formula keys are unique in the user's 10x10x10 library; preserve all
        # rows just in case there are multiple injections or accidental duplicates.
        for tr in train_by_key[key]:
            for ta in target_by_key[key]:
                gt = _gradient_values(tr, points, delay_min)
                ga = _gradient_values(ta, points, delay_min)
                ratio_fold = None
                if tr.ratio is not None and tr.ratio > 0 and ta.ratio is not None and ta.ratio > 0:
                    ratio_fold = float(ta.ratio / tr.ratio)
                out.append({
                    "ABC_Formula_Key": key,
                    "Combo_training": tr.combo,
                    "Combo_target": ta.combo,
                    "Training_source_file": _source_file(tr),
                    "Training_source_sheet": _source_sheet(tr),
                    "Training_source_row": _source_row(tr),
                    "Target_source_file": _source_file(ta),
                    "Target_source_sheet": _source_sheet(ta),
                    "Target_source_row": _source_row(ta),
                    "Training_Apex_RT_min": gt.apex_rt_min,
                    "Target_Apex_RT_min": ga.apex_rt_min,
                    "Delta_RT_target_minus_training_min": (
                        None if gt.apex_rt_min is None or ga.apex_rt_min is None else float(ga.apex_rt_min - gt.apex_rt_min)
                    ),
                    "Training_A_pct": gt.a_pct,
                    "Training_B_pct": gt.b_pct,
                    "Target_A_pct": ga.a_pct,
                    "Target_B_pct": ga.b_pct,
                    "Delta_B_target_minus_training_pct": (
                        None if gt.b_pct is None or ga.b_pct is None else float(ga.b_pct - gt.b_pct)
                    ),
                    "Training_ratio": tr.ratio,
                    "Target_ratio": ta.ratio,
                    "Target_to_training_ratio_fold": ratio_fold,
                    "Training_actual_concentration": tr.concentration,
                    "Structure_hash_training": getattr(tr, "structure_hash", ""),
                    "Structure_hash_target": getattr(ta, "structure_hash", ""),
                })
    return out


def _style_workbook(wb) -> None:
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception:
        return
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="E2F0D9")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row > 1 and ws.max_column > 1 and ws.title not in {"Descriptor_Definitions", "Gradient_Program"}:
            ws.auto_filter.ref = ws.dimensions
        if ws.max_row:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col in range(1, min(ws.max_column, 100) + 1):
            max_len = 8
            for cell in ws.iter_cols(min_col=col, max_col=col, min_row=1, max_row=min(ws.max_row, 200)):
                for c in cell:
                    max_len = max(max_len, min(len(str(c.value or "")), 48))
            ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 50)
    # Definitions sheet gets section highlighting.
    if "Descriptor_Definitions" in wb.sheetnames:
        ws = wb["Descriptor_Definitions"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value and not row[1].value:
                for cell in row:
                    cell.fill = section_fill
                    cell.font = Font(bold=True)


def _append_records(ws, records: Sequence[Dict[str, object]], headers: Sequence[str]) -> None:
    ws.append(list(headers))
    for record in records:
        ws.append([record.get(h, "") for h in headers])


def _public_record_headers(
    records: Sequence[Dict[str, object]],
    preferred_headers: Sequence[str],
) -> List[str]:
    """Return stable output headers while excluding private/internal fields.

    Some optional model branches append public metadata to records after the
    initial descriptor header list is built (for example Published_* fields).
    Conversely, private helper fields such as ``__Product_SMILES_Private`` must
    never be written to CSV/Excel in privacy mode.  Building the final header
    list from the records prevents ``csv.DictWriter`` from failing when optional
    branches add fields, while preserving the preferred column order.
    """
    out: List[str] = []
    seen = set()

    def add(name: object) -> None:
        key = str(name or "")
        if not key or key.startswith("_") or key in seen:
            return
        seen.add(key)
        out.append(key)

    for header in preferred_headers:
        add(header)
    for record in records:
        for key in record.keys():
            add(key)
    return out


def _append_simple_dict_sheet(wb, title: str, rows: Sequence[Dict[str, object]]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    if rows:
        headers: List[str] = []
        seen = set()
        for row in rows:
            for key in row.keys():
                if key not in seen:
                    seen.add(key); headers.append(str(key))
    else:
        headers = ["Message"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def run_esi_descriptor_export(
    product_master_file: Path,
    out_xlsx: Path,
    *,
    calibration_file: Optional[Path] = None,
    target_file: Optional[Path] = None,
    product_master_sheet: str = "Combo_SMILES_Master",
    product_master_combo_col: str = "Combo",
    product_master_smiles_col: str = "Product_SMILES",
    calibration_sheet: str = "",
    target_sheet: str = "",
    combo_col: str = "",
    formula_col: str = "",
    calibration_rt_col: str = "",
    target_rt_col: str = "",
    calibration_group_col: str = "",
    gradient_points: Optional[Sequence[GradientPoint]] = None,
    gradient_delay_min: float = 0.0,
    mobile_phase_a_name: str = "Mobile phase A",
    mobile_phase_b_name: str = "Mobile phase B",
    use_3d: bool = False,
    include_extended: bool = False,
    include_engineered: bool = True,
    privacy_mode: bool = True,
    output_language: str = "en",
    output_csv: bool = True,
    run_models: bool = False,
    model_calibration_mode: str = "training_original",
    model_objective: str = "trend",
    model_cv_splits: int = 5,
    model_cv_repeats: int = 3,
    model_max_features: int = 40,
    model_random_state: int = 42,
    model_auto_remove_outliers: bool = True,
    model_outlier_mode: str = "conservative",
    model_outlier_min_fold_error: float = 8.0,
    model_outlier_max_fraction_pct: float = 5.0,
    model_outlier_consensus_pct: float = 75.0,
    model_auto_select_features: bool = True,
    model_use_categorical_features: bool = False,
    model_deep_tuning: bool = False,
    model_feature_stability_threshold_pct: float = 50.0,
    model_min_selected_features: int = 6,
    model_correlation_threshold: float = 0.97,
    trend_optimizer_enabled: bool = True,
    trend_optimizer_trials: int = 30,
    trend_optimizer_pair_min_fold: float = 1.5,
    trend_optimizer_min_features: int = 6,
    trend_optimizer_max_features: int = 24,
    trend_optimizer_permutations: int = 500,
    fixed_feature_search_enabled: bool = True,
    fixed_feature_model: str = "auto_previous",
    fixed_feature_previous_best_model: str = "",
    fixed_feature_trials: int = 100,
    fixed_feature_min_features: int = 4,
    fixed_feature_max_features: int = 20,
    manual_feature_pool: Sequence[str] = (),
    level_classification_enabled: bool = True,
    dual_ie_enabled: bool = True,
    published_ie_mode: str = "negative_rf_literature",
    published_ion_mode: str = "negative",
    published_adduct: str = "[M-H]-",
    published_organic_modifier: str = "acetonitrile",
    published_aqueous_phase_name: str = "Water + 0.1% formic acid",
    published_organic_phase_name: str = "Acetonitrile",
    published_formic_acid_pct: float = 0.1,
    published_organic_formic_acid_pct: float = 0.0,
    published_aqueous_pH: float = 2.7,
    published_nh4_present: bool = False,
    published_negative_proxies_enabled: bool = True,
    published_dynamic_range_mode: str = "auto",
    published_external_logie_file: Optional[Path] = None,
    published_external_logie_sheet: str = "",
    published_external_combo_col: str = "",
    published_external_logie_col: str = "",
    published_rscript_path: str = "",
    published_instrument_name: str = "Thermo Scientific Q Exactive HF",
    published_concentration_unit: str = "mol_L",
    reuse_descriptor_workbook: Optional[Path] = None,
    excluded_injection_groups: Sequence[str] = (),
) -> EsiDescriptorReport:
    warnings: List[str] = []
    if reuse_descriptor_workbook:
        from .model_rerun import rerun_models_from_descriptor_workbook
        from .published_ie_benchmark import PublishedIEConfig
        cfg = PublishedIEConfig(
            enabled=bool(dual_ie_enabled),
            mode=str(published_ie_mode or "negative_rf_literature"),
            ion_mode=str(published_ion_mode or "negative"),
            adduct=str(published_adduct or "[M-H]-"),
            organic_modifier=str(published_organic_modifier or "acetonitrile"),
            aqueous_phase_name=str(published_aqueous_phase_name or "Water + 0.1% formic acid"),
            organic_phase_name=str(published_organic_phase_name or "Acetonitrile"),
            formic_acid_pct=float(published_formic_acid_pct),
            organic_formic_acid_pct=float(published_organic_formic_acid_pct),
            aqueous_pH=float(published_aqueous_pH),
            nh4_present=bool(published_nh4_present),
            negative_proxies_enabled=bool(published_negative_proxies_enabled),
            dynamic_range_mode=str(published_dynamic_range_mode or "auto"),
            cv_splits=max(2, int(model_cv_splits)),
            cv_repeats=max(1, int(model_cv_repeats)),
            random_state=int(model_random_state),
            output_language=str(output_language or "en"),
            external_logie_file=(Path(published_external_logie_file) if published_external_logie_file else None),
            external_logie_sheet=str(published_external_logie_sheet or ""),
            external_combo_col=str(published_external_combo_col or ""),
            external_logie_col=str(published_external_logie_col or ""),
            rscript_path=str(published_rscript_path or ""),
            instrument_name=str(published_instrument_name or "Thermo Scientific Q Exactive HF"),
            concentration_unit=str(published_concentration_unit or "mol_L"),
        )
        return rerun_models_from_descriptor_workbook(
            Path(reuse_descriptor_workbook), Path(out_xlsx),
            output_language=str(output_language or "en"),
            model_objective=str(model_objective or "trend"),
            model_cv_splits=int(model_cv_splits), model_cv_repeats=int(model_cv_repeats),
            model_max_features=int(model_max_features), model_random_state=int(model_random_state),
            model_auto_remove_outliers=bool(model_auto_remove_outliers),
            model_outlier_mode=str(model_outlier_mode),
            model_outlier_min_fold_error=float(model_outlier_min_fold_error),
            model_outlier_max_fraction_pct=float(model_outlier_max_fraction_pct),
            model_outlier_consensus_pct=float(model_outlier_consensus_pct),
            model_auto_select_features=bool(model_auto_select_features),
            model_use_categorical_features=bool(model_use_categorical_features),
            model_deep_tuning=bool(model_deep_tuning),
            model_feature_stability_threshold_pct=float(model_feature_stability_threshold_pct),
            model_min_selected_features=int(model_min_selected_features),
            model_correlation_threshold=float(model_correlation_threshold),
            trend_optimizer_enabled=bool(trend_optimizer_enabled),
            trend_optimizer_trials=int(trend_optimizer_trials),
            trend_optimizer_pair_min_fold=float(trend_optimizer_pair_min_fold),
            trend_optimizer_min_features=int(trend_optimizer_min_features),
            trend_optimizer_max_features=int(trend_optimizer_max_features),
            trend_optimizer_permutations=int(trend_optimizer_permutations),
            fixed_feature_search_enabled=bool(fixed_feature_search_enabled),
            fixed_feature_model=str(fixed_feature_model or "auto_previous"),
            fixed_feature_previous_best_model=str(fixed_feature_previous_best_model or ""),
            fixed_feature_trials=int(fixed_feature_trials),
            fixed_feature_min_features=int(fixed_feature_min_features),
            fixed_feature_max_features=int(fixed_feature_max_features),
            manual_feature_pool=tuple(manual_feature_pool or ()),
            level_classification_enabled=bool(level_classification_enabled),
            dual_ie_enabled=bool(dual_ie_enabled), published_config=cfg,
            excluded_injection_groups=excluded_injection_groups,
        )
    points = list(gradient_points or DEFAULT_GRADIENT)
    if len(points) < 2:
        raise ValueError("Gradient program needs at least two points")
    points = sorted(points, key=lambda p: p.time_min)

    mapping, master_warnings, master_meta = load_product_smiles_master(
        Path(product_master_file),
        sheet_name=product_master_sheet,
        combo_col=product_master_combo_col,
        smiles_col=product_master_smiles_col,
    )
    warnings.extend(master_warnings)
    entries = _flatten_master_entries(mapping)
    master_rows = _master_prepared_rows(entries)

    train_rows, w1 = _prepare_dataset(
        Path(calibration_file) if calibration_file else None,
        sheet_name=calibration_sheet,
        is_training=True,
        combo_col=combo_col,
        formula_col=formula_col,
        rt_col=calibration_rt_col,
        group_col=calibration_group_col,
    )
    target_rows, w2 = _prepare_dataset(
        Path(target_file) if target_file else None,
        sheet_name=target_sheet,
        is_training=False,
        combo_col=combo_col,
        formula_col=formula_col,
        rt_col=target_rt_col,
    )
    warnings.extend(w1)
    warnings.extend(w2)

    matched_train, _ = apply_product_smiles_master(train_rows, mapping, formula_verification_mode="review")
    matched_target, _ = apply_product_smiles_master(target_rows, mapping, formula_verification_mode="review")

    # Product_SMILES is already embedded in each master row.  Compute each unique
    # product only once, then copy its descriptor evidence to training/target
    # rows through the stable ordered A/B/C formula key.  This roughly halves
    # 3D work for the common 1000-master + 98-training + 1000-target workflow.
    empty_lib = StructureLibrary()
    enrich_rows(master_rows, empty_lib, use_3d=bool(use_3d), privacy_mode=bool(privacy_mode), product_smiles_col="Product_SMILES")
    master_by_key = {_combo_formula_key(row.combo): row for row in master_rows if _combo_formula_key(row.combo)}

    def copy_from_master(rows):
        missing = []
        copied = 0
        for row in rows:
            src = master_by_key.get(_combo_formula_key(row.combo))
            if src is None or not src.structure_status:
                missing.append(row)
                continue
            row.core_id = src.core_id
            row.structure_status = src.structure_status
            row.structure_method = src.structure_method
            row.structure_hash = src.structure_hash
            row.product_formula_calc = src.product_formula_calc
            row.formula_match = src.formula_match
            row.three_d_status = src.three_d_status
            row.structure_features = dict(src.structure_features)
            row.structure_fingerprint = src.structure_fingerprint
            row.component_fingerprints = dict(src.component_fingerprints)
            row.structure_warnings = list(src.structure_warnings)
            row.product_smiles = src.product_smiles
            copied += 1
        if missing:
            enrich_rows(missing, empty_lib, use_3d=bool(use_3d), privacy_mode=bool(privacy_mode), product_smiles_col="Product_SMILES")
        return copied, len(missing)

    copied_training, missing_training = copy_from_master(train_rows)
    copied_target, missing_target = copy_from_master(target_rows)
    warnings.append(
        f"Descriptor reuse within run: training copied={copied_training}, fallback={missing_training}; "
        f"target copied={copied_target}, fallback={missing_target}."
    )

    all_rows = list(master_rows) + list(train_rows) + list(target_rows)
    if include_extended:
        for row in all_rows:
            smi = str(row.raw.get("Product_SMILES", "") or "").strip()
            if smi:
                row.structure_features.update(extended_descriptors_from_smiles(smi))
    descriptor_names = _descriptor_header(all_rows, include_extended=bool(include_extended))
    if include_engineered:
        for name in ENGINEERED_DESCRIPTOR_NAMES:
            if name not in descriptor_names:
                descriptor_names.append(name)
    published_config = None
    if bool(dual_ie_enabled):
        from .published_ie_benchmark import PublishedIEConfig, PUBLISHED_ELUENT_FEATURES
        from .negative_esi_literature import NEGATIVE_ESI_FEATURES
        published_config = PublishedIEConfig(
            enabled=True,
            mode=str(published_ie_mode or "negative_rf_literature"),
            ion_mode=str(published_ion_mode or "negative"),
            adduct=str(published_adduct or "[M-H]-"),
            organic_modifier=str(published_organic_modifier or "acetonitrile"),
            aqueous_phase_name=str(published_aqueous_phase_name or "Water + 0.1% formic acid"),
            organic_phase_name=str(published_organic_phase_name or "Acetonitrile"),
            formic_acid_pct=float(published_formic_acid_pct),
            organic_formic_acid_pct=float(published_organic_formic_acid_pct),
            aqueous_pH=float(published_aqueous_pH),
            nh4_present=bool(published_nh4_present),
            negative_proxies_enabled=bool(published_negative_proxies_enabled),
            dynamic_range_mode=str(published_dynamic_range_mode or "auto"),
            cv_splits=max(2, int(model_cv_splits)),
            cv_repeats=max(1, int(model_cv_repeats)),
            random_state=int(model_random_state),
            output_language=str(output_language or "en"),
            external_logie_file=(Path(published_external_logie_file) if published_external_logie_file else None),
            external_logie_sheet=str(published_external_logie_sheet or ""),
            external_combo_col=str(published_external_combo_col or ""),
            external_logie_col=str(published_external_logie_col or ""),
            rscript_path=str(published_rscript_path or ""),
            instrument_name=str(published_instrument_name or "Thermo Scientific Q Exactive HF"),
            concentration_unit=str(published_concentration_unit or "mol_L"),
        )
        for name in PUBLISHED_ELUENT_FEATURES:
            if name not in descriptor_names:
                descriptor_names.append(name)
        if published_config.negative_proxies_enabled and (
            published_config.ion_mode == "negative" or published_config.mode == "negative_rf_literature"
        ):
            for name in NEGATIVE_ESI_FEATURES:
                if name not in descriptor_names:
                    descriptor_names.append(name)
    base_headers = [
        "Dataset", "Source_file", "Source_sheet", "Source_row", "Source_index", "Name", "Formula", "Combo",
        "A_Formula", "B_Formula", "C_Formula", "Injection_Group", "Injection_Group_Column", "Measured_ratio", "Actual_concentration", "Exact_mass", "DBE",
        "Apex_RT_min", "Gradient_delay_min", "Effective_gradient_time_min", "Mobile_phase_A_pct", "Mobile_phase_B_pct",
        "B_slope_pct_per_min", "Gradient_segment_start_min", "Gradient_segment_end_min", "Gradient_state",
        "Gradient_range_status", "Structure_status", "Structure_method", "Structure_hash", "Product_formula_calc",
        "Formula_match", "3D_status", "Descriptor_count", "Product_Master_Match", "Product_Master_Formula_Key",
    ]
    if not privacy_mode:
        base_headers.append("Product_SMILES")
    all_headers = base_headers + descriptor_names + ["Warnings"]

    master_records = [_row_to_record("Master", r, descriptor_names, points, gradient_delay_min, privacy_mode=privacy_mode, include_engineered=bool(include_engineered)) for r in master_rows]
    training_records = [_row_to_record("Training", r, descriptor_names, points, gradient_delay_min, privacy_mode=privacy_mode, include_engineered=bool(include_engineered)) for r in train_rows]
    target_records = [_row_to_record("Target", r, descriptor_names, points, gradient_delay_min, privacy_mode=privacy_mode, include_engineered=bool(include_engineered)) for r in target_rows]
    if published_config is not None:
        from .published_ie_benchmark import add_published_eluent_features
        from .negative_esi_literature import add_negative_esi_proxy_features
        add_published_eluent_features(master_records, published_config)
        add_published_eluent_features(training_records, published_config)
        add_published_eluent_features(target_records, published_config)
        if published_config.negative_proxies_enabled and (
            published_config.ion_mode == "negative" or published_config.mode == "negative_rf_literature"
        ):
            for dataset_records in (master_records, training_records, target_records):
                add_negative_esi_proxy_features(
                    dataset_records,
                    aqueous_pH=published_config.aqueous_pH,
                    aqueous_formic_acid_pct=published_config.formic_acid_pct,
                    organic_formic_acid_pct=published_config.organic_formic_acid_pct,
                )
    records: List[Dict[str, object]] = list(master_records) + list(training_records) + list(target_records)
    # Optional published/negative-ESI branches add public metadata after the
    # initial header list is assembled.  Rebuild the final public header list
    # here so both Excel and CSV include those fields, while private keys are
    # excluded unconditionally.
    all_headers = _public_record_headers(records, all_headers)
    bridge = _overlap_records(train_rows, target_rows, points, gradient_delay_min)

    try:
        from openpyxl import Workbook
    except Exception as e:
        raise RuntimeError("openpyxl is required to write the descriptor workbook") from e

    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_key = wb.active
    ws_key.title = "Key_ESI_Summary"
    key_headers = [
        "Dataset", "Source_file", "Source_sheet", "Source_row", "Name", "Formula", "Combo",
        "Injection_Group", "Measured_ratio", "Actual_concentration", "Exact_mass", "DBE", "Apex_RT_min",
        "Mobile_phase_A_pct", "Mobile_phase_B_pct", "B_slope_pct_per_min", "Gradient_state",
        "Structure_status", "Formula_match", "AcidicSiteCount_proxy", "BasicSiteCount_proxy",
        "IonizableSiteCount_proxy", "HBD", "HBA", "FormalCharge", "GasteigerChargeRange",
        "MolLogP", "TPSA", "PolarSurfaceFraction_proxy", "FractionCSP3", "AromaticRingCount",
        "RotatableBonds", "MolWt", "LabuteASA", "MolMR", "MolVolume3D", "RadiusOfGyration",
        "Asphericity", "Published_Viscosity_mPa_s", "Published_Surface_Tension_mN_m",
        "Published_Polarity_Index", "Published_Aqueous_pH", "Published_NH4_Present",
        "Structure_hash", "Warnings",
    ]
    _append_records(ws_key, records, key_headers)

    ws = wb.create_sheet("ESI_Features")
    _append_records(ws, records, all_headers)

    ws_bridge = wb.create_sheet("Bridge_Overlap")
    bridge_headers = list(bridge[0].keys()) if bridge else [
        "ABC_Formula_Key", "Combo_training", "Combo_target", "Training_Apex_RT_min", "Target_Apex_RT_min",
        "Delta_RT_target_minus_training_min", "Training_B_pct", "Target_B_pct",
        "Delta_B_target_minus_training_pct", "Training_ratio", "Target_ratio", "Target_to_training_ratio_fold",
    ]
    _append_records(ws_bridge, bridge, bridge_headers)

    ws_grad = wb.create_sheet("Gradient_Program")
    ws_grad.append(["Time_min", "Mobile_phase_A_pct", "Mobile_phase_B_pct", "A_name", "B_name"])
    for p in points:
        ws_grad.append([p.time_min, p.a_pct, p.b_pct, mobile_phase_a_name, mobile_phase_b_name])
    ws_grad.append([])
    ws_grad.append(["Setting", "Value", "Meaning"])
    ws_grad.append(["Gradient_delay_min", float(gradient_delay_min), "Effective gradient time = Apex_RT - Gradient_delay_min"])
    ws_grad.append(["Interpolation", "linear", "A/B percentages are linearly interpolated between programmed time points"])
    if published_config is not None:
        ws_grad.append(["Published_IE_Aqueous_phase", published_config.aqueous_phase_name, "Dual IE benchmark setting"])
        ws_grad.append(["Published_IE_Organic_phase", published_config.organic_phase_name, "Verify this solvent; default is acetonitrile"])
        ws_grad.append(["Published_IE_Aqueous_formic_acid_pct", published_config.formic_acid_pct, "Aqueous-phase additive concentration"])
        ws_grad.append(["Published_IE_Organic_formic_acid_pct", published_config.organic_formic_acid_pct, "Organic-phase additive concentration"])
        ws_grad.append(["Published_IE_Dynamic_range_mode", published_config.dynamic_range_mode, "Cross-fitted high-range response calibration"])
        ws_grad.append(["Published_IE_Aqueous_pH", published_config.aqueous_pH, "Editable approximation; measured pH preferred"])
        ws_grad.append(["Published_IE_NH4_present", published_config.nh4_present, "Binary published eluent descriptor"])

    ws_def = wb.create_sheet("Descriptor_Definitions")
    ws_def.append(["Category", "Descriptor", "Interpretation / limitation"])
    definitions = [
        ("Important", "Acid/base proxies", "SMARTS and Gasteiger structural proxies only; they are not predicted pKa values."),
        ("Important", "Gradient composition", "Programmed composition at corrected Apex RT; set delay to account for known system dwell time."),
        ("Important", "3D descriptors", "One deterministic ETKDG conformer; not a solution-phase conformational ensemble."),
        ("Acid/base", "AcidicSiteCount_proxy", "Count of selected acidic motifs: carboxylic, sulfonic, phosphoric OH, phenol, thiol and imide NH."),
        ("Acid/base", "BasicSiteCount_proxy", "Count of selected basic motifs: aliphatic amines, pyridine-like N, amidine/guanidine, imine N and quaternary ammonium."),
        ("Acid/base", "GasteigerCharge*", "Approximate atomic partial-charge summaries; useful as relative descriptors, not quantum-chemical charges."),
        ("Hydrophobicity", "MolLogP", "RDKit Crippen cLogP estimate."),
        ("Polarity", "TPSA", "Topological polar surface area."),
        ("Polarity", "PolarSurfaceFraction_proxy", "TPSA divided by LabuteASA."),
        ("Size/volume", "LabuteASA", "Approximate molecular surface area."),
        ("Size/volume", "MolMR", "Crippen molar refractivity; correlates with molecular volume/polarizability."),
        ("Size/volume", "MolVolume3D", "RDKit grid-based volume from the generated 3D conformer; blank when 3D is disabled or embedding fails."),
        ("Gradient", "Mobile_phase_A_pct / Mobile_phase_B_pct", "Linearly interpolated from the supplied gradient table."),
        ("Gradient", "B_slope_pct_per_min", "Local programmed rate of change of mobile phase B."),
        ("Bridge", "Bridge_Overlap", "Matches compounds present in both training and target by ordered A/B/C component formulas to inspect RT and composition shifts."),
    ]
    for row in definitions:
        ws_def.append(list(row))
    for group, names in DESCRIPTOR_GROUPS:
        for name in names:
            if name in descriptor_names and not any(r[1] == name for r in definitions):
                ws_def.append([group, name, "RDKit-derived molecular descriptor; inspect alongside the category-level definitions above."])
    if include_extended:
        ws_def.append(["Extended", "PEOE_VSA / SlogP_VSA / SMR_VSA / EState_VSA / BCUT2D", "Extended charge, lipophilicity, refractivity, E-state and topology descriptors for later model benchmarking."])

    # A compact bilingual dictionary is always available.  Technical keys stay
    # stable regardless of the selected output language.
    ws_guide = wb.create_sheet("Descriptor_Guide")
    guide_headers = ["Descriptor", 'Display_name', "English_name", 'Definition', "English_definition", "Unit_or_type", "Source", "Requires_3D"]
    ws_guide.append(guide_headers)
    guide_names = list(dict.fromkeys([
        "Exact_mass", "DBE", "Apex_RT_min", "Effective_gradient_time_min",
        "Mobile_phase_A_pct", "Mobile_phase_B_pct", "B_slope_pct_per_min", "Gradient_state",
    ] + list(descriptor_names)))
    for name in guide_names:
        meta = descriptor_meta(name)
        ws_guide.append([
            name, meta["Name_zh"], meta["Name_en"], meta["Definition_zh"],
            meta["Definition_en"], meta["Unit_or_type"], meta["Source"], meta["Requires_3D"],
        ])

    ws_qc = wb.create_sheet("Diagnostics")
    ws_qc.append(["Metric", "Value"])
    metrics = {
        "RDKit_available": bool(RDKIT_AVAILABLE),
        "Product_master_rows": int(master_meta.get("rows", 0) or 0),
        "Product_master_usable": int(master_meta.get("usable", 0) or 0),
        "Master_descriptor_rows": len(master_rows),
        "Training_rows": len(train_rows),
        "Training_structure_matches": matched_train,
        "Training_rows_with_RT": sum(1 for r in train_rows if _guess_rt(r) is not None),
        "Target_rows": len(target_rows),
        "Target_structure_matches": matched_target,
        "Target_rows_with_RT": sum(1 for r in target_rows if _guess_rt(r) is not None),
        "Bridge_overlap_rows": len(bridge),
        "Descriptor_columns": len(descriptor_names),
        "3D_enabled": bool(use_3d),
        "Extended_descriptors_enabled": bool(include_extended),
        "Engineered_physical_proxies_enabled": bool(include_engineered),
        "Output_language": str(output_language),
        "Privacy_mode": bool(privacy_mode),
        "Multi_model_requested": bool(run_models),
        "Model_calibration_mode_requested": str(model_calibration_mode),
        "Model_objective_requested": str(model_objective),
        "Calibration_group_column_requested": str(calibration_group_col),
        "Model_CV_splits": int(model_cv_splits),
        "Model_CV_repeats": int(model_cv_repeats),
        "Model_max_selected_features": int(model_max_features),
        "Model_auto_remove_outliers": bool(model_auto_remove_outliers),
        "Model_outlier_mode": str(model_outlier_mode),
        "Model_outlier_min_fold_error": float(model_outlier_min_fold_error),
        "Model_outlier_max_fraction_pct": float(model_outlier_max_fraction_pct),
        "Model_outlier_consensus_pct": float(model_outlier_consensus_pct),
        "Model_auto_select_features": bool(model_auto_select_features),
        "Model_use_categorical_features": bool(model_use_categorical_features),
        "Model_deep_tuning": bool(model_deep_tuning),
        "Model_feature_stability_threshold_pct": float(model_feature_stability_threshold_pct),
        "Model_min_selected_features": int(model_min_selected_features),
        "Model_correlation_threshold": float(model_correlation_threshold),
        "Trend_optimizer_enabled": bool(trend_optimizer_enabled),
        "Trend_optimizer_trials": int(trend_optimizer_trials),
        "Trend_optimizer_pair_min_fold": float(trend_optimizer_pair_min_fold),
        "Trend_optimizer_min_features": int(trend_optimizer_min_features),
        "Trend_optimizer_max_features": int(trend_optimizer_max_features),
        "Trend_optimizer_permutations": int(trend_optimizer_permutations),
        "Fixed_feature_search_enabled": bool(fixed_feature_search_enabled),
        "Fixed_feature_model_requested": str(fixed_feature_model),
        "Fixed_feature_previous_best_model": str(fixed_feature_previous_best_model),
        "Fixed_feature_trials": int(fixed_feature_trials),
        "Fixed_feature_count_range": f"{int(fixed_feature_min_features)}-{int(fixed_feature_max_features)}",
        "Manual_feature_pool_count": len(tuple(manual_feature_pool or ())),
        "Level_classification_enabled": bool(level_classification_enabled),
        "Dual_IE_enabled": bool(dual_ie_enabled),
        "Published_IE_mode": str(published_ie_mode),
        "Published_ion_mode": str(published_ion_mode),
        "Published_adduct": str(published_adduct),
        "Published_organic_modifier": str(published_organic_modifier),
        "Published_aqueous_phase": str(published_aqueous_phase_name),
        "Published_organic_phase": str(published_organic_phase_name),
        "Published_formic_acid_pct": float(published_formic_acid_pct),
        "Published_aqueous_pH": float(published_aqueous_pH),
        "Published_NH4_present": bool(published_nh4_present),
        "Published_Rscript_path": str(published_rscript_path),
        "Published_instrument": str(published_instrument_name),
        "Published_concentration_unit": str(published_concentration_unit),
    }
    for k, v in metrics.items():
        ws_qc.append([k, v])
    ws_qc.append([])
    ws_qc.append(["Warning"])
    for w in warnings:
        ws_qc.append([w])

    benchmark_result = None
    trend_result = None
    fixed_result = None
    level_result = None
    published_result = None
    dual_result = None
    if run_models:
        try:
            from .esi_model_benchmark import run_model_benchmark, append_benchmark_to_workbook
            from .model_rerun import parse_excluded_groups, _filter_groups
            excluded_groups = parse_excluded_groups(excluded_injection_groups)
            model_training_records, group_exclusion_audit = _filter_groups(training_records, excluded_groups)
            if excluded_groups:
                warnings.append("Excluded standard injection group(s) from model fitting: " + ", ".join(excluded_groups))
                _append_simple_dict_sheet(wb, "Group_Exclusion_Audit", group_exclusion_audit)
                ws_qc.append(["Excluded_injection_groups", "; ".join(excluded_groups)])
                ws_qc.append(["Training_rows_after_group_exclusion", len(model_training_records)])
            benchmark_result = run_model_benchmark(
                model_training_records, target_records, descriptor_names, out_xlsx,
                calibration_mode="training_original",
                model_objective=str(model_objective or "trend"),
                cv_splits=max(2, int(model_cv_splits)),
                cv_repeats=max(1, int(model_cv_repeats)),
                max_features=max(3, int(model_max_features)),
                random_state=int(model_random_state),
                expected_descriptor_names=_key_descriptor_names(),
                three_d_enabled=bool(use_3d),
                output_language=str(output_language or "en"),
                auto_remove_outliers=bool(model_auto_remove_outliers),
                outlier_mode=str(model_outlier_mode or "conservative"),
                outlier_min_fold_error=max(1.01, float(model_outlier_min_fold_error)),
                outlier_max_fraction_pct=max(0.0, float(model_outlier_max_fraction_pct)),
                outlier_consensus_pct=max(0.0, min(100.0, float(model_outlier_consensus_pct))),
                auto_select_features=bool(model_auto_select_features),
                use_categorical_features=bool(model_use_categorical_features),
                deep_tuning=bool(model_deep_tuning),
                feature_stability_threshold_pct=float(model_feature_stability_threshold_pct),
                min_selected_features=max(2, int(model_min_selected_features)),
                correlation_threshold=float(model_correlation_threshold),
            )
            append_benchmark_to_workbook(wb, benchmark_result)
            if bool(trend_optimizer_enabled):
                try:
                    from .trend_rank_optimizer import run_trend_optimizer, append_trend_optimizer_to_workbook
                    trend_result = run_trend_optimizer(
                        benchmark_result.calibration_records,
                        target_records,
                        descriptor_names,
                        out_xlsx,
                        trials=max(3, int(trend_optimizer_trials)),
                        pair_min_fold=max(1.01, float(trend_optimizer_pair_min_fold)),
                        min_features=max(2, int(trend_optimizer_min_features)),
                        max_features=max(3, int(trend_optimizer_max_features)),
                        cv_splits=max(2, int(model_cv_splits)),
                        cv_repeats=max(1, int(model_cv_repeats)),
                        permutations=max(20, int(trend_optimizer_permutations)),
                        random_state=int(model_random_state),
                        correlation_threshold=float(model_correlation_threshold),
                        output_language=str(output_language or "en"),
                        concentration_unit=str(published_concentration_unit or ""),
                    )
                    append_trend_optimizer_to_workbook(wb, trend_result)
                    ws_qc.append(["Trend_optimizer", "completed"])
                    ws_qc.append(["Trend_optimizer_best_model", trend_result.best_model])
                    ws_qc.append(["Trend_optimizer_decision", trend_result.decision])
                    ws_qc.append(["Trend_optimizer_summary", trend_result.decision_summary])
                    for warning in trend_result.warnings:
                        warnings.append("Trend optimizer: " + warning)
                except Exception as trend_exc:
                    warnings.append(f"Trend optimizer did not complete: {trend_exc}")
                    ws_qc.append(["Trend_optimizer", "failed"])
                    ws_qc.append(["Trend_optimizer_error", str(trend_exc)])
            if bool(fixed_feature_search_enabled):
                try:
                    from .fixed_feature_search import run_fixed_feature_search, append_fixed_feature_search_to_workbook
                    previous_best = str(fixed_feature_previous_best_model or "").strip()
                    if not previous_best:
                        previous_best = trend_result.best_model if trend_result is not None else benchmark_result.best_model
                    fixed_result = run_fixed_feature_search(
                        benchmark_result.calibration_records,
                        target_records,
                        descriptor_names,
                        out_xlsx,
                        fixed_model=str(fixed_feature_model or "auto_previous"),
                        previous_best_model=previous_best,
                        manual_feature_pool=tuple(manual_feature_pool or ()),
                        trial_count=max(5, int(fixed_feature_trials)),
                        min_features=max(1, int(fixed_feature_min_features)),
                        max_features=max(1, int(fixed_feature_max_features)),
                        cv_splits=max(2, int(model_cv_splits)),
                        cv_repeats=max(1, int(model_cv_repeats)),
                        random_state=int(model_random_state),
                        correlation_threshold=float(model_correlation_threshold),
                        output_language=str(output_language or "en"),
                        concentration_unit=str(published_concentration_unit or ""),
                    )
                    append_fixed_feature_search_to_workbook(wb, fixed_result)
                    ws_qc.append(["Fixed_feature_search", "completed"])
                    ws_qc.append(["Fixed_feature_model", fixed_result.model])
                    ws_qc.append(["Fixed_feature_candidate_count", len(fixed_result.candidate_features)])
                    ws_qc.append(["Fixed_feature_best_absolute_trial", fixed_result.top5_absolute[0].get("Trial", "") if fixed_result.top5_absolute else ""])
                    ws_qc.append(["Fixed_feature_best_trend_trial", fixed_result.top5_trend[0].get("Trial", "") if fixed_result.top5_trend else ""])
                    for warning in fixed_result.warnings:
                        warnings.append("Fixed feature search: " + str(warning))
                except Exception as fixed_exc:
                    warnings.append(f"Fixed-model feature search did not complete: {fixed_exc}")
                    ws_qc.append(["Fixed_feature_search", "failed"])
                    ws_qc.append(["Fixed_feature_search_error", str(fixed_exc)])
            if bool(level_classification_enabled):
                try:
                    from .concentration_classification import (
                        run_concentration_level_classification,
                        append_concentration_level_to_workbook,
                    )
                    selected_names = [
                        str(row.get("Feature", ""))
                        for row in benchmark_result.selected_descriptors
                        if str(row.get("Feature", "")).strip()
                    ]
                    level_result = run_concentration_level_classification(
                        benchmark_result.calibration_records,
                        target_records,
                        descriptor_names,
                        out_xlsx,
                        selected_feature_names=selected_names,
                        manual_feature_pool=tuple(manual_feature_pool or ()),
                        cv_splits=max(2, int(model_cv_splits)),
                        cv_repeats=max(1, int(model_cv_repeats)),
                        max_features=max(3, min(30, int(model_max_features))),
                        random_state=int(model_random_state),
                        level_count=10,
                        output_language=str(output_language or "en"),
                        concentration_unit=str(published_concentration_unit or ""),
                    )
                    append_concentration_level_to_workbook(wb, level_result)
                    ws_qc.append(["Level_classification", "completed"])
                    ws_qc.append(["Level_classification_best_model", level_result.best_model])
                    ws_qc.append(["Level_classification_decision", level_result.decision])
                    best_level_row = next((r for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), {})
                    ws_qc.append(["Level_classification_exact_pct", best_level_row.get("Exact_level_accuracy_pct", "")])
                    ws_qc.append(["Level_classification_within1_pct", best_level_row.get("Within_1_level_pct", "")])
                    ws_qc.append(["Level_classification_all_hit", best_level_row.get("All_exactly_hit", False)])
                    for warning in level_result.warnings:
                        warnings.append("Level classification: " + str(warning))
                except Exception as level_exc:
                    warnings.append(f"Ten-level concentration classification did not complete: {level_exc}")
                    ws_qc.append(["Level_classification", "failed"])
                    ws_qc.append(["Level_classification_error", str(level_exc)])
            if bool(dual_ie_enabled) and published_config is not None:
                from .published_ie_benchmark import (
                    run_published_ie_benchmark, run_dual_ie_validation, append_published_ie_to_workbook,
                )
                valid_target_keys = {str(r.get("ABC_Formula_Key", "")) for r in benchmark_result.target_predictions}
                published_targets = [
                    r for r in target_records
                    if str(r.get("Product_Master_Formula_Key", "")) in valid_target_keys
                    and _optional_float(r.get("Measured_ratio")) is not None
                    and float(r.get("Measured_ratio")) > 0
                ]
                published_result = run_published_ie_benchmark(
                    benchmark_result.calibration_records, published_targets, descriptor_names, out_xlsx, published_config,
                )
                dual_result = run_dual_ie_validation(
                    benchmark_result, published_result, out_xlsx, output_language=str(output_language or "en"),
                )
                append_published_ie_to_workbook(wb, published_result, dual_result)
                ws_qc.append(["Dual_IE_validation", "completed"])
                ws_qc.append(["Published_IE_method", published_result.method_code])
                ws_qc.append(["Published_IE_Spearman", published_result.metrics.get("Trend_Spearman_r", "")])
                ws_qc.append(["Published_IE_Median_Fold", published_result.metrics.get("Median_Fold_Error", "")])
                ws_qc.append(["Dual_target_rank_Spearman", dual_result.target_rank_spearman if dual_result.target_rank_spearman is not None else ""])
                ws_qc.append(["Dual_IE_decision", dual_result.decision_label])
                for warning in dual_result.warnings:
                    warnings.append("DUAL_IE: " + str(warning))
                    ws_qc.append(["Dual_IE_warning", str(warning)])
            try:
                from .simple_prediction_outputs import build_simple_prediction_rows, append_simple_outputs_to_workbook
                simple_known_rows, simple_target_rows = build_simple_prediction_rows(
                    benchmark_result,
                    trend_result,
                    level_result,
                    target_records,
                    concentration_unit=str(published_concentration_unit or ""),
                )
                append_simple_outputs_to_workbook(
                    wb, simple_known_rows, simple_target_rows, out_xlsx,
                    output_language=str(output_language or "en"),
                )
                ws_qc.append(["Known_Standards_Predictions_rows", len(simple_known_rows)])
                ws_qc.append(["Unknown_Sample_Predictions_rows", len(simple_target_rows)])
            except Exception as simple_exc:
                warnings.append(f"Concise prediction tables/plot did not complete: {simple_exc}")
                ws_qc.append(["Concise_prediction_outputs", "failed"])
                ws_qc.append(["Concise_prediction_error", str(simple_exc)])
            ws_qc.append([])
            ws_qc.append(["Multi-model benchmark", "completed"])
            ws_qc.append(["Calibration_mode_used", benchmark_result.calibration_mode_used])
            ws_qc.append(["Model_objective_used", benchmark_result.model_objective_used])
            ws_qc.append(["Unknown_target_used_for_supervision", False])
            ws_qc.append(["Model_calibration_rows", benchmark_result.n_calibration])
            ws_qc.append(["Model_target_rows", benchmark_result.n_target_valid])
            ws_qc.append(["Best_model", benchmark_result.best_model])
            ws_qc.append(["Best_model_rating", benchmark_result.best_rating])
            ws_qc.append(["Model_decision", benchmark_result.decision_label])
            ws_qc.append(["Model_decision_summary", benchmark_result.decision_summary])
            ws_qc.append(["Best_Median_Fold_Error", benchmark_result.best_metrics.get("Median_Fold_Error", "")])
            ws_qc.append(["Best_P80_Fold_Error", benchmark_result.best_metrics.get("P80_Fold_Error", "")])
            ws_qc.append(["Best_Within_2x_pct", benchmark_result.best_metrics.get("Within_2x_pct", "")])
            ws_qc.append(["Best_Within_5x_pct", benchmark_result.best_metrics.get("Within_5x_pct", "")])
            ws_qc.append(["Best_R2_log", benchmark_result.best_metrics.get("R2_log_concentration", "")])
            ws_qc.append(["Best_RMSE_log10", benchmark_result.best_metrics.get("RMSE_log10", "")])
            ws_qc.append(["Improvement_vs_GlobalMedian_pct", benchmark_result.best_metrics.get("Median_Fold_Improvement_vs_Global_pct", "")])
            ws_qc.append(["Corrected_Spearman", benchmark_result.best_metrics.get("Trend_Spearman_r", "")])
            ws_qc.append(["Raw_ratio_Spearman", benchmark_result.best_metrics.get("Raw_Trend_Spearman_r", "")])
            ws_qc.append(["Delta_Spearman_vs_raw", benchmark_result.best_metrics.get("Delta_Trend_Spearman_r_vs_raw", "")])
            ws_qc.append(["Pairwise_concordance_pct", benchmark_result.best_metrics.get("Trend_Pairwise_Concordance_pct", "")])
            ws_qc.append(["Top20_overlap_pct", benchmark_result.best_metrics.get("Trend_Top20_Overlap_pct", "")])
            ws_qc.append(["Calibration_rows_before_QC", benchmark_result.n_calibration_before_qc])
            ws_qc.append(["Calibration_outliers_auto_excluded", benchmark_result.n_outliers_removed])
            ws_qc.append(["Selected_raw_descriptors", benchmark_result.n_selected_descriptors])
            ws_qc.append(["Unavailable_or_too_sparse_descriptors", benchmark_result.n_unavailable_descriptors])
            for warning in benchmark_result.warnings:
                warnings.append("MODEL: " + str(warning))
                ws_qc.append(["Model_warning", str(warning)])
        except Exception as e:
            warning = f"Multi-model benchmark failed; descriptor workbook was still created: {e}"
            warnings.append(warning)
            ws_qc.append([])
            ws_qc.append(["Multi-model benchmark", "FAILED"])
            ws_qc.append(["Reason", str(e)])

    _style_workbook(wb)
    if "ESI_Features" in wb.sheetnames:
        wb["ESI_Features"].sheet_state = "hidden"
    if "Known_Standards_Predictions" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Known_Standards_Predictions")
    elif level_result is not None and "Level_Model_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Level_Model_Decision")
    elif dual_result is not None and "Dual_Model_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Dual_Model_Decision")
    elif benchmark_result is not None and "Model_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Model_Decision")
    else:
        wb.active = wb.sheetnames.index("Key_ESI_Summary")
    wb.save(str(out_xlsx))

    out_csv: Optional[Path] = None
    if output_csv:
        out_csv = out_xlsx.with_name(out_xlsx.stem + "__ESI_Features.csv")
        with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
            # ``extrasaction="ignore"`` is an additional safeguard against
            # future optional branches adding internal-only fields after the
            # public header list is built.
            writer = csv.DictWriter(f, fieldnames=all_headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(records)

    return EsiDescriptorReport(
        output_xlsx=out_xlsx,
        output_csv=out_csv,
        n_master=len(master_rows),
        n_training=len(train_rows),
        n_target=len(target_rows),
        n_training_matched=matched_train,
        n_target_matched=matched_target,
        n_rt_training=sum(1 for r in train_rows if _guess_rt(r) is not None),
        n_rt_target=sum(1 for r in target_rows if _guess_rt(r) is not None),
        warnings=warnings,
        model_benchmark_enabled=benchmark_result is not None,
        model_calibration_rows=(benchmark_result.n_calibration if benchmark_result is not None else 0),
        model_target_rows=(benchmark_result.n_target_valid if benchmark_result is not None else 0),
        best_model=(benchmark_result.best_model if benchmark_result is not None else ""),
        best_model_rating=(benchmark_result.best_rating if benchmark_result is not None else ""),
        model_decision=(benchmark_result.decision_label if benchmark_result is not None else ""),
        model_decision_summary=(benchmark_result.decision_summary if benchmark_result is not None else ""),
        best_median_fold_error=(_optional_float(benchmark_result.best_metrics.get("Median_Fold_Error")) if benchmark_result is not None else None),
        best_p80_fold_error=(_optional_float(benchmark_result.best_metrics.get("P80_Fold_Error")) if benchmark_result is not None else None),
        best_within_2x_pct=(_optional_float(benchmark_result.best_metrics.get("Within_2x_pct")) if benchmark_result is not None else None),
        best_within_5x_pct=(_optional_float(benchmark_result.best_metrics.get("Within_5x_pct")) if benchmark_result is not None else None),
        best_r2_log=(_optional_float(benchmark_result.best_metrics.get("R2_log_concentration")) if benchmark_result is not None else None),
        best_rmse_log10=(_optional_float(benchmark_result.best_metrics.get("RMSE_log10")) if benchmark_result is not None else None),
        improvement_vs_global_pct=(_optional_float(benchmark_result.best_metrics.get("Median_Fold_Improvement_vs_Global_pct")) if benchmark_result is not None else None),
        model_outliers_removed=(benchmark_result.n_outliers_removed if benchmark_result is not None else 0),
        model_selected_descriptors=(benchmark_result.n_selected_descriptors if benchmark_result is not None else 0),
        model_unavailable_descriptors=(benchmark_result.n_unavailable_descriptors if benchmark_result is not None else 0),
        model_objective=(benchmark_result.model_objective_used if benchmark_result is not None else ""),
        corrected_spearman=(_optional_float(benchmark_result.best_metrics.get("Trend_Spearman_r")) if benchmark_result is not None else None),
        raw_spearman=(_optional_float(benchmark_result.best_metrics.get("Raw_Trend_Spearman_r")) if benchmark_result is not None else None),
        delta_spearman=(_optional_float(benchmark_result.best_metrics.get("Delta_Trend_Spearman_r_vs_raw")) if benchmark_result is not None else None),
        pairwise_concordance_pct=(_optional_float(benchmark_result.best_metrics.get("Trend_Pairwise_Concordance_pct")) if benchmark_result is not None else None),
        top20_overlap_pct=(_optional_float(benchmark_result.best_metrics.get("Trend_Top20_Overlap_pct")) if benchmark_result is not None else None),
        injection_group_spearman=(
            _optional_float(next((r.get("Trend_Spearman_r") for r in benchmark_result.injection_group_summary if str(r.get("Injection_Group")) == "ALL_GROUPS"), None))
            if benchmark_result is not None else None
        ),
        dual_ie_enabled=(dual_result is not None),
        dual_ie_decision=(dual_result.decision_label if dual_result is not None else ""),
        dual_ie_summary=(dual_result.decision_summary if dual_result is not None else ""),
        dual_ie_published_method=(published_result.method_code if published_result is not None else ""),
        dual_ie_target_rank_spearman=(dual_result.target_rank_spearman if dual_result is not None else None),
        published_ie_spearman=(_optional_float(published_result.metrics.get("Trend_Spearman_r")) if published_result is not None else None),
        published_ie_median_fold=(_optional_float(published_result.metrics.get("Median_Fold_Error")) if published_result is not None else None),
        trend_optimizer_enabled=(trend_result is not None),
        trend_optimizer_best_model=(trend_result.best_model if trend_result is not None else ""),
        trend_optimizer_decision=(trend_result.decision if trend_result is not None else ""),
        trend_optimizer_summary=(trend_result.decision_summary if trend_result is not None else ""),
        trend_optimizer_spearman=(
            _optional_float(next((r.get("Trend_Spearman_r") for r in trend_result.comparison if str(r.get("Model")) == trend_result.best_model), None))
            if trend_result is not None else None
        ),
        trend_optimizer_delta_spearman=(
            _optional_float(next((r.get("Delta_Spearman_vs_raw") for r in trend_result.comparison if str(r.get("Model")) == trend_result.best_model), None))
            if trend_result is not None else None
        ),
        trend_optimizer_permutation_p=(
            _optional_float(next((r.get("Spearman_permutation_p") for r in trend_result.comparison if str(r.get("Model")) == trend_result.best_model), None))
            if trend_result is not None else None
        ),
        fixed_feature_search_enabled=(fixed_result is not None),
        fixed_feature_model=(fixed_result.model if fixed_result is not None else ""),
        fixed_feature_best_absolute_trial=(
            int(fixed_result.top5_absolute[0].get("Trial")) if fixed_result is not None and fixed_result.top5_absolute else None
        ),
        fixed_feature_best_trend_trial=(
            int(fixed_result.top5_trend[0].get("Trial")) if fixed_result is not None and fixed_result.top5_trend else None
        ),
        level_classification_enabled=(level_result is not None),
        level_classification_best_model=(level_result.best_model if level_result is not None else ""),
        level_classification_decision=(level_result.decision if level_result is not None else ""),
        level_classification_exact_pct=(
            _optional_float(next((r.get("Exact_level_accuracy_pct") for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), None))
            if level_result is not None else None
        ),
        level_classification_within1_pct=(
            _optional_float(next((r.get("Within_1_level_pct") for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), None))
            if level_result is not None else None
        ),
        level_classification_all_hit=(
            bool(next((r.get("All_exactly_hit") for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), False))
            if level_result is not None else False
        ),
    )
