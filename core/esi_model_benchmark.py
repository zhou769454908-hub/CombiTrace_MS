"""Small-sample multi-model benchmark for ESI response-factor prediction.

The model target is the logarithm of the relative response factor (RRF):

    RRF = (Area / Internal Standard Area) / Actual Concentration
    y   = log10(RRF)

After predicting ``y`` for a target compound, concentration is recovered as:

    Predicted Concentration = Measured Ratio / 10**y_pred

The calibration set is strictly the known standard set: concentration,
area/internal-standard ratio, retention time and gradient descriptors must come
from the same standard injection.  The unknown 1000-compound sample is never
used to construct a supervised calibration label, even when it contains some of
the same compounds.

The primary objective may be response correction and concentration-trend
recovery rather than exact point concentration.  Models still predict log10 RRF
for physical interpretability, but are selected and evaluated using out-of-fold
rank agreement, pairwise ordering and top-abundance recovery when trend mode is
chosen. Fold-level estimators fit their own preprocessing and feature selection.
Initial descriptor availability/redundancy screening uses the current calibration
set, and optional outlier screening conditions the subsequent CV results; see
docs/METHODS.md for the validation boundary.
"""
from __future__ import annotations

import hashlib
import math
import re
import statistics
import tempfile
import warnings as pywarnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np

from .esi_descriptor_meta import descriptor_meta, display_name, group_display, model_compatibility, status_display
from .concentration_levels import add_decile_columns

SKLEARN_AVAILABLE = False
SKLEARN_ERROR = ""
try:
    from sklearn.base import BaseEstimator, TransformerMixin, clone
    from sklearn.compose import ColumnTransformer
    from sklearn.cross_decomposition import PLSRegression
    from sklearn.ensemble import ExtraTreesRegressor, HistGradientBoostingRegressor, RandomForestRegressor
    from sklearn.feature_selection import SelectKBest, f_regression
    from sklearn.gaussian_process import GaussianProcessRegressor
    from sklearn.gaussian_process.kernels import ConstantKernel, RBF, WhiteKernel
    from sklearn.impute import SimpleImputer
    from sklearn.inspection import permutation_importance
    from sklearn.linear_model import BayesianRidge
    from sklearn.model_selection import GridSearchCV, KFold, RepeatedKFold
    from sklearn.neighbors import KNeighborsRegressor
    from sklearn.pipeline import Pipeline
    from sklearn.preprocessing import OneHotEncoder, StandardScaler
    from sklearn.svm import SVR
    SKLEARN_AVAILABLE = True
except Exception as exc:  # optional dependency; descriptor export remains usable
    SKLEARN_ERROR = str(exc)


CALIBRATION_MODE_OPTIONS: Tuple[Tuple[str, str], ...] = (
    ('Known standards only: concentration and standard ratio/RT', "training_original"),
)
CALIBRATION_MODE_LABELS = [x[0] for x in CALIBRATION_MODE_OPTIONS]
CALIBRATION_MODE_MAP = {x[0]: x[1] for x in CALIBRATION_MODE_OPTIONS}

MODEL_OBJECTIVE_OPTIONS: Tuple[Tuple[str, str], ...] = (
    ('Response correction and ranking', "trend"),
    ('Concentration ranges (fold-error priority)', "range"),
    ('Combined: ranking with fold-error constraints', "hybrid"),
)
MODEL_OBJECTIVE_LABELS = [x[0] for x in MODEL_OBJECTIVE_OPTIONS]
MODEL_OBJECTIVE_MAP = {x[0]: x[1] for x in MODEL_OBJECTIVE_OPTIONS}

OUTLIER_MODE_OPTIONS: Tuple[Tuple[str, str], ...] = (
    ('Off (audit only)', "off"),
    ('Conservative: robust MAD and directional consensus', "conservative"),
    ('Fold-error threshold and directional consensus', "threshold_consensus"),
    ('Ranked CV errors with an exclusion cap (exploratory)', "top_fraction"),
)
OUTLIER_MODE_LABELS = [x[0] for x in OUTLIER_MODE_OPTIONS]
OUTLIER_MODE_MAP = {x[0]: x[1] for x in OUTLIER_MODE_OPTIONS}


@dataclass
class ModelBenchmarkResult:
    enabled: bool
    calibration_mode_requested: str
    calibration_mode_used: str
    n_calibration: int
    n_target_valid: int
    best_model: str
    best_rating: str
    comparison: List[Dict[str, object]] = field(default_factory=list)
    cv_predictions: List[Dict[str, object]] = field(default_factory=list)
    calibration_records: List[Dict[str, object]] = field(default_factory=list)
    target_predictions: List[Dict[str, object]] = field(default_factory=list)
    calibration_checks: List[Dict[str, object]] = field(default_factory=list)
    target_checks: List[Dict[str, object]] = field(default_factory=list)
    feature_manifest: List[Dict[str, object]] = field(default_factory=list)
    feature_importance: List[Dict[str, object]] = field(default_factory=list)
    selected_descriptors: List[Dict[str, object]] = field(default_factory=list)
    unavailable_descriptors: List[Dict[str, object]] = field(default_factory=list)
    descriptor_selection_rows: List[Dict[str, object]] = field(default_factory=list)
    descriptor_selection_curve: List[Dict[str, object]] = field(default_factory=list)
    outlier_audit: List[Dict[str, object]] = field(default_factory=list)
    comparison_before_qc: List[Dict[str, object]] = field(default_factory=list)
    outlier_summary: List[Dict[str, object]] = field(default_factory=list)
    model_input_audit: List[Dict[str, object]] = field(default_factory=list)
    descriptor_guide: List[Dict[str, object]] = field(default_factory=list)
    tuning_rows: List[Dict[str, object]] = field(default_factory=list)
    leave_component_rows: List[Dict[str, object]] = field(default_factory=list)
    leave_component_summary: List[Dict[str, object]] = field(default_factory=list)
    plot_paths: List[Path] = field(default_factory=list)
    accuracy_band_rows: List[Dict[str, object]] = field(default_factory=list)
    decision_rows: List[Dict[str, object]] = field(default_factory=list)
    best_metrics: Dict[str, object] = field(default_factory=dict)
    baseline_metrics: Dict[str, object] = field(default_factory=dict)
    decision_label: str = ""
    decision_summary: str = ""
    n_calibration_before_qc: int = 0
    n_outliers_removed: int = 0
    n_selected_descriptors: int = 0
    n_unavailable_descriptors: int = 0
    model_objective_requested: str = "trend"
    model_objective_used: str = "trend"
    trend_comparison: List[Dict[str, object]] = field(default_factory=list)
    raw_vs_corrected_rows: List[Dict[str, object]] = field(default_factory=list)
    injection_group_rows: List[Dict[str, object]] = field(default_factory=list)
    injection_group_summary: List[Dict[str, object]] = field(default_factory=list)
    best_trend_metrics: Dict[str, object] = field(default_factory=dict)
    raw_trend_metrics: Dict[str, object] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)


if SKLEARN_AVAILABLE:
    class SafeSelectKBest(BaseEstimator, TransformerMixin):
        """Select at most ``k`` features without failing on small folds."""

        def __init__(self, k: int = 40):
            self.k = k

        def fit(self, X, y):
            n_features = int(X.shape[1])
            kk = max(1, min(int(self.k), n_features))
            self.selector_ = SelectKBest(score_func=f_regression, k=kk)
            with pywarnings.catch_warnings():
                pywarnings.simplefilter("ignore")
                self.selector_.fit(X, y)
            return self

        def transform(self, X):
            return self.selector_.transform(X)

        def get_support(self, indices: bool = False):
            return self.selector_.get_support(indices=indices)


# ---------------------------------------------------------------------------
# Basic utilities
# ---------------------------------------------------------------------------

def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _num(value: object) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def _median(values: Iterable[float]) -> Optional[float]:
    vals = [float(x) for x in values if x is not None and math.isfinite(float(x))]
    return float(statistics.median(vals)) if vals else None


def _key(record: Dict[str, object]) -> str:
    return _text(record.get("Product_Master_Formula_Key") or record.get("ABC_Formula_Key"))


def _source_label(record: Dict[str, object]) -> str:
    return f"{_text(record.get('Source_file'))}|{_text(record.get('Source_sheet'))}|row {_text(record.get('Source_row'))}"


def _valid_ratio(record: Dict[str, object]) -> Optional[float]:
    x = _num(record.get("Measured_ratio"))
    return x if x is not None and x > 0 else None


def _valid_concentration(record: Dict[str, object]) -> Optional[float]:
    x = _num(record.get("Actual_concentration"))
    return x if x is not None and x > 0 else None


def _model_rating(metrics: Dict[str, object]) -> str:
    r2 = _num(metrics.get("R2_log_concentration"))
    med_fold = _num(metrics.get("Median_Fold_Error"))
    within2 = _num(metrics.get("Within_2x_pct"))
    within5 = _num(metrics.get("Within_5x_pct"))
    if r2 is not None and med_fold is not None and within2 is not None:
        if r2 >= 0.50 and med_fold <= 1.50 and within2 >= 80:
            return "Good / potentially quantitative"
        if r2 >= 0.20 and med_fold <= 2.00 and within2 >= 60:
            return "Semi-quantitative / useful"
    if med_fold is not None and within5 is not None and med_fold <= 3.50 and within5 >= 80:
        return "Screening / range estimation"
    return "Poor / screening only"


def _average_ranks(values: np.ndarray) -> np.ndarray:
    """Return average ranks (1-based) with deterministic tie handling."""
    values = np.asarray(values, dtype=float)
    order = np.argsort(values, kind="mergesort")
    ranks = np.empty(len(values), dtype=float)
    i = 0
    while i < len(order):
        j = i + 1
        while j < len(order) and values[order[j]] == values[order[i]]:
            j += 1
        avg_rank = (i + 1 + j) / 2.0
        ranks[order[i:j]] = avg_rank
        i = j
    return ranks


def _safe_corr(x: np.ndarray, y: np.ndarray) -> float:
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    if len(x) < 2 or float(np.std(x)) <= 0 or float(np.std(y)) <= 0:
        return float("nan")
    return float(np.corrcoef(x, y)[0, 1])


def _metrics(actual: Sequence[float], predicted: Sequence[float]) -> Dict[str, object]:
    pairs = [
        (float(a), float(p)) for a, p in zip(actual, predicted)
        if a is not None and p is not None
        and math.isfinite(float(a)) and math.isfinite(float(p))
        and float(a) > 0 and float(p) > 0
    ]
    if not pairs:
        return {"N": 0}

    a = np.asarray([x[0] for x in pairs], dtype=float)
    p = np.asarray([x[1] for x in pairs], dtype=float)
    la = np.log10(a)
    lp = np.log10(p)
    log_error = lp - la
    abs_log_error = np.abs(log_error)
    ape = np.abs((p - a) / a) * 100.0
    fold = np.maximum(p / a, a / p)

    denom = float(np.sum((la - np.mean(la)) ** 2))
    r2 = float(1.0 - np.sum((lp - la) ** 2) / denom) if denom > 0 else float("nan")
    pearson = _safe_corr(la, lp)
    spearman = _safe_corr(_average_ranks(la), _average_ranks(lp))

    if len(a) >= 2 and float(np.std(la)) > 0:
        slope, intercept = np.polyfit(la, lp, 1)
        slope = float(slope)
        intercept = float(intercept)
    else:
        slope = float("nan")
        intercept = float("nan")

    result = {
        "N": int(len(a)),
        "MAE_log10": float(np.mean(abs_log_error)),
        "Median_AE_log10": float(np.median(abs_log_error)),
        "RMSE_log10": float(np.sqrt(np.mean(log_error ** 2))),
        "R2_log_concentration": r2,
        "Pearson_r_log": pearson,
        "Spearman_r_log": spearman,
        "Calibration_slope_log": slope,
        "Calibration_intercept_log": intercept,
        "Mean_log10_bias": float(np.mean(log_error)),
        "Median_log10_bias": float(np.median(log_error)),
        "Geometric_bias_pred_over_actual": float(10.0 ** np.mean(log_error)),
        "MAPE_pct": float(np.mean(ape)),
        "Median_APE_pct": float(np.median(ape)),
        "P90_APE_pct": float(np.percentile(ape, 90)),
        "Median_Fold_Error": float(np.median(fold)),
        "Mean_Fold_Error": float(np.mean(fold)),
        "P80_Fold_Error": float(np.percentile(fold, 80)),
        "P90_Fold_Error": float(np.percentile(fold, 90)),
        "Max_Fold_Error": float(np.max(fold)),
        "Within_1.25x_pct": float(np.mean(fold <= 1.25) * 100.0),
        "Within_1.5x_pct": float(np.mean(fold <= 1.5) * 100.0),
        "Within_2x_pct": float(np.mean(fold <= 2.0) * 100.0),
        "Within_3x_pct": float(np.mean(fold <= 3.0) * 100.0),
        "Within_5x_pct": float(np.mean(fold <= 5.0) * 100.0),
        "Over_5x_pct": float(np.mean(fold > 5.0) * 100.0),
    }
    result["Rating"] = _model_rating(result)
    return result



def _kendall_tau_b(x: Sequence[float], y: Sequence[float]) -> float:
    """Small-sample Kendall tau-b without an optional scipy dependency."""
    xx = np.asarray(x, dtype=float)
    yy = np.asarray(y, dtype=float)
    if len(xx) < 2 or len(xx) != len(yy):
        return float("nan")
    concordant = discordant = tie_x = tie_y = 0
    for i in range(len(xx) - 1):
        for j in range(i + 1, len(xx)):
            dx = float(xx[i] - xx[j])
            dy = float(yy[i] - yy[j])
            if dx == 0 and dy == 0:
                continue
            if dx == 0:
                tie_x += 1
            elif dy == 0:
                tie_y += 1
            elif dx * dy > 0:
                concordant += 1
            else:
                discordant += 1
    denom = math.sqrt((concordant + discordant + tie_x) * (concordant + discordant + tie_y))
    return float((concordant - discordant) / denom) if denom > 0 else float("nan")


def _pairwise_concordance_pct(actual: Sequence[float], score: Sequence[float]) -> float:
    """Percentage of non-tied actual pairs ordered correctly by a score.

    A predicted tie receives half credit.  This metric directly answers whether
    the response-corrected score gets the high/low ordering right.
    """
    a = np.asarray(actual, dtype=float)
    s = np.asarray(score, dtype=float)
    total = 0
    credit = 0.0
    for i in range(len(a) - 1):
        for j in range(i + 1, len(a)):
            da = float(a[i] - a[j])
            if da == 0:
                continue
            ds = float(s[i] - s[j])
            total += 1
            if ds == 0:
                credit += 0.5
            elif da * ds > 0:
                credit += 1.0
    return float(credit / total * 100.0) if total else float("nan")


def _top_fraction_overlap_pct(actual: np.ndarray, score: np.ndarray, fraction: float) -> float:
    n = len(actual)
    if n == 0:
        return float("nan")
    k = max(1, min(n, int(math.ceil(n * float(fraction)))))
    actual_top = set(np.argsort(actual, kind="mergesort")[-k:].tolist())
    score_top = set(np.argsort(score, kind="mergesort")[-k:].tolist())
    return float(len(actual_top.intersection(score_top)) / k * 100.0)


def _tertile_rank_accuracy_pct(actual: np.ndarray, score: np.ndarray) -> float:
    """Rank-tertile agreement; robust to arbitrary concentration units."""
    n = len(actual)
    if n < 3:
        return float("nan")

    def labels(values: np.ndarray) -> np.ndarray:
        order = np.argsort(values, kind="mergesort")
        out = np.zeros(n, dtype=int)
        for rank, idx in enumerate(order):
            out[idx] = min(2, int(rank * 3 / n))
        return out

    return float(np.mean(labels(actual) == labels(score)) * 100.0)


def _trend_metrics(actual: Sequence[float], score: Sequence[float]) -> Dict[str, object]:
    pairs = [
        (float(a), float(s)) for a, s in zip(actual, score)
        if a is not None and s is not None
        and math.isfinite(float(a)) and math.isfinite(float(s))
        and float(a) > 0 and float(s) > 0
    ]
    if not pairs:
        return {"Trend_N": 0}
    a = np.asarray([x[0] for x in pairs], dtype=float)
    s = np.asarray([x[1] for x in pairs], dtype=float)
    la = np.log10(a)
    ls = np.log10(s)
    spearman = _safe_corr(_average_ranks(la), _average_ranks(ls))
    kendall = _kendall_tau_b(la, ls)
    pairwise = _pairwise_concordance_pct(la, ls)
    top10 = _top_fraction_overlap_pct(la, ls, 0.10)
    top20 = _top_fraction_overlap_pct(la, ls, 0.20)
    tertile = _tertile_rank_accuracy_pct(la, ls)
    pearson = _safe_corr(la, ls)
    # Composite score is only used to order candidate models in trend mode.
    # The individual metrics remain the reportable evidence.
    components = [
        0.40 * (spearman if math.isfinite(spearman) else -1.0),
        0.20 * (kendall if math.isfinite(kendall) else -1.0),
        0.20 * ((pairwise / 50.0 - 1.0) if math.isfinite(pairwise) else -1.0),
        0.10 * ((top20 / 50.0 - 1.0) if math.isfinite(top20) else -1.0),
        0.10 * ((tertile / (100.0 / 3.0) - 1.0) if math.isfinite(tertile) else -1.0),
    ]
    return {
        "Trend_N": int(len(a)),
        "Trend_Pearson_r_log": pearson,
        "Trend_Spearman_r": spearman,
        "Trend_Kendall_tau": kendall,
        "Trend_Pairwise_Concordance_pct": pairwise,
        "Trend_Top10_Overlap_pct": top10,
        "Trend_Top20_Overlap_pct": top20,
        "Trend_Tertile_Accuracy_pct": tertile,
        "Trend_Composite_Score": float(sum(components)),
    }


def _augment_comparison_with_trend(
    comparison: List[Dict[str, object]],
    cv_aggregated: Sequence[Dict[str, object]],
) -> Tuple[List[Dict[str, object]], Dict[str, object]]:
    """Attach response-correction trend metrics to every CV model row."""
    raw_actual: List[float] = []
    raw_ratio: List[float] = []
    # Use one row per calibration sample; all models share the same measured ratio.
    seen = set()
    for row in cv_aggregated:
        idx_value = _num(row.get("Calibration_Index"))
        idx = int(idx_value) if idx_value is not None else -1
        if idx in seen:
            continue
        actual = _num(row.get("Actual_concentration"))
        ratio = _num(row.get("Measured_ratio"))
        if actual is not None and ratio is not None and actual > 0 and ratio > 0:
            seen.add(idx)
            raw_actual.append(float(actual))
            raw_ratio.append(float(ratio))
    raw = _trend_metrics(raw_actual, raw_ratio)

    for model_row in comparison:
        name = str(model_row.get("Model", ""))
        subset = [r for r in cv_aggregated if str(r.get("Model", "")) == name]
        actual = [float(r["Actual_concentration"]) for r in subset]
        predicted = [float(r["Predicted_concentration"]) for r in subset]
        trend = _trend_metrics(actual, predicted)
        model_row.update(trend)
        model_row.update({f"Raw_{k}": v for k, v in raw.items()})
        for metric in (
            "Trend_Spearman_r", "Trend_Kendall_tau",
            "Trend_Pairwise_Concordance_pct", "Trend_Top10_Overlap_pct",
            "Trend_Top20_Overlap_pct", "Trend_Tertile_Accuracy_pct",
            "Trend_Composite_Score",
        ):
            value = _num(trend.get(metric))
            raw_key = "Raw_" + metric
            raw_value = _num(model_row.get(raw_key))
            model_row["Delta_" + metric + "_vs_raw"] = (
                float(value - raw_value) if value is not None and raw_value is not None else ""
            )
    return comparison, raw


def _sort_comparison_for_objective(
    comparison: List[Dict[str, object]], objective: str,
) -> List[Dict[str, object]]:
    objective = str(objective or "trend").lower()
    if objective == "range":
        comparison.sort(key=lambda r: (
            float("inf") if _num(r.get("Median_Fold_Error")) is None else float(r["Median_Fold_Error"]),
            -(_num(r.get("Within_2x_pct")) or 0.0),
            float("inf") if _num(r.get("RMSE_log10")) is None else float(r["RMSE_log10"]),
        ))
    elif objective == "hybrid":
        comparison.sort(key=lambda r: (
            -(_num(r.get("Trend_Composite_Score")) or -999.0),
            float("inf") if _num(r.get("Median_Fold_Error")) is None else float(r["Median_Fold_Error"]),
            -(_num(r.get("Within_5x_pct")) or 0.0),
        ))
    else:
        comparison.sort(key=lambda r: (
            -(_num(r.get("Trend_Spearman_r")) or -999.0),
            -(_num(r.get("Trend_Kendall_tau")) or -999.0),
            -(_num(r.get("Trend_Pairwise_Concordance_pct")) or -999.0),
            -(_num(r.get("Trend_Top20_Overlap_pct")) or -999.0),
            float("inf") if _num(r.get("Median_Fold_Error")) is None else float(r["Median_Fold_Error"]),
        ))
    for rank, row in enumerate(comparison, start=1):
        row["Rank"] = rank
        row["Selection_Objective"] = objective
    return comparison


# ---------------------------------------------------------------------------
# Calibration construction
# ---------------------------------------------------------------------------

def _group_by_key(records: Sequence[Dict[str, object]]) -> Dict[str, List[Dict[str, object]]]:
    out: Dict[str, List[Dict[str, object]]] = {}
    for record in records:
        key = _key(record)
        if key:
            out.setdefault(key, []).append(record)
    return out


def _representative_by_ratio(rows: Sequence[Dict[str, object]]) -> Optional[Dict[str, object]]:
    valid = [(r, _valid_ratio(r)) for r in rows]
    valid = [(r, x) for r, x in valid if x is not None]
    if not valid:
        return None
    med = float(statistics.median(x for _r, x in valid))
    return min(valid, key=lambda pair: abs(pair[1] - med))[0]


def build_calibration_records(
    training_records: Sequence[Dict[str, object]],
    target_records: Sequence[Dict[str, object]],
    *,
    requested_mode: str = "training_original",
) -> Tuple[List[Dict[str, object]], List[Dict[str, object]], str, List[str]]:
    """Build supervised calibration rows from the known-standard injections only.

    ``target_records`` are inspected only to report structural/formula overlap with the
    unknown real sample.  Their ratios, retention times and other measured values are
    never copied into the supervised calibration table.
    """
    warnings: List[str] = []
    checks: List[Dict[str, object]] = []
    from .training_limits import validate_training_count
    validate_training_count(training_records)
    train_groups = _group_by_key(training_records)
    target_groups = _group_by_key(target_records)

    valid_train_keys = []
    for key, rows in train_groups.items():
        if any(_valid_concentration(r) is not None for r in rows):
            valid_train_keys.append(key)
    target_overlap = sum(
        1 for key in valid_train_keys
        if any(_valid_ratio(r) is not None for r in target_groups.get(key, []))
    )
    valid_train_n = len(valid_train_keys)

    requested = str(requested_mode or "training_original").strip().lower()
    mode = "training_original"
    if requested not in {"training_original", "standard_only", "original"}:
        warnings.append(
            f"Calibration mode {requested!r} was overridden: supervised calibration is fixed to the known standard injections. "
            "The unknown 1000-compound sample is used only for final prediction and overlap diagnostics."
        )
    warnings.append(
        f"Standard-only supervision: {valid_train_n} known-standard formula keys; "
        f"unknown-sample overlap={target_overlap}/{valid_train_n} is diagnostic only and contributes no labels or measured ratios."
    )

    calibration: List[Dict[str, object]] = []
    for record in training_records:
        ratio = _valid_ratio(record)
        conc = _valid_concentration(record)
        reason = ""
        if ratio is None:
            reason = "invalid or missing measured ratio in known-standard injection"
        elif conc is None:
            reason = "invalid or missing actual concentration in known-standard injection"
        if reason:
            checks.append({
                "ABC_Formula_Key": _key(record),
                "Status": "SKIPPED",
                "Reason": reason,
                "Training_rows": _source_label(record),
                "Target_rows": "",
            })
            continue
        row = dict(record)
        row["Measured_ratio"] = float(ratio)
        row["Actual_concentration"] = float(conc)
        row["Calibration_source"] = "known_standard_injection"
        row["Calibration_training_sources"] = _source_label(record)
        row["Calibration_target_sources"] = ""
        row["Calibration_training_count"] = 1
        row["Calibration_target_count"] = 0
        calibration.append(row)
        checks.append({
            "ABC_Formula_Key": _key(record),
            "Status": "USED",
            "Reason": "known concentration, ratio, RT and descriptors all from the same standard injection",
            "Actual_concentration": float(conc),
            "Measured_ratio": float(ratio),
            "Training_rows": _source_label(record),
            "Target_rows": "",
        })

    return calibration, checks, mode, warnings

def _prepare_target_records(records: Sequence[Dict[str, object]]) -> Tuple[List[Dict[str, object]], List[Dict[str, object]]]:
    valid: List[Dict[str, object]] = []
    checks: List[Dict[str, object]] = []
    for record in records:
        ratio = _valid_ratio(record)
        key = _key(record)
        reason = ""
        if not key:
            reason = "missing ordered A/B/C Formula key (often internal-standard or summary row)"
        elif ratio is None:
            reason = "invalid or missing measured ratio"
        elif not _text(record.get("Structure_status")):
            reason = "missing structure features"
        if reason:
            checks.append({
                "Source_file": record.get("Source_file", ""),
                "Source_sheet": record.get("Source_sheet", ""),
                "Source_row": record.get("Source_row", ""),
                "Combo": record.get("Combo", ""),
                "ABC_Formula_Key": key,
                "Status": "SKIPPED",
                "Reason": reason,
            })
        else:
            valid.append(record)
            checks.append({
                "Source_file": record.get("Source_file", ""),
                "Source_sheet": record.get("Source_sheet", ""),
                "Source_row": record.get("Source_row", ""),
                "Combo": record.get("Combo", ""),
                "ABC_Formula_Key": key,
                "Status": "PREDICTABLE",
                "Reason": "valid ratio and structure/gradient features",
            })
    return valid, checks


# ---------------------------------------------------------------------------
# Feature matrix and model definitions
# ---------------------------------------------------------------------------

def _finite_fraction(records: Sequence[Dict[str, object]], name: str) -> Tuple[int, float]:
    vals = [_num(r.get(name)) for r in records]
    good = [x for x in vals if x is not None]
    if not good:
        return 0, 0.0
    return len(good), float(np.var(np.asarray(good, dtype=float)))


_THREE_D_FEATURES = {
    "MolVolume3D", "RadiusOfGyration", "Asphericity", "Eccentricity", "SpherocityIndex",
    "PMI1", "PMI2", "PMI3", "NPR1", "NPR2", "InertialShapeFactor", "PBF",
}


def _feature_group(name: str) -> str:
    name = str(name or "")
    if name in {"A_Formula", "B_Formula", "C_Formula"}:
        return "Component identity"
    if name in {
        "Apex_RT_min", "Effective_gradient_time_min", "Mobile_phase_B_pct",
        "B_slope_pct_per_min", "Gradient_state",
    }:
        return "LC gradient / RT"
    if name.startswith("Published_"):
        return "Published IE eluent descriptors"
    if name.startswith(("Formula_", "FormulaFrac_")) or name in {"Exact_mass", "DBE"}:
        return "Formula / exact mass"
    if name in _THREE_D_FEATURES:
        return "3D size / shape"
    if any(token in name for token in (
        "Acid", "Basic", "Amine", "Amidine", "Guanidine", "Imine", "Ammonium",
        "Gasteiger", "FormalCharge", "Ionizable", "Phenol", "Thiol", "Imide",
        "Carboxylic", "Sulfonic", "Phosphoric", "HBD", "HBA", "NHOH", "NOCount",
    )):
        return "Acid/base / charge"
    if name.endswith("_proxy") and name in {
        "HydrophobicGradientExposure_proxy", "PolarAqueousExposure_proxy",
        "IonizableAqueousExposure_proxy", "BasicAqueousExposure_proxy",
        "AcidicAqueousExposure_proxy", "ChargeDensitySurface_proxy",
        "PolarMassDensity_proxy", "HydrophobicSurface_proxy",
        "AromaticHydrophobicity_proxy", "IonizableSiteDensity_proxy",
        "Mass_per_IonizableSite_proxy", "GradientChangeExposure_proxy",
        "LogP_TPSA_balance_proxy",
    } or name in {"DBE_per_C", "Hetero_to_C_ratio"}:
        return "Engineered physical proxy"
    if name.startswith(("PEOE_VSA", "SlogP_VSA", "SMR_VSA", "EState_VSA", "BCUT2D_")):
        return "Extended RDKit"
    if any(token in name for token in (
        "LogP", "TPSA", "Surface", "ASA", "MolMR", "FractionCSP3", "Aromatic",
        "HeteroAtomFraction", "CarbonAtomFraction", "RingCount", "Rotatable",
    )):
        return "Hydrophobicity / polarity"
    return "Size / topology"


def _is_auto_generated_feature(name: str) -> bool:
    return name not in {
        "A_Formula", "B_Formula", "C_Formula", "Apex_RT_min",
        "Effective_gradient_time_min", "Mobile_phase_B_pct", "B_slope_pct_per_min",
        "Gradient_state",
    }


def _manifest_row(
    name: str,
    feature_type: str,
    n: int,
    variance: object,
    total: int,
    status: str,
    reason: str,
) -> Dict[str, object]:
    missing_pct = 100.0 if total <= 0 else max(0.0, (total - int(n)) / total * 100.0)
    return {
        "Feature": name,
        "Group": _feature_group(name),
        "Type": feature_type,
        "Auto_generated": "Yes" if _is_auto_generated_feature(name) else "Input/derived",
        "Valid_N": int(n),
        "Calibration_N": int(total),
        "Missing_pct": float(missing_pct),
        "Variance": variance,
        "Initial_Status": status,
        "Final_Status": status,
        "Selected_Frequency_pct": 0.0,
        "Selected_in_final_model": "No",
        "Redundant_with": "",
        "Permutation_Importance_Mean": "",
        "Reason": reason,
    }


def choose_features(
    calibration: Sequence[Dict[str, object]],
    descriptor_names: Sequence[str],
    *,
    expected_descriptor_names: Sequence[str] = (),
    three_d_enabled: bool = False,
    correlation_threshold: float = 0.97,
    include_categorical: bool = False,
) -> Tuple[List[str], List[str], List[Dict[str, object]]]:
    base_numeric = [
        "Exact_mass", "DBE", "Apex_RT_min", "Effective_gradient_time_min",
        "Mobile_phase_B_pct", "B_slope_pct_per_min",
    ]
    # ``expected_descriptor_names`` deliberately includes descriptors that may
    # be unavailable (for example 3D descriptors when 3D generation is off),
    # so the output can explain what was not generated instead of silently
    # omitting the column.
    candidates = base_numeric + list(expected_descriptor_names) + list(descriptor_names)
    numeric: List[str] = []
    manifest: List[Dict[str, object]] = []
    minimum = max(5, int(math.ceil(len(calibration) * 0.50)))
    seen = set()
    for name in candidates:
        if name in seen:
            continue
        seen.add(name)
        n, var = _finite_fraction(calibration, name)
        if n == 0:
            if name in _THREE_D_FEATURES and not three_d_enabled:
                status = "Disabled"
                reason = "3D generation is disabled by the current setting"
            elif name in _THREE_D_FEATURES:
                status = "Generation failed"
                reason = "3D descriptor could not be generated for any calibration row"
            else:
                status = "Unavailable"
                reason = "descriptor was not generated or contained no numeric values"
            include = False
        elif n < minimum:
            status = "Too sparse"
            reason = f"valid values {n} < required {minimum}"
            include = False
        elif var <= 1e-14:
            status = "Constant"
            if name.startswith(("Formula_", "FormulaFrac_")):
                reason = "element is absent or constant in this chemical library; not informative for this dataset"
            else:
                reason = "zero or near-zero variance in the calibration set"
            include = False
        else:
            status = "Candidate"
            reason = "available for automatic selection"
            include = True
        manifest.append(_manifest_row(name, "numeric", n, var, len(calibration), status, reason))
        if include:
            numeric.append(name)
    categorical = ["A_Formula", "B_Formula", "C_Formula", "Gradient_state"]
    for name in categorical:
        nonblank = sum(bool(_text(r.get(name))) for r in calibration)
        available = nonblank >= max(5, int(len(calibration) * 0.50))
        include = bool(include_categorical) and available
        if include:
            status, reason = "Candidate", "optional categorical feature; one-hot encoded to numeric columns"
        elif available:
            status, reason = "Disabled", "categorical one-hot features disabled; numeric molecular descriptors are used"
        else:
            status, reason = "Too sparse", "too many blank values"
        manifest.append(_manifest_row(name, "categorical", nonblank, "", len(calibration), status, reason))
    categorical = [x for x in categorical if any(m["Feature"] == x and m["Initial_Status"] == "Candidate" for m in manifest)]

    # Remove near-duplicate numeric descriptors without consulting the target
    # variable.  This prevents dozens of almost identical VSA/size columns from
    # overwhelming a 90-100 row calibration set while avoiding supervised data
    # leakage before cross-validation.
    if len(numeric) >= 2 and 0.0 < float(correlation_threshold) < 1.0:
        columns: Dict[str, np.ndarray] = {}
        for name in numeric:
            values = np.asarray([
                np.nan if _num(r.get(name)) is None else float(_num(r.get(name)))
                for r in calibration
            ], dtype=float)
            finite = values[np.isfinite(values)]
            fill = float(np.median(finite)) if finite.size else 0.0
            values = np.where(np.isfinite(values), values, fill)
            columns[name] = values
        kept: List[str] = []
        dropped: Dict[str, str] = {}
        for name in numeric:
            redundant = ""
            for prior in kept:
                x, z = columns[name], columns[prior]
                if float(np.std(x)) <= 0 or float(np.std(z)) <= 0:
                    continue
                corr = abs(float(np.corrcoef(x, z)[0, 1]))
                if math.isfinite(corr) and corr >= float(correlation_threshold):
                    redundant = prior
                    break
            if redundant:
                dropped[name] = redundant
            else:
                kept.append(name)
        numeric = kept
        for row in manifest:
            feature = str(row.get("Feature", ""))
            if feature in dropped:
                row["Final_Status"] = "Redundant"
                row["Redundant_with"] = dropped[feature]
                row["Reason"] = f"absolute correlation >= {float(correlation_threshold):.2f} with {dropped[feature]}"
    return numeric, categorical, manifest


def _matrix(records: Sequence[Dict[str, object]], numeric: Sequence[str], categorical: Sequence[str]) -> np.ndarray:
    rows: List[List[object]] = []
    for record in records:
        row: List[object] = []
        for name in numeric:
            value = _num(record.get(name))
            row.append(np.nan if value is None else float(value))
        for name in categorical:
            row.append(_text(record.get(name)) or "<MISSING>")
        rows.append(row)
    return np.asarray(rows, dtype=object)


def _categories(calibration: Sequence[Dict[str, object]], categorical: Sequence[str]) -> List[List[str]]:
    out: List[List[str]] = []
    for name in categorical:
        vals = sorted({_text(r.get(name)) or "<MISSING>" for r in calibration})
        out.append(vals or ["<MISSING>"])
    return out


def _preprocessor(n_num: int, n_cat: int, categories: Sequence[Sequence[str]], *, scale: bool):
    num_steps = [("impute", SimpleImputer(strategy="median", keep_empty_features=True))]
    if scale:
        num_steps.append(("scale", StandardScaler()))
    num_pipe = Pipeline(num_steps)
    transformers = []
    if n_num:
        transformers.append(("num", num_pipe, list(range(n_num))))
    if n_cat:
        cat_pipe = Pipeline([
            ("impute", SimpleImputer(strategy="most_frequent")),
            ("onehot", OneHotEncoder(categories=list(categories), handle_unknown="ignore", sparse_output=False)),
        ])
        transformers.append(("cat", cat_pipe, list(range(n_num, n_num + n_cat))))
    return ColumnTransformer(transformers=transformers, remainder="drop", sparse_threshold=0.0)


def _pipeline(model, n_num: int, n_cat: int, categories, max_features: int, *, scale: bool):
    return Pipeline([
        ("preprocess", _preprocessor(n_num, n_cat, categories, scale=scale)),
        ("select", SafeSelectKBest(k=max_features)),
        ("model", model),
    ])


def build_models(
    n_samples: int,
    n_num: int,
    n_cat: int,
    categories: Sequence[Sequence[str]],
    *,
    max_features: int,
    random_state: int,
):
    if not SKLEARN_AVAILABLE:
        raise RuntimeError("scikit-learn is required for multi-model benchmarking: " + SKLEARN_ERROR)
    neighbors = max(2, min(7, n_samples - 1))
    kernel = ConstantKernel(1.0, (1e-2, 1e2)) * RBF(length_scale=1.0, length_scale_bounds=(1e-2, 1e2)) + WhiteKernel(noise_level=0.1, noise_level_bounds=(1e-5, 1e1))
    return {
        "BayesianRidge": _pipeline(BayesianRidge(), n_num, n_cat, categories, max_features, scale=True),
        "SVR_RBF": _pipeline(SVR(C=10.0, epsilon=0.10, gamma="scale"), n_num, n_cat, categories, max_features, scale=True),
        "RandomForest": _pipeline(RandomForestRegressor(
            n_estimators=300, max_depth=6, min_samples_leaf=3, max_features=0.5,
            bootstrap=True, random_state=random_state, n_jobs=-1,
        ), n_num, n_cat, categories, max_features, scale=False),
        "ExtraTrees": _pipeline(ExtraTreesRegressor(
            n_estimators=300, max_depth=7, min_samples_leaf=2, max_features=0.7,
            random_state=random_state, n_jobs=-1,
        ), n_num, n_cat, categories, max_features, scale=False),
        "HistGradientBoosting": _pipeline(HistGradientBoostingRegressor(
            max_iter=250, learning_rate=0.05, max_leaf_nodes=7, min_samples_leaf=10,
            l2_regularization=1.0, random_state=random_state,
        ), n_num, n_cat, categories, max_features, scale=True),
        "GaussianProcess": _pipeline(GaussianProcessRegressor(
            kernel=kernel, alpha=1e-6, normalize_y=True, n_restarts_optimizer=0, random_state=random_state,
        ), n_num, n_cat, categories, max_features, scale=True),
        "KNN_Local": _pipeline(KNeighborsRegressor(n_neighbors=neighbors, weights="distance", p=2), n_num, n_cat, categories, max_features, scale=True),
        "PLS": _pipeline(PLSRegression(n_components=3, scale=False, max_iter=500), n_num, n_cat, categories, max_features, scale=True),
    }


def _predict_estimator(estimator, X) -> np.ndarray:
    pred = estimator.predict(X)
    return np.asarray(pred, dtype=float).reshape(-1)


def _model_param_grid(model_name: str, selected_k: int, n_train: int) -> Dict[str, Sequence[object]]:
    """Compact model-specific grids for optional nested tuning.

    The grids are intentionally small because the calibration set is only about
    90-100 rows.  Tuning is performed strictly inside the outer training fold.
    """
    if model_name == "SVR_RBF":
        return {
            "model__C": [1.0, 10.0, 100.0],
            "model__gamma": ["scale", 0.05],
            "model__epsilon": [0.05, 0.15],
        }
    if model_name == "BayesianRidge":
        return {
            "model__alpha_1": [1e-6, 1e-4],
            "model__lambda_1": [1e-6, 1e-4],
        }
    if model_name == "RandomForest":
        return {
            "model__max_depth": [3, 6, None],
            "model__min_samples_leaf": [2, 5],
            "model__max_features": [0.3, 0.7],
        }
    if model_name == "PLS":
        upper = max(1, min(int(selected_k), int(n_train) - 1, 6))
        values = sorted({x for x in (1, 2, 3, 4, 5, upper) if 1 <= x <= upper})
        return {"model__n_components": values}
    return {}


def _fit_with_optional_tuning(
    estimator,
    model_name: str,
    X_train: np.ndarray,
    y_train: np.ndarray,
    *,
    selected_k: int,
    deep_tuning: bool,
    random_state: int,
):
    """Fit a pipeline, optionally using compact inner-CV hyperparameter tuning."""
    estimator = clone(estimator)
    estimator.set_params(select__k=int(selected_k))
    grid = _model_param_grid(model_name, int(selected_k), len(y_train)) if deep_tuning else {}
    if not grid:
        estimator.fit(X_train, y_train)
        return estimator, {}, ""
    inner_splits = max(2, min(3, len(y_train) // 8 if len(y_train) >= 16 else 2))
    inner_cv = KFold(n_splits=inner_splits, shuffle=True, random_state=int(random_state))
    search = GridSearchCV(
        estimator, grid, scoring="neg_mean_absolute_error", cv=inner_cv,
        n_jobs=1, refit=True, error_score=np.nan, return_train_score=False,
    )
    search.fit(X_train, y_train)
    best = search.best_estimator_
    params = {str(k): v for k, v in (search.best_params_ or {}).items()}
    score = "" if not math.isfinite(float(search.best_score_)) else float(-search.best_score_)
    return best, params, score


# ---------------------------------------------------------------------------
# Cross-validation and final fit
# ---------------------------------------------------------------------------

def _aggregate_cv_rows(rows: Sequence[Dict[str, object]]) -> Tuple[List[Dict[str, object]], Dict[str, Dict[int, float]]]:
    by_model_sample: Dict[Tuple[str, int], List[Dict[str, object]]] = {}
    for row in rows:
        by_model_sample.setdefault((str(row["Model"]), int(row["Calibration_Index"])), []).append(row)
    aggregated: List[Dict[str, object]] = []
    pred_map: Dict[str, Dict[int, float]] = {}
    for (model, idx), group in sorted(by_model_sample.items()):
        preds = [float(x["Predicted_concentration"]) for x in group if _num(x.get("Predicted_concentration")) is not None]
        if not preds:
            continue
        pred = float(statistics.median(preds))
        actual = float(group[0]["Actual_concentration"])
        ratio = float(group[0]["Measured_ratio"])
        fold = max(pred / actual, actual / pred)
        row = dict(group[0])
        row.update({
            "Repeat": "median_over_repeats", "Fold": "", "Predicted_concentration": pred,
            "APE_pct": abs(pred - actual) / actual * 100.0, "Fold_error": fold,
            "Predicted_log10_RRF": float(math.log10(ratio / pred)),
            "Prediction_count": len(preds),
        })
        aggregated.append(row)
        pred_map.setdefault(model, {})[idx] = pred
    return aggregated, pred_map


def _raw_feature_name(transformed_name: str, numeric: Sequence[str], categorical: Sequence[str]) -> str:
    """Map a transformed ColumnTransformer name back to the raw descriptor.

    ColumnTransformer was built from integer column positions, so sklearn emits
    names such as ``num__x3`` and ``cat__x12_value``.  v18.29 treated ``x3`` as
    a descriptor name, which made all stability frequencies appear as zero.
    """
    text = str(transformed_name or "")
    match = re.match(r"^(?:num|cat)__x(\d+)(?:_|$)", text)
    if match:
        idx = int(match.group(1))
        if idx < len(numeric):
            return str(numeric[idx])
        cat_idx = idx - len(numeric)
        if 0 <= cat_idx < len(categorical):
            return str(categorical[cat_idx])
    if text.startswith("num__"):
        suffix = text.split("__", 1)[1]
        if suffix in numeric:
            return suffix
    if text.startswith("cat__"):
        suffix = text.split("__", 1)[1]
        for name in sorted(categorical, key=len, reverse=True):
            if suffix == name or suffix.startswith(name + "_"):
                return name
    if text in numeric or text in categorical:
        return text
    return text


def _selected_raw_features(
    estimator,
    numeric: Sequence[str],
    categorical: Sequence[str],
) -> Tuple[List[str], List[str]]:
    """Return raw and expanded features selected by a fitted pipeline."""
    try:
        pre = estimator.named_steps["preprocess"]
        selector = estimator.named_steps["select"]
        expanded = list(pre.get_feature_names_out())
        support = np.asarray(selector.get_support(), dtype=bool)
        chosen_expanded = [expanded[i] for i in range(min(len(expanded), len(support))) if support[i]]
        raw: List[str] = []
        for name in chosen_expanded:
            base = _raw_feature_name(name, numeric, categorical)
            if base and base not in raw:
                raw.append(base)
        return raw, chosen_expanded
    except Exception:
        return [], []


def _candidate_feature_counts(min_features: int, max_features: int) -> List[int]:
    lo = max(2, int(min_features))
    hi = max(lo, int(max_features))
    candidates = {lo, hi}
    for x in (5, 8, 12, 16, 24, 32, 48, 64):
        if lo <= x <= hi:
            candidates.add(x)
    return sorted(candidates)


def _choose_k_inner(
    X_train: np.ndarray,
    y_train: np.ndarray,
    models: Dict[str, object],
    *,
    min_features: int,
    max_features: int,
    random_state: int,
) -> Tuple[int, List[Dict[str, object]], List[str]]:
    """Choose descriptor count inside the outer training fold.

    A compact Bayesian-Ridge/SVR consensus is used so selection is not tied to
    one tree or one nonlinear learner.  The smallest feature count within 2%
    of the best inner-CV MAE is selected to favour parsimonious models.
    """
    warnings: List[str] = []
    candidates = _candidate_feature_counts(min_features, max_features)
    n = len(y_train)
    inner_splits = max(2, min(3, n // 6 if n >= 12 else 2))
    kfold = KFold(n_splits=inner_splits, shuffle=True, random_state=int(random_state))
    selector_models = [name for name in ("BayesianRidge", "SVR_RBF") if name in models]
    if not selector_models:
        selector_models = list(models.keys())[:1]
    curve: List[Dict[str, object]] = []
    score_by_k: Dict[int, float] = {}
    for k in candidates:
        fold_errors: List[float] = []
        failures = 0
        for inner_no, (tr, va) in enumerate(kfold.split(X_train, y_train), start=1):
            for model_name in selector_models:
                try:
                    est = clone(models[model_name])
                    est.set_params(select__k=int(k))
                    with pywarnings.catch_warnings():
                        pywarnings.simplefilter("ignore")
                        est.fit(X_train[tr], y_train[tr])
                        pred = _predict_estimator(est, X_train[va])
                    fold_errors.extend(np.abs(pred - y_train[va]).tolist())
                except Exception as exc:
                    failures += 1
                    if failures <= 2:
                        warnings.append(f"inner feature-count fit failed for k={k}, {model_name}: {exc}")
        score = float(np.mean(fold_errors)) if fold_errors else float("inf")
        score_by_k[int(k)] = score
        curve.append({
            "Feature_Count": int(k),
            "Inner_MAE_log10_RRF": score if math.isfinite(score) else "",
            "Selector_Models": "+".join(selector_models),
            "Inner_Fits": len(fold_errors),
            "Failed_Fits": failures,
        })
    finite_scores = [x for x in score_by_k.values() if math.isfinite(x)]
    if not finite_scores:
        return max(2, min(int(max_features), X_train.shape[1])), curve, warnings
    best = min(finite_scores)
    tolerance = max(0.005, best * 0.02)
    eligible = [k for k, score in score_by_k.items() if math.isfinite(score) and score <= best + tolerance]
    chosen = min(eligible) if eligible else min(score_by_k, key=score_by_k.get)
    return int(chosen), curve, warnings


def _run_cv(
    calibration: Sequence[Dict[str, object]],
    X: np.ndarray,
    y: np.ndarray,
    models: Dict[str, object],
    *,
    splits: int,
    repeats: int,
    random_state: int,
    numeric: Sequence[str] = (),
    categorical: Sequence[str] = (),
    auto_select_features: bool = False,
    min_features: int = 6,
    max_features: int = 40,
    deep_tuning: bool = False,
    model_objective: str = "trend",
) -> Tuple[
    List[Dict[str, object]],
    List[Dict[str, object]],
    List[str],
    List[Dict[str, object]],
    List[Dict[str, object]],
    List[Dict[str, object]],
]:
    warnings: List[str] = []
    n = len(calibration)
    n_splits = max(2, min(int(splits), n))
    cv = RepeatedKFold(n_splits=n_splits, n_repeats=max(1, int(repeats)), random_state=int(random_state))
    split_list = list(cv.split(X, y))
    rows: List[Dict[str, object]] = []
    selection_events: List[Dict[str, object]] = []
    selection_curve: List[Dict[str, object]] = []
    tuning_events: List[Dict[str, object]] = []
    model_names = ["GlobalMedian"] + list(models.keys())
    failures: Dict[str, int] = {name: 0 for name in model_names}

    for split_no, (train_idx, test_idx) in enumerate(split_list, start=1):
        repeat_no = (split_no - 1) // n_splits + 1
        fold_no = (split_no - 1) % n_splits + 1
        y_train = y[train_idx]
        selected_k = int(max_features)
        if auto_select_features:
            selected_k, curve_rows, inner_warnings = _choose_k_inner(
                X[train_idx], y_train, models,
                min_features=min_features,
                max_features=max_features,
                random_state=int(random_state) + split_no * 101,
            )
            warnings.extend(inner_warnings[:3])
            for curve_row in curve_rows:
                row = dict(curve_row)
                row.update({
                    "Outer_Repeat": repeat_no,
                    "Outer_Fold": fold_no,
                    "Chosen": "Yes" if int(curve_row["Feature_Count"]) == int(selected_k) else "No",
                })
                selection_curve.append(row)
        model_predictions: Dict[str, Optional[np.ndarray]] = {
            "GlobalMedian": np.full(len(test_idx), float(np.median(y_train)), dtype=float)
        }
        selection_captured = False
        for name, base_estimator in models.items():
            try:
                with pywarnings.catch_warnings():
                    pywarnings.simplefilter("ignore")
                    est, best_params, inner_mae = _fit_with_optional_tuning(
                        base_estimator, name, X[train_idx], y_train,
                        selected_k=int(selected_k),
                        deep_tuning=bool(deep_tuning),
                        random_state=int(random_state) + split_no * 1009 + len(tuning_events),
                    )
                    pred = _predict_estimator(est, X[test_idx])
                model_predictions[name] = pred
                tuning_events.append({
                    "Outer_Repeat": repeat_no,
                    "Outer_Fold": fold_no,
                    "Model": name,
                    "Deep_tuning": "Yes" if deep_tuning and bool(_model_param_grid(name, int(selected_k), len(train_idx))) else "No",
                    "Selected_transformed_feature_count": int(selected_k),
                    "Inner_MAE_log10_RRF": inner_mae,
                    "Best_params": "; ".join(f"{k}={v}" for k, v in sorted(best_params.items())),
                })
                if not selection_captured:
                    raw_selected, expanded_selected = _selected_raw_features(est, numeric, categorical)
                    for feature in list(numeric) + list(categorical):
                        selection_events.append({
                            "Outer_Repeat": repeat_no,
                            "Outer_Fold": fold_no,
                            "Chosen_Feature_Count": int(selected_k),
                            "Feature": feature,
                            "Selected": "Yes" if feature in raw_selected else "No",
                            "Selected_Expanded_Levels": sum(
                                1 for x in expanded_selected
                                if _raw_feature_name(x, numeric, categorical) == feature
                            ),
                        })
                    selection_captured = True
            except Exception as exc:
                failures[name] += 1
                model_predictions[name] = None
                if failures[name] <= 3:
                    warnings.append(f"{name} CV fold failed (repeat {repeat_no}, fold {fold_no}): {exc}")

        for name, preds in model_predictions.items():
            if preds is None:
                continue
            for local_pos, sample_idx in enumerate(test_idx):
                record = calibration[int(sample_idx)]
                ratio = float(record["Measured_ratio"])
                actual = float(record["Actual_concentration"])
                pred_log_rrf = float(preds[local_pos])
                try:
                    pred_conc = float(ratio / (10.0 ** pred_log_rrf))
                except Exception:
                    continue
                if not math.isfinite(pred_conc) or pred_conc <= 0:
                    continue
                rows.append({
                    "Model": name,
                    "Calibration_Index": int(sample_idx),
                    "Repeat": repeat_no,
                    "Fold": fold_no,
                    "ABC_Formula_Key": _key(record),
                    "Combo": record.get("Combo", ""),
                    "A_Formula": record.get("A_Formula", ""),
                    "B_Formula": record.get("B_Formula", ""),
                    "C_Formula": record.get("C_Formula", ""),
                    "Injection_Group": record.get("Injection_Group", ""),
                    "Measured_ratio": ratio,
                    "Actual_concentration": actual,
                    "Actual_log10_RRF": float(y[int(sample_idx)]),
                    "Predicted_log10_RRF": pred_log_rrf,
                    "Predicted_concentration": pred_conc,
                    "APE_pct": abs(pred_conc - actual) / actual * 100.0,
                    "Fold_error": max(pred_conc / actual, actual / pred_conc),
                    "Calibration_source": record.get("Calibration_source", ""),
                    "Auto_Selected_Feature_Count": int(selected_k),
                    "Deep_tuning": "Yes" if deep_tuning else "No",
                    "Source_file": record.get("Source_file", ""),
                    "Source_sheet": record.get("Source_sheet", ""),
                    "Source_row": record.get("Source_row", ""),
                })

    aggregated, _pred_map = _aggregate_cv_rows(rows)
    comparison: List[Dict[str, object]] = []
    for name in model_names:
        subset = [r for r in aggregated if r["Model"] == name]
        metrics = _metrics(
            [float(r["Actual_concentration"]) for r in subset],
            [float(r["Predicted_concentration"]) for r in subset],
        )
        metrics.update({
            "Model": name,
            "CV_scheme": f"Repeated {n_splits}-fold x {max(1, int(repeats))}",
            "Expected_sample_predictions": n,
            "Actual_sample_predictions": len(subset),
            "Failed_folds": failures.get(name, 0),
        })
        comparison.append(metrics)

    comparison, _raw_trend = _augment_comparison_with_trend(comparison, aggregated)
    baseline = next((x for x in comparison if x.get("Model") == "GlobalMedian"), None)
    base_med = _num(baseline.get("Median_Fold_Error")) if baseline else None
    for row in comparison:
        med = _num(row.get("Median_Fold_Error"))
        if base_med is not None and med is not None and base_med > 0:
            row["Median_Fold_Improvement_vs_Global_pct"] = (base_med - med) / base_med * 100.0
        else:
            row["Median_Fold_Improvement_vs_Global_pct"] = ""
    comparison = _sort_comparison_for_objective(comparison, model_objective)
    return rows, comparison, warnings, selection_events, selection_curve, tuning_events


def _summarise_feature_selection(
    events: Sequence[Dict[str, object]],
    manifest: List[Dict[str, object]],
    numeric_candidates: Sequence[str],
    categorical_candidates: Sequence[str],
    *,
    auto_select: bool,
    stability_threshold_pct: float,
    min_features: int,
    max_features: int,
) -> Tuple[List[str], List[str], List[Dict[str, object]], List[Dict[str, object]], List[Dict[str, object]]]:
    feature_order = list(numeric_candidates) + list(categorical_candidates)
    fold_keys = {
        (int(_num(r.get("Outer_Repeat")) or 0), int(_num(r.get("Outer_Fold")) or 0))
        for r in events
    }
    n_folds = max(1, len(fold_keys))
    selected_counts: Dict[str, int] = {name: 0 for name in feature_order}
    expanded_counts: Dict[str, int] = {name: 0 for name in feature_order}
    for row in events:
        name = str(row.get("Feature", ""))
        if name not in selected_counts:
            continue
        if str(row.get("Selected", "")).lower() == "yes":
            selected_counts[name] += 1
            expanded_counts[name] += int(_num(row.get("Selected_Expanded_Levels")) or 0)

    ranked = sorted(
        feature_order,
        key=lambda name: (-selected_counts.get(name, 0), feature_order.index(name)),
    )
    if auto_select and events:
        chosen = [
            name for name in ranked
            if selected_counts.get(name, 0) / n_folds * 100.0 >= float(stability_threshold_pct)
        ]
        if len(chosen) < max(2, int(min_features)):
            chosen = ranked[: min(len(ranked), max(2, int(min_features)))]
        chosen = chosen[: min(len(chosen), max(2, int(max_features)))]
    else:
        chosen = feature_order[: min(len(feature_order), max(2, int(max_features)))]

    selected_set = set(chosen)
    manifest_by_name = {str(r.get("Feature", "")): r for r in manifest}
    selection_rows: List[Dict[str, object]] = []
    selected_rows: List[Dict[str, object]] = []
    for rank, name in enumerate(ranked, start=1):
        base = manifest_by_name.get(name, {})
        freq = selected_counts.get(name, 0) / n_folds * 100.0 if events else (100.0 if name in selected_set else 0.0)
        row = {
            "Rank_by_stability": rank,
            "Feature": name,
            "Group": base.get("Group", _feature_group(name)),
            "Type": base.get("Type", ""),
            "Selection_frequency_pct": float(freq),
            "Selected_folds": selected_counts.get(name, 0),
            "Total_outer_folds": n_folds if events else "",
            "Mean_selected_expanded_levels": (
                expanded_counts.get(name, 0) / max(1, selected_counts.get(name, 0))
                if selected_counts.get(name, 0) else 0.0
            ),
            "Final_selected": "Yes" if name in selected_set else "No",
            "Valid_N": base.get("Valid_N", ""),
            "Missing_pct": base.get("Missing_pct", ""),
            "Reason": (
                f"selection frequency >= {float(stability_threshold_pct):.1f}%"
                if name in selected_set and freq >= float(stability_threshold_pct)
                else "retained to satisfy minimum descriptor count"
                if name in selected_set
                else "low selection stability / no added CV value"
            ),
        }
        selection_rows.append(row)
        if name in selected_set:
            selected_rows.append(row)
        if name in manifest_by_name:
            m = manifest_by_name[name]
            m["Selected_Frequency_pct"] = float(freq)
            m["Selected_in_final_model"] = "Yes" if name in selected_set else "No"
            if name in selected_set:
                m["Final_Status"] = "Selected"
                m["Reason"] = row["Reason"]
            elif str(m.get("Final_Status")) == "Candidate":
                m["Final_Status"] = "Not selected"
                m["Reason"] = "low nested-CV selection stability or no incremental value"

    unavailable = [
        dict(row) for row in manifest
        if str(row.get("Final_Status")) in {"Unavailable", "Disabled", "Generation failed", "Too sparse"}
    ]
    selected_numeric = [x for x in numeric_candidates if x in selected_set]
    selected_categorical = [x for x in categorical_candidates if x in selected_set]
    return selected_numeric, selected_categorical, selection_rows, selected_rows, unavailable


def _build_model_input_audit(
    calibration: Sequence[Dict[str, object]],
    selected_numeric: Sequence[str],
    selected_categorical: Sequence[str],
    *,
    language: str = "en",
) -> List[Dict[str, object]]:
    """Document how raw descriptors become a purely numeric model matrix."""
    rows: List[Dict[str, object]] = []
    for name in selected_numeric:
        meta = descriptor_meta(name)
        rows.append({
            "Descriptor": name,
            "Display_name": display_name(name, language),
            "Raw_type": "numeric",
            "Transformation": "median imputation; standardization for BayesianRidge/SVR/KNN/PLS/GP; original scale for tree models",
            "Expanded_numeric_columns": 1,
            "Final_model_input_is_numeric": "Yes",
            "Model_compatibility": model_compatibility("numeric", _feature_group(name)),
            "Definition": meta["Definition_zh"] if str(language).startswith("zh") else meta["Definition_en"],
        })
    for name in selected_categorical:
        levels = sorted({_text(r.get(name)) or "<MISSING>" for r in calibration})
        meta = descriptor_meta(name)
        rows.append({
            "Descriptor": name,
            "Display_name": display_name(name, language),
            "Raw_type": "categorical",
            "Transformation": "one-hot encoding to 0/1 indicator columns",
            "Expanded_numeric_columns": len(levels),
            "Final_model_input_is_numeric": "Yes",
            "Model_compatibility": model_compatibility("categorical", _feature_group(name)),
            "Definition": meta["Definition_zh"] if str(language).startswith("zh") else meta["Definition_en"],
        })
    return rows


def _build_descriptor_guide(
    manifest: Sequence[Dict[str, object]],
    *,
    language: str = "en",
) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    for item in manifest:
        name = str(item.get("Feature", ""))
        meta = descriptor_meta(name)
        rows.append({
            "Descriptor": name,
            "Display_name": display_name(name, language),
            'Display_name': meta["Name_zh"],
            "English_name": meta["Name_en"],
            "Category": group_display(str(item.get("Group", "")), language),
            "Definition": meta["Definition_zh"] if str(language).startswith("zh") else meta["Definition_en"],
            'Definition': meta["Definition_zh"],
            "English_definition": meta["Definition_en"],
            "Unit_or_type": meta["Unit_or_type"],
            "Source": meta["Source"],
            "Requires_3D": meta["Requires_3D"],
            "Raw_type": item.get("Type", ""),
            "Current_status": status_display(str(item.get("Final_Status", "")), language),
            "Technical_status": item.get("Final_Status", ""),
            "Selected_in_final_model": item.get("Selected_in_final_model", ""),
            "Selection_frequency_pct": item.get("Selected_Frequency_pct", ""),
            "Valid_N": item.get("Valid_N", ""),
            "Missing_pct": item.get("Missing_pct", ""),
            "Variance": item.get("Variance", ""),
            "Redundant_with": item.get("Redundant_with", ""),
            "Permutation_Importance_Mean": item.get("Permutation_Importance_Mean", ""),
            "Reason": item.get("Reason", ""),
            "Model_compatibility": model_compatibility(str(item.get("Type", "numeric")), str(item.get("Group", ""))),
        })
    return rows


def _detect_consensus_outliers(
    calibration: Sequence[Dict[str, object]],
    cv_aggregated: Sequence[Dict[str, object]],
    *,
    enabled: bool,
    mode: str = "conservative",
    min_fold_error: float,
    max_fraction_pct: float,
    consensus_pct: float = 75.0,
) -> Tuple[List[int], List[Dict[str, object]], List[Dict[str, object]]]:
    """Audit and optionally exclude calibration rows using repeated-CV residuals.

    Modes
    -----
    off
        Audit only; no row is excluded.
    conservative
        Original v18.29 rule: the user fold threshold is only a lower bound;
        a robust median+3.5*MAD threshold and same-direction model consensus are
        also required.  This explains why setting 2x could still yield an
        effective threshold near 19x.
    threshold_consensus
        Apply the user-specified fold threshold directly, while still requiring
        a minimum proportion of models to err in the same direction.
    top_fraction
        Rank rows by cross-model median fold error and exclude up to the user
        cap among rows above the fold threshold.  This is exploratory and does
        not require same-direction errors.
    """
    mode = str(mode or "conservative").strip().lower()
    if mode not in {"off", "conservative", "threshold_consensus", "top_fraction"}:
        mode = "conservative"
    if not enabled:
        mode = "off"

    by_index: Dict[int, List[Dict[str, object]]] = {}
    for row in cv_aggregated:
        if str(row.get("Model")) == "GlobalMedian":
            continue
        idx_value = _num(row.get("Calibration_Index"))
        idx = int(idx_value) if idx_value is not None else -1
        if 0 <= idx < len(calibration):
            by_index.setdefault(idx, []).append(row)

    scores: List[float] = []
    prelim: Dict[int, Dict[str, object]] = {}
    for idx, record in enumerate(calibration):
        rows = by_index.get(idx, [])
        signed: List[float] = []
        model_folds: List[str] = []
        for row in rows:
            actual = _num(row.get("Actual_concentration"))
            pred = _num(row.get("Predicted_concentration"))
            if actual is None or pred is None or actual <= 0 or pred <= 0:
                continue
            err = math.log10(pred / actual)
            signed.append(err)
            model_folds.append(f"{row.get('Model')}={max(pred/actual, actual/pred):.2f}x")
        if signed:
            med_signed = float(np.median(signed))
            med_abs = float(np.median(np.abs(signed)))
            pos = sum(1 for x in signed if x > 0)
            neg = sum(1 for x in signed if x < 0)
            agreement = max(pos, neg) / len(signed) * 100.0
            score = abs(med_signed)
            scores.append(score)
        else:
            med_signed = med_abs = agreement = score = float("nan")
        prelim[idx] = {
            "Calibration_Index": idx,
            "Source_file": record.get("Source_file", ""),
            "Source_sheet": record.get("Source_sheet", ""),
            "Source_row": record.get("Source_row", ""),
            "Name": record.get("Name", ""),
            "Formula": record.get("Formula", ""),
            "Combo": record.get("Combo", ""),
            "ABC_Formula_Key": _key(record),
            "Measured_ratio": record.get("Measured_ratio", ""),
            "Actual_concentration": record.get("Actual_concentration", ""),
            "Actual_log10_RRF": (
                math.log10(float(record["Measured_ratio"]) / float(record["Actual_concentration"]))
                if _valid_ratio(record) is not None and _valid_concentration(record) is not None else ""
            ),
            "Model_count": len(signed),
            "Median_signed_log10_error": med_signed if math.isfinite(med_signed) else "",
            "Median_abs_log10_error": med_abs if math.isfinite(med_abs) else "",
            "Consensus_direction_pct": agreement if math.isfinite(agreement) else "",
            "Consensus_fold_error": 10.0 ** score if math.isfinite(score) else "",
            "Per_model_fold_errors": "; ".join(model_folds),
        }

    finite_scores = np.asarray([x for x in scores if math.isfinite(x)], dtype=float)
    if finite_scores.size:
        center = float(np.median(finite_scores))
        mad = float(np.median(np.abs(finite_scores - center)))
        robust_sigma = 1.4826 * mad
        robust_cut = center + 3.5 * robust_sigma
    else:
        center = mad = robust_sigma = robust_cut = 0.0
    user_threshold_log = math.log10(max(1.01, float(min_fold_error)))
    if mode == "conservative":
        threshold_log = max(user_threshold_log, robust_cut)
    else:
        threshold_log = user_threshold_log
    threshold_fold = 10.0 ** threshold_log

    candidates: List[Tuple[int, float]] = []
    for idx, row in prelim.items():
        med_abs = _num(row.get("Median_abs_log10_error"))
        med_signed = _num(row.get("Median_signed_log10_error"))
        agreement = _num(row.get("Consensus_direction_pct"))
        model_count = int(_num(row.get("Model_count")) or 0)
        above = med_abs is not None and float(med_abs) >= threshold_log
        if mode == "conservative":
            candidate = (
                above and med_signed is not None and abs(float(med_signed)) >= threshold_log
                and agreement is not None and agreement >= float(consensus_pct)
                and model_count >= 3
            )
        elif mode == "threshold_consensus":
            candidate = (
                above and agreement is not None and agreement >= float(consensus_pct)
                and model_count >= 3
            )
        elif mode == "top_fraction":
            candidate = above and model_count >= 3
        else:
            candidate = False
        if candidate:
            candidates.append((idx, float(med_abs or 0.0)))

    max_remove = int(math.floor(len(calibration) * max(0.0, float(max_fraction_pct)) / 100.0))
    if mode != "off" and float(max_fraction_pct) > 0 and max_remove < 1 and len(calibration) >= 20:
        max_remove = 1
    candidates.sort(key=lambda item: item[1], reverse=True)
    removed = [idx for idx, _score in candidates[:max_remove]] if mode != "off" else []
    removed_set = set(removed)
    candidate_set = {idx for idx, _ in candidates}

    audit: List[Dict[str, object]] = []
    for idx in range(len(calibration)):
        row = dict(prelim[idx])
        is_candidate = idx in candidate_set
        if idx in removed_set:
            decision = "AUTO_EXCLUDED"
        elif is_candidate and mode == "off":
            decision = "FLAGGED_AUDIT_ONLY"
        elif is_candidate:
            decision = "FLAGGED_BUT_NOT_EXCLUDED_CAP_REACHED"
        else:
            decision = "KEPT"
        row.update({
            "Outlier_mode": mode,
            "Robust_center_abs_log10_error": center,
            "Robust_MAD": mad,
            "Robust_sigma": robust_sigma,
            "Robust_threshold_fold": 10.0 ** robust_cut,
            "User_threshold_fold": float(min_fold_error),
            "Effective_threshold_fold": threshold_fold,
            "Consensus_required_pct": float(consensus_pct),
            "Outlier_candidate": "Yes" if is_candidate else "No",
            "Auto_excluded": "Yes" if idx in removed_set else "No",
            "Decision": decision,
            "Interpretation": (
                "Large cross-model out-of-fold residual. Verify raw XIC, internal-standard peak, concentration entry, saturation, integration and sample identity."
                if is_candidate else
                "No outlier flag under the selected rule."
            ),
        })
        audit.append(row)

    mode_note = {
        "off": "Audit only; no automatic exclusion.",
        "conservative": "User fold threshold is a floor; robust MAD threshold and same-direction consensus are also required.",
        "threshold_consensus": "User fold threshold is applied directly with same-direction model consensus.",
        "top_fraction": "Rows above the user threshold are ranked by CV residual; up to the removal cap are excluded. Exploratory use only.",
    }[mode]
    summary = [
        {"Metric": "Outlier_mode", "Value": mode},
        {"Metric": "Mode_explanation", "Value": mode_note},
        {"Metric": "Auto_outlier_removal_enabled", "Value": mode != "off"},
        {"Metric": "Calibration_rows_before_QC", "Value": len(calibration)},
        {"Metric": "Outlier_candidates", "Value": len(candidates)},
        {"Metric": "Rows_auto_excluded", "Value": len(removed)},
        {"Metric": "Maximum_removal_fraction_pct", "Value": float(max_fraction_pct)},
        {"Metric": "Maximum_rows_by_cap", "Value": max_remove},
        {"Metric": "Minimum_fold_error_setting", "Value": float(min_fold_error)},
        {"Metric": "Robust_MAD_threshold_fold", "Value": 10.0 ** robust_cut},
        {"Metric": "Effective_threshold_fold", "Value": threshold_fold},
        {"Metric": "Consensus_direction_required_pct", "Value": float(consensus_pct)},
        {"Metric": "Caution", "Value": "Automatic exclusion is a QC hypothesis, not proof. Compare pre/post CV and manually inspect every excluded row."},
    ]
    return removed, audit, summary


def _build_accuracy_band_rows(
    cv_aggregated: Sequence[Dict[str, object]],
    comparison: Sequence[Dict[str, object]],
) -> List[Dict[str, object]]:
    """Summarise repeated-CV accuracy bands after per-sample median aggregation."""
    rank_map = {str(r.get("Model")): r.get("Rank", "") for r in comparison}
    out: List[Dict[str, object]] = []
    models = [str(r.get("Model")) for r in comparison]
    for model in models:
        subset = [r for r in cv_aggregated if str(r.get("Model")) == model]
        folds = [
            float(r["Fold_error"]) for r in subset
            if _num(r.get("Fold_error")) is not None and float(r["Fold_error"]) >= 1.0
        ]
        n = len(folds)
        if not n:
            continue

        def count_le(limit: float) -> int:
            return int(sum(1 for x in folds if x <= limit))

        c125 = count_le(1.25)
        c15 = count_le(1.5)
        c2 = count_le(2.0)
        c3 = count_le(3.0)
        c5 = count_le(5.0)
        out.append({
            "Rank": rank_map.get(model, ""),
            "Model": model,
            "N": n,
            "<=1.25x_n": c125,
            "<=1.25x_pct": c125 / n * 100.0,
            "<=1.5x_n": c15,
            "<=1.5x_pct": c15 / n * 100.0,
            "<=2x_n": c2,
            "<=2x_pct": c2 / n * 100.0,
            "<=3x_n": c3,
            "<=3x_pct": c3 / n * 100.0,
            "<=5x_n": c5,
            "<=5x_pct": c5 / n * 100.0,
            ">5x_n": n - c5,
            ">5x_pct": (n - c5) / n * 100.0,
            "Median_Fold_Error": float(np.median(folds)),
            "P80_Fold_Error": float(np.percentile(folds, 80)),
            "P90_Fold_Error": float(np.percentile(folds, 90)),
        })
    return out


def _build_model_decision(
    comparison: Sequence[Dict[str, object]],
    best_model: str,
    leave_component_summary: Sequence[Dict[str, object]],
) -> Tuple[str, str, List[Dict[str, object]], Dict[str, object], Dict[str, object]]:
    """Translate CV metrics into a conservative, auditable go/no-go decision."""
    best = dict(next((r for r in comparison if str(r.get("Model")) == str(best_model)), {}))
    baseline = dict(next((r for r in comparison if str(r.get("Model")) == "GlobalMedian"), {}))

    expected = _num(best.get("Expected_sample_predictions")) or 0.0
    actual_n = _num(best.get("Actual_sample_predictions")) or 0.0
    coverage = (actual_n / expected * 100.0) if expected > 0 else 0.0
    failed_folds = int(_num(best.get("Failed_folds")) or 0)
    med_fold = _num(best.get("Median_Fold_Error"))
    p80 = _num(best.get("P80_Fold_Error"))
    within2 = _num(best.get("Within_2x_pct"))
    within5 = _num(best.get("Within_5x_pct"))
    r2 = _num(best.get("R2_log_concentration"))
    rmse = _num(best.get("RMSE_log10"))
    improvement = _num(best.get("Median_Fold_Improvement_vs_Global_pct"))

    leave_valid = [r for r in leave_component_summary if _num(r.get("Median_Fold_Error")) is not None]
    leave_worst_med = max((_num(r.get("Median_Fold_Error")) or 0.0 for r in leave_valid), default=float("nan"))
    leave_min_within2 = min((_num(r.get("Within_2x_pct")) or 0.0 for r in leave_valid), default=float("nan"))

    criteria: List[Dict[str, object]] = []

    def add(name: str, value: object, threshold: str, passed: bool, meaning: str) -> None:
        criteria.append({
            "Criterion": name,
            "Observed": value,
            "Recommended_threshold": threshold,
            "Pass": "Yes" if bool(passed) else "No",
            "Interpretation": meaning,
        })

    add(
        "CV prediction coverage",
        coverage,
        ">=95%",
        coverage >= 95.0,
        "Nearly every calibration compound must receive an out-of-fold prediction.",
    )
    add(
        "Failed CV folds",
        failed_folds,
        "=0",
        failed_folds == 0,
        "Model fitting should not silently fail in any repeated-CV fold.",
    )
    add(
        "Improvement over GlobalMedian",
        improvement if improvement is not None else "",
        ">=10% preferred; >=5% minimum",
        improvement is not None and improvement >= 5.0,
        "A complex model must outperform the no-structure global response-factor baseline.",
    )
    add(
        "Median fold error",
        med_fold if med_fold is not None else "",
        "<=2.0 for semiquantitative; <=3.5 for range estimation",
        med_fold is not None and med_fold <= 3.5,
        "Typical multiplicative concentration error.",
    )
    add(
        "P80 fold error",
        p80 if p80 is not None else "",
        "<=3.0 preferred",
        p80 is not None and p80 <= 3.0,
        "80% of predictions should remain within this multiplicative range.",
    )
    add(
        "Within 2x",
        within2 if within2 is not None else "",
        ">=60% preferred",
        within2 is not None and within2 >= 60.0,
        "Share of calibration compounds predicted within two-fold error.",
    )
    add(
        "Within 5x",
        within5 if within5 is not None else "",
        ">=80% minimum for screening/range use",
        within5 is not None and within5 >= 80.0,
        "Large-error containment for exploratory use.",
    )
    add(
        "R2 on log concentration",
        r2 if r2 is not None else "",
        ">0 minimum; >=0.2 preferred",
        r2 is not None and r2 > 0.0,
        "A negative value means the model is worse than predicting the mean log concentration.",
    )
    add(
        "RMSE on log10 concentration",
        rmse if rmse is not None else "",
        "<=0.30 preferred (~2-fold)",
        rmse is not None and rmse <= 0.30,
        "0.301 log10 units corresponds approximately to a two-fold error.",
    )
    if leave_valid:
        add(
            "Worst leave-component-out median fold error",
            leave_worst_med,
            "<=3.5 preferred",
            math.isfinite(leave_worst_med) and leave_worst_med <= 3.5,
            "Tests transfer to an unseen A, B or C component formula.",
        )
        add(
            "Worst leave-component-out Within 2x",
            leave_min_within2,
            ">=40% preferred",
            math.isfinite(leave_min_within2) and leave_min_within2 >= 40.0,
            "Low values indicate poor extrapolation to unseen component chemistry.",
        )

    complete = coverage >= 95.0 and failed_folds == 0
    semiquant = (
        complete
        and improvement is not None and improvement >= 10.0
        and med_fold is not None and med_fold <= 2.0
        and p80 is not None and p80 <= 3.0
        and within2 is not None and within2 >= 60.0
        and r2 is not None and r2 >= 0.20
    )
    range_only = (
        complete
        and improvement is not None and improvement >= 5.0
        and med_fold is not None and med_fold <= 3.5
        and within5 is not None and within5 >= 80.0
    )

    if semiquant:
        label = "GO \u2013 semi-quantitative within the validated applicability domain"
    elif range_only:
        label = "CAUTION \u2013 concentration range estimation / prioritisation only"
    else:
        label = "NO-GO \u2013 screening only; do not report as quantitative concentration"

    metric_parts = []
    if med_fold is not None:
        metric_parts.append(f"median fold error={med_fold:.2f}x")
    if p80 is not None:
        metric_parts.append(f"P80={p80:.2f}x")
    if within2 is not None:
        metric_parts.append(f"Within2x={within2:.1f}%")
    if r2 is not None:
        metric_parts.append(f"R2log={r2:.3f}")
    if improvement is not None:
        metric_parts.append(f"vs GlobalMedian={improvement:+.1f}%")
    summary = (
        f"{best_model}: " + ", ".join(metric_parts) + ". "
        + (
            "The model passes the conservative semiquantitative gate."
            if semiquant else
            "The model improves the baseline but should be interpreted as a broad range estimator."
            if range_only else
            "Cross-validation does not demonstrate reliable concentration prediction."
        )
    )

    criteria.insert(0, {
        "Criterion": "OVERALL DECISION",
        "Observed": label,
        "Recommended_threshold": "",
        "Pass": "",
        "Interpretation": summary,
    })
    criteria.insert(1, {
        "Criterion": "Selected model",
        "Observed": best_model,
        "Recommended_threshold": "",
        "Pass": "",
        "Interpretation": "Selected only from repeated out-of-fold performance, not training fit.",
    })
    return label, summary, criteria, best, baseline





def _build_raw_vs_corrected_rows(
    cv_aggregated: Sequence[Dict[str, object]], best_model: str,
) -> List[Dict[str, object]]:
    subset = [r for r in cv_aggregated if str(r.get("Model")) == str(best_model)]
    if not subset:
        return []
    actual = np.asarray([float(r["Actual_concentration"]) for r in subset], dtype=float)
    raw = np.asarray([float(r["Measured_ratio"]) for r in subset], dtype=float)
    corrected = np.asarray([float(r["Predicted_concentration"]) for r in subset], dtype=float)
    n = len(subset)

    def rank_info(values: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
        order = np.argsort(values, kind="mergesort")
        ranks = np.empty(n, dtype=int)
        for rank, idx in enumerate(order, start=1):
            ranks[int(idx)] = rank
        return ranks, ranks.astype(float) / n * 100.0

    actual_rank, actual_pct = rank_info(actual)
    raw_rank, raw_pct = rank_info(raw)
    corrected_rank, corrected_pct = rank_info(corrected)
    def rank_class(rank: int) -> str:
        frac = float(rank) / max(1, n)
        return "Low" if frac <= 1.0 / 3.0 else "Medium" if frac <= 2.0 / 3.0 else "High"

    rows: List[Dict[str, object]] = []
    for i, src in enumerate(subset):
        rows.append({
            "Calibration_Index": src.get("Calibration_Index", ""),
            "Injection_Group": src.get("Injection_Group", ""),
            "Source_file": src.get("Source_file", ""),
            "Source_sheet": src.get("Source_sheet", ""),
            "Source_row": src.get("Source_row", ""),
            "Combo": src.get("Combo", ""),
            "ABC_Formula_Key": src.get("ABC_Formula_Key", ""),
            "Actual_concentration": float(actual[i]),
            "Raw_area_IS_ratio": float(raw[i]),
            "OOF_response_corrected_score_log10": float(math.log10(corrected[i])),
            "OOF_estimated_concentration": float(corrected[i]),
            "Actual_rank_low_to_high": int(actual_rank[i]),
            "Raw_rank_low_to_high": int(raw_rank[i]),
            "Corrected_rank_low_to_high": int(corrected_rank[i]),
            "Actual_percentile": float(actual_pct[i]),
            "Raw_percentile": float(raw_pct[i]),
            "Corrected_percentile": float(corrected_pct[i]),
            "Actual_class": rank_class(int(actual_rank[i])),
            "Raw_rank_class": rank_class(int(raw_rank[i])),
            "Corrected_rank_class": rank_class(int(corrected_rank[i])),
            "Raw_absolute_rank_error": int(abs(int(raw_rank[i]) - int(actual_rank[i]))),
            "Corrected_absolute_rank_error": int(abs(int(corrected_rank[i]) - int(actual_rank[i]))),
        })
    return rows


def _build_trend_decision(
    comparison: Sequence[Dict[str, object]],
    best_model: str,
    injection_summary: Sequence[Dict[str, object]],
    *,
    hybrid: bool = False,
) -> Tuple[str, str, List[Dict[str, object]], Dict[str, object], Dict[str, object]]:
    best = dict(next((r for r in comparison if str(r.get("Model")) == str(best_model)), {}))
    baseline = dict(next((r for r in comparison if str(r.get("Model")) == "GlobalMedian"), {}))
    spearman = _num(best.get("Trend_Spearman_r"))
    raw_spearman = _num(best.get("Raw_Trend_Spearman_r"))
    delta_s = _num(best.get("Delta_Trend_Spearman_r_vs_raw"))
    kendall = _num(best.get("Trend_Kendall_tau"))
    pairwise = _num(best.get("Trend_Pairwise_Concordance_pct"))
    top20 = _num(best.get("Trend_Top20_Overlap_pct"))
    tertile = _num(best.get("Trend_Tertile_Accuracy_pct"))
    med_fold = _num(best.get("Median_Fold_Error"))
    p80 = _num(best.get("P80_Fold_Error"))
    coverage = _num(best.get("Actual_sample_predictions")) or 0.0
    expected = _num(best.get("Expected_sample_predictions")) or 0.0
    coverage_pct = coverage / expected * 100.0 if expected else 0.0

    group_all = next((r for r in injection_summary if str(r.get("Injection_Group")) == "ALL_GROUPS"), {})
    group_s = _num(group_all.get("Trend_Spearman_r"))
    group_delta = _num(group_all.get("Delta_Spearman_vs_raw"))

    rows: List[Dict[str, object]] = []

    def add(name: str, observed: object, threshold: str, passed: bool, interpretation: str) -> None:
        rows.append({
            "Criterion": name,
            "Observed": observed,
            "Recommended_threshold": threshold,
            "Pass": "Yes" if passed else "No",
            "Interpretation": interpretation,
        })

    add("Out-of-fold prediction coverage", coverage_pct, ">=95%", coverage_pct >= 95.0,
        "Trend evidence must use out-of-fold predictions for nearly every standard.")
    add("Corrected Spearman rank correlation", spearman if spearman is not None else "", ">=0.50 strong; >=0.30 useful",
        spearman is not None and spearman >= 0.30,
        "Primary evidence that response correction restores concentration ordering.")
    add("Raw-ratio Spearman baseline", raw_spearman if raw_spearman is not None else "", "reference only", True,
        "Raw area/internal-standard ratio before response correction.")
    add("Spearman improvement vs raw ratio", delta_s if delta_s is not None else "", ">=+0.10 strong; >=+0.05 useful",
        delta_s is not None and delta_s >= 0.05,
        "The corrected score should outperform the uncorrected signal.")
    add("Kendall tau", kendall if kendall is not None else "", ">=0.25 useful", kendall is not None and kendall >= 0.25,
        "Agreement of pairwise ordering including ties.")
    add("Pairwise concordance", pairwise if pairwise is not None else "", ">=60% useful; >=65% strong",
        pairwise is not None and pairwise >= 60.0,
        "Share of standard pairs whose high/low ordering is correct.")
    add("Top-20% overlap", top20 if top20 is not None else "", ">=50% useful; >=60% strong",
        top20 is not None and top20 >= 50.0,
        "Recovery of the truly highest-concentration standards.")
    add("Low/medium/high rank-tertile accuracy", tertile if tertile is not None else "", ">=50% useful",
        tertile is not None and tertile >= 50.0,
        "Coarse concentration-class agreement.")
    if group_all:
        add("Leave-one-injection-group-out Spearman", group_s if group_s is not None else "", ">=0.25 preferred",
            group_s is not None and group_s >= 0.25,
            "Checks transfer to an entirely held-out standard-mixture injection.")
        add("Injection-group Spearman improvement vs raw", group_delta if group_delta is not None else "", ">=0 preferred",
            group_delta is not None and group_delta >= 0.0,
            "Response correction should not lose the raw trend when an injection group is held out.")
    if hybrid:
        add("Median fold error", med_fold if med_fold is not None else "", "<=3.5 for broad range use",
            med_fold is not None and med_fold <= 3.5,
            "Hybrid mode also constrains approximate concentration magnitude.")
        add("P80 fold error", p80 if p80 is not None else "", "<=5 preferred",
            p80 is not None and p80 <= 5.0,
            "80% empirical concentration range.")

    strong = (
        coverage_pct >= 95.0
        and spearman is not None and spearman >= 0.50
        and delta_s is not None and delta_s >= 0.10
        and pairwise is not None and pairwise >= 65.0
        and top20 is not None and top20 >= 60.0
    )
    useful = (
        coverage_pct >= 95.0
        and spearman is not None and spearman >= 0.30
        and delta_s is not None and delta_s >= 0.05
        and pairwise is not None and pairwise >= 58.0
    )
    if hybrid:
        strong = strong and med_fold is not None and med_fold <= 3.0
        useful = useful and med_fold is not None and med_fold <= 4.0

    if strong:
        label = "TREND-GO \u2013 response correction supports concentration ranking within the validated domain"
    elif useful:
        label = "TREND-CAUTION \u2013 useful for broad ranking/classes; exact concentration remains uncertain"
    else:
        label = "TREND-NO-GO \u2013 no reproducible improvement over raw area ratio"

    summary = (
        f"{best_model}: corrected Spearman={spearman if spearman is not None else float('nan'):.3f}, "
        f"raw Spearman={raw_spearman if raw_spearman is not None else float('nan'):.3f}, "
        f"delta={delta_s if delta_s is not None else float('nan'):+.3f}, "
        f"pairwise={pairwise if pairwise is not None else float('nan'):.1f}%, "
        f"Top20 overlap={top20 if top20 is not None else float('nan'):.1f}%. "
        "The unknown 1000-compound sample was not used as a supervised calibration label."
    )
    rows.insert(0, {
        "Criterion": "OVERALL TREND DECISION", "Observed": label,
        "Recommended_threshold": "", "Pass": "", "Interpretation": summary,
    })
    rows.insert(1, {
        "Criterion": "Selected model", "Observed": best_model,
        "Recommended_threshold": "", "Pass": "", "Interpretation": "Selected from out-of-fold trend performance.",
    })
    return label, summary, rows, best, baseline


def _applicability_domain_rows(
    calibration: Sequence[Dict[str, object]],
    targets: Sequence[Dict[str, object]],
    numeric: Sequence[str],
    categorical: Sequence[str],
) -> List[Dict[str, object]]:
    """Simple auditable applicability-domain estimate in selected feature space."""
    if not targets:
        return []
    n_cal = len(calibration)
    if n_cal == 0:
        return [{} for _ in targets]
    if numeric:
        cal_num = np.asarray([
            [np.nan if _num(r.get(name)) is None else float(_num(r.get(name))) for name in numeric]
            for r in calibration
        ], dtype=float)
        tgt_num = np.asarray([
            [np.nan if _num(r.get(name)) is None else float(_num(r.get(name))) for name in numeric]
            for r in targets
        ], dtype=float)
        med = np.nanmedian(cal_num, axis=0)
        med = np.where(np.isfinite(med), med, 0.0)
        cal_num = np.where(np.isfinite(cal_num), cal_num, med)
        tgt_num = np.where(np.isfinite(tgt_num), tgt_num, med)
        scale = np.nanstd(cal_num, axis=0)
        scale = np.where(np.isfinite(scale) & (scale > 1e-12), scale, 1.0)
        cal_num = (cal_num - med) / scale
        tgt_num = (tgt_num - med) / scale
    else:
        cal_num = np.zeros((n_cal, 0), dtype=float)
        tgt_num = np.zeros((len(targets), 0), dtype=float)

    def distance_vector(target_index: int) -> np.ndarray:
        if cal_num.shape[1]:
            d = np.sqrt(np.mean((cal_num - tgt_num[target_index]) ** 2, axis=1))
        else:
            d = np.zeros(n_cal, dtype=float)
        if categorical:
            penalties = np.zeros(n_cal, dtype=float)
            for j, cal_row in enumerate(calibration):
                mismatches = sum(
                    1 for name in categorical
                    if _text(cal_row.get(name)) != _text(targets[target_index].get(name))
                )
                penalties[j] = mismatches / max(1, len(categorical))
            d = d + penalties
        return d

    # Reference distribution: nearest other calibration sample.
    loo_nearest: List[float] = []
    for i in range(n_cal):
        if cal_num.shape[1]:
            d = np.sqrt(np.mean((cal_num - cal_num[i]) ** 2, axis=1))
        else:
            d = np.zeros(n_cal, dtype=float)
        if categorical:
            for j in range(n_cal):
                mismatches = sum(
                    1 for name in categorical
                    if _text(calibration[j].get(name)) != _text(calibration[i].get(name))
                )
                d[j] += mismatches / max(1, len(categorical))
        d[i] = np.inf
        finite = d[np.isfinite(d)]
        if finite.size:
            loo_nearest.append(float(np.min(finite)))
    q50 = float(np.percentile(loo_nearest, 50)) if loo_nearest else 1.0
    q90 = float(np.percentile(loo_nearest, 90)) if loo_nearest else 2.0
    q90 = max(q90, q50 + 1e-12)

    out: List[Dict[str, object]] = []
    for i, _target in enumerate(targets):
        d = distance_vector(i)
        order = np.argsort(d, kind="mergesort")[: min(5, len(d))]
        nearest = float(d[order[0]]) if len(order) else float("nan")
        if not math.isfinite(nearest):
            level = "Unknown"
        elif nearest <= q50:
            level = "High"
        elif nearest <= q90:
            level = "Medium"
        else:
            level = "Low"
        summaries = []
        for idx in order[:3]:
            cal = calibration[int(idx)]
            summaries.append(
                f"{_text(cal.get('Combo')) or _key(cal)} | C={_text(cal.get('Actual_concentration'))} | d={float(d[idx]):.3f}"
            )
        out.append({
            "Applicability_domain": level,
            "Nearest_calibration_distance": nearest if math.isfinite(nearest) else "",
            "AD_reference_median_nearest_distance": q50,
            "AD_reference_P90_nearest_distance": q90,
            "Nearest_calibrants": "; ".join(summaries),
        })
    return out


def _fit_best_and_predict(
    best_name: str,
    calibration: Sequence[Dict[str, object]],
    targets: Sequence[Dict[str, object]],
    X: np.ndarray,
    X_target: np.ndarray,
    y: np.ndarray,
    models: Dict[str, object],
    cv_aggregated: Sequence[Dict[str, object]],
    p80_fold: float,
    *,
    selected_k: int,
    selected_numeric: Sequence[str] = (),
    selected_categorical: Sequence[str] = (),
    deep_tuning: bool = False,
    random_state: int = 42,
) -> Tuple[List[Dict[str, object]], Optional[object], List[str], Dict[str, object]]:
    warnings: List[str] = []
    fitted = None
    final_tuning: Dict[str, object] = {}
    fitted_model_name = best_name
    if best_name == "GlobalMedian":
        pred_y = np.full(len(targets), float(np.median(y)), dtype=float)
    else:
        try:
            with pywarnings.catch_warnings():
                pywarnings.simplefilter("ignore")
                fitted, best_params, inner_mae = _fit_with_optional_tuning(
                    models[best_name], best_name, X, y,
                    selected_k=int(selected_k),
                    deep_tuning=bool(deep_tuning),
                    random_state=int(random_state) + 991,
                )
                pred_y = _predict_estimator(fitted, X_target)
            final_tuning = {
                "Model": best_name,
                "Stage": "full_standard_calibration_refit",
                "Deep_tuning": "Yes" if deep_tuning else "No",
                "Selected_transformed_feature_count": int(selected_k),
                "Inner_MAE_log10_RRF": inner_mae,
                "Best_params": "; ".join(f"{k}={v}" for k, v in sorted(best_params.items())),
            }
        except Exception as exc:
            warnings.append(f"Best model {best_name} failed on full fit; GlobalMedian used: {exc}")
            fitted_model_name = "GlobalMedian"
            fitted = None
            pred_y = np.full(len(targets), float(np.median(y)), dtype=float)

    standard_cv: Dict[str, Dict[str, object]] = {}
    for row in cv_aggregated:
        if str(row.get("Model")) != str(best_name):
            continue
        key = _text(row.get("ABC_Formula_Key"))
        if key:
            standard_cv[key] = {
                "CV_predicted_standard_concentration": row.get("Predicted_concentration", ""),
                "CV_standard_fold_error": row.get("Fold_error", ""),
                "CV_standard_actual_concentration": row.get("Actual_concentration", ""),
            }

    ad_rows = _applicability_domain_rows(
        calibration, targets, selected_numeric, selected_categorical,
    )
    fold = max(1.0, float(p80_fold or 1.0))
    out: List[Dict[str, object]] = []
    predicted_values: List[float] = []
    raw_ratios: List[float] = []
    for i, record in enumerate(targets):
        ratio = float(record["Measured_ratio"])
        log_rrf = float(pred_y[i])
        corrected_score = float(math.log10(ratio) - log_rrf)
        pred = float(10.0 ** corrected_score)
        predicted_values.append(pred)
        raw_ratios.append(ratio)
        key = _key(record)
        cv_ref = standard_cv.get(key, {})
        result = {
            "Best_Model": fitted_model_name,
            "Source_file": record.get("Source_file", ""),
            "Source_sheet": record.get("Source_sheet", ""),
            "Source_row": record.get("Source_row", ""),
            "Name": record.get("Name", ""),
            "Formula": record.get("Formula", ""),
            "Combo": record.get("Combo", ""),
            "ABC_Formula_Key": key,
            "A_Formula": record.get("A_Formula", ""),
            "B_Formula": record.get("B_Formula", ""),
            "C_Formula": record.get("C_Formula", ""),
            "Raw_area_IS_ratio": ratio,
            "Raw_log10_ratio": math.log10(ratio),
            "Predicted_log10_RRF": log_rrf,
            "Predicted_RRF": 10.0 ** log_rrf,
            "Response_corrected_score_log10": corrected_score,
            "Estimated_concentration": pred,
            "Estimated_concentration_P80_lower": pred / fold,
            "Estimated_concentration_P80_upper": pred * fold,
            "P80_fold_factor": fold,
            "Overlaps_standard_structure": "Yes" if key in standard_cv else "No",
            "Standard_CV_predicted_concentration_reference": cv_ref.get("CV_predicted_standard_concentration", ""),
            "Standard_CV_fold_error_reference": cv_ref.get("CV_standard_fold_error", ""),
            "Note_on_overlap": (
                "The matching standard concentration belongs to the separate standard injection and is not the unknown sample concentration."
                if key in standard_cv else ""
            ),
            "Apex_RT_min": record.get("Apex_RT_min", ""),
            "Mobile_phase_A_pct": record.get("Mobile_phase_A_pct", ""),
            "Mobile_phase_B_pct": record.get("Mobile_phase_B_pct", ""),
            "B_slope_pct_per_min": record.get("B_slope_pct_per_min", ""),
            "Structure_hash": record.get("Structure_hash", ""),
            "Formula_match": record.get("Formula_match", ""),
        }
        if i < len(ad_rows):
            result.update(ad_rows[i])
        out.append(result)

    if out:
        n = len(out)
        pred_order = np.argsort(np.asarray(predicted_values, dtype=float), kind="mergesort")
        raw_order = np.argsort(np.asarray(raw_ratios, dtype=float), kind="mergesort")
        pred_rank = np.empty(n, dtype=int)
        raw_rank = np.empty(n, dtype=int)
        for rank, idx in enumerate(pred_order, start=1):
            pred_rank[int(idx)] = rank
        for rank, idx in enumerate(raw_order, start=1):
            raw_rank[int(idx)] = rank
        q1, q2 = np.quantile(np.asarray(predicted_values, dtype=float), [1.0 / 3.0, 2.0 / 3.0])
        med_pred = float(np.median(predicted_values)) if predicted_values else 1.0
        for i, row in enumerate(out):
            value = float(predicted_values[i])
            row["Corrected_rank_low_to_high"] = int(pred_rank[i])
            row["Corrected_percentile"] = float(pred_rank[i] / n * 100.0)
            row["Raw_ratio_rank_low_to_high"] = int(raw_rank[i])
            row["Raw_ratio_percentile"] = float(raw_rank[i] / n * 100.0)
            row["Concentration_class"] = "Low" if value <= q1 else "Medium" if value <= q2 else "High"
            row["Relative_to_target_median_pct"] = float(value / med_pred * 100.0) if med_pred > 0 else ""
        add_decile_columns(out, "Estimated_concentration", prefix="Concentration")
    return out, fitted, warnings, final_tuning


def _leave_component_out(
    best_name: str,
    calibration: Sequence[Dict[str, object]],
    X: np.ndarray,
    y: np.ndarray,
    models: Dict[str, object],
) -> Tuple[List[Dict[str, object]], List[Dict[str, object]], List[str]]:
    rows: List[Dict[str, object]] = []
    warnings: List[str] = []
    for role in ("A_Formula", "B_Formula", "C_Formula"):
        groups = sorted({_text(r.get(role)) for r in calibration if _text(r.get(role))})
        for group in groups:
            test_idx = np.asarray([i for i, r in enumerate(calibration) if _text(r.get(role)) == group], dtype=int)
            train_idx = np.asarray([i for i in range(len(calibration)) if i not in set(test_idx.tolist())], dtype=int)
            if len(test_idx) == 0 or len(train_idx) < 10:
                continue
            try:
                if best_name == "GlobalMedian":
                    pred_y = np.full(len(test_idx), float(np.median(y[train_idx])), dtype=float)
                else:
                    est = clone(models[best_name])
                    with pywarnings.catch_warnings():
                        pywarnings.simplefilter("ignore")
                        est.fit(X[train_idx], y[train_idx])
                        pred_y = _predict_estimator(est, X[test_idx])
                for pos, idx in enumerate(test_idx):
                    record = calibration[int(idx)]
                    ratio = float(record["Measured_ratio"])
                    actual = float(record["Actual_concentration"])
                    pred = float(ratio / (10.0 ** float(pred_y[pos])))
                    rows.append({
                        "Role": role.replace("_Formula", ""), "Held_out_Formula": group,
                        "Calibration_Index": int(idx), "ABC_Formula_Key": _key(record),
                        "Combo": record.get("Combo", ""), "Actual_concentration": actual,
                        "Predicted_concentration": pred,
                        "Fold_error": max(pred / actual, actual / pred),
                        "APE_pct": abs(pred - actual) / actual * 100.0,
                    })
            except Exception as exc:
                warnings.append(f"Leave-{role}-out failed for {group}: {exc}")
    summary: List[Dict[str, object]] = []
    for role in ("A", "B", "C"):
        subset = [r for r in rows if r["Role"] == role]
        metrics = _metrics(
            [float(r["Actual_concentration"]) for r in subset],
            [float(r["Predicted_concentration"]) for r in subset],
        )
        metrics["Role"] = role
        metrics["Groups"] = len({r["Held_out_Formula"] for r in subset})
        summary.append(metrics)
    return rows, summary, warnings



def _leave_injection_group_out(
    best_name: str,
    calibration: Sequence[Dict[str, object]],
    X: np.ndarray,
    y: np.ndarray,
    models: Dict[str, object],
) -> Tuple[List[Dict[str, object]], List[Dict[str, object]], List[str]]:
    """Hold out each standard-mixture injection group as a whole."""
    warnings: List[str] = []
    groups = sorted({_text(r.get("Injection_Group")) for r in calibration if _text(r.get("Injection_Group"))})
    if len(groups) < 2:
        return [], [], [
            "Leave-one-injection-group-out was skipped: fewer than two nonblank injection groups were available. "
            "Set the standard injection-group column in the GUI if the training table contains RAW/batch identifiers."
        ]
    rows: List[Dict[str, object]] = []
    for group in groups:
        test_idx = np.asarray([i for i, r in enumerate(calibration) if _text(r.get("Injection_Group")) == group], dtype=int)
        test_set = set(test_idx.tolist())
        train_idx = np.asarray([i for i in range(len(calibration)) if i not in test_set], dtype=int)
        if len(test_idx) == 0 or len(train_idx) < 10:
            warnings.append(f"Injection group {group!r} skipped: train={len(train_idx)}, test={len(test_idx)}")
            continue
        try:
            if best_name == "GlobalMedian":
                pred_y = np.full(len(test_idx), float(np.median(y[train_idx])), dtype=float)
            else:
                est = clone(models[best_name])
                with pywarnings.catch_warnings():
                    pywarnings.simplefilter("ignore")
                    est.fit(X[train_idx], y[train_idx])
                    pred_y = _predict_estimator(est, X[test_idx])
            for pos, idx in enumerate(test_idx):
                record = calibration[int(idx)]
                ratio = float(record["Measured_ratio"])
                actual = float(record["Actual_concentration"])
                corrected_score = math.log10(ratio) - float(pred_y[pos])
                pred = float(10.0 ** corrected_score)
                rows.append({
                    "Injection_Group": group,
                    "Calibration_Index": int(idx),
                    "ABC_Formula_Key": _key(record),
                    "Combo": record.get("Combo", ""),
                    "Measured_ratio": ratio,
                    "Actual_concentration": actual,
                    "Predicted_log10_RRF": float(pred_y[pos]),
                    "Response_corrected_score_log10": corrected_score,
                    "Predicted_concentration": pred,
                    "Fold_error": max(pred / actual, actual / pred),
                    "Source_file": record.get("Source_file", ""),
                    "Source_sheet": record.get("Source_sheet", ""),
                    "Source_row": record.get("Source_row", ""),
                })
        except Exception as exc:
            warnings.append(f"Leave-injection-group-out failed for {group}: {exc}")

    summary: List[Dict[str, object]] = []
    for group in groups + ["ALL_GROUPS"]:
        subset = rows if group == "ALL_GROUPS" else [r for r in rows if r["Injection_Group"] == group]
        if not subset:
            continue
        actual = [float(r["Actual_concentration"]) for r in subset]
        predicted = [float(r["Predicted_concentration"]) for r in subset]
        ratio = [float(r["Measured_ratio"]) for r in subset]
        metrics = _metrics(actual, predicted)
        metrics.update(_trend_metrics(actual, predicted))
        raw = _trend_metrics(actual, ratio)
        metrics.update({f"Raw_{k}": v for k, v in raw.items()})
        metrics["Delta_Spearman_vs_raw"] = (
            (_num(metrics.get("Trend_Spearman_r")) or 0.0) - (_num(raw.get("Trend_Spearman_r")) or 0.0)
        )
        metrics["Injection_Group"] = group
        metrics["Group_Count"] = len({r["Injection_Group"] for r in subset})
        summary.append(metrics)
    return rows, summary, warnings


def _feature_importance(fitted, X: np.ndarray, y: np.ndarray, names: Sequence[str], random_state: int) -> List[Dict[str, object]]:
    if fitted is None or len(names) == 0:
        return []
    try:
        with pywarnings.catch_warnings():
            pywarnings.simplefilter("ignore")
            imp = permutation_importance(
                fitted, X, y, scoring="neg_mean_absolute_error", n_repeats=12,
                random_state=int(random_state), n_jobs=1,
            )
        rows = [
            {"Feature": names[i], "Permutation_Importance_Mean": float(imp.importances_mean[i]), "Permutation_Importance_SD": float(imp.importances_std[i])}
            for i in range(len(names))
        ]
        rows.sort(key=lambda r: float(r["Permutation_Importance_Mean"]), reverse=True)
        for rank, row in enumerate(rows, start=1):
            row["Rank"] = rank
        return rows
    except Exception:
        return []


def _safe_plot_directory(output_xlsx: Path) -> Tuple[Path, List[str]]:
    """Use a deliberately short plot path to avoid Windows MAX_PATH failures."""
    warnings: List[str] = []
    token = hashlib.sha1(str(Path(output_xlsx)).encode("utf-8", errors="ignore")).hexdigest()[:8]
    candidates = [
        Path(output_xlsx).parent / f"_ctplots_{token}",
        Path(tempfile.gettempdir()) / f"ctplots_{token}",
    ]
    last_error = ""
    for candidate in candidates:
        try:
            candidate.mkdir(parents=True, exist_ok=True)
            for old in candidate.glob("*.png"):
                try:
                    old.unlink()
                except Exception:
                    pass
            if candidate != candidates[0]:
                warnings.append(
                    f"Model plots were written to temporary short path because the output path was too long or unavailable: {candidate}"
                )
            return candidate, warnings
        except Exception as exc:
            last_error = str(exc)
    raise OSError(f"Could not create a usable short plot directory: {last_error}")


def _make_plots(
    output_xlsx: Path,
    comparison: Sequence[Dict[str, object]],
    cv_aggregated: Sequence[Dict[str, object]],
    best_model: str,
    feature_importance: Sequence[Dict[str, object]],
    leave_component_summary: Sequence[Dict[str, object]],
    *,
    descriptor_manifest: Sequence[Dict[str, object]] = (),
    descriptor_selection: Sequence[Dict[str, object]] = (),
    descriptor_selection_curve: Sequence[Dict[str, object]] = (),
    outlier_audit: Sequence[Dict[str, object]] = (),
    comparison_before_qc: Sequence[Dict[str, object]] = (),
    raw_vs_corrected: Sequence[Dict[str, object]] = (),
    injection_group_summary: Sequence[Dict[str, object]] = (),
    model_objective: str = "trend",
    output_language: str = "en",
) -> Tuple[List[Path], List[str]]:
    """Create independent diagnostic figures without allowing one failed plot to abort the benchmark."""
    warnings: List[str] = []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        if str(output_language).lower().startswith("zh"):
            matplotlib.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Arial Unicode MS", "DejaVu Sans"]
            matplotlib.rcParams["axes.unicode_minus"] = False
    except Exception as exc:
        return [], [f"Matplotlib unavailable; model tables were created without plots: {exc}"]

    zh = str(output_language).lower().startswith("zh")
    def L(en: str, cn: str) -> str:
        return cn if zh else en

    try:
        plot_dir, path_warnings = _safe_plot_directory(Path(output_xlsx))
        warnings.extend(path_warnings)
    except Exception as exc:
        return [], [f"Could not initialise plot directory; model tables were still created: {exc}"]

    paths: List[Path] = []

    def save(fig, filename: str, label: str) -> Optional[Path]:
        target = plot_dir / filename
        try:
            fig.tight_layout()
            try:
                fig.savefig(str(target), dpi=170, bbox_inches="tight")
                paths.append(target)
                return target
            except Exception as first_exc:
                token = hashlib.sha1(
                    (str(output_xlsx) + "|" + filename).encode("utf-8", errors="ignore")
                ).hexdigest()[:10]
                fallback_dir = Path(tempfile.gettempdir()) / "combitrace_plots"
                fallback_dir.mkdir(parents=True, exist_ok=True)
                fallback = fallback_dir / f"{token}_{filename}"
                try:
                    fig.savefig(str(fallback), dpi=170, bbox_inches="tight")
                    paths.append(fallback)
                    warnings.append(
                        f"{label} plot used a short temporary path because the requested path failed: {first_exc}"
                    )
                    return fallback
                except Exception as second_exc:
                    warnings.append(
                        f"{label} plot skipped. Requested path error: {first_exc}; "
                        f"temporary-path error: {second_exc}"
                    )
                    return None
        finally:
            try:
                plt.close(fig)
            except Exception:
                pass

    ranked = sorted(
        [r for r in comparison if _num(r.get("Median_Fold_Error")) is not None],
        key=lambda r: int(_num(r.get("Rank")) or 999),
    )
    model_names = [str(r.get("Model")) for r in ranked]

    # 1. Median / P80 / P90 fold error.
    if ranked:
        try:
            fig, ax = plt.subplots(figsize=(10.0, 5.8))
            x = np.arange(len(ranked), dtype=float)
            width = 0.25
            med = [float(r["Median_Fold_Error"]) for r in ranked]
            p80 = [float(r["P80_Fold_Error"]) for r in ranked]
            p90 = [float(r["P90_Fold_Error"]) for r in ranked]
            ax.bar(x - width, med, width=width, label=L("Median", 'Median'))
            ax.bar(x, p80, width=width, label="P80")
            ax.bar(x + width, p90, width=width, label="P90")
            ax.set_xticks(x)
            ax.set_xticklabels(model_names, rotation=35, ha="right")
            ax.set_ylabel(L("Fold error", 'Fold error'))
            ax.set_title(L("Repeated-CV fold-error comparison (lower is better)", 'Repeated-CV fold-error comparison (lower is better)'))
            ax.legend()
            ax.grid(axis="y", alpha=0.25)
            save(fig, "01_fold.png", "Fold-error comparison")
        except Exception as exc:
            warnings.append(f"Fold-error comparison plot skipped: {exc}")

    # 2. Accuracy bands.
    if ranked:
        try:
            fig, ax = plt.subplots(figsize=(10.0, 5.8))
            x = np.arange(len(ranked), dtype=float)
            width = 0.34
            within2 = [float(r.get("Within_2x_pct") or 0.0) for r in ranked]
            within5 = [float(r.get("Within_5x_pct") or 0.0) for r in ranked]
            ax.bar(x - width / 2, within2, width=width, label=L("Within 2x", 'Within 2x'))
            ax.bar(x + width / 2, within5, width=width, label=L("Within 5x", 'Within 5x'))
            ax.set_xticks(x)
            ax.set_xticklabels(model_names, rotation=35, ha="right")
            ax.set_ylabel(L("CV predictions (%)", 'CV predictions (%)'))
            ax.set_ylim(0, 105)
            ax.set_title(L("Repeated-CV accuracy bands (higher is better)", 'Repeated-CV accuracy bands (higher is better)'))
            ax.legend()
            ax.grid(axis="y", alpha=0.25)
            save(fig, "02_accuracy.png", "Accuracy-band comparison")
        except Exception as exc:
            warnings.append(f"Accuracy-band comparison plot skipped: {exc}")

    # 3. R2 on log concentration.
    if ranked:
        try:
            fig, ax = plt.subplots(figsize=(10.0, 5.4))
            vals = [float(r.get("R2_log_concentration") or 0.0) for r in ranked]
            ax.bar(range(len(ranked)), vals)
            ax.axhline(0.0, linestyle="--", linewidth=1.0)
            ax.set_xticks(range(len(ranked)))
            ax.set_xticklabels(model_names, rotation=35, ha="right")
            ax.set_ylabel(L("R\u00B2 on log10 concentration", 'R² on log10 concentration'))
            ax.set_title(L("Repeated-CV trend agreement (positive is required)", 'Repeated-CV trend agreement (positive is required)'))
            ax.grid(axis="y", alpha=0.25)
            save(fig, "03_r2.png", "R2 comparison")
        except Exception as exc:
            warnings.append(f"R2 comparison plot skipped: {exc}")

    # 4. RMSE on log10 concentration.
    if ranked:
        try:
            fig, ax = plt.subplots(figsize=(10.0, 5.4))
            vals = [float(r.get("RMSE_log10") or 0.0) for r in ranked]
            ax.bar(range(len(ranked)), vals)
            ax.axhline(math.log10(2.0), linestyle=":", linewidth=1.0, label=L("2-fold \u2248 0.301", '2-fold ≈ 0.301'))
            ax.axhline(math.log10(5.0), linestyle="--", linewidth=1.0, label=L("5-fold \u2248 0.699", '5-fold ≈ 0.699'))
            ax.set_xticks(range(len(ranked)))
            ax.set_xticklabels(model_names, rotation=35, ha="right")
            ax.set_ylabel(L("RMSE (log10 concentration)", 'RMSE (log10 concentration)'))
            ax.set_title(L("Repeated-CV log error (lower is better)", 'Repeated-CV log error (lower is better)'))
            ax.legend()
            ax.grid(axis="y", alpha=0.25)
            save(fig, "04_rmse.png", "RMSE comparison")
        except Exception as exc:
            warnings.append(f"RMSE comparison plot skipped: {exc}")

    best_rows = [r for r in cv_aggregated if str(r.get("Model")) == str(best_model)]
    best_metric = next((r for r in comparison if str(r.get("Model")) == str(best_model)), {})

    # 5. Actual vs out-of-fold prediction.
    if best_rows:
        try:
            fig, ax = plt.subplots(figsize=(6.8, 6.2))
            actual = np.asarray([float(r["Actual_concentration"]) for r in best_rows], dtype=float)
            pred = np.asarray([float(r["Predicted_concentration"]) for r in best_rows], dtype=float)
            ax.scatter(actual, pred, alpha=0.78)
            lo = max(min(actual.min(), pred.min()), 1e-15)
            hi = max(actual.max(), pred.max())
            ax.plot([lo, hi], [lo, hi], linestyle="-", linewidth=1.2, label="1x")
            ax.plot([lo, hi], [lo * 2, hi * 2], linestyle=":", linewidth=1.0, label="2x")
            ax.plot([lo, hi], [lo / 2, hi / 2], linestyle=":", linewidth=1.0)
            ax.plot([lo, hi], [lo * 5, hi * 5], linestyle="--", linewidth=0.9, label="5x")
            ax.plot([lo, hi], [lo / 5, hi / 5], linestyle="--", linewidth=0.9)
            ax.set_xscale("log")
            ax.set_yscale("log")
            ax.set_xlabel(L("Actual concentration", 'Actual concentration'))
            ax.set_ylabel(L("Repeated-CV predicted concentration", 'Repeated-CV predicted concentration'))
            ax.set_title(f"{best_model}: " + L("actual vs out-of-fold prediction", 'actual vs out-of-fold prediction'))
            ax.grid(alpha=0.25)
            ax.legend()
            annotation = (
                f"Median fold={_num(best_metric.get('Median_Fold_Error')) or float('nan'):.2f}x\n"
                f"P80={_num(best_metric.get('P80_Fold_Error')) or float('nan'):.2f}x\n"
                f"Within2x={_num(best_metric.get('Within_2x_pct')) or 0.0:.1f}%\n"
                f"R\u00B2log={_num(best_metric.get('R2_log_concentration')) or float('nan'):.3f}"
            )
            ax.text(
                0.03, 0.97, annotation, transform=ax.transAxes, va="top", ha="left",
                bbox={"boxstyle": "round", "alpha": 0.15},
            )
            save(fig, "05_actual.png", "Actual-vs-predicted")
        except Exception as exc:
            warnings.append(f"Actual-vs-predicted plot skipped: {exc}")

    # 6. Signed log-error distribution.
    if best_rows:
        try:
            fig, ax = plt.subplots(figsize=(8.2, 5.2))
            errors = np.asarray([
                math.log10(float(r["Predicted_concentration"]) / float(r["Actual_concentration"]))
                for r in best_rows
                if float(r["Predicted_concentration"]) > 0 and float(r["Actual_concentration"]) > 0
            ])
            bins = max(8, min(24, int(round(math.sqrt(len(errors)) * 2))))
            ax.hist(errors, bins=bins, alpha=0.8)
            ax.axvline(0.0, linestyle="-", linewidth=1.0, label=L("No bias", 'No bias'))
            ax.axvline(math.log10(2.0), linestyle=":", linewidth=1.0, label="\u00B12x")
            ax.axvline(-math.log10(2.0), linestyle=":", linewidth=1.0)
            ax.axvline(math.log10(5.0), linestyle="--", linewidth=1.0, label="\u00B15x")
            ax.axvline(-math.log10(5.0), linestyle="--", linewidth=1.0)
            ax.set_xlabel(L("log10(predicted / actual)", 'log10(predicted / actual)'))
            ax.set_ylabel(L("Count", 'Count'))
            ax.set_title(f"{best_model}: " + L("repeated-CV signed error distribution", 'repeated-CV signed error distribution'))
            ax.legend()
            ax.grid(axis="y", alpha=0.25)
            save(fig, "06_error.png", "Error distribution")
        except Exception as exc:
            warnings.append(f"Error-distribution plot skipped: {exc}")

    # 7. Empirical cumulative fold-error profile: selected model versus baseline.
    try:
        baseline_rows = [r for r in cv_aggregated if str(r.get("Model")) == "GlobalMedian"]
        if best_rows:
            fig, ax = plt.subplots(figsize=(8.0, 5.4))
            for label, rows in ((best_model, best_rows), ("GlobalMedian", baseline_rows)):
                folds = sorted(
                    float(r["Fold_error"]) for r in rows
                    if _num(r.get("Fold_error")) is not None and float(r["Fold_error"]) >= 1.0
                )
                if not folds:
                    continue
                yvals = np.arange(1, len(folds) + 1, dtype=float) / len(folds) * 100.0
                ax.plot(folds, yvals, linewidth=1.8, label=label)
            ax.axvline(2.0, linestyle=":", linewidth=1.0)
            ax.axvline(5.0, linestyle="--", linewidth=1.0)
            ax.set_xscale("log")
            ax.set_xlabel(L("Fold error", 'Fold error'))
            ax.set_ylabel(L("CV predictions at or below error (%)", 'CV predictions at or below error (%)'))
            ax.set_title(L("Cumulative fold-error profile", 'Cumulative fold-error profile'))
            ax.legend()
            ax.grid(alpha=0.25)
            save(fig, "07_ecdf.png", "Fold-error ECDF")
    except Exception as exc:
        warnings.append(f"Fold-error ECDF plot skipped: {exc}")

    # 8. Exploratory permutation importance.
    if feature_importance:
        try:
            fig, ax = plt.subplots(figsize=(8.8, 6.8))
            top = list(feature_importance[:20])[::-1]
            ax.barh(range(len(top)), [float(r["Permutation_Importance_Mean"]) for r in top])
            ax.set_yticks(range(len(top)))
            ax.set_yticklabels([display_name(str(r["Feature"]), output_language) for r in top])
            ax.set_xlabel(L("Increase in MAE(log10 RRF) after permutation", 'Increase in MAE(log10 RRF) after permutation'))
            ax.set_title(f"{best_model}: " + L("exploratory feature importance", 'exploratory feature importance'))
            ax.grid(axis="x", alpha=0.25)
            save(fig, "08_importance.png", "Feature importance")
        except Exception as exc:
            warnings.append(f"Feature-importance plot skipped: {exc}")

    # 9. Leave-one-component-formula-out transfer performance.
    valid_leave = [
        r for r in leave_component_summary
        if _num(r.get("Median_Fold_Error")) is not None
    ]
    if valid_leave:
        try:
            fig, ax = plt.subplots(figsize=(7.4, 5.0))
            labels = [str(r.get("Role")) for r in valid_leave]
            vals = [float(r["Median_Fold_Error"]) for r in valid_leave]
            ax.bar(range(len(labels)), vals)
            ax.set_xticks(range(len(labels)))
            ax.set_xticklabels(labels)
            ax.set_ylabel(L("Median fold error", 'Median fold error'))
            ax.set_title(f"{best_model}: " + L("leave-one-component-formula-out", 'leave-one-component-formula-out'))
            ax.axhline(2.0, linestyle=":", linewidth=1.0)
            ax.axhline(5.0, linestyle="--", linewidth=1.0)
            ax.grid(axis="y", alpha=0.25)
            save(fig, "09_leave.png", "Leave-component-out")
        except Exception as exc:
            warnings.append(f"Leave-component-out plot skipped: {exc}")

    # 10. Descriptor status overview.
    if descriptor_manifest:
        try:
            counts: Dict[str, int] = {}
            for row in descriptor_manifest:
                status = str(row.get("Final_Status") or row.get("Initial_Status") or "Unknown")
                counts[status] = counts.get(status, 0) + 1
            labels = sorted(counts, key=lambda x: (-counts[x], x))
            fig, ax = plt.subplots(figsize=(8.6, 5.2))
            ax.bar(range(len(labels)), [counts[x] for x in labels])
            ax.set_xticks(range(len(labels)))
            ax.set_xticklabels([status_display(x, output_language) for x in labels], rotation=30, ha="right")
            ax.set_ylabel(L("Descriptor count", 'Descriptor count'))
            ax.set_title(L("Descriptor generation and automatic-selection status", 'Descriptor generation and automatic-selection status'))
            ax.grid(axis="y", alpha=0.25)
            save(fig, "10_desc_status.png", "Descriptor status")
        except Exception as exc:
            warnings.append(f"Descriptor-status plot skipped: {exc}")

    # 11. Nested-CV selection stability.
    if descriptor_selection:
        try:
            ranked_desc = sorted(
                [r for r in descriptor_selection if _num(r.get("Selection_frequency_pct")) is not None],
                key=lambda r: float(r.get("Selection_frequency_pct") or 0.0),
                reverse=True,
            )[:30]
            if ranked_desc:
                top = ranked_desc[::-1]
                fig, ax = plt.subplots(figsize=(9.2, 7.4))
                ax.barh(range(len(top)), [float(r["Selection_frequency_pct"]) for r in top])
                ax.set_yticks(range(len(top)))
                ax.set_yticklabels([display_name(str(r.get("Feature")), output_language) for r in top])
                ax.set_xlim(0, 105)
                ax.set_xlabel(L("Outer-fold selection frequency (%)", 'Outer-fold selection frequency (%)'))
                ax.set_title(L("Descriptors retained by nested-CV stability selection", 'Descriptors retained by nested-CV stability selection'))
                ax.grid(axis="x", alpha=0.25)
                save(fig, "11_desc_stability.png", "Descriptor stability")
        except Exception as exc:
            warnings.append(f"Descriptor-stability plot skipped: {exc}")

    # 12. Inner-CV feature-count tuning curve.
    if descriptor_selection_curve:
        try:
            by_k: Dict[int, List[float]] = {}
            for row in descriptor_selection_curve:
                k = int(_num(row.get("Feature_Count")) or 0)
                score = _num(row.get("Inner_MAE_log10_RRF"))
                if k > 0 and score is not None:
                    by_k.setdefault(k, []).append(float(score))
            if by_k:
                ks = sorted(by_k)
                means = [float(np.mean(by_k[k])) for k in ks]
                sds = [float(np.std(by_k[k])) for k in ks]
                fig, ax = plt.subplots(figsize=(8.0, 5.0))
                ax.errorbar(ks, means, yerr=sds, marker="o", capsize=3)
                ax.set_xlabel(L("Selected transformed-feature count", 'Selected transformed-feature count'))
                ax.set_ylabel(L("Inner-CV MAE on log10 RRF", 'Inner-CV MAE on log10 RRF'))
                ax.set_title(L("Automatic descriptor-count tuning", 'Automatic descriptor-count tuning'))
                ax.grid(alpha=0.25)
                save(fig, "12_desc_count.png", "Descriptor-count tuning")
        except Exception as exc:
            warnings.append(f"Descriptor-count tuning plot skipped: {exc}")

    # 13. Outlier audit.  Only rows with a consensus score are shown.
    if outlier_audit:
        try:
            valid = [r for r in outlier_audit if _num(r.get("Consensus_fold_error")) is not None]
            valid.sort(key=lambda r: float(r.get("Consensus_fold_error") or 1.0), reverse=True)
            if valid:
                vals = [float(r["Consensus_fold_error"]) for r in valid]
                threshold = _num(valid[0].get("Effective_threshold_fold")) or 1.0
                fig, ax = plt.subplots(figsize=(9.2, 5.4))
                ax.scatter(range(1, len(vals) + 1), vals, alpha=0.8)
                removed_x = [i + 1 for i, r in enumerate(valid) if str(r.get("Auto_excluded")) == "Yes"]
                removed_y = [vals[i] for i, r in enumerate(valid) if str(r.get("Auto_excluded")) == "Yes"]
                if removed_x:
                    ax.scatter(removed_x, removed_y, marker="x", s=80, label=L("Auto-excluded", 'Auto-excluded'))
                ax.axhline(float(threshold), linestyle="--", linewidth=1.0, label=L(f"Threshold {threshold:.2f}x", f'Threshold: {threshold:.2f}x'))
                ax.set_yscale("log")
                ax.set_xlabel(L("Calibration rows sorted by consensus residual", 'Calibration rows sorted by consensus residual'))
                ax.set_ylabel(L("Cross-model consensus fold error", 'Cross-model consensus fold error'))
                ax.set_title(L("Training-set outlier audit (manual verification required)", 'Training-set outlier audit (manual verification required)'))
                ax.legend()
                ax.grid(alpha=0.25)
                save(fig, "13_outliers.png", "Outlier audit")
        except Exception as exc:
            warnings.append(f"Outlier-audit plot skipped: {exc}")

    # 14. Before/after QC comparison for the best model in each stage.
    if comparison_before_qc:
        try:
            before = comparison_before_qc[0] if comparison_before_qc else {}
            after = comparison[0] if comparison else {}
            if before and after:
                labels = [L("Before QC", 'Before QC'), L("After QC", 'After QC')]
                med = [float(before.get("Median_Fold_Error") or 0.0), float(after.get("Median_Fold_Error") or 0.0)]
                p80 = [float(before.get("P80_Fold_Error") or 0.0), float(after.get("P80_Fold_Error") or 0.0)]
                fig, ax = plt.subplots(figsize=(7.2, 4.8))
                x = np.arange(2, dtype=float)
                width = 0.34
                ax.bar(x - width / 2, med, width=width, label=L("Median fold", 'Median fold error'))
                ax.bar(x + width / 2, p80, width=width, label=L("P80 fold", 'P80 fold'))
                ax.set_xticks(x)
                ax.set_xticklabels(labels)
                ax.set_ylabel(L("Fold error", 'Fold error'))
                ax.set_title(L("Automatic training-QC impact", 'Automatic training-QC impact'))
                ax.legend()
                ax.grid(axis="y", alpha=0.25)
                save(fig, "14_qc_compare.png", "QC before/after")
        except Exception as exc:
            warnings.append(f"QC before/after plot skipped: {exc}")

    # 15. Names of descriptors that could not be generated or were too sparse.
    unavailable_named = [
        r for r in descriptor_manifest
        if str(r.get("Final_Status")) in {"Unavailable", "Disabled", "Generation failed", "Too sparse"}
    ]
    if unavailable_named:
        try:
            unavailable_named = sorted(
                unavailable_named,
                key=lambda r: (
                    -float(_num(r.get("Missing_pct")) or 0.0),
                    str(r.get("Feature", "")),
                ),
            )[:30]
            top = unavailable_named[::-1]
            fig, ax = plt.subplots(figsize=(9.4, max(4.8, 0.30 * len(top) + 1.8)))
            values = [float(_num(r.get("Missing_pct")) or 0.0) for r in top]
            ax.barh(range(len(top)), values)
            ax.set_yticks(range(len(top)))
            ax.set_yticklabels([display_name(str(r.get("Feature", "")), output_language) for r in top])
            ax.set_xlim(0, 105)
            ax.set_xlabel(L("Missing / unavailable calibration rows (%)", 'Missing / unavailable calibration rows (%)'))
            ax.set_title(L("Descriptors not automatically available", 'Descriptors not automatically available'))
            ax.grid(axis="x", alpha=0.25)
            save(fig, "15_unavailable.png", "Unavailable descriptors")
        except Exception as exc:
            warnings.append(f"Unavailable-descriptor plot skipped: {exc}")


    # 16. Trend agreement by model: the main evidence in response-correction mode.
    if comparison:
        try:
            valid = [r for r in comparison if _num(r.get("Trend_Spearman_r")) is not None]
            if valid:
                labels = [str(r.get("Model")) for r in valid]
                corrected = [float(r.get("Trend_Spearman_r") or 0.0) for r in valid]
                raw = [float(r.get("Raw_Trend_Spearman_r") or 0.0) for r in valid]
                fig, ax = plt.subplots(figsize=(10.4, 5.4))
                x = np.arange(len(labels), dtype=float)
                width = 0.36
                ax.bar(x - width / 2, raw, width=width, label=L("Raw ratio", 'Raw ratio'))
                ax.bar(x + width / 2, corrected, width=width, label=L("Response-corrected", 'Response-corrected'))
                ax.axhline(0.0, linewidth=0.8)
                ax.set_xticks(x)
                ax.set_xticklabels(labels, rotation=35, ha="right")
                ax.set_ylim(-1.0, 1.0)
                ax.set_ylabel(L("Spearman rank correlation", 'Spearman rank correlation'))
                ax.set_title(L("Out-of-fold concentration-trend recovery", 'Out-of-fold concentration-trend recovery'))
                ax.legend()
                ax.grid(axis="y", alpha=0.25)
                save(fig, "16_trend_spearman.png", "Trend Spearman")
        except Exception as exc:
            warnings.append(f"Trend-Spearman plot skipped: {exc}")

    # 17. Actual rank versus raw and corrected ranks for the selected model.
    if raw_vs_corrected:
        try:
            ordered = sorted(raw_vs_corrected, key=lambda r: int(_num(r.get("Actual_rank_low_to_high")) or 0))
            actual_rank = np.asarray([float(r.get("Actual_rank_low_to_high") or 0.0) for r in ordered])
            raw_rank = np.asarray([float(r.get("Raw_rank_low_to_high") or 0.0) for r in ordered])
            corrected_rank = np.asarray([float(r.get("Corrected_rank_low_to_high") or 0.0) for r in ordered])
            fig, ax = plt.subplots(figsize=(7.0, 6.2))
            ax.plot(actual_rank, actual_rank, linestyle="--", linewidth=1.0, label=L("Ideal", 'Ideal'))
            ax.scatter(actual_rank, raw_rank, alpha=0.65, label=L("Raw ratio rank", 'Raw ratio rank'))
            ax.scatter(actual_rank, corrected_rank, alpha=0.75, label=L("Corrected rank", 'Corrected rank'))
            ax.set_xlabel(L("Actual concentration rank", 'Actual concentration rank'))
            ax.set_ylabel(L("Signal-derived rank", 'Signal-derived rank'))
            ax.set_title(f"{best_model}: " + L("raw versus response-corrected ranking", 'raw versus response-corrected ranking'))
            ax.legend()
            ax.grid(alpha=0.25)
            save(fig, "17_rank_recovery.png", "Rank recovery")
        except Exception as exc:
            warnings.append(f"Rank-recovery plot skipped: {exc}")

    # 18. Top-fraction recovery curve for high-concentration standards.
    if raw_vs_corrected:
        try:
            actual = np.asarray([float(r["Actual_concentration"]) for r in raw_vs_corrected], dtype=float)
            raw = np.asarray([float(r["Raw_area_IS_ratio"]) for r in raw_vs_corrected], dtype=float)
            corrected = np.asarray([float(r["OOF_estimated_concentration"]) for r in raw_vs_corrected], dtype=float)
            fractions = np.asarray([0.05, 0.10, 0.15, 0.20, 0.30, 0.40, 0.50])
            raw_overlap = [_top_fraction_overlap_pct(actual, raw, f) for f in fractions]
            cor_overlap = [_top_fraction_overlap_pct(actual, corrected, f) for f in fractions]
            fig, ax = plt.subplots(figsize=(7.8, 5.2))
            ax.plot(fractions * 100.0, raw_overlap, marker="o", label=L("Raw ratio", 'Raw ratio'))
            ax.plot(fractions * 100.0, cor_overlap, marker="o", label=L("Corrected score", 'Corrected score'))
            ax.plot(fractions * 100.0, fractions * 100.0, linestyle="--", linewidth=1.0, label=L("Random expectation", 'Random expectation'))
            ax.set_xlabel(L("Top concentration fraction (%)", 'Top concentration fraction (%)'))
            ax.set_ylabel(L("Overlap / recall (%)", 'Overlap / recall (%)'))
            ax.set_ylim(0, 105)
            ax.set_title(L("Recovery of high-concentration standards", 'Recovery of high-concentration standards'))
            ax.legend()
            ax.grid(alpha=0.25)
            save(fig, "18_topk_recovery.png", "Top-k recovery")
        except Exception as exc:
            warnings.append(f"Top-k recovery plot skipped: {exc}")

    # 19. Entire standard-mixture injection held out.
    group_rows = [r for r in injection_group_summary if str(r.get("Injection_Group")) != "ALL_GROUPS"]
    if group_rows:
        try:
            labels = [str(r.get("Injection_Group")) for r in group_rows]
            raw_s = [float(_num(r.get("Raw_Trend_Spearman_r")) or 0.0) for r in group_rows]
            cor_s = [float(_num(r.get("Trend_Spearman_r")) or 0.0) for r in group_rows]
            fig, ax = plt.subplots(figsize=(max(8.0, 0.65 * len(labels) + 3.0), 5.2))
            x = np.arange(len(labels), dtype=float)
            width = 0.36
            ax.bar(x - width / 2, raw_s, width=width, label=L("Raw ratio", 'Raw ratio'))
            ax.bar(x + width / 2, cor_s, width=width, label=L("Corrected", 'Corrected'))
            ax.axhline(0.0, linewidth=0.8)
            ax.set_xticks(x)
            ax.set_xticklabels(labels, rotation=35, ha="right")
            ax.set_ylim(-1.0, 1.0)
            ax.set_ylabel(L("Spearman correlation", 'Spearman correlation'))
            ax.set_title(L("Leave-one-standard-injection-group-out trend validation", 'Leave-one-standard-injection-group-out trend validation'))
            ax.legend()
            ax.grid(axis="y", alpha=0.25)
            save(fig, "19_injection_group.png", "Injection-group validation")
        except Exception as exc:
            warnings.append(f"Injection-group plot skipped: {exc}")


    return paths, warnings


def run_model_benchmark(
    training_records: Sequence[Dict[str, object]],
    target_records: Sequence[Dict[str, object]],
    descriptor_names: Sequence[str],
    output_xlsx: Path,
    *,
    calibration_mode: str = "training_original",
    model_objective: str = "trend",
    cv_splits: int = 5,
    cv_repeats: int = 3,
    max_features: int = 40,
    random_state: int = 42,
    expected_descriptor_names: Sequence[str] = (),
    three_d_enabled: bool = False,
    output_language: str = "en",
    auto_remove_outliers: bool = True,
    outlier_mode: str = "conservative",
    outlier_min_fold_error: float = 8.0,
    outlier_max_fraction_pct: float = 5.0,
    outlier_consensus_pct: float = 75.0,
    auto_select_features: bool = True,
    use_categorical_features: bool = False,
    deep_tuning: bool = False,
    feature_stability_threshold_pct: float = 50.0,
    min_selected_features: int = 6,
    correlation_threshold: float = 0.97,
) -> ModelBenchmarkResult:
    if not SKLEARN_AVAILABLE:
        raise RuntimeError(
            "Multi-model benchmark needs scikit-learn. For Python 3.8 install: pip install scikit-learn==1.3.2. "
            + SKLEARN_ERROR
        )

    objective = str(model_objective or "trend").strip().lower()
    if objective not in {"trend", "range", "hybrid"}:
        objective = "trend"
    calibration_all, cal_checks, mode_used, warnings = build_calibration_records(
        training_records, target_records, requested_mode="training_original",
    )
    if str(calibration_mode or "training_original").lower() != "training_original":
        warnings.append("Requested calibration mode was ignored; v18.32 always uses the separate known-standard injections for supervision.")
    targets, target_checks = _prepare_target_records(target_records)
    if len(calibration_all) < 20:
        raise ValueError(
            f"Only {len(calibration_all)} usable calibration rows; at least 20 are required for multi-model comparison"
        )

    # ------------------------------------------------------------------
    # Preliminary repeated-CV: used only for conservative consensus-outlier
    # auditing.  Descriptor selection is already nested inside each outer
    # training fold, so the test fold does not decide its own feature count.
    # ------------------------------------------------------------------
    pre_numeric, pre_categorical, pre_manifest = choose_features(
        calibration_all,
        descriptor_names,
        expected_descriptor_names=expected_descriptor_names,
        three_d_enabled=three_d_enabled,
        correlation_threshold=correlation_threshold,
        include_categorical=bool(use_categorical_features),
    )
    if len(pre_numeric) + len(pre_categorical) < 3:
        raise ValueError("Too few usable structure/gradient features for modelling")
    X_pre = _matrix(calibration_all, pre_numeric, pre_categorical)
    y_pre = np.asarray([
        math.log10(float(r["Measured_ratio"]) / float(r["Actual_concentration"]))
        for r in calibration_all
    ], dtype=float)
    pre_cats = _categories(calibration_all, pre_categorical)
    pre_models = build_models(
        len(calibration_all), len(pre_numeric), len(pre_categorical), pre_cats,
        max_features=max(3, int(max_features)), random_state=int(random_state),
    )
    pre_cv_rows, pre_comparison, pre_cv_warnings, pre_selection_events, pre_selection_curve, pre_tuning_events = _run_cv(
        calibration_all,
        X_pre,
        y_pre,
        pre_models,
        splits=cv_splits,
        repeats=cv_repeats,
        random_state=random_state,
        numeric=pre_numeric,
        categorical=pre_categorical,
        auto_select_features=bool(auto_select_features),
        min_features=max(2, int(min_selected_features)),
        max_features=max(3, int(max_features)),
        deep_tuning=False,
        model_objective=objective,
    )
    warnings.extend(pre_cv_warnings)
    pre_cv_agg, _ = _aggregate_cv_rows(pre_cv_rows)

    removed_indices, outlier_audit, outlier_summary = _detect_consensus_outliers(
        calibration_all,
        pre_cv_agg,
        enabled=bool(auto_remove_outliers),
        mode=str(outlier_mode or "conservative"),
        min_fold_error=max(1.01, float(outlier_min_fold_error)),
        max_fraction_pct=max(0.0, float(outlier_max_fraction_pct)),
        consensus_pct=max(0.0, min(100.0, float(outlier_consensus_pct))),
    )
    removed_set = set(removed_indices)
    calibration = [row for i, row in enumerate(calibration_all) if i not in removed_set]
    for idx in removed_indices:
        record = calibration_all[idx]
        cal_checks.append({
            "ABC_Formula_Key": _key(record),
            "Status": "AUTO_EXCLUDED_OUTLIER",
            "Reason": "strong cross-model repeated-CV residual consensus; inspect Outlier_Audit",
            "Actual_concentration": record.get("Actual_concentration", ""),
            "Measured_ratio": record.get("Measured_ratio", ""),
            "Training_rows": _source_label(record),
            "Target_rows": record.get("Calibration_target_sources", ""),
        })
    if len(calibration) < 20:
        warnings.append(
            "Automatic outlier exclusion would leave fewer than 20 calibration rows; no rows were removed."
        )
        calibration = list(calibration_all)
        removed_indices = []
        removed_set = set()
        for row in outlier_audit:
            if row.get("Auto_excluded") == "Yes":
                row["Auto_excluded"] = "No"
                row["Decision"] = "KEPT_MINIMUM_SAMPLE_GUARD"
        for row in outlier_summary:
            if row.get("Metric") == "Rows_auto_excluded":
                row["Value"] = 0

    # Re-estimate descriptor availability/correlation after QC because even a
    # few removed rows can change missingness, variance and redundancy.
    numeric, categorical, manifest = choose_features(
        calibration,
        descriptor_names,
        expected_descriptor_names=expected_descriptor_names,
        three_d_enabled=three_d_enabled,
        correlation_threshold=correlation_threshold,
        include_categorical=bool(use_categorical_features),
    )
    if len(numeric) + len(categorical) < 3:
        raise ValueError("Too few usable structure/gradient features after training-set QC")

    if removed_indices or bool(deep_tuning):
        X = _matrix(calibration, numeric, categorical)
        y = np.asarray([
            math.log10(float(r["Measured_ratio"]) / float(r["Actual_concentration"]))
            for r in calibration
        ], dtype=float)
        cats = _categories(calibration, categorical)
        models = build_models(
            len(calibration), len(numeric), len(categorical), cats,
            max_features=max(3, int(max_features)), random_state=int(random_state),
        )
        cv_rows, comparison, cv_warnings, selection_events, selection_curve, tuning_events = _run_cv(
            calibration,
            X,
            y,
            models,
            splits=cv_splits,
            repeats=cv_repeats,
            random_state=int(random_state) + 17,
            numeric=numeric,
            categorical=categorical,
            auto_select_features=bool(auto_select_features),
            min_features=max(2, int(min_selected_features)),
            max_features=max(3, int(max_features)),
            deep_tuning=bool(deep_tuning),
            model_objective=objective,
        )
        warnings.extend(cv_warnings)
    else:
        # Reuse the preliminary run when QC did not remove anything.
        X = X_pre
        y = y_pre
        cats = pre_cats
        models = pre_models
        cv_rows = pre_cv_rows
        comparison = pre_comparison
        selection_events = pre_selection_events
        selection_curve = pre_selection_curve
        tuning_events = pre_tuning_events
        manifest = pre_manifest
        numeric = pre_numeric
        categorical = pre_categorical

    cv_agg, _pred_map = _aggregate_cv_rows(cv_rows)
    selected_numeric, selected_categorical, selection_rows, selected_rows, unavailable_rows = _summarise_feature_selection(
        selection_events,
        manifest,
        numeric,
        categorical,
        auto_select=bool(auto_select_features),
        stability_threshold_pct=float(feature_stability_threshold_pct),
        min_features=max(2, int(min_selected_features)),
        max_features=max(3, int(max_features)),
    )
    if len(selected_numeric) + len(selected_categorical) < 3:
        warnings.append("Automatic descriptor selection returned fewer than 3 raw features; all usable candidates were restored.")
        selected_numeric = list(numeric)
        selected_categorical = list(categorical)

    # Final target fit uses only the stable raw descriptors; verbose unused
    # descriptor columns remain available in the hidden ESI_Features sheet.
    X_final = _matrix(calibration, selected_numeric, selected_categorical)
    X_target_final = _matrix(targets, selected_numeric, selected_categorical)
    final_cats = _categories(calibration, selected_categorical)
    chosen_k_values = [
        int(_num(r.get("Feature_Count")) or 0)
        for r in selection_curve
        if str(r.get("Chosen", "")).lower() == "yes" and int(_num(r.get("Feature_Count")) or 0) > 0
    ]
    final_transformed_k = (
        int(round(statistics.median(chosen_k_values)))
        if chosen_k_values else max(3, int(max_features))
    )
    final_models = build_models(
        len(calibration), len(selected_numeric), len(selected_categorical), final_cats,
        max_features=max(3, final_transformed_k),
        random_state=int(random_state),
    )

    best_row = comparison[0] if comparison else {
        "Model": "GlobalMedian", "P80_Fold_Error": 5.0, "Rating": "Poor / screening only"
    }
    best_model = str(best_row.get("Model", "GlobalMedian"))
    p80 = _num(best_row.get("P80_Fold_Error")) or 5.0
    target_predictions, fitted, fit_warnings, final_tuning = _fit_best_and_predict(
        best_model,
        calibration,
        targets,
        X_final,
        X_target_final,
        y,
        final_models,
        cv_agg,
        p80,
        selected_k=max(1, int(final_transformed_k)),
        selected_numeric=selected_numeric,
        selected_categorical=selected_categorical,
        deep_tuning=bool(deep_tuning),
        random_state=int(random_state),
    )
    warnings.extend(fit_warnings)
    if final_tuning:
        tuning_events.append(final_tuning)
    raw_feature_names = list(selected_numeric) + list(selected_categorical)
    importance = _feature_importance(fitted, X_final, y, raw_feature_names, random_state)
    importance_map = {str(r.get("Feature")): r for r in importance}
    for row in manifest:
        feature = str(row.get("Feature", ""))
        if feature in importance_map:
            row["Permutation_Importance_Mean"] = importance_map[feature].get("Permutation_Importance_Mean", "")
    for row in selected_rows:
        imp = importance_map.get(str(row.get("Feature", "")), {})
        row["Permutation_Importance_Mean"] = imp.get("Permutation_Importance_Mean", "")
        row["Permutation_Importance_SD"] = imp.get("Permutation_Importance_SD", "")
        row["Display_name"] = display_name(str(row.get("Feature", "")), output_language)
        meta = descriptor_meta(str(row.get("Feature", "")))
        row["Definition"] = meta["Definition_zh"] if str(output_language).startswith("zh") else meta["Definition_en"]
        row["Unit_or_type"] = meta["Unit_or_type"]

    model_input_audit = _build_model_input_audit(
        calibration, selected_numeric, selected_categorical, language=str(output_language or "en"),
    )
    descriptor_guide = _build_descriptor_guide(manifest, language=str(output_language or "en"))

    leave_rows, leave_summary, leave_warnings = _leave_component_out(
        best_model, calibration, X_final, y, final_models,
    )
    warnings.extend(leave_warnings)
    injection_rows, injection_summary, injection_warnings = _leave_injection_group_out(
        best_model, calibration, X_final, y, final_models,
    )
    warnings.extend(injection_warnings)

    accuracy_bands = _build_accuracy_band_rows(cv_agg, comparison)
    raw_vs_corrected_rows = _build_raw_vs_corrected_rows(cv_agg, best_model)
    if objective == "range":
        decision_label, decision_summary, decision_rows, best_metrics, baseline_metrics = _build_model_decision(
            comparison, best_model, leave_summary,
        )
    else:
        decision_label, decision_summary, decision_rows, best_metrics, baseline_metrics = _build_trend_decision(
            comparison, best_model, injection_summary, hybrid=(objective == "hybrid"),
        )
    if removed_indices:
        qc_note = (
            f"{len(removed_indices)} calibration row(s) were automatically excluded after strong cross-model OOF-residual consensus. "
            "The post-QC CV is conditional on this data-driven screen; manually verify the listed rows before publication."
        )
        warnings.append(qc_note)
        decision_rows.insert(2, {
            "Criterion": "Automatic calibration outlier QC",
            "Observed": len(removed_indices),
            "Recommended_threshold": "manual verification required",
            "Pass": "Caution",
            "Interpretation": qc_note,
        })
        decision_summary += " " + qc_note
        if decision_label.startswith("GO"):
            decision_label = "CAUTION \u2013 post-QC semiquantitative candidate; independent/manual validation required"
        if decision_rows and str(decision_rows[0].get("Criterion")) == "OVERALL DECISION":
            decision_rows[0]["Observed"] = decision_label
            decision_rows[0]["Interpretation"] = decision_summary

    # Before/after QC summary is easier to interpret than two full long tables.
    pre_best = pre_comparison[0] if pre_comparison else {}
    post_best = comparison[0] if comparison else {}
    outlier_summary.extend([
        {"Metric": "Best_model_before_QC", "Value": pre_best.get("Model", "")},
        {"Metric": "Median_fold_before_QC", "Value": pre_best.get("Median_Fold_Error", "")},
        {"Metric": "P80_fold_before_QC", "Value": pre_best.get("P80_Fold_Error", "")},
        {"Metric": "R2_log_before_QC", "Value": pre_best.get("R2_log_concentration", "")},
        {"Metric": "Best_model_after_QC", "Value": post_best.get("Model", "")},
        {"Metric": "Median_fold_after_QC", "Value": post_best.get("Median_Fold_Error", "")},
        {"Metric": "P80_fold_after_QC", "Value": post_best.get("P80_Fold_Error", "")},
        {"Metric": "R2_log_after_QC", "Value": post_best.get("R2_log_concentration", "")},
        {"Metric": "Selected_raw_descriptors", "Value": len(selected_rows)},
        {"Metric": "Final_transformed_feature_count", "Value": final_transformed_k},
        {"Metric": "Categorical_formula_features_enabled", "Value": bool(use_categorical_features)},
        {"Metric": "Deep_model_tuning_enabled", "Value": bool(deep_tuning)},
        {"Metric": "Output_language", "Value": str(output_language)},
        {"Metric": "Model_objective", "Value": objective},
        {"Metric": "Calibration_supervision", "Value": "known standard injections only"},
        {"Metric": "Unknown_target_used_for_supervision", "Value": False},
        {"Metric": "Best_corrected_Spearman", "Value": best_metrics.get("Trend_Spearman_r", "")},
        {"Metric": "Raw_ratio_Spearman", "Value": best_metrics.get("Raw_Trend_Spearman_r", "")},
        {"Metric": "Delta_Spearman_vs_raw", "Value": best_metrics.get("Delta_Trend_Spearman_r_vs_raw", "")},
        {"Metric": "Unavailable_or_too_sparse_descriptors", "Value": len(unavailable_rows)},
    ])

    plots, plot_warnings = _make_plots(
        Path(output_xlsx),
        comparison,
        cv_agg,
        best_model,
        importance,
        leave_summary,
        descriptor_manifest=manifest,
        descriptor_selection=selection_rows,
        descriptor_selection_curve=selection_curve,
        outlier_audit=outlier_audit,
        comparison_before_qc=pre_comparison,
        raw_vs_corrected=raw_vs_corrected_rows,
        injection_group_summary=injection_summary,
        model_objective=objective,
        output_language=str(output_language or "en"),
    )
    warnings.extend(plot_warnings)

    best_med = _num(best_metrics.get("Median_Fold_Error"))
    base_med = _num(baseline_metrics.get("Median_Fold_Error"))
    if best_med is not None and base_med is not None and best_med >= base_med * 0.95:
        warnings.append(
            "Best multi-model result does not improve median fold error by at least 5% over GlobalMedian; "
            "the descriptors/available calibrants do not yet support a useful predictive model."
        )

    # Add a stage column to the tuning rows so repeated runs are auditable.
    for row in selection_curve:
        row.setdefault("Stage", "after_QC" if removed_indices else "all_calibration")

    return ModelBenchmarkResult(
        enabled=True,
        calibration_mode_requested=calibration_mode,
        calibration_mode_used=mode_used,
        n_calibration=len(calibration),
        n_target_valid=len(targets),
        best_model=best_model,
        best_rating=str(best_row.get("Rating", "")),
        comparison=comparison,
        comparison_before_qc=pre_comparison,
        cv_predictions=cv_agg,
        calibration_records=calibration,
        target_predictions=target_predictions,
        calibration_checks=cal_checks,
        target_checks=target_checks,
        feature_manifest=manifest,
        feature_importance=importance,
        selected_descriptors=selected_rows,
        unavailable_descriptors=unavailable_rows,
        descriptor_selection_rows=selection_rows,
        descriptor_selection_curve=selection_curve,
        outlier_audit=outlier_audit,
        outlier_summary=outlier_summary,
        model_input_audit=model_input_audit,
        descriptor_guide=descriptor_guide,
        tuning_rows=tuning_events,
        leave_component_rows=leave_rows,
        leave_component_summary=leave_summary,
        plot_paths=plots,
        accuracy_band_rows=accuracy_bands,
        decision_rows=decision_rows,
        best_metrics=best_metrics,
        baseline_metrics=baseline_metrics,
        decision_label=decision_label,
        decision_summary=decision_summary,
        n_calibration_before_qc=len(calibration_all),
        n_outliers_removed=len(removed_indices),
        n_selected_descriptors=len(selected_rows),
        n_unavailable_descriptors=len(unavailable_rows),
        model_objective_requested=str(model_objective or "trend"),
        model_objective_used=objective,
        trend_comparison=comparison,
        raw_vs_corrected_rows=raw_vs_corrected_rows,
        injection_group_rows=injection_rows,
        injection_group_summary=injection_summary,
        best_trend_metrics={k: v for k, v in best_metrics.items() if str(k).startswith("Trend_") or str(k).startswith("Delta_Trend_")},
        raw_trend_metrics={k: v for k, v in best_metrics.items() if str(k).startswith("Raw_Trend_")},
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Workbook output
# ---------------------------------------------------------------------------

def _append_dict_sheet(wb, title: str, rows: Sequence[Dict[str, object]], headers: Optional[Sequence[str]] = None):
    ws = wb.create_sheet(title)
    if headers is None:
        headers = list(rows[0].keys()) if rows else ["Message"]
    ws.append(list(headers))
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    return ws


def append_benchmark_to_workbook(wb, result: ModelBenchmarkResult) -> None:
    _append_dict_sheet(wb, "Model_Decision", result.decision_rows)
    _append_dict_sheet(wb, "Model_Comparison", result.comparison)
    _append_dict_sheet(wb, "Model_Comparison_Before_QC", result.comparison_before_qc)
    _append_dict_sheet(wb, "Accuracy_Bands", result.accuracy_band_rows)
    _append_dict_sheet(wb, "Outlier_QC_Summary", result.outlier_summary)
    _append_dict_sheet(wb, "Outlier_Audit", result.outlier_audit)
    _append_dict_sheet(wb, "Selected_Descriptors", result.selected_descriptors)
    _append_dict_sheet(wb, "Descriptor_Guide", result.descriptor_guide)
    _append_dict_sheet(wb, "Model_Input_Audit", result.model_input_audit)
    _append_dict_sheet(wb, "Hyperparameter_Tuning", result.tuning_rows)
    _append_dict_sheet(wb, "Unavailable_Descriptors", result.unavailable_descriptors)
    _append_dict_sheet(
        wb,
        "Excluded_Descriptors",
        [r for r in result.feature_manifest if str(r.get("Final_Status")) != "Selected"],
    )
    _append_dict_sheet(wb, "Descriptor_Selection", result.descriptor_selection_rows)
    _append_dict_sheet(wb, "Descriptor_Tuning", result.descriptor_selection_curve)
    _append_dict_sheet(wb, "Model_CV_Predictions", result.cv_predictions)
    _append_dict_sheet(wb, "Raw_vs_Corrected", result.raw_vs_corrected_rows)
    _append_dict_sheet(wb, "Response_Corrected_Targets", result.target_predictions)
    _append_dict_sheet(wb, "Best_Model_Predictions", result.target_predictions)
    _append_dict_sheet(wb, "Standard_Calibration", result.calibration_records)
    _append_dict_sheet(wb, "Injection_Group_Out", result.injection_group_rows)
    _append_dict_sheet(wb, "Injection_Group_Summary", result.injection_group_summary)
    _append_dict_sheet(wb, "Calibration_Row_Check", result.calibration_checks)
    _append_dict_sheet(wb, "Model_Target_Row_Check", result.target_checks)
    _append_dict_sheet(wb, "Descriptor_Status_Full", result.feature_manifest)
    _append_dict_sheet(wb, "Feature_Importance", result.feature_importance)
    _append_dict_sheet(wb, "Leave_Component_Out", result.leave_component_rows)
    _append_dict_sheet(wb, "Leave_Component_Summary", result.leave_component_summary)

    best = result.best_metrics or {}
    baseline = result.baseline_metrics or {}
    ws = wb.create_sheet("Model_Readme")
    rows = [
        ["Item", "Value"],
        ["OVERALL DECISION", result.decision_label],
        ["Decision summary", result.decision_summary],
        ["Model objective requested", result.model_objective_requested],
        ["Model objective used", result.model_objective_used],
        ["Model target", "log10(RRF), where RRF=(standard Area/Internal Standard Area)/known standard concentration"],
        ["Response-corrected score", "log10(sample ratio) - predicted log10(RRF); equals log10 estimated concentration"],
        ["Concentration recovery", "Estimated concentration = unknown-sample measured ratio / predicted RRF"],
        ["Calibration mode requested", result.calibration_mode_requested],
        ["Calibration mode used", result.calibration_mode_used],
        ["Supervised calibration guardrail", "Only the separate known-standard injections are used. Unknown 1000-compound sample ratios are never paired with standard concentrations for training."],
        ["Usable calibration rows", result.n_calibration],
        ["Calibration rows before automatic QC", result.n_calibration_before_qc],
        ["Automatically excluded calibration rows", result.n_outliers_removed],
        ["Predictable target rows", result.n_target_valid],
        ["Selected raw descriptors", result.n_selected_descriptors],
        ["Unavailable/too-sparse descriptors", result.n_unavailable_descriptors],
        ["Best model", result.best_model],
        ["Best model rating", result.best_rating],
        ["Best median fold error", best.get("Median_Fold_Error", "")],
        ["Best P80 fold error", best.get("P80_Fold_Error", "")],
        ["Best P90 fold error", best.get("P90_Fold_Error", "")],
        ["Best Within 2x (%)", best.get("Within_2x_pct", "")],
        ["Best Within 5x (%)", best.get("Within_5x_pct", "")],
        ["Best R2 log concentration", best.get("R2_log_concentration", "")],
        ["Best RMSE log10", best.get("RMSE_log10", "")],
        ["Improvement vs GlobalMedian (%)", best.get("Median_Fold_Improvement_vs_Global_pct", "")],
        ["GlobalMedian median fold error", baseline.get("Median_Fold_Error", "")],
        ["Important", "Model comparison uses repeated cross-validation. Full-fit predictions for known calibrants are not validation results; use CV columns."],
        ["Automatic outlier QC", "Only extreme cross-model out-of-fold residual consensus is auto-excluded, under a strict removal cap. Every row remains listed in Outlier_Audit and must be checked manually."],
        ["Automatic descriptor selection", "Descriptor count is tuned inside each outer training fold; final descriptors are retained by selection stability. Unavailable, disabled, constant, sparse and highly correlated descriptors are listed separately."],
        ["Numeric model matrix", "All classic models receive numeric matrices. Numeric descriptors are imputed (and scaled where appropriate); optional A/B/C formula categories are one-hot encoded to 0/1 columns. Raw SMILES/text is never passed directly to the models."],
        ["Deep model tuning", "When enabled, compact model-specific hyperparameter grids are selected inside each outer training fold. This is slower but uses the retained descriptors more fully without training-set-only optimisation."],
        ["Descriptor guide", "Descriptor_Guide provides bilingual names, definitions, units, generation sources, 3D requirements, current status and selection results."],
        ["Decision rule", "Only the Model_Decision sheet determines whether the result supports semiquantitative use, range estimation, or screening only."],
        ["Primary trend evidence", "Raw_vs_Corrected compares out-of-fold corrected scores with raw area/internal-standard ratios using Spearman, Kendall, pairwise ordering, Top-10/20% recovery and concentration classes."],
        ["Injection-group validation", "Injection_Group_Summary holds out each complete standard-mixture injection when a RAW/batch/group column is available."],
        ["Target interpretation", "Response_Corrected_Targets reports corrected score, approximate concentration, rank, percentile, high/medium/low class, empirical range and applicability domain."],
        ["Overlap limitation", "A compound occurring in both the standard set and the unknown sample does not make the unknown-sample concentration known. Overlap is diagnostic only."],
        ["Limitation", "Structure and gradient descriptors cannot fully represent ESI matrix suppression, source conditions, pKa, saturation or unmeasured adduct distributions."],
    ]
    for row in rows:
        ws.append(row)
    ws.append([])
    ws.append(["Warnings"])
    for warning in result.warnings:
        ws.append([warning])

    # Keep the workbook readable: the concise decision/QC/descriptor sheets are
    # visible, while long row-level audit tables remain available but hidden.
    for title in (
        "Model_CV_Predictions", "Best_Model_Predictions", "Standard_Calibration",
        "Calibration_Row_Check", "Model_Target_Row_Check", "Descriptor_Status_Full",
        "Injection_Group_Out",
        "Leave_Component_Out", "Descriptor_Tuning", "Descriptor_Selection",
    ):
        if title in wb.sheetnames:
            wb[title].sheet_state = "hidden"

    if result.plot_paths:
        ps = wb.create_sheet("Model_Diagnostic_Plots")
        ps.column_dimensions["A"].width = 110
        row = 1
        image_buffers = []
        try:
            from io import BytesIO
            from openpyxl.drawing.image import Image as XLImage
            for path in result.plot_paths:
                path = Path(path)
                ps.cell(row, 1, path.name)
                if not path.exists():
                    ps.cell(row + 1, 1, f"Plot file missing: {path}")
                    result.warnings.append(f"Plot file missing during workbook embedding: {path}")
                    row += 4
                    continue
                try:
                    # Read into memory so openpyxl is not exposed to long Windows paths during wb.save().
                    buffer = BytesIO(path.read_bytes())
                    image_buffers.append(buffer)
                    img = XLImage(buffer)
                    original_w = float(getattr(img, "width", 760) or 760)
                    original_h = float(getattr(img, "height", 480) or 480)
                    img.width = 760
                    img.height = max(300, int(original_h * 760.0 / max(original_w, 1.0)))
                    img.anchor = f"A{row + 1}"
                    ps.add_image(img)
                    row += max(24, int(img.height / 19.0) + 4)
                except Exception as exc:
                    ps.cell(row + 1, 1, f"Could not embed plot: {exc}")
                    result.warnings.append(f"Could not embed plot {path.name}: {exc}")
                    row += 4
            # Keep buffers alive until Workbook.save() has completed.
            setattr(ps, "_combitrace_image_buffers", image_buffers)
        except Exception as exc:
            ps.cell(row, 1, f"Plot embedding unavailable: {exc}")
            result.warnings.append(f"Plot embedding unavailable: {exc}")
