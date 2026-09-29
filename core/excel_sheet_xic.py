"""Worksheet-based component enumeration and triplicate XIC summaries."""

from __future__ import annotations

import csv
import math
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl.utils import get_column_letter

from .abc_enumerator import ComponentRow, enumerate_abc_detailed
from .chemistry import format_formula_hill, monoisotopic_mass, parse_formula
from .targets_csv import TargetSpec


_SUBSCRIPT_TRANS = str.maketrans("\u2080\u2081\u2082\u2083\u2084\u2085\u2086\u2087\u2088\u2089", "0123456789")

def _cell_coordinate(row: int, column: int) -> str:
    """Return a coordinate even for read-only EmptyCell objects."""
    return f"{get_column_letter(int(column))}{int(row)}"



@dataclass(frozen=True)
class ExcelEnumTarget:
    no: int
    sheet_name: str
    name: str
    formula: str
    exact_mass: float
    combo: str
    duplicate_index: int
    duplicate_count: int


@dataclass
class ExcelSheetPlan:
    sheet_name: str
    headers: Tuple[str, str, str]
    a_rows: List[ComponentRow]
    b_rows: List[ComponentRow]
    c_rows: List[ComponentRow]
    targets: List[ExcelEnumTarget]
    skipped_invalid: int = 0
    warnings: List[str] = field(default_factory=list)


@dataclass(frozen=True)
class RawResolution:
    replicate_label: str
    expected_stem: str
    raw_path: Optional[Path]
    match_mode: str  # exact | normalized | fuzzy | missing | ambiguous
    note: str = ""


def _safe_text(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip().translate(_SUBSCRIPT_TRANS)
    
    if s.startswith("'"):
        s = s[1:].strip()
    return s


def normalize_internal_standard_formula(formula: str) -> Tuple[str, float]:

    text = _safe_text(formula)
    if not text:
        return "", float("nan")
    counts = parse_formula(text)
    hill = format_formula_hill(counts)
    mass = float(monoisotopic_mass(counts))
    if not hill or not math.isfinite(mass):
        raise ValueError(f'Invalid internal-standard formula: {formula!r}')
    return hill, mass


def _component_from_formula(formula: str, *, header: str, excel_row: int) -> ComponentRow:

    counts = parse_formula(formula)
    hill = format_formula_hill(counts)
    mass = monoisotopic_mass(counts)
    label = f"{header}[R{int(excel_row)}]:{hill}"
    return ComponentRow(name=label, formula_input=hill, counts=counts, exact_mass=float(mass))


def _component_from_named_formula(
    formula: str,
    *,
    reagent_name: str,
    header: str,
    excel_row: int,
) -> ComponentRow:

    counts = parse_formula(formula)
    hill = format_formula_hill(counts)
    mass = monoisotopic_mass(counts)
    reagent = _safe_text(reagent_name)
    if reagent:
        label = f"{reagent}:{hill}"
    else:
        label = f"{header}[R{int(excel_row)}]:{hill}"
    return ComponentRow(name=label, formula_input=hill, counts=counts, exact_mass=float(mass))


def _read_formula_column(ws, col_idx: int, *, header_row: int, data_start_row: int) -> Tuple[str, List[ComponentRow], List[str]]:
    header = _safe_text(ws.cell(row=int(header_row), column=int(col_idx)).value) or f'Column {int(col_idx)}'
    rows: List[ComponentRow] = []
    warnings: List[str] = []

    max_row = int(ws.max_row or data_start_row)
    for r in range(int(data_start_row), max_row + 1):
        formula = _safe_text(ws.cell(row=r, column=int(col_idx)).value)
        if not formula:
            continue
        try:
            rows.append(_component_from_formula(formula, header=header, excel_row=r))
        except Exception as e:
            warnings.append(f'{ws.title}!{_cell_coordinate(r, int(col_idx))}: could not parse {formula!r}: {e}')

    return header, rows, warnings


def _read_named_triplet_group(
    ws,
    *,
    name_col: int,
    formula_col: int,
    header_row: int,
    data_start_row: int,
) -> Tuple[str, List[ComponentRow], List[str]]:

    header = _safe_text(ws.cell(row=int(header_row), column=int(name_col)).value) or f'Column {int(name_col)}'
    rows: List[ComponentRow] = []
    warnings: List[str] = []

    max_row = int(ws.max_row or data_start_row)
    for r in range(int(data_start_row), max_row + 1):
        reagent = _safe_text(ws.cell(row=r, column=int(name_col)).value)
        formula = _safe_text(ws.cell(row=r, column=int(formula_col)).value)
        if not formula and not reagent:
            continue
        if not formula:
            warnings.append(
                f'{ws.title}!{_cell_coordinate(r, int(formula_col))} is empty; skipped component {reagent!r}'
            )
            continue
        try:
            rows.append(_component_from_named_formula(formula, reagent_name=reagent, header=header, excel_row=r))
        except Exception as e:
            warnings.append(
                f'{ws.title}!{_cell_coordinate(r, int(formula_col))}: could not parse {formula!r} (component={reagent!r}): {e}'
            )

    return header, rows, warnings


def _read_legacy_pairs(ws, *, formula_columns: Sequence[int], header_row: int, data_start_row: int):
    headers: List[str] = []
    all_rows: List[List[ComponentRow]] = []
    warnings: List[str] = []
    for col_idx in formula_columns:
        h, rows, ws_warn = _read_formula_column(
            ws,
            int(col_idx),
            header_row=int(header_row),
            data_start_row=int(data_start_row),
        )
        headers.append(h)
        all_rows.append(rows)
        warnings.extend(ws_warn)
    return headers, all_rows, warnings


def _read_named_triplets(ws, *, header_row: int, data_start_row: int):
    headers: List[str] = []
    all_rows: List[List[ComponentRow]] = []
    warnings: List[str] = []
    
    for name_col, formula_col in ((1, 2), (4, 5), (7, 8)):
        h, rows, ws_warn = _read_named_triplet_group(
            ws,
            name_col=name_col,
            formula_col=formula_col,
            header_row=int(header_row),
            data_start_row=int(data_start_row),
        )
        headers.append(h)
        all_rows.append(rows)
        warnings.extend(ws_warn)
    return headers, all_rows, warnings


def _read_named_quartets(ws, *, header_row: int, data_start_row: int):
    headers: List[str] = []
    all_rows: List[List[ComponentRow]] = []
    warnings: List[str] = []
    for name_col, formula_col in ((1, 2), (5, 6), (9, 10)):
        h, rows, ws_warn = _read_named_triplet_group(
            ws, name_col=name_col, formula_col=formula_col,
            header_row=int(header_row), data_start_row=int(data_start_row),
        )
        headers.append(h); all_rows.append(rows); warnings.extend(ws_warn)
    return headers, all_rows, warnings


def _add_duplicate_metadata(sheet_name: str, detailed_rows) -> List[ExcelEnumTarget]:
    counts: Dict[str, int] = {}
    for r in detailed_rows:
        counts[r.formula] = counts.get(r.formula, 0) + 1

    seen: Dict[str, int] = {}
    out: List[ExcelEnumTarget] = []
    for no, r in enumerate(detailed_rows, start=1):
        idx = seen.get(r.formula, 0) + 1
        seen[r.formula] = idx
        dup_count = counts.get(r.formula, 1)
        name = f"{sheet_name}_{no:04d}"
        out.append(
            ExcelEnumTarget(
                no=int(no),
                sheet_name=sheet_name,
                name=name,
                formula=r.formula,
                exact_mass=float(r.exact_mass),
                combo=r.combo,
                duplicate_index=int(idx),
                duplicate_count=int(dup_count),
            )
        )
    return out


def load_excel_sheet_plans(
    workbook_path: Path,
    *,
    formula_columns: Sequence[int] = (1, 3, 5),
    header_row: int = 1,
    data_start_row: int = 2,
    only_formula: bool = False,
    include_hidden_sheets: bool = False,
    layout: str = "auto",
) -> Tuple[List[ExcelSheetPlan], List[str]]:

    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError('Excel input requires openpyxl. Install it in the active environment: python -m pip install openpyxl') from e

    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(str(workbook_path))

    wb = load_workbook(filename=str(workbook_path), read_only=True, data_only=True)
    plans: List[ExcelSheetPlan] = []
    global_warnings: List[str] = []

    try:
        for ws in wb.worksheets:
            if not include_hidden_sheets and getattr(ws, "sheet_state", "visible") != "visible":
                global_warnings.append(f'Skipped hidden sheet: {ws.title}')
                continue

            layout_code = str(layout or "auto").strip().lower()
            if layout_code not in {"auto", "named_quartets", "named_triplets", "legacy_pairs"}:
                layout_code = "auto"

            headers: List[str]
            all_rows: List[List[ComponentRow]]
            warnings: List[str]
            used_layout = layout_code

            if layout_code == "legacy_pairs":
                headers, all_rows, warnings = _read_legacy_pairs(
                    ws, formula_columns=formula_columns, header_row=header_row, data_start_row=data_start_row
                )
            elif layout_code == "named_quartets":
                headers, all_rows, warnings = _read_named_quartets(
                    ws, header_row=header_row, data_start_row=data_start_row
                )
            elif layout_code == "named_triplets":
                headers, all_rows, warnings = _read_named_triplets(
                    ws, header_row=header_row, data_start_row=data_start_row
                )
            else:
                headers_q, rows_q, warn_q = _read_named_quartets(
                    ws, header_row=header_row, data_start_row=data_start_row
                )
                if len(rows_q) == 3 and all(rows_q):
                    headers, all_rows, warnings = headers_q, rows_q, warn_q
                    used_layout = "named_quartets"
                else:
                    headers_t, rows_t, warn_t = _read_named_triplets(
                        ws, header_row=header_row, data_start_row=data_start_row
                    )
                    if len(rows_t) == 3 and all(rows_t):
                        headers, all_rows, warnings = headers_t, rows_t, warn_t
                        used_layout = "named_triplets"
                    else:
                        headers, all_rows, warnings = _read_legacy_pairs(
                            ws, formula_columns=formula_columns, header_row=header_row, data_start_row=data_start_row
                        )
                        used_layout = "legacy_pairs"

            if len(all_rows) != 3:
                global_warnings.append(f'Skipped sheet {ws.title}: expected three component groups')
                continue

            a_rows, b_rows, c_rows = all_rows
            if not a_rows or not b_rows or not c_rows:
                if used_layout == "named_quartets":
                    hint = 'Four-column layout: A:D, E:H and I:L contain name, formula, mass and SMILES.'
                elif used_layout == "named_triplets":
                    hint = 'Three-column layout: A:C, D:F and G:I contain name, formula and mass.'
                else:
                    hint = 'Legacy layout: molecular formulas in columns A, C and E, starting at row 2.'
                global_warnings.append(
                    f'Skipped sheet {ws.title}: at least one component group has no valid formulas (A={len(a_rows)}, B={len(b_rows)}, C={len(c_rows)}；{hint})'
                )
                global_warnings.extend(warnings)
                continue

            detailed, skipped = enumerate_abc_detailed(
                a_rows,
                b_rows,
                c_rows,
                only_formula=bool(only_formula),
            )
            targets = _add_duplicate_metadata(ws.title, detailed)
            plans.append(
                ExcelSheetPlan(
                    sheet_name=ws.title,
                    headers=(headers[0], headers[1], headers[2]),
                    a_rows=a_rows,
                    b_rows=b_rows,
                    c_rows=c_rows,
                    targets=targets,
                    skipped_invalid=int(skipped),
                    warnings=warnings,
                )
            )
            global_warnings.extend(warnings)
    finally:
        wb.close()

    return plans, global_warnings


def parse_replicate_suffixes(text: str) -> List[str]:

    vals = [x.strip() for x in re.split(r"[,\uFF0C;\uFF1B\s]+", str(text or "")) if x.strip()]
    if not vals:
        vals = ["-1", "-2", "-3"]
    
    out: List[str] = []
    for x in vals:
        if x not in out:
            out.append(x)
    return out


def _normalize_stem(s: str) -> str:
    s = str(s or "").strip().lower()
    if s.endswith(".raw"):
        s = s[:-4]
    
    s = s.replace("\u2013", "-").replace("\u2014", "-").replace("\u2212", "-").replace("\uFF3F", "_")
    s = re.sub(r"[\s_]+", "-", s)
    s = re.sub(r"-+", "-", s)
    return s.strip("-")


def expected_raw_stems(sheet_name: str, suffixes: Sequence[str]) -> List[Tuple[str, str]]:

    out: List[Tuple[str, str]] = []
    for i, suffix in enumerate(suffixes, start=1):
        suffix = str(suffix or "").strip()
        label = suffix.lstrip("-_ ") or str(i)
        out.append((label, f"{sheet_name}{suffix}"))
    return out


def resolve_sheet_raws(
    sheet_name: str,
    raw_paths: Sequence[Path],
    *,
    suffixes: Sequence[str] = ("-1", "-2", "-3"),
) -> List[RawResolution]:

    raw_paths = [Path(p) for p in raw_paths]
    by_lower: Dict[str, List[Path]] = {}
    by_norm: Dict[str, List[Path]] = {}
    for p in raw_paths:
        stem = p.stem if p.suffix.lower() == ".raw" else p.name
        by_lower.setdefault(stem.lower(), []).append(p)
        by_norm.setdefault(_normalize_stem(stem), []).append(p)

    resolutions: List[RawResolution] = []
    for label, expected in expected_raw_stems(sheet_name, suffixes):
        exact = by_lower.get(expected.lower(), [])
        if len(exact) == 1:
            resolutions.append(RawResolution(label, expected, exact[0], "exact"))
            continue
        if len(exact) > 1:
            resolutions.append(RawResolution(label, expected, None, "ambiguous", f'Exact-name RAW matches: {len(exact)} files'))
            continue

        norm_key = _normalize_stem(expected)
        norm = by_norm.get(norm_key, [])
        if len(norm) == 1:
            resolutions.append(RawResolution(label, expected, norm[0], "normalized"))
            continue
        if len(norm) > 1:
            resolutions.append(RawResolution(label, expected, None, "ambiguous", f'Normalized-name RAW matches: {len(norm)} files'))
            continue

        
        fuzzy: List[Path] = []
        for p in raw_paths:
            stem = p.stem if p.suffix.lower() == ".raw" else p.name
            n = _normalize_stem(stem)
            if n == norm_key or n.endswith("-" + norm_key):
                fuzzy.append(p)
        if len(fuzzy) == 1:
            resolutions.append(RawResolution(label, expected, fuzzy[0], "fuzzy", 'RAW name has an additional prefix'))
            continue
        if len(fuzzy) > 1:
            resolutions.append(RawResolution(label, expected, None, "ambiguous", f'Partial-name matches: {len(fuzzy)} RAW files'))
            continue

        
        rep_num = label
        replaced = re.sub(r"(?i)x", str(rep_num), sheet_name)
        if replaced != sheet_name:
            alt_candidates = [replaced, f"{replaced}{suffixes[max(0, len(resolutions))] if len(suffixes) > len(resolutions) else ''}"]
            alt_hits: List[Path] = []
            for alt in alt_candidates:
                alt_hits.extend(by_norm.get(_normalize_stem(alt), []))
            uniq = list({str(p.resolve()): p for p in alt_hits}.values())
            if len(uniq) == 1:
                resolutions.append(RawResolution(label, expected, uniq[0], "fuzzy", f'Using legacy x->{rep_num} name matching'))
                continue

        resolutions.append(RawResolution(label, expected, None, "missing", 'No matching RAW file'))

    return resolutions


def targets_for_raw(
    plan: ExcelSheetPlan,
    raw_stem: str,
    *,
    internal_standard_formula: str = "",
) -> List[TargetSpec]:

    key = str(raw_stem or "").strip().lower()
    out: List[TargetSpec] = []
    for t in plan.targets:
        raw_row = {
            "sheet": plan.sheet_name,
            "combo": t.combo,
            "duplicate_index": str(t.duplicate_index),
            "duplicate_count": str(t.duplicate_count),
            "exact_mass": f"{t.exact_mass:.6f}",
            "is_internal_standard": "0",
        }
        out.append(
            TargetSpec(
                file_key=key,
                name=t.name,
                formula=t.formula,
                adduct="",
                polarity="",
                theoretical_mz=None,
                ppm=None,
                raw_row=raw_row,
            )
        )

    is_formula, is_mass = normalize_internal_standard_formula(internal_standard_formula)
    if is_formula:
        out.append(
            TargetSpec(
                file_key=key,
                name="Internal Standard",
                formula=is_formula,
                adduct="",
                polarity="",
                theoretical_mz=None,
                ppm=None,
                raw_row={
                    "sheet": plan.sheet_name,
                    "combo": "Internal Standard",
                    "duplicate_index": "1",
                    "duplicate_count": "1",
                    "exact_mass": f"{is_mass:.6f}",
                    "is_internal_standard": "1",
                },
            )
        )
    return out


def _sanitize_filename(name: str, max_len: int = 80) -> str:
    s = re.sub(r"[\\/:*?\"<>|]+", "_", str(name or ""))
    s = re.sub(r"\s+", "_", s).strip("._")
    return (s or "sheet")[:max_len]


def write_plan_csv(plan: ExcelSheetPlan, out_csv: Path, resolutions: Optional[Sequence[RawResolution]] = None) -> Path:
    out_csv = Path(out_csv)
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    raw_names = [r.raw_path.name if r.raw_path else r.expected_stem + ".RAW" for r in (resolutions or [])]
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "sheet",
                "No",
                "name",
                "formula",
                "exact_mass",
                "count",
                "duplicate_index",
                "duplicate_count",
                "combo",
                "raw_1",
                "raw_2",
                "raw_3",
            ]
        )
        for t in plan.targets:
            vals = raw_names[:3] + [""] * max(0, 3 - len(raw_names))
            w.writerow(
                [
                    plan.sheet_name,
                    t.no,
                    t.name,
                    t.formula,
                    f"{t.exact_mass:.6f}",
                    1,
                    t.duplicate_index,
                    t.duplicate_count,
                    t.combo,
                    vals[0],
                    vals[1],
                    vals[2],
                ]
            )
    return out_csv


def write_all_targets_csv(
    plans_and_raws: Sequence[Tuple[ExcelSheetPlan, Sequence[RawResolution]]],
    out_csv: Path,
    *,
    internal_standard_formula: str = "",
) -> Path:

    out_csv = Path(out_csv)
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    is_formula, is_mass = normalize_internal_standard_formula(internal_standard_formula)
    with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
        fields = [
            "raw",
            "name",
            "formula",
            "sheet",
            "replicate",
            "combo",
            "duplicate_index",
            "duplicate_count",
            "exact_mass",
            "is_internal_standard",
        ]
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for plan, resolutions in plans_and_raws:
            for rr in resolutions:
                raw_key = rr.raw_path.stem if rr.raw_path else rr.expected_stem
                for t in plan.targets:
                    w.writerow(
                        {
                            "raw": raw_key,
                            "name": t.name,
                            "formula": t.formula,
                            "sheet": plan.sheet_name,
                            "replicate": rr.replicate_label,
                            "combo": t.combo,
                            "duplicate_index": t.duplicate_index,
                            "duplicate_count": t.duplicate_count,
                            "exact_mass": f"{t.exact_mass:.6f}",
                            "is_internal_standard": 0,
                        }
                    )
                if is_formula:
                    w.writerow(
                        {
                            "raw": raw_key,
                            "name": "Internal Standard",
                            "formula": is_formula,
                            "sheet": plan.sheet_name,
                            "replicate": rr.replicate_label,
                            "combo": "Internal Standard",
                            "duplicate_index": 1,
                            "duplicate_count": 1,
                            "exact_mass": f"{is_mass:.6f}",
                            "is_internal_standard": 1,
                        }
                    )
    return out_csv


def read_quant_csv(path: Path) -> List[Dict[str, str]]:
    path = Path(path)
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return [{str(k): ("" if v is None else str(v)) for k, v in row.items()} for row in csv.DictReader(f)]


def _to_float(v: object) -> Optional[float]:
    try:
        if v is None or str(v).strip() == "":
            return None
        x = float(v)
        if not math.isfinite(x):
            return None
        return x
    except Exception:
        return None


def _to_bool(v: object) -> bool:
    return str(v or "").strip().lower() in {'1', 'true', 'yes', 'y', 'Yes', '是'}


def _mean_sd_rsd(values: Sequence[float]) -> Tuple[Optional[float], Optional[float], Optional[float]]:
    vals = [float(v) for v in values if v is not None and math.isfinite(float(v))]
    if not vals:
        return None, None, None
    mean = float(sum(vals) / len(vals))
    sd = float(statistics.stdev(vals)) if len(vals) >= 2 else 0.0
    rsd = float(sd / mean * 100.0) if mean != 0 else None
    return mean, sd, rsd


def _internal_standard_qc(
    rep_is_rows: Sequence[Dict[str, str]],
    *,
    min_rel_to_median: float = 0.10,
    max_rel_to_median: float = 10.0,
) -> Tuple[List[bool], List[str], Optional[float]]:
    """Check whether the internal-standard peak is usable for normalization.

    Self-check rationale:
    - The previous summary treated any found internal-standard area > 0 as valid.
      If an interference/noise peak was picked as the internal standard, a very small
      IS area could produce unrealistically large target/IS ratios.
    - We therefore flag obvious IS outliers within the three parallel RAW files.
      This does not modify the RAW-level XIC CSV; it controls whether that
      replicate's target/IS ratio is used in the triplicate summary.
    """

    areas: List[Optional[float]] = []
    valid_for_median: List[float] = []
    for row in rep_is_rows:
        found = _to_bool(row.get("Found", ""))
        area = _to_float(row.get("Area", ""))
        areas.append(area)
        if found and area is not None and area > 0:
            valid_for_median.append(float(area))

    median_area: Optional[float] = None
    if valid_for_median:
        median_area = float(statistics.median(valid_for_median))

    flags: List[bool] = []
    notes: List[str] = []
    n_valid = len(valid_for_median)
    for row, area in zip(rep_is_rows, areas):
        found = _to_bool(row.get("Found", ""))
        if not found:
            flags.append(False)
            notes.append("IS not found")
            continue
        if area is None or area <= 0:
            flags.append(False)
            notes.append("IS area missing/<=0")
            continue

        # With fewer than 2 valid IS peaks, a cross-replicate outlier check is not reliable.
        if median_area is None or median_area <= 0 or n_valid < 2:
            flags.append(True)
            notes.append("OK (no median QC)")
            continue

        rel = float(area) / float(median_area)
        if rel < float(min_rel_to_median):
            flags.append(False)
            notes.append(f"IS area too low ({rel:.3g}x median)")
        elif rel > float(max_rel_to_median):
            flags.append(False)
            notes.append(f"IS area too high ({rel:.3g}x median)")
        else:
            flags.append(True)
            notes.append(f"OK ({rel:.3g}x median)")

    return flags, notes, median_area


def write_parallel_summary_xlsx(
    out_xlsx: Path,
    run_records: Sequence[Dict[str, object]],
    *,
    internal_standard_formula: str = "",
    average_csv_path: Optional[Path] = None,
) -> Path:
    '``{"plan": ExcelSheetPlan, "resolutions": [...], "quant_rows": {rep_label: rows}, "status": ...}``'

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except Exception as e:
        raise RuntimeError('Summary export requires openpyxl. Install it in the active environment: python -m pip install openpyxl') from e

    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_run = wb.active
    ws_run.title = 'Run Summary'
    
    ws_avg = wb.create_sheet('Triplicate Average')
    ws_wide = wb.create_sheet('Triplicate XIC Summary')
    ws_long = wb.create_sheet('XIC Long Table')
    ws_is = wb.create_sheet('Internal Standard')

    is_formula, is_mass = normalize_internal_standard_formula(internal_standard_formula)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    run_headers = [
        "Sheet",
        'A Count',
        'B Count',
        'C Count',
        'Combination Count (Duplicates Kept)',
        'Invalid Combination Count',
        'Rep1 RAW',
        'Rep1 Status',
        'Rep2 RAW',
        'Rep2 Status',
        'Rep3 RAW',
        'Rep3 Status',
        'Common internal-standard formula',
        'Note',
    ]
    ws_run.append(run_headers)

    
    wide_headers = [
        "Sheet",
        "No",
        "Name",
        "Formula",
        "Exact_mass",
        "Duplicate_index",
        "Duplicate_count",
        "Combo",
    ]
    for i in range(1, 4):
        wide_headers.extend(
            [
                f"Rep{i}_RAW",
                f"Rep{i}_Found",
                f"Rep{i}_Adduct",
                f"Rep{i}_XIC_mz",
                f"Rep{i}_Apex_RT",
                f"Rep{i}_RT_start",
                f"Rep{i}_RT_end",
                f"Rep{i}_Peak_height",
                f"Rep{i}_Area",
                f"Rep{i}_Ratio_%",
                f"Rep{i}_IS_Area",
                f"Rep{i}_IS_Found",
                f"Rep{i}_IS_QC",
                f"Rep{i}_IS_QC_note",
            ]
        )
    wide_headers.extend([
        "Found_n",
        "Area_mean(found)",
        "Area_SD(found)",
        "Area_RSD_%(found)",
        "Apex_RT_mean(found)",
        "Ratio_mean(vs_IS)",
        "Ratio_SD(vs_IS)",
        "Ratio_RSD_%(vs_IS)",
    ])
    ws_wide.append(wide_headers)

    
    
    
    ratio_mean_header = 'Mean_Area_to_IS_Ratio_%' if is_formula else 'Mean_Percent_of_Total_Area_%'
    avg_headers = [
        "Sheet",
        "No",
        "Name",
        "Formula",
        "Exact_mass",
        "Duplicate_index",
        "Duplicate_count",
        "Combo",
    ]
    for i in range(1, 4):
        avg_headers.extend(
            [
                f"Rep{i}_Found",
                f"Rep{i}_Area",
                f"Rep{i}_Ratio_%",
                f"Rep{i}_Apex_RT",
                f"Rep{i}_IS_Found",
                f"Rep{i}_IS_QC",
                f"Rep{i}_IS_Area",
                f"Rep{i}_IS_Apex_RT",
            ]
        )
    avg_headers.extend(
        [
            'Valid_Area_Replicate_Count',
            'Valid_Ratio_Replicate_Count',
            'Triplicates_Complete',
            'Mean_Peak_Area',
            'Peak_Area_SD',
            'Peak_Area_RSD_%',
            ratio_mean_header,
            'Ratio_SD',
            'Ratio_RSD_%',
            'Mean_Apex_RT',
            "Observed_Fragility_Index",
            "Adduct_Cluster_Proneness_Index",
            "Primary_Ion_Fraction",
            "Ion_Form_Diversity_Count",
            "Accepted_Channel_Count",
            "Summed_Channel_Count",
            "Br_Fragment_Fraction",
            "Negative_Ion_Panel",
            "Accepted_Channels",
            'Final_Triplicate_Mean_Result',
            'Strict_Triplicate_Mean_Result_3of3',
            'Result_Type',
            'Valid_Result_Replicate_Count',
            'Averaging_Rule',
            'IS_QC_Notes',
        ]
    )
    ws_avg.append(avg_headers)
    average_csv_rows: List[List[object]] = []

    long_headers = [
        "Sheet",
        "Replicate",
        "RAW",
        "No",
        "Name",
        "Formula",
        "Exact_mass",
        "Duplicate_index",
        "Duplicate_count",
        "Combo",
        "Found",
        "Adduct",
        "XIC_mz",
        "Apex_RT",
        "RT_start",
        "RT_end",
        "Peak_height",
        "Area",
        "MH_Area",
        "Formate_Accepted",
        "Formate_Area",
        "Formate_Area_Fraction_%",
        "Formate_Apex_RT",
        "Formate_RT_Delta_min",
        "Formate_Shape_Correlation",
        "Adduct_Channels_Used",
        "Formate_Note",
        "Negative_Channel_Mode",
        "Negative_Ion_Panel",
        "Accepted_Channels",
        "Accepted_Channel_Count",
        "Summed_Channel_Count",
        "Additional_Summed_Area",
        "Ion_Form_Diversity_Count",
        "Observed_Fragility_Index",
        "Adduct_Cluster_Proneness_Index",
        "Primary_Ion_Fraction",
        "Br_Fragment_Fraction",
        "Ratio_%",
        "Ratio_mode",
        "Internal_standard",
        "Internal_area",
        "Internal_standard_found",
        "XIC_PNG",
    ]
    ws_long.append(long_headers)

    is_headers = [
        "Sheet",
        "Replicate",
        "RAW",
        "Internal_standard_formula",
        "Exact_mass",
        "Found",
        "Adduct",
        "XIC_mz",
        "Apex_RT",
        "RT_start",
        "RT_end",
        "Peak_height",
        "Area",
        "MH_Area",
        "Formate_Accepted",
        "Formate_Area",
        "Formate_Area_Fraction_%",
        "Adduct_Channels_Used",
        "Formate_Note",
        "Negative_Channel_Mode",
        "Negative_Ion_Panel",
        "Accepted_Channels",
        "Accepted_Channel_Count",
        "Summed_Channel_Count",
        "Additional_Summed_Area",
        "Ion_Form_Diversity_Count",
        "Observed_Fragility_Index",
        "Adduct_Cluster_Proneness_Index",
        "Primary_Ion_Fraction",
        "Br_Fragment_Fraction",
        "QC_valid",
        "QC_note",
        "Median_IS_area",
        "IS_Diagnostic_Status",
        "IS_Accepted_By",
        "IS_Dedicated_Adduct",
        "IS_PPM",
        "IS_Candidate_Apex_RT",
        "IS_Candidate_RT_Start",
        "IS_Candidate_RT_End",
        "IS_Candidate_Area",
        "IS_Candidate_Raw_Apex",
        "IS_Candidate_Baseline",
        "IS_Candidate_Net_Height",
        "IS_Noise_Sigma_Proxy",
        "IS_SNR_Proxy",
        "IS_Global_Min_Peak_Height",
        "IS_Adaptive_Min_Height",
        "IS_Expected_RT",
        "IS_Expected_RT_Tolerance_Min",
        "IS_Diagnostic_Note",
        "XIC_PNG",
        "IS_Diagnostic_PNG",
        "IS_XIC_Plot_In_Plots_Folder",
        "IS_Diagnostic_CSV",
        "IS_Trace_CSV",
        "IS_All_Channels_PNG",
    ]
    ws_is.append(is_headers)

    for rec in run_records:
        plan: ExcelSheetPlan = rec["plan"]  # type: ignore[assignment]
        resolutions: Sequence[RawResolution] = rec.get("resolutions", [])  # type: ignore[assignment]
        quant_rows: Dict[str, List[Dict[str, str]]] = rec.get("quant_rows", {})  # type: ignore[assignment]
        note = str(rec.get("note", "") or "")

        run_row: List[object] = [
            plan.sheet_name,
            len(plan.a_rows),
            len(plan.b_rows),
            len(plan.c_rows),
            len(plan.targets),
            int(plan.skipped_invalid),
        ]
        for i in range(3):
            rr = resolutions[i] if i < len(resolutions) else None
            if rr is None:
                run_row.extend(["", 'Not configured'])
            else:
                run_row.extend([
                    str(rr.raw_path) if rr.raw_path else rr.expected_stem + ".RAW",
                    rr.match_mode,
                ])
        run_row.append(is_formula)
        run_row.append(note)
        ws_run.append(run_row)

        
        rep_maps: List[Dict[int, Dict[str, str]]] = []
        rep_is_rows: List[Dict[str, str]] = []
        rep_raw_names: List[str] = []
        rep_labels: List[str] = []
        for i in range(3):
            rr = resolutions[i] if i < len(resolutions) else None
            label = rr.replicate_label if rr else str(i + 1)
            m: Dict[int, Dict[str, str]] = {}
            is_qr: Dict[str, str] = {}
            for qr in quant_rows.get(label, []):
                if _to_bool(qr.get("Is_internal_standard", "")):
                    is_qr = qr
                    
                    continue
                try:
                    no = int(float(qr.get("No", "")))
                except Exception:
                    continue
                m[no] = qr
            rep_maps.append(m)
            rep_is_rows.append(is_qr)
            rep_raw_names.append(rr.raw_path.name if rr and rr.raw_path else (rr.expected_stem + ".RAW" if rr else ""))
            rep_labels.append(label)

        is_qc_flags, is_qc_notes, is_area_median = _internal_standard_qc(rep_is_rows)

        
        
        for i in range(3):
            rr = resolutions[i] if i < len(resolutions) else None
            is_qr = rep_is_rows[i] if i < len(rep_is_rows) else {}
            ws_is.append(
                [
                    plan.sheet_name,
                    rep_labels[i] if i < len(rep_labels) else str(i + 1),
                    rep_raw_names[i] if i < len(rep_raw_names) else "",
                    is_formula,
                    float(is_mass) if is_formula and math.isfinite(is_mass) else None,
                    _to_bool(is_qr.get("Found", "")),
                    is_qr.get("Adduct", ""),
                    _to_float(is_qr.get("XIC_mz", "")),
                    _to_float(is_qr.get("Apex_RT", "")),
                    _to_float(is_qr.get("RT_start", "")),
                    _to_float(is_qr.get("RT_end", "")),
                    _to_float(is_qr.get("Peak_height", "")),
                    _to_float(is_qr.get("Area", "")),
                    _to_float(is_qr.get("MH_Area", "")),
                    _to_bool(is_qr.get("Formate_Accepted", "")),
                    _to_float(is_qr.get("Formate_Area", "")),
                    _to_float(is_qr.get("Formate_Area_Fraction_%", "")),
                    is_qr.get("Adduct_Channels_Used", ""),
                    is_qr.get("Formate_Note", ""),
                    is_qr.get("Negative_Channel_Mode", ""),
                    is_qr.get("Negative_Ion_Panel", ""),
                    is_qr.get("Accepted_Channels", ""),
                    _to_float(is_qr.get("Accepted_Channel_Count", "")),
                    _to_float(is_qr.get("Summed_Channel_Count", "")),
                    _to_float(is_qr.get("Additional_Summed_Area", "")),
                    _to_float(is_qr.get("Ion_Form_Diversity_Count", "")),
                    _to_float(is_qr.get("Observed_Fragility_Index", "")),
                    _to_float(is_qr.get("Adduct_Cluster_Proneness_Index", "")),
                    _to_float(is_qr.get("Primary_Ion_Fraction", "")),
                    _to_float(is_qr.get("Br_Fragment_Fraction", "")),
                    bool(is_qc_flags[i]) if i < len(is_qc_flags) else False,
                    is_qc_notes[i] if i < len(is_qc_notes) else "",
                    is_area_median,
                    is_qr.get("IS_Diagnostic_Status", ""),
                    is_qr.get("IS_Accepted_By", ""),
                    is_qr.get("IS_Dedicated_Adduct", ""),
                    _to_float(is_qr.get("IS_PPM", "")),
                    _to_float(is_qr.get("IS_Candidate_Apex_RT", "")),
                    _to_float(is_qr.get("IS_Candidate_RT_Start", "")),
                    _to_float(is_qr.get("IS_Candidate_RT_End", "")),
                    _to_float(is_qr.get("IS_Candidate_Area", "")),
                    _to_float(is_qr.get("IS_Candidate_Raw_Apex", "")),
                    _to_float(is_qr.get("IS_Candidate_Baseline", "")),
                    _to_float(is_qr.get("IS_Candidate_Net_Height", "")),
                    _to_float(is_qr.get("IS_Noise_Sigma_Proxy", "")),
                    _to_float(is_qr.get("IS_SNR_Proxy", "")),
                    _to_float(is_qr.get("IS_Global_Min_Peak_Height", "")),
                    _to_float(is_qr.get("IS_Adaptive_Min_Height", "")),
                    _to_float(is_qr.get("IS_Expected_RT", "")),
                    _to_float(is_qr.get("IS_Expected_RT_Tolerance_Min", "")),
                    is_qr.get("IS_Diagnostic_Note", ""),
                    is_qr.get("XIC_PNG", ""),
                    is_qr.get("IS_Diagnostic_PNG", ""),
                    is_qr.get("IS_XIC_Plot_In_Plots_Folder", ""),
                    is_qr.get("IS_Diagnostic_CSV", ""),
                    is_qr.get("IS_Trace_CSV", ""),
                    is_qr.get("IS_All_Channels_PNG", ""),
                ]
            )

        for t in plan.targets:
            row: List[object] = [
                plan.sheet_name,
                t.no,
                t.name,
                t.formula,
                float(t.exact_mass),
                t.duplicate_index,
                t.duplicate_count,
                t.combo,
            ]
            found_areas: List[float] = []
            found_rts: List[float] = []
            found_ratios: List[float] = []
            fragility_values: List[float] = []
            cluster_values: List[float] = []
            primary_fraction_values: List[float] = []
            diversity_values: List[float] = []
            accepted_count_values: List[float] = []
            summed_count_values: List[float] = []
            br_fragment_values: List[float] = []
            panel_names: List[str] = []
            accepted_channel_labels: List[str] = []
            avg_rep_values: List[object] = []

            for i in range(3):
                rr = resolutions[i] if i < len(resolutions) else None
                qr = rep_maps[i].get(t.no, {})
                found = _to_bool(qr.get("Found", ""))
                area = _to_float(qr.get("Area", ""))
                apex = _to_float(qr.get("Apex_RT", ""))
                ratio_val = _to_float(qr.get("Ratio_%", ""))
                is_area = _to_float(qr.get("Internal_area", ""))
                is_found = _to_bool(qr.get("Internal_standard_found", ""))
                
                
                
                is_qc_valid = bool(is_qc_flags[i]) if i < len(is_qc_flags) else bool(is_found and is_area is not None and is_area > 0)
                is_qc_note = is_qc_notes[i] if i < len(is_qc_notes) else ""
                is_apex = _to_float(rep_is_rows[i].get("Apex_RT", "")) if i < len(rep_is_rows) else None

                
                
                
                ratio_valid = bool(found and ratio_val is not None)
                if is_formula:
                    ratio_valid = bool(ratio_valid and is_qc_valid and is_area is not None and is_area > 0)
                ratio_for_average = ratio_val if ratio_valid else None

                if found and area is not None:
                    found_areas.append(area)
                if found and apex is not None:
                    found_rts.append(apex)
                if ratio_for_average is not None:
                    found_ratios.append(ratio_for_average)
                if found:
                    for source_key, dest in (
                        ("Observed_Fragility_Index", fragility_values),
                        ("Adduct_Cluster_Proneness_Index", cluster_values),
                        ("Primary_Ion_Fraction", primary_fraction_values),
                        ("Ion_Form_Diversity_Count", diversity_values),
                        ("Accepted_Channel_Count", accepted_count_values),
                        ("Summed_Channel_Count", summed_count_values),
                        ("Br_Fragment_Fraction", br_fragment_values),
                    ):
                        val = _to_float(qr.get(source_key, ""))
                        if val is not None:
                            dest.append(float(val))
                    panel_name = str(qr.get("Negative_Ion_Panel", "") or "").strip()
                    if panel_name and panel_name not in panel_names:
                        panel_names.append(panel_name)
                    channel_label = str(qr.get("Accepted_Channels", "") or "").strip()
                    if channel_label and channel_label not in accepted_channel_labels:
                        accepted_channel_labels.append(channel_label)

                avg_rep_values.extend(
                    [
                        found,
                        area if found else None,
                        ratio_for_average,
                        apex if found else None,
                        is_found,
                        is_qc_valid,
                        is_area if is_found else None,
                        is_apex if is_found else None,
                    ]
                )

                raw_name = rr.raw_path.name if rr and rr.raw_path else (rr.expected_stem + ".RAW" if rr else "")
                row.extend(
                    [
                        raw_name,
                        found,
                        qr.get("Adduct", ""),
                        _to_float(qr.get("XIC_mz", "")),
                        apex,
                        _to_float(qr.get("RT_start", "")),
                        _to_float(qr.get("RT_end", "")),
                        _to_float(qr.get("Peak_height", "")),
                        area,
                        ratio_for_average,
                        is_area,
                        is_found,
                        is_qc_valid,
                        is_qc_note,
                    ]
                )

                
                ws_long.append(
                    [
                        plan.sheet_name,
                        rr.replicate_label if rr else str(i + 1),
                        raw_name,
                        t.no,
                        t.name,
                        t.formula,
                        float(t.exact_mass),
                        t.duplicate_index,
                        t.duplicate_count,
                        t.combo,
                        found,
                        qr.get("Adduct", ""),
                        _to_float(qr.get("XIC_mz", "")),
                        apex,
                        _to_float(qr.get("RT_start", "")),
                        _to_float(qr.get("RT_end", "")),
                        _to_float(qr.get("Peak_height", "")),
                        area,
                        _to_float(qr.get("MH_Area", "")),
                        _to_bool(qr.get("Formate_Accepted", "")),
                        _to_float(qr.get("Formate_Area", "")),
                        _to_float(qr.get("Formate_Area_Fraction_%", "")),
                        _to_float(qr.get("Formate_Apex_RT", "")),
                        _to_float(qr.get("Formate_RT_Delta_min", "")),
                        _to_float(qr.get("Formate_Shape_Correlation", "")),
                        qr.get("Adduct_Channels_Used", ""),
                        qr.get("Formate_Note", ""),
                        qr.get("Negative_Channel_Mode", ""),
                        qr.get("Negative_Ion_Panel", ""),
                        qr.get("Accepted_Channels", ""),
                        _to_float(qr.get("Accepted_Channel_Count", "")),
                        _to_float(qr.get("Summed_Channel_Count", "")),
                        _to_float(qr.get("Additional_Summed_Area", "")),
                        _to_float(qr.get("Ion_Form_Diversity_Count", "")),
                        _to_float(qr.get("Observed_Fragility_Index", "")),
                        _to_float(qr.get("Adduct_Cluster_Proneness_Index", "")),
                        _to_float(qr.get("Primary_Ion_Fraction", "")),
                        _to_float(qr.get("Br_Fragment_Fraction", "")),
                        ratio_val,
                        qr.get("Ratio_mode", ""),
                        qr.get("Internal_standard", ""),
                        is_area,
                        is_found,
                        qr.get("XIC_PNG", ""),
                    ]
                )

            area_mean, area_sd, area_rsd = _mean_sd_rsd(found_areas)
            rt_mean, _, _ = _mean_sd_rsd(found_rts)
            ratio_mean, ratio_sd, ratio_rsd = _mean_sd_rsd(found_ratios)
            fragility_mean, _, _ = _mean_sd_rsd(fragility_values)
            cluster_mean, _, _ = _mean_sd_rsd(cluster_values)
            primary_fraction_mean, _, _ = _mean_sd_rsd(primary_fraction_values)
            diversity_mean, _, _ = _mean_sd_rsd(diversity_values)
            accepted_count_mean, _, _ = _mean_sd_rsd(accepted_count_values)
            summed_count_mean, _, _ = _mean_sd_rsd(summed_count_values)
            br_fragment_mean, _, _ = _mean_sd_rsd(br_fragment_values)
            row.extend([
                len(found_areas),
                area_mean,
                area_sd,
                area_rsd,
                rt_mean,
                ratio_mean,
                ratio_sd,
                ratio_rsd,
            ])
            ws_wide.append(row)

            
            
            
            use_is_result = bool(is_formula)
            result_mean = ratio_mean if use_is_result else area_mean
            result_type = 'Area / Internal Standard Area (%)' if use_is_result else 'Peak Area'
            result_valid_n = len(found_ratios) if use_is_result else len(found_areas)
            complete_triplicate = result_valid_n == 3
            strict_triplicate_mean = result_mean if complete_triplicate else None

            avg_row: List[object] = [
                plan.sheet_name,
                t.no,
                t.name,
                t.formula,
                float(t.exact_mass),
                t.duplicate_index,
                t.duplicate_count,
                t.combo,
            ]
            avg_row.extend(avg_rep_values)
            avg_row.extend(
                [
                    len(found_areas),
                    len(found_ratios),
                    'Yes' if complete_triplicate else 'No',
                    area_mean,
                    area_sd,
                    area_rsd,
                    ratio_mean,
                    ratio_sd,
                    ratio_rsd,
                    rt_mean,
                    fragility_mean,
                    cluster_mean,
                    primary_fraction_mean,
                    diversity_mean,
                    accepted_count_mean,
                    summed_count_mean,
                    br_fragment_mean,
                    "; ".join(panel_names),
                    " || ".join(accepted_channel_labels),
                    result_mean,
                    strict_triplicate_mean,
                    result_type,
                    result_valid_n,
                    'The mean of valid replicates includes only Found=True records with valid numerical values. In common internal-standard mode, the internal standard must also have Found=True, area >0 and pass IS QC. The strict triplicate mean is reported only when all three replicates are valid; missing values are not treated as zero.',
                    "; ".join(f"Rep{i+1}: {is_qc_notes[i]}" for i in range(min(3, len(is_qc_notes)))) if is_formula else "",
                ]
            )
            ws_avg.append(avg_row)
            average_csv_rows.append(avg_row)

        
        
        
        if is_formula:
            is_area_values: List[float] = []
            is_area_values_qc: List[float] = []
            is_rt_values: List[float] = []
            is_avg_rep_values: List[object] = []
            is_wide_row: List[object] = [
                plan.sheet_name,
                "IS",
                "Internal Standard",
                is_formula,
                float(is_mass) if math.isfinite(is_mass) else None,
                1,
                1,
                "Internal Standard",
            ]

            for i in range(3):
                rr = resolutions[i] if i < len(resolutions) else None
                is_qr = rep_is_rows[i] if i < len(rep_is_rows) else {}
                is_found = _to_bool(is_qr.get("Found", ""))
                is_area = _to_float(is_qr.get("Area", ""))
                is_apex = _to_float(is_qr.get("Apex_RT", ""))
                is_qc_valid = bool(is_qc_flags[i]) if i < len(is_qc_flags) else False
                is_qc_note = is_qc_notes[i] if i < len(is_qc_notes) else ""
                raw_name = rep_raw_names[i] if i < len(rep_raw_names) else (rr.expected_stem + ".RAW" if rr else "")

                if is_found and is_area is not None and is_area > 0:
                    is_area_values.append(float(is_area))
                if is_qc_valid and is_area is not None and is_area > 0:
                    is_area_values_qc.append(float(is_area))
                if is_found and is_apex is not None:
                    is_rt_values.append(float(is_apex))

                
                
                is_ratio = 100.0 if is_qc_valid else None
                is_avg_rep_values.extend([is_found, is_area if is_found else None, is_ratio, is_apex if is_found else None, is_found, is_qc_valid, is_area if is_found else None, is_apex if is_found else None])

                is_wide_row.extend(
                    [
                        raw_name,
                        is_found,
                        is_qr.get("Adduct", ""),
                        _to_float(is_qr.get("XIC_mz", "")),
                        is_apex,
                        _to_float(is_qr.get("RT_start", "")),
                        _to_float(is_qr.get("RT_end", "")),
                        _to_float(is_qr.get("Peak_height", "")),
                        is_area,
                        is_ratio,
                        is_area,
                        is_found,
                        is_qc_valid,
                        is_qc_note,
                    ]
                )

                
                ws_long.append(
                    [
                        plan.sheet_name,
                        rep_labels[i] if i < len(rep_labels) else str(i + 1),
                        raw_name,
                        "IS",
                        "Internal Standard",
                        is_formula,
                        float(is_mass) if math.isfinite(is_mass) else None,
                        1,
                        1,
                        "Internal Standard",
                        is_found,
                        is_qr.get("Adduct", ""),
                        _to_float(is_qr.get("XIC_mz", "")),
                        is_apex,
                        _to_float(is_qr.get("RT_start", "")),
                        _to_float(is_qr.get("RT_end", "")),
                        _to_float(is_qr.get("Peak_height", "")),
                        is_area,
                        _to_float(is_qr.get("MH_Area", "")),
                        _to_bool(is_qr.get("Formate_Accepted", "")),
                        _to_float(is_qr.get("Formate_Area", "")),
                        _to_float(is_qr.get("Formate_Area_Fraction_%", "")),
                        _to_float(is_qr.get("Formate_Apex_RT", "")),
                        _to_float(is_qr.get("Formate_RT_Delta_min", "")),
                        _to_float(is_qr.get("Formate_Shape_Correlation", "")),
                        is_qr.get("Adduct_Channels_Used", ""),
                        is_qr.get("Formate_Note", ""),
                        is_qr.get("Negative_Channel_Mode", ""),
                        is_qr.get("Negative_Ion_Panel", ""),
                        is_qr.get("Accepted_Channels", ""),
                        _to_float(is_qr.get("Accepted_Channel_Count", "")),
                        _to_float(is_qr.get("Summed_Channel_Count", "")),
                        _to_float(is_qr.get("Additional_Summed_Area", "")),
                        _to_float(is_qr.get("Ion_Form_Diversity_Count", "")),
                        _to_float(is_qr.get("Observed_Fragility_Index", "")),
                        _to_float(is_qr.get("Adduct_Cluster_Proneness_Index", "")),
                        _to_float(is_qr.get("Primary_Ion_Fraction", "")),
                        _to_float(is_qr.get("Br_Fragment_Fraction", "")),
                        is_ratio,
                        "internal_standard_self_check",
                        "Internal Standard",
                        is_area,
                        is_found,
                        is_qr.get("XIC_PNG", ""),
                    ]
                )

            is_area_mean, is_area_sd, is_area_rsd = _mean_sd_rsd(is_area_values)
            is_area_qc_mean, is_area_qc_sd, is_area_qc_rsd = _mean_sd_rsd(is_area_values_qc)
            is_rt_mean, _, _ = _mean_sd_rsd(is_rt_values)
            is_complete = len(is_area_values_qc) == 3

            is_wide_row.extend(
                [
                    len(is_area_values),
                    is_area_mean,
                    is_area_sd,
                    is_area_rsd,
                    is_rt_mean,
                    100.0 if is_area_values_qc else None,
                    0.0 if is_area_values_qc else None,
                    0.0 if is_area_values_qc else None,
                ]
            )
            ws_wide.append(is_wide_row)

            is_avg_row: List[object] = [
                plan.sheet_name,
                "IS",
                "Internal Standard",
                is_formula,
                float(is_mass) if math.isfinite(is_mass) else None,
                1,
                1,
                "Internal Standard",
            ]
            is_avg_row.extend(is_avg_rep_values)
            is_avg_row.extend(
                [
                    len(is_area_values),
                    len(is_area_values_qc),
                    'Yes' if is_complete else 'No',
                    is_area_mean,
                    is_area_sd,
                    is_area_rsd,
                    100.0 if is_area_values_qc else None,
                    0.0 if is_area_values_qc else None,
                    0.0 if is_area_values_qc else None,
                    is_rt_mean,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "; ".join(sorted({str(r.get("Negative_Ion_Panel", "") or "") for r in rep_is_rows if str(r.get("Negative_Ion_Panel", "") or "")})),
                    " || ".join(sorted({str(r.get("Accepted_Channels", "") or "") for r in rep_is_rows if str(r.get("Accepted_Channels", "") or "")})),
                    is_area_qc_mean if is_area_qc_mean is not None else is_area_mean,
                    is_area_qc_mean if is_complete else None,
                    'Internal-standard peak area (QC reference)',
                    len(is_area_values_qc),
                    'Internal-standard row: peak areas, RT and QC for the common internal standard across three replicates. Target ratios are target Area / IS Area x 100%. Replicates with invalid internal-standard results are excluded from the mean target-to-IS ratio.',
                    "; ".join(f"Rep{i+1}: {is_qc_notes[i]}" for i in range(min(3, len(is_qc_notes)))) if is_formula else "",
                ]
            )
            ws_avg.append(is_avg_row)
            average_csv_rows.append(is_avg_row)

    for ws in (ws_run, ws_avg, ws_wide, ws_long, ws_is):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=False)

    
    for ws in (ws_avg, ws_wide, ws_long, ws_is):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                h = ws.cell(row=1, column=cell.column).value
                hs = str(h or "")
                if "mz" in hs.lower() or "Exact_mass" in hs:
                    cell.number_format = "0.000000"
                elif any(k in hs for k in ("Apex_RT", "RT_start", "RT_end")):
                    cell.number_format = "0.000"
                elif "Area" in hs or "Peak_height" in hs or 'Peak Area' in hs:
                    cell.number_format = "0.000"
                elif "Ratio_%" in hs or "RSD_%" in hs or 'ratio' in hs.lower():
                    cell.number_format = "0.00"

    
    width_caps = {
        "Combo": 55,
        'Note': 40,
        "XIC_PNG": 45,
        "IS_Diagnostic_PNG": 45,
        "IS_XIC_Plot_In_Plots_Folder": 45,
        "IS_Diagnostic_CSV": 45,
        "IS_Trace_CSV": 45,
        "IS_All_Channels_PNG": 45,
        "IS_Diagnostic_Note": 52,
        "IS_Diagnostic_Status": 22,
        "IS_Accepted_By": 24,
    }
    for ws in (ws_run, ws_avg, ws_wide, ws_long, ws_is):
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(row=1, column=col_idx).value or "")
            max_len = len(header)
            for r in range(2, min(ws.max_row, 300) + 1):
                v = ws.cell(row=r, column=col_idx).value
                if v is not None:
                    max_len = max(max_len, len(str(v)))
            cap = width_caps.get(header, 28)
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(cap, max_len + 2))

    
    avg_header_to_col = {
        str(ws_avg.cell(row=1, column=c).value or ""): c
        for c in range(1, ws_avg.max_column + 1)
    }
    final_fill = PatternFill("solid", fgColor="E2F0D9")
    final_header_fill = PatternFill("solid", fgColor="548235")
    for header in (
        'Triplicates_Complete',
        'Final_Triplicate_Mean_Result',
        'Strict_Triplicate_Mean_Result_3of3',
        'Result_Type',
        'Valid_Result_Replicate_Count',
    ):
        col_idx = avg_header_to_col.get(header)
        if not col_idx:
            continue
        hcell = ws_avg.cell(row=1, column=col_idx)
        hcell.fill = final_header_fill
        hcell.font = header_font
        for row_idx in range(2, ws_avg.max_row + 1):
            cell = ws_avg.cell(row=row_idx, column=col_idx)
            cell.fill = final_fill
            if header in ('Final_Triplicate_Mean_Result', 'Strict_Triplicate_Mean_Result_3of3', '最终三平行平均结果', 'final_triplicate_mean_result', '严格三平行平均结果(3/3)', 'strict_triplicate_mean_result_3of3'):
                cell.font = Font(bold=True)
                cell.number_format = "0.0000"

    
    wb.active = wb.index(ws_avg)
    wb.save(str(out_xlsx))

    if average_csv_path is not None:
        average_csv_path = Path(average_csv_path)
        average_csv_path.parent.mkdir(parents=True, exist_ok=True)
        with average_csv_path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(avg_headers)
            w.writerows(average_csv_rows)

    return out_xlsx


def safe_sheet_folder_name(sheet_name: str) -> str:
    return _sanitize_filename(sheet_name)
