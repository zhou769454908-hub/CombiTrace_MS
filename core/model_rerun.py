"""Model-only rerun from an existing ESI descriptor workbook.

This module allows slow 2D/3D descriptor generation to be reused.  The hidden
``ESI_Features`` sheet is treated as an immutable descriptor cache; only model,
QC, ranking and visualization sheets are rebuilt.
"""
from __future__ import annotations

import math
import shutil
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _num(value: object) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def _norm_group(value: object) -> str:
    return " ".join(_text(value).lower().split())


def parse_excluded_groups(values: Sequence[str] | str | None) -> List[str]:
    if values is None:
        return []
    if isinstance(values, str):
        raw = values.replace("\n", ",").replace(";", ",")
        parts = raw.split(",")
    else:
        parts = []
        for value in values:
            parts.extend(str(value).replace("\n", ",").replace(";", ",").split(","))
    out: List[str] = []
    seen = set()
    for value in parts:
        item = _text(value)
        key = _norm_group(item)
        if item and key not in seen:
            seen.add(key)
            out.append(item)
    return out


def load_descriptor_records(path: Path) -> Tuple[List[Dict[str, object]], List[Dict[str, object]], List[Dict[str, object]], List[str], List[str]]:
    """Read records and descriptor names from a previous ESI output workbook."""
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for model-only rerun") from exc
    wb = load_workbook(str(path), read_only=True, data_only=True)
    if "ESI_Features" not in wb.sheetnames:
        wb.close()
        raise ValueError("The selected workbook does not contain the hidden ESI_Features sheet")
    ws = wb["ESI_Features"]
    row_iter = ws.iter_rows(values_only=True)
    try:
        header_values = next(row_iter)
    except StopIteration:
        wb.close()
        raise ValueError("ESI_Features is empty")
    headers = [str(v or "").strip() for v in header_values]
    if "Dataset" not in headers:
        wb.close()
        raise ValueError("ESI_Features does not contain the Dataset column")
    try:
        start = headers.index("Product_Master_Formula_Key") + 1
    except ValueError:
        start = 0
    try:
        end = headers.index("Warnings")
    except ValueError:
        end = len(headers)
    descriptor_names = [h for h in headers[start:end] if h]
    master: List[Dict[str, object]] = []
    training: List[Dict[str, object]] = []
    target: List[Dict[str, object]] = []
    for values in row_iter:
        record = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers)) if headers[i]}
        dataset = _text(record.get("Dataset")).lower()
        if dataset == "master":
            master.append(record)
        elif dataset == "training":
            training.append(record)
        elif dataset == "target":
            target.append(record)
    wb.close()
    warnings: List[str] = []
    if not training:
        warnings.append("No Training rows were found in ESI_Features")
    if not target:
        warnings.append("No Target rows were found in ESI_Features")
    return master, training, target, descriptor_names, warnings


_MODEL_SHEET_EXACT = {
    "Diagnostics", "Model_Decision", "Model_Comparison", "Model_CV_Predictions",
    "Best_Model_Predictions", "Response_Corrected_Targets", "Model_Diagnostic_Plots",
    "Accuracy_Bands", "Raw_vs_Corrected", "Feature_Importance", "Hyperparameter_Tuning",
    "Calibration_Row_Check", "Model_Target_Row_Check", "Selected_Descriptors",
    "Excluded_Descriptors", "Unavailable_Descriptors", "Descriptor_Selection",
    "Descriptor_Status", "Model_Input_Audit", "Injection_Group_Out", "Injection_Group_Summary",
    "Outlier_Audit", "Outlier_QC_Summary", "Model_Comparison_Before_QC",
    "Trend_Optimizer_Decision", "Trend_Model_Comparison", "Trend_CV_Predictions",
    "Trend_Targets", "Trend_Feature_Stability", "Trend_Feature_Groups",
    "Trend_Random_Trials", "Trend_Permutation", "Trend_Final_Configs", "Trend_Settings",
    "Trend_Optimizer_Plots", "Published_IE_CV", "Published_IE_Targets", "Published_IE_Eluent",
    "Published_IE_Features", "External_logIE_Audit", "Dual_Model_Decision",
    "Dual_Model_Comparison", "Dual_Model_Targets", "Dual_Model_Methods", "Dual_Model_Plots",
    "NegESI_Decision", "Negative_ESI_Features", "Dynamic_Range_Comparison", "Dynamic_Range_CV",
    "Dynamic_Range_Targets", "High_Range_Bias", "Literature_Context", "Literature_vs_Current",
    "MS2Quant_Environment", "MS2Quant_IE_Audit", "MS2Quant_QEHF_Transfer",
    "Group_Exclusion_Audit", "Group_Exclusion_Impact", "Model_Rerun_Diagnostics", "Concentration_Ratio_Comparison",
    "Fixed_Search_Decision", "Fixed_Trial_All", "Fixed_Top5_Absolute", "Fixed_Top5_Trend",
    "Fixed_Top5_Absolute_CV", "Fixed_Top5_Trend_CV", "Fixed_Top5_Targets",
    "Fixed_Feature_Effects", "Manual_Feature_Selection", "Fixed_Search_Settings", "Fixed_Search_Plots",
    "Level_Model_Decision", "Level_Model_Comparison", "Level_Known_Standards",
    "Level_Target_Predictions", "Level_Confusion_Matrix", "Level_Definitions",
    "Level_Group_Validation", "Level_Features", "Level_Plots",
    "Known_Standards_Predictions", "Unknown_Sample_Predictions", "Simple_Result_Plots",
}


def load_previous_metrics(path: Path) -> Dict[str, object]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception:
        return {}
    metrics: Dict[str, object] = {}
    for sheet_name in ("Model_Rerun_Diagnostics", "Diagnostics", "Descriptor_Run_Diagnostics"):
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for values in ws.iter_rows(values_only=True):
            if not values:
                continue
            key = _text(values[0])
            if key and len(values) > 1 and key not in metrics:
                metrics[key] = values[1]
    if "Trend_Optimizer_Decision" in wb.sheetnames:
        ws = wb["Trend_Optimizer_Decision"]
        for values in ws.iter_rows(values_only=True):
            if values and len(values) > 1 and _text(values[0]):
                metrics["Trend_" + _text(values[0])] = values[1]
    if "Trend_Model_Comparison" in wb.sheetnames:
        ws = wb["Trend_Model_Comparison"]
        values = list(ws.iter_rows(values_only=True))
        if values:
            headers = [_text(x) for x in values[0]]
            rows = [dict(zip(headers, row)) for row in values[1:] if row]
            best_name = _text(metrics.get("Trend_Best_model"))
            best_row = next((r for r in rows if _text(r.get("Model")) == best_name), None)
            if best_row is None and rows:
                best_row = sorted(rows, key=lambda r: float(_num(r.get("Rank")) or 999.0))[0]
            if best_row:
                for key in ("Model", "Trend_Spearman_r", "Delta_Spearman_vs_raw", "Trend_Pairwise_Concordance_pct", "Trend_Top20_Overlap_pct", "Median_Fold_Error", "P80_Fold_Error"):
                    metrics["TrendComparison_" + key] = best_row.get(key, "")
    wb.close()
    return metrics


def _remove_stale_model_sheets(wb) -> None:
    prefixes = (
        "Trend_", "Published_", "Dual_", "NegESI_", "Dynamic_", "MS2Quant_",
        "Model_", "Outlier_", "Injection_Group_", "Literature_", "Fixed_", "Manual_Feature_", "Level_",
    )
    for name in list(wb.sheetnames):
        if name in _MODEL_SHEET_EXACT or name.startswith(prefixes):
            del wb[name]


def _append_dict_sheet(wb, title: str, rows: Sequence[Dict[str, object]]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    if rows:
        headers: List[str] = []
        seen = set()
        for row in rows:
            for key in row.keys():
                if key not in seen:
                    seen.add(key); headers.append(str(key))
    else:
        headers = ["Message"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _filter_groups(records: Sequence[Dict[str, object]], excluded_groups: Sequence[str]) -> Tuple[List[Dict[str, object]], List[Dict[str, object]]]:
    excluded_map = {_norm_group(x): x for x in excluded_groups if _norm_group(x)}
    kept: List[Dict[str, object]] = []
    audit: List[Dict[str, object]] = []
    counts: Dict[str, int] = {}
    for record in records:
        group = _text(record.get("Injection_Group"))
        counts[group] = counts.get(group, 0) + 1
        key = _norm_group(group)
        is_excluded = bool(key and key in excluded_map)
        audit.append({
            "Record_type": "Row",
            "Injection_Group": group,
            "Excluded": is_excluded,
            "Reason": "manual group exclusion" if is_excluded else "retained",
            "Source_file": record.get("Source_file", ""),
            "Source_sheet": record.get("Source_sheet", ""),
            "Source_row": record.get("Source_row", ""),
            "Combo": record.get("Combo", ""),
            "Actual_concentration": record.get("Actual_concentration", ""),
            "Measured_ratio": record.get("Measured_ratio", ""),
        })
        if not is_excluded:
            kept.append(dict(record))
    summary = [{
        "Record_type": "Group summary",
        "Injection_Group": group,
        "Rows": count,
        "Excluded": _norm_group(group) in excluded_map,
        "Reason": "manual group exclusion" if _norm_group(group) in excluded_map else "retained",
    } for group, count in sorted(counts.items(), key=lambda x: str(x[0]))]
    return kept, summary + audit


def rerun_models_from_descriptor_workbook(
    source_xlsx: Path,
    out_xlsx: Path,
    *,
    output_language: str,
    model_objective: str,
    model_cv_splits: int,
    model_cv_repeats: int,
    model_max_features: int,
    model_random_state: int,
    model_auto_remove_outliers: bool,
    model_outlier_mode: str,
    model_outlier_min_fold_error: float,
    model_outlier_max_fraction_pct: float,
    model_outlier_consensus_pct: float,
    model_auto_select_features: bool,
    model_use_categorical_features: bool,
    model_deep_tuning: bool,
    model_feature_stability_threshold_pct: float,
    model_min_selected_features: int,
    model_correlation_threshold: float,
    trend_optimizer_enabled: bool,
    trend_optimizer_trials: int,
    trend_optimizer_pair_min_fold: float,
    trend_optimizer_min_features: int,
    trend_optimizer_max_features: int,
    trend_optimizer_permutations: int,
    fixed_feature_search_enabled: bool,
    fixed_feature_model: str,
    fixed_feature_previous_best_model: str,
    fixed_feature_trials: int,
    fixed_feature_min_features: int,
    fixed_feature_max_features: int,
    manual_feature_pool: Sequence[str],
    level_classification_enabled: bool,
    dual_ie_enabled: bool,
    published_config,
    excluded_injection_groups: Sequence[str] = (),
):
    """Reuse a previous descriptor workbook and rebuild only model outputs."""
    from openpyxl import load_workbook
    from .esi_descriptors import EsiDescriptorReport
    from .esi_model_benchmark import run_model_benchmark, append_benchmark_to_workbook

    source_xlsx = Path(source_xlsx)
    out_xlsx = Path(out_xlsx)
    if not source_xlsx.exists():
        raise FileNotFoundError(source_xlsx)
    master, training_all, targets, descriptor_names, warnings = load_descriptor_records(source_xlsx)
    from .training_limits import validate_training_count
    validate_training_count(training_all, label="Cached training table")
    previous_metrics = load_previous_metrics(source_xlsx)
    excluded = parse_excluded_groups(excluded_injection_groups)
    training, group_audit = _filter_groups(training_all, excluded)
    available_group_keys = {_norm_group(r.get("Injection_Group")) for r in training_all if _norm_group(r.get("Injection_Group"))}
    missing_groups = [g for g in excluded if _norm_group(g) not in available_group_keys]
    if excluded:
        warnings.append(
            "Model-only rerun excluded standard injection group(s): " + ", ".join(g for g in excluded if g not in missing_groups)
        )
    if missing_groups:
        warnings.append("Requested exclusion group(s) were not found: " + ", ".join(missing_groups))
    if len(training) < 30:
        raise ValueError(f"Too few calibration rows after group exclusion: {len(training)}")

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    if source_xlsx.resolve() == out_xlsx.resolve():
        temp = out_xlsx.with_name(out_xlsx.stem + "__rerun_tmp.xlsx")
        shutil.copy2(source_xlsx, temp)
        work_path = temp
    else:
        shutil.copy2(source_xlsx, out_xlsx)
        work_path = out_xlsx
    wb = load_workbook(str(work_path))
    _remove_stale_model_sheets(wb)
    if "Diagnostics" in wb.sheetnames:
        wb["Diagnostics"].title = "Descriptor_Run_Diagnostics"

    benchmark = run_model_benchmark(
        training, targets, descriptor_names, out_xlsx,
        calibration_mode="training_original",
        model_objective=str(model_objective or "trend"),
        cv_splits=max(2, int(model_cv_splits)),
        cv_repeats=max(1, int(model_cv_repeats)),
        max_features=max(3, int(model_max_features)),
        random_state=int(model_random_state),
        expected_descriptor_names=(),
        three_d_enabled=any(_text(r.get("3D_status")) not in {"", "not_requested"} for r in master),
        output_language=str(output_language or "en"),
        auto_remove_outliers=bool(model_auto_remove_outliers),
        outlier_mode=str(model_outlier_mode or "conservative"),
        outlier_min_fold_error=max(1.01, float(model_outlier_min_fold_error)),
        outlier_max_fraction_pct=max(0.0, float(model_outlier_max_fraction_pct)),
        outlier_consensus_pct=max(0.0, min(100.0, float(model_outlier_consensus_pct))),
        auto_select_features=bool(model_auto_select_features),
        use_categorical_features=bool(model_use_categorical_features),
        deep_tuning=bool(model_deep_tuning),
        feature_stability_threshold_pct=float(model_feature_stability_threshold_pct),
        min_selected_features=max(2, int(model_min_selected_features)),
        correlation_threshold=float(model_correlation_threshold),
    )
    append_benchmark_to_workbook(wb, benchmark)

    trend = None
    if trend_optimizer_enabled:
        from .trend_rank_optimizer import run_trend_optimizer, append_trend_optimizer_to_workbook
        trend = run_trend_optimizer(
            benchmark.calibration_records,
            targets,
            descriptor_names,
            out_xlsx,
            trials=max(3, int(trend_optimizer_trials)),
            pair_min_fold=max(1.01, float(trend_optimizer_pair_min_fold)),
            min_features=max(2, int(trend_optimizer_min_features)),
            max_features=max(3, int(trend_optimizer_max_features)),
            cv_splits=max(2, int(model_cv_splits)),
            cv_repeats=max(1, int(model_cv_repeats)),
            permutations=max(20, int(trend_optimizer_permutations)),
            random_state=int(model_random_state),
            correlation_threshold=float(model_correlation_threshold),
            output_language=str(output_language or "en"),
            concentration_unit=str(getattr(published_config, "concentration_unit", "") if published_config else ""),
        )
        append_trend_optimizer_to_workbook(wb, trend)

    fixed_result = None
    if fixed_feature_search_enabled:
        try:
            from .fixed_feature_search import run_fixed_feature_search, append_fixed_feature_search_to_workbook
            previous_best = str(fixed_feature_previous_best_model or "").strip()
            if not previous_best:
                previous_best = (
                    (trend.best_model if trend is not None else "")
                    or _text(previous_metrics.get("Trend_Best_model"))
                    or _text(previous_metrics.get("Best_model"))
                    or benchmark.best_model
                )
            fixed_result = run_fixed_feature_search(
                benchmark.calibration_records,
                targets,
                descriptor_names,
                out_xlsx,
                fixed_model=str(fixed_feature_model or "auto_previous"),
                previous_best_model=previous_best,
                manual_feature_pool=tuple(manual_feature_pool or ()),
                trial_count=max(5, int(fixed_feature_trials)),
                min_features=max(1, int(fixed_feature_min_features)),
                max_features=max(1, int(fixed_feature_max_features)),
                cv_splits=max(2, int(model_cv_splits)),
                cv_repeats=max(1, int(model_cv_repeats)),
                random_state=int(model_random_state),
                correlation_threshold=float(model_correlation_threshold),
                output_language=str(output_language or "en"),
                concentration_unit=str(getattr(published_config, "concentration_unit", "") if published_config else ""),
            )
            append_fixed_feature_search_to_workbook(wb, fixed_result)
        except Exception as exc:
            warnings.append(f"Fixed-model feature search did not complete: {exc}")

    level_result = None
    if level_classification_enabled:
        try:
            from .concentration_classification import (
                run_concentration_level_classification,
                append_concentration_level_to_workbook,
            )
            selected_names = [
                str(row.get("Feature", ""))
                for row in benchmark.selected_descriptors
                if str(row.get("Feature", "")).strip()
            ]
            level_result = run_concentration_level_classification(
                benchmark.calibration_records,
                targets,
                descriptor_names,
                out_xlsx,
                selected_feature_names=selected_names,
                manual_feature_pool=tuple(manual_feature_pool or ()),
                cv_splits=max(2, int(model_cv_splits)),
                cv_repeats=max(1, int(model_cv_repeats)),
                max_features=max(3, min(30, int(model_max_features))),
                random_state=int(model_random_state),
                level_count=10,
                output_language=str(output_language or "en"),
                concentration_unit=str(getattr(published_config, "concentration_unit", "") if published_config else ""),
            )
            append_concentration_level_to_workbook(wb, level_result)
        except Exception as exc:
            warnings.append(f"Ten-level concentration classification did not complete: {exc}")

    published_result = None
    dual_result = None
    if dual_ie_enabled and published_config is not None:
        from .published_ie_benchmark import (
            run_published_ie_benchmark, run_dual_ie_validation, append_published_ie_to_workbook,
        )
        valid_target_keys = {str(r.get("ABC_Formula_Key", "")) for r in benchmark.target_predictions}
        published_targets = [
            r for r in targets
            if str(r.get("Product_Master_Formula_Key", "")) in valid_target_keys
            and _num(r.get("Measured_ratio")) is not None and float(r.get("Measured_ratio")) > 0
        ]
        published_result = run_published_ie_benchmark(
            benchmark.calibration_records, published_targets, descriptor_names, out_xlsx, published_config,
        )
        dual_result = run_dual_ie_validation(
            benchmark, published_result, out_xlsx, output_language=str(output_language or "en"),
        )
        append_published_ie_to_workbook(wb, published_result, dual_result)

    try:
        from .simple_prediction_outputs import build_simple_prediction_rows, append_simple_outputs_to_workbook
        simple_known_rows, simple_target_rows = build_simple_prediction_rows(
            benchmark,
            trend,
            level_result,
            targets,
            concentration_unit=str(getattr(published_config, "concentration_unit", "") if published_config else ""),
        )
        append_simple_outputs_to_workbook(
            wb, simple_known_rows, simple_target_rows, out_xlsx,
            output_language=str(output_language or "en"),
        )
    except Exception as exc:
        warnings.append(f"Concise prediction tables/plot did not complete: {exc}")

    _append_dict_sheet(wb, "Group_Exclusion_Audit", group_audit)
    impact_rows = [
        {"Metric": "Excluded_injection_groups", "Before": "", "After": "; ".join(excluded)},
        {"Metric": "Calibration_rows", "Before": previous_metrics.get("Model_calibration_rows", len(training_all)), "After": benchmark.n_calibration},
        {"Metric": "Corrected_Spearman", "Before": previous_metrics.get("Corrected_Spearman", ""), "After": benchmark.best_metrics.get("Trend_Spearman_r", "")},
        {"Metric": "Median_Fold_Error", "Before": previous_metrics.get("Best_Median_Fold_Error", ""), "After": benchmark.best_metrics.get("Median_Fold_Error", "")},
        {"Metric": "P80_Fold_Error", "Before": previous_metrics.get("Best_P80_Fold_Error", ""), "After": benchmark.best_metrics.get("P80_Fold_Error", "")},
        {"Metric": "Pairwise_concordance_pct", "Before": previous_metrics.get("Pairwise_concordance_pct", ""), "After": benchmark.best_metrics.get("Trend_Pairwise_Concordance_pct", "")},
        {"Metric": "Top20_overlap_pct", "Before": previous_metrics.get("Top20_overlap_pct", ""), "After": benchmark.best_metrics.get("Trend_Top20_Overlap_pct", "")},
        {"Metric": "Best_model", "Before": previous_metrics.get("Best_model", ""), "After": benchmark.best_model},
        {"Metric": "Trend_best_model", "Before": previous_metrics.get("TrendComparison_Model", previous_metrics.get("Trend_Best_model", "")), "After": trend.best_model if trend else ""},
        {"Metric": "Trend_Spearman", "Before": previous_metrics.get("TrendComparison_Trend_Spearman_r", ""), "After": next((r.get("Trend_Spearman_r", "") for r in trend.comparison if r.get("Model") == trend.best_model), "") if trend else ""},
        {"Metric": "Trend_Delta_vs_raw", "Before": previous_metrics.get("TrendComparison_Delta_Spearman_vs_raw", ""), "After": next((r.get("Delta_Spearman_vs_raw", "") for r in trend.comparison if r.get("Model") == trend.best_model), "") if trend else ""},
        {"Metric": "Trend_pairwise_pct", "Before": previous_metrics.get("TrendComparison_Trend_Pairwise_Concordance_pct", ""), "After": next((r.get("Trend_Pairwise_Concordance_pct", "") for r in trend.comparison if r.get("Model") == trend.best_model), "") if trend else ""},
        {"Metric": "Trend_decision", "Before": previous_metrics.get("Trend_Decision", ""), "After": trend.decision if trend else ""},
        {"Metric": "Interpretation", "Before": "previous descriptor-workbook result", "After": "model-only rerun after selected group exclusions and current settings"},
    ]
    _append_dict_sheet(wb, "Group_Exclusion_Impact", impact_rows)
    diagnostics = [
        {"Metric": "Model_only_rerun", "Value": True},
        {"Metric": "Reused_descriptor_workbook", "Value": str(source_xlsx)},
        {"Metric": "Master_descriptor_rows", "Value": len(master)},
        {"Metric": "Training_rows_before_group_exclusion", "Value": len(training_all)},
        {"Metric": "Training_rows_after_group_exclusion", "Value": len(training)},
        {"Metric": "Excluded_injection_groups", "Value": "; ".join(excluded)},
        {"Metric": "Target_rows", "Value": len(targets)},
        {"Metric": "Descriptor_columns", "Value": len(descriptor_names)},
        {"Metric": "Concentration_unit", "Value": getattr(published_config, "concentration_unit", "") if published_config else ""},
        {"Metric": "Best_model", "Value": benchmark.best_model},
        {"Metric": "Model_decision", "Value": benchmark.decision_label},
        {"Metric": "Trend_optimizer_best_model", "Value": trend.best_model if trend else ""},
        {"Metric": "Trend_optimizer_decision", "Value": trend.decision if trend else ""},
        {"Metric": "Fixed_feature_search", "Value": bool(fixed_result)},
        {"Metric": "Fixed_feature_model", "Value": fixed_result.model if fixed_result else ""},
        {"Metric": "Fixed_feature_candidate_count", "Value": len(fixed_result.candidate_features) if fixed_result else 0},
        {"Metric": "Fixed_feature_best_absolute_trial", "Value": fixed_result.top5_absolute[0].get("Trial", "") if fixed_result and fixed_result.top5_absolute else ""},
        {"Metric": "Fixed_feature_best_trend_trial", "Value": fixed_result.top5_trend[0].get("Trial", "") if fixed_result and fixed_result.top5_trend else ""},
        {"Metric": "Manual_feature_pool_count", "Value": len(tuple(manual_feature_pool or ()))},
        {"Metric": "Level_classification", "Value": bool(level_result)},
        {"Metric": "Level_classification_best_model", "Value": level_result.best_model if level_result else ""},
        {"Metric": "Level_classification_decision", "Value": level_result.decision if level_result else ""},
    ]
    for warning in warnings + list(benchmark.warnings) + (list(trend.warnings) if trend else []) + (list(fixed_result.warnings) if fixed_result else []):
        diagnostics.append({"Metric": "Warning", "Value": warning})
    _append_dict_sheet(wb, "Model_Rerun_Diagnostics", diagnostics)

    # Keep the new decision page easy to find.
    if "Known_Standards_Predictions" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Known_Standards_Predictions")
    elif level_result is not None and "Level_Model_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Level_Model_Decision")
    elif fixed_result is not None and "Fixed_Search_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Fixed_Search_Decision")
    elif trend is not None and "Trend_Optimizer_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Trend_Optimizer_Decision")
    elif "Model_Decision" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Model_Decision")
    wb.save(str(work_path))
    if work_path != out_xlsx:
        shutil.move(str(work_path), str(out_xlsx))

    best = benchmark.best_metrics
    def opt(key):
        return _num(best.get(key))
    return EsiDescriptorReport(
        output_xlsx=out_xlsx,
        output_csv=None,
        n_master=len(master),
        n_training=len(training_all),
        n_target=len(targets),
        n_training_matched=sum(1 for r in training_all if _text(r.get("Structure_status"))),
        n_target_matched=sum(1 for r in targets if _text(r.get("Structure_status"))),
        n_rt_training=sum(1 for r in training_all if _num(r.get("Apex_RT_min")) is not None),
        n_rt_target=sum(1 for r in targets if _num(r.get("Apex_RT_min")) is not None),
        warnings=warnings + list(benchmark.warnings) + (list(trend.warnings) if trend else []) + (list(fixed_result.warnings) if fixed_result else []),
        model_benchmark_enabled=True,
        model_calibration_rows=benchmark.n_calibration,
        model_target_rows=benchmark.n_target_valid,
        best_model=benchmark.best_model,
        best_model_rating=benchmark.best_rating,
        model_decision=benchmark.decision_label,
        model_decision_summary=benchmark.decision_summary,
        best_median_fold_error=opt("Median_Fold_Error"),
        best_p80_fold_error=opt("P80_Fold_Error"),
        best_within_2x_pct=opt("Within_2x_pct"),
        best_within_5x_pct=opt("Within_5x_pct"),
        best_r2_log=opt("R2_log_concentration"),
        best_rmse_log10=opt("RMSE_log10"),
        improvement_vs_global_pct=opt("Median_Fold_Improvement_vs_Global_pct"),
        model_outliers_removed=benchmark.n_outliers_removed,
        model_selected_descriptors=benchmark.n_selected_descriptors,
        model_unavailable_descriptors=benchmark.n_unavailable_descriptors,
        model_objective=benchmark.model_objective_used,
        corrected_spearman=opt("Trend_Spearman_r"),
        raw_spearman=opt("Raw_Trend_Spearman_r"),
        delta_spearman=opt("Delta_Trend_Spearman_r_vs_raw"),
        pairwise_concordance_pct=opt("Trend_Pairwise_Concordance_pct"),
        top20_overlap_pct=opt("Trend_Top20_Overlap_pct"),
        dual_ie_enabled=dual_result is not None,
        dual_ie_decision=dual_result.decision_label if dual_result else "",
        dual_ie_summary=dual_result.decision_summary if dual_result else "",
        dual_ie_published_method=published_result.method_code if published_result else "",
        dual_ie_target_rank_spearman=dual_result.target_rank_spearman if dual_result else None,
        published_ie_spearman=_num(published_result.metrics.get("Trend_Spearman_r")) if published_result else None,
        published_ie_median_fold=_num(published_result.metrics.get("Median_Fold_Error")) if published_result else None,
        trend_optimizer_enabled=trend is not None,
        trend_optimizer_best_model=trend.best_model if trend else "",
        trend_optimizer_decision=trend.decision if trend else "",
        trend_optimizer_summary=trend.decision_summary if trend else "",
        trend_optimizer_spearman=_num(next((r.get("Trend_Spearman_r") for r in trend.comparison if r.get("Model") == trend.best_model), None)) if trend else None,
        trend_optimizer_delta_spearman=_num(next((r.get("Delta_Spearman_vs_raw") for r in trend.comparison if r.get("Model") == trend.best_model), None)) if trend else None,
        trend_optimizer_permutation_p=_num(next((r.get("Spearman_permutation_p") for r in trend.comparison if r.get("Model") == trend.best_model), None)) if trend else None,
        fixed_feature_search_enabled=fixed_result is not None,
        fixed_feature_model=fixed_result.model if fixed_result else "",
        fixed_feature_best_absolute_trial=(int(fixed_result.top5_absolute[0].get("Trial")) if fixed_result and fixed_result.top5_absolute else None),
        fixed_feature_best_trend_trial=(int(fixed_result.top5_trend[0].get("Trial")) if fixed_result and fixed_result.top5_trend else None),
        level_classification_enabled=level_result is not None,
        level_classification_best_model=level_result.best_model if level_result else "",
        level_classification_decision=level_result.decision if level_result else "",
        level_classification_exact_pct=_num(next((r.get("Exact_level_accuracy_pct") for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), None)) if level_result else None,
        level_classification_within1_pct=_num(next((r.get("Within_1_level_pct") for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), None)) if level_result else None,
        level_classification_all_hit=bool(next((r.get("All_exactly_hit") for r in level_result.comparison if str(r.get("Model")) == level_result.best_model), False)) if level_result else False,
    )
