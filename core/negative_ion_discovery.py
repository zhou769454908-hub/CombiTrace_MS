"""Aggregate negative-ion channel evidence and generate an experiment-specific panel."""
from __future__ import annotations

import csv
import math
from collections import defaultdict
from pathlib import Path
from statistics import median
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .negative_ion_channels import (
    CHANNELS,
    CHANNEL_BY_ID,
    PUBLIC_CHANNEL_IDS,
    NegativeIonPanel,
    PanelChannel,
    channel_registry_rows,
    normalise_action,
    write_panel_csv,
)


def _bool(v: object) -> bool:
    s = str(v or "").strip().lower()
    return s in {"1", "true", "yes", "y", "accepted"}


def _float(v: object) -> float:
    try:
        x = float(v)
        return x if math.isfinite(x) else float("nan")
    except Exception:
        return float("nan")


def _median(values: Iterable[object]) -> object:
    xs = [_float(v) for v in values]
    xs = [x for x in xs if math.isfinite(x)]
    return median(xs) if xs else ""


def _pct(n: int, d: int) -> float:
    return (100.0 * float(n) / float(d)) if d else 0.0


def _read_csv(path: Path) -> List[Dict[str, str]]:
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        return [dict(r) for r in csv.DictReader(f)]


def _autosize(ws, max_width: int = 52) -> None:
    for col in range(1, ws.max_column + 1):
        width = 10
        for row in range(1, min(ws.max_row, 500) + 1):
            value = ws.cell(row, col).value
            if value is not None:
                width = max(width, len(str(value)) + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(max_width, width)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_sheet(ws, rows: Sequence[Mapping[str, object]]) -> None:
    if not rows:
        ws.append(["No data"])
        return
    headers: List[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(str(key))
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    _autosize(ws)


def aggregate_channel_evidence(
    evidence_csv_paths: Sequence[Path],
    *,
    out_xlsx: Path,
    out_panel_csv: Path,
    prevalence_threshold_pct: float = 20.0,
    min_isotope_support_pct: float = 50.0,
    min_median_area_fraction_pct: float = 0.5,
    panel_name: str = "experiment_negative_ion_panel",
) -> Tuple[Path, Path, List[Dict[str, object]]]:
    """Aggregate evidence from multiple RAWs and build a conservative editable panel.

    Suggested actions:
    - frequent 1:1 quantifiable adduct/substitution with sufficient isotope support -> sum;
    - clusters and less frequent 1:1 channels -> search;
    - dimer/fragment/transformation/diagnostic -> evidence;
    - absent channels -> exclude.
    """
    evidence_rows: List[Dict[str, object]] = []
    for p in evidence_csv_paths:
        pp = Path(p)
        if not pp.exists():
            continue
        sample = pp.name.split("__negative_ion_channels", 1)[0]
        for row in _read_csv(pp):
            row["Source_File"] = str(pp)
            row["Sample"] = sample
            evidence_rows.append(row)

    by_channel: Dict[str, List[Dict[str, object]]] = defaultdict(list)
    for row in evidence_rows:
        cid = str(row.get("Channel_ID", "") or "").strip().upper()
        if cid:
            by_channel[cid].append(row)

    summary: List[Dict[str, object]] = []
    panel_rows: List[PanelChannel] = []
    for cid in PUBLIC_CHANNEL_IDS:
        defn = CHANNEL_BY_ID[cid]
        rows = by_channel.get(cid, [])
        eligible = [r for r in rows if _bool(r.get("Eligible"))]
        accepted = [r for r in eligible if _bool(r.get("Accepted"))]
        high_conf = [r for r in accepted if _bool(r.get("High_Confidence"))]
        isotope_applicable = [r for r in accepted if str(r.get("Isotope_Support_OK", "")).strip() != ""]
        isotope_ok = [r for r in isotope_applicable if _bool(r.get("Isotope_Support_OK"))]
        detection_pct = _pct(len(accepted), len(eligible))
        high_conf_pct = _pct(len(high_conf), len(eligible))
        isotope_pct = _pct(len(isotope_ok), len(isotope_applicable)) if isotope_applicable else 100.0
        med_area_pct = _median(r.get("Area_Fraction_vs_MH_Pct") for r in accepted)
        med_area_num = float(med_area_pct) if med_area_pct != "" else 0.0
        med_rt = _median(r.get("RT_Delta_Min") for r in accepted)
        med_corr = _median(r.get("Shape_Correlation") for r in accepted)

        present = bool(
            len(accepted) > 0
            and detection_pct >= float(prevalence_threshold_pct)
            and med_area_num >= float(min_median_area_fraction_pct)
            and isotope_pct >= float(min_isotope_support_pct)
        )
        if not accepted:
            present_label = "Absent"
        elif present:
            present_label = "Present"
        else:
            present_label = "Occasional / review"

        if defn.category in {"dimer", "fragment", "transformation", "diagnostic"}:
            action = "evidence" if accepted else "exclude"
        elif not present:
            action = "search" if accepted else "exclude"
        elif defn.category in {"adduct", "substitution"} and defn.quantifiable:
            action = "sum"
        else:
            action = "search"

        cfg = PanelChannel(
            channel_id=cid,
            enabled=action != "exclude",
            quant_action=action,
            min_rel_height_pct=1.0,
            min_shape_correlation=0.30,
            rt_tolerance_min=0.10,
            note=f"auto suggestion: {present_label}; detection={detection_pct:.1f}%",
        )
        panel_rows.append(cfg)
        summary.append({
            "Channel_ID": cid,
            "Channel": defn.display,
            "Category": defn.category,
            "Exact_Delta_From_[M-H]-": (
                f"{defn.exact_delta_from_mh:+.6f}" if defn.exact_delta_from_mh is not None else defn.mz_kind
            ),
            "Source_Hypothesis": defn.source,
            "Priority_Stars": defn.priority_stars,
            "Eligible_Rows": len(eligible),
            "Accepted_Rows": len(accepted),
            "Detection_Rate_Pct": detection_pct,
            "High_Confidence_Rate_Pct": high_conf_pct,
            "Median_Area_vs_MH_Pct": med_area_pct,
            "Median_RT_Delta_Min": med_rt,
            "Median_Shape_Correlation": med_corr,
            "Isotope_Support_Rate_Pct": isotope_pct if isotope_applicable else "NA",
            "Presence_Assessment": present_label,
            "Suggested_Action": action,
            "Quantifiable": defn.quantifiable,
            "Note": defn.note,
        })

    panel = NegativeIonPanel(panel_rows, panel_name=panel_name, require_primary_mh=True, version="1.0")
    summary_map: Dict[str, Dict[str, object]] = {}
    for row in summary:
        cid = str(row.get("Channel_ID", ""))
        summary_map[cid] = {
            "Eligible_Count": row.get("Eligible_Rows", ""),
            "Detected_Count": row.get("Accepted_Rows", ""),
            "Detection_Rate_Pct": row.get("Detection_Rate_Pct", ""),
            "Median_Area_Fraction_vs_MH_Pct": row.get("Median_Area_vs_MH_Pct", ""),
            "Median_RT_Delta_Min": row.get("Median_RT_Delta_Min", ""),
            "Median_Shape_Correlation": row.get("Median_Shape_Correlation", ""),
            "Isotope_Support_Rate_Pct": row.get("Isotope_Support_Rate_Pct", ""),
            "Suggested_Status": row.get("Presence_Assessment", ""),
        }
    write_panel_csv(panel, Path(out_panel_csv), summary=summary_map)

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet("Channel_Summary"), summary)
    _write_sheet(wb.create_sheet("Per_Target_Evidence"), evidence_rows)
    panel_table: List[Dict[str, object]] = []
    for cfg in panel.channels:
        defn = CHANNEL_BY_ID[cfg.channel_id]
        panel_table.append({
            "Channel_ID": cfg.channel_id,
            "Enabled": cfg.enabled,
            "Quant_Action": cfg.quant_action,
            "Min_Rel_Height_Pct": cfg.min_rel_height_pct,
            "Min_Shape_Correlation": cfg.min_shape_correlation,
            "RT_Tolerance_Min": cfg.rt_tolerance_min,
            "Channel": defn.display,
            "Category": defn.category,
            "Note": cfg.note,
        })
    _write_sheet(wb.create_sheet("Suggested_Panel"), panel_table)
    _write_sheet(wb.create_sheet("Channel_Registry"), channel_registry_rows())
    methods = [{
        "Item": "Interpretation",
        "Value": "Discovery reports all hypotheses but does not alter quantitative area. Only a fixed panel marked 'sum' contributes to Area.",
    }, {
        "Item": "Conservative rule",
        "Value": "Dimer, fragment, transformation, and diagnostic channels remain evidence-only by default.",
    }, {
        "Item": "Isotope support",
        "Value": "Chloride and bromide hypotheses require their characteristic isotope-support channel for high-confidence acceptance/summation.",
    }]
    _write_sheet(wb.create_sheet("Methods"), methods)
    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx, Path(out_panel_csv), summary
