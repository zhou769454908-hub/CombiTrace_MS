"""Response-factor based concentration prediction.

This module adds a practical calibration/prediction layer for CombiTrace-MS
results.  A small set of compounds with known actual concentration and measured
area/internal-standard ratio is used to estimate compound-dependent response
correction factors.  The fitted factors are then transferred to a larger result
list by combo/provenance-aware similarity.

The key quantity is:

    correction_factor = actual_concentration / measured_ratio
    predicted_concentration = measured_ratio * predicted_correction_factor

The measured ratio can be a raw Area/IS value or Area/IS*100 %.  The model only
requires that the calibration table and the target table use the same ratio
scale.
"""

from __future__ import annotations
from .legacy_schema import canonical_header, canonical_sheet, alias_record

import csv
import hashlib
import math
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np

from .chemistry import monoisotopic_mass, parse_formula


@dataclass
class TableData:
    path: Path
    sheet_name: str
    headers: List[str]
    rows: List[Dict[str, object]]


@dataclass
class ColumnConfig:
    combo_col: str = ""
    formula_col: str = ""
    ratio_col: str = ""
    concentration_col: str = ""
    name_col: str = ""
    exact_mass_col: str = ""


@dataclass
class PreparedRow:
    source_index: int
    raw: Dict[str, object]
    combo: str
    formula: str
    name: str
    ratio: Optional[float]
    concentration: Optional[float]
    exact_mass: Optional[float]
    correction_factor: Optional[float]
    log_correction_factor: Optional[float]
    combo_parts: Dict[str, Dict[str, str]] = field(default_factory=dict)
    element_counts: Dict[str, int] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)
    # Optional local/private structure evidence used by CombiTrace-IE.
    structure_features: Dict[str, float] = field(default_factory=dict)
    structure_fingerprint: object = None
    component_fingerprints: Dict[str, object] = field(default_factory=dict)
    core_id: str = ""
    structure_status: str = ""
    structure_method: str = ""
    structure_hash: str = ""
    product_formula_calc: str = ""
    formula_match: str = ""
    three_d_status: str = ""
    structure_warnings: List[str] = field(default_factory=list)
    product_smiles: str = ""


@dataclass
class PredictionResult:
    row: PreparedRow
    predicted_concentration: Optional[float]
    predicted_correction_factor: Optional[float]
    method: str
    evidence: str
    nearest_count: int
    similarity_mean: Optional[float]
    nearest_summary: str
    ridge_log_cf: Optional[float] = None
    knn_log_cf: Optional[float] = None
    warnings: List[str] = field(default_factory=list)


@dataclass
class RunReport:
    output_xlsx: Path
    output_csv: Optional[Path]
    n_train_total: int
    n_train_used: int
    n_target_total: int
    n_predicted: int
    warnings: List[str]
    qc: Dict[str, object]


MODEL_LABEL_TO_CODE = {
    'Hybrid (KNN + ridge)': "blend",
    'KNN response transfer (Combo/formula similarity)': "knn",
    'Ridge response transfer (structural features)': "ridge",
    'Global median response factor': "global",
}


# -----------------------------
# Table reading and column guess
# -----------------------------

def _norm_header(s: str) -> str:
    s = canonical_header(s)
    s = str(s or "").strip().lower()
    s = s.replace("\uFF05", "%")
    s = re.sub(r"[\s\-_()/\\]+", "", s)
    return s


def _safe_text(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _read_csv_table(path: Path) -> TableData:
    encodings = ["utf-8-sig", "utf-8", "gbk", "gb18030"]
    last_err: Optional[Exception] = None
    for enc in encodings:
        try:
            with Path(path).open("r", encoding=enc, newline="") as f:
                reader = csv.DictReader(f)
                headers = [str(h or "").strip() for h in (reader.fieldnames or [])]
                rows = []
                for line_no, row in enumerate(reader, start=2):
                    rr = {str(k or "").strip(): ("" if v is None else v) for k, v in row.items()}
                    rr = alias_record(rr)
                    rr["__source_row__"] = line_no
                    rr["__source_sheet__"] = ""
                    rr["__source_file__"] = str(Path(path).name)
                    rows.append(rr)
                return TableData(path=Path(path), sheet_name="", headers=headers, rows=rows)
        except Exception as e:
            last_err = e
    raise RuntimeError(f'Could not read CSV: {path} ({last_err})')


def _read_xlsx_table(path: Path, sheet_name: str = "") -> TableData:
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError('Excel input requires openpyxl. Install it in the active environment: python -m pip install openpyxl') from e

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif sheet_name and sheet_name not in wb.sheetnames:
            raise ValueError(f'Worksheet not found: {sheet_name}')
        else:
            
            candidates = [
                "Triplicate Average",
                "Final Triplicate Mean",
                "Predictions",
            ]
            ws = None
            for c in candidates:
                matching = [s for s in wb.sheetnames if canonical_sheet(s).casefold() == canonical_sheet(c).casefold()]
                if c in wb.sheetnames:
                    ws = wb[c]
                    break
                if len(matching) > 1:
                    raise ValueError("Multiple equivalent summary sheets; select a sheet explicitly: " + ", ".join(matching))
                if matching:
                    ws = wb[matching[0]]
                    break
            if ws is None:
                ws = wb.active

        values = list(ws.iter_rows(values_only=True))
        if not values:
            return TableData(path=Path(path), sheet_name=ws.title, headers=[], rows=[])
        
        header_idx = 0
        for i, row in enumerate(values[:20]):
            if row and sum(1 for v in row if _safe_text(v)) >= 2:
                header_idx = i
                break
        headers = [_safe_text(v) for v in values[header_idx]]
        
        last = 0
        for i, h in enumerate(headers, start=1):
            if h:
                last = i
        headers = headers[:last]
        out_rows: List[Dict[str, object]] = []
        for offset, row in enumerate(values[header_idx + 1 :], start=1):
            vals = list(row[: len(headers)])
            if not any(_safe_text(v) for v in vals):
                continue
            rr = {headers[i]: (vals[i] if i < len(vals) else "") for i in range(len(headers))}
            rr = alias_record(rr)
            rr["__source_row__"] = int(header_idx + 1 + offset)  # Excel row number, 1-based
            rr["__source_sheet__"] = ws.title
            rr["__source_file__"] = str(Path(path).name)
            out_rows.append(rr)
        return TableData(path=Path(path), sheet_name=ws.title, headers=headers, rows=out_rows)
    finally:
        wb.close()


def read_table(path: Path, sheet_name: str = "") -> TableData:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(str(path))
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _read_xlsx_table(path, sheet_name=sheet_name)
    return _read_csv_table(path)


def _pick_col(headers: Sequence[str], candidates: Sequence[str], *, contains: bool = False) -> str:
    norm_to_orig = {_norm_header(h): h for h in headers}
    for c in candidates:
        nc = _norm_header(c)
        if nc in norm_to_orig:
            return norm_to_orig[nc]
    if contains:
        for h in headers:
            nh = _norm_header(h)
            for c in candidates:
                if _norm_header(c) and _norm_header(c) in nh:
                    return h
    return ""


def guess_columns(table: TableData, *, need_concentration: bool) -> ColumnConfig:
    h = table.headers
    combo = _pick_col(h, ["combo", 'Combo', 'Source_Combination', "sourcecombo", 'Source_Combination'], contains=True)
    formula = _pick_col(h, ["formula", "molecularformula", "chemicalformula", 'Formula', "Formula"], contains=True)
    exact_mass = _pick_col(h, ["exact_mass", "Exact_mass", "exactmass", "Exact Mass", 'Exact_mass', 'Monoisotopic_mass'], contains=True)
    name = _pick_col(h, ["name", "Name", "compound", "compoundname", 'Compound', 'Name'], contains=True)

    ratio_candidates = [
        'Mean_Area_to_IS_Ratio_%',
        'Mean_Area_to_IS_Ratio_%',
        "Ratio_mean(vs_IS)",
        "Ratio_mean",
        "Ratio_%",
        "ratio",
        "ratio%",
        "area_is_ratio",
        "area/is",
        "area_to_is",
        "area/internalstandard",
        'Area_to_IS_Ratio',
        'Area_to_IS_Ratio',
        'Area_to_IS_Ratio',
        'Measured_ratio',
        'Ratio',
        'Final_Triplicate_Mean_Result',
        "Final Triplicate Mean",
        "finalresult",
    ]
    ratio = _pick_col(h, ratio_candidates, contains=True)

    conc = ""
    if need_concentration:
        conc_candidates = [
            "actual_concentration",
            "actualconc",
            "actual concentration",
            "known_concentration",
            "knownconc",
            "concentration",
            "conc",
            'Actual concentration',
            'Actual_concentration',
            'Known_concentration',
            'Standard_concentration',
            'Concentration',
        ]
        conc = _pick_col(h, conc_candidates, contains=True)

    return ColumnConfig(combo_col=combo, formula_col=formula, ratio_col=ratio, concentration_col=conc, name_col=name, exact_mass_col=exact_mass)


# -----------------------------
# Row preparation and features
# -----------------------------

def parse_number(v: object) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x if math.isfinite(x) else None
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("\uFF0C", ",").replace("\uFF05", "%")
    # remove units / commas, keep exponent and sign.
    s = s.replace(",", "")
    m = re.search(r"[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?", s)
    if not m:
        return None
    try:
        x = float(m.group(0))
        return x if math.isfinite(x) else None
    except Exception:
        return None


def parse_combo(combo: str) -> Dict[str, Dict[str, str]]:
    out: Dict[str, Dict[str, str]] = {}
    text = str(combo or "")
    for part in re.split(r"\s*\|\s*", text):
        part = part.strip()
        m = re.match(r"^([ABC])#(\d+)\s*:\s*(.+)$", part, flags=re.I)
        if not m:
            continue
        group = m.group(1).upper()
        idx = m.group(2)
        rest = m.group(3).strip()
        label = rest
        formula = ""
        if ":" in rest:
            label, formula = rest.rsplit(":", 1)
        out[group] = {"index": idx, "label": label.strip(), "formula": formula.strip()}
    return out


def _counts_from_formula(formula: str) -> Tuple[Dict[str, int], List[str]]:
    warnings: List[str] = []
    formula = str(formula or "").strip()
    if not formula:
        return {}, warnings
    try:
        counts = parse_formula(formula)
        return {str(k): int(v) for k, v in counts.items()}, warnings
    except Exception as e:
        warnings.append(f"Formula parse failed: {formula!r}: {e}")
        return {}, warnings


def prepare_rows(table: TableData, config: ColumnConfig, *, is_training: bool) -> Tuple[List[PreparedRow], List[str]]:
    warnings: List[str] = []
    if is_training:
        from .training_limits import validate_training_count
        validate_training_count(table.rows, label=str(table.path.name))
    if not config.combo_col:
        warnings.append(f'{table.path.name}: no Combo column detected; formula/mass fallback matching will be used.')
    if not config.ratio_col:
        warnings.append(f'{table.path.name}: no target-to-internal-standard ratio column detected.')
    if is_training and not config.concentration_col:
        warnings.append(f'{table.path.name}: no known-concentration column detected.')

    rows: List[PreparedRow] = []
    for i, raw in enumerate(table.rows, start=1):
        combo = _safe_text(raw.get(config.combo_col, "")) if config.combo_col else ""
        formula = _safe_text(raw.get(config.formula_col, "")) if config.formula_col else ""
        name = _safe_text(raw.get(config.name_col, "")) if config.name_col else ""
        ratio = parse_number(raw.get(config.ratio_col, "")) if config.ratio_col else None
        conc = parse_number(raw.get(config.concentration_col, "")) if config.concentration_col else None
        exact_mass = parse_number(raw.get(config.exact_mass_col, "")) if config.exact_mass_col else None
        combo_parts = parse_combo(combo)
        elem_counts, ws = _counts_from_formula(formula)
        if exact_mass is None and elem_counts:
            try:
                exact_mass = float(monoisotopic_mass(elem_counts))
            except Exception:
                pass
        cf: Optional[float] = None
        lcf: Optional[float] = None
        rwarn = list(ws)
        if ratio is not None and ratio > 0 and conc is not None and conc > 0:
            cf = float(conc / ratio)
            if cf > 0 and math.isfinite(cf):
                lcf = float(math.log(cf))
        elif is_training:
            rwarn.append('Missing or invalid ratio/concentration; record excluded from training.')
        rows.append(
            PreparedRow(
                source_index=i,
                raw=raw,
                combo=combo,
                formula=formula,
                name=name,
                ratio=ratio,
                concentration=conc,
                exact_mass=exact_mass,
                correction_factor=cf,
                log_correction_factor=lcf,
                combo_parts=combo_parts,
                element_counts=elem_counts,
                warnings=rwarn,
            )
        )
    return rows, warnings


def _element_keys(rows: Sequence[PreparedRow]) -> List[str]:
    keys = set()
    for r in rows:
        keys.update(r.element_counts.keys())
        for p in r.combo_parts.values():
            fc, _ = _counts_from_formula(p.get("formula", ""))
            keys.update(fc.keys())
    preferred = ["C", "H", "N", "O", "P", "S", "Na", "Cl", "Br", "I", "F", "K"]
    return [k for k in preferred if k in keys] + sorted(k for k in keys if k not in preferred)


def _formula_distance(a: Dict[str, int], b: Dict[str, int], keys: Sequence[str]) -> float:
    if not a and not b:
        return 1.0
    denom = 1.0 + sum(abs(a.get(k, 0)) + abs(b.get(k, 0)) for k in keys)
    dist = sum(abs(a.get(k, 0) - b.get(k, 0)) for k in keys) / denom
    return float(dist)


def combo_similarity(a: PreparedRow, b: PreparedRow, element_keys: Sequence[str]) -> float:
    # Exact combo should dominate when available.
    if a.combo and b.combo and a.combo.strip() == b.combo.strip():
        return 1.0

    # Component/provenance match score.
    comp_score = 0.0
    comp_n = 0
    for g in ("A", "B", "C"):
        pa = a.combo_parts.get(g, {})
        pb = b.combo_parts.get(g, {})
        if pa or pb:
            comp_n += 1
            # Exact label match is more useful than row index because row order may change.
            if pa.get("label") and pb.get("label") and pa.get("label") == pb.get("label"):
                comp_score += 1.0
            elif pa.get("index") and pb.get("index") and pa.get("index") == pb.get("index"):
                comp_score += 0.45
            elif pa.get("formula") and pb.get("formula") and pa.get("formula") == pb.get("formula"):
                comp_score += 0.35
    comp_sim = comp_score / comp_n if comp_n else 0.0

    fdist = _formula_distance(a.element_counts, b.element_counts, element_keys)
    formula_sim = math.exp(-7.0 * fdist)

    mass_sim = 0.0
    if a.exact_mass is not None and b.exact_mass is not None:
        mass_sim = math.exp(-abs(float(a.exact_mass) - float(b.exact_mass)) / 80.0)

    # Weighted combination. Provenance is very important, but formula/mass gives fallback.
    score = 0.52 * comp_sim + 0.36 * formula_sim + 0.12 * mass_sim
    return float(max(0.0, min(0.999, score)))


# -----------------------------
# Models
# -----------------------------

def _valid_training(rows: Sequence[PreparedRow]) -> List[PreparedRow]:
    return [r for r in rows if r.log_correction_factor is not None and r.ratio is not None and r.ratio > 0 and r.concentration is not None and r.concentration > 0]


def _weighted_mean(vals: Sequence[float], weights: Sequence[float]) -> Optional[float]:
    pairs = [(float(v), float(w)) for v, w in zip(vals, weights) if math.isfinite(float(v)) and math.isfinite(float(w)) and float(w) > 0]
    if not pairs:
        return None
    sw = sum(w for _, w in pairs)
    if sw <= 0:
        return None
    return float(sum(v * w for v, w in pairs) / sw)


def predict_log_cf_knn(target: PreparedRow, train_rows: Sequence[PreparedRow], element_keys: Sequence[str], *, k: int = 15) -> Tuple[Optional[float], str, int, Optional[float], str]:
    train = _valid_training(train_rows)
    if not train:
        return None, "no_training", 0, None, ""

    # Exact same combo: use these only.
    exact = [r for r in train if target.combo and r.combo and target.combo.strip() == r.combo.strip()]
    if exact:
        vals = [float(r.log_correction_factor) for r in exact if r.log_correction_factor is not None]
        y = float(statistics.median(vals)) if vals else None
        summary = "; ".join(_short_combo(r.combo) for r in exact[:3])
        return y, "exact_combo", len(vals), 1.0, summary

    scored: List[Tuple[float, PreparedRow]] = []
    for r in train:
        s = combo_similarity(target, r, element_keys)
        if s > 0:
            scored.append((s, r))
    scored.sort(key=lambda x: x[0], reverse=True)
    top = scored[: max(1, int(k))]
    if not top:
        vals = [float(r.log_correction_factor) for r in train if r.log_correction_factor is not None]
        return float(statistics.median(vals)), "global_median", len(vals), None, ""
    vals = [float(r.log_correction_factor) for s, r in top if r.log_correction_factor is not None]
    weights = [max(s, 1e-6) ** 2 for s, r in top if r.log_correction_factor is not None]
    y = _weighted_mean(vals, weights)
    sim_mean = float(sum(s for s, _ in top) / len(top)) if top else None
    summary = "; ".join(f"{_short_combo(r.combo)} (sim={s:.2f})" for s, r in top[:3])
    return y, "nearest_combo_formula", len(top), sim_mean, summary


def _short_combo(combo: str, max_len: int = 80) -> str:
    s = str(combo or "").replace("\n", " ").strip()
    return s[: max_len - 1] + "\u2026" if len(s) > max_len else s


@dataclass
class RidgeModel:
    beta: Optional[np.ndarray]
    mu: np.ndarray
    sd: np.ndarray
    element_keys: List[str]
    dim_hash: int
    lambda_: float
    y_median: float


def _stable_hash(s: str, mod: int) -> int:
    h = hashlib.md5(s.encode("utf-8", errors="ignore")).hexdigest()
    return int(h[:8], 16) % int(mod)


def _base_feature_vector(row: PreparedRow, element_keys: Sequence[str], dim_hash: int = 64) -> np.ndarray:
    cont: List[float] = []
    total_atoms = sum(max(0, int(v)) for v in row.element_counts.values())
    for k in element_keys:
        cont.append(float(row.element_counts.get(k, 0)))
    cont.append(float(total_atoms))
    cont.append(float(row.exact_mass or 0.0))
    hetero = sum(float(row.element_counts.get(k, 0)) for k in element_keys if k not in {"C", "H"})
    cont.append(float(hetero))

    hashes = np.zeros(int(dim_hash), dtype=float)
    for g in ("A", "B", "C"):
        p = row.combo_parts.get(g, {})
        for key in ("label", "formula"):
            val = p.get(key, "")
            if val:
                hashes[_stable_hash(f"{g}:{key}:{val}", dim_hash)] += 1.0
    # Pair provenance features often help transfer response factor.
    for pair in (("A", "B"), ("A", "C"), ("B", "C")):
        v1 = row.combo_parts.get(pair[0], {}).get("label", "")
        v2 = row.combo_parts.get(pair[1], {}).get("label", "")
        if v1 and v2:
            hashes[_stable_hash(f"{pair[0]}{pair[1]}:{v1}|{v2}", dim_hash)] += 1.0

    return np.concatenate([np.array(cont, dtype=float), hashes])


def fit_ridge(train_rows: Sequence[PreparedRow], element_keys: Sequence[str], *, lambda_: float = 1.0, dim_hash: int = 64) -> RidgeModel:
    train = _valid_training(train_rows)
    ys = np.array([float(r.log_correction_factor) for r in train], dtype=float)
    if len(train) == 0:
        return RidgeModel(None, np.zeros(1), np.ones(1), list(element_keys), dim_hash, float(lambda_), 0.0)
    y_median = float(np.median(ys))
    X0 = np.vstack([_base_feature_vector(r, element_keys, dim_hash=dim_hash) for r in train])
    mu = X0.mean(axis=0)
    sd = X0.std(axis=0)
    sd[sd == 0] = 1.0
    X = (X0 - mu) / sd
    X = np.hstack([np.ones((X.shape[0], 1)), X])
    lam = float(lambda_ if lambda_ is not None else 1.0)
    I = np.eye(X.shape[1])
    I[0, 0] = 0.0
    try:
        beta = np.linalg.solve(X.T @ X + lam * I, X.T @ ys)
    except Exception:
        beta = np.linalg.pinv(X.T @ X + lam * I) @ (X.T @ ys)
    return RidgeModel(beta=beta, mu=mu, sd=sd, element_keys=list(element_keys), dim_hash=dim_hash, lambda_=lam, y_median=y_median)


def predict_log_cf_ridge(row: PreparedRow, model: RidgeModel) -> Optional[float]:
    if model.beta is None:
        return None
    x0 = _base_feature_vector(row, model.element_keys, dim_hash=model.dim_hash)
    x = (x0 - model.mu) / model.sd
    x = np.concatenate([np.ones(1), x])
    try:
        return float(x @ model.beta)
    except Exception:
        return model.y_median


def predict_one(
    row: PreparedRow,
    train_rows: Sequence[PreparedRow],
    element_keys: Sequence[str],
    *,
    mode: str = "blend",
    k: int = 15,
    ridge_model: Optional[RidgeModel] = None,
) -> PredictionResult:
    warnings: List[str] = []
    if row.ratio is None or row.ratio <= 0:
        return PredictionResult(row, None, None, mode, "no_valid_ratio", 0, None, "", warnings=['No valid target-to-internal-standard area ratio.'])

    train = _valid_training(train_rows)
    if not train:
        return PredictionResult(row, None, None, mode, "no_training", 0, None, "", warnings=['No usable training records.'])

    global_log = float(statistics.median([float(r.log_correction_factor) for r in train if r.log_correction_factor is not None]))
    knn_log, evidence, n_near, sim_mean, near_summary = predict_log_cf_knn(row, train, element_keys, k=k)
    ridge_log = predict_log_cf_ridge(row, ridge_model) if ridge_model is not None else None

    code = str(mode or "blend").lower()
    if code == "global":
        y = global_log
        evidence = "global_median"
    elif code == "knn":
        y = knn_log if knn_log is not None else global_log
    elif code == "ridge":
        y = ridge_log if ridge_log is not None else global_log
        if ridge_log is None:
            evidence = "global_median"
        else:
            evidence = "ridge_model"
    else:
        # If exact combo exists, trust exact KNN. Otherwise blend ridge with nearest-neighbour transfer.
        if evidence == "exact_combo" and knn_log is not None:
            y = knn_log
        elif knn_log is not None and ridge_log is not None:
            # Similarity-dependent blend: when close reference exists, rely more on KNN.
            w_knn = 0.65
            if sim_mean is not None:
                w_knn = max(0.45, min(0.85, 0.35 + 0.55 * float(sim_mean)))
            y = w_knn * knn_log + (1.0 - w_knn) * ridge_log
            evidence = "blend_knn_ridge"
        elif knn_log is not None:
            y = knn_log
        elif ridge_log is not None:
            y = ridge_log
            evidence = "ridge_model"
        else:
            y = global_log
            evidence = "global_median"

    if y is None or not math.isfinite(float(y)):
        return PredictionResult(row, None, None, mode, "failed", n_near, sim_mean, near_summary, ridge_log, knn_log, warnings=['Response-factor prediction failed.'])
    cf = float(math.exp(float(y)))
    pred = float(row.ratio * cf)
    return PredictionResult(row, pred, cf, code, evidence, n_near, sim_mean, near_summary, ridge_log, knn_log, warnings=warnings)


# -----------------------------
# Evaluation and output
# -----------------------------

def _metrics(y_true: Sequence[float], y_pred: Sequence[float]) -> Dict[str, object]:
    pairs = [(float(a), float(b)) for a, b in zip(y_true, y_pred) if a is not None and b is not None and a > 0 and b > 0 and math.isfinite(a) and math.isfinite(b)]
    if not pairs:
        return {}
    yt = np.array([a for a, _ in pairs], dtype=float)
    yp = np.array([b for _, b in pairs], dtype=float)
    err = yp - yt
    ratio = yp / yt
    # Fold error is symmetric: 2 means either 2x high or 2x low.
    fold_error = np.maximum(ratio, 1.0 / ratio)
    ape = np.abs(err / yt) * 100.0
    log_err = np.log(yp) - np.log(yt)
    mae = float(np.mean(np.abs(err)))
    rmse = float(np.sqrt(np.mean(err * err)))
    mape = float(np.mean(ape))
    medape = float(np.median(ape))
    p90ape = float(np.percentile(ape, 90))
    maxape = float(np.max(ape))
    log_rmse = float(np.sqrt(np.mean(log_err * log_err)))
    ss_res = float(np.sum((np.log(yt) - np.log(yp)) ** 2))
    ss_tot = float(np.sum((np.log(yt) - np.mean(np.log(yt))) ** 2))
    r2_log = float(1.0 - ss_res / ss_tot) if ss_tot > 0 else None
    within2x = float(np.mean(fold_error <= 2.0) * 100.0)
    within5x = float(np.mean(fold_error <= 5.0) * 100.0)
    within10x = float(np.mean(fold_error <= 10.0) * 100.0)
    median_fold = float(np.median(fold_error))
    p90_fold = float(np.percentile(fold_error, 90))
    return {
        "N": len(pairs),
        "MAE": mae,
        "RMSE": rmse,
        "MAPE_%": mape,
        "Median_APE_%": medape,
        "P90_APE_%": p90ape,
        "Max_APE_%": maxape,
        "Log_RMSE": log_rmse,
        "R2_log": r2_log,
        "Within_2x_%": within2x,
        "Within_5x_%": within5x,
        "Within_10x_%": within10x,
        "Median_fold_error": median_fold,
        "P90_fold_error": p90_fold,
    }


def _as_float_or_none(v: object) -> Optional[float]:
    try:
        if v is None or v == "":
            return None
        x = float(v)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def evaluate_model_qc(qc: Dict[str, object]) -> Dict[str, object]:
    """Translate numerical LOOCV metrics into a readable model-quality judgement."""
    medape = _as_float_or_none(qc.get("Median_APE_%"))
    within2 = _as_float_or_none(qc.get("Within_2x_%"))
    r2 = _as_float_or_none(qc.get("R2_log"))
    n = int(_as_float_or_none(qc.get("N")) or 0)

    if n < 20:
        rating = "Insufficient calibration"
        conclusion = "Calibration set is too small for reliable response-factor transfer."
        recommendation = "Increase the number and chemical diversity of known-concentration calibration compounds."
    elif medape is not None and within2 is not None and r2 is not None and medape <= 30 and within2 >= 70 and r2 >= 0.5:
        rating = "Good"
        conclusion = "LOOCV supports approximate quantitative prediction for chemically similar candidates."
        recommendation = "Use predictions with the reported confidence/evidence columns; still verify key compounds with standards."
    elif medape is not None and within2 is not None and r2 is not None and medape <= 60 and within2 >= 50 and r2 >= 0:
        rating = "Semi-quantitative / caution"
        conclusion = "Model captures part of the response trend but individual concentrations can deviate substantially."
        recommendation = "Use as semi-quantitative ranking and prioritize high-confidence candidates for validation."
    else:
        rating = "Poor / screening only"
        conclusion = "LOOCV does not support reliable absolute concentration prediction. Results should be treated as exploratory screening estimates."
        recommendation = "Do not use predicted concentrations as final quantitative values. Expand calibration coverage, inspect outliers, stratify by reagent class/polarity, or use class-specific/internal-standard calibration."

    out = dict(qc)
    out["Model_rating"] = rating
    out["Model_conclusion"] = conclusion
    out["Recommendation"] = recommendation
    return out


def _prediction_reliability(pr: PredictionResult, qc: Dict[str, object]) -> Tuple[str, str]:
    """Row-level qualitative reliability derived from evidence plus global QC."""
    rating = str(qc.get("Model_rating", "")).lower()
    if pr.predicted_concentration is None:
        return "No prediction", "No valid ratio or no usable response-factor estimate."
    if pr.evidence == "exact_combo":
        level = "High"
        reason = "Exact Combo exists in the calibration set."
    elif pr.evidence == "global_median":
        level = "Very low"
        reason = "No close reference; global median response factor was used."
    else:
        sim = pr.similarity_mean if pr.similarity_mean is not None else 0.0
        if sim >= 0.70 and pr.nearest_count >= 5:
            level = "Medium-high"
            reason = f"Nearest calibration compounds are reasonably similar (mean similarity={sim:.2f})."
        elif sim >= 0.45 and pr.nearest_count >= 3:
            level = "Medium"
            reason = f"Moderate nearest-neighbour support (mean similarity={sim:.2f})."
        else:
            level = "Low"
            reason = f"Weak nearest-neighbour support (mean similarity={sim:.2f})."
    if "poor" in rating:
        if level in {"High", "Medium-high", "Medium"}:
            return "Exploratory", reason + " Global LOOCV QC is poor, so use as screening only."
        return level, reason + " Global LOOCV QC is poor."
    return level, reason


def _augment_cv_row(row: Dict[str, object]) -> Dict[str, object]:
    actual = _as_float_or_none(row.get("Actual_concentration"))
    pred = _as_float_or_none(row.get("Predicted_concentration_LOOCV"))
    out = dict(row)
    if actual is not None and pred is not None and actual > 0 and pred > 0:
        err_pct = (pred - actual) / actual * 100.0
        ape = abs(err_pct)
        pa = pred / actual
        fold = max(pa, 1.0 / pa)
        if fold <= 2:
            band = "within 2x"
        elif fold <= 5:
            band = "2-5x"
        elif fold <= 10:
            band = "5-10x"
        else:
            band = ">10x"
        out.update({
            "Error_%": err_pct,
            "Abs_Error_%": ape,
            "Pred/Actual": pa,
            "Fold_error": fold,
            "Accuracy_band": band,
            "log10_actual": math.log10(actual),
            "log10_predicted": math.log10(pred),
        })
    return out


def _diagnostic_arrays(cv_rows: Sequence[Dict[str, object]]) -> Tuple[np.ndarray, np.ndarray]:
    yt: List[float] = []
    yp: List[float] = []
    for row in cv_rows:
        actual = _as_float_or_none(row.get("Actual_concentration"))
        pred = _as_float_or_none(row.get("Predicted_concentration_LOOCV"))
        if actual is not None and pred is not None and actual > 0 and pred > 0:
            yt.append(float(actual))
            yp.append(float(pred))
    return np.array(yt, dtype=float), np.array(yp, dtype=float)


def _make_diagnostic_plots(cv_rows: Sequence[Dict[str, object]], out_xlsx: Path) -> List[Tuple[str, Path]]:
    """Create PNG diagnostics next to the workbook. Returned files are embedded when possible."""
    yt, yp = _diagnostic_arrays(cv_rows)
    if yt.size == 0:
        return []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return []

    plot_dir = Path(out_xlsx).with_suffix("").with_name(Path(out_xlsx).stem + "__diagnostics")
    plot_dir.mkdir(parents=True, exist_ok=True)
    outputs: List[Tuple[str, Path]] = []

    ratio = yp / yt
    fold = np.maximum(ratio, 1.0 / ratio)
    ape = np.abs(yp - yt) / yt * 100.0
    log_ratio = np.log10(ratio)

    # 1. Observed vs predicted on log scale.
    fig = plt.figure(figsize=(6.2, 5.0), dpi=160)
    ax = fig.add_subplot(111)
    ax.scatter(np.log10(yt), np.log10(yp), s=22, alpha=0.75)
    lo = float(min(np.log10(yt).min(), np.log10(yp).min()))
    hi = float(max(np.log10(yt).max(), np.log10(yp).max()))
    pad = max(0.05, (hi - lo) * 0.08)
    lo -= pad
    hi += pad
    xs = np.linspace(lo, hi, 100)
    ax.plot(xs, xs, linestyle="-", linewidth=1.0, label="ideal")
    ax.plot(xs, xs + math.log10(2), linestyle="--", linewidth=0.9, label="2x")
    ax.plot(xs, xs - math.log10(2), linestyle="--", linewidth=0.9)
    ax.set_xlabel("log10(actual concentration)")
    ax.set_ylabel("log10(LOOCV predicted concentration)")
    ax.set_title("LOOCV observed vs predicted")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    p1 = plot_dir / "loocv_observed_vs_predicted.png"
    fig.savefig(p1)
    plt.close(fig)
    outputs.append(("LOOCV observed vs predicted", p1))

    # 2. Log fold-error histogram.
    fig = plt.figure(figsize=(6.2, 4.2), dpi=160)
    ax = fig.add_subplot(111)
    ax.hist(log_ratio, bins=min(30, max(8, int(math.sqrt(len(log_ratio)) * 2))), alpha=0.85)
    ax.axvline(0, linestyle="-", linewidth=1.0, label="ideal")
    ax.axvline(math.log10(2), linestyle="--", linewidth=0.9, label="2x")
    ax.axvline(-math.log10(2), linestyle="--", linewidth=0.9)
    ax.set_xlabel("log10(predicted / actual)")
    ax.set_ylabel("Count")
    ax.set_title("LOOCV error distribution")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    p2 = plot_dir / "loocv_log_error_histogram.png"
    fig.savefig(p2)
    plt.close(fig)
    outputs.append(("LOOCV log error distribution", p2))

    # 3. Absolute percent error distribution, capped only for visualization.
    fig = plt.figure(figsize=(6.2, 4.2), dpi=160)
    ax = fig.add_subplot(111)
    cap = float(np.percentile(ape, 95)) if len(ape) >= 10 else float(np.max(ape))
    cap = max(100.0, min(cap, 1000.0))
    shown = np.minimum(ape, cap)
    ax.hist(shown, bins=min(30, max(8, int(math.sqrt(len(shown)) * 2))), alpha=0.85)
    ax.axvline(50, linestyle="--", linewidth=0.9, label="50%")
    ax.axvline(100, linestyle="--", linewidth=0.9, label="100%")
    ax.set_xlabel(f"Absolute percent error (%) (capped at {cap:.0f}% for display)")
    ax.set_ylabel("Count")
    ax.set_title("LOOCV absolute percent error")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(True, alpha=0.25)
    fig.tight_layout()
    p3 = plot_dir / "loocv_ape_histogram.png"
    fig.savefig(p3)
    plt.close(fig)
    outputs.append(("LOOCV absolute percent error", p3))

    # 4. Fold-error sorted curve.
    fig = plt.figure(figsize=(6.2, 4.2), dpi=160)
    ax = fig.add_subplot(111)
    sf = np.sort(fold)
    ax.plot(np.arange(1, len(sf) + 1), sf, marker="o", markersize=2, linewidth=1.0)
    ax.axhline(2, linestyle="--", linewidth=0.9, label="2x")
    ax.axhline(5, linestyle="--", linewidth=0.9, label="5x")
    ax.set_xlabel("LOOCV cases sorted by fold error")
    ax.set_ylabel("Fold error (symmetric)")
    ax.set_yscale("log")
    ax.set_title("LOOCV fold-error profile")
    ax.legend(frameon=False, fontsize=8)
    ax.grid(True, which="both", alpha=0.25)
    fig.tight_layout()
    p4 = plot_dir / "loocv_fold_error_profile.png"
    fig.savefig(p4)
    plt.close(fig)
    outputs.append(("LOOCV fold-error profile", p4))
    return outputs


def _add_evaluation_and_plot_sheets(wb, qc: Dict[str, object], cv_rows: Sequence[Dict[str, object]], out_xlsx: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws_eval = wb.create_sheet("Model_Evaluation", 0)
    ws_eval.append(["Item", "Value", "Interpretation / action"])
    rows = [
        ("Overall rating", qc.get("Model_rating", ""), qc.get("Model_conclusion", "")),
        ("Recommendation", "", qc.get("Recommendation", "")),
        ("N", qc.get("N", ""), "Number of LOOCV calibration rows used for model evaluation."),
        ("Median_APE_%", qc.get("Median_APE_%", ""), "Typical relative error. Lower is better; >60% means weak quantitative performance."),
        ("MAPE_%", qc.get("MAPE_%", ""), "Mean relative error; can be dominated by outliers or very low-concentration samples."),
        ("P90_APE_%", qc.get("P90_APE_%", ""), "90th percentile absolute percent error; shows tail risk."),
        ("R2_log", qc.get("R2_log", ""), "Log-scale coefficient of determination. Negative means worse than predicting the calibration mean."),
        ("Within_2x_%", qc.get("Within_2x_%", ""), "Percent of LOOCV predictions within a factor of 2 of actual concentration."),
        ("Within_5x_%", qc.get("Within_5x_%", ""), "Percent within a factor of 5."),
        ("Median_fold_error", qc.get("Median_fold_error", ""), "Symmetric median fold error; 1 is perfect."),
        ("Training rows used", qc.get("Training rows used", ""), "Calibration rows with valid ratio and concentration."),
        ("Target rows predicted", qc.get("Predicted rows", ""), "Rows for which a prediction was produced."),
    ]
    for row in rows:
        ws_eval.append(list(row))
    for cell in ws_eval[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws_eval.column_dimensions["A"].width = 28
    ws_eval.column_dimensions["B"].width = 24
    ws_eval.column_dimensions["C"].width = 95
    for row in ws_eval.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    # Standalone figures + embedded images when Pillow is available.
    plots = _make_diagnostic_plots(cv_rows, out_xlsx)
    ws_plots = wb.create_sheet("Diagnostic_Plots")
    ws_plots.append(["Diagnostic plots", "PNG file"])
    for title, path in plots:
        ws_plots.append([title, str(path)])
    ws_plots.column_dimensions["A"].width = 34
    ws_plots.column_dimensions["B"].width = 90
    try:
        from openpyxl.drawing.image import Image as XLImage
        anchors = ["A4", "K4", "A32", "K32"]
        for (title, path), anchor in zip(plots, anchors):
            img = XLImage(str(path))
            img.width = 520
            img.height = 390
            ws_plots.add_image(img, anchor)
    except Exception as e:
        ws_plots.append(["Image embedding skipped", f"Install Pillow if Excel image embedding is needed. Plot PNGs were still saved. ({e})"])


def loocv(train_rows: Sequence[PreparedRow], *, mode: str, k: int, ridge_lambda: float, element_keys: Sequence[str]) -> Tuple[Dict[str, object], List[Dict[str, object]]]:
    train = _valid_training(train_rows)
    rows_out: List[Dict[str, object]] = []
    y_true: List[float] = []
    y_pred: List[float] = []
    if len(train) < 3:
        return {"N": len(train), "warning": 'Fewer than three calibration records; LOOCV was not performed.'}, []
    for idx, row in enumerate(train):
        others = [r for j, r in enumerate(train) if j != idx]
        ek = list(element_keys)
        ridge = fit_ridge(others, ek, lambda_=ridge_lambda)
        pr = predict_one(row, others, ek, mode=mode, k=k, ridge_model=ridge)
        actual = float(row.concentration) if row.concentration is not None else None
        pred = pr.predicted_concentration
        if actual and pred and actual > 0 and pred > 0:
            y_true.append(actual)
            y_pred.append(pred)
        rows_out.append(
            {
                "Combo": row.combo,
                "Formula": row.formula,
                "Ratio": row.ratio,
                "Actual_concentration": actual,
                "Predicted_concentration_LOOCV": pred,
                "Error_%": ((pred - actual) / actual * 100.0) if actual and pred else None,
                "Evidence": pr.evidence,
                "Nearest_count": pr.nearest_count,
            }
        )
    return _metrics(y_true, y_pred), rows_out


def _write_rows_csv(path: Path, headers: Sequence[str], rows: Sequence[Sequence[object]]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(list(headers))
        for row in rows:
            w.writerow(list(row))
    return path


def _add_autofilter_style(wb, sheets: Sequence[object]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    key_fill = PatternFill("solid", fgColor="E2F0D9")
    for ws in sheets:
        if ws.max_row >= 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=False)
                    header = str(ws.cell(row=1, column=cell.column).value or "")
                    if any(k.lower() in header.lower() for k in ("concentration", "response_factor", "ratio", "area", "error", "mape")):
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = "0.0000"
            headers = {str(ws.cell(row=1, column=c).value or ""): c for c in range(1, ws.max_column + 1)}
            for h in ("Predicted_concentration", "Predicted_response_factor", "Confidence", "Evidence"):
                c = headers.get(h)
                if c:
                    ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="548235")
                    for r in range(2, ws.max_row + 1):
                        ws.cell(row=r, column=c).fill = key_fill
            for col_idx in range(1, ws.max_column + 1):
                header = str(ws.cell(row=1, column=col_idx).value or "")
                max_len = len(header)
                for r in range(2, min(ws.max_row, 300) + 1):
                    v = ws.cell(row=r, column=col_idx).value
                    if v is not None:
                        max_len = max(max_len, len(str(v)))
                cap = 60 if header in {"Combo", "Nearest_train_top3"} else 28
                ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(cap, max_len + 2))


def write_prediction_workbook(
    out_xlsx: Path,
    predictions: Sequence[PredictionResult],
    train_rows: Sequence[PreparedRow],
    qc: Dict[str, object],
    cv_rows: Sequence[Dict[str, object]],
    *,
    model_mode: str,
    ratio_note: str = "",
    output_csv: bool = True,
) -> Tuple[Path, Optional[Path]]:
    try:
        from openpyxl import Workbook
    except Exception as e:
        raise RuntimeError('Excel export requires openpyxl. Install it in the active environment: python -m pip install openpyxl') from e

    out_xlsx = Path(out_xlsx)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_pred = wb.active
    ws_pred.title = "Predictions"
    ws_cal = wb.create_sheet("Calibration")
    ws_qc = wb.create_sheet("Model_QC")
    ws_cv = wb.create_sheet("LOOCV")
    ws_readme = wb.create_sheet("Readme")

    pred_headers = [
        "Source_row",
        "Name",
        "Formula",
        "Combo",
        "Measured_ratio",
        "Predicted_concentration",
        "Predicted_response_factor",
        "Model",
        "Evidence",
        "Nearest_count",
        "Similarity_mean",
        "Nearest_train_top3",
        "Reliability_level",
        "Reliability_reason",
        "A_label",
        "B_label",
        "C_label",
        "Exact_mass",
        "Warnings",
    ]
    # Keep original target columns after prediction columns for traceability.
    orig_headers: List[str] = []
    for pr in predictions:
        for h in pr.row.raw.keys():
            if h not in orig_headers:
                orig_headers.append(h)
    ws_pred.append(pred_headers + [f"orig__{h}" for h in orig_headers])
    pred_csv_rows: List[List[object]] = []
    for pr in predictions:
        parts = pr.row.combo_parts
        reliability_level, reliability_reason = _prediction_reliability(pr, qc)
        row_vals = [
            pr.row.source_index,
            pr.row.name,
            pr.row.formula,
            pr.row.combo,
            pr.row.ratio,
            pr.predicted_concentration,
            pr.predicted_correction_factor,
            pr.method,
            pr.evidence,
            pr.nearest_count,
            pr.similarity_mean,
            pr.nearest_summary,
            reliability_level,
            reliability_reason,
            parts.get("A", {}).get("label", ""),
            parts.get("B", {}).get("label", ""),
            parts.get("C", {}).get("label", ""),
            pr.row.exact_mass,
            "; ".join(pr.warnings + pr.row.warnings),
        ] + [pr.row.raw.get(h, "") for h in orig_headers]
        ws_pred.append(row_vals)
        pred_csv_rows.append(row_vals)

    cal_headers = [
        "Source_row",
        "Name",
        "Formula",
        "Combo",
        "Measured_ratio",
        "Actual_concentration",
        "Response_factor(conc/ratio)",
        "log_Response_factor",
        "A_label",
        "B_label",
        "C_label",
        "Exact_mass",
        "Warnings",
    ]
    ws_cal.append(cal_headers)
    for r in train_rows:
        parts = r.combo_parts
        ws_cal.append(
            [
                r.source_index,
                r.name,
                r.formula,
                r.combo,
                r.ratio,
                r.concentration,
                r.correction_factor,
                r.log_correction_factor,
                parts.get("A", {}).get("label", ""),
                parts.get("B", {}).get("label", ""),
                parts.get("C", {}).get("label", ""),
                r.exact_mass,
                "; ".join(r.warnings),
            ]
        )

    ws_qc.append(["Metric", "Value"])
    for k, v in qc.items():
        ws_qc.append([k, v])

    cv_headers = [
        "Combo",
        "Formula",
        "Ratio",
        "Actual_concentration",
        "Predicted_concentration_LOOCV",
        "Error_%",
        "Abs_Error_%",
        "Pred/Actual",
        "Fold_error",
        "Accuracy_band",
        "log10_actual",
        "log10_predicted",
        "Evidence",
        "Nearest_count",
    ]
    ws_cv.append(cv_headers)
    cv_rows_aug = [_augment_cv_row(row) for row in cv_rows]
    for row in cv_rows_aug:
        ws_cv.append([row.get(h, "") for h in cv_headers])

    readme_lines = [
        ["Purpose", "Predict actual concentration of large candidate set using a smaller calibration set with known concentration and measured area/internal-standard ratio."],
        ["Formula", "response_factor = actual_concentration / measured_ratio; predicted_concentration = measured_ratio * predicted_response_factor."],
        ["Model", model_mode],
        ["Ratio scale", ratio_note or "Calibration and prediction tables must use the same ratio scale (e.g. Area/IS or Area/IS*100%)."],
        ["Combo usage", "Combo is parsed into A/B/C reagent labels and formulas; exact combo match is preferred; otherwise formula/provenance similarity is used."],
        ["Warning", "Predictions are response-corrected estimates, not substitute for standards. Use calibration coverage and LOOCV metrics to judge reliability."],
    ]
    ws_readme.append(["Item", "Description"])
    for line in readme_lines:
        ws_readme.append(line)

    _add_evaluation_and_plot_sheets(wb, qc, cv_rows_aug if 'cv_rows_aug' in locals() else cv_rows, out_xlsx)
    # Style all sheets after adding evaluation and diagnostic sheets.
    _add_autofilter_style(wb, list(wb.worksheets))
    # Open on the evaluation sheet by default so the user sees the verdict first.
    try:
        wb.active = wb.index(wb["Model_Evaluation"])
    except Exception:
        wb.active = wb.index(ws_pred)
    wb.save(str(out_xlsx))

    out_csv: Optional[Path] = None
    if output_csv:
        out_csv = out_xlsx.with_suffix("").with_name(out_xlsx.stem + "__predictions.csv")
        _write_rows_csv(out_csv, pred_headers + [f"orig__{h}" for h in orig_headers], pred_csv_rows)
    return out_xlsx, out_csv


def run_response_prediction(
    calibration_path: Path,
    target_path: Path,
    output_xlsx: Path,
    *,
    calibration_sheet: str = "",
    target_sheet: str = "",
    combo_col: str = "",
    formula_col: str = "",
    ratio_col: str = "",
    concentration_col: str = "",
    target_ratio_col: str = "",
    model_mode: str = "blend",
    k_neighbors: int = 15,
    ridge_lambda: float = 1.0,
    output_csv: bool = True,
) -> RunReport:
    warnings: List[str] = []
    cal_table = read_table(Path(calibration_path), calibration_sheet)
    tar_table = read_table(Path(target_path), target_sheet)

    cal_cfg = guess_columns(cal_table, need_concentration=True)
    tar_cfg = guess_columns(tar_table, need_concentration=False)
    # User overrides. Use same overrides for both if target-specific empty.
    if combo_col:
        cal_cfg.combo_col = combo_col
        tar_cfg.combo_col = combo_col if combo_col in tar_table.headers else tar_cfg.combo_col
    if formula_col:
        cal_cfg.formula_col = formula_col
        tar_cfg.formula_col = formula_col if formula_col in tar_table.headers else tar_cfg.formula_col
    if ratio_col:
        cal_cfg.ratio_col = ratio_col
        if not target_ratio_col and ratio_col in tar_table.headers:
            tar_cfg.ratio_col = ratio_col
    if target_ratio_col:
        tar_cfg.ratio_col = target_ratio_col
    if concentration_col:
        cal_cfg.concentration_col = concentration_col

    cal_rows, ws1 = prepare_rows(cal_table, cal_cfg, is_training=True)
    tar_rows, ws2 = prepare_rows(tar_table, tar_cfg, is_training=False)
    warnings.extend(ws1)
    warnings.extend(ws2)

    train_used = _valid_training(cal_rows)
    element_keys = _element_keys(list(train_used) + list(tar_rows))
    mode_code = MODEL_LABEL_TO_CODE.get(model_mode, model_mode or "blend")
    k = max(1, int(k_neighbors or 15))
    ridge = fit_ridge(train_used, element_keys, lambda_=float(ridge_lambda or 1.0))
    predictions: List[PredictionResult] = []
    for row in tar_rows:
        predictions.append(predict_one(row, train_used, element_keys, mode=mode_code, k=k, ridge_model=ridge))

    qc, cv_rows = loocv(train_used, mode=mode_code, k=k, ridge_lambda=float(ridge_lambda or 1.0), element_keys=element_keys)
    qc = dict(qc)
    qc.update(
        {
            "Calibration file": str(calibration_path),
            "Calibration sheet": cal_table.sheet_name,
            "Target file": str(target_path),
            "Target sheet": tar_table.sheet_name,
            "Calibration combo column": cal_cfg.combo_col,
            "Calibration ratio column": cal_cfg.ratio_col,
            "Calibration concentration column": cal_cfg.concentration_col,
            "Target combo column": tar_cfg.combo_col,
            "Target ratio column": tar_cfg.ratio_col,
            "Training rows total": len(cal_rows),
            "Training rows used": len(train_used),
            "Target rows total": len(tar_rows),
            "Predicted rows": sum(1 for p in predictions if p.predicted_concentration is not None),
            "Model mode": mode_code,
            "K neighbors": k,
            "Ridge lambda": float(ridge_lambda or 1.0),
        }
    )
    qc = evaluate_model_qc(qc)

    out_xlsx, out_csv = write_prediction_workbook(
        Path(output_xlsx),
        predictions,
        cal_rows,
        qc,
        cv_rows,
        model_mode=mode_code,
        output_csv=bool(output_csv),
    )
    return RunReport(
        output_xlsx=out_xlsx,
        output_csv=out_csv,
        n_train_total=len(cal_rows),
        n_train_used=len(train_used),
        n_target_total=len(tar_rows),
        n_predicted=sum(1 for p in predictions if p.predicted_concentration is not None),
        warnings=warnings,
        qc=qc,
    )
