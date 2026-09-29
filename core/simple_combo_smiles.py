"""Simple A/B/C + fixed CORE mapped-SMILES product builder.

One compact workbook is used instead of separate Structures/ReactionRules
sheets.  The input layout is four columns per component group:

    A:D   = A_ID / A_Formula / A_Mass / A_SMILES
    E:H   = B_ID / B_Formula / B_Mass / B_SMILES
    I:L   = C_ID / C_Formula / C_Mass / C_SMILES

The fixed CORE mapped SMILES is entered once in the GUI.  A/B/C fragments and
CORE must carry matching dummy maps (for example CORE [*:1]/[*:2]/[*:3], A
[*:1], B [*:2], C [*:3]).  The module enumerates A\u00D7B\u00D7C, stitches each product,
checks the existing formula rule, and writes one Combo\u2192Product_SMILES master
sheet that can be joined to both calibration and target tables by Combo.
"""
from __future__ import annotations

import csv
import hashlib
import itertools
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

from .chemistry import format_formula_hill, monoisotopic_mass, parse_formula
from .smiles_builder import RDKIT_AVAILABLE, RDKIT_ERROR, _draw_product, _dummy_stitch_candidate
from .smiles_diagnostics import (
    DiagnosticsReport,
    SmilesDiagnostic,
    diagnose_smiles,
    diagnostics_text,
    headers as diagnostic_headers,
    row_values as diagnostic_row_values,
    validate_pairing,
    normalize_smiles,
    write_txt as write_diagnostics_txt,
    write_xlsx as write_diagnostics_xlsx,
)


@dataclass
class SimpleComponent:
    role: str
    item_index: int
    excel_row: int
    item_id: str
    formula: str
    counts: Dict[str, int]
    exact_mass: float
    smiles: str
    normalized_smiles: str = ""
    canonical_smiles: str = ""
    smiles_hash: str = ""
    observed_heavy_formula: str = ""
    expected_heavy_formula: str = ""
    heavy_formula_delta: str = ""
    heavy_formula_check: str = "unknown"
    heavy_delta_counts: Dict[str, int] = field(default_factory=dict)
    integrity_flags: List[str] = field(default_factory=list)


@dataclass
class SimpleProductResult:
    source_sheet: str
    no: int
    combo: str
    a: SimpleComponent
    b: SimpleComponent
    c: SimpleComponent
    expected_formula: str
    expected_exact_mass: Optional[float]
    duplicate_index: int = 1
    duplicate_count: int = 1
    status: str = "FAILED"
    product_smiles: str = ""
    product_formula: str = ""
    product_exact_mass: Optional[float] = None
    formula_match: str = ""
    inchikey: str = ""
    structure_png: str = ""
    product_formula_delta: str = ""
    suspected_component: str = ""
    warnings: List[str] = field(default_factory=list)


@dataclass
class SimpleBuildReport:
    output_xlsx: Path
    output_csv: Optional[Path]
    image_dir: Optional[Path]
    n_sheets: int
    n_total: int
    n_success: int
    n_formula_match: int
    n_failed: int
    warnings: List[str]
    results: List[SimpleProductResult] = field(default_factory=list)
    diagnostics: Optional[DiagnosticsReport] = None


def _text(v: object) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s[1:].strip() if s.startswith("'") else s


def _safe_file_part(text: str, max_len: int = 70) -> str:
    s = re.sub(r"[^0-9A-Za-z._-]+", "_", str(text or "").strip())
    return (s.strip("._-") or "product")[:max_len]


def _add_counts(*parts: Dict[str, int]) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for part in parts:
        for element, count in part.items():
            out[element] = out.get(element, 0) + int(count)
    return out


def _without_h(counts: Dict[str, int]) -> Dict[str, int]:
    return {str(k): int(v) for k, v in counts.items() if str(k) != "H" and int(v) != 0}


def _delta_counts(actual: Dict[str, int], expected: Dict[str, int]) -> Dict[str, int]:
    keys = sorted(set(actual) | set(expected), key=lambda x: (x != "C", x != "H", x))
    return {k: int(actual.get(k, 0)) - int(expected.get(k, 0)) for k in keys if int(actual.get(k, 0)) != int(expected.get(k, 0))}


def _delta_text(delta: Dict[str, int]) -> str:
    if not delta:
        return "0"
    return "; ".join(f"{el}:{value:+d}" for el, value in delta.items())


def _component_expected_heavy(role: str, formula_counts: Dict[str, int]) -> Tuple[Dict[str, int], str]:
    expected = _without_h(formula_counts)
    note = "Listed Formula heavy atoms"
    if str(role).upper() == "B":
        # The user's reaction rule states that B loses N2 before/while coupling.
        expected["N"] = int(expected.get("N", 0)) - 2
        if expected["N"] == 0:
            expected.pop("N", None)
        note = "Listed Formula heavy atoms after B loss N2"
    return expected, note


def _analyze_component_smiles(role: str, smiles: str, formula_counts: Dict[str, int]):
    normalized, _changes = normalize_smiles(smiles)
    expected, _note = _component_expected_heavy(role, formula_counts)
    expected_formula = format_formula_hill(expected) if expected else ""
    if not RDKIT_AVAILABLE or not normalized:
        return normalized, "", "", "", expected_formula, "", "unknown", {}, []
    try:
        from rdkit import Chem
        mol = Chem.MolFromSmiles(normalized)
        if mol is None:
            return normalized, "", "", "", expected_formula, "", "unknown", {}, []
        Chem.SanitizeMol(mol)
        canonical = Chem.MolToSmiles(mol, canonical=True, isomericSmiles=True)
        observed: Dict[str, int] = {}
        for atom in mol.GetAtoms():
            if atom.GetAtomicNum() <= 0:
                continue
            symbol = str(atom.GetSymbol())
            if symbol == "H":
                continue
            observed[symbol] = observed.get(symbol, 0) + 1
        observed_formula = format_formula_hill(observed) if observed else ""
        delta = _delta_counts(observed, expected)
        check = "yes" if not delta else "no"
        digest = hashlib.sha256(canonical.encode("utf-8", errors="ignore")).hexdigest()[:16]
        flags: List[str] = []
        if delta:
            flags.append("MAPPED_SMILES_HEAVY_FORMULA_MISMATCH")
        return normalized, canonical, digest, observed_formula, expected_formula, _delta_text(delta), check, delta, flags
    except Exception:
        return normalized, "", "", "", expected_formula, "", "unknown", {}, []


def _mark_duplicate_component_smiles(rows: Sequence[SimpleComponent], warnings: List[str], sheet_name: str) -> None:
    by_hash: Dict[str, List[SimpleComponent]] = {}
    for comp in rows:
        if comp.smiles_hash:
            by_hash.setdefault(comp.smiles_hash, []).append(comp)
    for digest, group in by_hash.items():
        if len(group) < 2:
            continue
        formulas = {x.formula for x in group}
        ids = ", ".join(f"{x.item_id}(row {x.excel_row}, {x.formula})" for x in group)
        if len(formulas) > 1:
            flag = "DUPLICATE_MAPPED_SMILES_DIFFERENT_FORMULA"
            for x in group:
                if flag not in x.integrity_flags:
                    x.integrity_flags.append(flag)
            warnings.append(
                f"{sheet_name}: {group[0].role} rows have identical mapped SMILES but different Formula: {ids}. "
                "This can make a later component behave like an earlier one (for example B2 still using B1 structure)."
            )
        else:
            flag = "DUPLICATE_MAPPED_SMILES"
            for x in group:
                if flag not in x.integrity_flags:
                    x.integrity_flags.append(flag)
            warnings.append(f"{sheet_name}: duplicate {group[0].role} mapped SMILES detected: {ids}")


def _product_formula_delta(actual_formula: str, expected_counts: Dict[str, int]) -> Tuple[str, Dict[str, int]]:
    if not actual_formula:
        return "", {}
    try:
        actual = {str(k): int(v) for k, v in parse_formula(actual_formula).items()}
    except Exception:
        return "", {}
    delta = _delta_counts(actual, expected_counts)
    return _delta_text(delta), delta


def _suspect_components(product_delta: Dict[str, int], components: Sequence[SimpleComponent]) -> str:
    if not product_delta:
        return ""
    heavy_product_delta = {k: v for k, v in product_delta.items() if k != "H"}
    exact: List[str] = []
    weak: List[str] = []
    for comp in components:
        if comp.heavy_formula_check == "no":
            weak.append(f"{comp.role}(row {comp.excel_row}, {comp.item_id})")
            if comp.heavy_delta_counts == heavy_product_delta:
                exact.append(f"{comp.role}(row {comp.excel_row}, {comp.item_id})")
        if "DUPLICATE_MAPPED_SMILES_DIFFERENT_FORMULA" in comp.integrity_flags:
            label = f"{comp.role}(row {comp.excel_row}, {comp.item_id}; duplicate SMILES)"
            if label not in weak:
                weak.append(label)
    return ", ".join(exact or weak)


def _apply_existing_formula_rule(base: Dict[str, int], *, only_formula: bool) -> Optional[Dict[str, int]]:
    d = dict(base)
    d["H"] = int(d.get("H", 0)) - 1
    d["N"] = int(d.get("N", 0)) - 2
    d["O"] = int(d.get("O", 0)) + 1
    d["P"] = int(d.get("P", 0)) + 1
    if only_formula:
        d = {k: int(d.get(k, 0)) for k in ("C", "H", "N", "O", "P")}
    if any(int(v) < 0 for v in d.values()):
        return None
    return d


def _read_group(ws, *, role: str, id_col: int, formula_col: int, smiles_col: int, data_start_row: int):
    rows: List[SimpleComponent] = []
    warnings: List[str] = []
    for excel_row in range(int(data_start_row), int(ws.max_row or data_start_row) + 1):
        item_id = _text(ws.cell(excel_row, id_col).value)
        formula = _text(ws.cell(excel_row, formula_col).value)
        smiles = _text(ws.cell(excel_row, smiles_col).value)
        if not item_id and not formula and not smiles:
            continue
        if not formula:
            warnings.append(f"{ws.title}!row {excel_row}: {role} Formula is empty; row skipped")
            continue
        try:
            counts = {str(k): int(v) for k, v in parse_formula(formula).items()}
            hill = format_formula_hill(counts)
            mass = float(monoisotopic_mass(counts))
        except Exception as e:
            warnings.append(f"{ws.title}!row {excel_row}: invalid {role} Formula {formula!r}: {e}")
            continue
        idx = len(rows) + 1
        item_id = item_id or f"{role}{idx}"
        if not smiles:
            warnings.append(f"{ws.title}!row {excel_row}: {role}/{item_id} has no SMILES; related products will fail")
        (normalized, canonical, digest, observed_heavy, expected_heavy, heavy_delta,
         heavy_check, heavy_delta_counts, integrity_flags) = _analyze_component_smiles(role, smiles, counts)
        comp = SimpleComponent(
            role, idx, excel_row, item_id, hill, counts, mass, smiles,
            normalized_smiles=normalized, canonical_smiles=canonical, smiles_hash=digest,
            observed_heavy_formula=observed_heavy, expected_heavy_formula=expected_heavy,
            heavy_formula_delta=heavy_delta, heavy_formula_check=heavy_check,
            heavy_delta_counts=heavy_delta_counts, integrity_flags=list(integrity_flags),
        )
        if heavy_check == "no":
            warnings.append(
                f"{ws.title}!row {excel_row}: {role}/{item_id} Formula={hill}, but mapped SMILES heavy atoms are "
                f"{observed_heavy or '-'}; expected {expected_heavy or '-'} "
                f"({'after N2 loss' if role == 'B' else 'same heavy atoms as listed Formula'}), delta={heavy_delta}."
            )
        rows.append(comp)
    _mark_duplicate_component_smiles(rows, warnings, ws.title)
    return rows, warnings


def load_component_sheets(
    workbook_path: Path,
    *,
    sheet_name: str = "",
    data_start_row: int = 2,
    include_hidden_sheets: bool = False,
):
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise RuntimeError("openpyxl is required. Please install openpyxl.") from e

    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(str(workbook_path))
    wb = load_workbook(str(workbook_path), read_only=True, data_only=True)
    out = []
    warnings: List[str] = []
    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet not found: {sheet_name}")
            worksheets = [wb[sheet_name]]
        else:
            worksheets = list(wb.worksheets)
        for ws in worksheets:
            if not include_hidden_sheets and getattr(ws, "sheet_state", "visible") != "visible":
                continue
            a_rows, wa = _read_group(ws, role="A", id_col=1, formula_col=2, smiles_col=4, data_start_row=data_start_row)
            b_rows, wbw = _read_group(ws, role="B", id_col=5, formula_col=6, smiles_col=8, data_start_row=data_start_row)
            c_rows, wc = _read_group(ws, role="C", id_col=9, formula_col=10, smiles_col=12, data_start_row=data_start_row)
            warnings.extend(wa + wbw + wc)
            if not a_rows or not b_rows or not c_rows:
                warnings.append(
                    f"Sheet {ws.title} skipped: A={len(a_rows)}, B={len(b_rows)}, C={len(c_rows)}. "
                    "Expected A:D, E:H, I:L = ID/Formula/Mass/SMILES."
                )
                continue
            out.append((ws.title, a_rows, b_rows, c_rows))
    finally:
        wb.close()
    if not out:
        raise ValueError("No valid worksheet found. Expected A:D, E:H, I:L = ID/Formula/Mass/SMILES.")
    return out, warnings


def _build_input_diagnostics(
    sheets,
    *,
    core_smiles: str,
) -> Tuple[DiagnosticsReport, SmilesDiagnostic, Dict[Tuple[str, str, int], SmilesDiagnostic]]:
    diagnostics: List[SmilesDiagnostic] = []
    core_diag = diagnose_smiles(
        "CORE", core_smiles, item_id="FIXED_CORE",
        source_label="GUI fixed CORE mapped SMILES text box", expected_maps=(1, 2, 3),
    )
    diagnostics.append(core_diag)
    lookup: Dict[Tuple[str, str, int], SmilesDiagnostic] = {}
    for sheet_name, a_rows, b_rows, c_rows in sheets:
        for role, rows, map_no in (("A", a_rows, 1), ("B", b_rows, 2), ("C", c_rows, 3)):
            for comp in rows:
                d = diagnose_smiles(
                    role, comp.smiles, item_id=comp.item_id, source_sheet=sheet_name,
                    excel_row=comp.excel_row, source_label=f"{sheet_name}!row {comp.excel_row}",
                    expected_maps=(map_no,),
                )
                diagnostics.append(d)
                lookup[(sheet_name, role, int(comp.excel_row))] = d
    ge, gw = validate_pairing(diagnostics)
    return DiagnosticsReport(diagnostics, ge, gw), core_diag, lookup


def diagnose_simple_combo_inputs(
    component_workbook: Path,
    *,
    core_smiles: str,
    sheet_name: str = "",
    data_start_row: int = 2,
    output_xlsx: Optional[Path] = None,
    output_txt: Optional[Path] = None,
    include_raw_smiles: bool = True,
) -> DiagnosticsReport:
    """Audit all private mapped SMILES locally without product enumeration."""
    if not RDKIT_AVAILABLE:
        raise RuntimeError(
            "RDKit is required. Install in the same environment: conda install -c conda-forge rdkit. "
            + str(RDKIT_ERROR)
        )
    sheets, load_warnings = load_component_sheets(
        Path(component_workbook), sheet_name=sheet_name, data_start_row=data_start_row
    )
    report, _, _ = _build_input_diagnostics(sheets, core_smiles=core_smiles)
    for w in load_warnings:
        if w not in report.global_warnings:
            report.global_warnings.append(w)
    if output_xlsx:
        write_diagnostics_xlsx(report, Path(output_xlsx), include_raw=include_raw_smiles)
    if output_txt:
        write_diagnostics_txt(report, Path(output_txt), include_raw=False)
    return report


def _diagnostic_messages(d: SmilesDiagnostic) -> List[str]:
    out: List[str] = []
    prefix = f"{d.role} [{d.location}]"
    out += [f"{prefix}: ERROR: {x}" for x in d.errors]
    out += [f"{prefix}: WARNING: {x}" for x in d.warnings]
    out += [f"{prefix}: FIX: {x}" for x in d.suggestions]
    return out


def _combo(a: SimpleComponent, b: SimpleComponent, c: SimpleComponent) -> str:
    return (
        f"A#{a.item_index}:{a.item_id}:{a.formula} | "
        f"B#{b.item_index}:{b.item_id}:{b.formula} | "
        f"C#{c.item_index}:{c.item_id}:{c.formula}"
    )


def _enumerate_sheet(
    sheet_name: str,
    a_rows: Sequence[SimpleComponent],
    b_rows: Sequence[SimpleComponent],
    c_rows: Sequence[SimpleComponent],
    *,
    core_smiles: str,
    only_formula: bool,
    core_diagnostic: Optional[SmilesDiagnostic] = None,
    component_diagnostics: Optional[Dict[Tuple[str, str, int], SmilesDiagnostic]] = None,
    progress: Optional[Callable[[str], None]] = None,
):
    out: List[SimpleProductResult] = []
    total = len(a_rows) * len(b_rows) * len(c_rows)
    component_diagnostics = component_diagnostics or {}
    core_diag = core_diagnostic or diagnose_smiles(
        "CORE", core_smiles, item_id="FIXED_CORE",
        source_label="GUI fixed CORE mapped SMILES text box", expected_maps=(1, 2, 3),
    )

    for n, (a, b, c) in enumerate(itertools.product(a_rows, b_rows, c_rows), start=1):
        final_counts = _apply_existing_formula_rule(_add_counts(a.counts, b.counts, c.counts), only_formula=only_formula)
        if final_counts is None:
            out.append(SimpleProductResult(
                sheet_name, n, _combo(a, b, c), a, b, c, "", None,
                status="INVALID_FORMULA", warnings=["Formula rule produced a negative element count"],
            ))
            continue
        formula = format_formula_hill(final_counts)
        mass = float(monoisotopic_mass(final_counts))
        r = SimpleProductResult(sheet_name, n, _combo(a, b, c), a, b, c, formula, mass)

        role_diags: List[SmilesDiagnostic] = [core_diag]
        for role, comp, map_no in (("A", a, 1), ("B", b, 2), ("C", c, 3)):
            d = component_diagnostics.get((sheet_name, role, int(comp.excel_row)))
            if d is None:
                d = diagnose_smiles(
                    role, comp.smiles, item_id=comp.item_id, source_sheet=sheet_name,
                    excel_row=comp.excel_row, source_label=f"{sheet_name}!row {comp.excel_row}",
                    expected_maps=(map_no,),
                )
            role_diags.append(d)

        bad = [d for d in role_diags if d.fatal]
        if bad:
            r.status = "INPUT_ERROR"
            for d in bad:
                r.warnings.extend(_diagnostic_messages(d))
            for d in role_diags:
                if not d.fatal and d.normalization_changes:
                    r.warnings.append(f"{d.role} [{d.location}]: AUTO_NORMALIZED: " + "; ".join(d.normalization_changes))
            if progress and (n == 1 or n == total or n % 50 == 0):
                progress(f"[{sheet_name}] {n}/{total}: INPUT_ERROR | {bad[0].one_line()}")
            out.append(r)
            continue

        # Use the immutable source component selected by this Cartesian-product
        # iteration directly.  Diagnostics are only a validation layer; they
        # are not used as a shared structure cache.  This prevents any chance
        # that B2 could accidentally reuse B1 because of a stale lookup.
        norm = {
            "CORE": core_diag.normalized_smiles,
            "A": a.normalized_smiles,
            "B": b.normalized_smiles,
            "C": c.normalized_smiles,
        }
        candidates, ws = _dummy_stitch_candidate(
            norm, ("CORE", "A", "B", "C"), formula,
        )
        for d in role_diags:
            if d.normalization_changes:
                r.warnings.append(f"{d.role} [{d.location}]: AUTO_NORMALIZED: " + "; ".join(d.normalization_changes))
        r.warnings.extend(ws)
        if candidates:
            selected = candidates[0]
            r.product_smiles = selected.smiles
            r.product_formula = selected.formula
            r.product_exact_mass = selected.exact_mass
            r.formula_match = selected.formula_match
            r.inchikey = selected.inchikey
            r.product_formula_delta, product_delta_counts = _product_formula_delta(selected.formula, final_counts)
            r.suspected_component = _suspect_components(product_delta_counts, (a, b, c))
            r.status = "SUCCESS" if selected.formula_match != "no" else "CHECK_FORMULA"
            if selected.formula_match == "no":
                r.warnings.append(
                    f"Product Formula delta (Product - Expected) = {r.product_formula_delta or 'unknown'}."
                )
                if r.suspected_component:
                    r.warnings.append(
                        "Probable component source: " + r.suspected_component + ". "
                        "Check that this row's mapped SMILES was not copied from another component."
                    )
                for comp in (a, b, c):
                    if comp.heavy_formula_check == "no":
                        r.warnings.append(
                            f"{comp.role} row {comp.excel_row}/{comp.item_id}: listed Formula={comp.formula}; "
                            f"mapped-SMILES heavy formula={comp.observed_heavy_formula or '-'}; "
                            f"expected={comp.expected_heavy_formula or '-'}; delta={comp.heavy_formula_delta}."
                        )
                    if "DUPLICATE_MAPPED_SMILES_DIFFERENT_FORMULA" in comp.integrity_flags:
                        r.warnings.append(
                            f"{comp.role} row {comp.excel_row}/{comp.item_id}: mapped SMILES hash {comp.smiles_hash} "
                            "is shared by rows with different Formula; this is a likely stale-copy error."
                        )
        elif not r.warnings:
            r.warnings.append("Mapped-dummy stitching produced no product for an unknown reason")
        if progress and (n == 1 or n == total or n % 50 == 0):
            detail = r.warnings[0] if r.warnings and not r.product_smiles else r.combo[:90]
            progress(f"[{sheet_name}] {n}/{total}: {r.status} | {detail}")
        out.append(r)

    formula_counts: Dict[str, int] = {}
    for r in out:
        if r.expected_formula:
            formula_counts[r.expected_formula] = formula_counts.get(r.expected_formula, 0) + 1
    seen: Dict[str, int] = {}
    for r in out:
        if not r.expected_formula:
            continue
        seen[r.expected_formula] = seen.get(r.expected_formula, 0) + 1
        r.duplicate_index = seen[r.expected_formula]
        r.duplicate_count = formula_counts.get(r.expected_formula, 1)
    return out


def _headers(include_fragment_smiles: bool, embed_images: bool):
    h = [
        "Source_Sheet", "No", "Formula", "Exact_Mass", "Combo",
        "A_ID", "B_ID", "C_ID",
        "A_Source_Row", "B_Source_Row", "C_Source_Row",
        "A_SMILES_Hash", "B_SMILES_Hash", "C_SMILES_Hash",
        "A_Heavy_Check", "B_Heavy_Check", "C_Heavy_Check",
        "Duplicate_Index", "Duplicate_Count",
        "Product_SMILES", "Product_Formula", "Product_Formula_Delta", "Suspected_Component",
        "Product_Exact_Mass", "Formula_Match", "Structure_Status", "InChIKey", "Structure_PNG", "Warnings",
    ]
    if include_fragment_smiles:
        h += ["A_SMILES", "B_SMILES", "C_SMILES"]
    if embed_images:
        h += ["2D_Preview"]
    return h


def _row_values(r: SimpleProductResult, include_fragment_smiles: bool, embed_images: bool):
    row = [
        r.source_sheet, r.no, r.expected_formula, r.expected_exact_mass, r.combo,
        r.a.item_id, r.b.item_id, r.c.item_id,
        r.a.excel_row, r.b.excel_row, r.c.excel_row,
        r.a.smiles_hash, r.b.smiles_hash, r.c.smiles_hash,
        r.a.heavy_formula_check, r.b.heavy_formula_check, r.c.heavy_formula_check,
        r.duplicate_index, r.duplicate_count,
        r.product_smiles, r.product_formula, r.product_formula_delta, r.suspected_component,
        r.product_exact_mass, r.formula_match, r.status, r.inchikey, r.structure_png,
        "; ".join(dict.fromkeys(r.warnings)),
    ]
    if include_fragment_smiles:
        row += [r.a.smiles, r.b.smiles, r.c.smiles]
    if embed_images:
        row += [""]
    return row


def _write_csv(results, path: Path, *, include_fragment_smiles: bool):
    h = _headers(include_fragment_smiles, False)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(h)
        for r in results:
            w.writerow(_row_values(r, include_fragment_smiles, False))


def _write_xlsx(
    results,
    path: Path,
    *,
    core_smiles: str,
    include_fragment_smiles: bool,
    embed_images: bool,
    embed_image_limit: int,
    diagnostics: Optional[DiagnosticsReport] = None,
):
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as e:
        raise RuntimeError("openpyxl is required to write the result workbook") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Combo_SMILES_Master"
    headers = _headers(include_fragment_smiles, embed_images)
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    image_col = headers.index("2D_Preview") + 1 if embed_images else 0
    embedded = 0
    for ridx, r in enumerate(results, start=2):
        ws.append(_row_values(r, include_fragment_smiles, embed_images))
        ws.cell(ridx, headers.index("Exact_Mass") + 1).number_format = "0.000000"
        ws.cell(ridx, headers.index("Product_Exact_Mass") + 1).number_format = "0.000000"
        if embed_images and r.structure_png and Path(r.structure_png).exists() and (embed_image_limit <= 0 or embedded < embed_image_limit):
            try:
                img = XLImage(r.structure_png)
                img.width, img.height = 300, 185
                img.anchor = ws.cell(ridx, image_col).coordinate
                ws.add_image(img)
                ws.row_dimensions[ridx].height = 145
                embedded += 1
            except Exception:
                pass

    width_map = {
        "Source_Sheet": 18, "No": 9, "Formula": 22, "Exact_Mass": 16, "Combo": 72,
        "A_ID": 16, "B_ID": 16, "C_ID": 16,
        "A_Source_Row": 12, "B_Source_Row": 12, "C_Source_Row": 12,
        "A_SMILES_Hash": 19, "B_SMILES_Hash": 19, "C_SMILES_Hash": 19,
        "A_Heavy_Check": 14, "B_Heavy_Check": 14, "C_Heavy_Check": 14,
        "Duplicate_Index": 14, "Duplicate_Count": 14,
        "Product_SMILES": 82, "Product_Formula": 22, "Product_Formula_Delta": 24,
        "Suspected_Component": 34, "Product_Exact_Mass": 18, "Formula_Match": 14,
        "Structure_Status": 18, "InChIKey": 30, "Structure_PNG": 42, "Warnings": 75,
        "A_SMILES": 60, "B_SMILES": 60, "C_SMILES": 60, "2D_Preview": 42,
    }
    for i, head in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width_map.get(head, 24)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # One row per A/B/C source component.  This is deliberately separate
    # from the 1000-row product table so stale copied SMILES can be found
    # without inspecting every combination.
    audit = wb.create_sheet("Component_SMILES_Audit")
    audit_headers = [
        "Source_Sheet", "Role", "Item_Index", "Excel_Row", "ID", "Listed_Formula",
        "Expected_Heavy_Formula", "Mapped_SMILES_Heavy_Formula", "Heavy_Formula_Delta",
        "Heavy_Formula_Check", "SMILES_Hash", "Canonical_Mapped_SMILES",
        "Integrity_Flags", "Probable_Issue",
    ]
    audit.append(audit_headers)
    for cell in audit[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    seen_components = set()
    for r in results:
        for comp in (r.a, r.b, r.c):
            key = (r.source_sheet, comp.role, int(comp.excel_row))
            if key in seen_components:
                continue
            seen_components.add(key)
            issue = ""
            if "DUPLICATE_MAPPED_SMILES_DIFFERENT_FORMULA" in comp.integrity_flags:
                issue = "Identical mapped SMILES is used by rows with different Formula; likely copied/stale structure."
            elif comp.heavy_formula_check == "no":
                issue = (
                    "Mapped SMILES heavy atoms do not match the listed component Formula"
                    + (" after the expected B loss N2." if comp.role == "B" else ".")
                )
            audit.append([
                r.source_sheet, comp.role, comp.item_index, comp.excel_row, comp.item_id, comp.formula,
                comp.expected_heavy_formula, comp.observed_heavy_formula, comp.heavy_formula_delta,
                comp.heavy_formula_check, comp.smiles_hash, comp.canonical_smiles,
                "; ".join(comp.integrity_flags), issue,
            ])
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = f"A1:{get_column_letter(len(audit_headers))}{max(1, audit.max_row)}"
    audit_widths = [18, 9, 12, 11, 18, 20, 24, 28, 22, 18, 20, 70, 45, 70]
    for i, width in enumerate(audit_widths, start=1):
        audit.column_dimensions[get_column_letter(i)].width = width
    for row in audit.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if diagnostics is not None:
        ds = wb.create_sheet("SMILES_Input_Diagnostics")
        dh = diagnostic_headers(include_raw=True)
        ds.append(dh)
        for cell in ds[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for d in diagnostics.diagnostics:
            ds.append(diagnostic_row_values(d, include_raw=True))
        ds.freeze_panes = "A2"
        ds.auto_filter.ref = f"A1:{get_column_letter(len(dh))}{max(1, ds.max_row)}"
        for i, head in enumerate(dh, start=1):
            width = 14 if head in {"Role", "Status", "Parse_OK", "Sanitize_OK", "Excel_Row"} else (24 if head in {"Source_Sheet", "ID", "Source", "Mapped_Dummy_Counts"} else 46)
            ds.column_dimensions[get_column_letter(i)].width = width
        for row in ds.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        sm = wb.create_sheet("SMILES_Diagnostic_Summary")
        for row in [
            ["Item", "Value"], ["Overall_Status", diagnostics.overall_status],
            ["Inputs_Checked", len(diagnostics.diagnostics)], ["Nonfatal", diagnostics.ok_count],
            ["Error_Count", diagnostics.error_count], ["Warning_Count", diagnostics.warning_count], [],
            ["Required_CORE", "[*:1], [*:2], [*:3] exactly once each"],
            ["Required_A", "[*:1] exactly once"], ["Required_B", "[*:2] exactly once"],
            ["Required_C", "[*:3] exactly once"],
        ]:
            sm.append(row)
        for x in diagnostics.global_errors:
            sm.append(["GLOBAL_ERROR", x])
        for x in diagnostics.global_warnings:
            sm.append(["GLOBAL_WARNING", x])
        for cell in sm[1]:
            cell.fill = fill
            cell.font = font
        sm.column_dimensions["A"].width = 24
        sm.column_dimensions["B"].width = 105
        for row in sm.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def run_simple_combo_smiles(
    component_workbook: Path,
    output_xlsx: Path,
    *,
    core_smiles: str,
    sheet_name: str = "",
    data_start_row: int = 2,
    only_formula: bool = False,
    generate_png: bool = True,
    image_limit: int = 100,
    embed_images: bool = True,
    embed_image_limit: int = 30,
    include_fragment_smiles: bool = False,
    output_csv: bool = True,
    progress: Optional[Callable[[str], None]] = None,
):
    if not RDKIT_AVAILABLE:
        raise RuntimeError(
            "RDKit is required. Install in the same environment: conda install -c conda-forge rdkit. "
            + str(RDKIT_ERROR)
        )
    component_workbook, output_xlsx = Path(component_workbook), Path(output_xlsx)
    core_smiles = _text(core_smiles)
    if not core_smiles:
        raise ValueError("Fixed CORE mapped SMILES is required")
    sheets, warnings = load_component_sheets(component_workbook, sheet_name=sheet_name, data_start_row=data_start_row)
    diagnostics, core_diag, component_diags = _build_input_diagnostics(sheets, core_smiles=core_smiles)
    for w in warnings:
        if w not in diagnostics.global_warnings:
            diagnostics.global_warnings.append(w)
    if progress:
        progress(f"Component workbook: {component_workbook}")
        progress(f"Valid sheets: {len(sheets)}")
        progress("Layout: A:D, E:H, I:L = ID / Formula / Mass / SMILES")
        progress(f"SMILES audit: {diagnostics.overall_status}; inputs={len(diagnostics.diagnostics)}, errors={diagnostics.error_count}, warnings={diagnostics.warning_count}")
        for d in diagnostics.diagnostics:
            if d.fatal:
                progress("INPUT ERROR: " + d.one_line())
                for fix in d.suggestions[:2]:
                    progress("  FIX: " + fix)

    results: List[SimpleProductResult] = []
    for sname, aa, bb, cc in sheets:
        if progress:
            progress(f"Sheet {sname}: A={len(aa)}, B={len(bb)}, C={len(cc)}, combinations={len(aa)*len(bb)*len(cc)}")
        results += _enumerate_sheet(
            sname, aa, bb, cc, core_smiles=core_smiles, only_formula=only_formula,
            core_diagnostic=core_diag, component_diagnostics=component_diags, progress=progress,
        )

    image_dir = None
    made = 0
    if generate_png:
        image_dir = output_xlsx.with_name(output_xlsx.stem + "__structures")
        image_dir.mkdir(parents=True, exist_ok=True)
        for old in image_dir.glob("*.png"):
            try:
                old.unlink()
            except Exception:
                pass
        for idx, r in enumerate(results, start=1):
            if not r.product_smiles or (image_limit > 0 and made >= image_limit):
                continue
            p = image_dir / f"{idx:04d}_{_safe_file_part(r.a.item_id)}_{_safe_file_part(r.b.item_id)}_{_safe_file_part(r.c.item_id)}.png"
            if _draw_product(r.product_smiles, p, legend=f"{r.a.item_id} | {r.b.item_id} | {r.c.item_id}"):
                r.structure_png = str(p)
                made += 1

    _write_xlsx(
        results, output_xlsx, core_smiles=core_smiles,
        include_fragment_smiles=include_fragment_smiles,
        embed_images=embed_images, embed_image_limit=embed_image_limit,
        diagnostics=diagnostics,
    )
    write_diagnostics_txt(
        diagnostics, output_xlsx.with_name(output_xlsx.stem + "__SMILES_Diagnostics.txt"), include_raw=False
    )
    out_csv = output_xlsx.with_suffix(".csv") if output_csv else None
    if out_csv:
        _write_csv(results, out_csv, include_fragment_smiles=include_fragment_smiles)

    n_success = sum(bool(r.product_smiles) for r in results)
    n_match = sum(r.formula_match == "yes" for r in results)
    n_failed = len(results) - n_success
    if progress:
        progress(f"Completed: total={len(results)}, success={n_success}, formula_match={n_match}, failed={n_failed}")
        progress(f"Output: {output_xlsx}")
    return SimpleBuildReport(
        output_xlsx, out_csv, image_dir, len(sheets), len(results), n_success, n_match, n_failed,
        warnings, results, diagnostics
    )


def preview_first_simple_combo(
    component_workbook: Path,
    *,
    core_smiles: str,
    sheet_name: str = "",
    data_start_row: int = 2,
    only_formula: bool = False,
    preview_dir: Optional[Path] = None,
):
    if not RDKIT_AVAILABLE:
        raise RuntimeError("RDKit is required. Install with: conda install -c conda-forge rdkit")
    sheets, _ = load_component_sheets(Path(component_workbook), sheet_name=sheet_name, data_start_row=data_start_row)
    core_smiles = _text(core_smiles)
    if not core_smiles:
        raise ValueError("Fixed CORE mapped SMILES is required")
    last = None
    for sname, aa, bb, cc in sheets:
        rows = _enumerate_sheet(sname, aa[:1], bb[:1], cc[:1], core_smiles=core_smiles, only_formula=only_formula)
        if not rows:
            continue
        last = rows[0]
        if last.product_smiles:
            if preview_dir:
                preview_dir = Path(preview_dir)
                preview_dir.mkdir(parents=True, exist_ok=True)
                p = preview_dir / "simple_combo_smiles_preview.png"
                if _draw_product(last.product_smiles, p, legend=f"{last.a.item_id} | {last.b.item_id} | {last.c.item_id}"):
                    last.structure_png = str(p)
            return last
    if last:
        return last
    raise ValueError("No combination could be previewed")
