"""Concise model-output tables and reusable prediction plots.

This module keeps the detailed model audit sheets intact while adding three
human-readable plots:

1. known standards: actual vs out-of-fold predicted concentration;
2. unknown sample: predicted concentration sorted from low to high;
3. unknown sample: the same predictions with empirical P80 error bars from the
   selected best model.

The plotting functions can also read an already generated result workbook, so
plots can be regenerated without recalculating RDKit/3D descriptors or
refitting the models.
"""
from __future__ import annotations

import math
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _num(value: object) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def _identity(record: Dict[str, object]) -> str:
    parts = [_text(record.get("Source_file")), _text(record.get("Source_sheet")), _text(record.get("Source_row"))]
    if any(parts):
        return "|".join(parts)
    return _text(record.get("Record_ID") or record.get("Product_Master_Formula_Key") or record.get("ABC_Formula_Key") or record.get("Combo"))


def _key(record: Dict[str, object]) -> str:
    return _text(record.get("Product_Master_Formula_Key") or record.get("ABC_Formula_Key"))


def _compound_name(record: Dict[str, object], fallback: str = "") -> str:
    return (
        _text(record.get("Compound"))
        or _text(record.get("Name"))
        or _text(record.get("Formula"))
        or _text(record.get("Combo"))
        or fallback
    )


def _append_dict_sheet(wb, title: str, rows: Sequence[Dict[str, object]]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    if rows:
        headers: List[str] = []
        seen = set()
        for row in rows:
            for key in row.keys():
                if key not in seen:
                    seen.add(key)
                    headers.append(str(key))
    else:
        headers = ["Message"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        width_by_header = {
            "Order": 9, "Compound": 22, "Formula": 18, "Combo": 48,
            "Injection_Group": 18, "Actual_concentration": 20,
            "Predicted_concentration": 22, "Concentration_unit": 18,
            "Fold_error": 14, "Measured_ratio": 17,
            "Actual_level_L01_L10": 20, "Predicted_level_L01_L10": 22,
            "Predicted_level_label": 20, "Exact_level_hit": 16,
            "Within_1_level": 16,
            "Predicted_concentration_rank_low_to_high": 32,
            "Predicted_concentration_rank_high_to_low": 33,
            "Predicted_concentration_P80_lower": 29,
            "Predicted_concentration_P80_upper": 29,
            "P80_fold_factor": 17,
            "Prediction_interval_basis": 34,
            "Corrected_score_log10": 22, "Prediction_model": 25,
            "Prediction_status": 20,
        }
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = width_by_header.get(header, 18)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=(cell.column in {2, 4}))
    except Exception:
        pass


def _best_p80_fold(benchmark_result, trend_result) -> float:
    """Return the empirical P80 fold error for the actually selected model."""
    if trend_result is not None and getattr(trend_result, "best_model", ""):
        model = str(getattr(trend_result, "best_model", ""))
        for row in getattr(trend_result, "comparison", []) or []:
            if str(row.get("Model")) == model:
                value = _num(row.get("P80_Fold_Error"))
                if value is not None and value >= 1.0:
                    return float(value)
    if benchmark_result is not None:
        metrics = getattr(benchmark_result, "best_metrics", {}) or {}
        value = _num(metrics.get("P80_Fold_Error"))
        if value is not None and value >= 1.0:
            return float(value)
        model = str(getattr(benchmark_result, "best_model", ""))
        for row in getattr(benchmark_result, "comparison", []) or []:
            if str(row.get("Model")) == model:
                value = _num(row.get("P80_Fold_Error"))
                if value is not None and value >= 1.0:
                    return float(value)
    # Conservative fallback for older workbooks or classifier-only output.
    return 5.0


def _prediction_diversity(rows: Sequence[Dict[str, object]], key: str) -> Dict[str, object]:
    values = [
        float(v) for v in (_num(row.get(key)) for row in rows)
        if v is not None and float(v) > 0
    ]
    if not values:
        return {"valid": 0, "unique": 0, "unique_fraction_pct": 0.0, "max_repeat": 0, "max_repeat_pct": 0.0}
    rounded = np.round(np.asarray(values, dtype=float), 12)
    unique_values, counts = np.unique(rounded, return_counts=True)
    max_repeat = int(np.max(counts)) if len(counts) else 0
    return {
        "valid": int(len(values)),
        "unique": int(len(unique_values)),
        "unique_fraction_pct": float(len(unique_values) / max(1, len(values)) * 100.0),
        "max_repeat": max_repeat,
        "max_repeat_pct": float(max_repeat / max(1, len(values)) * 100.0),
    }


def _choose_continuous_target_column(rows: Sequence[Dict[str, object]], preferred_model: str) -> Tuple[str, str, bool, Dict[str, object]]:
    """Choose a genuinely continuous target-concentration column.

    Pairwise ranking models in older workbooks were converted to concentration
    with isotonic regression.  That mapping may contain only a few plateaus.
    When this is detected, the response-factor ExtraTrees estimate is used for
    concentration magnitude while the pairwise model remains the ranking model.
    """
    candidates: List[Tuple[str, str]] = [
        ("Preferred_estimated_concentration", preferred_model or "Preferred trend model"),
        ("Continuous_estimated_concentration", "Continuous response-factor model"),
        ("ResponseFactorExtraTrees_estimated_concentration", "ResponseFactorExtraTrees"),
        ("Estimated_concentration", "Continuous benchmark model"),
    ]
    stats_by_key = {key: _prediction_diversity(rows, key) for key, _ in candidates}
    preferred_key, preferred_label = candidates[0]
    preferred_stats = stats_by_key[preferred_key]
    n_valid = int(preferred_stats.get("valid", 0))
    minimum_unique = max(12, int(math.ceil(math.sqrt(max(1, n_valid))))) if n_valid else 12
    plateau = bool(n_valid >= 20 and int(preferred_stats.get("unique", 0)) < minimum_unique)
    if not plateau:
        return preferred_key, preferred_label, False, preferred_stats

    best_key, best_label, best_stats = preferred_key, preferred_label, preferred_stats
    for key, label in candidates[1:]:
        stats = stats_by_key[key]
        # Require broadly comparable coverage and substantially greater value
        # diversity before replacing the preferred concentration column.
        if int(stats.get("valid", 0)) < max(10, int(0.8 * n_valid)):
            continue
        if int(stats.get("unique", 0)) > int(best_stats.get("unique", 0)):
            best_key, best_label, best_stats = key, label, stats
    applied = best_key != preferred_key
    return best_key, best_label, applied, best_stats


def _continuous_maps(benchmark_result, trend_result):
    known_map: Dict[str, Dict[str, object]] = {}
    target_map: Dict[str, Dict[str, object]] = {}
    source = ""
    model = ""
    default_fold = _best_p80_fold(benchmark_result, trend_result)

    if trend_result is not None and getattr(trend_result, "best_model", ""):
        model = str(trend_result.best_model)
        source = "Trend optimizer"
        for row in getattr(trend_result, "cv_predictions", []):
            if str(row.get("Model")) != model:
                continue
            ident = _identity(row)
            if ident:
                known_map[ident] = {
                    "Predicted_concentration": row.get("Predicted_concentration", ""),
                    "Fold_error": row.get("Fold_error", ""),
                    "Predicted_log10_RRF": row.get("Predicted_log10_RRF", ""),
                    "Response_corrected_score_log10": row.get("Predicted_log10_concentration", ""),
                    "Prediction_model": model,
                    "Prediction_source": source,
                }
        trend_target_rows = list(getattr(trend_result, "target_predictions", []))
        concentration_key, concentration_model, plateau_fallback, diversity = _choose_continuous_target_column(
            trend_target_rows, model
        )
        for row in trend_target_rows:
            ident = _identity(row)
            pred = _num(row.get(concentration_key))
            if ident:
                target_map[ident] = {
                    "Predicted_concentration": pred if pred is not None else "",
                    "Corrected_score_log10": (
                        math.log10(pred) if pred is not None and pred > 0 else ""
                    ),
                    "Predicted_level_from_continuous": row.get("Preferred_concentration_decile_1_to_10", ""),
                    "Predicted_level_label_from_continuous": row.get("Preferred_concentration_level", ""),
                    "Predicted_concentration_P80_lower": (pred / default_fold if pred is not None and pred > 0 else ""),
                    "Predicted_concentration_P80_upper": (pred * default_fold if pred is not None and pred > 0 else ""),
                    "P80_fold_factor": default_fold,
                    "Prediction_interval_basis": "Best-model cross-validated P80 fold error",
                    "Prediction_model": concentration_model,
                    "Trend_rank_model": model,
                    "Continuous_prediction_column": concentration_key,
                    "Plateau_fallback_applied": plateau_fallback,
                    "Target_unique_concentration_values": diversity.get("unique", ""),
                    "Target_unique_value_fraction_pct": diversity.get("unique_fraction_pct", ""),
                    "Prediction_source": source,
                }
    elif benchmark_result is not None:
        model = str(getattr(benchmark_result, "best_model", ""))
        source = "Multi-model benchmark"
        for row in getattr(benchmark_result, "cv_predictions", []):
            if str(row.get("Model")) != model:
                continue
            ident = _identity(row)
            if ident:
                known_map[ident] = {
                    "Predicted_concentration": row.get("Predicted_concentration", ""),
                    "Fold_error": row.get("Fold_error", ""),
                    "Predicted_log10_RRF": row.get("Predicted_log10_RRF", ""),
                    "Response_corrected_score_log10": row.get("Predicted_log10_concentration", ""),
                    "Prediction_model": model,
                    "Prediction_source": source,
                }
        for row in getattr(benchmark_result, "target_predictions", []):
            ident = _identity(row)
            pred = _num(row.get("Estimated_concentration"))
            fold = _num(row.get("P80_fold_factor")) or default_fold
            lower = _num(row.get("Estimated_concentration_P80_lower"))
            upper = _num(row.get("Estimated_concentration_P80_upper"))
            if pred is not None and pred > 0:
                if lower is None or lower <= 0:
                    lower = pred / max(float(fold), 1.0)
                if upper is None or upper <= 0:
                    upper = pred * max(float(fold), 1.0)
            if ident:
                target_map[ident] = {
                    "Predicted_concentration": pred if pred is not None else "",
                    "Corrected_score_log10": row.get("Response_corrected_score_log10", ""),
                    "Predicted_level_from_continuous": row.get("Concentration_decile_1_to_10", ""),
                    "Predicted_level_label_from_continuous": row.get("Concentration_level", ""),
                    "Predicted_concentration_P80_lower": lower if lower is not None else "",
                    "Predicted_concentration_P80_upper": upper if upper is not None else "",
                    "P80_fold_factor": fold,
                    "Prediction_interval_basis": "Best-model cross-validated P80 fold error",
                    "Prediction_model": model,
                    "Prediction_source": source,
                }
    return known_map, target_map, model, source, default_fold


def build_simple_prediction_rows(
    benchmark_result,
    trend_result,
    level_result,
    target_records: Sequence[Dict[str, object]],
    *,
    concentration_unit: str = "",
) -> Tuple[List[Dict[str, object]], List[Dict[str, object]]]:
    known_cont, target_cont, preferred_model, preferred_source, default_fold = _continuous_maps(benchmark_result, trend_result)
    known_level = {_identity(r): r for r in getattr(level_result, "known_predictions", [])} if level_result is not None else {}
    target_level = {_identity(r): r for r in getattr(level_result, "target_predictions", [])} if level_result is not None else {}

    calibration = list(getattr(benchmark_result, "calibration_records", []) if benchmark_result is not None else [])
    known_rows: List[Dict[str, object]] = []
    for record in calibration:
        ident = _identity(record)
        cont = known_cont.get(ident, {})
        level = known_level.get(ident, {})
        # Direct level classification is kept in separate level columns.  It
        # must not be converted back into a pseudo-continuous concentration,
        # otherwise the unknown sample can collapse to only a few repeated
        # representative values.
        predicted_concentration = cont.get("Predicted_concentration", "")
        actual = _num(record.get("Actual_concentration"))
        pred = _num(predicted_concentration)
        known_rows.append({
            "Order": 0,
            "Compound": record.get("Name", ""),
            "Formula": record.get("Formula", ""),
            "Combo": record.get("Combo", ""),
            "Injection_Group": record.get("Injection_Group", ""),
            "Actual_concentration": actual if actual is not None else "",
            "Predicted_concentration": pred if pred is not None else "",
            "Concentration_unit": str(concentration_unit or ""),
            "Fold_error": (max(pred / actual, actual / pred) if actual and pred and actual > 0 and pred > 0 else cont.get("Fold_error", "")),
            "Actual_level_L01_L10": level.get("Actual_level_1_to_10", ""),
            "Predicted_level_L01_L10": level.get("Predicted_level_1_to_10", ""),
            "Exact_level_hit": level.get("Exact_level_hit", ""),
            "Within_1_level": level.get("Within_1_level", ""),
            "Prediction_model": cont.get("Prediction_model", preferred_model) or level.get("Level_model", ""),
            "Prediction_status": "PREDICTED" if pred is not None or level.get("Predicted_level_1_to_10", "") != "" else "NO PREDICTION",
        })
    known_rows.sort(key=lambda r: (float(_num(r.get("Actual_concentration")) or 1e300), _text(r.get("Compound") or r.get("Formula") or r.get("Combo"))))
    for idx, row in enumerate(known_rows, start=1):
        row["Order"] = idx

    # Keep all target rows, including rows that could not be predicted.
    target_rows: List[Dict[str, object]] = []
    for idx, record in enumerate(target_records, start=1):
        ident = _identity(record)
        cont = target_cont.get(ident, {})
        level = target_level.get(ident, {})
        # Continuous concentration and level classification are independent
        # outputs.  Never use a class representative as the continuous value.
        pred = _num(cont.get("Predicted_concentration"))
        lower = _num(cont.get("Predicted_concentration_P80_lower"))
        upper = _num(cont.get("Predicted_concentration_P80_upper"))
        fold = _num(cont.get("P80_fold_factor")) or default_fold
        if pred is not None and pred > 0:
            if lower is None or lower <= 0:
                lower = pred / max(float(fold), 1.0)
            if upper is None or upper <= 0:
                upper = pred * max(float(fold), 1.0)
        if pred is not None:
            status = "PREDICTED"
        elif level.get("Predicted_level_1_to_10", "") != "":
            status = "LEVEL_ONLY"
        else:
            status = level.get("Prediction_status", "SKIPPED")
        target_rows.append({
            "Order": idx,
            "Compound": record.get("Name", ""),
            "Formula": record.get("Formula", ""),
            "Combo": record.get("Combo", ""),
            "Measured_ratio": record.get("Measured_ratio", ""),
            "Predicted_concentration": pred if pred is not None else "",
            "Concentration_unit": str(concentration_unit or ""),
            "Predicted_concentration_P80_lower": lower if lower is not None else "",
            "Predicted_concentration_P80_upper": upper if upper is not None else "",
            "P80_fold_factor": fold if pred is not None else "",
            "Prediction_interval_basis": cont.get("Prediction_interval_basis", "Best-model cross-validated P80 fold error" if pred is not None else ""),
            "Predicted_level_L01_L10": level.get("Predicted_level_1_to_10", cont.get("Predicted_level_from_continuous", "")),
            "Predicted_level_label": level.get("Predicted_level_label", cont.get("Predicted_level_label_from_continuous", "")),
            "Predicted_concentration_rank_low_to_high": "",
            "Predicted_concentration_rank_high_to_low": "",
            "Corrected_score_log10": cont.get("Corrected_score_log10", ""),
            "Prediction_model": cont.get("Prediction_model", preferred_model) or level.get("Level_model", ""),
            "Trend_rank_model": cont.get("Trend_rank_model", ""),
            "Continuous_prediction_column": cont.get("Continuous_prediction_column", ""),
            "Plateau_fallback_applied": cont.get("Plateau_fallback_applied", False),
            "Target_unique_concentration_values": cont.get("Target_unique_concentration_values", ""),
            "Target_unique_value_fraction_pct": cont.get("Target_unique_value_fraction_pct", ""),
            "Prediction_status": status,
        })

    valid = [(i, _num(r.get("Predicted_concentration"))) for i, r in enumerate(target_rows)]
    valid = [(i, v) for i, v in valid if v is not None]
    low_sorted = sorted(valid, key=lambda x: x[1])
    n_valid = len(low_sorted)
    for low_rank, (i, _) in enumerate(low_sorted, start=1):
        target_rows[i]["Predicted_concentration_rank_low_to_high"] = low_rank
        target_rows[i]["Predicted_concentration_rank_high_to_low"] = n_valid - low_rank + 1
    return known_rows, target_rows


def _use_log_axis(values: np.ndarray) -> bool:
    if values.size == 0 or not np.all(np.isfinite(values)) or not np.all(values > 0):
        return False
    minimum = float(np.min(values))
    maximum = float(np.max(values))
    return minimum > 0 and maximum / minimum >= 30.0


def _set_sparse_compound_ticks(ax, names: Sequence[str], *, max_labels: int = 80) -> None:
    n = len(names)
    if n <= 0:
        return
    step = max(1, int(math.ceil(n / max(1, int(max_labels)))))
    positions = list(range(0, n, step))
    if positions[-1] != n - 1:
        positions.append(n - 1)
    labels = [names[i] for i in positions]
    ax.set_xticks(positions)
    ax.set_xticklabels(labels, rotation=90, fontsize=6.5)


def make_known_actual_predicted_plot(
    known_rows: Sequence[Dict[str, object]],
    output_path: Path,
    *,
    output_language: str = "en",
) -> Path:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    rows = [r for r in known_rows if _num(r.get("Actual_concentration")) is not None and _num(r.get("Predicted_concentration")) is not None]
    rows.sort(key=lambda r: (float(r["Actual_concentration"]), _compound_name(r)))
    if not rows:
        raise ValueError("No rows contain both actual and predicted concentrations")
    names = [_compound_name(r, str(i + 1)) for i, r in enumerate(rows)]
    actual = np.asarray([float(r["Actual_concentration"]) for r in rows], dtype=float)
    predicted = np.asarray([float(r["Predicted_concentration"]) for r in rows], dtype=float)
    x = np.arange(len(rows), dtype=float)
    zh = str(output_language or "en").lower().startswith("zh")

    fig, ax = plt.subplots(figsize=(max(14.0, len(rows) * 0.18), 7.2))
    ax.scatter(x, actual, s=24, label="Actual concentration" if not zh else 'Actual concentration')
    ax.scatter(x, predicted, s=24, label="Predicted concentration" if not zh else 'Predicted concentration')
    combined = np.concatenate([actual, predicted])
    if _use_log_axis(combined):
        ax.set_yscale("log")
    ax.set_xticks(x)
    ax.set_xticklabels(names, rotation=90, fontsize=7)
    ax.set_xlabel("Known standards sorted by actual concentration" if not zh else 'Known standards sorted by actual concentration')
    unit = _text(rows[0].get("Concentration_unit"))
    ylabel = ("Concentration" if not zh else 'Concentration') + (f" ({unit})" if unit else "")
    ax.set_ylabel(ylabel)
    ax.set_title("Actual and out-of-fold predicted concentrations" if not zh else 'Actual and out-of-fold predicted concentrations')
    ax.legend()
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=200, bbox_inches="tight")
    plt.close(fig)
    return output_path


def _valid_unknown_rows(target_rows: Sequence[Dict[str, object]]) -> List[Dict[str, object]]:
    rows = []
    for row in target_rows:
        pred = _num(row.get("Predicted_concentration"))
        if pred is None or pred <= 0:
            continue
        item = dict(row)
        item["Predicted_concentration"] = pred
        rows.append(item)
    rows.sort(key=lambda r: (float(r["Predicted_concentration"]), _compound_name(r)))
    return rows


def make_unknown_predicted_sorted_plot(
    target_rows: Sequence[Dict[str, object]],
    output_path: Path,
    *,
    output_language: str = "en",
) -> Path:
    """Plot valid unknown-sample predictions from low to high."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    rows = _valid_unknown_rows(target_rows)
    if not rows:
        raise ValueError("No unknown-sample rows contain a positive predicted concentration")
    predicted = np.asarray([float(r["Predicted_concentration"]) for r in rows], dtype=float)
    names = [_compound_name(r, str(i + 1)) for i, r in enumerate(rows)]
    x = np.arange(len(rows), dtype=float)
    zh = str(output_language or "en").lower().startswith("zh")

    width = min(36.0, max(16.0, len(rows) * 0.035))
    fig, ax = plt.subplots(figsize=(width, 7.6))
    ax.scatter(x, predicted, s=12, label="Predicted concentration" if not zh else 'Predicted concentration')
    if _use_log_axis(predicted):
        ax.set_yscale("log")
    _set_sparse_compound_ticks(ax, names, max_labels=80)
    unit = _text(rows[0].get("Concentration_unit"))
    ylabel = ("Predicted concentration" if not zh else 'Predicted concentration') + (f" ({unit})" if unit else "")
    ax.set_ylabel(ylabel)
    ax.set_xlabel(
        f"Unknown compounds sorted by predicted concentration (valid n={len(rows)})"
        if not zh else f'Unknown compounds sorted by predicted concentration (valid predictions n={len(rows)}）'
    )
    ax.set_title(
        "Unknown-sample predicted concentrations, low to high"
        if not zh else 'Unknown-sample predicted concentrations, low to high'
    )
    ax.legend(loc="best")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return output_path


def make_unknown_prediction_errorbar_plot(
    target_rows: Sequence[Dict[str, object]],
    output_path: Path,
    *,
    output_language: str = "en",
) -> Path:
    """Plot one empirical best-model P80 error bar for every valid target.

    The interval is multiplicative.  When a row already contains lower/upper
    bounds they are used directly; otherwise the row's P80 fold factor is
    applied as [prediction/fold, prediction*fold].
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    rows = _valid_unknown_rows(target_rows)
    if not rows:
        raise ValueError("No unknown-sample rows contain a positive predicted concentration")

    pred_values: List[float] = []
    lower_values: List[float] = []
    upper_values: List[float] = []
    fold_values: List[float] = []
    names: List[str] = []
    kept_rows: List[Dict[str, object]] = []
    for i, row in enumerate(rows):
        pred = float(row["Predicted_concentration"])
        fold = _num(row.get("P80_fold_factor")) or 5.0
        fold = max(1.0, float(fold))
        lower = _num(row.get("Predicted_concentration_P80_lower"))
        upper = _num(row.get("Predicted_concentration_P80_upper"))
        if lower is None or lower <= 0:
            lower = pred / fold
        if upper is None or upper < pred:
            upper = pred * fold
        lower = min(float(lower), pred)
        upper = max(float(upper), pred)
        pred_values.append(pred)
        lower_values.append(lower)
        upper_values.append(upper)
        fold_values.append(fold)
        names.append(_compound_name(row, str(i + 1)))
        kept_rows.append(row)

    predicted = np.asarray(pred_values, dtype=float)
    lower = np.asarray(lower_values, dtype=float)
    upper = np.asarray(upper_values, dtype=float)
    yerr = np.vstack([predicted - lower, upper - predicted])
    x = np.arange(len(predicted), dtype=float)
    zh = str(output_language or "en").lower().startswith("zh")
    width = min(36.0, max(16.0, len(predicted) * 0.035))

    fig, ax = plt.subplots(figsize=(width, 8.0))
    ax.errorbar(
        x,
        predicted,
        yerr=yerr,
        fmt="o",
        markersize=2.4,
        elinewidth=0.45,
        capsize=1.2,
        capthick=0.45,
        alpha=0.65,
        label="Prediction with P80 empirical interval" if not zh else 'Prediction with P80 empirical interval',
    )
    combined = np.concatenate([lower, upper])
    if _use_log_axis(combined):
        ax.set_yscale("log")
    _set_sparse_compound_ticks(ax, names, max_labels=80)
    unit = _text(kept_rows[0].get("Concentration_unit"))
    ylabel = ("Predicted concentration" if not zh else 'Predicted concentration') + (f" ({unit})" if unit else "")
    ax.set_ylabel(ylabel)
    ax.set_xlabel(
        f"Unknown compounds sorted by predicted concentration (valid n={len(predicted)})"
        if not zh else f'Unknown compounds sorted by predicted concentration (valid predictions n={len(predicted)}）'
    )
    median_fold = float(np.median(np.asarray(fold_values, dtype=float))) if fold_values else 5.0
    ax.set_title(
        f"Best-model empirical uncertainty for every prediction (median P80 factor {median_fold:.2f}x)"
        if not zh else f'Empirical error bars for the selected model (median P80 factor {median_fold:.2f}×）'
    )
    ax.legend(loc="best")
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, dpi=180, bbox_inches="tight")
    plt.close(fig)
    return output_path


def _add_plot_image(ws, path: Path, title: str, anchor_row: int) -> None:
    from openpyxl.drawing.image import Image as XLImage
    ws.cell(anchor_row, 1, title)
    image = XLImage(str(path))
    scale = min(1.0, 1400.0 / max(1, image.width), 760.0 / max(1, image.height))
    image.width = int(image.width * scale)
    image.height = int(image.height * scale)
    ws.add_image(image, f"A{anchor_row + 1}")


def _prediction_value_audit_rows(
    known_rows: Sequence[Dict[str, object]],
    target_rows: Sequence[Dict[str, object]],
) -> List[Dict[str, object]]:
    rows_out: List[Dict[str, object]] = []
    for label, rows in (("Known standards", known_rows), ("Unknown sample", target_rows)):
        stats = _prediction_diversity(rows, "Predicted_concentration")
        valid = int(stats.get("valid", 0))
        unique = int(stats.get("unique", 0))
        minimum_unique = max(12, int(math.ceil(math.sqrt(max(1, valid))))) if valid else 12
        rows_out.append({
            "Dataset": label,
            "Valid_predictions": valid,
            "Unique_concentration_values": unique,
            "Unique_value_fraction_pct": stats.get("unique_fraction_pct", 0.0),
            "Most_repeated_value_count": stats.get("max_repeat", 0),
            "Most_repeated_value_pct": stats.get("max_repeat_pct", 0.0),
            "Minimum_reasonable_unique_values": minimum_unique,
            "Plateau_warning": bool(valid >= 20 and unique < minimum_unique),
            "Interpretation": (
                "Discrete/plateaued concentration output detected; inspect mapping/source columns."
                if valid >= 20 and unique < minimum_unique
                else "Continuous concentration diversity is acceptable."
            ),
        })
    return rows_out


def append_simple_outputs_to_workbook(
    wb,
    known_rows: Sequence[Dict[str, object]],
    target_rows: Sequence[Dict[str, object]],
    output_xlsx: Path,
    *,
    output_language: str = "en",
) -> List[Path]:
    _append_dict_sheet(wb, "Known_Standards_Predictions", known_rows)
    _append_dict_sheet(wb, "Unknown_Sample_Predictions", target_rows)
    _append_dict_sheet(wb, "Prediction_Value_Audit", _prediction_value_audit_rows(known_rows, target_rows))
    sorted_target_rows = sorted(
        [dict(row) for row in target_rows],
        key=lambda row: (
            _num(row.get("Predicted_concentration")) is None,
            float(_num(row.get("Predicted_concentration")) or 0.0),
            _compound_name(row),
        ),
    )
    _append_dict_sheet(wb, "Unknown_Predictions_Sorted", sorted_target_rows)
    paths: List[Path] = []
    warnings: List[Dict[str, object]] = []
    plot_dir = Path(output_xlsx).with_name("_ctsimple_" + str(abs(hash(str(output_xlsx))))[-8:])
    plot_dir.mkdir(parents=True, exist_ok=True)

    plot_specs = [
        (
            "Known standards: actual vs predicted",
            lambda: make_known_actual_predicted_plot(
                known_rows,
                plot_dir / "32_known_actual_vs_predicted_sorted.png",
                output_language=output_language,
            ),
        ),
        (
            "Unknown sample: predicted concentration, low to high",
            lambda: make_unknown_predicted_sorted_plot(
                target_rows,
                plot_dir / "33_unknown_predicted_sorted.png",
                output_language=output_language,
            ),
        ),
        (
            "Unknown sample: best-model P80 error bars",
            lambda: make_unknown_prediction_errorbar_plot(
                target_rows,
                plot_dir / "34_unknown_prediction_errorbars.png",
                output_language=output_language,
            ),
        ),
    ]

    if "Simple_Result_Plots" in wb.sheetnames:
        del wb["Simple_Result_Plots"]
    ws = wb.create_sheet("Simple_Result_Plots")
    anchor_rows = [1, 60, 119]
    for (title, maker), anchor_row in zip(plot_specs, anchor_rows):
        try:
            path = maker()
            paths.append(path)
            _add_plot_image(ws, path, title, anchor_row)
        except Exception as exc:
            warnings.append({"Plot": title, "Warning": str(exc)})
            ws.cell(anchor_row, 1, title)
            ws.cell(anchor_row + 1, 1, f"Plot skipped: {exc}")

    if warnings:
        _append_dict_sheet(wb, "Simple_Plot_Warnings", warnings)
    return paths


def _read_sheet_rows(wb, sheet_name: str) -> List[Dict[str, object]]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    headers = [_text(x) for x in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if row and any(x not in (None, "") for x in row)]


def _find_global_p80_fold(wb) -> float:
    """Find a best-model P80 fold factor in an existing workbook."""
    candidates = []
    for sheet_name in ("Trend_Model_Comparison", "Model_Comparison", "Dual_Model_Comparison"):
        for row in _read_sheet_rows(wb, sheet_name):
            rank = _num(row.get("Rank"))
            p80 = _num(row.get("P80_Fold_Error") or row.get("Best_P80_Fold_Error") or row.get("P80_fold"))
            if p80 is not None and p80 >= 1.0:
                candidates.append((rank if rank is not None else 999.0, float(p80)))
    if candidates:
        candidates.sort(key=lambda x: x[0])
        return candidates[0][1]
    for sheet_name in ("Diagnostics", "Model_Decision", "Trend_Optimizer_Decision"):
        for row in _read_sheet_rows(wb, sheet_name):
            key = _text(row.get("Metric") or row.get("Item") or row.get("Setting"))
            if "P80" in key and "Fold" in key:
                p80 = _num(row.get("Value"))
                if p80 is not None and p80 >= 1.0:
                    return float(p80)
    return 5.0


def _normalize_known_rows(rows: Sequence[Dict[str, object]]) -> List[Dict[str, object]]:
    normalized: List[Dict[str, object]] = []
    for row in rows:
        actual = row.get("Actual_concentration", row.get("Actual concentration", row.get("Known_concentration", "")))
        predicted = row.get(
            "Predicted_concentration",
            row.get("OOF_estimated_concentration", row.get("Predicted concentration", row.get("OOF_Predicted_concentration", ""))),
        )
        normalized.append({
            "Compound": row.get("Compound", row.get("Name", "")),
            "Name": row.get("Name", row.get("Compound", "")),
            "Formula": row.get("Formula", ""),
            "Combo": row.get("Combo", ""),
            "Actual_concentration": actual,
            "Predicted_concentration": predicted,
            "Concentration_unit": row.get("Concentration_unit", ""),
        })
    return normalized


def _normalize_target_rows(rows: Sequence[Dict[str, object]], default_fold: float) -> List[Dict[str, object]]:
    rows = list(rows)
    candidate_keys = [
        "Predicted_concentration",
        "Continuous_estimated_concentration",
        "ResponseFactorExtraTrees_estimated_concentration",
        "Estimated_concentration",
        "Preferred_estimated_concentration",
        "Consensus_estimated_concentration",
        "PhysicalPairwiseRanker_estimated_concentration",
    ]
    stats = {key: _prediction_diversity(rows, key) for key in candidate_keys}
    preferred_key = candidate_keys[0]
    preferred = stats[preferred_key]
    n_valid = int(preferred.get("valid", 0))
    minimum_unique = max(12, int(math.ceil(math.sqrt(max(1, n_valid))))) if n_valid else 12
    plateau = bool(n_valid >= 20 and int(preferred.get("unique", 0)) < minimum_unique)
    chosen_key = preferred_key
    if n_valid == 0 or plateau:
        best_score = (-1, -1)
        for key in candidate_keys:
            item = stats[key]
            valid = int(item.get("valid", 0))
            unique = int(item.get("unique", 0))
            if valid < max(1, int(0.8 * max(1, n_valid))) and n_valid > 0:
                continue
            score = (unique, valid)
            if score > best_score:
                best_score = score
                chosen_key = key

    normalized: List[Dict[str, object]] = []
    chosen_stats = stats.get(chosen_key, {})
    for row in rows:
        pred = row.get(chosen_key, "")
        pred_num = _num(pred)
        fold = _num(row.get("P80_fold_factor") or row.get("P80_Fold_Factor")) or default_fold
        lower = row.get("Predicted_concentration_P80_lower", row.get("Estimated_concentration_P80_lower", ""))
        upper = row.get("Predicted_concentration_P80_upper", row.get("Estimated_concentration_P80_upper", ""))
        lower_num = _num(lower)
        upper_num = _num(upper)
        if pred_num is not None and pred_num > 0:
            if lower_num is None or lower_num <= 0:
                lower_num = pred_num / max(float(fold), 1.0)
            if upper_num is None or upper_num < pred_num:
                upper_num = pred_num * max(float(fold), 1.0)
        normalized.append({
            "Compound": row.get("Compound", row.get("Name", "")),
            "Name": row.get("Name", row.get("Compound", "")),
            "Formula": row.get("Formula", ""),
            "Combo": row.get("Combo", ""),
            "Predicted_concentration": pred_num if pred_num is not None else "",
            "Concentration_unit": row.get("Concentration_unit", ""),
            "Predicted_concentration_P80_lower": lower_num if lower_num is not None else "",
            "Predicted_concentration_P80_upper": upper_num if upper_num is not None else "",
            "P80_fold_factor": fold if pred_num is not None else "",
            "Prediction_status": row.get("Prediction_status", ""),
            "Continuous_prediction_column": chosen_key,
            "Plateau_fallback_applied": chosen_key != preferred_key,
            "Target_unique_concentration_values": chosen_stats.get("unique", ""),
            "Target_unique_value_fraction_pct": chosen_stats.get("unique_fraction_pct", ""),
        })
    return normalized


def replot_all_prediction_plots_from_workbook(
    workbook_path: Path,
    *,
    output_language: str = "en",
) -> List[Path]:
    """Regenerate all three concise plots from an existing result workbook."""
    from openpyxl import load_workbook

    workbook_path = Path(workbook_path)
    wb = load_workbook(str(workbook_path), read_only=True, data_only=True)
    known_sheet = "Known_Standards_Predictions"
    if known_sheet not in wb.sheetnames:
        for candidate in ("Level_Known_Standards", "Concentration_Ratio_Comparison", "Raw_vs_Corrected"):
            if candidate in wb.sheetnames:
                known_sheet = candidate
                break
    target_sheet = "Unknown_Sample_Predictions"
    if target_sheet not in wb.sheetnames:
        for candidate in ("Response_Corrected_Targets", "Best_Model_Predictions", "Trend_Targets", "Dual_Model_Targets"):
            if candidate in wb.sheetnames:
                target_sheet = candidate
                break

    known_rows_raw = _read_sheet_rows(wb, known_sheet) if known_sheet in wb.sheetnames else []
    target_rows_raw = _read_sheet_rows(wb, target_sheet) if target_sheet in wb.sheetnames else []
    default_fold = _find_global_p80_fold(wb)

    known_rows = _normalize_known_rows(known_rows_raw)
    target_rows = _normalize_target_rows(target_rows_raw, default_fold)
    current_stats = _prediction_diversity(target_rows, "Predicted_concentration")
    current_valid = int(current_stats.get("valid", 0))
    current_min_unique = max(12, int(math.ceil(math.sqrt(max(1, current_valid))))) if current_valid else 12
    current_plateau = bool(current_valid >= 20 and int(current_stats.get("unique", 0)) < current_min_unique)
    if current_plateau and "Trend_Targets" in wb.sheetnames and target_sheet != "Trend_Targets":
        alternative = _normalize_target_rows(_read_sheet_rows(wb, "Trend_Targets"), default_fold)
        alt_stats = _prediction_diversity(alternative, "Predicted_concentration")
        if int(alt_stats.get("unique", 0)) > int(current_stats.get("unique", 0)):
            target_rows = alternative
    wb.close()
    output_dir = workbook_path.with_name("_replot_" + workbook_path.stem[:40])
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: List[Path] = []
    if known_rows:
        paths.append(make_known_actual_predicted_plot(
            known_rows,
            output_dir / "known_actual_vs_predicted_sorted.png",
            output_language=output_language,
        ))
    if target_rows:
        paths.append(make_unknown_predicted_sorted_plot(
            target_rows,
            output_dir / "unknown_predicted_sorted.png",
            output_language=output_language,
        ))
        paths.append(make_unknown_prediction_errorbar_plot(
            target_rows,
            output_dir / "unknown_prediction_errorbars.png",
            output_language=output_language,
        ))
    if not paths:
        raise ValueError("No compatible known-standard or unknown-sample prediction sheet was found")
    return paths


# Preferred public entry point used by the Tkinter application.
replot_simple_result_plots_from_workbook = replot_all_prediction_plots_from_workbook


def replot_known_standards_from_workbook(
    workbook_path: Path,
    *,
    output_path: Optional[Path] = None,
    output_language: str = "en",
) -> Path:
    """Backward-compatible helper that regenerates only the known-standard plot."""
    from openpyxl import load_workbook
    workbook_path = Path(workbook_path)
    wb = load_workbook(str(workbook_path), read_only=True, data_only=True)
    sheet_name = "Known_Standards_Predictions"
    if sheet_name not in wb.sheetnames:
        for candidate in ("Level_Known_Standards", "Concentration_Ratio_Comparison", "Raw_vs_Corrected"):
            if candidate in wb.sheetnames:
                sheet_name = candidate
                break
        else:
            wb.close()
            raise ValueError("No known-standard prediction sheet was found in the selected workbook")
    rows = _normalize_known_rows(_read_sheet_rows(wb, sheet_name))
    wb.close()
    if output_path is None:
        output_path = workbook_path.with_name(workbook_path.stem + "__known_actual_vs_predicted.png")
    return make_known_actual_predicted_plot(rows, Path(output_path), output_language=output_language)
