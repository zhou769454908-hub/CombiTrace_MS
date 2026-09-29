"""Standalone/private SMILES assembly and preview workflow.

This module intentionally does not call the response-factor model.  Its only
job is to resolve A/B/C/CORE structures for each Combo, generate candidate
product structures locally, and write an auditable preview workbook.

Two chemically explicit assembly routes are supported:

1. Reaction SMARTS (recommended)
   The private reaction template defines connectivity and atom loss.
2. Mapped-dummy stitching
   Matching dummy atoms such as ``[*:1]`` on the core and a fragment are
   removed and their neighbouring atoms are connected.  In this mode the
   supplied fragment SMILES must already represent the desired reactive/post-
   loss fragments.  Stoichiometric losses (for example B loses N2 and C loses
   H) are used only for formula validation and cannot, by themselves, determine
   atom connectivity.

All structures are processed locally.  Nothing is uploaded.
"""
from __future__ import annotations

import csv
import hashlib
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from .chemistry import monoisotopic_mass, parse_formula
from .response_predictor import PreparedRow, _norm_header, parse_combo, read_table
from .structure_features import (
    RDKIT_AVAILABLE,
    RDKIT_ERROR,
    StructureLibrary,
    calculate_product_formula,
    formula_equal,
    load_structure_library,
    resolve_core_id,
)

if RDKIT_AVAILABLE:
    from rdkit import Chem
    from rdkit.Chem import Descriptors, Draw, rdDepictor, rdMolDescriptors
    from rdkit.Chem.rdChemReactions import ReactionFromSmarts
else:  # pragma: no cover - handled by runtime guard
    Chem = None
    Descriptors = None
    Draw = None
    rdDepictor = None
    rdMolDescriptors = None
    ReactionFromSmarts = None


BUILD_MODE_OPTIONS: Tuple[Tuple[str, str], ...] = (
    ('Automatic: product SMILES, reaction SMARTS, then mapped-dummy assembly', "auto"),
    ('Reaction SMARTS', "reaction_smarts"),
    ('Mapped-dummy assembly [*:1]', "dummy_stitch"),
    ('Validate supplied product SMILES only', "direct_product"),
)
BUILD_MODE_LABELS = [x[0] for x in BUILD_MODE_OPTIONS]
BUILD_MODE_MAP = {x[0]: x[1] for x in BUILD_MODE_OPTIONS}


@dataclass
class ProductCandidate:
    rank: int
    smiles: str
    formula: str
    exact_mass: Optional[float]
    inchikey: str
    heavy_atoms: int
    formula_match: str
    method: str
    score: Tuple[float, float, float, str]
    warnings: List[str] = field(default_factory=list)


@dataclass
class BuildResult:
    source_file: str
    source_sheet: str
    source_row: object
    source_index: int
    combo: str
    expected_formula_input: str
    expected_formula_calculated: str
    expected_formula_used: str
    formula_input_vs_calculated: str
    core_id: str
    labels: Dict[str, str]
    formulas: Dict[str, str]
    smiles: Dict[str, str]
    smiles_hashes: Dict[str, str]
    reaction_rule_id: str
    reaction_smarts_present: bool
    reactant_order: Tuple[str, ...]
    requested_mode: str
    selected_method: str
    status: str
    candidates: List[ProductCandidate]
    selected: Optional[ProductCandidate]
    image_path: str = ""
    warnings: List[str] = field(default_factory=list)


@dataclass
class SmilesBuildReport:
    output_xlsx: Path
    output_csv: Optional[Path]
    image_dir: Optional[Path]
    n_total: int
    n_success: int
    n_failed: int
    n_formula_match: int
    warnings: List[str]
    results: List[BuildResult] = field(default_factory=list)


def _text(v: object) -> str:
    return "" if v is None else str(v).strip()


def _safe_file_part(text: str, max_len: int = 80) -> str:
    s = re.sub(r"[^0-9A-Za-z._-]+", "_", str(text or "").strip())
    s = s.strip("._-") or "product"
    return s[:max_len]


def _hash_text(text: str) -> str:
    if not text:
        return ""
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()[:16]


def _row_value(raw: Dict[str, object], names: Sequence[str]) -> str:
    norm = {_norm_header(k): k for k in raw if not str(k).startswith("__source_")}
    for name in names:
        n = _norm_header(name)
        if n in norm:
            v = _text(raw.get(norm[n], ""))
            if v:
                return v
    return ""


def _guess_col(headers: Sequence[str], names: Sequence[str], *, contains: bool = True) -> str:
    norm = {_norm_header(h): h for h in headers}
    for name in names:
        n = _norm_header(name)
        if n in norm:
            return norm[n]
    if contains:
        for h in headers:
            nh = _norm_header(h)
            for name in names:
                n = _norm_header(name)
                if n and n in nh:
                    return h
    return ""


def _prepared_row(raw: Dict[str, object], index: int, combo_col: str, formula_col: str) -> PreparedRow:
    combo = _text(raw.get(combo_col, "")) if combo_col else _row_value(raw, ["Combo", 'Combo', 'Source_Combination'])
    formula = _text(raw.get(formula_col, "")) if formula_col else _row_value(raw, ["Formula", "Molecular Formula", 'Formula'])
    counts: Dict[str, int] = {}
    if formula:
        try:
            counts = {str(k): int(v) for k, v in parse_formula(formula).items()}
        except Exception:
            pass
    exact_mass = None
    if counts:
        try:
            exact_mass = float(monoisotopic_mass(counts))
        except Exception:
            pass
    return PreparedRow(
        source_index=index,
        raw=raw,
        combo=combo,
        formula=formula,
        name=_row_value(raw, ["Name", "Compound", 'Name']),
        ratio=None,
        concentration=None,
        exact_mass=exact_mass,
        correction_factor=None,
        log_correction_factor=None,
        combo_parts=parse_combo(combo),
        element_counts=counts,
        warnings=[],
    )


def _parse_mol(smiles: str):
    if not RDKIT_AVAILABLE or not smiles:
        return None
    try:
        mol = Chem.MolFromSmiles(smiles)
        if mol is None:
            return None
        Chem.SanitizeMol(mol)
        return mol
    except Exception:
        return None


def _canonical(mol) -> str:
    try:
        return Chem.MolToSmiles(mol, canonical=True, isomericSmiles=True)
    except Exception:
        return ""


def _split_fragments(mol) -> List[object]:
    try:
        frags = list(Chem.GetMolFrags(mol, asMols=True, sanitizeFrags=True))
        return frags if frags else [mol]
    except Exception:
        return [mol]


def _candidate_from_mol(
    mol,
    *,
    method: str,
    expected_formula: str,
    rank: int = 0,
    warnings: Optional[Sequence[str]] = None,
) -> Optional[ProductCandidate]:
    if mol is None:
        return None
    try:
        Chem.SanitizeMol(mol)
        smi = _canonical(mol)
        if not smi:
            return None
        formula = str(rdMolDescriptors.CalcMolFormula(mol))
        mass = float(Descriptors.ExactMolWt(mol))
        heavy = int(mol.GetNumHeavyAtoms())
        try:
            inchikey = str(Chem.MolToInchiKey(mol))
        except Exception:
            inchikey = ""
        if expected_formula:
            fmatch = "yes" if formula_equal(formula, expected_formula) else "no"
        else:
            fmatch = "unknown"
        mass_delta = 1e9
        if expected_formula:
            try:
                em = float(monoisotopic_mass(parse_formula(expected_formula)))
                mass_delta = abs(mass - em)
            except Exception:
                pass
        score = (
            2.0 if fmatch == "yes" else (1.0 if fmatch == "unknown" else 0.0),
            -float(mass_delta),
            float(heavy),
            smi,
        )
        return ProductCandidate(
            rank=rank,
            smiles=smi,
            formula=formula,
            exact_mass=mass,
            inchikey=inchikey,
            heavy_atoms=heavy,
            formula_match=fmatch,
            method=method,
            score=score,
            warnings=list(warnings or []),
        )
    except Exception:
        return None


def _deduplicate_candidates(candidates: Iterable[ProductCandidate]) -> List[ProductCandidate]:
    best: Dict[str, ProductCandidate] = {}
    for c in candidates:
        old = best.get(c.smiles)
        if old is None or c.score > old.score:
            best[c.smiles] = c
    out = sorted(best.values(), key=lambda x: x.score, reverse=True)
    for i, c in enumerate(out, start=1):
        c.rank = i
    return out


def _reaction_candidates(
    smiles_by_role: Dict[str, str],
    reaction_smarts: str,
    reactant_order: Sequence[str],
    expected_formula: str,
    *,
    max_candidates: int = 100,
) -> Tuple[List[ProductCandidate], List[str]]:
    warnings: List[str] = []
    if not reaction_smarts:
        return [], ["Reaction SMARTS is empty"]
    reactants = []
    for role in reactant_order:
        role_u = str(role or "").strip().upper()
        smi = smiles_by_role.get(role_u, "")
        mol = _parse_mol(smi)
        if mol is None:
            return [], [f"Missing or invalid {role_u}_SMILES required by Reactant_Order"]
        reactants.append(mol)
    try:
        rxn = ReactionFromSmarts(reaction_smarts)
        if rxn is None:
            return [], ["Reaction SMARTS could not be parsed"]
        try:
            rxn.Initialize()
        except Exception:
            pass
        product_sets = rxn.RunReactants(tuple(reactants))
    except Exception as e:
        return [], [f"Reaction SMARTS execution failed: {e}"]

    raw_candidates: List[ProductCandidate] = []
    n_sets = len(product_sets)
    if n_sets == 0:
        warnings.append("Reaction SMARTS produced no products")
    for pset_idx, pset in enumerate(product_sets, start=1):
        for product_idx, product in enumerate(pset, start=1):
            # Keep both the product object and its disconnected fragments as
            # candidates.  Formula matching normally identifies the intended
            # major product while leaving all alternatives auditable.
            mols = [product]
            frags = _split_fragments(product)
            if len(frags) > 1:
                mols.extend(frags)
            for mol in mols:
                cand = _candidate_from_mol(
                    mol,
                    method="reaction_smarts",
                    expected_formula=expected_formula,
                    warnings=[f"product_set={pset_idx}; product={product_idx}"],
                )
                if cand is not None:
                    raw_candidates.append(cand)
                    if len(raw_candidates) >= max(1, int(max_candidates)) * 4:
                        break
            if len(raw_candidates) >= max(1, int(max_candidates)) * 4:
                break
        if len(raw_candidates) >= max(1, int(max_candidates)) * 4:
            break
    candidates = _deduplicate_candidates(raw_candidates)[: max(1, int(max_candidates))]
    if n_sets > 1:
        warnings.append(f"Reaction SMARTS generated {n_sets} product sets; candidates were ranked by formula match and molecular size")
    return candidates, warnings


def _dummy_stitch_candidate(
    smiles_by_role: Dict[str, str],
    reactant_order: Sequence[str],
    expected_formula: str,
) -> Tuple[List[ProductCandidate], List[str]]:
    warnings: List[str] = []
    combined = None
    for role in reactant_order:
        role_u = str(role or "").strip().upper()
        smi = smiles_by_role.get(role_u, "")
        mol = _parse_mol(smi)
        if mol is None:
            return [], [f"Missing or invalid {role_u}_SMILES for mapped-dummy stitching"]
        combined = Chem.Mol(mol) if combined is None else Chem.CombineMols(combined, mol)
    if combined is None:
        return [], ["No reactants available"]

    rw = Chem.RWMol(combined)
    groups: Dict[int, List[int]] = {}
    unmapped: List[int] = []
    for atom in rw.GetAtoms():
        if atom.GetAtomicNum() != 0:
            continue
        amap = int(atom.GetAtomMapNum())
        if amap <= 0:
            unmapped.append(atom.GetIdx())
        else:
            groups.setdefault(amap, []).append(atom.GetIdx())
    if unmapped:
        warnings.append(f"Found {len(unmapped)} unmapped dummy atoms; every attachment dummy should use [*:number]")
    if not groups:
        return [], warnings + ["No mapped dummy atoms such as [*:1] were found"]

    to_remove: List[int] = []
    for amap in sorted(groups):
        indices = groups[amap]
        if len(indices) != 2:
            return [], warnings + [f"Attachment map {amap} occurs {len(indices)} times; exactly 2 occurrences are required"]
        info = []
        for idx in indices:
            atom = rw.GetAtomWithIdx(idx)
            nbrs = list(atom.GetNeighbors())
            if len(nbrs) != 1:
                return [], warnings + [f"Dummy [*:{amap}] must have exactly one neighbour"]
            nbr_idx = int(nbrs[0].GetIdx())
            bond = rw.GetBondBetweenAtoms(idx, nbr_idx)
            btype = bond.GetBondType() if bond is not None else Chem.BondType.SINGLE
            info.append((idx, nbr_idx, btype))
        (_, n1, b1), (_, n2, b2) = info
        if n1 == n2:
            return [], warnings + [f"Attachment map {amap} points to the same atom on both sides"]
        if rw.GetBondBetweenAtoms(n1, n2) is not None:
            warnings.append(f"Attachment map {amap}: neighbour atoms were already bonded")
        else:
            # A dummy-aromatic bond is normally single; prefer a non-single
            # type only when both sides agree.
            btype = b1 if b1 == b2 else Chem.BondType.SINGLE
            try:
                rw.AddBond(n1, n2, btype)
            except Exception:
                rw.AddBond(n1, n2, Chem.BondType.SINGLE)
                warnings.append(f"Attachment map {amap}: bond type was normalized to single")
        to_remove.extend(indices)

    for idx in sorted(set(to_remove + unmapped), reverse=True):
        try:
            rw.RemoveAtom(int(idx))
        except Exception:
            pass
    try:
        product = rw.GetMol()
        Chem.SanitizeMol(product)
    except Exception as e:
        return [], warnings + [f"Mapped-dummy product sanitization failed: {e}"]
    cand = _candidate_from_mol(product, method="dummy_stitch", expected_formula=expected_formula, warnings=warnings)
    return ([cand] if cand else []), warnings


def _resolve_components(
    row: PreparedRow,
    lib: StructureLibrary,
    core_id: str,
) -> Tuple[Dict[str, str], Dict[str, str], Dict[str, str], List[str]]:
    labels: Dict[str, str] = {}
    formulas: Dict[str, str] = {}
    smiles: Dict[str, str] = {}
    warnings: List[str] = []

    direct_names = {
        "A": ["A_SMILES", "A SMILES", "SMILES_A", 'A_SMILES'],
        "B": ["B_SMILES", "B SMILES", "SMILES_B", 'B_SMILES'],
        "C": ["C_SMILES", "C SMILES", "SMILES_C", 'C_SMILES'],
        "CORE": ["CORE_SMILES", "Core SMILES", "Middle_SMILES", 'Core_SMILES', 'Intermediate_SMILES'],
    }
    for role in ("A", "B", "C"):
        part = row.combo_parts.get(role, {})
        label = _text(part.get("label", ""))
        labels[role] = label
        entry = lib.entries.get((role, label)) if label else None
        formulas[role] = (entry.formula if entry else "") or _text(part.get("formula", ""))
        smiles[role] = _row_value(row.raw, direct_names[role]) or (entry.smiles if entry else "")
        if not label:
            warnings.append(f"Combo does not contain {role} component")
        elif entry is None and not smiles[role]:
            warnings.append(f"No structure-library entry or direct SMILES for {role}/{label}")
    labels["CORE"] = core_id
    core = lib.cores.get(core_id) if core_id else None
    formulas["CORE"] = core.formula if core else ""
    smiles["CORE"] = _row_value(row.raw, direct_names["CORE"]) or (core.smiles if core else "")
    if not core_id:
        warnings.append("Core_ID could not be resolved")
    elif core is None and not smiles["CORE"]:
        warnings.append(f"No structure-library entry or direct CORE_SMILES for {core_id}")
    return labels, formulas, smiles, warnings


def build_one(
    raw: Dict[str, object],
    source_index: int,
    lib: StructureLibrary,
    *,
    combo_col: str = "",
    formula_col: str = "",
    core_col: str = "",
    mode: str = "auto",
    reaction_smarts_override: str = "",
    reactant_order_override: str = "",
    max_candidates: int = 100,
) -> BuildResult:
    row = _prepared_row(raw, source_index, combo_col, formula_col)
    warnings: List[str] = list(row.warnings)
    source_file = _text(raw.get("__source_file__", ""))
    source_sheet = _text(raw.get("__source_sheet__", ""))
    source_row = raw.get("__source_row__", source_index)

    direct_core_id = _text(raw.get(core_col, "")) if core_col else _row_value(raw, ["Core_ID", "Middle_ID", "Fixed_Core_ID", 'Core_ID'])
    if direct_core_id:
        core_id = direct_core_id
    else:
        core_id, ws = resolve_core_id(row, lib)
        warnings.extend(ws)

    if lib.entries or lib.cores:
        calc_formula, ws = calculate_product_formula(row, lib, core_id)
        warnings.extend(ws)
    else:
        calc_formula = ""
    input_formula = row.formula
    expected_formula = input_formula or calc_formula
    if input_formula and calc_formula:
        input_vs_calc = "yes" if formula_equal(input_formula, calc_formula) else "no"
        if input_vs_calc == "no":
            warnings.append(f"Input Formula {input_formula} differs from calculated stoichiometric Formula {calc_formula}")
    else:
        input_vs_calc = "unknown"

    labels, formulas, smiles, ws = _resolve_components(row, lib, core_id)
    warnings.extend(ws)
    smiles_hashes = {role: _hash_text(smi) for role, smi in smiles.items()}

    rule = lib.rules.get(core_id)
    rule_id = rule.rule_id if rule else ""
    reaction_smarts = str(reaction_smarts_override or (rule.reaction_smarts if rule else "") or "").strip()
    if reactant_order_override.strip():
        reactant_order = tuple(x.strip().upper() for x in re.split(r"[,;|>]", reactant_order_override) if x.strip())
    elif rule and rule.reactant_order:
        reactant_order = tuple(str(x).upper() for x in rule.reactant_order)
    else:
        reactant_order = ("CORE", "A", "B", "C")

    direct_product = _row_value(raw, ["Product_SMILES", "Product SMILES", 'Product_SMILES'])
    if not direct_product and row.combo:
        product_entry = lib.products.get(re.sub(r"\s+", "", row.combo).lower())
        if product_entry:
            direct_product = product_entry.product_smiles or product_entry.smiles

    requested_mode = str(mode or "auto").strip().lower()
    selected_method = ""
    candidates: List[ProductCandidate] = []

    def use_direct() -> bool:
        nonlocal candidates, selected_method
        if not direct_product:
            return False
        mol = _parse_mol(direct_product)
        if mol is None:
            warnings.append("Product_SMILES exists but cannot be parsed")
            return False
        cand = _candidate_from_mol(mol, method="direct_product", expected_formula=expected_formula)
        if cand:
            candidates = [cand]
            selected_method = "direct_product"
            return True
        return False

    def use_reaction() -> bool:
        nonlocal candidates, selected_method
        cs, ws2 = _reaction_candidates(
            smiles,
            reaction_smarts,
            reactant_order,
            expected_formula,
            max_candidates=max_candidates,
        )
        warnings.extend(ws2)
        if cs:
            candidates = cs
            selected_method = "reaction_smarts"
            return True
        return False

    def use_dummy() -> bool:
        nonlocal candidates, selected_method
        cs, ws2 = _dummy_stitch_candidate(smiles, reactant_order, expected_formula)
        warnings.extend(ws2)
        if cs:
            candidates = cs
            selected_method = "dummy_stitch"
            return True
        return False

    if requested_mode == "direct_product":
        use_direct()
    elif requested_mode == "reaction_smarts":
        use_reaction()
    elif requested_mode == "dummy_stitch":
        use_dummy()
    else:
        if not use_direct():
            if reaction_smarts:
                if not use_reaction():
                    use_dummy()
            else:
                use_dummy()

    candidates = _deduplicate_candidates(candidates)
    selected = candidates[0] if candidates else None
    if selected is not None:
        status = "SUCCESS"
        if selected.formula_match == "no":
            status = "CHECK_FORMULA"
            warnings.append("Selected product formula does not match the expected formula")
    else:
        status = "FAILED"
        if not warnings:
            warnings.append("No product candidate was generated")

    return BuildResult(
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=source_row,
        source_index=source_index,
        combo=row.combo,
        expected_formula_input=input_formula,
        expected_formula_calculated=calc_formula,
        expected_formula_used=expected_formula,
        formula_input_vs_calculated=input_vs_calc,
        core_id=core_id,
        labels=labels,
        formulas=formulas,
        smiles=smiles,
        smiles_hashes=smiles_hashes,
        reaction_rule_id=rule_id,
        reaction_smarts_present=bool(reaction_smarts),
        reactant_order=reactant_order,
        requested_mode=requested_mode,
        selected_method=selected_method,
        status=status,
        candidates=candidates,
        selected=selected,
        warnings=list(dict.fromkeys([str(x) for x in warnings if str(x).strip()])),
    )


def _draw_product(smiles: str, path: Path, legend: str = "") -> bool:
    mol = _parse_mol(smiles)
    if mol is None:
        return False
    try:
        rdDepictor.Compute2DCoords(mol)
        img = Draw.MolToImage(mol, size=(900, 560), legend=legend[:120])
        path.parent.mkdir(parents=True, exist_ok=True)
        img.save(str(path))
        return True
    except Exception:
        return False


def _write_csv(results: Sequence[BuildResult], path: Path) -> None:
    headers = [
        "Source_File", "Source_Sheet", "Source_Row", "Combo", "Core_ID",
        "A_ID", "B_ID", "C_ID", "Build_Method", "Status", "Candidate_Count",
        "Selected_Product_SMILES", "Product_Formula", "Formula_Match", "Exact_Mass",
        "InChIKey", "Structure_PNG", "Warnings",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        for r in results:
            s = r.selected
            w.writerow({
                "Source_File": r.source_file,
                "Source_Sheet": r.source_sheet,
                "Source_Row": r.source_row,
                "Combo": r.combo,
                "Core_ID": r.core_id,
                "A_ID": r.labels.get("A", ""),
                "B_ID": r.labels.get("B", ""),
                "C_ID": r.labels.get("C", ""),
                "Build_Method": r.selected_method,
                "Status": r.status,
                "Candidate_Count": len(r.candidates),
                "Selected_Product_SMILES": s.smiles if s else "",
                "Product_Formula": s.formula if s else "",
                "Formula_Match": s.formula_match if s else "",
                "Exact_Mass": f"{s.exact_mass:.6f}" if s and s.exact_mass is not None else "",
                "InChIKey": s.inchikey if s else "",
                "Structure_PNG": r.image_path,
                "Warnings": "; ".join(r.warnings),
            })


def _write_xlsx(
    results: Sequence[BuildResult],
    path: Path,
    *,
    include_input_smiles: bool,
    include_all_candidates: bool,
    embed_images: bool,
    embed_image_limit: int,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as e:
        raise RuntimeError("openpyxl is required to write the SMILES preview workbook") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Generated_Products"
    headers = [
        "Source_File", "Source_Sheet", "Source_Row", "Source_Index", "Combo",
        "Expected_Formula_Input", "Expected_Formula_Calculated", "Expected_Formula_Used",
        "Input_vs_Calculated_Formula", "Core_ID", "A_ID", "B_ID", "C_ID",
        "Requested_Mode", "Build_Method", "Status", "Candidate_Count",
        "Selected_Product_SMILES", "Product_Formula", "Formula_Match", "Exact_Mass",
        "InChIKey", "Heavy_Atoms", "Reaction_Rule_ID", "Reaction_SMARTS_Present",
        "Reactant_Order", "A_SMILES_Hash", "B_SMILES_Hash", "C_SMILES_Hash",
        "CORE_SMILES_Hash", "Structure_PNG", "Warnings",
    ]
    if include_input_smiles:
        headers.extend(["A_SMILES", "B_SMILES", "C_SMILES", "CORE_SMILES"])
    if embed_images:
        headers.append("2D_Preview")
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    image_col = headers.index("2D_Preview") + 1 if embed_images else 0
    embedded = 0
    for i, r in enumerate(results, start=2):
        s = r.selected
        row = [
            r.source_file, r.source_sheet, r.source_row, r.source_index, r.combo,
            r.expected_formula_input, r.expected_formula_calculated, r.expected_formula_used,
            r.formula_input_vs_calculated, r.core_id, r.labels.get("A", ""), r.labels.get("B", ""), r.labels.get("C", ""),
            r.requested_mode, r.selected_method, r.status, len(r.candidates),
            s.smiles if s else "", s.formula if s else "", s.formula_match if s else "",
            s.exact_mass if s else None, s.inchikey if s else "", s.heavy_atoms if s else None,
            r.reaction_rule_id, "Yes" if r.reaction_smarts_present else "No", ",".join(r.reactant_order),
            r.smiles_hashes.get("A", ""), r.smiles_hashes.get("B", ""), r.smiles_hashes.get("C", ""), r.smiles_hashes.get("CORE", ""),
            r.image_path, "; ".join(r.warnings),
        ]
        if include_input_smiles:
            row.extend([r.smiles.get("A", ""), r.smiles.get("B", ""), r.smiles.get("C", ""), r.smiles.get("CORE", "")])
        if embed_images:
            row.append("")
        ws.append(row)
        ws.cell(i, headers.index("Exact_Mass") + 1).number_format = "0.000000"
        if embed_images and r.image_path and Path(r.image_path).exists() and (embed_image_limit <= 0 or embedded < embed_image_limit):
            try:
                img = XLImage(r.image_path)
                img.width = 300
                img.height = 185
                img.anchor = ws.cell(i, image_col).coordinate
                ws.add_image(img)
                ws.row_dimensions[i].height = 145
                embedded += 1
            except Exception:
                pass

    widths = {
        "A": 18, "B": 18, "C": 12, "D": 12, "E": 70,
        "F": 24, "G": 28, "H": 24, "I": 16, "J": 16,
        "K": 16, "L": 16, "M": 16, "N": 18, "O": 20, "P": 16,
        "Q": 16, "R": 70, "S": 22, "T": 14, "U": 16, "V": 30,
        "W": 14, "X": 18, "Y": 18, "Z": 24,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for col in range(1, len(headers) + 1):
        if ws.column_dimensions[get_column_letter(col)].width is None:
            ws.column_dimensions[get_column_letter(col)].width = 18
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if include_all_candidates:
        wc = wb.create_sheet("All_Candidates")
        cheaders = [
            "Source_Index", "Source_Row", "Combo", "Candidate_Rank", "Selected", "Method",
            "Product_SMILES", "Product_Formula", "Formula_Match", "Exact_Mass", "InChIKey",
            "Heavy_Atoms", "Warnings",
        ]
        wc.append(cheaders)
        for cell in wc[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in results:
            for c in r.candidates:
                wc.append([
                    r.source_index, r.source_row, r.combo, c.rank,
                    "Yes" if r.selected and c.smiles == r.selected.smiles else "No",
                    c.method, c.smiles, c.formula, c.formula_match, c.exact_mass,
                    c.inchikey, c.heavy_atoms, "; ".join(c.warnings),
                ])
        wc.freeze_panes = "A2"
        wc.auto_filter.ref = f"A1:{get_column_letter(len(cheaders))}{max(1, wc.max_row)}"
        for col, width in enumerate([12, 12, 70, 14, 12, 20, 80, 22, 14, 16, 30, 14, 50], start=1):
            wc.column_dimensions[get_column_letter(col)].width = width
        for row in wc.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.column == 10:
                    cell.number_format = "0.000000"

    wa = wb.create_sheet("Input_Audit")
    aheaders = [
        "Source_Index", "Source_Row", "Combo", "Core_ID", "A_ID", "B_ID", "C_ID",
        "A_Formula", "B_Formula", "C_Formula", "CORE_Formula",
        "A_SMILES_Present", "B_SMILES_Present", "C_SMILES_Present", "CORE_SMILES_Present",
        "Reaction_Rule_ID", "Reaction_SMARTS_Present", "Reactant_Order", "Status", "Warnings",
    ]
    wa.append(aheaders)
    for cell in wa[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in results:
        wa.append([
            r.source_index, r.source_row, r.combo, r.core_id,
            r.labels.get("A", ""), r.labels.get("B", ""), r.labels.get("C", ""),
            r.formulas.get("A", ""), r.formulas.get("B", ""), r.formulas.get("C", ""), r.formulas.get("CORE", ""),
            "Yes" if r.smiles.get("A") else "No", "Yes" if r.smiles.get("B") else "No",
            "Yes" if r.smiles.get("C") else "No", "Yes" if r.smiles.get("CORE") else "No",
            r.reaction_rule_id, "Yes" if r.reaction_smarts_present else "No",
            ",".join(r.reactant_order), r.status, "; ".join(r.warnings),
        ])
    wa.freeze_panes = "A2"
    wa.auto_filter.ref = f"A1:{get_column_letter(len(aheaders))}{max(1, wa.max_row)}"
    for col in range(1, len(aheaders) + 1):
        wa.column_dimensions[get_column_letter(col)].width = 18
    wa.column_dimensions["C"].width = 70
    wa.column_dimensions["T"].width = 65
    for row in wa.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wi = wb.create_sheet("Instructions")
    instructions = [
        ("Purpose", "Generate and inspect product SMILES only. This module does not run the response-factor model."),
        ("Recommended route", "Use a private Reaction SMARTS to define the actual bonds and atom losses. Stoichiometric losses alone cannot determine a unique structure."),
        ("Mapped-dummy route", "Place matching mapped dummy atoms such as [*:1] on two attachment sites. The software removes the two dummies and connects their neighbours."),
        ("B/C loss", "B_Loss=N2 and C_Loss=H are used for formula validation. In dummy-stitch mode, supply post-loss/reactive fragments or encode the leaving groups in Reaction SMARTS."),
        ("Direct columns", "A_SMILES, B_SMILES, C_SMILES, CORE_SMILES and Product_SMILES in the Combo table override the structure library for that row."),
        ("Selection", "When multiple products are generated, formula match is ranked first, followed by exact-mass agreement and heavy-atom count. All candidates remain in All_Candidates."),
        ("Privacy", "All SMILES and Reaction SMARTS are processed locally. Input fragment/core SMILES are omitted from the result unless explicitly enabled."),
        ("Validation", "A successful generated SMILES is a computational assembly candidate, not structural confirmation. Review the 2D image, formula match and chemistry before downstream use."),
    ]
    wi.append(["Item", "Description"])
    for cell in wi[1]:
        cell.fill = header_fill
        cell.font = header_font
    for item, desc in instructions:
        wi.append([item, desc])
    wi.column_dimensions["A"].width = 24
    wi.column_dimensions["B"].width = 110
    for row in wi.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def run_smiles_builder(
    combo_file: Path,
    output_xlsx: Path,
    *,
    structure_file: Optional[Path] = None,
    combo_sheet: str = "",
    structure_sheet: str = "Structures",
    reaction_sheet: str = "ReactionRules",
    combo_col: str = "",
    formula_col: str = "",
    core_col: str = "",
    mode: str = "auto",
    reaction_smarts_override: str = "",
    reactant_order_override: str = "",
    generate_png: bool = True,
    image_limit: int = 100,
    embed_images: bool = True,
    embed_image_limit: int = 30,
    include_input_smiles: bool = False,
    include_all_candidates: bool = True,
    max_candidates: int = 100,
    output_csv: bool = True,
    progress: Optional[Callable[[str], None]] = None,
) -> SmilesBuildReport:
    if not RDKIT_AVAILABLE:
        raise RuntimeError(
            "RDKit is required for SMILES generation. Install it in the same environment with: "
            "conda install -c conda-forge rdkit. Details: " + str(RDKIT_ERROR)
        )
    combo_file = Path(combo_file)
    output_xlsx = Path(output_xlsx)
    if not combo_file.exists():
        raise FileNotFoundError(str(combo_file))

    table = read_table(combo_file, sheet_name=combo_sheet)
    ccol = combo_col or _guess_col(table.headers, ["Combo", 'Combo', 'Source_Combination', "Source Combo"])
    fcol = formula_col or _guess_col(table.headers, ["Formula", "Molecular Formula", 'Formula'])
    corecol = core_col or _guess_col(table.headers, ["Core_ID", "Middle_ID", "Fixed_Core_ID", 'Core_ID'])
    if not ccol:
        raise ValueError("Combo column could not be detected. Specify the Combo column name in the interface.")

    lib = load_structure_library(Path(structure_file) if structure_file else None, structure_sheet, reaction_sheet)
    warnings = list(lib.warnings)
    if progress:
        progress(f"RDKit_available: {RDKIT_AVAILABLE}")
        progress(f"Combo rows: {len(table.rows)} | sheet={table.sheet_name or '(CSV)'} | Combo column={ccol}")
        progress(f"Structure entries: {len(lib.entries)} | cores={len(lib.cores)} | rules={len(lib.rules)}")

    image_dir: Optional[Path] = None
    if generate_png:
        image_dir = output_xlsx.with_name(output_xlsx.stem + "__structures")
        image_dir.mkdir(parents=True, exist_ok=True)
        for old in image_dir.glob("*.png"):
            try:
                old.unlink()
            except Exception:
                pass

    results: List[BuildResult] = []
    n_total = len(table.rows)
    made_images = 0
    for idx, raw in enumerate(table.rows, start=1):
        result = build_one(
            raw,
            idx,
            lib,
            combo_col=ccol,
            formula_col=fcol,
            core_col=corecol,
            mode=mode,
            reaction_smarts_override=reaction_smarts_override,
            reactant_order_override=reactant_order_override,
            max_candidates=max_candidates,
        )
        if image_dir is not None and result.selected is not None and (image_limit <= 0 or made_images < image_limit):
            name = f"{idx:04d}_{_safe_file_part(result.labels.get('A',''))}_{_safe_file_part(result.labels.get('B',''))}_{_safe_file_part(result.labels.get('C',''))}_{_hash_text(result.selected.smiles)[:8]}.png"
            ip = image_dir / name
            if _draw_product(result.selected.smiles, ip, legend=f"{result.labels.get('A','')} | {result.labels.get('B','')} | {result.labels.get('C','')}"):
                result.image_path = str(ip)
                made_images += 1
        results.append(result)
        if progress and (idx == 1 or idx == n_total or idx % 20 == 0):
            progress(f"[{idx}/{n_total}] {result.status}: {result.combo[:100]}")

    _write_xlsx(
        results,
        output_xlsx,
        include_input_smiles=include_input_smiles,
        include_all_candidates=include_all_candidates,
        embed_images=embed_images,
        embed_image_limit=embed_image_limit,
    )
    out_csv = output_xlsx.with_suffix(".csv") if output_csv else None
    if out_csv is not None:
        _write_csv(results, out_csv)

    n_success = sum(1 for r in results if r.selected is not None)
    n_failed = n_total - n_success
    n_formula_match = sum(1 for r in results if r.selected is not None and r.selected.formula_match == "yes")
    if progress:
        progress(f"Completed: success={n_success}, failed={n_failed}, formula_match={n_formula_match}")
        progress(f"Excel: {output_xlsx}")
        if out_csv:
            progress(f"CSV: {out_csv}")
        if image_dir:
            progress(f"2D images: {image_dir} ({made_images})")
    return SmilesBuildReport(
        output_xlsx=output_xlsx,
        output_csv=out_csv,
        image_dir=image_dir,
        n_total=n_total,
        n_success=n_success,
        n_failed=n_failed,
        n_formula_match=n_formula_match,
        warnings=warnings,
        results=results,
    )


def preview_first_product(
    combo_file: Path,
    *,
    structure_file: Optional[Path] = None,
    combo_sheet: str = "",
    structure_sheet: str = "Structures",
    reaction_sheet: str = "ReactionRules",
    combo_col: str = "",
    formula_col: str = "",
    core_col: str = "",
    mode: str = "auto",
    reaction_smarts_override: str = "",
    reactant_order_override: str = "",
    preview_dir: Optional[Path] = None,
    max_candidates: int = 100,
) -> BuildResult:
    if not RDKIT_AVAILABLE:
        raise RuntimeError(
            "RDKit is required for SMILES generation. Install it with: conda install -c conda-forge rdkit"
        )
    table = read_table(Path(combo_file), sheet_name=combo_sheet)
    ccol = combo_col or _guess_col(table.headers, ["Combo", 'Combo', 'Source_Combination', "Source Combo"])
    fcol = formula_col or _guess_col(table.headers, ["Formula", "Molecular Formula", 'Formula'])
    corecol = core_col or _guess_col(table.headers, ["Core_ID", "Middle_ID", "Fixed_Core_ID", 'Core_ID'])
    if not ccol:
        raise ValueError("Combo column could not be detected")
    lib = load_structure_library(Path(structure_file) if structure_file else None, structure_sheet, reaction_sheet)
    last: Optional[BuildResult] = None
    for idx, raw in enumerate(table.rows, start=1):
        r = build_one(
            raw,
            idx,
            lib,
            combo_col=ccol,
            formula_col=fcol,
            core_col=corecol,
            mode=mode,
            reaction_smarts_override=reaction_smarts_override,
            reactant_order_override=reactant_order_override,
            max_candidates=max_candidates,
        )
        last = r
        if r.selected is not None:
            if preview_dir is not None:
                preview_dir = Path(preview_dir)
                preview_dir.mkdir(parents=True, exist_ok=True)
                ip = preview_dir / "smiles_builder_preview.png"
                if _draw_product(r.selected.smiles, ip, legend=f"{r.labels.get('A','')} | {r.labels.get('B','')} | {r.labels.get('C','')}"):
                    r.image_path = str(ip)
            return r
    if last is not None:
        return last
    raise ValueError("The Combo table contains no data rows")
