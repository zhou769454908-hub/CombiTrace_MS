"""Detailed local checks for mapped-dummy SMILES used by the simple builder.

Nothing is uploaded. The diagnostics identify the exact GUI field, worksheet,
row, component ID, map number, RDKit parse/sanitisation error and suggested fix.
"""
from __future__ import annotations

import hashlib
import logging
import re
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from .structure_features import RDKIT_AVAILABLE, RDKIT_ERROR

if RDKIT_AVAILABLE:
    from rdkit import Chem, rdBase
else:  # pragma: no cover
    Chem = None
    rdBase = None

EXPECTED_MAPS: Dict[str, Tuple[int, ...]] = {
    "CORE": (1, 2, 3),
    "A": (1,),
    "B": (2,),
    "C": (3,),
}

REPLACEMENTS: Dict[str, str] = {
    "\ufeff": "", "\u200b": "", "\u200c": "", "\u200d": "", "\u2060": "",
    "\uFF1A": ":", "\uFF3B": "[", "\uFF3D": "]", "\u3010": "[", "\u3011": "]",
    "\uFF08": "(", "\uFF09": ")", "\uFF0A": "*", "\uFF0B": "+", "\uFF0D": "-",
    "\u2212": "-", "\u2013": "-", "\u2014": "-", "\uFF1D": "=", "\uFF0C": ",", "\u3002": ".",
    "\u2018": "'", "\u2019": "'", "\u201C": '"', "\u201D": '"',
}

LOG_LOCK = threading.RLock()


@dataclass
class SmilesDiagnostic:
    role: str
    item_id: str = ""
    source_sheet: str = ""
    excel_row: Optional[int] = None
    source_label: str = ""
    raw_smiles: str = ""
    normalized_smiles: str = ""
    normalization_changes: List[str] = field(default_factory=list)
    status: str = "ERROR"
    parse_ok: bool = False
    sanitize_ok: bool = False
    canonical_smiles: str = ""
    smiles_hash: str = ""
    atom_count: int = 0
    heavy_atom_count: int = 0
    fragment_count: int = 0
    mapped_dummy_counts: Dict[int, int] = field(default_factory=dict)
    unmapped_dummy_count: int = 0
    real_atom_map_counts: Dict[int, int] = field(default_factory=dict)
    dummy_details: List[str] = field(default_factory=list)
    rdkit_message: str = ""
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    suggestions: List[str] = field(default_factory=list)

    @property
    def fatal(self) -> bool:
        return bool(self.errors)

    @property
    def location(self) -> str:
        if self.role == "CORE":
            return self.source_label or "GUI fixed CORE text box"
        sheet = self.source_sheet or "unknown sheet"
        row = f"row {self.excel_row}" if self.excel_row is not None else "unknown row"
        ident = f"/{self.item_id}" if self.item_id else ""
        return f"{sheet}!{row} {self.role}{ident}"

    def one_line(self) -> str:
        first = self.errors[0] if self.errors else (self.warnings[0] if self.warnings else "valid")
        return f"{self.role} [{self.location}] {self.status}: {first}"


@dataclass
class DiagnosticsReport:
    diagnostics: List[SmilesDiagnostic]
    global_errors: List[str] = field(default_factory=list)
    global_warnings: List[str] = field(default_factory=list)
    output_xlsx: Optional[Path] = None
    output_txt: Optional[Path] = None

    @property
    def error_count(self) -> int:
        return sum(1 for d in self.diagnostics if d.fatal) + len(self.global_errors)

    @property
    def warning_count(self) -> int:
        return sum(1 for d in self.diagnostics if d.warnings) + len(self.global_warnings)

    @property
    def ok_count(self) -> int:
        return sum(1 for d in self.diagnostics if not d.fatal)

    @property
    def overall_status(self) -> str:
        if self.error_count:
            return "FAILED"
        if self.warning_count:
            return "PASS_WITH_WARNINGS"
        return "PASS"


def add_unique(seq: List[str], value: str) -> None:
    value = str(value or "").strip()
    if value and value not in seq:
        seq.append(value)


def normalize_smiles(value: object) -> Tuple[str, List[str]]:
    raw = "" if value is None else str(value)
    s = raw
    changes: List[str] = []
    if s.startswith("'") and not s.endswith("'"):
        s = s[1:]
        changes.append("removed leading Excel text apostrophe")
    if s.strip() != s:
        s = s.strip()
        changes.append("trimmed leading/trailing whitespace")
    if len(s) >= 2 and ((s[0] == "`" and s[-1] == "`") or (s[0] == '"' and s[-1] == '"')):
        s = s[1:-1].strip()
        changes.append("removed surrounding quote/backtick characters")
    m = re.match(r"^\s*(?:CORE[_ ]?SMILES|A[_ ]?SMILES|B[_ ]?SMILES|C[_ ]?SMILES|SMILES)\s*[:=]\s*(.+)$", s, flags=re.I | re.S)
    if m:
        s = m.group(1).strip()
        changes.append("removed pasted field label such as 'SMILES:'")
    for old, new in REPLACEMENTS.items():
        if old in s:
            n = s.count(old)
            s = s.replace(old, new)
            changes.append(f"replaced {n} occurrence(s) of {old!r} (U+{ord(old):04X}) with {new!r}")
    compact = re.sub(r"\s+", "", s)
    if compact != s:
        s = compact
        changes.append("removed internal spaces/newlines/tabs")
    return s, changes


def clean_rdkit_log(text: str) -> str:
    out: List[str] = []
    for raw in str(text or "").splitlines():
        line = re.sub(r"^\[[0-9:]+\]\s*", "", raw).strip()
        if line and line not in out:
            out.append(line)
    return " | ".join(out)


def parse_with_messages(smiles: str):
    if not RDKIT_AVAILABLE:
        return None, False, "RDKit unavailable: " + str(RDKIT_ERROR)
    if not smiles:
        return None, False, "SMILES is empty"

    messages: List[str] = []

    class Capture(logging.Handler):
        def emit(self, record):  # type: ignore[override]
            try:
                messages.append(self.format(record))
            except Exception:
                pass

    handler = Capture()
    handler.setFormatter(logging.Formatter("%(message)s"))
    logger = logging.getLogger("rdkit")
    with LOG_LOCK:
        try:
            rdBase.LogToPythonLogger()
        except Exception:
            pass
        old_level, old_prop = logger.level, logger.propagate
        logger.addHandler(handler)
        logger.setLevel(logging.DEBUG)
        logger.propagate = False
        try:
            mol = Chem.MolFromSmiles(smiles, sanitize=False)
            if mol is None:
                return None, False, clean_rdkit_log("\n".join(messages)) or "RDKit could not parse this SMILES"
            try:
                Chem.SanitizeMol(mol)
            except Exception as e:
                msg = clean_rdkit_log("\n".join(messages))
                tail = f"Sanitization error: {type(e).__name__}: {e}"
                return mol, False, " | ".join(x for x in (msg, tail) if x)
            return mol, True, clean_rdkit_log("\n".join(messages))
        finally:
            logger.removeHandler(handler)
            logger.setLevel(old_level)
            logger.propagate = old_prop


def diagnose_smiles(
    role: str,
    raw_smiles: object,
    *,
    item_id: str = "",
    source_sheet: str = "",
    excel_row: Optional[int] = None,
    source_label: str = "",
    expected_maps: Optional[Sequence[int]] = None,
) -> SmilesDiagnostic:
    role = str(role or "").strip().upper() or "UNKNOWN"
    expected = tuple(int(x) for x in (expected_maps if expected_maps is not None else EXPECTED_MAPS.get(role, ())))
    raw = "" if raw_smiles is None else str(raw_smiles)
    normalized, changes = normalize_smiles(raw)
    d = SmilesDiagnostic(
        role=role, item_id=str(item_id or ""), source_sheet=str(source_sheet or ""),
        excel_row=excel_row, source_label=str(source_label or ""), raw_smiles=raw,
        normalized_smiles=normalized, normalization_changes=changes,
        smiles_hash=hashlib.sha256(normalized.encode("utf-8", errors="ignore")).hexdigest()[:16] if normalized else "",
    )
    if not raw.strip():
        d.errors.append("SMILES is blank")
        if role == "CORE":
            d.suggestions.append("Paste the fixed CORE SMILES into the GUI text box; it must contain [*:1], [*:2] and [*:3].")
        elif expected:
            d.suggestions.append(f"Fill this Excel cell with the {role} mapped SMILES containing [*:{expected[0]}].")
        return d

    non_ascii = [(i, ch) for i, ch in enumerate(raw, start=1) if ord(ch) > 127]
    for i, ch in non_ascii:
        add_unique(d.warnings, f"non-ASCII character at position {i}: {ch!r} (U+{ord(ch):04X})")
    if re.search(r"\[\s*\*\s*\d+\s*\]", normalized):
        add_unique(d.errors, "dummy map syntax omits ':'; use [*:number], not [*number]")
        add_unique(d.suggestions, "For example, change [*1] to [*:1].")
    if re.search(r"(?:^|[^A-Za-z0-9])R[123](?:$|[^A-Za-z0-9])", normalized, flags=re.I):
        add_unique(d.warnings, "R1/R2/R3 labels were detected; ChemDraw R-group labels are not mapped dummy atoms")
        add_unique(d.suggestions, "Replace the attachment label with a terminal mapped dummy such as [*:1].")

    mol, sanitized, rdmsg = parse_with_messages(normalized)
    d.rdkit_message = rdmsg
    d.parse_ok = mol is not None
    d.sanitize_ok = bool(mol is not None and sanitized)
    if mol is None:
        add_unique(d.errors, "RDKit could not parse the SMILES")
        add_unique(d.errors, rdmsg)
        low = rdmsg.lower()
        if "unclosed ring" in low:
            add_unique(d.suggestions, "Check that every ring number occurs exactly twice, e.g. c1ccccc1.")
        if "parenth" in low or "branch" in low:
            add_unique(d.suggestions, "Check that every '(' has a matching ')'.")
        if "syntax error" in low:
            add_unique(d.suggestions, "Check the exact position shown by RDKit; full-width punctuation or a pasted 'SMILES:' label is a common cause.")
        if changes:
            add_unique(d.suggestions, "Inspect the Normalized_SMILES in the diagnostic file and replace the source text with it.")
        d.status = "ERROR"
        return d

    try:
        d.canonical_smiles = Chem.MolToSmiles(mol, canonical=True, isomericSmiles=True)
    except Exception:
        pass
    d.atom_count = int(mol.GetNumAtoms())
    d.heavy_atom_count = int(mol.GetNumHeavyAtoms())
    try:
        d.fragment_count = len(Chem.GetMolFrags(mol))
    except Exception:
        d.fragment_count = 1

    if not sanitized:
        add_unique(d.errors, "SMILES parsed but failed chemical valence/aromaticity sanitisation")
        add_unique(d.errors, rdmsg)
        low = rdmsg.lower()
        if "valence" in low:
            add_unique(d.suggestions, "The dummy may be attached to an atom already at full valence. Put RealAtom-[*:n] at a site where H/leaving group is replaced, or correct charge/bond order.")
        if "kekul" in low:
            add_unique(d.suggestions, "Check aromatic lower-case atoms and ring bonds near the attachment site.")

    mapped: Dict[int, int] = {}
    real_maps: Dict[int, int] = {}
    unmapped = 0
    for atom in mol.GetAtoms():
        amap = int(atom.GetAtomMapNum())
        if atom.GetAtomicNum() == 0:
            if amap <= 0:
                unmapped += 1
                d.dummy_details.append(f"dummy atom idx={atom.GetIdx()}: unmapped, degree={atom.GetDegree()}")
            else:
                mapped[amap] = mapped.get(amap, 0) + 1
                nbrs = list(atom.GetNeighbors())
                nbr = ",".join(f"{n.GetSymbol()}(idx={n.GetIdx()})" for n in nbrs) or "none"
                d.dummy_details.append(f"[*:{amap}] idx={atom.GetIdx()}: degree={atom.GetDegree()}, neighbour={nbr}")
                if atom.GetDegree() != 1:
                    add_unique(d.errors, f"dummy [*:{amap}] has {atom.GetDegree()} neighbours; exactly one is required")
                    add_unique(d.suggestions, f"Draw the attachment as RealAtom-[*:{amap}] with the dummy at the end of a bond.")
        elif amap > 0:
            real_maps[amap] = real_maps.get(amap, 0) + 1

    d.mapped_dummy_counts = mapped
    d.unmapped_dummy_count = unmapped
    d.real_atom_map_counts = real_maps
    if unmapped:
        add_unique(d.errors, f"found {unmapped} unmapped dummy atom(s) '*' or '[*]'")
        add_unique(d.suggestions, "Every attachment dummy needs a colon map number, such as [*:1].")
    for map_no, count in sorted(real_maps.items()):
        if map_no in expected:
            add_unique(d.errors, f"map number {map_no} is on {count} real atom(s), not on a dummy atom")
            add_unique(d.suggestions, f"Use a separate terminal [*:{map_no}] atom; ordinary atom-map labels on C/N/O are not stitching points.")
        else:
            add_unique(d.warnings, f"real atom map {map_no} is present and ignored by dummy stitching")

    for no in expected:
        count = mapped.get(no, 0)
        if count == 0:
            add_unique(d.errors, f"required mapped dummy [*:{no}] is missing")
            add_unique(d.suggestions, f"Add exactly one terminal [*:{no}] at the {role} connection site.")
        elif count > 1:
            add_unique(d.errors, f"mapped dummy [*:{no}] occurs {count} times in {role}; exactly one is required")
            add_unique(d.suggestions, f"Keep only one [*:{no}] in this {role} SMILES.")
    extras = sorted(no for no in mapped if no not in expected)
    if extras:
        add_unique(d.errors, f"unexpected mapped dummy number(s) in {role}: {', '.join(map(str, extras))}")
        if role == "CORE":
            add_unique(d.suggestions, "In simple mode CORE must contain only maps 1, 2 and 3, once each.")
        elif expected:
            add_unique(d.suggestions, f"In simple mode {role} must contain only [*:{expected[0]}].")
    if not mapped and not unmapped:
        add_unique(d.errors, "no dummy attachment atom was found")
        if expected:
            add_unique(d.suggestions, "Use mapped dummy syntax: " + ", ".join(f"[*:{x}]" for x in expected) + ".")
    if d.fragment_count > 1:
        add_unique(d.warnings, f"SMILES contains {d.fragment_count} disconnected fragments separated by '.'")
        add_unique(d.suggestions, "Check whether salts/counterions are intentional; the simple stitcher keeps all fragments.")
    if changes and d.parse_ok:
        add_unique(d.warnings, "copy/paste characters were automatically normalized in memory")
        add_unique(d.suggestions, "Replace the source text with Normalized_SMILES to avoid ambiguity on the next run.")

    if d.errors:
        d.status = "ERROR"
    elif d.warnings:
        d.status = "AUTO_FIXED" if changes else "WARNING"
    else:
        d.status = "OK"
    return d


def validate_pairing(diags: Sequence[SmilesDiagnostic]) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []
    cores = [d for d in diags if d.role == "CORE"]
    if not cores:
        errors.append("CORE diagnostic is missing")
    elif len(cores) > 1:
        warnings.append(f"{len(cores)} CORE inputs found; simple mode expects one fixed CORE")
    if cores:
        core = cores[0]
        for no in (1, 2, 3):
            if core.mapped_dummy_counts.get(no, 0) != 1:
                errors.append(f"CORE does not contain exactly one [*:{no}]")
    for role, no in (("A", 1), ("B", 2), ("C", 3)):
        rows = [d for d in diags if d.role == role]
        if not rows:
            errors.append(f"No {role} rows found")
            continue
        bad = [d for d in rows if d.fatal or d.mapped_dummy_counts.get(no, 0) != 1]
        if bad:
            errors.append(f"{len(bad)}/{len(rows)} {role} row(s) fail [*:{no}] validation")
    return errors, warnings


def diagnostics_text(report: DiagnosticsReport, *, include_raw: bool = False) -> str:
    lines = [
        "Mapped-SMILES input diagnostics",
        "=" * 76,
        f"Overall status: {report.overall_status}",
        f"Inputs checked: {len(report.diagnostics)} | nonfatal: {report.ok_count} | errors: {report.error_count} | warnings: {report.warning_count}",
        "",
        "Required: CORE=[*:1],[*:2],[*:3] once each; A=[*:1]; B=[*:2]; C=[*:3].",
        "Every dummy must be terminal and have exactly one real-atom neighbour.",
        "",
    ]
    if report.global_errors:
        lines.append("GLOBAL ERRORS")
        lines += ["  - " + x for x in report.global_errors]
        lines.append("")
    if report.global_warnings:
        lines.append("GLOBAL WARNINGS")
        lines += ["  - " + x for x in report.global_warnings]
        lines.append("")
    for i, d in enumerate(report.diagnostics, start=1):
        lines.append(f"[{i}] {d.role} | {d.location} | STATUS={d.status}")
        lines.append(f"    Parse_OK={d.parse_ok}; Sanitize_OK={d.sanitize_ok}; maps={d.mapped_dummy_counts or '{}'}; unmapped={d.unmapped_dummy_count}")
        if include_raw:
            lines.append(f"    Raw_SMILES={d.raw_smiles}")
        if d.normalized_smiles:
            lines.append(f"    Normalized_SMILES={d.normalized_smiles}")
        if d.canonical_smiles:
            lines.append(f"    Canonical_SMILES={d.canonical_smiles}")
        lines += ["    NORMALIZED: " + x for x in d.normalization_changes]
        lines += ["    ERROR: " + x for x in d.errors]
        lines += ["    WARNING: " + x for x in d.warnings]
        lines += ["    FIX: " + x for x in d.suggestions]
        lines += ["    DUMMY: " + x for x in d.dummy_details]
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def headers(include_raw: bool = True) -> List[str]:
    h = [
        "Role", "Source_Sheet", "Excel_Row", "ID", "Source", "Status",
        "Parse_OK", "Sanitize_OK", "Mapped_Dummy_Counts", "Unmapped_Dummy_Count",
        "Real_Atom_Map_Counts", "Atom_Count", "Heavy_Atom_Count", "Fragment_Count",
        "Normalized_SMILES", "Canonical_SMILES", "SMILES_Hash", "Normalization_Changes",
        "RDKit_Message", "Errors", "Warnings", "Suggested_Fixes", "Dummy_Details",
    ]
    if include_raw:
        h.insert(14, "Raw_SMILES")
    return h


def row_values(d: SmilesDiagnostic, include_raw: bool = True) -> List[object]:
    row: List[object] = [
        d.role, d.source_sheet, d.excel_row, d.item_id, d.source_label, d.status,
        "Yes" if d.parse_ok else "No", "Yes" if d.sanitize_ok else "No",
        "; ".join(f"{k}:{v}" for k, v in sorted(d.mapped_dummy_counts.items())), d.unmapped_dummy_count,
        "; ".join(f"{k}:{v}" for k, v in sorted(d.real_atom_map_counts.items())),
        d.atom_count, d.heavy_atom_count, d.fragment_count, d.normalized_smiles,
        d.canonical_smiles, d.smiles_hash, "; ".join(d.normalization_changes), d.rdkit_message,
        "; ".join(d.errors), "; ".join(d.warnings), "; ".join(d.suggestions), "; ".join(d.dummy_details),
    ]
    if include_raw:
        row.insert(14, d.raw_smiles)
    return row


def write_xlsx(report: DiagnosticsReport, path: Path, *, include_raw: bool = True) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "SMILES_Input_Diagnostics"
    hs = headers(include_raw)
    ws.append(hs)
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for d in report.diagnostics:
        ws.append(row_values(d, include_raw))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hs))}{max(1, ws.max_row)}"
    for i, h in enumerate(hs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if h in {"Role", "Status", "Parse_OK", "Sanitize_OK", "Excel_Row"} else (24 if h in {"Source_Sheet", "ID", "Source", "Mapped_Dummy_Counts"} else 46)
    for rr in ws.iter_rows(min_row=2):
        for c in rr:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    sm = wb.create_sheet("Summary")
    rows = [
        ["Item", "Value"], ["Overall_Status", report.overall_status],
        ["Inputs_Checked", len(report.diagnostics)], ["Nonfatal", report.ok_count],
        ["Error_Count", report.error_count], ["Warning_Count", report.warning_count], [],
        ["Required_CORE", "[*:1], [*:2], [*:3] exactly once each"],
        ["Required_A", "[*:1] exactly once"], ["Required_B", "[*:2] exactly once"],
        ["Required_C", "[*:3] exactly once"],
    ]
    for r in rows:
        sm.append(r)
    for x in report.global_errors:
        sm.append(["GLOBAL_ERROR", x])
    for x in report.global_warnings:
        sm.append(["GLOBAL_WARNING", x])
    for c in sm[1]:
        c.fill = fill
        c.font = Font(color="FFFFFF", bold=True)
    sm.column_dimensions["A"].width = 24
    sm.column_dimensions["B"].width = 105
    for rr in sm.iter_rows():
        for c in rr:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(str(path))
    report.output_xlsx = path
    return path


def write_txt(report: DiagnosticsReport, path: Path, *, include_raw: bool = False) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(diagnostics_text(report, include_raw=include_raw), encoding="utf-8-sig")
    report.output_txt = path
    return path
