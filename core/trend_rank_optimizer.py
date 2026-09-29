"""Trend-first response correction for small standard sets.

This module adds two deliberately narrow models:

1. PhysicalPairwiseRanker -- keeps the coefficient of log(area/IS) fixed at
   +1 and learns only a structure/mobile-phase response correction.  The
   corrected score is log(area/IS) - predicted log(response factor).
2. ResponseFactorExtraTrees -- predicts log(response factor), never
   concentration directly, and restores concentration with the same fixed
   physical equation.

Both branches use an out-of-fold positive-slope continuous calibration from
corrected score to log concentration.  Isotonic regression is retained only as
a plateau diagnostic, because stepwise isotonic mapping can collapse hundreds
of unknown compounds into only a few repeated concentration values.  The
continuous calibration preserves ordering while keeping the unknown sample
outside all supervision.

Feature subsets and model parameters are selected by Monte-Carlo search inside
outer cross-validation folds.  The unknown sample is never used for feature
selection, model fitting or calibration.
"""
from __future__ import annotations

import math
import statistics
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np

from .concentration_levels import add_decile_columns

from .esi_descriptor_meta import display_name
from .esi_model_benchmark import (
    _average_ranks,
    _feature_group,
    _key,
    _metrics,
    _num,
    _pairwise_concordance_pct,
    _safe_corr,
    _text,
    _top_fraction_overlap_pct,
    _trend_metrics,
    choose_features,
)

SKLEARN_AVAILABLE = False
SKLEARN_ERROR = ""
try:
    from sklearn.ensemble import ExtraTreesRegressor
    from sklearn.impute import SimpleImputer
    from sklearn.isotonic import IsotonicRegression
    from sklearn.linear_model import HuberRegressor
    from sklearn.model_selection import KFold, RepeatedKFold, GroupKFold, LeaveOneGroupOut
    from sklearn.preprocessing import StandardScaler
    from scipy.optimize import minimize
    SKLEARN_AVAILABLE = True
except Exception as exc:  # pragma: no cover - optional dependency
    SKLEARN_ERROR = str(exc)


@dataclass
class TrendOptimizerResult:
    enabled: bool
    best_model: str = ""
    decision: str = ""
    decision_summary: str = ""
    comparison: List[Dict[str, object]] = field(default_factory=list)
    cv_predictions: List[Dict[str, object]] = field(default_factory=list)
    target_predictions: List[Dict[str, object]] = field(default_factory=list)
    trial_rows: List[Dict[str, object]] = field(default_factory=list)
    feature_rows: List[Dict[str, object]] = field(default_factory=list)
    group_rows: List[Dict[str, object]] = field(default_factory=list)
    permutation_rows: List[Dict[str, object]] = field(default_factory=list)
    settings_rows: List[Dict[str, object]] = field(default_factory=list)
    final_config_rows: List[Dict[str, object]] = field(default_factory=list)
    plot_paths: List[Path] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    concentration_unit: str = ""


# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------

def _valid_ratio(record: Dict[str, object]) -> Optional[float]:
    value = _num(record.get("Measured_ratio"))
    return float(value) if value is not None and value > 0 else None


def _valid_concentration(record: Dict[str, object]) -> Optional[float]:
    value = _num(record.get("Actual_concentration"))
    return float(value) if value is not None and value > 0 else None


def _candidate_records(records: Sequence[Dict[str, object]], require_concentration: bool) -> List[Dict[str, object]]:
    out: List[Dict[str, object]] = []
    for row in records:
        ratio = _valid_ratio(row)
        conc = _valid_concentration(row)
        if ratio is None:
            continue
        if require_concentration and conc is None:
            continue
        if not _key(row):
            continue
        if not _text(row.get("Structure_status")):
            continue
        out.append(dict(row))
    return out


def _matrix(records: Sequence[Dict[str, object]], feature_names: Sequence[str]) -> np.ndarray:
    rows: List[List[float]] = []
    for record in records:
        current: List[float] = []
        for name in feature_names:
            if name == "Log_Measured_Ratio":
                ratio = _valid_ratio(record)
                current.append(math.log10(ratio) if ratio is not None else np.nan)
            else:
                value = _num(record.get(name))
                current.append(float(value) if value is not None else np.nan)
        rows.append(current)
    return np.asarray(rows, dtype=float)


def _prepare_numeric(X_train: np.ndarray, X_other: np.ndarray, *, scale: bool) -> Tuple[np.ndarray, np.ndarray, SimpleImputer, Optional[StandardScaler]]:
    imputer = SimpleImputer(strategy="median", keep_empty_features=True)
    train = imputer.fit_transform(X_train)
    other = imputer.transform(X_other)
    scaler: Optional[StandardScaler] = None
    if scale:
        scaler = StandardScaler()
        train = scaler.fit_transform(train)
        other = scaler.transform(other)
    return np.asarray(train, dtype=float), np.asarray(other, dtype=float), imputer, scaler


def _spearman(y: Sequence[float], score: Sequence[float]) -> float:
    yy = np.asarray(y, dtype=float)
    ss = np.asarray(score, dtype=float)
    if len(yy) < 3 or len(yy) != len(ss):
        return float("nan")
    return _safe_corr(_average_ranks(yy), _average_ranks(ss))


def _trend_trial_metrics(y_log: np.ndarray, pred_log: np.ndarray, n_features: int) -> Dict[str, float]:
    actual = 10.0 ** np.asarray(y_log, dtype=float)
    predicted = 10.0 ** np.asarray(pred_log, dtype=float)
    trend = _trend_metrics(actual, predicted)
    range_metrics = _metrics(actual, predicted)
    rho = float(_num(trend.get("Trend_Spearman_r")) or -1.0)
    pairwise = float(_num(trend.get("Trend_Pairwise_Concordance_pct")) or 0.0)
    top20 = float(_num(trend.get("Trend_Top20_Overlap_pct")) or 0.0)
    med_fold = float(_num(range_metrics.get("Median_Fold_Error")) or 999.0)
    # Random baselines are 0 for Spearman, 50% for pairwise and 20% for Top20.
    composite = (
        0.60 * rho
        + 0.25 * ((pairwise - 50.0) / 50.0)
        + 0.15 * ((top20 - 20.0) / 80.0)
        - 0.025 * min(2.0, max(0.0, math.log10(max(1.0, med_fold))))
        - 0.0015 * max(0, int(n_features) - 8)
    )
    return {
        "Spearman": rho,
        "Pairwise_pct": pairwise,
        "Top20_pct": top20,
        "Median_Fold": med_fold,
        "Composite": float(composite),
    }


def _fit_range_calibrator(base_pred: np.ndarray, actual_log: np.ndarray) -> Dict[str, float]:
    """Fit a robust continuous intercept/slope calibration on OOF predictions.

    Pairwise-ranking scores have an arbitrary numerical scale.  A positive
    Huber slope maps that scale back to log10 concentration without introducing
    the plateaus created by isotonic regression.  The slope is allowed to be
    small because rank scores may be much wider than the actual concentration
    range; only non-positive or non-finite slopes are rejected.
    """
    x = np.asarray(base_pred, dtype=float)
    y = np.asarray(actual_log, dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    x = x[mask]
    y = y[mask]
    if len(x) < 8 or float(np.std(x)) <= 1e-10:
        center = float(np.nanmedian(y)) if len(y) else 0.0
        return {
            "kind": "constant_fallback",
            "intercept": center,
            "slope": 0.0,
            "raw_slope": 0.0,
            "y_min": float(np.nanmin(y)) if len(y) else center,
            "y_max": float(np.nanmax(y)) if len(y) else center,
            "clip_margin": 0.25,
        }
    try:
        model = HuberRegressor(epsilon=1.35, alpha=1e-3, max_iter=500)
        model.fit(x.reshape(-1, 1), y)
        raw_slope = float(model.coef_[0])
        intercept = float(model.intercept_)
    except Exception:
        raw_slope, intercept = np.polyfit(x, y, 1)
        raw_slope = float(raw_slope)
        intercept = float(intercept)
    if not math.isfinite(raw_slope) or raw_slope <= 0:
        raw_slope = 0.05
    # Ranking-model scores may need strong compression; do not force a 0.5
    # minimum slope as earlier releases did.  Keep only a very small positive
    # lower bound and a conservative upper bound.
    slope = float(min(3.0, max(0.05, raw_slope)))
    intercept = float(np.median(y - slope * x))
    return {
        "kind": "robust_continuous_huber",
        "intercept": intercept,
        "slope": slope,
        "raw_slope": raw_slope,
        "y_min": float(np.min(y)),
        "y_max": float(np.max(y)),
        "clip_margin": 0.25,
    }


def _apply_calibrator(pred: np.ndarray, info: Dict[str, float]) -> np.ndarray:
    out = float(info.get("intercept", 0.0)) + float(info.get("slope", 1.0)) * np.asarray(pred, dtype=float)
    y_min = _num(info.get("y_min"))
    y_max = _num(info.get("y_max"))
    if y_min is not None and y_max is not None and y_max >= y_min:
        margin = float(_num(info.get("clip_margin")) or 0.25)
        out = np.clip(out, float(y_min) - margin, float(y_max) + margin)
    return np.asarray(out, dtype=float)


# ---------------------------------------------------------------------------
# Physically constrained response-correction models
# ---------------------------------------------------------------------------

class _PhysicalPairwiseModel:
    """Linear response-correction model fitted with a pairwise hinge objective.

    The measured log area ratio is *not* a free model input.  Its coefficient
    is fixed at +1 in the corrected concentration score::

        corrected_log_score = log10(area_ratio) - X @ coef

    The model therefore learns only the structure/mobile-phase correction that
    must be subtracted from the analytical signal.
    """

    def __init__(self, coef: np.ndarray):
        self.coef_ = np.asarray(coef, dtype=float).reshape(-1)

    def corrected_score(self, X: np.ndarray, log_ratio: np.ndarray) -> np.ndarray:
        return np.asarray(log_ratio, dtype=float) - np.asarray(X, dtype=float).dot(self.coef_)


def _pairwise_physical_data(
    X: np.ndarray,
    y_log_concentration: np.ndarray,
    log_ratio: np.ndarray,
    min_fold: float,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, int]:
    """Build pairwise constraints for the fixed-ratio response model.

    For a pair i,j the desired concentration order is enforced on
    ``delta(log ratio) - delta(predicted log RRF)``.  Pairs whose true
    concentrations are closer than ``min_fold`` are ignored.
    """
    threshold = math.log10(max(1.0001, float(min_fold)))
    z_rows: List[np.ndarray] = []
    offsets: List[float] = []
    weights: List[float] = []
    n = len(y_log_concentration)
    for i in range(n - 1):
        for j in range(i + 1, n):
            dy = float(y_log_concentration[i] - y_log_concentration[j])
            if abs(dy) < threshold:
                continue
            sign = 1.0 if dy > 0 else -1.0
            # Correct order margin = sign * (delta log ratio - delta X @ w).
            z_rows.append(sign * (X[i] - X[j]))
            offsets.append(sign * float(log_ratio[i] - log_ratio[j]))
            weights.append(min(4.0, max(1.0, abs(dy) / max(threshold, 1e-6))))
    if not z_rows:
        raise ValueError("No usable concentration pairs after the pair-difference threshold")
    return (
        np.asarray(z_rows, dtype=float),
        np.asarray(offsets, dtype=float),
        np.asarray(weights, dtype=float),
        len(z_rows),
    )


def _fit_physical_pairwise(
    X_train: np.ndarray,
    y_log_concentration: np.ndarray,
    log_ratio: np.ndarray,
    *,
    alpha: float,
    pair_min_fold: float,
    margin: float = 0.10,
) -> Tuple[_PhysicalPairwiseModel, int]:
    z, offsets, weights, pair_count = _pairwise_physical_data(
        X_train, y_log_concentration, log_ratio, pair_min_fold
    )
    alpha = max(1e-8, float(alpha))
    margin = max(0.0, float(margin))
    weight_sum = float(np.sum(weights)) or 1.0

    def objective(w: np.ndarray) -> Tuple[float, np.ndarray]:
        # desired: offset - z @ w >= margin
        residual = margin - offsets + z.dot(w)
        active = residual > 0
        if np.any(active):
            rr = residual[active]
            ww = weights[active]
            loss = float(np.sum(ww * rr * rr) / weight_sum)
            grad = 2.0 * np.sum((ww * rr)[:, None] * z[active], axis=0) / weight_sum
        else:
            loss = 0.0
            grad = np.zeros(z.shape[1], dtype=float)
        loss += alpha * float(np.dot(w, w))
        grad = grad + 2.0 * alpha * w
        return loss, np.asarray(grad, dtype=float)

    result = minimize(
        lambda w: objective(w)[0],
        np.zeros(X_train.shape[1], dtype=float),
        jac=lambda w: objective(w)[1],
        method="L-BFGS-B",
        options={"maxiter": 1500, "ftol": 1e-10},
    )
    if not result.success and not np.all(np.isfinite(result.x)):
        raise RuntimeError("Physical pairwise optimisation failed: " + str(result.message))
    return _PhysicalPairwiseModel(np.asarray(result.x, dtype=float)), int(pair_count)


def _splitter_for_inner(
    X: np.ndarray,
    y: np.ndarray,
    *,
    groups: Optional[np.ndarray],
    n_splits: int,
    random_state: int,
):
    unique_groups = sorted({str(x) for x in groups if str(x).strip()}) if groups is not None else []
    if groups is not None and len(unique_groups) >= 2 and all(str(x).strip() for x in groups):
        splitter = GroupKFold(n_splits=max(2, min(int(n_splits), len(unique_groups))))
        return splitter.split(X, y, groups)
    splitter = KFold(
        n_splits=max(2, min(int(n_splits), len(y) // 6 if len(y) >= 12 else 2)),
        shuffle=True,
        random_state=int(random_state),
    )
    return splitter.split(X)


def _crossfit_physical_scores(
    X: np.ndarray,
    y_log: np.ndarray,
    log_ratio: np.ndarray,
    *,
    subset: Sequence[int],
    alpha: float,
    pair_min_fold: float,
    margin: float,
    random_state: int,
    n_splits: int = 3,
    groups: Optional[np.ndarray] = None,
) -> Tuple[np.ndarray, int]:
    out = np.full(len(y_log), np.nan, dtype=float)
    pair_count = 0
    split_iter = _splitter_for_inner(
        X, y_log, groups=groups, n_splits=n_splits, random_state=random_state
    )
    for fold_no, (tr, va) in enumerate(split_iter, start=1):
        Xtr, Xva, _, _ = _prepare_numeric(X[tr][:, subset], X[va][:, subset], scale=True)
        model, pairs = _fit_physical_pairwise(
            Xtr, y_log[tr], log_ratio[tr], alpha=alpha,
            pair_min_fold=pair_min_fold, margin=margin,
        )
        out[va] = model.corrected_score(Xva, log_ratio[va])
        pair_count += pairs
    return out, int(pair_count)


# ---------------------------------------------------------------------------
# Response-factor ExtraTrees
# ---------------------------------------------------------------------------

def _make_extra_trees(params: Dict[str, object], random_state: int, *, final: bool = False):
    return ExtraTreesRegressor(
        n_estimators=500 if final else 160,
        max_depth=params.get("max_depth"),
        min_samples_leaf=int(params.get("min_samples_leaf", 3)),
        max_features=params.get("max_features", 0.7),
        bootstrap=False,
        random_state=int(random_state),
        n_jobs=-1,
    )


def _crossfit_rrf_extra_scores(
    X: np.ndarray,
    y_log: np.ndarray,
    log_ratio: np.ndarray,
    *,
    subset: Sequence[int],
    params: Dict[str, object],
    random_state: int,
    n_splits: int = 3,
    groups: Optional[np.ndarray] = None,
) -> np.ndarray:
    """Cross-fitted corrected log concentration from predicted log RRF."""
    out = np.full(len(y_log), np.nan, dtype=float)
    y_log_rrf = np.asarray(log_ratio, dtype=float) - np.asarray(y_log, dtype=float)
    split_iter = _splitter_for_inner(
        X, y_log_rrf, groups=groups, n_splits=n_splits, random_state=random_state
    )
    for fold_no, (tr, va) in enumerate(split_iter, start=1):
        Xtr, Xva, _, _ = _prepare_numeric(X[tr][:, subset], X[va][:, subset], scale=False)
        model = _make_extra_trees(params, random_state + fold_no, final=False)
        model.fit(Xtr, y_log_rrf[tr])
        pred_log_rrf = np.asarray(model.predict(Xva), dtype=float)
        out[va] = log_ratio[va] - pred_log_rrf
    return out


# ---------------------------------------------------------------------------
# Monotonic score-to-concentration calibration
# ---------------------------------------------------------------------------

def _fit_monotonic_calibrator(base_pred: np.ndarray, actual_log: np.ndarray) -> Dict[str, object]:
    """Fit a continuous positive-slope calibrator and audit isotonic plateaus.

    Earlier releases used isotonic regression as the deployed concentration
    map.  With fewer than one hundred standards, isotonic regression can form
    only a few blocks, so hundreds of target compounds may receive exactly the
    same concentration.  We now deploy a robust continuous Huber map and keep
    isotonic regression only to report how severe that plateau risk would have
    been.
    """
    x = np.asarray(base_pred, dtype=float)
    y = np.asarray(actual_log, dtype=float)
    mask = np.isfinite(x) & np.isfinite(y)
    x = x[mask]
    y = y[mask]
    info: Dict[str, object] = dict(_fit_range_calibrator(x, y))
    info["calibration_sample_count"] = int(len(x))
    info["isotonic_level_count"] = 0
    info["isotonic_unique_fraction_pct"] = 0.0
    info["isotonic_plateau_detected"] = False
    if len(x) >= 8 and float(np.std(x)) > 1e-10:
        try:
            iso = IsotonicRegression(
                increasing=True,
                out_of_bounds="clip",
                y_min=float(np.min(y)),
                y_max=float(np.max(y)),
            )
            iso.fit(x, y)
            fitted = np.asarray(iso.predict(x), dtype=float)
            level_count = int(len(np.unique(np.round(fitted, 10))))
            expected_min = int(max(8, math.ceil(math.sqrt(len(x)))))
            info["isotonic_level_count"] = level_count
            info["isotonic_unique_fraction_pct"] = float(level_count / max(1, len(x)) * 100.0)
            info["isotonic_plateau_detected"] = bool(level_count < expected_min)
            info["isotonic_expected_min_levels"] = expected_min
        except Exception as exc:
            info["isotonic_diagnostic_error"] = str(exc)
    return info


def _apply_monotonic_calibrator(pred: np.ndarray, info: Dict[str, object]) -> np.ndarray:
    # Deployed target concentrations are always continuous.  The old isotonic
    # branch is intentionally not used here; isotonic statistics are retained
    # in ``info`` only for diagnostics.
    return _apply_calibrator(np.asarray(pred, dtype=float), info)  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# Monte-Carlo feature/parameter search
# ---------------------------------------------------------------------------

def _random_subset(feature_names: Sequence[str], rng: np.random.Generator, min_features: int, max_features: int) -> List[int]:
    n = len(feature_names)
    upper = max(1, min(int(max_features), n))
    lower = max(1, min(int(min_features), upper))
    k = int(rng.integers(lower, upper + 1))
    groups: Dict[str, List[int]] = {}
    for idx, name in enumerate(feature_names):
        groups.setdefault(_feature_group(name), []).append(idx)
    group_names = list(groups)
    rng.shuffle(group_names)
    n_groups = int(rng.integers(1, max(2, min(len(group_names), 6)) + 1)) if group_names else 0
    active = group_names[:n_groups]
    pool = [idx for group in active for idx in groups[group]]
    if len(pool) < k:
        pool = list(range(n))
    selected = rng.choice(np.asarray(pool, dtype=int), size=min(k, len(pool)), replace=False).tolist() if pool else []
    if len(selected) < k:
        left = [i for i in range(n) if i not in selected]
        if left:
            selected.extend(rng.choice(np.asarray(left, dtype=int), size=min(k - len(selected), len(left)), replace=False).tolist())
    return sorted(set(int(x) for x in selected))


def _trial_config(model_name: str, feature_names: Sequence[str], rng: np.random.Generator, min_features: int, max_features: int, pair_min_fold: float) -> Dict[str, object]:
    subset = _random_subset(feature_names, rng, min_features, max_features)
    if model_name == "PhysicalPairwiseRanker":
        threshold_choices = sorted(set([max(1.05, pair_min_fold), 1.25, 1.5, 2.0, 3.0]))
        return {
            "subset": subset,
            "alpha": float(10.0 ** rng.uniform(-4.0, 1.0)),
            "pair_min_fold": float(rng.choice(threshold_choices)),
            "margin": float(rng.choice([0.0, 0.05, 0.10, 0.20])),
        }
    return {
        "subset": subset,
        "max_depth": rng.choice([3, 5, 8, None]),
        "min_samples_leaf": int(rng.choice([2, 3, 5, 8])),
        "max_features": float(rng.choice([0.35, 0.5, 0.7, 1.0])),
    }


def _evaluate_trial(
    model_name: str,
    X: np.ndarray,
    y_log: np.ndarray,
    log_ratio: np.ndarray,
    config: Dict[str, object],
    random_state: int,
    groups: Optional[np.ndarray] = None,
) -> Tuple[Dict[str, float], np.ndarray, Dict[str, object], int]:
    subset = list(config["subset"])
    if model_name == "PhysicalPairwiseRanker":
        base, pair_count = _crossfit_physical_scores(
            X, y_log, log_ratio, subset=subset, alpha=float(config["alpha"]),
            pair_min_fold=float(config["pair_min_fold"]), margin=float(config.get("margin", 0.10)),
            random_state=random_state, groups=groups,
        )
    else:
        base = _crossfit_rrf_extra_scores(
            X, y_log, log_ratio, subset=subset, params=config,
            random_state=random_state, groups=groups,
        )
        pair_count = 0
    mask = np.isfinite(base) & np.isfinite(y_log)
    if int(np.sum(mask)) < max(10, len(y_log) // 2):
        raise ValueError("Insufficient inner-CV predictions")
    calibrator = _fit_monotonic_calibrator(base[mask], y_log[mask])
    calibrated = np.full(len(y_log), np.nan, dtype=float)
    calibrated[mask] = _apply_monotonic_calibrator(base[mask], calibrator)
    metrics = _trend_trial_metrics(y_log[mask], calibrated[mask], len(subset))
    return metrics, calibrated, calibrator, pair_count


def _search_model(
    model_name: str,
    X: np.ndarray,
    y_log: np.ndarray,
    log_ratio: np.ndarray,
    feature_names: Sequence[str],
    *,
    trials: int,
    min_features: int,
    max_features: int,
    pair_min_fold: float,
    random_state: int,
    stage: str,
    outer_repeat: int = 0,
    outer_fold: int = 0,
    groups: Optional[np.ndarray] = None,
) -> Tuple[Dict[str, object], List[Dict[str, object]]]:
    rng = np.random.default_rng(int(random_state))
    rows: List[Dict[str, object]] = []
    best: Optional[Dict[str, object]] = None
    for trial_no in range(1, max(1, int(trials)) + 1):
        config = _trial_config(model_name, feature_names, rng, min_features, max_features, pair_min_fold)
        try:
            metrics, _, calibrator, pair_count = _evaluate_trial(
                model_name, X, y_log, log_ratio, config, random_state + trial_no * 17, groups=groups
            )
            status = "OK"; error = ""
        except Exception as exc:
            metrics = {"Spearman": -1.0, "Pairwise_pct": 0.0, "Top20_pct": 0.0, "Median_Fold": 999.0, "Composite": -999.0}
            calibrator = {"kind": "failed", "intercept": 0.0, "slope": 1.0, "raw_slope": 1.0}
            pair_count = 0; status = "FAILED"; error = str(exc)
        subset = list(config["subset"])
        feature_list = [feature_names[i] for i in subset]
        row = {
            "Stage": stage, "Outer_Repeat": outer_repeat, "Outer_Fold": outer_fold,
            "Model": model_name, "Trial": trial_no, "Status": status,
            "Composite_score": metrics["Composite"], "Inner_Spearman": metrics["Spearman"],
            "Inner_Pairwise_pct": metrics["Pairwise_pct"], "Inner_Top20_pct": metrics["Top20_pct"],
            "Inner_Median_Fold": metrics["Median_Fold"], "Feature_count": len(feature_list),
            "Features": ";".join(feature_list), "Calibration_method": calibrator.get("kind", ""),
            "Calibration_method": calibrator.get("kind", ""),
                "Calibration_slope": calibrator.get("slope", 1.0),
            "Calibration_raw_slope": calibrator.get("raw_slope", 1.0),
            "Pair_count": pair_count, "Error": error,
        }
        for key, value in config.items():
            if key != "subset":
                row["Param_" + str(key)] = value
        rows.append(row)
        candidate = dict(config)
        candidate.update({"metrics": metrics, "calibrator": calibrator, "features": feature_list, "trial_no": trial_no})
        if best is None or float(metrics["Composite"]) > float(best["metrics"]["Composite"]):
            best = candidate
    if best is None:
        raise RuntimeError(f"No usable {model_name} trial")
    return best, rows


def _fit_predict_config(
    model_name: str,
    X_train: np.ndarray,
    y_train: np.ndarray,
    log_ratio_train: np.ndarray,
    X_test: np.ndarray,
    log_ratio_test: np.ndarray,
    config: Dict[str, object],
    *,
    random_state: int,
    groups: Optional[np.ndarray] = None,
) -> Tuple[np.ndarray, np.ndarray, Dict[str, object], int]:
    subset = list(config["subset"])
    if model_name == "PhysicalPairwiseRanker":
        inner_base, pair_count = _crossfit_physical_scores(
            X_train, y_train, log_ratio_train, subset=subset,
            alpha=float(config["alpha"]), pair_min_fold=float(config["pair_min_fold"]),
            margin=float(config.get("margin", 0.10)), random_state=random_state + 100,
            groups=groups,
        )
        calibrator = _fit_monotonic_calibrator(inner_base, y_train)
        Xtr, Xte, _, _ = _prepare_numeric(X_train[:, subset], X_test[:, subset], scale=True)
        model, fitted_pairs = _fit_physical_pairwise(
            Xtr, y_train, log_ratio_train, alpha=float(config["alpha"]),
            pair_min_fold=float(config["pair_min_fold"]), margin=float(config.get("margin", 0.10)),
        )
        base = model.corrected_score(Xte, log_ratio_test)
        pair_count = max(pair_count, fitted_pairs)
    else:
        inner_base = _crossfit_rrf_extra_scores(
            X_train, y_train, log_ratio_train, subset=subset, params=config,
            random_state=random_state + 100, groups=groups,
        )
        calibrator = _fit_monotonic_calibrator(inner_base, y_train)
        Xtr, Xte, _, _ = _prepare_numeric(X_train[:, subset], X_test[:, subset], scale=False)
        model = _make_extra_trees(config, random_state, final=True)
        y_log_rrf = log_ratio_train - y_train
        model.fit(Xtr, y_log_rrf)
        pred_log_rrf = np.asarray(model.predict(Xte), dtype=float)
        base = log_ratio_test - pred_log_rrf
        pair_count = 0
    calibrated = _apply_monotonic_calibrator(base, calibrator)
    return np.asarray(base, dtype=float), np.asarray(calibrated, dtype=float), calibrator, int(pair_count)


# ---------------------------------------------------------------------------
# Outer CV, final fit and diagnostics
# ---------------------------------------------------------------------------

def _aggregate_cv(rows: Sequence[Dict[str, object]]) -> List[Dict[str, object]]:
    grouped: Dict[Tuple[str, int], List[Dict[str, object]]] = {}
    for row in rows:
        grouped.setdefault((str(row["Model"]), int(row["Calibration_Index"])), []).append(row)
    out: List[Dict[str, object]] = []
    for (model, idx), group in sorted(grouped.items()):
        pred_log = float(statistics.median(float(r["Predicted_log10_concentration"]) for r in group))
        base_log = float(statistics.median(float(r["Base_log10_prediction"]) for r in group))
        actual = float(group[0]["Actual_concentration"])
        ratio = float(group[0]["Measured_ratio"])
        pred = 10.0 ** pred_log
        out.append({
            **{k: v for k, v in group[0].items() if k not in {"Repeat", "Fold"}},
            "Repeat": "median_over_repeats",
            "Fold": "",
            "Base_log10_prediction": base_log,
            "Predicted_log10_concentration": pred_log,
            "Predicted_concentration": pred,
            "Predicted_log10_RRF": math.log10(ratio) - pred_log,
            "Fold_error": max(pred / actual, actual / pred),
            "Prediction_count": len(group),
        })
    return out


def _comparison(aggregated: Sequence[Dict[str, object]], calibration: Sequence[Dict[str, object]]) -> List[Dict[str, object]]:
    raw_actual = [float(r["Actual_concentration"]) for r in calibration]
    raw_ratio = [float(r["Measured_ratio"]) for r in calibration]
    raw_trend = _trend_metrics(raw_actual, raw_ratio)
    rows: List[Dict[str, object]] = []
    for model in ("PhysicalPairwiseRanker", "ResponseFactorExtraTrees"):
        subset = [r for r in aggregated if str(r.get("Model")) == model]
        actual = [float(r["Actual_concentration"]) for r in subset]
        pred = [float(r["Predicted_concentration"]) for r in subset]
        metrics = _metrics(actual, pred)
        trend = _trend_metrics(actual, pred)
        row = {"Model": model, **metrics, **trend}
        row["Raw_Spearman"] = raw_trend.get("Trend_Spearman_r", "")
        row["Delta_Spearman_vs_raw"] = (
            float(trend.get("Trend_Spearman_r", 0.0)) - float(raw_trend.get("Trend_Spearman_r", 0.0))
            if _num(trend.get("Trend_Spearman_r")) is not None and _num(raw_trend.get("Trend_Spearman_r")) is not None else ""
        )
        rows.append(row)
    rows.sort(key=lambda r: (
        -float(_num(r.get("Trend_Spearman_r")) or -999.0),
        -float(_num(r.get("Trend_Pairwise_Concordance_pct")) or -999.0),
        float(_num(r.get("Median_Fold_Error")) or 999.0),
    ))
    for rank, row in enumerate(rows, start=1):
        row["Rank"] = rank
    return rows


def _bootstrap_spearman(actual: np.ndarray, predicted: np.ndarray, random_state: int, n_boot: int = 1000) -> Tuple[float, float]:
    rng = np.random.default_rng(int(random_state))
    vals: List[float] = []
    n = len(actual)
    for _ in range(max(100, int(n_boot))):
        idx = rng.integers(0, n, size=n)
        rho = _spearman(actual[idx], predicted[idx])
        if math.isfinite(rho):
            vals.append(rho)
    if not vals:
        return float("nan"), float("nan")
    return float(np.quantile(vals, 0.025)), float(np.quantile(vals, 0.975))


def _permutation_test(actual: np.ndarray, predicted: np.ndarray, random_state: int, n_perm: int) -> Tuple[float, List[float]]:
    rng = np.random.default_rng(int(random_state))
    observed = _spearman(actual, predicted)
    null: List[float] = []
    for _ in range(max(20, int(n_perm))):
        null.append(_spearman(rng.permutation(actual), predicted))
    p = (1.0 + sum(1 for x in null if abs(x) >= abs(observed))) / (len(null) + 1.0)
    return float(p), null


def _feature_stability(trials: Sequence[Dict[str, object]], selected_configs: Sequence[Dict[str, object]], feature_names: Sequence[str], language: str) -> Tuple[List[Dict[str, object]], List[Dict[str, object]]]:
    valid = [r for r in trials if str(r.get("Status")) == "OK" and _num(r.get("Composite_score")) is not None]
    if not valid:
        return [], []
    scores = np.asarray([float(r["Composite_score"]) for r in valid], dtype=float)
    top_cut = float(np.quantile(scores, 0.90))
    rows: List[Dict[str, object]] = []
    groups: Dict[str, Dict[str, List[float]]] = {}
    selected_feature_lists = [set(str(c.get("features", "")).split(";")) for c in selected_configs]
    for name in feature_names:
        included = [r for r in valid if name in set(str(r.get("Features", "")).split(";"))]
        excluded = [r for r in valid if name not in set(str(r.get("Features", "")).split(";"))]
        top_included = [r for r in included if float(r["Composite_score"]) >= top_cut]
        outer_frequency = sum(1 for s in selected_feature_lists if name in s) / max(1, len(selected_feature_lists)) * 100.0
        mean_in = float(np.mean([float(r["Composite_score"]) for r in included])) if included else float("nan")
        mean_out = float(np.mean([float(r["Composite_score"]) for r in excluded])) if excluded else float("nan")
        delta = mean_in - mean_out if math.isfinite(mean_in) and math.isfinite(mean_out) else float("nan")
        group = _feature_group(name)
        row = {
            "Feature": name,
            "Display_name": display_name(name, language),
            "Group": group,
            "Trial_inclusion_pct": len(included) / len(valid) * 100.0,
            "Top10_trial_inclusion_pct": len(top_included) / max(1, sum(1 for r in valid if float(r["Composite_score"]) >= top_cut)) * 100.0,
            "Outer_selected_pct": outer_frequency,
            "Mean_score_when_included": mean_in if math.isfinite(mean_in) else "",
            "Mean_score_when_excluded": mean_out if math.isfinite(mean_out) else "",
            "Inclusion_score_delta": delta if math.isfinite(delta) else "",
            "Interpretation": (
                "stable positive contribution" if outer_frequency >= 50 and math.isfinite(delta) and delta > 0
                else "unstable or model-dependent"
            ),
        }
        rows.append(row)
        bucket = groups.setdefault(group, {"in": [], "out": [], "outer": []})
        if math.isfinite(delta):
            bucket["in"].append(delta)
        bucket["outer"].append(outer_frequency)
    rows.sort(key=lambda r: (-float(_num(r.get("Outer_selected_pct")) or 0.0), -float(_num(r.get("Inclusion_score_delta")) or -999.0)))
    group_rows: List[Dict[str, object]] = []
    for group, data in groups.items():
        group_rows.append({
            "Group": group,
            "Mean_outer_selected_pct": float(np.mean(data["outer"])) if data["outer"] else "",
            "Mean_feature_inclusion_score_delta": float(np.mean(data["in"])) if data["in"] else "",
            "Feature_count": len(data["outer"]),
        })
    group_rows.sort(key=lambda r: -float(_num(r.get("Mean_feature_inclusion_score_delta")) or -999.0))
    return rows, group_rows


def _make_plots(output_xlsx: Path, comparison: Sequence[Dict[str, object]], aggregated: Sequence[Dict[str, object]], best_model: str, feature_rows: Sequence[Dict[str, object]], trial_rows: Sequence[Dict[str, object]], permutation_null: Sequence[float], actual_log: np.ndarray, pred_log: np.ndarray, language: str) -> Tuple[List[Path], List[str]]:
    warnings: List[str] = []
    paths: List[Path] = []
    try:
        import matplotlib.pyplot as plt
    except Exception as exc:
        return [], [f"Trend optimizer plots unavailable: {exc}"]
    plot_dir = Path(output_xlsx).parent / ("_trendplots_" + str(abs(hash(str(output_xlsx))) % 10**8).zfill(8))
    plot_dir.mkdir(parents=True, exist_ok=True)
    zh = str(language).lower().startswith("zh")

    try:
        names = [str(r["Model"]) for r in comparison]
        spearman = [float(_num(r.get("Trend_Spearman_r")) or 0.0) for r in comparison]
        pairwise = [float(_num(r.get("Trend_Pairwise_Concordance_pct")) or 0.0) / 100.0 for r in comparison]
        fig, ax = plt.subplots(figsize=(9, 5.5))
        x = np.arange(len(names))
        ax.bar(x - 0.18, spearman, width=0.36, label="Spearman")
        ax.bar(x + 0.18, pairwise, width=0.36, label="Pairwise")
        ax.axhline(0, linewidth=1)
        ax.set_xticks(x, names, rotation=20, ha="right")
        ax.set_ylim(-0.2, 1.0)
        ax.set_title('Trend-first model comparison' if zh else "Trend-first model comparison")
        ax.legend()
        fig.tight_layout()
        p = plot_dir / "19_trend_models.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Trend model plot failed: {exc}")

    try:
        fig, ax = plt.subplots(figsize=(6.8, 6.2))
        ax.scatter(_average_ranks(actual_log), _average_ranks(pred_log), alpha=0.75)
        lim = [1, len(actual_log)]
        ax.plot(lim, lim, linestyle="--")
        ax.set_xlabel('Actual concentration rank' if zh else "Actual concentration rank")
        ax.set_ylabel('Out-of-fold predicted rank' if zh else "Out-of-fold predicted rank")
        ax.set_title((best_model + ': rank recovery') if zh else (best_model + ": rank recovery"))
        fig.tight_layout()
        p = plot_dir / "20_rank_scatter.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Rank scatter failed: {exc}")

    try:
        top = list(feature_rows[:20])[::-1]
        fig, ax = plt.subplots(figsize=(9, 7))
        ax.barh([str(r["Display_name"]) for r in top], [float(_num(r.get("Outer_selected_pct")) or 0.0) for r in top])
        ax.set_xlabel('Outer-CV selection frequency (%)' if zh else "Outer-CV selection frequency (%)")
        ax.set_title('Stable features from random-subset search' if zh else "Stable features from random-subset search")
        fig.tight_layout()
        p = plot_dir / "21_feature_stability.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Feature stability plot failed: {exc}")

    try:
        fig, ax = plt.subplots(figsize=(8, 5))
        for model in ("PhysicalPairwiseRanker", "ResponseFactorExtraTrees"):
            vals = [float(r["Composite_score"]) for r in trial_rows if str(r.get("Model")) == model and _num(r.get("Composite_score")) is not None and float(r["Composite_score"]) > -100]
            if vals:
                ax.hist(vals, bins=20, alpha=0.5, label=model)
        ax.set_xlabel('Inner-CV trend composite score' if zh else "Inner-CV trend composite score")
        ax.set_ylabel('Trials' if zh else "Trials")
        ax.set_title('Random parameter/descriptor subset search' if zh else "Random parameter/descriptor subset search")
        ax.legend()
        fig.tight_layout()
        p = plot_dir / "22_trial_scores.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Trial distribution plot failed: {exc}")

    try:
        observed = _spearman(actual_log, pred_log)
        fig, ax = plt.subplots(figsize=(8, 5))
        ax.hist(permutation_null, bins=30, alpha=0.75)
        ax.axvline(observed, linestyle="--", label=f"observed={observed:.3f}")
        ax.set_xlabel('Permuted Spearman' if zh else "Permuted Spearman")
        ax.set_title('Permutation test of OOF trend association' if zh else "Permutation test of OOF trend association")
        ax.legend()
        fig.tight_layout()
        p = plot_dir / "23_permutation.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Permutation plot failed: {exc}")

    try:
        slope, intercept = np.polyfit(pred_log, actual_log, 1) if len(pred_log) >= 2 and np.std(pred_log) > 0 else (float("nan"), float("nan"))
        fig, ax = plt.subplots(figsize=(7, 6))
        ax.scatter(pred_log, actual_log, alpha=0.75)
        lo = min(float(np.min(pred_log)), float(np.min(actual_log)))
        hi = max(float(np.max(pred_log)), float(np.max(actual_log)))
        ax.plot([lo, hi], [lo, hi], linestyle="--", label="1x")
        if math.isfinite(slope):
            xx = np.linspace(lo, hi, 100)
            ax.plot(xx, intercept + slope * xx, label=f"calibration slope={slope:.2f}")
        ax.set_xlabel('OOF predicted log10 concentration' if zh else "OOF predicted log10 concentration")
        ax.set_ylabel('Actual log10 concentration' if zh else "Actual log10 concentration")
        ax.set_title('Prediction-range compression diagnostic' if zh else "Prediction-range compression diagnostic")
        ax.legend()
        fig.tight_layout()
        p = plot_dir / "24_compression.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Compression plot failed: {exc}")

    # Requested direct value comparison: actual/theoretical concentration vs OOF prediction.
    try:
        best_rows = [r for r in aggregated if str(r.get("Model")) == best_model]
        actual = np.asarray([float(r["Actual_concentration"]) for r in best_rows], dtype=float)
        predicted = np.asarray([float(r["Predicted_concentration"]) for r in best_rows], dtype=float)
        mask = np.isfinite(actual) & np.isfinite(predicted) & (actual > 0) & (predicted > 0)
        actual = actual[mask]; predicted = predicted[mask]
        fig, ax = plt.subplots(figsize=(7, 6.4))
        ax.scatter(actual, predicted, alpha=0.78)
        lo = min(float(np.min(actual)), float(np.min(predicted)))
        hi = max(float(np.max(actual)), float(np.max(predicted)))
        xx = np.logspace(math.log10(lo), math.log10(hi), 200)
        ax.plot(xx, xx, linestyle="-", label="1x")
        ax.plot(xx, xx * 2.0, linestyle=":", label="2x")
        ax.plot(xx, xx / 2.0, linestyle=":")
        ax.plot(xx, xx * 5.0, linestyle="--", label="5x")
        ax.plot(xx, xx / 5.0, linestyle="--")
        ax.set_xscale("log"); ax.set_yscale("log")
        ax.set_xlabel('Actual / theoretical concentration' if zh else "Actual / theoretical concentration")
        ax.set_ylabel('Out-of-fold predicted concentration' if zh else "Out-of-fold predicted concentration")
        ax.set_title((best_model + ': known vs predicted concentration') if zh else (best_model + ": actual vs predicted concentration"))
        ax.legend()
        fig.tight_layout()
        p = plot_dir / "25_actual_vs_predicted.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Actual-vs-predicted plot failed: {exc}")

    # Requested direct comparison on an area-ratio-like scale after response normalization.
    try:
        best_rows = [r for r in aggregated if str(r.get("Model")) == best_model]
        raw = np.asarray([float(r.get("Raw_area_ratio", r.get("Measured_ratio"))) for r in best_rows], dtype=float)
        corrected = np.asarray([float(r.get("Response_normalized_area_ratio")) for r in best_rows], dtype=float)
        mask = np.isfinite(raw) & np.isfinite(corrected) & (raw > 0) & (corrected > 0)
        raw = raw[mask]; corrected = corrected[mask]
        fig, ax = plt.subplots(figsize=(7, 6.4))
        ax.scatter(raw, corrected, alpha=0.78)
        lo = min(float(np.min(raw)), float(np.min(corrected)))
        hi = max(float(np.max(raw)), float(np.max(corrected)))
        xx = np.logspace(math.log10(lo), math.log10(hi), 200)
        ax.plot(xx, xx, linestyle="--", label="unchanged")
        ax.set_xscale("log"); ax.set_yscale("log")
        ax.set_xlabel('Raw peak-area / internal-standard ratio' if zh else "Raw peak-area / internal-standard ratio")
        ax.set_ylabel('Response-normalized area ratio' if zh else "Response-normalized area ratio")
        ax.set_title('Raw vs response-normalized area ratio' if zh else "Raw vs response-normalized area ratio")
        ax.legend()
        fig.tight_layout()
        p = plot_dir / "26_raw_vs_corrected_ratio.png"; fig.savefig(p, dpi=180); plt.close(fig); paths.append(p)
    except Exception as exc:
        warnings.append(f"Raw-vs-corrected ratio plot failed: {exc}")
    return paths, warnings


def run_trend_optimizer(
    calibration_records: Sequence[Dict[str, object]],
    target_records: Sequence[Dict[str, object]],
    descriptor_names: Sequence[str],
    output_xlsx: Path,
    *,
    trials: int = 30,
    pair_min_fold: float = 1.5,
    min_features: int = 6,
    max_features: int = 24,
    cv_splits: int = 5,
    cv_repeats: int = 3,
    permutations: int = 500,
    random_state: int = 42,
    correlation_threshold: float = 0.97,
    output_language: str = "en",
    concentration_unit: str = "",
) -> TrendOptimizerResult:
    if not SKLEARN_AVAILABLE:
        raise RuntimeError("Trend optimizer needs scikit-learn: " + SKLEARN_ERROR)
    calibration = _candidate_records(calibration_records, require_concentration=True)
    targets = _candidate_records(target_records, require_concentration=False)
    if len(calibration) < 30:
        raise ValueError(f"Trend optimizer requires at least 30 calibration rows; got {len(calibration)}")

    numeric, _categorical, _manifest = choose_features(
        calibration, descriptor_names, expected_descriptor_names=(), three_d_enabled=True,
        correlation_threshold=float(correlation_threshold), include_categorical=False,
    )
    # The analytical signal is handled outside the machine-learning feature
    # matrix with a fixed coefficient of +1.  Models see only molecular and
    # mobile-phase descriptors and learn the response-factor correction.
    blocked = {"Measured_ratio", "Log_Measured_Ratio", "Actual_concentration"}
    feature_names = [x for x in numeric if x not in blocked and not x.lower().startswith("measured_ratio")]
    feature_names = list(dict.fromkeys(feature_names))
    if len(feature_names) < 3:
        raise ValueError("Too few numeric structure/gradient features for the physical trend optimizer")
    max_features = max(3, min(int(max_features), len(feature_names)))
    min_features = max(2, min(int(min_features), max_features))

    X = _matrix(calibration, feature_names)
    y_log = np.asarray([math.log10(float(r["Actual_concentration"])) for r in calibration], dtype=float)
    log_ratio = np.asarray([math.log10(float(r["Measured_ratio"])) for r in calibration], dtype=float)
    group_values = np.asarray([_text(r.get("Injection_Group")) for r in calibration], dtype=object)
    nonblank_groups = sorted({str(x) for x in group_values if str(x).strip()})
    use_group_cv = len(nonblank_groups) >= 3 and all(str(x).strip() for x in group_values)
    if use_group_cv:
        splitter = LeaveOneGroupOut()
        split_iter = list(splitter.split(X, y_log, group_values))
        outer_cv_label = f"Leave-one-injection-group-out ({len(nonblank_groups)} groups)"
        effective_splits = len(split_iter)
    else:
        splitter = RepeatedKFold(
            n_splits=max(2, min(int(cv_splits), len(calibration) // 8 if len(calibration) >= 16 else 2)),
            n_repeats=max(1, int(cv_repeats)), random_state=int(random_state),
        )
        split_iter = list(splitter.split(X))
        outer_cv_label = f"Repeated {cv_splits}-fold x {cv_repeats}"
        effective_splits = max(2, min(int(cv_splits), len(calibration) // 8 if len(calibration) >= 16 else 2))

    raw_rows: List[Dict[str, object]] = []
    trial_rows: List[Dict[str, object]] = []
    selected_configs: List[Dict[str, object]] = []
    split_no = 0
    for tr, te in split_iter:
        split_no += 1
        repeat_no = 1 if use_group_cv else (split_no - 1) // effective_splits + 1
        fold_no = split_no if use_group_cv else (split_no - 1) % effective_splits + 1
        for model_no, model_name in enumerate(("PhysicalPairwiseRanker", "ResponseFactorExtraTrees"), start=1):
            seed = int(random_state + split_no * 1009 + model_no * 100000)
            best, trials_for_fold = _search_model(
                model_name, X[tr], y_log[tr], log_ratio[tr], feature_names,
                trials=max(3, int(trials)), min_features=min_features, max_features=max_features,
                pair_min_fold=float(pair_min_fold), random_state=seed,
                stage="outer_inner_search", outer_repeat=repeat_no, outer_fold=fold_no,
                groups=(group_values[tr] if use_group_cv else None),
            )
            trial_rows.extend(trials_for_fold)
            base, pred_log, calibrator, pair_count = _fit_predict_config(
                model_name, X[tr], y_log[tr], log_ratio[tr], X[te], log_ratio[te], best, random_state=seed + 77,
                groups=(group_values[tr] if use_group_cv else None),
            )
            selected_configs.append({
                "Model": model_name,
                "Outer_Repeat": repeat_no,
                "Outer_Fold": fold_no,
                "features": ";".join(best["features"]),
                "Composite_score": best["metrics"]["Composite"],
                "Calibration_slope": calibrator.get("slope", 1.0),
                **{f"Param_{k}": v for k, v in best.items() if k not in {"subset", "metrics", "calibrator", "features", "trial_no"}},
            })
            for pos, idx in enumerate(te):
                record = calibration[int(idx)]
                actual = float(record["Actual_concentration"])
                ratio = float(record["Measured_ratio"])
                pred = float(10.0 ** pred_log[pos])
                raw_rows.append({
                    "Model": model_name,
                    "Calibration_Index": int(idx),
                    "Repeat": repeat_no,
                    "Fold": fold_no,
                    "ABC_Formula_Key": _key(record),
                    "Combo": record.get("Combo", ""),
                    "Source_file": record.get("Source_file", ""),
                    "Source_sheet": record.get("Source_sheet", ""),
                    "Source_row": record.get("Source_row", ""),
                    "Measured_ratio": ratio,
                    "Actual_concentration": actual,
                    "Base_log10_prediction": float(base[pos]),
                    "Predicted_log10_concentration": float(pred_log[pos]),
                    "Predicted_concentration": pred,
                    "Predicted_log10_RRF": math.log10(ratio) - float(pred_log[pos]),
                    "Fold_error": max(pred / actual, actual / pred),
                    "Feature_count": len(best["features"]),
                    "Features": ";".join(best["features"]),
                    "Calibration_slope": calibrator.get("slope", 1.0),
                    "Calibration_raw_slope": calibrator.get("raw_slope", 1.0),
                    "Pair_count": pair_count,
                })

    aggregated = _aggregate_cv(raw_rows)
    # Put response-corrected values back onto an area-ratio-like scale.  The
    # median predicted RRF is used as a reference so raw and corrected ratios
    # can be inspected on the same order of magnitude.
    for _model_name in sorted({str(r.get("Model")) for r in aggregated}):
        _rows = [r for r in aggregated if str(r.get("Model")) == _model_name]
        _rrfs = [10.0 ** float(r["Predicted_log10_RRF"]) for r in _rows if _num(r.get("Predicted_log10_RRF")) is not None]
        _median_rrf = float(statistics.median(_rrfs)) if _rrfs else 1.0
        for _row in _rows:
            _pred_rrf = 10.0 ** float(_row["Predicted_log10_RRF"])
            _raw_ratio = float(_row["Measured_ratio"])
            _row["Predicted_RRF"] = _pred_rrf
            _row["Reference_median_RRF"] = _median_rrf
            _row["Raw_area_ratio"] = _raw_ratio
            _row["Response_normalized_area_ratio"] = _raw_ratio / max(_pred_rrf, 1e-300) * _median_rrf
            _row["Concentration_unit"] = str(concentration_unit or "")
    comparison = _comparison(aggregated, calibration)
    best_model = str(comparison[0]["Model"])
    best_rows = [r for r in aggregated if str(r.get("Model")) == best_model]
    actual_log = np.asarray([math.log10(float(r["Actual_concentration"])) for r in best_rows], dtype=float)
    pred_log = np.asarray([float(r["Predicted_log10_concentration"]) for r in best_rows], dtype=float)
    ci_low, ci_high = _bootstrap_spearman(actual_log, pred_log, random_state + 919)
    p_value, null = _permutation_test(actual_log, pred_log, random_state + 1237, permutations)
    best_row = comparison[0]
    best_row["Spearman_bootstrap_CI_low"] = ci_low
    best_row["Spearman_bootstrap_CI_high"] = ci_high
    best_row["Spearman_permutation_p"] = p_value
    compression_slope = float(np.polyfit(pred_log, actual_log, 1)[0]) if len(pred_log) >= 2 and float(np.std(pred_log)) > 0 else float("nan")
    best_row["Actual_on_predicted_calibration_slope"] = compression_slope

    rho = float(_num(best_row.get("Trend_Spearman_r")) or -1.0)
    delta = float(_num(best_row.get("Delta_Spearman_vs_raw")) or 0.0)
    pairwise = float(_num(best_row.get("Trend_Pairwise_Concordance_pct")) or 0.0)
    if rho >= 0.40 and delta >= 0.15 and pairwise >= 60.0 and p_value < 0.05 and ci_low > 0:
        decision = "TREND-GO"
    elif rho >= 0.20 and delta >= 0.10 and pairwise >= 55.0 and p_value < 0.10:
        decision = "TREND-CAUTION"
    else:
        decision = "TREND-NO-GO"
    decision_summary = (
        f"{best_model}: OOF Spearman={rho:.3f} (95% bootstrap CI {ci_low:.3f} to {ci_high:.3f}), "
        f"permutation p={p_value:.4g}, delta vs raw={delta:+.3f}, pairwise={pairwise:.1f}%, "
        f"actual-on-predicted calibration slope={compression_slope:.2f}."
    )

    # Full-data search and prediction for both candidate models.
    X_target = _matrix(targets, feature_names)
    log_ratio_target = np.asarray([math.log10(float(r["Measured_ratio"])) for r in targets], dtype=float)
    final_config_rows: List[Dict[str, object]] = []
    per_model_target: Dict[str, np.ndarray] = {}
    per_model_calibrator: Dict[str, Dict[str, object]] = {}
    for model_no, model_name in enumerate(("PhysicalPairwiseRanker", "ResponseFactorExtraTrees"), start=1):
        seed = int(random_state + 700000 + model_no * 10000)
        best, final_trials = _search_model(
            model_name, X, y_log, log_ratio, feature_names, trials=max(5, int(trials) * 2),
            min_features=min_features, max_features=max_features, pair_min_fold=float(pair_min_fold),
            random_state=seed, stage="full_calibration_search",
            groups=(group_values if use_group_cv else None),
        )
        trial_rows.extend(final_trials)
        base, pred, calibrator, pair_count = _fit_predict_config(
            model_name, X, y_log, log_ratio, X_target, log_ratio_target, best,
            random_state=seed + 33,
            groups=(group_values if use_group_cv else None),
        )
        per_model_target[model_name] = pred
        per_model_calibrator[model_name] = dict(calibrator)
        finite_pred = np.asarray(pred, dtype=float)
        finite_pred = finite_pred[np.isfinite(finite_pred)]
        unique_log_count = int(len(np.unique(np.round(finite_pred, 12)))) if len(finite_pred) else 0
        final_config_rows.append({
            "Model": model_name,
            "Features": ";".join(best["features"]),
            "Feature_count": len(best["features"]),
            "Inner_Spearman": best["metrics"]["Spearman"],
            "Inner_Composite_score": best["metrics"]["Composite"],
            "Calibration_method": calibrator.get("kind", ""),
            "Calibration_slope": calibrator.get("slope", 1.0),
            "Calibration_raw_slope": calibrator.get("raw_slope", 1.0),
            "Isotonic_diagnostic_level_count": calibrator.get("isotonic_level_count", ""),
            "Isotonic_plateau_detected": calibrator.get("isotonic_plateau_detected", ""),
            "Target_unique_log_concentration_values": unique_log_count,
            "Target_unique_value_fraction_pct": float(unique_log_count / max(1, len(finite_pred)) * 100.0),
            "Pair_count": pair_count,
            **{f"Param_{k}": v for k, v in best.items() if k not in {"subset", "metrics", "calibrator", "features", "trial_no"}},
        })

    target_rows: List[Dict[str, object]] = []
    model_a = per_model_target.get("PhysicalPairwiseRanker")
    model_b = per_model_target.get("ResponseFactorExtraTrees")
    rank_a = _average_ranks(model_a) if model_a is not None else np.full(len(targets), np.nan)
    rank_b = _average_ranks(model_b) if model_b is not None else np.full(len(targets), np.nan)
    consensus_rank = (rank_a + rank_b) / 2.0
    rank_corr = _spearman(rank_a, rank_b) if model_a is not None and model_b is not None else float("nan")
    for i, record in enumerate(targets):
        row = {
            "Source_file": record.get("Source_file", ""),
            "Source_sheet": record.get("Source_sheet", ""),
            "Source_row": record.get("Source_row", ""),
            "Combo": record.get("Combo", ""),
            "Formula": record.get("Formula", ""),
            "ABC_Formula_Key": _key(record),
            "Measured_ratio": record.get("Measured_ratio", ""),
            "PhysicalPairwiseRanker_estimated_concentration": float(10.0 ** model_a[i]) if model_a is not None else "",
            "PhysicalPairwiseRanker_rank": float(rank_a[i]) if model_a is not None else "",
            "ResponseFactorExtraTrees_estimated_concentration": float(10.0 ** model_b[i]) if model_b is not None else "",
            "ResponseFactorExtraTrees_rank": float(rank_b[i]) if model_b is not None else "",
            "Consensus_rank": float(consensus_rank[i]),
            "Consensus_percentile": float(consensus_rank[i] / max(1, len(targets)) * 100.0),
            "Model_rank_Spearman_target": rank_corr,
            "Preferred_model": best_model,
            "Preferred_estimated_concentration": float(10.0 ** per_model_target[best_model][i]),
            "Preferred_concentration_mapping": per_model_calibrator.get(best_model, {}).get("kind", ""),
            "Preferred_mapping_slope": per_model_calibrator.get(best_model, {}).get("slope", ""),
            "Preferred_isotonic_diagnostic_levels": per_model_calibrator.get(best_model, {}).get("isotonic_level_count", ""),
            "Preferred_isotonic_plateau_detected": per_model_calibrator.get(best_model, {}).get("isotonic_plateau_detected", ""),
            "Concentration_unit": str(concentration_unit or ""),
        }
        preferred_pred_log = float(per_model_target[best_model][i])
        preferred_log_rrf = math.log10(float(record["Measured_ratio"])) - preferred_pred_log
        preferred_rrf = 10.0 ** preferred_log_rrf
        train_rrfs = [
            10.0 ** float(r["Predicted_log10_RRF"])
            for r in aggregated if str(r.get("Model")) == best_model and _num(r.get("Predicted_log10_RRF")) is not None
        ]
        reference_rrf = float(statistics.median(train_rrfs)) if train_rrfs else 1.0
        row["Preferred_predicted_log10_RRF"] = preferred_log_rrf
        row["Preferred_predicted_RRF"] = preferred_rrf
        row["Raw_area_ratio"] = float(record["Measured_ratio"])
        row["Reference_median_RRF"] = reference_rrf
        row["Response_normalized_area_ratio"] = float(record["Measured_ratio"]) / max(preferred_rrf, 1e-300) * reference_rrf
        target_rows.append(row)
    add_decile_columns(target_rows, "Preferred_estimated_concentration", prefix="Preferred_concentration")
    if target_rows:
        n_target_levels = max(1, len(target_rows))
        for row in target_rows:
            try:
                rank_value = float(row.get("Consensus_rank", ""))
                level = min(10, max(1, int(math.ceil(rank_value * 10.0 / n_target_levels))))
                row["Consensus_decile_1_to_10"] = level
                row["Consensus_level"] = f"L{level:02d}"
            except Exception:
                row["Consensus_decile_1_to_10"] = ""
                row["Consensus_level"] = ""
    target_rows.sort(key=lambda r: float(r["Consensus_rank"]), reverse=True)

    preferred_values = [
        float(r["Preferred_estimated_concentration"])
        for r in target_rows
        if _num(r.get("Preferred_estimated_concentration")) is not None
        and float(r["Preferred_estimated_concentration"]) > 0
    ]
    preferred_unique_count = int(len(np.unique(np.round(np.asarray(preferred_values, dtype=float), 12)))) if preferred_values else 0
    preferred_unique_fraction = float(preferred_unique_count / max(1, len(preferred_values)) * 100.0)

    feature_rows, group_rows = _feature_stability(trial_rows, selected_configs, feature_names, output_language)
    permutation_rows = [{"Item": "Observed_Spearman", "Value": rho}, {"Item": "Permutation_p", "Value": p_value}, {"Item": "Bootstrap_CI_low", "Value": ci_low}, {"Item": "Bootstrap_CI_high", "Value": ci_high}]
    for i, value in enumerate(null, start=1):
        permutation_rows.append({"Item": f"Permutation_{i}", "Value": value})
    settings_rows = [
        {"Setting": "Models", "Value": "PhysicalPairwiseRanker; ResponseFactorExtraTrees"},
        {"Setting": "Monte_Carlo_trials_per_outer_fold", "Value": int(trials)},
        {"Setting": "Pair_minimum_concentration_fold", "Value": float(pair_min_fold)},
        {"Setting": "Outer_CV", "Value": outer_cv_label},
        {"Setting": "Group_aware_CV", "Value": bool(use_group_cv)},
        {"Setting": "Feature_count_range", "Value": f"{min_features}-{max_features}"},
        {"Setting": "Permutation_count", "Value": int(permutations)},
        {"Setting": "Unknown_target_used_for_search", "Value": False},
        {"Setting": "Selection_guardrail", "Value": "All feature and hyperparameter search is performed inside each outer training fold."},
        {"Setting": "Pairwise_guardrail", "Value": "Concentration pairs closer than the selected fold threshold are ignored as potentially meaningless orders."},
        {"Setting": "Physical_constraint", "Value": "log concentration = log(area/IS) - predicted log RRF; area-ratio coefficient fixed at +1."},
        {"Setting": "Score_mapping", "Value": "Out-of-fold robust positive-slope continuous Huber mapping; isotonic regression is diagnostic only to detect plateau risk."},
        {"Setting": "Preferred_target_unique_concentration_values", "Value": preferred_unique_count},
        {"Setting": "Preferred_target_unique_value_fraction_pct", "Value": preferred_unique_fraction},
        {"Setting": "Concentration_unit", "Value": str(concentration_unit or "")},
    ]
    plot_paths, plot_warnings = _make_plots(Path(output_xlsx), comparison, aggregated, best_model, feature_rows, trial_rows, null, actual_log, pred_log, output_language)
    warnings = list(plot_warnings)
    minimum_reasonable_unique = max(20, int(math.ceil(math.sqrt(max(1, len(preferred_values))))))
    if preferred_values and preferred_unique_count < minimum_reasonable_unique:
        warnings.append(
            f"Target concentration diversity remains low ({preferred_unique_count} unique values for {len(preferred_values)} predictions). "
            "Review Trend_Final_Configs and use the response-factor continuous model for concentration magnitude if needed."
        )
    if decision != "TREND-GO":
        warnings.append("Trend optimizer has not yet demonstrated robust rank recovery; use target ranks as exploratory only.")
    return TrendOptimizerResult(
        enabled=True,
        best_model=best_model,
        decision=decision,
        decision_summary=decision_summary,
        comparison=comparison,
        cv_predictions=aggregated,
        target_predictions=target_rows,
        trial_rows=trial_rows,
        feature_rows=feature_rows,
        group_rows=group_rows,
        permutation_rows=permutation_rows,
        settings_rows=settings_rows,
        final_config_rows=final_config_rows,
        plot_paths=plot_paths,
        warnings=warnings,
        concentration_unit=str(concentration_unit or ""),
    )


def _append_dict_sheet(wb, title: str, rows: Sequence[Dict[str, object]]) -> None:
    ws = wb.create_sheet(title)
    headers = list(rows[0].keys()) if rows else ["Message"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def append_trend_optimizer_to_workbook(wb, result: TrendOptimizerResult) -> None:
    decision_rows = [
        {"Metric": "Decision", "Value": result.decision},
        {"Metric": "Best_model", "Value": result.best_model},
        {"Metric": "Summary", "Value": result.decision_summary},
    ] + [{"Metric": "Warning", "Value": w} for w in result.warnings]
    _append_dict_sheet(wb, "Trend_Optimizer_Decision", decision_rows)
    _append_dict_sheet(wb, "Trend_Model_Comparison", result.comparison)
    _append_dict_sheet(wb, "Trend_CV_Predictions", result.cv_predictions)
    _append_dict_sheet(wb, "Trend_Targets", result.target_predictions)
    _append_dict_sheet(wb, "Trend_Feature_Stability", result.feature_rows)
    _append_dict_sheet(wb, "Trend_Feature_Groups", result.group_rows)
    _append_dict_sheet(wb, "Trend_Random_Trials", result.trial_rows)
    _append_dict_sheet(wb, "Trend_Permutation", result.permutation_rows)
    _append_dict_sheet(wb, "Trend_Final_Configs", result.final_config_rows)
    _append_dict_sheet(wb, "Trend_Settings", result.settings_rows)
    comparison_rows = []
    for row in result.cv_predictions:
        if str(row.get("Model")) != result.best_model:
            continue
        comparison_rows.append({
            "Combo": row.get("Combo", ""),
            "Actual_concentration": row.get("Actual_concentration", ""),
            "Predicted_concentration": row.get("Predicted_concentration", ""),
            "Concentration_unit": row.get("Concentration_unit", result.concentration_unit),
            "Raw_area_ratio": row.get("Raw_area_ratio", row.get("Measured_ratio", "")),
            "Response_normalized_area_ratio": row.get("Response_normalized_area_ratio", ""),
            "Predicted_RRF": row.get("Predicted_RRF", ""),
            "Reference_median_RRF": row.get("Reference_median_RRF", ""),
            "Fold_error": row.get("Fold_error", ""),
            "Source_file": row.get("Source_file", ""),
            "Source_sheet": row.get("Source_sheet", ""),
            "Source_row": row.get("Source_row", ""),
        })
    _append_dict_sheet(wb, "Concentration_Ratio_Comparison", comparison_rows)

    if result.plot_paths:
        try:
            from openpyxl.drawing.image import Image as XLImage
            ws = wb.create_sheet("Trend_Optimizer_Plots")
            row = 1
            for path in result.plot_paths:
                if not Path(path).exists():
                    continue
                ws.cell(row=row, column=1, value=Path(path).name)
                img = XLImage(str(path))
                img.width = min(img.width, 900)
                img.height = min(img.height, 650)
                ws.add_image(img, f"A{row + 1}")
                row += 35
        except Exception:
            pass
