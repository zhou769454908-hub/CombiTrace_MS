"""Reproducible final-table example using labelled synthetic data; no RAW access."""
from __future__ import annotations
import argparse
import json
import math
import sys
from datetime import datetime
from pathlib import Path
from dataclasses import replace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT/'examples'/'demo_output')
    args = parser.parse_args()
    folder = args.output / datetime.now().strftime('synthetic_%Y%m%d_%H%M%S_%f')
    folder.mkdir(parents=True, exist_ok=False)
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from core.final_table_postprocess import autoconfig, execute, Options, theoretical_value, sha256
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt

    workbook = Workbook()
    workbook.remove(workbook.active)
    formulas = ['C2H6O', 'C3H8O', 'C4H10O', 'C5H12O']
    offsets = [1.25, -2.5, 0.0, 8.0]
    headers = ['Name', 'Formula', 'Observed_mz', 'Adduct', 'Combo', 'Sheet', 'Data_origin']
    for sheet, group, size in [('Known','SYNTHETIC_STD',3),('Unknown','SYNTHETIC_LIB',4)]:
        ws = workbook.create_sheet(sheet)
        ws.append(headers)
        for i in range(size):
            name = '%s_%04d' % (group, i+1)
            ion = float(theoretical_value(formulas[i], ion=True, adduct='[M-H]-'))
            measured = ion * (1 + offsets[i]*1e-6)
            ws.append([name,formulas[i],measured,'[M-H]-','DEMO_KEY_%d'%i,group,'SYNTHETIC_NOT_EXPERIMENTAL'])
            for rep in (1,2,3):
                out = folder/'plots'/group/('%s-%d__XIC_plots'%(group,rep))
                out.mkdir(parents=True, exist_ok=True)
                times = [k*.01 for k in range(301)]
                apex = 1.0 + i*.15 + rep*.015
                intensity = [5 + (90+5*rep)*math.exp(-.5*((t-apex)/.07)**2) for t in times]
                fig,ax=plt.subplots(figsize=(7.2,2.3))
                ax.plot(times,intensity,linewidth=1.5)
                ax.set(xlabel='Retention time (min)',ylabel='Synthetic intensity',
                       title='SYNTHETIC SOFTWARE EXAMPLE - NOT EXPERIMENTAL')
                fig.tight_layout()
                fig.savefig(out/('%03d__%s__mz%.5f.png'%(i+1,name,ion)),dpi=130)
                plt.close(fig)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for col in 'ABCDEFG': ws.column_dimensions[col].width=25
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24465C')
            cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=32
        for row in ws.iter_rows(min_row=2): row[2].number_format='0.000000'
    input_path=folder/'synthetic_input.xlsx';workbook.save(input_path)
    before=sha256(input_path)
    known=replace(autoconfig(input_path,'Known'),mass_mode='observed_mz',image_dir=str(folder/'plots'/'SYNTHETIC_STD'))
    target=replace(autoconfig(input_path,'Unknown',label='1000'),mass_mode='observed_mz',image_dir=str(folder/'plots'/'SYNTHETIC_LIB'))
    out,summary=execute(known,target,folder/'results',Options(image_width=520,image_height=480,
        xic_display_mode='all',xic_caption_replicate=True))
    assert before==sha256(input_path),'Input workbook changed'
    assert summary['tables']['100']['images_embedded']==3
    assert summary['tables']['1000']['images_embedded']==4
    check=load_workbook(summary['outputs'][0],data_only=True)
    for sheet in ('Known','Unknown'):
        ws=check[sheet];columns={str(c.value):c.column for c in ws[1]}
        for i in range(ws.max_row-1):
            value=float(ws.cell(i+2,columns['Mass_difference_ppm']).value)
            assert abs(value-offsets[i])<1e-5,(sheet,i,value)
    check.close()
    (folder/'DEMO_CHECK.json').write_text(json.dumps({'status':'passed','data':'synthetic',
        'known_rows':3,'target_rows':4,'source_images':21,'input_unchanged':True,
        'ppm_offsets_verified':offsets},indent=2),encoding='utf-8')
    print('Synthetic demonstration passed. No experimental performance was evaluated.')
    print(out)

if __name__=='__main__':
    main()
