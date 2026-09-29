"""Synthetic regression tests. These are NOT performance claims on laboratory data."""
from __future__ import annotations
import copy
import csv
import math
import tempfile
import unittest
from pathlib import Path

import numpy as np

from core.continuous_rrf import (FoldModel, ModelSpec, default_specs, feature_names,
                                 independent_groups, matrix, metrics, run_audit)
from core.continuous_io import refresh_records, execute, sha256
from core.training_limits import validate_training_count
from core.response_predictor import read_table, guess_columns, prepare_rows

FAST = [ModelSpec('RRF_median', 'median', features=0),
        ModelSpec('Ridge_standard', 'ridge', alpha=.1, features=4),
        ModelSpec('Ridge_robust', 'ridge', scaler='robust', alpha=1, features=4)]


def fixture(n=98, seed=1707, target_n=12):
    rng = np.random.RandomState(seed)
    x = rng.normal(size=(n + target_n, 6))
    logc = rng.uniform(-2, 1, size=len(x))
    logrf = .9 * x[:, 0] - .6 * x[:, 1] + rng.normal(0, .015, size=len(x))
    rows = []
    for i in range(len(x)):
        r = {'Dataset': 'Training' if i < n else 'Target',
             'Source_file': 'synthetic_only.xlsx', 'Source_sheet': 'Data', 'Source_row': i + 2,
             'Name': 'synthetic_%03d' % i, 'Formula': 'C10H20O2',
             'Combo': 'A#%d:a:C2H4 | B#1:b:C3H6 | C#1:c:C5H10O2' % (i + 1),
             'Structure_hash': 'synthetic_structure_%03d' % i,
             'Structure_status': 'synthetic_only', 'Formula_match': 'True',
             'Product_Master_Formula_Key': 'synthetic_key_%03d' % i,
             'Injection_Group': 'Injection_%02d' % (i // 10),
             'Measured_ratio': float(10 ** (logc[i] + logrf[i])),
             'Actual_concentration': float(10 ** logc[i]) if i < n else None,
             'Apex_RT_min': 4.0, 'Warnings': ''}
        r.update({'d%d' % j: float(v) for j, v in enumerate(x[i])})
        rows.append(r)
    return rows[:n], rows[n:], ['d%d' % j for j in range(6)]


def write_cache(path, training, targets, names):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = 'Summary'
    wb.active.append(['Fixture', 'Synthetic only - not laboratory data'])
    ws = wb.create_sheet('ESI_Features')
    before = ['Dataset', 'Source_file', 'Source_sheet', 'Source_row', 'Name', 'Formula',
              'Combo', 'Structure_hash', 'Structure_status', 'Formula_match', 'Injection_Group',
              'Measured_ratio', 'Actual_concentration', 'Apex_RT_min', 'Product_Master_Formula_Key']
    headers = before + names + ['Warnings']
    ws.append(headers)
    for r in training + targets:
        ws.append([r.get(h) for h in headers])
    ws.sheet_state = 'hidden'
    wb.save(path)


def write_current(path, rows, blank=False):
    fields = ['Combo', 'Formula', 'Name', 'Measured_ratio', 'Actual_concentration', 'Apex_RT_min', 'Injection_Group']
    with Path(path).open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for r in rows:
            writer.writerow({key: r.get(key) for key in fields})
        if blank:
            writer.writerow({})


class LimitsAndMatching(unittest.TestCase):
    def test_98_99_100_and_101(self):
        for n in (98, 99, 100):
            self.assertEqual(validate_training_count([{'Name': str(i)} for i in range(n)]), n)
        with self.assertRaisesRegex(ValueError, 'maximum of 100'):
            validate_training_count([{'Name': str(i)} for i in range(101)])
        self.assertEqual(validate_training_count([{'Name': 'x'}, {'Name': '', '__source_row__': 12}]), 1)

    def test_current_table_authoritative_and_reordered(self):
        train, _, _ = fixture(100)
        subset = list(reversed(train[1:99]))
        subset[0] = dict(subset[0], Actual_concentration=.12345)
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / 'training.csv'
            write_current(p, subset, blank=True)
            rows, audit, mapping = refresh_records(train, p)
            self.assertEqual(len(rows), 98)
            self.assertEqual(rows[0]['Combo'], subset[0]['Combo'])
            self.assertEqual(rows[0]['Actual_concentration'], .12345)
            self.assertEqual(audit[0]['Cached_source_row'], subset[0]['Source_row'])
            table = read_table(p)
            prepared, _ = prepare_rows(table, guess_columns(table, need_concentration=True), is_training=True)
            self.assertEqual(sum(r.ratio is not None for r in prepared), 98)

    def test_ambiguous_rt_formula_and_unmatched_fail(self):
        train, _, _ = fixture(98)
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / 'training.csv'
            write_current(p, train[:1])
            with self.assertRaises(ValueError):
                refresh_records(train + train[:1], p)
            write_current(p, [dict(train[0], Apex_RT_min=8)])
            with self.assertRaisesRegex(ValueError, 'RT'):
                refresh_records(train, p)
            write_current(p, [dict(train[0], Formula='C100H100')])
            with self.assertRaisesRegex(ValueError, 'Formula mismatch'):
                refresh_records(train, p)
            write_current(p, [dict(train[0], Combo='No such identity')])
            with self.assertRaises(ValueError):
                refresh_records(train, p)

    def test_predicted_label_column_rejected(self):
        train, _, _ = fixture(98)
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / 'predicted.csv'
            with p.open('w', encoding='utf-8-sig', newline='') as f:
                w = csv.writer(f)
                w.writerow(['Combo', 'Measured_ratio', 'Predicted concentration'])
                w.writerow([train[0]['Combo'], 1, .1])
            with self.assertRaisesRegex(ValueError, 'training labels'):
                refresh_records(train, p, column_overrides={'concentration': 'Predicted concentration'})

    def test_101_current_and_legacy_prepare_rejected(self):
        train, _, _ = fixture(101)
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / 'training.csv'
            write_current(p, train)
            with self.assertRaises(ValueError):
                refresh_records(train, p)
            table = read_table(p)
            with self.assertRaises(ValueError):
                prepare_rows(table, guess_columns(table, need_concentration=True), is_training=True)


class NumericalTests(unittest.TestCase):
    def test_standardization_invariant_to_feature_unit(self):
        rng = np.random.RandomState(12)
        X = rng.normal(size=(60, 4))
        y = X[:, 0] - X[:, 1] * .3
        spec = ModelSpec('ridge', 'ridge', features=4)
        a = FoldModel(spec).fit(X, y)
        scaled = X.copy()
        scaled[:, 0] *= 1e6
        b = FoldModel(spec).fit(scaled, y)
        np.testing.assert_allclose(a.predict_rrf(X), b.predict_rrf(scaled), rtol=1e-8, atol=1e-8)

    def test_preprocessor_does_not_learn_validation_statistics(self):
        X = np.arange(40, dtype=float).reshape(20, 2)
        y = np.arange(20, dtype=float)
        model = FoldModel(ModelSpec('r', 'ridge')).fit(X, y)
        medians, means = model.medians.copy(), model.scaler.mean_.copy()
        model.predict_rrf(np.asarray([[1e12, -1e12]]))
        np.testing.assert_array_equal(model.medians, medians)
        np.testing.assert_array_equal(model.scaler.mean_, means)

    def test_no_label_or_response_in_feature_matrix(self):
        rows = [{'d0': 1, 'Actual_concentration': 50, 'Measured_ratio': .2,
                 'Predicted_concentration': 9, 'Source_row': 12, 'log10_RRF': 4, 'Additional_Summed_Area': 1e9}]
        self.assertEqual(feature_names(rows, list(rows[0])), ['d0'])

    def test_fixed_response_coefficient_and_no_clipping(self):
        train, targets, names = fixture()
        X = matrix(train, names)
        logr = np.log10([r['Measured_ratio'] for r in train])
        y = np.log10([r['Actual_concentration'] for r in train])
        model = FoldModel(ModelSpec('r', 'ridge', alpha=.1)).fit(X, logr - y)
        testx = matrix(targets[:1], names)
        pred = model.predict_logc(testx, np.array([0.0]))
        shifted = model.predict_logc(testx, np.array([4.0]))
        np.testing.assert_allclose(shifted - pred, 4.0, atol=1e-12)

    def test_constant_midpoint_is_flagged_by_diagnostics(self):
        y = np.linspace(-3, 2, 98)
        report = metrics(y, np.full(98, np.mean(y)))
        self.assertEqual(report['Spearman'], 0)
        self.assertAlmostEqual(report['Pred_on_actual_slope'], 0)
        self.assertAlmostEqual(report['Log_spread_ratio'], 0)
        self.assertGreater(report['Low20_log_bias'], 0)
        self.assertLess(report['High20_log_bias'], 0)

    def test_duplicate_compounds_connect_injection_groups(self):
        rows = [{'Combo': 'A', 'Injection_Group': 'g1'},
                {'Combo': 'A', 'Injection_Group': 'g2'},
                {'Combo': 'B', 'Injection_Group': 'g2'},
                {'Combo': 'C', 'Injection_Group': 'g3'}]
        groups, _ = independent_groups(rows)
        self.assertEqual(groups[0], groups[2])
        self.assertNotEqual(groups[0], groups[3])

    def test_nested_recovery_target_independence_and_row_audit(self):
        train, target, names = fixture(100)
        train[3]['Measured_ratio'] = 0
        train[50]['Actual_concentration'] = None
        before = copy.deepcopy(train)
        a = run_audit(train, target, names, specs=FAST, folds=3)
        self.assertEqual(a['summary']['Training_used'], 98)
        self.assertEqual(len(a['known']), 100)
        self.assertGreater(a['summary']['Nested_OOF_metrics']['Log_R2'], .98)
        self.assertGreater(a['summary']['Nested_OOF_metrics']['Pred_on_actual_slope'], .95)
        self.assertEqual(train, before)
        poisoned = [dict(r, Actual_concentration=1e90, d0=1e5) for r in target]
        poisoned[0]['Measured_ratio'] = 0
        poisoned[1].update({name: None for name in names})
        b = run_audit(train, poisoned, names, specs=FAST, folds=3)
        np.testing.assert_array_equal(a['predicted_log'], b['predicted_log'])
        self.assertEqual(a['summary']['Final_model'], b['summary']['Final_model'])
        self.assertEqual(len(b['unknown']), len(target))
        self.assertIn('Skipped', b['unknown'][0]['Prediction_status'])
        self.assertIn('Skipped', b['unknown'][1]['Prediction_status'])
        for f in range(1, 4):
            tr = {r['Independent_group'] for r in b['cv_folds'] if r['Fold'] == f and r['Partition'] == 'Train'}
            va = {r['Independent_group'] for r in b['cv_folds'] if r['Fold'] == f and r['Partition'] == 'Validation'}
            self.assertFalse(tr & va)

    def test_no_signal_not_declared_valid(self):
        train, target, names = fixture(98)
        rng = np.random.RandomState(53)
        for r in train:
            r['Measured_ratio'] = float(10 ** rng.normal(0, 4))
        result = run_audit(train, target, names, specs=FAST, folds=3)
        self.assertEqual(result['summary']['Validation_status'], 'NOT_VALIDATED')


class ExportTests(unittest.TestCase):
    def test_end_to_end_source_immutable_distinct_outputs(self):
        train, targets, names = fixture(98)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            cache = root / 'cache.xlsx'
            write_cache(cache, train, targets, names)
            original = sha256(cache)
            folder, result = execute(cache, root / 'runs', specs=FAST, folds=3, concentration_unit='synthetic unit')
            self.assertEqual(sha256(cache), original)
            self.assertNotEqual(folder / 'continuous_concentration_audit.xlsx', cache)
            self.assertTrue((folder / 'completion.json').is_file())
            self.assertTrue((folder / '02_known_sorted.png').is_file())
            self.assertTrue((folder / 'fitted_model.joblib').is_file())
            import joblib
            saved = joblib.load(folder / 'fitted_model.joblib')
            X = matrix(targets, saved['feature_names'])
            p = saved['model'].predict_logc(X, np.log10([r['Measured_ratio'] for r in targets]))
            np.testing.assert_allclose(p, [r['Predicted_log10_concentration'] for r in result['unknown']])
            from openpyxl import load_workbook
            wb = load_workbook(folder / 'continuous_concentration_audit.xlsx', read_only=True, data_only=True)
            self.assertEqual(wb['Known_Standards_Predictions'].max_row, 99)
            self.assertEqual(wb['Unknown_Sample_Predictions'].max_row, len(targets) + 1)
            wb.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
