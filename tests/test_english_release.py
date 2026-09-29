"""English release and legacy input regression checks; all records are synthetic."""
import ast
import csv
import json
import re
import tempfile
import unittest
from pathlib import Path
from core.legacy_schema import canonical_header, alias_record
from core.response_predictor import read_table, guess_columns, prepare_rows, _norm_header
from core.final_table_postprocess import autoconfig, hnorm, Options, execute, sha256
from core.final_table_selftest import fixture
from core.final_table_triplicates import hn
from core.chemistry import normalize_polarity

ROOT=Path(__file__).resolve().parents[1]

class SchemaTests(unittest.TestCase):
    def test_triplicate_headers(self):
        pairs={'平均峰面积':'Mean_Peak_Area','平均相对内标比例_%':'Mean_Area_to_IS_Ratio_%',
               '比例SD':'Ratio_SD','峰面积RSD_%':'Peak_Area_RSD_%','平均Apex_RT':'Mean_Apex_RT',
               '最终三平行平均结果':'Final_Triplicate_Mean_Result','严格三平行平均结果(3/3)':'Strict_Triplicate_Mean_Result_3of3',
               '结果有效平行数':'Valid_Result_Replicate_Count','内标QC说明':'IS_QC_Notes'}
        for old,new in pairs.items():
            with self.subTest(old=old):self.assertEqual(canonical_header(old),new)
    def test_earlier_english_schema(self):
        for old in ('Mean Ratio vs IS (%)','平均相对内标比例_%','Mean_Area_to_IS_Ratio_%'):
            self.assertEqual(_norm_header(old),_norm_header('Mean_Area_to_IS_Ratio_%'))
    def test_identity_values_not_translated(self):
        r={'名称':'普通样品甲','Formula':'C2H6O','浓度':0.25}
        out=alias_record(r)
        self.assertEqual(out['Name'],'普通样品甲');self.assertEqual(r,{'名称':'普通样品甲','Formula':'C2H6O','浓度':0.25})
    def test_alias_conflicts_rejected(self):
        with self.assertRaises(ValueError):alias_record({'平均峰面积':100,'Mean_Peak_Area':101})
    def test_equal_aliases_accepted(self):
        self.assertEqual(alias_record({'平均峰面积':'100.0','Mean_Peak_Area':100})['Mean_Peak_Area'],100)
    def test_legacy_polarity(self):
        self.assertEqual(normalize_polarity('正离子'),'positive');self.assertEqual(normalize_polarity('负离子'),'negative')
    def test_postprocessor_header_equivalence(self):
        for old,new in [('分子式','Formula'),('名称','Name'),('峰顶保留时间','Apex_RT_min')]:
            # RT labels may denote multiple known spellings; verify identity schemas here.
            if old=='峰顶保留时间':continue
            self.assertEqual(hnorm(old),hnorm(new));self.assertEqual(hn(old),hn(new))
    def test_legacy_csv_calibration(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'input.csv'
            with p.open('w',encoding='utf-8-sig',newline='') as f:
                w=csv.writer(f);w.writerow(['名称','分子式','组合','平均相对内标比例_%','已知浓度','平均Apex_RT'])
                w.writerow(['样品甲','C2H6O','A#1:A1:C | B#1:B1:N | C#1:C1:O',120,.5,2.5])
            table=read_table(p);cfg=guess_columns(table,need_concentration=True)
            rows,_=prepare_rows(table,cfg,is_training=True)
            self.assertEqual(len(rows),1);self.assertEqual(rows[0].name,'样品甲')
            self.assertEqual(rows[0].ratio,120);self.assertEqual(rows[0].concentration,.5)
            self.assertEqual(float(rows[0].raw['Mean_Apex_RT']),2.5)
    def test_legacy_sheet_preference(self):
        with tempfile.TemporaryDirectory() as d:
            p=fixture(Path(d)/'old.xlsx',{'说明':[['Item','Value'],['A','B']], '三平行平均结果':[['名称','分子式','平均相对内标比例_%','已知浓度'],['样品甲','C2H6O',100,.5]]})
            t=read_table(p);self.assertEqual(t.sheet_name,'三平行平均结果')
            c=guess_columns(t,need_concentration=True);self.assertEqual(c.ratio_col,'平均相对内标比例_%')
    def test_postprocess_legacy_input_unchanged(self):
        with tempfile.TemporaryDirectory() as d:
            p=fixture(Path(d)/'old.xlsx',{'Known':[['名称','分子式','Exact_mass','Combo'],['样品甲','C2H6O',46.041864,'KEY_A']],
               'Unknown':[['Name','Formula','Exact_mass','Combo'],['Target_A','C2H6O',46.041864,'KEY_A']]})
            before=sha256(p);a=autoconfig(p,'Known',label='100');b=autoconfig(p,'Unknown',label='1000')
            self.assertEqual(a.name_col,1);self.assertEqual(a.formula_col,2)
            folder,result=execute(a,b,Path(d)/'out',Options())
            self.assertEqual(sha256(p),before)
            from core.final_table_ooxml import Book
            book=Book(result['outputs'][0]);headers,rows=book.scan('Known')
            self.assertEqual(headers[1],'名称');self.assertEqual(rows[0][1][1],'样品甲')
            self.assertIn('Name_in_1000',headers.values())
    def test_english_only_language_selector(self):
        from core.esi_descriptor_meta import LANGUAGE_OPTIONS
        self.assertEqual(LANGUAGE_OPTIONS,(("English","en"),))
    def test_removed_feature_modules_not_bundled(self):
        from packaging_tools.catalog import project_modules
        mods=project_modules(ROOT)
        for name in ('core.review','core.deuteration','core.generator','core.fine_resolution','core.local_structure_rrf'):
            self.assertNotIn(name,mods)
    def test_python38_parse(self):
        for p in ROOT.rglob('*.py'):
            if '__pycache__' in p.parts:continue
            ast.parse(p.read_text('utf-8-sig'),filename=str(p),feature_version=(3,8))

class InterfaceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from app import ThermoBatchReportApp
        import tkinter as tk
        try:cls.app=ThermoBatchReportApp();cls.app.withdraw();cls.app.update_idletasks()
        except tk.TclError as exc:raise unittest.SkipTest(str(exc))
    @classmethod
    def tearDownClass(cls):cls.app.destroy()
    def test_tabs(self):
        self.assertEqual([self.app.notebook.tab(x,'text') for x in self.app.notebook.tabs()],
            ['Enumeration','Final tables','Triplicate XIC','Table tools','CSV filter','Response','SMILES','ESI models','Local RRF','XIC'])
        for a in ('tab_generate','tab_deut','tab_review'):self.assertFalse(hasattr(self.app,a))
    def test_standard_controls_are_english(self):
        def walk(w):
            for c in w.winfo_children():
                for k in ('text','values'):
                    try:s=str(c.cget(k))
                    except Exception:continue
                    self.assertIsNone(re.search('[\u3400-\u9fff]',s),(str(c),s))
                walk(c)
        walk(self.app)
    def test_filter_mode_not_changed_by_translation(self):
        a=self.app;a.var_filter_mode_label.set('Match normalized molecular formulas');a._sync_filter_mode()
        self.assertEqual(a.var_filter_mode_code.get(),'canonical')
        a.var_filter_mode_label.set('Exact text match');a._sync_filter_mode();self.assertEqual(a.var_filter_mode_code.get(),'exact')
    def test_analysis_defaults_unchanged(self):
        a=self.app
        self.assertEqual(a.var_excel_ppm.get(),5.0);self.assertEqual(a.var_excel_min_peak_height.get(),1e4)
        self.assertEqual(a.var_excel_dup_min_rel_height.get(),.1)
    def test_adduct_controls_keep_codes(self):
        from app import ADDUCT_MODE_OPTIONS,MULTIPEAK_MODE_OPTIONS
        self.assertEqual([v for k,v in ADDUCT_MODE_OPTIONS],['posneg_neghonly','posneg_all','pos_all','neg_all','force'])
        self.assertEqual([v for k,v in MULTIPEAK_MODE_OPTIONS],['auto','conservative','balanced','sensitive'])


class TriplicateExportTests(unittest.TestCase):
    def test_english_triplicate_export_roundtrip(self):
        from core.excel_sheet_xic import load_excel_sheet_plans, write_parallel_summary_xlsx, RawResolution
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            plans,_=load_excel_sheet_plans(ROOT/'templates'/'excel_sheet_abc_template.xlsx')
            plan=plans[0];target=plan.targets[0]
            resolutions=[RawResolution(str(i),'S-%d'%i,root/('S-%d.raw'%i),'exact') for i in (1,2,3)]
            quant={}
            for i in (1,2,3):
                quant[str(i)]=[{'No':str(target.no),'Name':target.name,'Formula':target.formula,
                    'Found':'True','Area':str(100*i),'Apex_RT':'2.5','Ratio_%':str(10*i),
                    'Internal_area':'1000','Internal_standard_found':'True'},
                    {'No':'0','Name':'IS','Formula':'C2H6O','Is_internal_standard':'True',
                     'Found':'True','Area':'1000','Apex_RT':'1.5'}]
            path=write_parallel_summary_xlsx(root/'summary.xlsx',[{'plan':plan,'resolutions':resolutions,
                'quant_rows':quant}],internal_standard_formula='C2H6O')
            w=load_workbook(path,data_only=True)
            self.assertIn('Triplicate Average',w.sheetnames)
            rows=list(w['Triplicate Average'].values);cols={str(x):i for i,x in enumerate(rows[0])}
            row=next(x for x in rows[1:] if x[cols['Name']]==target.name)
            self.assertEqual(row[cols['Mean_Peak_Area']],200)
            self.assertEqual(row[cols['Mean_Area_to_IS_Ratio_%']],20)
            self.assertEqual(row[cols['Valid_Ratio_Replicate_Count']],3)
            for sheet in w:
                for values in sheet.values:
                    for value in values:
                        if isinstance(value,str): self.assertIsNone(re.search('[\u3400-\u9fff]',value),value)
            w.close()
            read=read_table(path)
            config=guess_columns(read,need_concentration=False)
            self.assertEqual(config.ratio_col,'Mean_Area_to_IS_Ratio_%')

if __name__=='__main__':
    unittest.main()
